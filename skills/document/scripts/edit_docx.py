#!/usr/bin/env python3
import argparse
import glob
import json
import os

from docx import Document

from create_docx import add_block


def parse_arguments():
    parser = argparse.ArgumentParser(description="Append content to an existing DOCX in place.")
    parser.add_argument("document_path", nargs="?", help="Path to the .docx; defaults to the newest .docx in ~/documents")
    parser.add_argument("--heading", action="append", default=[], metavar="TEXT", help="Append a level-1 heading")
    parser.add_argument("--paragraph", action="append", default=[], metavar="TEXT", help="Append a paragraph")
    parser.add_argument("--bullet", action="append", default=[], metavar="TEXT", help="Append a bullet item")
    parser.add_argument("--blocks", metavar="JSON_PATH", help="Optional JSON file with an array of blocks (same schema as create_docx)")
    return parser.parse_args()


def resolve_document_path(given_path):
    if given_path:
        return os.path.expanduser(given_path)
    documents = sorted(
        glob.glob(os.path.expanduser("~/documents/*.docx")),
        key=os.path.getmtime,
        reverse=True,
    )
    if not documents:
        raise SystemExit("no .docx found in ~/documents; pass the document path explicitly")
    return documents[0]


def load_blocks(blocks_path):
    with open(os.path.expanduser(blocks_path), "r", encoding="utf-8") as blocks_file:
        blocks = json.load(blocks_file)
    if not isinstance(blocks, list):
        raise ValueError("blocks must be an array")
    return blocks


def main():
    arguments = parse_arguments()
    document_path = resolve_document_path(arguments.document_path)
    document = Document(document_path)
    for heading_text in arguments.heading:
        document.add_heading(heading_text, level=1)
    for paragraph_text in arguments.paragraph:
        document.add_paragraph(paragraph_text)
    for bullet_text in arguments.bullet:
        document.add_paragraph(bullet_text, style="List Bullet")
    if arguments.blocks:
        for block in load_blocks(arguments.blocks):
            add_block(document, block)
    document.save(document_path)
    print(document_path)


if __name__ == "__main__":
    main()
