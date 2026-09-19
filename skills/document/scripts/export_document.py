#!/usr/bin/env python3
import argparse
import os
import re
from pathlib import Path

from skill_runtime import cache_home_path, ensure_requirements

INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def parse_arguments():
    parser = argparse.ArgumentParser(description="Render a markdown source of truth into a .docx deliverable")
    parser.add_argument("markdown_path", help="path to content.md")
    parser.add_argument("--output", help="output path; defaults next to the markdown")
    parser.add_argument("--format", default="docx", choices=["docx", "pdf"], help="deliverable format")
    parser.add_argument("--font", default="맑은 고딕", help="base font family name for docx")
    parser.add_argument("--font-size", type=float, default=10.5)
    parser.add_argument("--font-path", default="", help="Korean-capable TTF for pdf output")
    return parser.parse_args()


def add_inline_runs(paragraph, text, point_class):
    for segment in INLINE_PATTERN.split(text):
        if not segment:
            continue
        if segment.startswith("**") and segment.endswith("**") and len(segment) > 4:
            run = paragraph.add_run(segment[2:-2])
            run.bold = True
        elif segment.startswith("*") and segment.endswith("*") and len(segment) > 2:
            run = paragraph.add_run(segment[1:-1])
            run.italic = True
        elif segment.startswith("`") and segment.endswith("`") and len(segment) > 2:
            run = paragraph.add_run(segment[1:-1])
            run.font.name = "Courier New"
            run.font.size = point_class(9.5)
        else:
            paragraph.add_run(segment)


def is_table_line(line):
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def is_divider_row(line):
    return bool(re.fullmatch(r"\|?[\s:|-]+\|?", line.strip())) and "-" in line


def split_table_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def render_table(document, table_lines, point_class):
    rows = [split_table_row(line) for line in table_lines if not is_divider_row(line)]
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            cell = table.rows[row_index].cells[column_index]
            text = row[column_index] if column_index < len(row) else ""
            cell.paragraphs[0].text = ""
            add_inline_runs(cell.paragraphs[0], text, point_class)
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def set_base_font(document, font_name, font_size, point_class):
    from docx.oxml.ns import qn

    style = document.styles["Normal"]
    style.font.name = font_name
    style.font.size = point_class(font_size)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def render_markdown(document, markdown_text, point_class):
    lines = markdown_text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        heading_match = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading_match:
            document.add_heading(heading_match.group(2).strip(), level=len(heading_match.group(1)))
            index += 1
            continue
        if is_table_line(line):
            table_lines = []
            while index < len(lines) and is_table_line(lines[index]):
                table_lines.append(lines[index])
                index += 1
            render_table(document, table_lines, point_class)
            continue
        bullet_match = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet_match.group(1), point_class)
            index += 1
            continue
        numbered_match = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered_match:
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, numbered_match.group(1), point_class)
            index += 1
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = point_class(18)
            run_text = stripped.lstrip("> ").strip()
            add_inline_runs(paragraph, run_text, point_class)
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue
        paragraph_lines = []
        while index < len(lines) and lines[index].strip() and not re.match(r"^(#{1,4})\s+", lines[index].strip()) and not is_table_line(lines[index]) and not re.match(r"^\s*([-*]|\d+[.)])\s+", lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = document.add_paragraph()
        add_inline_runs(paragraph, " ".join(paragraph_lines), point_class)


PDF_FONT_CANDIDATES = [
    Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
    Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
]


def is_embeddable_font(font_path):
    try:
        from fontTools.ttLib import TTFont
    except ImportError:
        return True
    try:
        font = TTFont(str(font_path), fontNumber=0, lazy=True)
    except Exception:
        return False
    return "OS/2" in font and "cmap" in font


def cached_font_paths():
    fonts_directory = cache_home_path(os.environ) / "fonts"
    return [fonts_directory / "NanumGothic.ttf", fonts_directory / "NotoSansKR-Regular.ttf"]


def resolve_pdf_font(font_path_argument):
    if font_path_argument:
        return Path(font_path_argument)
    for candidate in cached_font_paths() + PDF_FONT_CANDIDATES:
        if candidate.exists() and is_embeddable_font(candidate):
            return candidate
    return None


def export_pdf(markdown_text, output_path, font_path_argument, font_size):
    from fpdf import FPDF

    font_path = resolve_pdf_font(font_path_argument)
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    family = "Helvetica"
    if font_path and font_path.exists():
        family = "DocumentFont"
        pdf.add_font(family, "", str(font_path))
        pdf.add_font(family, "B", str(font_path))
        pdf.add_font(family, "I", str(font_path))
    elif any(ord(character) > 0x2000 for character in markdown_text):
        raise SystemExit("non-Latin PDF text requires --font-path or an installed Korean-capable font")

    def write_line(text, size, style="", indent=0, spacing=2):
        pdf.set_font(family, style, size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(0, size * 0.55, strip_inline_markers(text))
        pdf.ln(spacing)

    lines = markdown_text.splitlines()
    index = 0
    heading_sizes = {1: 20, 2: 15, 3: 12.5, 4: 11}
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        heading_match = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading_match:
            write_line(heading_match.group(2), heading_sizes[len(heading_match.group(1))], "B", spacing=3)
            index += 1
            continue
        if is_table_line(lines[index]):
            table_lines = []
            while index < len(lines) and is_table_line(lines[index]):
                table_lines.append(lines[index])
                index += 1
            rows = [split_table_row(line) for line in table_lines if not is_divider_row(line)]
            pdf.set_font(family, "", font_size)
            with pdf.table() as table:
                for row in rows:
                    table_row = table.row()
                    for cell in row:
                        table_row.cell(strip_inline_markers(cell))
            pdf.ln(3)
            continue
        list_match = re.match(r"^\s*(?:[-*]|\d+[.)])\s+(.*)$", lines[index])
        if list_match:
            marker = "• " if re.match(r"^\s*[-*]", lines[index]) else re.match(r"^\s*(\d+[.)])", lines[index]).group(1) + " "
            write_line(marker + list_match.group(1), font_size, indent=4, spacing=1)
            index += 1
            continue
        if stripped.startswith(">"):
            write_line(stripped.lstrip("> ").strip(), font_size, "I", indent=8)
            index += 1
            continue
        paragraph_lines = []
        while index < len(lines) and lines[index].strip() and not re.match(r"^(#{1,4})\s+", lines[index].strip()) and not is_table_line(lines[index]) and not re.match(r"^\s*([-*]|\d+[.)])\s+", lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        write_line(" ".join(paragraph_lines), font_size)
    pdf.output(str(output_path))


def strip_inline_markers(text):
    return re.sub(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`", lambda match: next(group for group in match.groups() if group is not None), text)


def main():
    arguments = parse_arguments()
    ensure_requirements("document")

    markdown_path = Path(arguments.markdown_path)
    markdown_text = markdown_path.read_text(encoding="utf-8")
    default_suffix = "." + arguments.format
    output_path = Path(arguments.output) if arguments.output else markdown_path.with_suffix(default_suffix)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if arguments.format == "pdf":
        export_pdf(markdown_text, output_path, arguments.font_path, arguments.font_size)
    else:
        from docx import Document
        from docx.shared import Pt

        document = Document()
        set_base_font(document, arguments.font, arguments.font_size, Pt)
        render_markdown(document, markdown_text, Pt)
        document.save(output_path)
    print(f"exported {output_path} from {markdown_path}")


if __name__ == "__main__":
    main()
