#!/usr/bin/env python3
import argparse
import json

from skill_runtime import ensure_requirements


def summarize_workbook(workbook_path):
    if not ensure_requirements("xlsx"):
        raise RuntimeError("xlsx dependencies are unavailable after bootstrap")

    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=False)
    sheets = []
    formula_cells = []
    warnings = []
    for worksheet in workbook.worksheets:
        sheet_summary = summarize_sheet(worksheet)
        sheets.append({
            "title": worksheet.title,
            "rows": worksheet.max_row,
            "columns": worksheet.max_column,
            "freezePanes": sheet_summary["freezePanes"],
            "autoFilter": sheet_summary["autoFilter"],
            "blankHeaderCount": sheet_summary["blankHeaderCount"],
            "errorFormulaCount": sheet_summary["errorFormulaCount"],
        })
        warnings.extend(sheet_warnings(worksheet.title, sheet_summary))
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append(f"{worksheet.title}!{cell.coordinate}")
    return {
        "sheetCount": len(workbook.worksheets),
        "sheets": sheets,
        "formulaCells": formula_cells[:50],
        "formulaCellCount": len(formula_cells),
        "warnings": warnings,
        "warningCount": len(warnings),
    }


def summarize_sheet(worksheet):
    header_row = header_row_index(worksheet)
    header_values = [cell.value for cell in worksheet[header_row]] if worksheet.max_row >= header_row else []
    blank_header_count = sum(1 for value in header_values if value is None or str(value).strip() == "")
    error_formula_count = 0
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and ("#REF!" in cell.value or "#VALUE!" in cell.value or "#DIV/0!" in cell.value):
                error_formula_count += 1
    return {
        "headerRow": header_row,
        "titleRowDetected": header_row == 2,
        "freezePanes": str(worksheet.freeze_panes) if worksheet.freeze_panes else None,
        "autoFilter": bool(worksheet.auto_filter.ref),
        "blankHeaderCount": blank_header_count,
        "errorFormulaCount": error_formula_count,
    }


def header_row_index(worksheet):
    if worksheet.max_row < 2:
        return 1
    first_row_values = [cell.value for cell in worksheet[1]]
    second_row_values = [cell.value for cell in worksheet[2]]
    if non_blank_count(first_row_values) == 1 and non_blank_count(second_row_values) > 1:
        return 2
    return 1


def non_blank_count(values):
    return sum(1 for value in values if value is not None and str(value).strip())



def sheet_warnings(title, sheet_summary):
    warnings = []
    if not sheet_summary["freezePanes"]:
        warnings.append(f"{title}: header row is not frozen")
    if not sheet_summary["autoFilter"]:
        warnings.append(f"{title}: auto filter is missing")
    if sheet_summary["blankHeaderCount"] > 0:
        warnings.append(f"{title}: {sheet_summary['blankHeaderCount']} blank header cells")
    if sheet_summary["errorFormulaCount"] > 0:
        warnings.append(f"{title}: {sheet_summary['errorFormulaCount']} formulas contain spreadsheet error markers")
    return warnings


def parse_arguments():
    parser = argparse.ArgumentParser(description="Validate and summarize an XLSX workbook.")
    parser.add_argument("workbook_path")
    return parser.parse_args()


def main():
    arguments = parse_arguments()
    summary = summarize_workbook(arguments.workbook_path)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
