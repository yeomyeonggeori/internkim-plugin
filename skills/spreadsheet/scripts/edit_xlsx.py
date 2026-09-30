#!/usr/bin/env python3
import argparse
import glob
import json
import os

from openpyxl import load_workbook

from cell_values import typed_cell_value


def parse_arguments():
    parser = argparse.ArgumentParser(description="Append rows to an existing XLSX workbook in place.")
    parser.add_argument("workbook_path", nargs="?", help="Path to the .xlsx; defaults to the newest .xlsx in ~/documents")
    parser.add_argument("--sheet", default=None, metavar="NAME", help="Sheet name (default: active sheet; created if missing)")
    parser.add_argument("--row", action="append", default=[], metavar="CELLS", help="Append one row; comma-separated cell values (repeatable)")
    parser.add_argument("--rows", metavar="JSON_PATH", help="Optional JSON file with an array of row arrays")
    return parser.parse_args()


def resolve_workbook_path(given_path):
    if given_path:
        return os.path.expanduser(given_path)
    documents = sorted(
        glob.glob(os.path.expanduser("~/documents/*.xlsx")),
        key=os.path.getmtime,
        reverse=True,
    )
    if not documents:
        raise SystemExit("no .xlsx found in ~/documents; pass the workbook path explicitly")
    return documents[0]


def load_rows_from_json(rows_path):
    with open(os.path.expanduser(rows_path), "r", encoding="utf-8") as rows_file:
        rows = json.load(rows_file)
    if not isinstance(rows, list):
        raise ValueError("rows JSON must be an array")
    return rows


def parse_row(row_string):
    return [typed_cell_value(cell.strip()) for cell in row_string.split(",")]


def resolve_worksheet(workbook, sheet_name):
    if sheet_name is None:
        return workbook.active
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return workbook.create_sheet(title=sheet_name)


def main():
    arguments = parse_arguments()
    workbook_path = resolve_workbook_path(arguments.workbook_path)
    workbook = load_workbook(workbook_path)
    worksheet = resolve_worksheet(workbook, arguments.sheet)
    for row_string in arguments.row:
        worksheet.append(parse_row(row_string))
    if arguments.rows:
        for row in load_rows_from_json(arguments.rows):
            if not isinstance(row, list):
                raise ValueError("each row in JSON must be an array")
            worksheet.append(row)
    workbook.save(workbook_path)
    print(workbook_path)


if __name__ == "__main__":
    main()
