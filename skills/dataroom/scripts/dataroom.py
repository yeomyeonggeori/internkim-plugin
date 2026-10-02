#!/usr/bin/env python3
import argparse
import csv
import datetime
import hashlib
import html.parser
import json
import re
import shutil
import subprocess
import sys
import unicodedata
import zipfile
from pathlib import Path

from skill_runtime import ensure_requirements, setup_command

TEMPLATE = json.loads((Path(__file__).resolve().parent.parent / "assets/template.json").read_text())


def category_path(category, categories):
    folder = category["code"] + "-" + category["slug"]
    if category["parent"] is None:
        return folder
    parent = next(item for item in categories if item["code"] == category["parent"])
    return parent["code"] + "-" + parent["slug"] + "/" + folder


README_LINES = [
    "# Data room",
    "",
    "- Start a session from `company.json` and `INDEX.md`.",
    "- Find with `company_document_search`, or `dataroom search` in a tree,",
    "  narrowed by `path` when browsing a category. Never list the tree.",
    "- Read the sidecar before the original, and the original only when the sidecar",
    "  cannot answer.",
    "- Never edit an existing document; add one with `supersedes`.",
    "- Add only through `dataroom ingest`, so the catalog is never out of step",
    "  with a directory.",
    "- A number comes from `company_metric_list`; the data room holds its evidence.",
    "  A metric record names its document by `id`.",
    "",
]
KINDS = ("contract", "policy", "report", "deck", "dataset", "image", "video")
STATUSES = ("current", "superseded", "draft")
REQUIRED_FIELDS = ("id", "title", "kind", "categoryCode", "date", "status", "summary", "language")
SUMMARY_LIMIT = 200
DIRECTORY_LIMIT = 40
INBOX_LIMIT_DAYS = 7
OPENING_LIMIT = 3000
TEXT_SUFFIXES = (".md", ".markdown", ".txt")
IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff", ".bmp", ".heic")
VIDEO_SUFFIXES = (".mp4", ".mov", ".mkv", ".webm", ".avi")
AUDIO_SUFFIXES = (".m4a", ".mp3", ".wav", ".aac", ".flac", ".ogg")
KIND_BY_SUFFIX = {
    ".pptx": "deck",
    ".key": "deck",
    ".xlsx": "dataset",
    ".csv": "dataset",
    **{suffix: "image" for suffix in IMAGE_SUFFIXES},
    **{suffix: "video" for suffix in VIDEO_SUFFIXES + AUDIO_SUFFIXES},
}
FRONTMATTER_PATTERN = re.compile(r"\A---\n(.*?)\n---\n?(.*)\Z", re.DOTALL)
INDEX_HASH_PATTERN = re.compile(r"<!-- frontmatter-hash: ([0-9a-f]{64}) -->")
XML_TEXT_PATTERN = re.compile(r"<(?:a|w):t(?:\s[^>]*)?>([^<]*)</(?:a|w):t>")
XML_PARAGRAPH_PATTERN = re.compile(r"<(?:a|w):p(?:\s[^>]*)?>(.*?)</(?:a|w):p>", re.DOTALL)
DOCX_HEADING_PATTERN = re.compile(r'<w:pStyle w:val="Heading')
SLIDE_PATTERN = re.compile(r"ppt/slides/slide(\d+)\.xml")
NOTES_PATTERN = re.compile(r"ppt/notesSlides/notesSlide(\d+)\.xml")


class Failure(Exception):
    pass


def main():
    if not ensure_requirements("dataroom"):
        print(f"error: the data room's Python packages are not prepared; run {setup_command()} once", file=sys.stderr)
        sys.exit(1)
    parser = build_parser()
    arguments = parser.parse_args()
    try:
        exit_code = arguments.handler(arguments)
    except Failure as failure:
        print(f"error: {failure}", file=sys.stderr)
        sys.exit(2)
    sys.exit(exit_code or 0)


def build_parser():
    parser = argparse.ArgumentParser(description="Keep a company data room in the standard shape.")
    subparsers = parser.add_subparsers(required=True)
    add_init_parser(subparsers)
    add_ingest_parser(subparsers)
    add_check_parser(subparsers)
    add_index_parser(subparsers)
    add_search_parser(subparsers)
    return parser


def add_init_parser(subparsers):
    parser = subparsers.add_parser("init", help="create an empty data room")
    parser.add_argument("directory")
    parser.add_argument("--slug", required=True)
    parser.add_argument("--name", action="append", default=[], help="<language>=<name>, repeatable")
    parser.add_argument("--country", default="KR")
    parser.add_argument("--locale", default="ko")
    parser.add_argument("--timezone", default="Asia/Seoul")
    parser.add_argument("--currency", default="KRW")
    parser.add_argument("--business", default="")
    parser.set_defaults(handler=run_init)


def add_ingest_parser(subparsers):
    parser = subparsers.add_parser("ingest", help="file a document with its sidecar and derived text")
    parser.add_argument("directory")
    parser.add_argument("file")
    parser.add_argument("--category", required=True, help="exact leaf mnemonic code, or X")
    parser.add_argument("--title")
    parser.add_argument("--kind", choices=KINDS)
    parser.add_argument("--date", help="YYYY-MM-DD, the date the document speaks from")
    parser.add_argument("--summary")
    parser.add_argument("--tags", help="comma separated")
    parser.add_argument("--supersedes", help="id of the version this one replaces")
    parser.add_argument("--language")
    parser.add_argument("--id")
    parser.add_argument("--period")
    parser.add_argument("--source")
    parser.set_defaults(handler=run_ingest)


def add_check_parser(subparsers):
    parser = subparsers.add_parser("check", help="report every departure from the standard")
    parser.add_argument("directory")
    parser.set_defaults(handler=run_check)


def add_index_parser(subparsers):
    parser = subparsers.add_parser("index", help="regenerate INDEX.md and every catalog.jsonl")
    parser.add_argument("directory")
    parser.set_defaults(handler=run_index)


def add_search_parser(subparsers):
    parser = subparsers.add_parser("search", help="search the catalogs")
    parser.add_argument("directory")
    parser.add_argument("query")
    parser.add_argument("--path", help="only documents whose path starts with this prefix")
    parser.set_defaults(handler=run_search)


def run_init(arguments):
    root = Path(arguments.directory).resolve()
    if (root / "company.json").exists():
        raise Failure(f"{root} already holds a data room")
    root.mkdir(parents=True, exist_ok=True)
    for folder, _ in category_folders(root, TEMPLATE["categories"]):
        (root / folder / "archive").mkdir(parents=True, exist_ok=True)
        (root / folder / "catalog.jsonl").touch()
    (root / "README.md").write_text("\n".join(README_LINES), encoding="utf-8")
    write_json(root / "company.json", build_company(arguments))
    write_generated_files(root)
    print(root)


def build_company(arguments):
    names = dict(parse_name(entry, arguments.locale) for entry in arguments.name)
    return {
        "schemaVersion": 2,
        "slug": arguments.slug,
        "country": arguments.country,
        "locale": arguments.locale,
        "timezone": arguments.timezone,
        "currencyCode": arguments.currency,
        "business": arguments.business,
        "profile": {"name": names},
        "dataroom": TEMPLATE,
    }


def parse_name(entry, locale):
    language, separator, name = entry.partition("=")
    if separator == "":
        return locale, entry
    return language, name


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def load_company(root):
    path = root / "company.json"
    if not path.exists():
        raise Failure(f"{root} holds no company.json; run init first")
    return json.loads(path.read_text(encoding="utf-8"))


def category_folders(root, categories=None):
    if categories is None:
        company = load_company(root)
        if company.get("schemaVersion") != 2:
            raise Failure("legacy data rooms need semantic reclassification; their permissions cannot be converted by folder name")
        categories = company["dataroom"]["categories"]
    parent_codes = {category["parent"] for category in categories}
    return [(category_path(category, categories), category["code"])
            for category in categories if category["code"] not in parent_codes]


def category_code(folder):
    return folder.rsplit("/", 1)[-1].split("-", 1)[0]


def resolve_category(root, requested):
    for folder, code in category_folders(root):
        if requested == code:
            return folder
    raise Failure(f"no filing category named {requested}")


def run_ingest(arguments):
    root = Path(arguments.directory).resolve()
    source = Path(arguments.file).resolve()
    if not source.is_file():
        raise Failure(f"{source} is not a file")
    folder = resolve_category(root, arguments.category)
    document_date = arguments.date or file_date(source)
    kind = arguments.kind or KIND_BY_SUFFIX.get(source.suffix.lower(), "report")
    destination_directory = destination_directory_for(root, folder)
    destination = destination_directory / f"{document_date}-{document_slug(arguments.title, source)}{source.suffix.lower()}"
    if destination.exists():
        raise Failure(f"{destination} exists; nothing is overwritten, pass another --title or --date")
    fields = build_fields(root, arguments, source, folder, kind, document_date)
    destination_directory.mkdir(parents=True, exist_ok=True)
    notes = file_document(source, destination, fields, arguments.summary)
    if arguments.supersedes:
        mark_superseded(root, arguments.supersedes)
    write_generated_files(root)
    report_ingest(root, destination, fields, notes)


def file_date(source):
    return datetime.date.fromtimestamp(source.stat().st_mtime).isoformat()


def destination_directory_for(root, folder):
    return root / folder


def document_slug(title, source):
    return slugify(title or "") or slugify(source.stem) or "document"


def slugify(value):
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")


def build_fields(root, arguments, source, folder, kind, document_date):
    fields = {
        "id": arguments.id or unique_id(root, folder, document_date, document_slug(arguments.title, source), source),
        "title": arguments.title or source.stem,
        "kind": kind,
        "categoryCode": category_code(folder),
        "date": document_date,
    }
    if arguments.period:
        fields["period"] = arguments.period
    fields["status"] = "current"
    if arguments.supersedes:
        fields["supersedes"] = arguments.supersedes
    if arguments.source:
        fields["source"] = arguments.source
    fields["language"] = arguments.language or load_company(root).get("locale", "en")
    fields["summary"] = arguments.summary or ""
    if arguments.tags:
        fields["tags"] = [tag.strip() for tag in arguments.tags.split(",") if tag.strip()]
    return fields


def unique_id(root, folder, document_date, slug, source):
    taken = {document.fields.get("id") for document in all_documents(root) if document.fields}
    base = f"{category_code(folder)}-{document_date[:4]}-{slug}"
    for candidate in (base, f"{base}-{source.suffix.lower().lstrip('.')}"):
        if candidate not in taken:
            return candidate
    return f"{base}-{sha256_of(source)[:8]}"


def file_document(source, destination, fields, summary):
    return file_binary_document(source, destination, fields, summary)


def file_binary_document(source, destination, fields, summary):
    shutil.copy2(source, destination)
    fields["sha256"] = sha256_of(destination)
    derived_directory = destination.parent / ".derived" / fields["sha256"]
    try:
        derivation = derive(destination, fields["kind"], derived_directory)
    except Failure:
        destination.unlink()
        shutil.rmtree(derived_directory, ignore_errors=True)
        raise
    derived_directory.mkdir(parents=True, exist_ok=True)
    (derived_directory / "content.txt").write_text(derivation["text"], encoding="utf-8")
    fields["summary"] = summary or opening_line(derivation["text"]) or fields["title"]
    sidecar_path(destination).write_text(dump_frontmatter(fields) + sidecar_body(fields, derivation, derived_directory), encoding="utf-8")
    notes = list(derivation["notes"])
    if not summary:
        notes.append("summary drafted from the derived text; review it")
    return notes


def sidecar_path(document_path):
    return document_path.with_name(document_path.name + ".md")


def opening_line(text):
    collapsed = " ".join(text.split())
    if len(collapsed) <= SUMMARY_LIMIT:
        return collapsed
    return collapsed[: SUMMARY_LIMIT - 1].rstrip() + "…"


def sidecar_body(fields, derivation, derived_directory):
    lines = [f"# {fields['title']}", "", fields["summary"], ""]
    if derivation["text"].strip():
        lines += ["## Text", "", derivation["text"].strip()[:OPENING_LIMIT], ""]
    parts = sorted(path.relative_to(derived_directory.parent.parent) for path in derived_directory.rglob("*") if path.is_file())
    if parts:
        lines += ["## Derived", ""] + [f"- {part}" for part in parts] + [""]
    if derivation["notes"]:
        lines += ["## Skipped", ""] + [f"- {note}" for note in derivation["notes"]] + [""]
    return "\n".join(lines)


def derive(path, kind, derived_directory):
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return derivation(path.read_text(encoding="utf-8", errors="replace"), [])
    if suffix == ".pdf":
        return derive_pdf(path, derived_directory)
    if suffix == ".pptx":
        return derive_pptx(path, derived_directory)
    if suffix == ".docx":
        return derive_docx(path, derived_directory)
    if suffix == ".xlsx":
        return derive_xlsx(path, derived_directory)
    if suffix in (".html", ".htm"):
        return derive_html(path, derived_directory)
    if suffix in IMAGE_SUFFIXES:
        return derive_image(path, derived_directory)
    if suffix in VIDEO_SUFFIXES + AUDIO_SUFFIXES:
        return derive_media(path, derived_directory, kind)
    return derivation("", [f"no text derivation for {suffix} files"])


def derivation(text, notes):
    return {"text": text, "notes": notes}


def write_parts(derived_directory, sections):
    text_directory = derived_directory / "text"
    filled = [section for section in sections if section[1].strip()]
    for number, (label, content) in enumerate(filled, start=1):
        part = text_directory / f"{number:02d}-{slugify(label) or 'part'}.md"
        part.parent.mkdir(parents=True, exist_ok=True)
        part.write_text(f"# {label}\n\n{content.strip()}\n", encoding="utf-8")


def joined_text(sections):
    return "\n\n".join(content.strip() for _, content in sections if content.strip())


def derive_pdf(path, derived_directory):
    if shutil.which("pdftotext") is None:
        return derivation("", ["pdftotext is not installed, so no text was derived"])
    output = run_command(["pdftotext", "-enc", "UTF-8", str(path), "-"])
    pages = [(f"page {number}", page) for number, page in enumerate(output.split("\f"), start=1)]
    write_parts(derived_directory, pages)
    notes = [] if joined_text(pages) else ["the PDF carries no text layer; OCR it to derive text"]
    return derivation(joined_text(pages), notes)


def derive_pptx(path, derived_directory):
    with zipfile.ZipFile(path) as archive:
        slides = numbered_members(archive, SLIDE_PATTERN)
        notes = dict(numbered_members(archive, NOTES_PATTERN))
        sections = [(f"slide {number}", slide_text(archive, member, notes.get(number))) for number, member in slides]
    write_parts(derived_directory, sections)
    return derivation(joined_text(sections), [])


def numbered_members(archive, pattern):
    matches = [(int(match.group(1)), match.group(0)) for match in map(pattern.fullmatch, archive.namelist()) if match]
    return sorted(matches)


def slide_text(archive, member, notes_member):
    lines = xml_paragraphs(archive.read(member).decode("utf-8", errors="replace"))
    note_lines = xml_paragraphs(archive.read(notes_member).decode("utf-8", errors="replace")) if notes_member else []
    if note_lines:
        lines += ["", "Notes:"] + note_lines
    return "\n".join(lines)


def xml_paragraphs(xml):
    paragraphs = ["".join(XML_TEXT_PATTERN.findall(paragraph)).strip() for paragraph in XML_PARAGRAPH_PATTERN.findall(xml)]
    return [html.unescape(paragraph) for paragraph in paragraphs if paragraph]


def derive_docx(path, derived_directory):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
    sections = docx_sections(xml)
    write_parts(derived_directory, sections)
    return derivation(joined_text(sections), [])


def docx_sections(xml):
    sections = [("document", [])]
    for paragraph in XML_PARAGRAPH_PATTERN.findall(xml):
        text = html.unescape("".join(XML_TEXT_PATTERN.findall(paragraph)).strip())
        if not text:
            continue
        if DOCX_HEADING_PATTERN.search(paragraph):
            sections.append((text, []))
            continue
        sections[-1][1].append(text)
    return [(label, "\n".join(lines)) for label, lines in sections]


def derive_xlsx(path, derived_directory):
    try:
        import openpyxl
    except ImportError:
        return derivation("", ["openpyxl is not installed, so no sheet was derived"])
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        csv_path = derived_directory / "sheets" / f"{slugify(sheet.title) or 'sheet'}.csv"
        row_count = write_sheet_csv(sheet, csv_path)
        lines.append(f"{sheet.title}: {row_count} rows, {sheet.max_column} columns")
    return derivation("\n".join(lines), [])


def write_sheet_csv(sheet, csv_path):
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    row_count = 0
    with csv_path.open("w", encoding="utf-8", newline="") as csv_file:
        writer = csv.writer(csv_file)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if cell is None else cell for cell in row])
            row_count += 1
    return row_count


class TextExtractor(html.parser.HTMLParser):
    SKIPPED_TAGS = ("script", "style", "noscript")
    BLOCK_TAGS = ("p", "div", "section", "li", "br", "h1", "h2", "h3", "h4", "tr")

    def __init__(self):
        super().__init__()
        self.pieces = []
        self.skip_depth = 0

    def handle_starttag(self, tag, attributes):
        if tag in self.SKIPPED_TAGS:
            self.skip_depth += 1
        if tag in self.BLOCK_TAGS:
            self.pieces.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIPPED_TAGS and self.skip_depth > 0:
            self.skip_depth -= 1
        if tag in self.BLOCK_TAGS:
            self.pieces.append("\n")

    def handle_data(self, data):
        if self.skip_depth == 0:
            self.pieces.append(data)


def derive_html(path, derived_directory):
    extractor = TextExtractor()
    extractor.feed(path.read_text(encoding="utf-8", errors="replace"))
    lines = [" ".join(line.split()) for line in "".join(extractor.pieces).splitlines()]
    text = "\n".join(line for line in lines if line)
    write_parts(derived_directory, [("body", text)])
    return derivation(text, [])


def derive_image(path, derived_directory):
    notes = []
    text = ""
    if shutil.which("tesseract") is None:
        notes.append("tesseract is not installed, so no OCR text was derived")
    else:
        text = run_command(["tesseract", str(path), "-"]).strip()
        if not text:
            notes.append("OCR found no text")
    write_parts(derived_directory, [("ocr", text)])
    notes += write_thumbnail(path, derived_directory / "thumbnail.png")
    return derivation(text, notes)


def write_thumbnail(path, thumbnail_path):
    thumbnail_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        from PIL import Image
    except ImportError:
        return write_thumbnail_with_ffmpeg(path, thumbnail_path)
    with Image.open(path) as image:
        image.thumbnail((512, 512))
        image.convert("RGB").save(thumbnail_path)
    return []


def write_thumbnail_with_ffmpeg(path, thumbnail_path):
    if shutil.which("ffmpeg") is None:
        return ["neither Pillow nor ffmpeg is available, so no thumbnail was made"]
    run_command(["ffmpeg", "-y", "-loglevel", "error", "-i", str(path), "-vf", "scale='min(512,iw)':-2", "-frames:v", "1", str(thumbnail_path)])
    return []


def derive_media(path, derived_directory, kind):
    if shutil.which("ffmpeg") is None:
        return derivation("", ["ffmpeg is not installed, so no transcript or contact sheet was derived"])
    notes = []
    if path.suffix.lower() in VIDEO_SUFFIXES:
        notes += write_contact_sheet(path, derived_directory / "contact-sheet.png")
    transcript, transcript_notes = transcribe(path, derived_directory)
    notes += transcript_notes
    sections = transcript_sections(transcript)
    write_parts(derived_directory, sections)
    return derivation(joined_text(sections), notes)


def write_contact_sheet(path, sheet_path):
    sheet_path.parent.mkdir(parents=True, exist_ok=True)
    run_command(["ffmpeg", "-y", "-loglevel", "error", "-i", str(path), "-vf", "fps=1/30,scale=320:-1,tile=4x4", "-frames:v", "1", str(sheet_path)])
    return []


def transcribe(path, derived_directory):
    if shutil.which("whisper") is None:
        return "", ["whisper is not installed, so no transcript was derived"]
    audio_path = derived_directory / "audio.wav"
    audio_path.parent.mkdir(parents=True, exist_ok=True)
    run_command(["ffmpeg", "-y", "-loglevel", "error", "-i", str(path), "-ac", "1", "-ar", "16000", str(audio_path)])
    run_command(["whisper", str(audio_path), "--output_format", "txt", "--output_dir", str(derived_directory)])
    audio_path.unlink()
    transcript_path = derived_directory / "audio.txt"
    transcript = transcript_path.read_text(encoding="utf-8") if transcript_path.exists() else ""
    return transcript, ([] if transcript else ["whisper produced no transcript"])


def transcript_sections(transcript, lines_per_part=60):
    lines = [line for line in transcript.splitlines() if line.strip()]
    chunks = [lines[start : start + lines_per_part] for start in range(0, len(lines), lines_per_part)]
    return [(f"transcript {number}", "\n".join(chunk)) for number, chunk in enumerate(chunks, start=1)]


def run_command(command):
    completed = subprocess.run(command, capture_output=True, text=True, errors="replace", check=False)
    if completed.returncode != 0:
        raise Failure(f"{command[0]} failed: {completed.stderr.strip()}")
    return completed.stdout


def mark_superseded(root, superseded_id):
    for document in all_documents(root):
        if document.fields and document.fields.get("id") == superseded_id:
            document.fields["status"] = "superseded"
            _, body = split_frontmatter(document.metadata_path.read_text(encoding="utf-8"))
            document.metadata_path.write_text(dump_frontmatter(document.fields) + body, encoding="utf-8")
            return
    raise Failure(f"--supersedes {superseded_id} names nothing")


def report_ingest(root, destination, fields, notes):
    print(f"filed {destination.relative_to(root)} as {fields['id']}")
    for note in notes:
        print(f"note: {note}")


class Document:
    def __init__(self, root, path, metadata_path, is_binary):
        self.path = path.relative_to(root).as_posix()
        self.metadata_path = metadata_path
        self.is_binary = is_binary
        self.fields = load_frontmatter(metadata_path) if metadata_path.exists() else None


def all_documents(root):
    documents = []
    for folder, _ in category_folders(root):
        documents += folder_documents(root, root / folder)
    return documents


def folder_documents(root, folder_path):
    if not folder_path.is_dir():
        return []
    candidates = sorted(path for path in folder_path.rglob("*") if is_document_candidate(path, folder_path))
    return [document_for(root, path) for path in candidates if not is_sidecar(path)]


def is_document_candidate(path, folder_path):
    if not path.is_file() or path.name.startswith(".") or path.name == "catalog.jsonl":
        return False
    relative_parts = path.relative_to(folder_path).parts
    return relative_parts[0] != "archive" and not any(part.startswith(".") for part in relative_parts)


def is_sidecar(path):
    return path.suffix == ".md" and path.with_suffix("").suffix != "" and path.with_suffix("").is_file()


def document_for(root, path):
    return Document(root, path, sidecar_path(path), True)


def split_frontmatter(text):
    match = FRONTMATTER_PATTERN.match(text)
    if match is None:
        return None, text
    return match.group(1), match.group(2)


def load_frontmatter(path):
    import yaml

    header, _ = split_frontmatter(path.read_text(encoding="utf-8", errors="replace"))
    if header is None:
        return None
    try:
        loaded = yaml.safe_load(header)
    except yaml.YAMLError:
        return None
    return normalize_dates(loaded) if isinstance(loaded, dict) else None


def normalize_dates(value):
    if isinstance(value, dict):
        return {key: normalize_dates(item) for key, item in value.items()}
    if isinstance(value, list):
        return [normalize_dates(item) for item in value]
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.isoformat()[:10]
    return value


def dump_frontmatter(fields):
    import yaml

    return "---\n" + yaml.safe_dump(with_date_objects(fields), allow_unicode=True, sort_keys=False, default_flow_style=None, width=1000) + "---\n"


def with_date_objects(value):
    if isinstance(value, dict):
        return {key: with_date_objects(item) for key, item in value.items()}
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return datetime.date.fromisoformat(value)
    return value


def sha256_of(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def frontmatter_hash(documents):
    payload = [[document.path, document.fields] for document in sorted(documents, key=lambda document: document.path)]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def run_index(arguments):
    root = Path(arguments.directory).resolve()
    load_company(root)
    written = write_generated_files(root)
    for path in written:
        print(f"wrote {path.relative_to(root)}")


def write_generated_files(root):
    written = []
    category_hashes = []
    index_rows = []
    for folder, code in category_folders(root):
        documents = folder_documents(root, root / folder)
        catalog_hash = frontmatter_hash(documents)
        write_catalog(root / folder / "catalog.jsonl", folder, catalog_hash, documents)
        written.append(root / folder / "catalog.jsonl")
        category_hashes.append(catalog_hash)
        index_rows.append(index_row(folder, code, documents))
    index_hash = hashlib.sha256("".join(category_hashes).encode("ascii")).hexdigest()
    (root / "INDEX.md").write_text(index_text(index_rows, index_hash), encoding="utf-8")
    return written + [root / "INDEX.md"]


def write_catalog(catalog_path, folder, catalog_hash, documents):
    catalog_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [json.dumps({"catalog": folder, "frontmatterHash": catalog_hash}, ensure_ascii=False)]
    for document in documents:
        if document.fields:
            lines.append(json.dumps({"path": document.path, **document.fields}, ensure_ascii=False))
    catalog_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def index_row(label, code, documents):
    dates = [document.fields.get("date", "") for document in documents if document.fields]
    last_change = max((date for date in dates if date), default="")
    return f"| {label} | {code} | {len(documents)} | {last_change} |"


def index_text(rows, index_hash):
    lines = ["# Index", "", "| Category | Code | Documents | Last change |", "|---|---|---|---|"]
    lines += rows
    lines += ["", f"<!-- frontmatter-hash: {index_hash} -->", ""]
    return "\n".join(lines)


def run_check(arguments):
    root = Path(arguments.directory).resolve()
    load_company(root)
    documents = all_documents(root)
    failures = []
    failures += document_failures(root, documents)
    failures += reference_failures(documents)
    failures += directory_failures(root)
    failures += inbox_failures(root)
    failures += generated_file_failures(root)
    for failure in failures:
        print(f"fail: {failure}")
    if failures:
        return 1
    print(f"ok {len(documents)} documents")
    return 0


def document_failures(root, documents):
    failures = []
    for document in documents:
        if document.fields is None:
            failures.append(f"{document.path}: {'binary without a sidecar' if document.is_binary else 'no frontmatter'}")
            continue
        failures += field_failures(document)
        failures += placement_failures(root, document)
        failures += hash_failures(root, document)
    return failures


def field_failures(document):
    fields = document.fields
    failures = [f"{document.path}: missing required field {field}" for field in REQUIRED_FIELDS if not fields.get(field)]
    if document.is_binary and not fields.get("sha256"):
        failures.append(f"{document.path}: missing required field sha256")
    if fields.get("kind") and fields["kind"] not in KINDS:
        failures.append(f"{document.path}: kind {fields['kind']} is not one of {', '.join(KINDS)}")
    if fields.get("status") and fields["status"] not in STATUSES:
        failures.append(f"{document.path}: status {fields['status']} is not one of {', '.join(STATUSES)}")
    if len(str(fields.get("summary", ""))) > SUMMARY_LIMIT:
        failures.append(f"{document.path}: summary over {SUMMARY_LIMIT} characters")
    return failures


def placement_failures(root, document):
    code = document.fields.get("categoryCode")
    folders = dict((code, folder) for folder, code in category_folders(root))
    expected = folders.get(code)
    if expected is None or not document.path.startswith(expected + "/"):
        return [f"{document.path}: category {code} disagrees with its path"]
    return []


def hash_failures(root, document):
    if not document.is_binary or not document.fields.get("sha256"):
        return []
    if sha256_of(root / document.path) == document.fields["sha256"]:
        return []
    return [f"{document.path}: sha256 in the sidecar does not match the file"]


def reference_failures(documents):
    ids = {}
    failures = []
    for document in documents:
        document_id = (document.fields or {}).get("id")
        if not document_id:
            continue
        if document_id in ids:
            failures.append(f"{document.path}: duplicate id {document_id} (also {ids[document_id]})")
        ids.setdefault(document_id, document.path)
    for document in documents:
        failures += dangling_reference_failures(document, ids)
    return failures


def dangling_reference_failures(document, ids):
    fields = document.fields or {}
    failures = []
    if fields.get("supersedes") and fields["supersedes"] not in ids:
        failures.append(f"{document.path}: supersedes {fields['supersedes']} which names nothing")
    published = fields.get("published")
    if isinstance(published, dict) and published.get("from") not in ids:
        failures.append(f"{document.path}: published from {published.get('from')} which names nothing")
    return failures


def directory_failures(root):
    counts = {}
    for document in all_documents(root):
        directory = document.path.rsplit("/", 1)[0]
        counts[directory] = counts.get(directory, 0) + 1
    return [f"{directory}/: {count} entries, over {DIRECTORY_LIMIT}" for directory, count in sorted(counts.items()) if count > DIRECTORY_LIMIT]


def inbox_failures(root):
    oldest_allowed = datetime.date.today() - datetime.timedelta(days=INBOX_LIMIT_DAYS)
    failures = []
    for document in folder_documents(root, root / "X-inbox"):
        arrived = datetime.date.fromtimestamp((root / document.path).stat().st_mtime)
        if arrived < oldest_allowed:
            failures.append(f"{document.path}: unclassified for over {INBOX_LIMIT_DAYS} days in X-inbox/")
    return failures


def generated_file_failures(root):
    failures = []
    category_hashes = []
    for folder, _ in category_folders(root):
        expected = frontmatter_hash(folder_documents(root, root / folder))
        category_hashes.append(expected)
        if catalog_hash(root / folder / "catalog.jsonl") != expected:
            failures.append(f"{folder}/catalog.jsonl: stale generated file")
    expected_index = hashlib.sha256("".join(category_hashes).encode("ascii")).hexdigest()
    if index_hash(root / "INDEX.md") != expected_index:
        failures.append("INDEX.md: stale generated file")
    return failures


def catalog_hash(catalog_path):
    if not catalog_path.is_file():
        return None
    first_line = catalog_path.read_text(encoding="utf-8").split("\n", 1)[0]
    try:
        return json.loads(first_line).get("frontmatterHash")
    except (json.JSONDecodeError, AttributeError):
        return None


def index_hash(index_path):
    if not index_path.is_file():
        return None
    match = INDEX_HASH_PATTERN.search(index_path.read_text(encoding="utf-8"))
    return match.group(1) if match else None


def run_search(arguments):
    root = Path(arguments.directory).resolve()
    catalogs = sorted(root.glob("*/catalog.jsonl"))
    lines = search_lines(arguments.query, catalogs)
    hits = [entry for entry in map(catalog_entry, lines) if entry and entry["path"].startswith(arguments.path or "")]
    for entry in hits:
        print("\t".join(str(entry.get(field, "")) for field in ("path", "id", "date", "title", "summary")))
    return 0 if hits else 1


def search_lines(query, catalogs):
    if not catalogs:
        return []
    if shutil.which("rg") is not None:
        completed = subprocess.run(["rg", "--no-filename", "--no-line-number", "-i", "-e", query, *map(str, catalogs)], capture_output=True, text=True, check=False)
        return completed.stdout.splitlines()
    pattern = re.compile(query, re.IGNORECASE)
    return [line for catalog in catalogs for line in catalog.read_text(encoding="utf-8").splitlines() if pattern.search(line)]


def catalog_entry(line):
    try:
        entry = json.loads(line)
    except json.JSONDecodeError:
        return None
    return entry if isinstance(entry, dict) and "path" in entry else None


if __name__ == "__main__":
    main()
