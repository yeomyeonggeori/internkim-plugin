from __future__ import annotations

import re

from openpyxl.utils import get_column_letter


MINIMUM_HEADER_CELLS = 3
MINIMUM_ITEM_ROWS = 2
NUMBERING = re.compile(r"\d{1,3}\.?")


def field_name(label: str, taken: set[str]) -> str:
    base = re.sub(r"\s+", "", label).strip(":：") or "field"
    name, counter = base, 2
    while name in taken:
        name, counter = f"{base}_{counter}", counter + 1
    taken.add(name)
    return name


def text_field(label: str, taken: set[str], at: dict) -> dict:
    return {"name": field_name(label, taken), "label": re.sub(r"\s+", " ", label).strip(), "type": "text", "at": at}


def list_field(label: str, headers: list[tuple[int, str]], taken: set[str], at: dict) -> dict:
    children_taken: set[str] = set()
    children = [{"name": field_name(header, children_taken), "label": header, "type": "text", "at": {"column": column}} for column, header in headers]
    return {"name": field_name(label, taken), "label": label, "type": "list", "at": at, "fields": children}


class GridTable:
    def __init__(self, texts: list[list[str]], identities: list[list[object]], formulas: set | None = None):
        self.texts = texts
        self.identities = identities
        self.formulas = formulas or set()

    def unique_columns(self, row: int) -> list[int]:
        identities = self.identities[row]
        return [column for column in range(len(identities)) if column == 0 or identities[column] is not identities[column - 1]]

    def continues_above(self, row: int, column: int) -> bool:
        return row > 0 and column < len(self.identities[row - 1]) and self.identities[row][column] is self.identities[row - 1][column]

    def is_empty(self, row: int, column: int) -> bool:
        return not self.texts[row][column].strip() and not self.continues_above(row, column) and (row, column) not in self.formulas


def item_rows(table: GridTable, header_row: int, numbering_column: int | None) -> int:
    count = 0
    for row in range(header_row + 1, len(table.texts)):
        filled = [column for column in table.unique_columns(row) if table.texts[row][column].strip() and column != numbering_column]
        numbered = numbering_column is None or NUMBERING.fullmatch(table.texts[row][numbering_column].strip() or "0")
        if filled or not numbered:
            break
        count += 1
    return count


def numbering_column(table: GridTable, header_row: int) -> int | None:
    if header_row + 1 >= len(table.texts):
        return None
    first = table.unique_columns(header_row + 1)[0]
    return first if NUMBERING.fullmatch(table.texts[header_row + 1][first].strip()) else None


def table_fields(table: GridTable, at_table: dict, list_label: str, taken: set[str]) -> list[dict]:
    fields, row = [], 0
    while row < len(table.texts):
        headers = [(column, table.texts[row][column].strip()) for column in table.unique_columns(row) if table.texts[row][column].strip()]
        numbered = numbering_column(table, row)
        rows = item_rows(table, row, numbered) if len(headers) >= MINIMUM_HEADER_CELLS else 0
        if rows >= MINIMUM_ITEM_ROWS:
            item_headers = [(column, header) for column, header in headers if column != numbered]
            field = list_field(list_label or item_headers[0][1], item_headers, taken, {**at_table, "firstRow": row + 1, "rows": rows})
            for child in field["fields"]:
                if (row + 1, child["at"]["column"]) in table.formulas:
                    child["type"] = "ignore"
                    child["description"] = "the form computes it with its own formula"
            fields.append(field)
            row += rows + 1
            continue
        fields.extend(row_fields(table, row, at_table, taken))
        row += 1
    return fields


def row_fields(table: GridTable, row: int, at_table: dict, taken: set[str]) -> list[dict]:
    fields = []
    columns = table.unique_columns(row)
    for index, column in enumerate(columns):
        label = table.texts[row][column].strip()
        if not label or table.continues_above(row, column):
            continue
        right = columns[index + 1] if index + 1 < len(columns) else None
        if right is not None and table.is_empty(row, right):
            fields.append(text_field(label, taken, {**at_table, "row": row, "column": right}))
        elif row + 1 < len(table.texts) and column < len(table.texts[row + 1]) and table.is_empty(row + 1, column) and not any(table.texts[row + 1][other].strip() and not table.continues_above(row + 1, other) for other in table.unique_columns(row + 1)):
            fields.append(text_field(label, taken, {**at_table, "row": row + 1, "column": column}))
    return fields


def docx_grid(table) -> GridTable:
    rows = [list(row.cells) for row in table.rows]
    return GridTable([[cell.text for cell in row] for row in rows], [[cell._tc for cell in row] for row in rows])


def preceding_text(document, table) -> str:
    previous = table._tbl.getprevious()
    while previous is not None and previous.tag.endswith("}p"):
        text = "".join(node.text or "" for node in previous.iter() if node.tag.endswith("}t")).strip()
        if text:
            return text
        previous = previous.getprevious()
    return ""


def docx_form_fields(path: str) -> list[dict]:
    from docx import Document

    document = Document(path)
    taken: set[str] = set()
    fields = []
    for index, table in enumerate(document.tables):
        fields.extend(table_fields(docx_grid(table), {"table": index}, preceding_text(document, table), taken))
    return fields


def xlsx_grid(worksheet) -> GridTable:
    owners = {}
    for merged in worksheet.merged_cells.ranges:
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                owners[(row, column)] = merged
    texts, identities = [], []
    for row in range(1, worksheet.max_row + 1):
        texts.append([cell_text(worksheet.cell(row, column).value) for column in range(1, worksheet.max_column + 1)])
        identities.append([owners.get((row, column)) or (row, column) for column in range(1, worksheet.max_column + 1)])
    for row_index, row in enumerate(identities):
        for column_index, identity in enumerate(row):
            if isinstance(identity, tuple):
                row[column_index] = object()
    for (row, column), merged in owners.items():
        texts[row - 1][column - 1] = cell_text(worksheet.cell(merged.min_row, merged.min_col).value)
    formulas = {(row - 1, column - 1) for row in range(1, worksheet.max_row + 1) for column in range(1, worksheet.max_column + 1) if isinstance(worksheet.cell(row, column).value, str) and worksheet.cell(row, column).value.startswith("=")}
    return GridTable(texts, identities, formulas)


def cell_text(value: object) -> str:
    if value is None or isinstance(value, str) and value.startswith("="):
        return ""
    return str(value)


def xlsx_form_fields(path: str) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    taken: set[str] = set()
    fields = []
    for worksheet in workbook.worksheets:
        fields.extend(table_fields(xlsx_grid(worksheet), {"sheet": worksheet.title}, "", taken))
    return fields


def cell_name(at: dict, row: int | None = None, column: int | None = None) -> str:
    return f"{get_column_letter((at['column'] if column is None else column) + 1)}{(at['row'] if row is None else row) + 1}"
