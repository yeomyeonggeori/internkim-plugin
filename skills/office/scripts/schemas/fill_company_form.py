from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl.utils import get_column_letter

from core.office_result import ERROR, IssueKind, OfficeFailure
from schemas.expression import Binary, Call, Number, Path as ExpressionPath
from schemas.resolution import Instance, formatted
from schemas.text_blanks import fill_paragraph, filled_text


FORM_ROWS_FULL = IssueKind("FORM_ROWS_FULL", ERROR, "the form has fewer item rows than the values hold", "split the items over two copies of the form, or ask the person which items belong on this one")


def placements(instance: Instance) -> dict:
    return instance.schema.extra.get("placements") or {}


def scalar_names(instance: Instance) -> list[str]:
    names = [field.name for field in instance.schema.fields if field.type != "list"]
    names += list(instance.schema.known)
    names += [derived.name for derived in instance.schema.derived if "." not in derived.name]
    return [name for name in names if name in placements(instance)]


def list_rows(instance: Instance, name: str, capacity: int) -> list[dict]:
    rows = instance.rows.get(name) or []
    if len(rows) > capacity:
        raise OfficeFailure(FORM_ROWS_FULL.issue(f"values.{name}: {len(rows)} items for the form's {capacity} rows", f"values.{name}"))
    return rows


def write_docx_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def fill_docx_form(instance: Instance, output_path: Path) -> list:
    from docx import Document

    document = Document(instance.schema.template)
    tables = document.tables
    for name in scalar_names(instance):
        at = placements(instance)[name]["at"]
        text = formatted(instance, name, None, None)
        if text is not None and "blank" not in at:
            write_docx_cell(tables[at["table"]].cell(at["row"], at["column"]), text)
    for (table, row, column, paragraph), filled in blanks_by_paragraph(instance).items():
        container = document if table is None else tables[table].cell(row, column)
        fill_paragraph(container.paragraphs[paragraph], filled)
    for field in instance.schema.fields:
        if not field.is_record_list or field.name not in placements(instance):
            continue
        placement = placements(instance)[field.name]
        table = tables[placement["at"]["table"]]
        for index, row in enumerate(list_rows(instance, field.name, placement["at"]["rows"])):
            for child, child_at in placement["fields"].items():
                text = formatted(instance, child, None, row) if child in row else None
                if text is not None:
                    write_docx_cell(table.cell(placement["at"]["firstRow"] + index, child_at["column"]), text)
    document.save(output_path)
    return []


def blank_value(instance: Instance, name: str) -> object:
    value = instance.value(name, None)
    return value if isinstance(value, date) else formatted(instance, name, None, None)


def blanks_by_paragraph(instance: Instance) -> dict[tuple, dict[int, object]]:
    grouped: dict[tuple, dict[int, object]] = {}
    for name in scalar_names(instance):
        at = placements(instance)[name]["at"]
        if "blank" not in at:
            continue
        key = (at.get("table", at.get("sheet")), at.get("row"), at.get("column"), at.get("paragraph", 0))
        grouped.setdefault(key, {})[at["blank"]] = blank_value(instance, name)
    return grouped


def cell_value(instance: Instance, name: str, row: dict | None):
    value = instance.value(name, row)
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, (date, int, float)) or value is None:
        return value
    return formatted(instance, name, None, row)


def live_formula(expression, references: dict) -> str | None:
    if isinstance(expression, Number):
        return str(expression.value)
    if isinstance(expression, ExpressionPath):
        return references.get(expression.name)
    if isinstance(expression, Binary):
        left, right = live_formula(expression.left, references), live_formula(expression.right, references)
        return None if left is None or right is None else f"({left}{expression.operator}{right})"
    if isinstance(expression, Call) and expression.function == "sum" and isinstance(expression.arguments[0], ExpressionPath):
        column_range = references.get(f"range:{expression.arguments[0].name}")
        return None if column_range is None else f"SUM({column_range})"
    return None


def fill_xlsx_form(instance: Instance, output_path: Path) -> list:
    from openpyxl import load_workbook

    workbook = load_workbook(instance.schema.template)
    references: dict[str, str] = {}
    for field in instance.schema.fields:
        if field.is_record_list and field.name in placements(instance):
            fill_xlsx_rows(instance, workbook, field.name, references)
    for (sheet, row, column, _paragraph), filled in blanks_by_paragraph(instance).items():
        cell = workbook[sheet].cell(row + 1, column + 1)
        cell.value = filled_text(str(cell.value or ""), filled)
    for name in scalar_names(instance):
        at = placements(instance)[name]["at"]
        if "blank" in at:
            continue
        worksheet = workbook[at["sheet"]]
        derived = instance.schema.derived_named(name)
        formula = live_formula(derived.expression, references) if derived is not None else None
        value = f"={formula}" if formula else cell_value(instance, name, None)
        if value is not None:
            worksheet.cell(at["row"] + 1, at["column"] + 1).value = value
    from sheet.formulas.cache import evaluation_issues, save_workbook_with_values

    return evaluation_issues(save_workbook_with_values(workbook, str(output_path)))


def fill_xlsx_rows(instance: Instance, workbook, name: str, references: dict) -> None:
    placement = placements(instance)[name]
    worksheet = workbook[placement["at"]["sheet"]]
    first_row = placement["at"]["firstRow"] + 1
    rows = list_rows(instance, name, placement["at"]["rows"])
    for child, child_at in placement["fields"].items():
        letter = get_column_letter(child_at["column"] + 1)
        references[f"range:{name}.{child}"] = f"{letter}{first_row}:{letter}{first_row + len(rows) - 1}" if rows else ""
    for index, row in enumerate(rows):
        row_number = first_row + index
        row_references = {f"{name}.{child}": f"{get_column_letter(child_at['column'] + 1)}{row_number}" for child, child_at in placement["fields"].items()}
        for child, child_at in placement["fields"].items():
            derived = instance.schema.derived_named(f"{name}.{child}")
            formula = live_formula(derived.expression, row_references) if derived is not None else None
            value = f"={formula}" if formula else cell_value(instance, child, row) if child in row else None
            if value is not None:
                worksheet.cell(row_number, child_at["column"] + 1).value = value
