#!/usr/bin/env python3
from __future__ import annotations

from pptx import Presentation

from core.office_arguments import route_arguments
from core.office_result import Result, run_command
from powerpoint.description import describe_presentation
from core.page_selection import select_pages


def main() -> Result:
    arguments = parse_arguments()
    presentation = Presentation(arguments.file)
    numbers = select_pages(arguments.pages, len(presentation.slides))
    details = describe_presentation(presentation, numbers, arguments.detail)
    return Result(summary=f"read {len(numbers)} of {len(presentation.slides)} slides from {arguments.file}", output_path=arguments.file, details=details)


def parse_arguments():
    return route_arguments("read", "pptx")


if __name__ == "__main__":
    raise SystemExit(run_command(main))
