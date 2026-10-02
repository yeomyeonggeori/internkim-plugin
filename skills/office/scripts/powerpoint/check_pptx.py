from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

from powerpoint.definitions import PPTX_NOT_RENDERED, SLIDE_COUNT_MISMATCH
from render.office_preview import drawn_page_details
from core.office_arguments import route_arguments
from core.office_outputs import preview_directory
from core.office_result import Issue, Result, run_command
from core.text_checks import text_presence_issues
from powerpoint.layout_audit import audit_presentation, substitutions
from powerpoint.preview.document import preview_document
from core.page_selection import select_pages
from render.renderer import RenderFailed, RendererUnavailable, draw_preview


PREVIEW_NAME = "preview.html"
SLIDE_SELECTOR = "section[data-slide]"
BOLD_WEIGHT = 700
REGULAR_WEIGHT = 400


def main() -> Result:
    arguments = route_arguments("check", "pptx")
    source_path = Path(arguments.file).expanduser()
    presentation = Presentation(str(source_path))
    content = slide_count_issues(presentation, arguments.slide_count, source_path.name) + text_presence_issues(visible_text(presentation), arguments.required_text, arguments.forbidden_text)
    return check_presentation(presentation, source_path, arguments.pages, arguments.output_directory, not arguments.no_preview, content)


def slide_count_issues(presentation, requested: int | None, name: str) -> list[Issue]:
    count = len(presentation.slides)
    if requested is None or requested == count:
        return []
    return [SLIDE_COUNT_MISMATCH.issue(f"{name} has {count} slides, but {requested} were requested", name)]


def visible_text(presentation) -> str:
    paragraphs = (paragraph for slide in presentation.slides for paragraph in slide.element.iter(qn("a:p")))
    return "\n".join("".join(text.text or "" for text in paragraph.iter(qn("a:t"))) for paragraph in paragraphs)


def check_presentation(presentation, source_path: Path, slide_selection: str | None, output_directory: str | None, preview: bool, content_issues: list[Issue]) -> Result:
    slides = list(presentation.slides)
    numbers = select_pages(slide_selection, len(slides))
    audit = audit_presentation(presentation, [(number, slides[number - 1]) for number in numbers])
    issues = content_issues + list(audit.issues)
    details = {"checkedSlides": numbers, "fontsMeasuredWith": substitutions(audit.faces)}
    if preview:
        preview_details, preview_issues = write_preview(presentation, numbers, preview_directory(source_path, output_directory))
        details.update(preview_details)
        issues.extend(preview_issues)
    return Result(summary=f"checked {len(numbers)} slides of {source_path}: {len(audit.issues)} layout issues", output_path=str(source_path), issues=tuple(issues), details=details)


def write_preview(presentation, numbers: list[int], directory: Path) -> tuple[dict, list]:
    preview = preview_document(presentation, numbers)
    directory.mkdir(parents=True, exist_ok=True)
    preview_path = directory / PREVIEW_NAME
    preview_path.write_text(preview.html, encoding="utf-8")
    fonts = preview_fonts(preview.faces)
    details = {"preview": str(preview_path), "previewFonts": fonts}
    try:
        rendered = draw_preview(preview_path, SLIDE_SELECTOR, fonts)
    except (RendererUnavailable, RenderFailed) as reason:
        return details | {"seen": False}, [PPTX_NOT_RENDERED.issue(f"the preview of {len(numbers)} slides was written as HTML to {preview_path}, and no image of it was drawn: {reason}", str(preview_path))]
    return details | drawn_page_details(rendered), []


def preview_fonts(faces: frozenset) -> list[dict]:
    files = {(face.family, face.path, face.index, BOLD_WEIGHT if face.bold else REGULAR_WEIGHT) for _, face in faces}
    return [{"family": family, "path": path, "index": index, "weight": weight} for family, path, index, weight in sorted(files)]


if __name__ == "__main__":
    raise SystemExit(run_command(main))
