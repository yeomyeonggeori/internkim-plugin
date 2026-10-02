from __future__ import annotations

from dataclasses import dataclass, field

from pptx import Presentation
from pptx.oxml.ns import qn

from core.office_operations import OPERATION_NOT_APPLICABLE, TARGET_NOT_FOUND
from core.office_result import OfficeFailure
from powerpoint.model.geometry import Frame, SLIDE_FRAME, child_frame
from powerpoint.model.shape_kinds import shape_address, shape_kind


AVAILABLE_SHAPES_LISTED = 24


@dataclass
class PptxEditing:
    presentation: Presentation
    original_package: bytes
    slides: list
    slide_id_elements: list
    deleted_slides: set = field(default_factory=set)
    touched_slides: set = field(default_factory=set)
    deleted_shapes: set = field(default_factory=set)
    edited_slides: list = field(default_factory=list)
    insertion_points: dict = field(default_factory=dict)
    new_slide_anchors: dict = field(default_factory=dict)
    structure_changed: bool = False

    def mark_edited(self, slide) -> None:
        if all(existing is not slide for existing in self.edited_slides):
            self.edited_slides.append(slide)


@dataclass(frozen=True)
class ShapeTarget:
    slide_number: int
    slide: object
    shape: object
    frame: Frame
    address: str

    @property
    def element(self):
        return self.shape._element

    @property
    def kind(self) -> str:
        return shape_kind(self.element)

    @property
    def label(self) -> str:
        return f"shape {self.address} of slide {self.slide_number}"


def load_editing(path: str) -> PptxEditing:
    with open(path, "rb") as package_file:
        original_package = package_file.read()
    presentation = Presentation(path)
    return PptxEditing(presentation, original_package, list(presentation.slides), list(presentation.slides._sldIdLst))


def resolve_slide(editing: PptxEditing, number: int, location: str):
    if number > len(editing.slides):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: slide {number} does not exist", location, f"use a slide number from 1 to {len(editing.slides)}"))
    if number in editing.deleted_slides:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: slide {number} is deleted by another operation in this batch", location))
    editing.touched_slides.add(number)
    return editing.slides[number - 1]


def resolve_shape(editing: PptxEditing, operation: dict, location: str, address_field: str = "shape") -> ShapeTarget:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")
    address = str(operation[address_field])
    shapes, frame, shape = slide.shapes, SLIDE_FRAME, None
    walked = []
    for step in address.split("."):
        if shape is not None:
            if shape_kind(shape._element) != "group":
                raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.{address_field}: shape {'.'.join(walked)} of slide {operation['slide']} is not a group", f"{location}.{address_field}"))
            frame = child_frame(frame, shape._element)
            shapes = shape.shapes
        shape = shape_at(shapes, int(step), operation["slide"], walked, f"{location}.{address_field}")
        walked.append(step)
    if id(shape._element) in editing.deleted_shapes:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.{address_field}: shape {address} of slide {operation['slide']} is deleted by another operation in this batch", f"{location}.{address_field}"))
    return ShapeTarget(operation["slide"], slide, shape, frame, address)


def shape_at(shapes, index: int, slide_number: int, walked: list[str], location: str):
    listed = list(shapes)
    if index < len(listed):
        return listed[index]
    prefix = ".".join(walked)
    owner = f"group {prefix} of slide {slide_number}" if walked else f"slide {slide_number}"
    raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: {owner} has no shape {index}", location, available_shapes(listed, prefix)))


def available_shapes(shapes: list, prefix: str = "") -> str:
    if not shapes:
        return "it holds no shapes; read the deck again"
    entries = [f"{shape_address(prefix, index)} {shape.name!r} ({shape_kind(shape._element)})" for index, shape in enumerate(shapes[:AVAILABLE_SHAPES_LISTED])]
    more = f", and {len(shapes) - AVAILABLE_SHAPES_LISTED} more" if len(shapes) > AVAILABLE_SHAPES_LISTED else ""
    return "use one of " + ", ".join(entries) + more


def require_kind(target: ShapeTarget, kinds: tuple[str, ...], location: str, purpose: str) -> None:
    if target.kind in kinds:
        return
    raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.shape: {target.label} is a {target.kind}; {purpose}", f"{location}.shape", shapes_of_kind_suggestion(target, kinds)))


def shapes_of_kind_suggestion(target: ShapeTarget, kinds: tuple[str, ...]) -> str:
    wanted = " or ".join(kinds)
    addresses = [str(index) for index, shape in enumerate(target.slide.shapes) if shape_kind(shape._element) in kinds]
    if addresses:
        return f"use shape {' or '.join(addresses)}, the {wanted} of slide {target.slide_number}"
    return f"slide {target.slide_number} holds no {wanted}; run office read to find the slide that does"


def text_body(element):
    return element.find(qn("p:txBody"))


def require_text(target: ShapeTarget, location: str):
    body = text_body(target.element)
    if body is None:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.shape: {target.label} is a {target.kind} that holds no text", f"{location}.shape", "use set_table_cell or format_table_cells for a table, or pick a shape office read shows with text"))
    return target.shape.text_frame


def resolve_paragraphs(text_frame, operation: dict, target: ShapeTarget, location: str) -> list:
    paragraphs = list(text_frame.paragraphs)
    index = operation.get("paragraph")
    if index is None:
        return paragraphs
    if index >= len(paragraphs):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.paragraph: {target.label} has no paragraph {index}", f"{location}.paragraph", f"use a paragraph from 0 to {len(paragraphs) - 1}"))
    return [paragraphs[index]]


def live_slides(editing: PptxEditing) -> list:
    return [slide for number, slide in enumerate(editing.slides, start=1) if number not in editing.deleted_slides]
