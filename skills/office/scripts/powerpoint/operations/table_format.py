from __future__ import annotations

from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

from core.office_operations import Change
from core.office_result import INVALID_VALUE, OfficeFailure
from powerpoint.operations.elements import rgb
from powerpoint.operations.tables import grid_columns, grow_frame, require_index, resolve_table, table_element, table_rows
from powerpoint.model.table_style_properties import set_table_style, style_named
from powerpoint.model.table_styles import PART_FLAGS
from powerpoint.operations.targets import PptxEditing, ShapeTarget
from powerpoint.operations.text import ALIGNMENTS, ANCHORS, RUN_STYLE_NAMES, apply_run_style, character_properties, require_any


CELL_LINE_TAGS = ("a:lnL", "a:lnR", "a:lnT", "a:lnB")
CELL_LINE_SUCCESSORS = ("a:lnTlToBr", "a:lnBlToTr", "a:cell3D", "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill", "a:headers", "a:extLst")
DEFAULT_BORDER_POINTS = 1
CELL_FORMAT_NAMES = ("fill", *RUN_STYLE_NAMES, "align", "anchor", "borderColor", "borderWidth")


def plan_set_table_style(editing: PptxEditing, operation: dict, location: str) -> Change:
    require_any(operation, ("style", *PART_FLAGS), location)
    target = resolve_table(editing, operation, location)

    def change() -> str:
        properties = table_element(target).get_or_add_tblPr()
        if operation.get("style") is not None:
            set_table_style(properties, style_named(operation["style"]))
        for flag in PART_FLAGS:
            if operation.get(flag) is not None:
                properties.set(flag, "1" if operation[flag] else "0")
        editing.mark_edited(target.slide)
        return f"restyled {target.label}"
    return change


def plan_format_table_cells(editing: PptxEditing, operation: dict, location: str) -> Change:
    require_any(operation, CELL_FORMAT_NAMES, location)
    if operation.get("borderWidth") is not None and operation.get("borderColor") is None:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.borderWidth: a width needs a borderColor to draw", f"{location}.borderWidth", "add borderColor"))
    target = resolve_table(editing, operation, location)
    rows = block_span(operation, "row", "rows", len(table_rows(target)), target, location)
    columns = block_span(operation, "column", "columns", len(grid_columns(target)), target, location)
    cells = [target.shape.table.cell(row, column) for row in rows for column in columns]

    def change() -> str:
        for cell in cells:
            format_cell(cell, operation)
        editing.mark_edited(target.slide)
        return f"formatted {len(cells)} cells of {target.label}"
    return change


def block_span(operation: dict, start_name: str, count_name: str, size: int, target: ShapeTarget, location: str) -> range:
    if operation.get(start_name) is None:
        if operation.get(count_name) is not None:
            raise OfficeFailure(INVALID_VALUE.issue(f"{location}.{count_name}: a count needs the {start_name} it starts at", f"{location}.{count_name}", f"add {start_name}, or leave out {count_name} for every {start_name}"))
        return range(size)
    start = operation[start_name]
    end = start + (operation.get(count_name) or 1)
    require_index(end - 1, size, start_name, target, f"{location}.{count_name if operation.get(count_name) else start_name}")
    return range(start, end)


def format_cell(cell, operation: dict) -> None:
    if operation.get("fill") is not None:
        fill_cell(cell, operation["fill"])
    for paragraph in cell.text_frame.paragraphs:
        for properties in character_properties(paragraph._p):
            apply_run_style(properties, operation)
        if operation.get("align") is not None:
            paragraph.alignment = ALIGNMENTS[operation["align"]]
    if operation.get("anchor") is not None:
        cell.vertical_anchor = ANCHORS[operation["anchor"]]
    if operation.get("borderColor") is not None:
        draw_cell_lines(cell._tc.get_or_add_tcPr(), operation["borderColor"], operation.get("borderWidth") or DEFAULT_BORDER_POINTS)


def fill_cell(cell, color: str) -> None:
    if color == "none":
        cell.fill.background()
        return
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(color)


def draw_cell_lines(cell_properties, color: str, width_points: float) -> None:
    for tag in CELL_LINE_TAGS:
        existing = cell_properties.find(qn(tag))
        if existing is not None:
            cell_properties.remove(existing)
        cell_properties.insert_element_before(cell_line(tag, color, width_points), *CELL_LINE_SUCCESSORS)


def cell_line(tag: str, color: str, width_points: float):
    line = OxmlElement(tag)
    line.set("w", str(Pt(width_points)))
    if color == "none":
        line.append(OxmlElement("a:noFill"))
        return line
    fill = OxmlElement("a:solidFill")
    color_element = OxmlElement("a:srgbClr")
    color_element.set("val", color.lstrip("#").upper())
    fill.append(color_element)
    line.append(fill)
    return line


def plan_set_table_column_width(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_table(editing, operation, location)
    require_index(operation["column"], len(grid_columns(target)), "column", target, f"{location}.column")
    column = target.shape.table.columns[operation["column"]]

    def change() -> str:
        grow_frame(target, "cx", operation["width"] - column.width)
        column.width = Emu(operation["width"])
        editing.mark_edited(target.slide)
        return f"set column {operation['column']} of {target.label} to {operation['width']} EMU wide"
    return change


def plan_set_table_row_height(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_table(editing, operation, location)
    require_index(operation["row"], len(table_rows(target)), "row", target, f"{location}.row")
    row = target.shape.table.rows[operation["row"]]

    def change() -> str:
        grow_frame(target, "cy", operation["height"] - row.height)
        row.height = Emu(operation["height"])
        editing.mark_edited(target.slide)
        return f"set row {operation['row']} of {target.label} to {operation['height']} EMU high"
    return change


TABLE_FORMAT_PLANNERS = {
    "set_table_style": plan_set_table_style,
    "format_table_cells": plan_format_table_cells,
    "set_table_column_width": plan_set_table_column_width,
    "set_table_row_height": plan_set_table_row_height,
}
