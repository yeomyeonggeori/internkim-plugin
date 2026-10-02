#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from core.office_arguments import route_arguments
from core.office_outputs import preview_directory
from core.office_result import Result, run_command
from core.page_selection import select_pages
from deck.check_pptx import write_preview


def main() -> Result:
    arguments = route_arguments("render", "pptx")
    source_path = Path(arguments.file).expanduser()
    presentation = Presentation(str(source_path))
    numbers = select_pages(arguments.pages, len(presentation.slides))
    directory = preview_directory(source_path, arguments.output_directory)
    details, issues = write_preview(presentation, numbers, directory)
    return Result(summary=f"drew {len(numbers)} of {len(presentation.slides)} slides of {source_path.name} in {directory}", output_path=str(directory), issues=tuple(issues), details=details)


if __name__ == "__main__":
    raise SystemExit(run_command(main))
