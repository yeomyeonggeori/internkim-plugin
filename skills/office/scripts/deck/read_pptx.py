#!/usr/bin/env python3
from __future__ import annotations

from pptx import Presentation

from office_inputs import office_file
from office_result import OfficeArgumentParser, Result, run_command
from pptx_description import describe_presentation
from pptx_slide_selection import select_slides


def main() -> Result:
    arguments = parse_arguments()
    presentation = Presentation(arguments.presentation_path)
    numbers = select_slides(arguments.slides, len(presentation.slides))
    details = describe_presentation(presentation, numbers, arguments.detail)
    return Result(summary=f"read {len(numbers)} of {len(presentation.slides)} slides from {arguments.presentation_path}", output_path=arguments.presentation_path, details=details)


def parse_arguments():
    parser = OfficeArgumentParser(description=(
        "Read a .pptx as numbered slides. Each shape has the index deck apply takes (3.1 is the second shape inside group 3), "
        "its id, name, kind, placeholder type, box in EMU and in percent of the slide, and its text with the effective font, size, bold and color. "
        "Shapes are listed back to front. Tables give their rows, charts their type, categories and series; "
        "links, transitions, comments and sections appear where the deck has them."
    ))
    parser.add_argument("presentation_path", type=office_file("pptx"))
    parser.add_argument("--slides", default="", help="slides to read, such as 2,4-6; default every slide")
    parser.add_argument("--detail", action="store_true", help="add each paragraph's runs with where every style value comes from (run, shape, layout, master, theme), fills, outlines, crops and animated shape ids")
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
