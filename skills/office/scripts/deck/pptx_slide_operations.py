from __future__ import annotations

import copy

from lxml import etree
from pptx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.parts.slide import SlidePart

from office_operations import OPERATION_NOT_APPLICABLE, TARGET_NOT_FOUND, Change
from office_result import INVALID_VALUE, OfficeFailure
from office_schema import closest_suggestion
from pptx_element_operations import rgb
from pptx_relationships import NOT_COPIED_RELATIONSHIP_TYPES, carry_relationships
from pptx_sections import remember_section, remove_from_custom_shows
from pptx_targets import PptxEditing, resolve_slide
from pptx_text_operations import replace_text


TITLE_PLACEHOLDER_TYPES = {"title", "ctrTitle"}
BODY_PLACEHOLDER_TYPES = {"body", "obj", "subTitle"}


def all_layouts(presentation) -> list:
    return [layout for master in presentation.slide_masters for layout in master.slide_layouts]


def resolve_layout(editing: PptxEditing, name: str, location: str):
    layouts = all_layouts(editing.presentation)
    match = next((layout for layout in layouts if layout.name == name), None)
    if match is not None:
        return match
    names = [layout.name for layout in layouts]
    fallback = "use one of " + ", ".join(repr(layout_name) for layout_name in names)
    raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: the deck has no layout {name!r}", location, closest_suggestion(name, names, "did you mean {match!r}?") or fallback))


def resolve_anchor(editing: PptxEditing, after: int | None, location: str):
    if after is None or after == 0:
        return after
    resolve_slide(editing, after, location)
    return after


def insert_slide_identifier(editing: PptxEditing, slide_part, after: int | None):
    presentation_part = editing.presentation.part
    relationship_id = presentation_part.relate_to(slide_part, RELATIONSHIP_TYPE.SLIDE)
    slide_list = editing.presentation.slides._sldIdLst
    slide_identifier = slide_list._add_sldId(id=next_slide_identifier(slide_list), rId=relationship_id)
    if after == 0:
        slide_list.insert(0, slide_identifier)
        remember_section(editing, slide_identifier.get("id"), None)
    elif after is not None:
        anchor = editing.insertion_points.get(after, editing.slide_id_elements[after - 1])
        anchor.addnext(slide_identifier)
        editing.insertion_points[after] = slide_identifier
        remember_section(editing, slide_identifier.get("id"), editing.slide_id_elements[after - 1].get("id"))
    else:
        remember_section(editing, slide_identifier.get("id"), slide_identifier.getprevious().get("id") if slide_identifier.getprevious() is not None else None)
    editing.structure_changed = True
    return slide_identifier


def next_slide_identifier(slide_list) -> int:
    return max((int(element.get("id")) for element in slide_list), default=255) + 1


def new_slide_part(editing: PptxEditing, layout, element=None):
    package = editing.presentation.part.package
    partname = package.next_partname("/ppt/slides/slide%d.xml")
    if element is None:
        return SlidePart.new(partname, package, layout.part)
    slide_part = SlidePart(partname, CONTENT_TYPE.PML_SLIDE, package, element)
    slide_part.relate_to(layout.part, RELATIONSHIP_TYPE.SLIDE_LAYOUT)
    return slide_part


def plan_add_slide(editing: PptxEditing, operation: dict, location: str) -> Change:
    layout = resolve_layout(editing, operation["layout"], f"{location}.layout")
    after = resolve_anchor(editing, operation.get("after"), f"{location}.after")

    def change() -> str:
        slide_part = new_slide_part(editing, layout)
        slide = slide_part.slide
        slide.shapes.clone_layout_placeholders(layout)
        insert_slide_identifier(editing, slide_part, after)
        fill_placeholders(slide, operation)
        editing.mark_edited(slide)
        return f"added a {layout.name!r} slide {placement(after)}"
    return change


def placement(after: int | None) -> str:
    if after is None:
        return "at the end"
    return "at the start" if after == 0 else f"after slide {after}"


def fill_placeholders(slide, operation: dict) -> None:
    placeholders = list(slide.placeholders)
    title = next((shape for shape in placeholders if placeholder_kind(shape) in TITLE_PLACEHOLDER_TYPES), None)
    body = next((shape for shape in placeholders if placeholder_kind(shape) in BODY_PLACEHOLDER_TYPES), None)
    for shape, text in ((title, operation.get("title")), (body, operation.get("body"))):
        if shape is not None and text:
            replace_text(shape.text_frame, None, text)


def placeholder_kind(shape) -> str:
    return shape._element.ph.get("type", "obj")


def plan_duplicate_slide(editing: PptxEditing, operation: dict, location: str) -> Change:
    source = resolve_slide(editing, operation["slide"], f"{location}.slide")
    after = resolve_anchor(editing, operation.get("after") if operation.get("after") is not None else operation["slide"], f"{location}.after")

    def change() -> str:
        element = parse_xml(etree.tostring(source._element))
        slide_part = new_slide_part(editing, source.slide_layout, element)
        insert_slide_identifier(editing, slide_part, after)
        carry_relationships(source.part, slide_part, element, NOT_COPIED_RELATIONSHIP_TYPES | {RELATIONSHIP_TYPE.SLIDE_LAYOUT})
        copy_notes(source, slide_part.slide)
        editing.mark_edited(slide_part.slide)
        return f"copied slide {operation['slide']} {placement(after)}"
    return change


def copy_notes(source, copy_slide) -> None:
    if not source.has_notes_slide:
        return
    source_placeholder = source.notes_slide.notes_placeholder
    target_placeholder = copy_slide.notes_slide.notes_placeholder
    if source_placeholder is None or target_placeholder is None:
        return
    target_body = target_placeholder._element.find(qn("p:txBody"))
    target_body.getparent().replace(target_body, copy.deepcopy(source_placeholder._element.find(qn("p:txBody"))))


def plan_delete_slide(editing: PptxEditing, operation: dict, location: str) -> Change:
    number = operation["slide"]
    if number in editing.touched_slides:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.slide: slide {number} is used by another operation in this batch", f"{location}.slide"))
    resolve_slide(editing, number, f"{location}.slide")
    editing.deleted_slides.add(number)
    slide_identifier = editing.slide_id_elements[number - 1]

    def change() -> str:
        relationship_id = slide_identifier.get(qn("r:id"))
        remove_from_custom_shows(editing.presentation.part._element, relationship_id)
        editing.presentation.part.drop_rel(relationship_id)
        slide_identifier.getparent().remove(slide_identifier)
        editing.structure_changed = True
        return f"deleted slide {number}"
    return change


def plan_reorder(editing: PptxEditing, operation: dict, location: str) -> Change:
    order = operation["order"]
    out_of_range = [number for number in order if number > len(editing.slides)]
    if out_of_range:
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.order: slide {out_of_range[0]} does not exist", f"{location}.order", f"use slide numbers from 1 to {len(editing.slides)}"))
    if len(set(order)) != len(order):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.order: a slide is listed twice", f"{location}.order"))

    def change() -> str:
        remaining = [number for number in range(1, len(editing.slides) + 1) if number not in editing.deleted_slides]
        if sorted(order) != remaining:
            raise OfficeFailure(INVALID_VALUE.issue(f"{location}.order: it must list exactly the slides that remain, {remaining}", f"{location}.order"))
        slide_list = editing.presentation.slides._sldIdLst
        children = list(slide_list)
        slots = sorted(children.index(editing.slide_id_elements[number - 1]) for number in remaining)
        for slot, number in zip(slots, order):
            children[slot] = editing.slide_id_elements[number - 1]
        for child in children:
            slide_list.append(child)
        editing.structure_changed = True
        return "reordered the slides to " + ", ".join(str(number) for number in order)
    return change


def plan_set_slide_hidden(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")

    def change() -> str:
        if operation["hidden"]:
            slide._element.set("show", "0")
        elif "show" in slide._element.attrib:
            del slide._element.attrib["show"]
        return f"{'hid' if operation['hidden'] else 'showed'} slide {operation['slide']}"
    return change


def plan_set_background(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")

    def change() -> str:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(operation["color"])
        editing.mark_edited(slide)
        return f"set the background of slide {operation['slide']} to {operation['color']}"
    return change


def plan_set_layout(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")
    layout = resolve_layout(editing, operation["layout"], f"{location}.layout")

    def change() -> str:
        relationship_id = next(identifier for identifier, relationship in slide.part.rels.items() if relationship.reltype == RELATIONSHIP_TYPE.SLIDE_LAYOUT)
        slide.part.rels.pop(relationship_id)
        slide.part.relate_to(layout.part, RELATIONSHIP_TYPE.SLIDE_LAYOUT)
        editing.mark_edited(slide)
        return f"switched slide {operation['slide']} to layout {layout.name!r}"
    return change


SLIDE_PLANNERS = {
    "add_slide": plan_add_slide,
    "duplicate_slide": plan_duplicate_slide,
    "delete_slide": plan_delete_slide,
    "reorder": plan_reorder,
    "set_slide_hidden": plan_set_slide_hidden,
    "set_background": plan_set_background,
    "set_layout": plan_set_layout,
}
