from __future__ import annotations

from lxml import etree
from pptx.oxml.ns import qn

from core.office_operations import OPERATION_NOT_APPLICABLE, Change
from core.office_result import INVALID_VALUE, OfficeFailure
from deck.pptx_element_operations import SHAPE_TAGS, current_box, next_shape_identifier, write_local_box
from deck.pptx_geometry import Box, child_frame, own_box, transform_of
from deck.pptx_shape_kinds import placeholder_of
from deck.pptx_targets import PptxEditing, ShapeTarget, require_kind, resolve_shape


SLIDE_REFERENCE = "slide"
MINIMUM_SELECTION = {"align_shapes": 2, "distribute_shapes": 3}
CROP_SIDES = ("left", "top", "right", "bottom")


def resolve_shapes(editing: PptxEditing, operation: dict, location: str) -> list[ShapeTarget]:
    addresses = operation["shapes"]
    if len(set(map(str, addresses))) != len(addresses):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.shapes: a shape is listed twice", f"{location}.shapes"))
    return [resolve_shape(editing, {"slide": operation["slide"], "shape": address}, f"{location}.shapes[{index}]") for index, address in enumerate(addresses)]


def require_selection(operation: dict, targets: list, location: str) -> None:
    needed = 1 if operation.get("to") == SLIDE_REFERENCE else MINIMUM_SELECTION[operation["op"]]
    if len(targets) < needed:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.shapes: {operation['op']} needs at least {needed} shapes here", f"{location}.shapes", 'list more shapes, or add "to": "slide"'))


def reference_box(editing: PptxEditing, operation: dict, boxes: list[Box]) -> Box:
    if operation.get("to") == SLIDE_REFERENCE:
        return Box(0, 0, editing.presentation.slide_width, editing.presentation.slide_height)
    left, top = min(box.x for box in boxes), min(box.y for box in boxes)
    return Box(left, top, max(box.right for box in boxes) - left, max(box.bottom for box in boxes) - top)


def aligned_box(box: Box, reference: Box, edge: str) -> Box:
    positions = {
        "left": (reference.x, box.y),
        "center": (reference.x + (reference.w - box.w) // 2, box.y),
        "right": (reference.right - box.w, box.y),
        "top": (box.x, reference.y),
        "middle": (box.x, reference.y + (reference.h - box.h) // 2),
        "bottom": (box.x, reference.bottom - box.h),
    }
    x, y = positions[edge]
    return Box(x, y, box.w, box.h)


def place(target: ShapeTarget, box: Box) -> None:
    write_local_box(target.element, target.frame.to_local(box))


def plan_align_shapes(editing: PptxEditing, operation: dict, location: str) -> Change:
    targets = resolve_shapes(editing, operation, location)
    require_selection(operation, targets, location)

    def change() -> str:
        boxes = [current_box(editing, target) for target in targets]
        reference = reference_box(editing, operation, boxes)
        for target, box in zip(targets, boxes):
            place(target, aligned_box(box, reference, operation["edge"]))
        editing.mark_edited(targets[0].slide)
        return f"aligned {len(targets)} shapes of slide {operation['slide']} on their {operation['edge']} to the {operation.get('to', 'selection')}"
    return change


def plan_distribute_shapes(editing: PptxEditing, operation: dict, location: str) -> Change:
    targets = resolve_shapes(editing, operation, location)
    require_selection(operation, targets, location)
    horizontal = operation["axis"] == "horizontal"

    def change() -> str:
        placed = sorted(((current_box(editing, target), target) for target in targets), key=lambda pair: pair[0].x if horizontal else pair[0].y)
        starts = distributed_starts([box for box, _ in placed], horizontal, editing, operation.get("to") == SLIDE_REFERENCE)
        for (box, target), start in zip(placed, starts):
            place(target, Box(start, box.y, box.w, box.h) if horizontal else Box(box.x, start, box.w, box.h))
        editing.mark_edited(targets[0].slide)
        return f"spaced {len(targets)} shapes of slide {operation['slide']} evenly {operation['axis']}ly"
    return change


def distributed_starts(boxes: list[Box], horizontal: bool, editing: PptxEditing, to_slide: bool) -> list[int]:
    lengths = [box.w if horizontal else box.h for box in boxes]
    if to_slide:
        first, last_end = 0, editing.presentation.slide_width if horizontal else editing.presentation.slide_height
        gap = (last_end - sum(lengths)) / (len(boxes) + 1)
        position = first + gap
    else:
        first = boxes[0].x if horizontal else boxes[0].y
        last_end = boxes[-1].right if horizontal else boxes[-1].bottom
        gap = (last_end - first - sum(lengths)) / (len(boxes) - 1)
        position = first
    starts = []
    for length in lengths:
        starts.append(round(position))
        position += length + gap
    return starts


def plan_group_shapes(editing: PptxEditing, operation: dict, location: str) -> Change:
    targets = resolve_shapes(editing, operation, location)
    if len(targets) < 2:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.shapes: a group needs at least two shapes", f"{location}.shapes"))
    nested = [target for target in targets if "." in target.address]
    if nested:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.shapes: {nested[0].label} is inside a group", f"{location}.shapes", "group shapes that sit directly on the slide, or ungroup first"))
    placeholders = [target for target in targets if placeholder_of(target.element) is not None]
    if placeholders:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.shapes: {placeholders[0].label} is a layout placeholder, which PowerPoint does not group", f"{location}.shapes", "leave placeholders out of the group"))

    def change() -> str:
        slide_element = targets[0].slide._element
        members = sorted(targets, key=lambda target: z_position(target.element))
        boxes = [current_box(editing, target) for target in members]
        bounds = reference_box(editing, {}, boxes)
        group = group_element(next_shape_identifier(slide_element), bounds)
        members[-1].element.addnext(group)
        for member in members:
            group.append(member.element)
        editing.mark_edited(targets[0].slide)
        return f"grouped {len(members)} shapes of slide {operation['slide']} into one group"
    return change


def z_position(element) -> int:
    return list(element.getparent()).index(element)


def group_element(identifier: int, bounds: Box):
    group = etree.Element(qn("p:grpSp"))
    properties = etree.SubElement(group, qn("p:nvGrpSpPr"))
    etree.SubElement(properties, qn("p:cNvPr"), id=str(identifier), name=f"Group {identifier}")
    etree.SubElement(properties, qn("p:cNvGrpSpPr"))
    etree.SubElement(properties, qn("p:nvPr"))
    transform = etree.SubElement(etree.SubElement(group, qn("p:grpSpPr")), qn("a:xfrm"))
    for tag, horizontal, vertical in (("a:off", "x", "y"), ("a:ext", "cx", "cy"), ("a:chOff", "x", "y"), ("a:chExt", "cx", "cy")):
        values = (bounds.x, bounds.y) if horizontal == "x" else (bounds.w, bounds.h)
        etree.SubElement(transform, qn(tag), {horizontal: str(values[0]), vertical: str(values[1])})
    return group


def plan_ungroup_shape(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_shape(editing, operation, location)
    require_kind(target, ("group",), location, "ungroup_shape dissolves a group")
    if transform_of(target.element) is not None and transform_of(target.element).get("rot", "0") != "0":
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.shape: {target.label} is rotated, and its shapes would lose that rotation", f"{location}.shape", "set its rotation to 0 with set_transform first"))

    def change() -> str:
        frame = child_frame(target.frame, target.element)
        children = [child for child in target.element if child.tag in SHAPE_TAGS]
        for child in children:
            box = own_box(child)
            target.element.addprevious(child)
            if box is not None:
                write_local_box(child, target.frame.to_local(frame.to_slide(box)))
        target.element.getparent().remove(target.element)
        editing.mark_edited(target.slide)
        return f"ungrouped {target.label} into {len(children)} shapes"
    return change


def plan_crop_picture(editing: PptxEditing, operation: dict, location: str) -> Change:
    sides = {side: operation[side] for side in CROP_SIDES if operation.get(side) is not None}
    if not sides:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: crop_picture changes nothing", location, "give at least one of left, top, right, bottom"))
    target = resolve_shape(editing, operation, location)
    require_kind(target, ("picture",), location, "crop_picture crops a picture")
    picture = target.shape
    crop = {side: sides.get(side, getattr(picture, f"crop_{side}")) for side in CROP_SIDES}
    for first, second in (("left", "right"), ("top", "bottom")):
        if crop[first] + crop[second] >= 1:
            raise OfficeFailure(INVALID_VALUE.issue(f"{location}: {first} and {second} together cut {crop[first] + crop[second]:.2f} of the image, leaving nothing", location, "keep each pair below 1, such as 0.1 and 0.1"))

    def change() -> str:
        for side, share in crop.items():
            setattr(picture, f"crop_{side}", share)
        editing.mark_edited(target.slide)
        return f"cropped {target.label} to " + ", ".join(f"{side} {share:g}" for side, share in crop.items())
    return change


ARRANGEMENT_PLANNERS = {
    "align_shapes": plan_align_shapes,
    "distribute_shapes": plan_distribute_shapes,
    "group_shapes": plan_group_shapes,
    "ungroup_shape": plan_ungroup_shape,
    "crop_picture": plan_crop_picture,
}
