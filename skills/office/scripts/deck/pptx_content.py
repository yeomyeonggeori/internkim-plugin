from __future__ import annotations

from pptx.oxml.ns import qn

from deck.pptx_shape_kinds import shape_kind


def text_frames_in(shapes) -> list:
    frames = []
    for shape in shapes:
        kind = shape_kind(shape._element)
        if kind == "group":
            frames.extend(text_frames_in(shape.shapes))
        elif kind == "table":
            frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
        elif shape._element.find(qn("p:txBody")) is not None:
            frames.append(shape.text_frame)
    return frames


def frame_text(text_frame) -> str:
    return text_frame.text.replace("\v", "\n")


def notes_text(slide) -> str:
    return frame_text(slide.notes_slide.notes_text_frame) if slide.has_notes_slide else ""
