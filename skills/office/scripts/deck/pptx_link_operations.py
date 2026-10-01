from __future__ import annotations

import copy

from pptx.opc.constants import RELATIONSHIP_TYPE
from pptx.oxml.ns import qn

from office_operations import TARGET_NOT_FOUND, Change
from office_result import INVALID_VALUE, OfficeFailure
from pptx_relationships import drop_unreferenced
from pptx_targets import PptxEditing, ShapeTarget, require_text, resolve_shape, resolve_slide


SLIDE_JUMP_ACTION = "ppaction://hlinksldjump"


def plan_set_link(editing: PptxEditing, operation: dict, location: str) -> Change:
    if (operation.get("url") is None) == (operation.get("toSlide") is None):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: set_link takes url or toSlide, one of them", location, 'give "url": "https://…", or "toSlide": 3; "url": "" removes the link'))
    target = resolve_shape(editing, operation, location)
    destination = resolve_slide(editing, operation["toSlide"], f"{location}.toSlide") if operation.get("toSlide") is not None else None
    runs = linked_runs(target, operation["text"], location) if operation.get("text") else None

    def change() -> str:
        if runs is None:
            link_shape(target, operation.get("url"), destination)
        else:
            for run in split_runs_at(*runs):
                link_run(target, run, operation.get("url"), destination)
        editing.mark_edited(target.slide)
        return f"{link_summary(operation)} on {target.label}" + (f" for {operation['text']!r}" if runs is not None else "")
    return change


def link_summary(operation: dict) -> str:
    if operation.get("toSlide") is not None:
        return f"linked to slide {operation['toSlide']}"
    return f"linked to {operation['url']}" if operation["url"] else "removed the link"


def link_shape(target: ShapeTarget, url: str | None, destination) -> None:
    action = target.shape.click_action
    if destination is not None:
        action.target_slide = destination
        return
    action.hyperlink.address = url or None


def linked_runs(target: ShapeTarget, text: str, location: str):
    text_frame = require_text(target, location)
    for paragraph in text_frame.paragraphs:
        runs = paragraph._p.findall(qn("a:r"))
        joined = "".join(run_text(run) for run in runs)
        start = joined.find(text)
        if start >= 0:
            return runs, start, start + len(text)
    raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.text: {target.label} has no paragraph holding {text!r}", f"{location}.text", "give text exactly as deck read shows it, inside one paragraph"))


def run_text(run) -> str:
    return "".join(node.text or "" for node in run.iter(qn("a:t")))


def split_runs_at(runs: list, start: int, end: int) -> list:
    covered = []
    position = 0
    for run in runs:
        length = len(run_text(run))
        run_start, run_end = position, position + length
        position = run_end
        if run_end <= start or run_start >= end:
            continue
        piece = run
        if run_start < start:
            piece = split_run(piece, start - run_start)
        if run_end > end:
            split_run(piece, len(run_text(piece)) - (run_end - end))
        covered.append(piece)
    return covered


def split_run(run, offset: int):
    text = run_text(run)
    tail = copy.deepcopy(run)
    set_run_text(run, text[:offset])
    set_run_text(tail, text[offset:])
    run.addnext(tail)
    return tail


def set_run_text(run, text: str) -> None:
    run.find(qn("a:t")).text = text


def link_run(target: ShapeTarget, run, url: str | None, destination) -> None:
    properties = run.find(qn("a:rPr"))
    if properties is None:
        properties = run.makeelement(qn("a:rPr"), {})
        run.insert(0, properties)
    previous = properties.find(qn("a:hlinkClick"))
    if previous is not None:
        properties.remove(previous)
        drop_unreferenced(target.slide.part, previous.get(qn("r:id")))
    if destination is None and not url:
        return
    part = target.slide.part
    if destination is not None:
        relationship_id = part.relate_to(destination.part, RELATIONSHIP_TYPE.SLIDE)
        attributes = {qn("r:id"): relationship_id, "action": SLIDE_JUMP_ACTION}
    else:
        relationship_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        attributes = {qn("r:id"): relationship_id}
    link = properties.makeelement(qn("a:hlinkClick"), attributes)
    insert_link(properties, link)


def insert_link(properties, link) -> None:
    following = next((child for child in properties if child.tag in (qn("a:hlinkMouseOver"), qn("a:rtl"), qn("a:extLst"))), None)
    if following is None:
        properties.append(link)
    else:
        following.addprevious(link)


LINK_PLANNERS = {
    "set_link": plan_set_link,
}
