from __future__ import annotations

from pptx.chart.axis import ValueAxis
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn

from charts.combo import combo_chart_space, lines_need_own_axis
from charts.kinds import COMBO_CHART_KIND, ROUND_CHART_KINDS, SCATTER_CHART_KIND, office_chart_type
from charts.look import LABEL_FLAG_NAMES, LABEL_FLAGS, labels_need_round_chart
from core.office_operations import OPERATION_NOT_APPLICABLE, Change
from core.office_result import INVALID_VALUE, OfficeFailure
from deck.pptx_element_operations import rgb
from deck.pptx_insert_operations import box_arguments, cell_text, chart_data, set_chart_title
from deck.pptx_targets import PptxEditing, resolve_slide


def plan_add_chart(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")
    require_chart_options(operation, location)
    data = new_chart_data(operation, location)

    def change() -> str:
        chart = slide.shapes.add_chart(drawn_chart_type(operation["type"]), *box_arguments(operation), data).chart
        if operation["type"] == COMBO_CHART_KIND:
            draw_lines_over_columns(chart, operation)
        show_legend(chart, operation)
        set_chart_title(chart, operation.get("title"))
        paint_series(chart, operation.get("colors") or [], operation["type"] in ROUND_CHART_KINDS)
        label_points(chart, operation.get("dataLabels"))
        title_axes(chart, operation)
        editing.mark_edited(slide)
        return f"added a {operation['type']} chart {len(slide.shapes) - 1} to slide {operation['slide']}"
    return change


def require_chart_options(operation: dict, location: str) -> None:
    kind = operation["type"]
    is_round = kind in ROUND_CHART_KINDS
    if is_round and len(operation["series"]) > 1:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.series: a {kind} chart shows one series", f"{location}.series", "give one series, or pick column or bar"))
    if is_round and (operation.get("xTitle") or operation.get("yTitle")):
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: a {kind} chart has no axes to title", location, "leave out xTitle and yTitle, or name the slices in the title"))
    if operation.get("dataLabels") and labels_need_round_chart(operation["dataLabels"]) and not is_round:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.dataLabels: {operation['dataLabels']} labels show each slice's share of a pie or doughnut, and PowerPoint draws none on other charts", f"{location}.dataLabels", 'on this chart use "dataLabels": "value"'))
    lines = [entry.get("line") for entry in operation["series"]]
    if any(lines) and kind != COMBO_CHART_KIND:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.series: line marks a series of a combo chart", f"{location}.series", 'pick "type": "combo", or leave out line'))
    if kind == COMBO_CHART_KIND and (all(lines) or not any(lines)):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.series: a combo chart draws columns and lines, so some series and not all take \"line\": true", f"{location}.series", 'mark the series drawn as a line with "line": true'))


def new_chart_data(operation: dict, location: str):
    if operation["type"] != SCATTER_CHART_KIND:
        return chart_data(operation["categories"], operation["series"], location)
    return scatter_data(operation, location)


def scatter_data(operation: dict, location: str) -> XyChartData:
    horizontal = operation["categories"]
    if not all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in horizontal):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.categories: a scatter chart places points by number, and the categories are its x values", f"{location}.categories", "give one number per point, or pick a column or line chart for named categories"))
    data = XyChartData()
    for index, entry in enumerate(operation["series"]):
        if len(entry["values"]) != len(horizontal):
            raise OfficeFailure(INVALID_VALUE.issue(f"{location}.series[{index}].values: {len(entry['values'])} values for {len(horizontal)} x values", f"{location}.series[{index}].values", "give one y value per x value"))
        series = data.add_series(entry["name"])
        for x_value, y_value in zip(horizontal, entry["values"]):
            series.add_data_point(x_value, y_value)
    return data


def drawn_chart_type(kind: str):
    if kind == SCATTER_CHART_KIND:
        return XL_CHART_TYPE.XY_SCATTER
    if kind == COMBO_CHART_KIND:
        return office_chart_type("column")
    return office_chart_type(kind)


def draw_lines_over_columns(chart, operation: dict) -> None:
    series = [(entry["name"], tuple(float(value) for value in entry["values"]), bool(entry.get("line"))) for entry in operation["series"]]
    secondary_axis = operation.get("secondaryAxis")
    combined = combo_chart_space(tuple(cell_text(category) for category in operation["categories"]), series, lines_need_own_axis(series) if secondary_axis is None else secondary_axis)
    plot_area = chart._chartSpace.find(f"{qn('c:chart')}/{qn('c:plotArea')}")
    plot_area.getparent().replace(plot_area, combined.find(f"{qn('c:chart')}/{qn('c:plotArea')}"))


def show_legend(chart, operation: dict) -> None:
    legend = operation.get("legend")
    chart.has_legend = legend if legend is not None else len(operation["series"]) > 1 or operation["type"] in ROUND_CHART_KINDS
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False


def paint_series(chart, colors: list[str], is_round: bool) -> None:
    every_series = [series for plot in chart.plots for series in plot.series]
    if is_round:
        for point_index, color in enumerate(colors):
            paint_fill(every_series[0].points[point_index].format.fill, color)
        return
    for series, color in zip(every_series, colors):
        if hasattr(series, "marker"):
            paint_marked_series(series, color)
        else:
            paint_fill(series.format.fill, color)


def paint_marked_series(series, color: str) -> None:
    if not has_hidden_line(series):
        series.format.line.color.rgb = rgb(color)
    paint_fill(series.marker.format.fill, color)
    series.marker.format.line.color.rgb = rgb(color)


def has_hidden_line(series) -> bool:
    return series._element.find(f"{qn('c:spPr')}/{qn('a:ln')}/{qn('a:noFill')}") is not None


def paint_fill(fill, color: str) -> None:
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def label_points(chart, mode: str | None) -> None:
    if mode is None:
        return
    for plot in chart.plots:
        plot.has_data_labels = mode != "none"
        if mode == "none":
            continue
        labels = plot._element.get_or_add_dLbls()
        for flag in LABEL_FLAG_NAMES:
            getattr(labels, f"get_or_add_{flag}")().set("val", "1" if flag in LABEL_FLAGS[mode] else "0")


def title_axes(chart, operation: dict) -> None:
    if not operation.get("xTitle") and not operation.get("yTitle"):
        return
    value_axes = chart._chartSpace.valAx_lst
    is_scatter = operation["type"] == SCATTER_CHART_KIND
    horizontal_axis = ValueAxis(value_axes[0]) if is_scatter else chart.category_axis
    vertical_axis = ValueAxis(value_axes[1] if is_scatter else value_axes[0])
    for axis, title in ((horizontal_axis, operation.get("xTitle")), (vertical_axis, operation.get("yTitle"))):
        if title:
            axis.axis_title.text_frame.text = title


CHART_INSERT_PLANNERS = {
    "add_chart": plan_add_chart,
}
