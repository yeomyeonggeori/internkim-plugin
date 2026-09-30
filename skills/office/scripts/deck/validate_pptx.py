#!/usr/bin/env python3
from __future__ import annotations

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deck_definitions import (
    DEFAULT_FONT_REMAINS,
    OVERLAY_OUT_OF_BOUNDS,
    OVERLAY_WITHOUT_BACKGROUND,
    SLIDE_EMPTY,
    SLIDE_TITLE_MISSING,
    THEME_FONT_INHERITED,
    TOO_MANY_SHAPES,
)
from office_result import Issue, OfficeArgumentParser, Result, run_command


DEFAULT_FONT_NAMES = {"Aptos", "Calibri"}
EXCESSIVE_SHAPE_COUNT = 40
CONTENT_SHAPE_TYPES = {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.MEDIA}


def main() -> Result:
    arguments = parse_arguments()
    presentation = Presentation(arguments.presentation_path)
    slides = []
    issues = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_summary, found_issues = summarize_slide(slide, index, presentation.slide_width, presentation.slide_height)
        slides.append(slide_summary)
        issues.extend(found_issues)
    details = {"slideCount": len(presentation.slides), "slides": slides}
    return Result(summary=f"checked {arguments.presentation_path}: {len(issues)} issues", output_path=arguments.presentation_path, issues=tuple(issues), details=details)


def summarize_slide(slide, index, slide_width, slide_height):
    text_entries = text_entries_from_slide(slide)
    title = slide_title(slide, text_entries)
    shape_count = len(slide.shapes)
    content_shape_count = sum(1 for shape in slide.shapes if holds_content(shape))
    explicit_default_fonts = sorted(default_fonts_from_entries(text_entries))
    inherited_font_runs = sum(entry["inheritedFontRuns"] for entry in text_entries)
    hybrid_summary = hybrid_slide_summary(slide)
    out_of_bounds_overlays = editable_overlays_out_of_bounds(slide, slide_width, slide_height)
    issues = slide_issues(f"slide {index}", title, text_entries, content_shape_count, explicit_default_fonts, inherited_font_runs, hybrid_summary, out_of_bounds_overlays)
    return {
        "index": index,
        "title": title,
        "shapeCount": shape_count,
        "contentShapeCount": content_shape_count,
        "textShapeCount": len(text_entries),
        "explicitDefaultFonts": explicit_default_fonts,
        "inheritedFontRuns": inherited_font_runs,
        "hybridBackgroundPresent": hybrid_summary["hasHybridBackground"],
        "editableOverlayCount": hybrid_summary["editableOverlayCount"],
        "outOfBoundsEditableOverlays": out_of_bounds_overlays,
    }, issues


def holds_content(shape) -> bool:
    if shape.shape_type in CONTENT_SHAPE_TYPES:
        return True
    return bool(getattr(shape, "has_text_frame", False) and shape.text.strip())


def text_entries_from_slide(slide):
    entries = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text.strip()
        font_names = []
        inherited_font_runs = 0
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                font_name = run.font.name
                if font_name:
                    font_names.append(font_name)
                elif run.text.strip():
                    inherited_font_runs += 1
        if text:
            entries.append({
                "text": text,
                "top": int(getattr(shape, "top", 0) or 0),
                "fontNames": font_names,
                "inheritedFontRuns": inherited_font_runs,
            })
    return entries


def slide_title(slide, text_entries):
    if slide.shapes.title and slide.shapes.title.text.strip():
        return slide.shapes.title.text.strip()
    if not text_entries:
        return ""
    first_entry = sorted(text_entries, key=lambda entry: entry["top"])[0]
    return first_entry["text"].splitlines()[0].strip()


def default_fonts_from_entries(text_entries):
    fonts = set()
    for entry in text_entries:
        for font_name in entry["fontNames"]:
            if font_name in DEFAULT_FONT_NAMES:
                fonts.add(font_name)
    return fonts


def hybrid_slide_summary(slide):
    return {
        "hasHybridBackground": any(getattr(shape, "name", "") == "Hybrid Background" for shape in slide.shapes),
        "editableOverlayCount": sum(1 for shape in slide.shapes if getattr(shape, "name", "") == "Editable Overlay"),
    }


def editable_overlays_out_of_bounds(slide, slide_width, slide_height):
    indexes = []
    for index, shape in enumerate(slide.shapes, start=1):
        if getattr(shape, "name", "") != "Editable Overlay":
            continue
        if shape.left < 0 or shape.top < 0 or shape.left + shape.width > slide_width or shape.top + shape.height > slide_height:
            indexes.append(index)
    return indexes


def slide_issues(location, title, text_entries, content_shape_count, explicit_default_fonts, inherited_font_runs, hybrid_summary, out_of_bounds_overlays) -> list[Issue]:
    issues = []
    if not text_entries:
        issues.append(SLIDE_EMPTY.issue("slide appears empty", location))
    if not title:
        issues.append(SLIDE_TITLE_MISSING.issue("slide is missing a title", location))
    if content_shape_count > EXCESSIVE_SHAPE_COUNT:
        issues.append(TOO_MANY_SHAPES.issue(f"slide has {content_shape_count} shapes that hold content", location))
    if explicit_default_fonts:
        issues.append(DEFAULT_FONT_REMAINS.issue("default font remains: " + ", ".join(explicit_default_fonts), location))
    if inherited_font_runs:
        issues.append(THEME_FONT_INHERITED.issue(f"{inherited_font_runs} text runs inherit the theme font", location))
    if hybrid_summary["editableOverlayCount"] and not hybrid_summary["hasHybridBackground"]:
        issues.append(OVERLAY_WITHOUT_BACKGROUND.issue("editable overlays are present without a hybrid background image", location))
    if out_of_bounds_overlays:
        issues.append(OVERLAY_OUT_OF_BOUNDS.issue("editable overlay out of bounds: " + ", ".join(str(index) for index in out_of_bounds_overlays), location))
    return issues


def parse_arguments():
    parser = OfficeArgumentParser(description="Validate and summarize a PPTX deck.")
    parser.add_argument("presentation_path")
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
