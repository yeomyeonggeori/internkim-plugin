from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from office_operations import Change
from office_result import INVALID_VALUE, OfficeFailure
from pptx_backdrop import readable_text_color
from pptx_connector_operations import plan_add_connector
from pptx_element_operations import readable_image_size, rgb
from pptx_geometry import Box
from pptx_targets import PptxEditing, resolve_slide
from pptx_text_operations import ALIGNMENTS, apply_run_style, character_properties, replace_text


SHAPE_GEOMETRIES = {
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "right_arrow": MSO_SHAPE.RIGHT_ARROW,
    "chevron": MSO_SHAPE.CHEVRON,
    "pentagon": MSO_SHAPE.PENTAGON,
    "diamond": MSO_SHAPE.DIAMOND,
}
CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
}
SINGLE_SERIES_CHARTS = {"pie", "doughnut"}
TABLE_ROW_HEIGHT_EMU = 365760


def box_arguments(operation: dict) -> tuple:
    return tuple(Emu(operation[name]) for name in ("x", "y", "w", "h"))


def style_new_text(text_frame, operation: dict) -> None:
    for paragraph in text_frame.paragraphs:
        for properties in character_properties(paragraph._p):
            apply_run_style(properties, operation)
        if operation.get("align") is not None:
            paragraph.alignment = ALIGNMENTS[operation["align"]]


def plan_add_text_box(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")

    def change() -> str:
        color = operation.get("color") or readable_text_color(editing.presentation, slide, Box(*(int(argument) for argument in box_arguments(operation))))
        shape = slide.shapes.add_textbox(*box_arguments(operation))
        shape.text_frame.word_wrap = True
        replace_text(shape.text_frame, None, operation["text"])
        style_new_text(shape.text_frame, {**operation, "color": color})
        editing.mark_edited(slide)
        return f"added text box {len(slide.shapes) - 1} to slide {operation['slide']}"
    return change


def plan_add_shape(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")

    def change() -> str:
        shape = slide.shapes.add_shape(SHAPE_GEOMETRIES[operation["kind"]], *box_arguments(operation))
        paint_new_shape(shape, operation)
        if operation.get("text"):
            replace_text(shape.text_frame, None, operation["text"])
        editing.mark_edited(slide)
        return f"added {operation['kind']} {len(slide.shapes) - 1} to slide {operation['slide']}"
    return change


def paint_new_shape(shape, operation: dict) -> None:
    if operation.get("fill") == "none":
        shape.fill.background()
    elif operation.get("fill") is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(operation["fill"])
    if operation.get("line") == "none":
        shape.line.fill.background()
    elif operation.get("line") is not None:
        shape.line.color.rgb = rgb(operation["line"])


def plan_add_picture(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")
    readable_image_size(operation["image"], f"{location}.image")

    def change() -> str:
        size = [Emu(operation[name]) if operation.get(name) is not None else None for name in ("w", "h")]
        slide.shapes.add_picture(operation["image"], Emu(operation["x"]), Emu(operation["y"]), *size)
        editing.mark_edited(slide)
        return f"added picture {len(slide.shapes) - 1} to slide {operation['slide']}"
    return change


def plan_add_table(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")
    rows = operation["rows"]
    column_count = max(len(row) for row in rows)
    if column_count == 0:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.rows: every row is empty", f"{location}.rows", "give at least one cell value"))

    def change() -> str:
        height = operation.get("h") or TABLE_ROW_HEIGHT_EMU * len(rows)
        frame = slide.shapes.add_table(len(rows), column_count, Emu(operation["x"]), Emu(operation["y"]), Emu(operation["w"]), Emu(height))
        for row, values in zip(frame.table.rows, rows):
            for cell, value in zip(row.cells, values):
                cell.text = cell_text(value)
        editing.mark_edited(slide)
        return f"added a {len(rows)} by {column_count} table {len(slide.shapes) - 1} to slide {operation['slide']}"
    return change


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def chart_data(categories: list, series: list[dict], location: str, number_format: str = "General") -> CategoryChartData:
    data = CategoryChartData(number_format=number_format)
    data.categories = [cell_text(category) for category in categories]
    for index, entry in enumerate(series):
        if len(entry["values"]) != len(categories):
            raise OfficeFailure(INVALID_VALUE.issue(f"{location}.series[{index}].values: {len(entry['values'])} values for {len(categories)} categories", f"{location}.series[{index}].values", "give one value per category"))
        data.add_series(entry["name"], entry["values"])
    return data


def plan_add_chart(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")
    data = chart_data(operation["categories"], operation["series"], location)
    if operation["type"] in SINGLE_SERIES_CHARTS and len(operation["series"]) > 1:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.series: a {operation['type']} chart shows one series", f"{location}.series", "give one series, or pick column or bar"))

    def change() -> str:
        chart = slide.shapes.add_chart(CHART_TYPES[operation["type"]], *box_arguments(operation), data).chart
        legend = operation.get("legend")
        chart.has_legend = legend if legend is not None else len(operation["series"]) > 1 or operation["type"] in SINGLE_SERIES_CHARTS
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        set_chart_title(chart, operation.get("title"))
        editing.mark_edited(slide)
        return f"added a {operation['type']} chart {len(slide.shapes) - 1} to slide {operation['slide']}"
    return change


def set_chart_title(chart, title: str | None) -> None:
    if title is None:
        return
    chart.has_title = bool(title)
    if title:
        chart.chart_title.text_frame.text = title


INSERT_PLANNERS = {
    "add_text_box": plan_add_text_box,
    "add_connector": plan_add_connector,
    "add_shape": plan_add_shape,
    "add_picture": plan_add_picture,
    "add_table": plan_add_table,
    "add_chart": plan_add_chart,
}
