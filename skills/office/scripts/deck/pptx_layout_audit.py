from __future__ import annotations

import copy
from dataclasses import dataclass
import math

from pptx.oxml.ns import qn

from deck.pptx_edit_definitions import BACKGROUND_SHARE_OF_SLIDE, CONTENT_OVERFLOW, DISTORTION_TOLERANCE, IMAGE_DISTORTED, OUT_OF_FRAME, OVERLAP_RATIO, TEXT_OVERLAP, ReviewCheck
from core.office_result import Issue
from deck.pptx_geometry import SLIDE_FRAME, Box, Frame, child_frame, local_box
from deck.pptx_inheritance import slide_context
from deck.pptx_shape_kinds import shape_address, shape_kind, shape_reference
from deck.pptx_text_measure import TextFit, grown_box, grows_with_text, largest_text_size, measure_text, wraps
from deck.pptx_text_operations import apply_run_style, character_properties
from core.units import EMU_PER_POINT
from core.image_formats import OFFICE_PICTURE_CONTENT_TYPES


EDGE_TOLERANCE = EMU_PER_POINT
OVERFLOW_TOLERANCE = 2 * EMU_PER_POINT
SUGGESTION_SLACK = 4 * EMU_PER_POINT
SMALLEST_SUGGESTED_SIZE = 10
SHRINK_STEP = 0.05
SHRINK_STEPS = 10
MEDIA_KINDS = {"picture", "chart", "table", "diagram", "media", "object"}
CROP_SCALE = 100000
NO_SINGLE_OPERATION = "no single operation fits this text at a legible size: shorten it, or split it across slides"


@dataclass(frozen=True)
class Entry:
    address: str
    name: str
    kind: str
    box: Box
    element: object
    fit: TextFit | None
    pixels: tuple[int, int] | None
    context: object
    hole: Box | None = None

    @property
    def has_text(self) -> bool:
        return self.fit is not None

    @property
    def visible_box(self) -> Box:
        return grown_box(self.element, self.box, self.fit)


@dataclass(frozen=True)
class SlideArea:
    number: int
    width: int
    height: int


@dataclass(frozen=True)
class Audit:
    issues: list[Issue]
    faces: frozenset


def audit_presentation(presentation, numbered_slides: list[tuple[int, object]]) -> Audit:
    issues, faces = [], set()
    for number, slide in numbered_slides:
        entries = slide_entries(presentation, slide)
        area = SlideArea(number, presentation.slide_width, presentation.slide_height)
        issues.extend(slide_issues(entries, area))
        faces.update(pair for entry in entries if entry.fit is not None for pair in entry.fit.faces)
    return Audit(issues, frozenset(faces))


def slide_entries(presentation, slide) -> list[Entry]:
    context = slide_context(presentation, slide)
    return list(walk_entries(slide.shapes, "", SLIDE_FRAME, context))


def walk_entries(shapes, prefix: str, frame: Frame, context):
    for index, shape in enumerate(shapes):
        element, address = shape._element, shape_address(prefix, index)
        kind = shape_kind(element)
        if kind == "group":
            yield from walk_entries(shape.shapes, address, child_frame(frame, element), context)
            continue
        box = frame.to_slide(local_box(element, context))
        fit = measure_text(context, element, box) if kind in ("text", "shape") else None
        yield Entry(address, shape.name, kind, box, element, fit, picture_pixels(shape, kind), context, doughnut_hole(shape, kind, box))


def doughnut_hole(shape, kind: str, box: Box) -> Box | None:
    if kind != "chart":
        return None
    chart_space = shape.chart._chartSpace
    hole_size = chart_space.find(f".//{qn('c:doughnutChart')}/{qn('c:holeSize')}")
    if hole_size is None or chart_space.find(f".//{qn('c:legend')}") is not None:
        return None
    side = int(min(box.w, box.h) * int(hole_size.get("val")) / 100 / math.sqrt(2))
    return Box(box.x + (box.w - side) // 2, box.y + (box.h - side) // 2, side, side)


def slide_issues(entries: list[Entry], area: SlideArea) -> list[Issue]:
    issues = []
    for entry in entries:
        issues.extend(frame_issues(entry, area))
        issues.extend(overflow_issues(entry, area, entries))
        issues.extend(distortion_issues(entry, area))
    issues.extend(overlap_issues(entries, area))
    issues.extend(card_spill_issues(entries, area))
    return issues


def card_spill_issues(entries: list[Entry], area: SlideArea) -> list[Issue]:
    issues = []
    for entry in entries:
        if entry.visible_box == entry.box:
            continue
        card = next((candidate for candidate in entries if candidate.kind == "shape" and not candidate.has_text and contains(candidate.box, entry.box)), None)
        if card is None or contains(card.box, entry.visible_box):
            continue
        grown = Box(card.box.x, card.box.y, max(card.box.w, entry.visible_box.right + (card.box.right - entry.box.right) - card.box.x), max(card.box.h, entry.visible_box.bottom + (card.box.bottom - entry.box.bottom) - card.box.y))
        text = f"{label(entry, area)} grows with its text past shape {card.address} {card.name!r} that holds it"
        issues.append(layout_issue(CONTENT_OVERFLOW, text, location(entry, area), growth_fix(entry, area) or transform_suggestion(card, area, inside(grown, area))))
    return issues


def layout_issue(check: ReviewCheck, text: str, where: str, remedy: dict | str | None) -> Issue:
    if isinstance(remedy, dict):
        return check.kind.issue(text, where, f"apply the {remedy['op']} in fix with office apply, then check the slide again", fix=[remedy])
    return check.kind.issue(text, where, remedy)


def contains(outer: Box, inner: Box) -> bool:
    return outer.x - EDGE_TOLERANCE <= inner.x and outer.y - EDGE_TOLERANCE <= inner.y and inner.right <= outer.right + EDGE_TOLERANCE and inner.bottom <= outer.bottom + EDGE_TOLERANCE


def label(entry: Entry, area: SlideArea) -> str:
    return f"slide {area.number} shape {entry.address} {entry.name!r}"


def location(entry: Entry, area: SlideArea) -> str:
    return f"slide {area.number} shape {entry.address}"


def transform_suggestion(entry: Entry, area: SlideArea, box: Box) -> dict:
    changed = {name: value for name, value in box.to_json().items() if value != getattr(entry.box, name)}
    return {"op": "set_transform", "slide": area.number, "shape": shape_reference(entry.address), **changed}


def inside(box: Box, area: SlideArea) -> Box:
    width, height = min(box.w, area.width), min(box.h, area.height)
    return Box(min(max(box.x, 0), area.width - width), min(max(box.y, 0), area.height - height), width, height)


def frame_issues(entry: Entry, area: SlideArea) -> list[Issue]:
    box = entry.visible_box
    overhangs = {
        "left": -box.x,
        "top": -box.y,
        "right": box.right - area.width,
        "bottom": box.bottom - area.height,
    }
    past = [f"{round(amount / EMU_PER_POINT)}pt past the {edge} edge" for edge, amount in overhangs.items() if amount > EDGE_TOLERANCE]
    if not past:
        return []
    whole = box.right < 0 or box.bottom < 0 or box.x > area.width or box.y > area.height
    text = f"{label(entry, area)} lies {'wholly outside the slide' if whole else 'partly outside the slide'}: {', '.join(past)}"
    fix = growth_fix(entry, area) or changed_transform(entry, area, inside(entry.box, area))
    return [layout_issue(OUT_OF_FRAME, text, location(entry, area), fix)]


def growth_fix(entry: Entry, area: SlideArea) -> dict | None:
    if entry.visible_box == entry.box:
        return None
    if not wraps(entry.element) and entry.fit.width_overflow > 0:
        return {"op": "set_text_frame", "slide": area.number, "shape": shape_reference(entry.address), "wrap": True}
    if entry.visible_box.bottom > area.height and entry.visible_box.h <= area.height:
        return {"op": "set_transform", "slide": area.number, "shape": shape_reference(entry.address), "y": area.height - entry.visible_box.h}
    if entry.visible_box.bottom > area.height:
        return smaller_text(entry, area, lambda fit: entry.box.y + fit.needed_height + entry.box.h - fit.available_height <= area.height)
    return None


def changed_transform(entry: Entry, area: SlideArea, box: Box) -> dict | None:
    return transform_suggestion(entry, area, box) if box != entry.box else None


def overflow_issues(entry: Entry, area: SlideArea, entries: list[Entry]) -> list[Issue]:
    fit = entry.fit
    if fit is None or grows_with_text(entry.element):
        return []
    issues = []
    if fit.height_overflow > OVERFLOW_TOLERANCE:
        text = f"{label(entry, area)}: its text needs {points(fit.needed_height)}pt of height and the box gives {points(fit.available_height)}pt{measured_with(entry)}"
        issues.append(layout_issue(CONTENT_OVERFLOW, text, location(entry, area), taller_box_or_smaller_text(entry, area, entries)))
    if fit.width_overflow > OVERFLOW_TOLERANCE:
        text = f"{label(entry, area)}: a line is {points(fit.widest_line)}pt wide and the box gives {points(fit.available_width)}pt{measured_with(entry)}"
        issues.append(layout_issue(CONTENT_OVERFLOW, text, location(entry, area), wider_box_or_smaller_text(entry, area, entries)))
    return issues


def points(emu: int) -> int:
    return round(emu / EMU_PER_POINT)


def taller_box_or_smaller_text(entry: Entry, area: SlideArea, entries: list[Entry]) -> dict | str:
    needed = entry.box.h + entry.fit.height_overflow
    grown = inside(Box(entry.box.x, entry.box.y, entry.box.w, needed + SUGGESTION_SLACK), area)
    if grown.h >= needed and not collides(grown, entry, entries):
        return transform_suggestion(entry, area, grown)
    return smaller_text(entry, area, lambda fit: fit.height_overflow <= 0)


def wider_box_or_smaller_text(entry: Entry, area: SlideArea, entries: list[Entry]) -> dict | str:
    needed = entry.box.w + entry.fit.width_overflow
    grown = inside(Box(entry.box.x, entry.box.y, needed + SUGGESTION_SLACK, entry.box.h), area)
    if grown.w >= needed and not collides(grown, entry, entries):
        return transform_suggestion(entry, area, grown)
    return smaller_text(entry, area, lambda fit: fit.width_overflow <= 0)


def collides(grown: Box, entry: Entry, entries: list[Entry]) -> bool:
    return any(
        other is not entry and is_content(other) and grown.intersection(other.box) is not None and entry.box.intersection(other.box) is None
        for other in entries
    )


def is_content(entry: Entry) -> bool:
    return entry.has_text or entry.kind in MEDIA_KINDS or entry.kind == "shape"


def smaller_text(entry: Entry, area: SlideArea, fits) -> dict | str:
    largest = largest_text_size(entry.context, entry.element)
    for step in range(1, SHRINK_STEPS + 1):
        size = math.floor(largest * (1 - step * SHRINK_STEP))
        if size < SMALLEST_SUGGESTED_SIZE:
            break
        if fits(measure_text(entry.context, restyled(entry.element, size), entry.box)):
            return {"op": "set_text_style", "slide": area.number, "shape": shape_reference(entry.address), "size": size}
    return NO_SINGLE_OPERATION


def restyled(element, size: int):
    copied = copy.deepcopy(element)
    for paragraph in copied.iter(qn("a:p")):
        for properties in character_properties(paragraph):
            apply_run_style(properties, {"size": size})
    return copied


def distortion_issues(entry: Entry, area: SlideArea) -> list[Issue]:
    if entry.pixels is None or entry.box.h <= 0 or entry.box.w <= 0:
        return []
    expected = cropped_ratio(entry.element, entry.pixels)
    actual = entry.box.w / entry.box.h
    distortion = abs(actual / expected - 1)
    if distortion <= DISTORTION_TOLERANCE:
        return []
    if actual > expected:
        corrected = Box(entry.box.x, entry.box.y, round(entry.box.h * expected), entry.box.h)
    else:
        corrected = Box(entry.box.x, entry.box.y, entry.box.w, round(entry.box.w / expected))
    text = f"{label(entry, area)} is stretched {round(distortion * 100)}% away from its image's ratio"
    return [layout_issue(IMAGE_DISTORTED, text, location(entry, area), transform_suggestion(entry, area, inside(corrected, area)))]


def cropped_ratio(element, pixels: tuple[int, int]) -> float:
    crop = element.find(f"{qn('p:blipFill')}/{qn('a:srcRect')}")
    sides = {side: int(crop.get(side, "0")) / CROP_SCALE if crop is not None else 0.0 for side in ("l", "t", "r", "b")}
    width = pixels[0] * (1 - sides["l"] - sides["r"])
    height = pixels[1] * (1 - sides["t"] - sides["b"])
    return width / height if width > 0 and height > 0 else pixels[0] / pixels[1]


def picture_pixels(shape, kind: str) -> tuple[int, int] | None:
    if kind != "picture":
        return None
    blip = shape._element.find(f"{qn('p:blipFill')}/{qn('a:blip')}")
    if blip is None or blip.get(qn("r:embed")) is None:
        return None
    image = shape.image
    return image.size if image.content_type in OFFICE_PICTURE_CONTENT_TYPES else None


def overlap_issues(entries: list[Entry], area: SlideArea) -> list[Issue]:
    slide_area = area.width * area.height
    content = [entry for entry in entries if (entry.has_text or entry.kind in MEDIA_KINDS) and entry.visible_box.area < slide_area * BACKGROUND_SHARE_OF_SLIDE]
    issues = []
    for index, first in enumerate(content):
        for second in content[index + 1:]:
            if not first.has_text and not second.has_text or sits_in_hole(first, second) or sits_in_hole(second, first):
                continue
            shared = first.visible_box.intersection(second.visible_box)
            if shared is None or shared.area < OVERLAP_RATIO * min(first.visible_box.area, second.visible_box.area):
                continue
            text = f"{label(first, area)} and shape {second.address} {second.name!r} overlap by {points(shared.w)}x{points(shared.h)}pt"
            fix = growth_fix(first, area) or growth_fix(second, area) or separation_fix(first, second, area, content)
            issues.append(layout_issue(TEXT_OVERLAP, text, location(first, area), fix))
    return issues


def sits_in_hole(text: Entry, chart: Entry) -> bool:
    hole, box = chart.hole, text.visible_box
    return hole is not None and hole.x <= box.x and hole.y <= box.y and box.right <= hole.right and box.bottom <= hole.bottom


def separation_fix(first: Entry, second: Entry, area: SlideArea, content: list[Entry]) -> dict:
    for moving, staying in ((second, first), (first, second)):
        for box in clear_positions(moving, staying):
            if fits_on_slide(box, area) and not any(other is not moving and box.intersection(other.visible_box) for other in content):
                return transform_suggestion(moving, area, box)
    text_entry = first if first.has_text else second
    return smaller_text(text_entry, area, lambda fit: fit.height_overflow <= 0 and fit.width_overflow <= 0)


def clear_positions(moving: Entry, staying: Entry) -> list[Box]:
    box, other = moving.box, staying.visible_box
    grown_height = moving.visible_box.h - box.h
    return [
        Box(box.x, other.bottom + SUGGESTION_SLACK, box.w, box.h),
        Box(other.right + SUGGESTION_SLACK, box.y, box.w, box.h),
        Box(box.x, other.y - SUGGESTION_SLACK - box.h - grown_height, box.w, box.h),
        Box(other.x - SUGGESTION_SLACK - box.w, box.y, box.w, box.h),
    ]


def fits_on_slide(box: Box, area: SlideArea) -> bool:
    return box.x >= 0 and box.y >= 0 and box.right <= area.width and box.bottom <= area.height


def substitutions(faces: frozenset) -> dict[str, str]:
    return {requested: face.family for requested, face in sorted(faces, key=lambda pair: pair[0]) if face.substituted}


def measured_with(entry: Entry) -> str:
    pairs = substitutions(entry.fit.faces)
    if not pairs:
        return ""
    return " (measured with " + ", ".join(f"{used} in place of {requested}" for requested, used in pairs.items()) + ", which are not installed)"
