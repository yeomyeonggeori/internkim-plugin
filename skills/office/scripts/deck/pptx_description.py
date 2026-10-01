from __future__ import annotations

from pptx.oxml.ns import qn

from pptx_content import frame_text, notes_text
from pptx_geometry import SLIDE_FRAME, Frame, child_frame, local_box, percent_of_slide, points, rotation_degrees
from pptx_inheritance import SlideContext, slide_context
from pptx_shape_kinds import placeholder_type, shape_address, shape_identifier, shape_kind
from pptx_style import resolve_color, run_style


ALIGNMENT_NAMES = {"l": "left", "ctr": "center", "r": "right", "just": "justify", "dist": "distributed"}
AUTOFIT_NAMES = {qn("a:noAutofit"): "none", qn("a:normAutofit"): "shrink", qn("a:spAutoFit"): "resize"}


def describe_presentation(presentation, numbers: list[int], detail: bool) -> dict:
    slides = list(presentation.slides)
    width, height = presentation.slide_width, presentation.slide_height
    return {
        "slideCount": len(slides),
        "slideSize": {"w": width, "h": height, "unit": "EMU", "emuPerPoint": 12700},
        "layouts": [layout.name for master in presentation.slide_masters for layout in master.slide_layouts],
        "slides": [describe_slide(presentation, slides[number - 1], number, detail) for number in numbers],
    }


def describe_slide(presentation, slide, number: int, detail: bool) -> dict:
    context = slide_context(presentation, slide)
    description = {"slide": number, "layout": slide.slide_layout.name}
    if slide._element.get("show") == "0":
        description["hidden"] = True
    description["shapes"] = describe_shapes(slide.shapes, "", SLIDE_FRAME, context, detail, (presentation.slide_width, presentation.slide_height))
    description["notes"] = notes_text(slide)
    if detail and slide._element.find(qn("p:timing")) is not None:
        description["animated"] = sorted({target.get("spid") for target in slide._element.iter(qn("p:spTgt"))})
    return description


def describe_shapes(shapes, prefix: str, frame: Frame, context: SlideContext, detail: bool, slide_size: tuple[int, int]) -> list[dict]:
    return [describe_shape(shape, shape_address(prefix, index), frame, context, detail, slide_size) for index, shape in enumerate(shapes)]


def describe_shape(shape, address: str, frame: Frame, context: SlideContext, detail: bool, slide_size: tuple[int, int]) -> dict:
    element = shape._element
    kind = shape_kind(element)
    box = frame.to_slide(local_box(element, context))
    description = {"index": int(address) if address.isdigit() else address, "id": shape_identifier(element), "name": shape.name, "kind": kind}
    if placeholder_type(element) is not None:
        description["placeholder"] = placeholder_type(element)
    description["box"] = box.to_json()
    description["percent"] = percent_of_slide(box, *slide_size)
    if detail:
        description["points"] = points(box)
        rotation = rotation_degrees(element)
        if rotation:
            description["rotation"] = rotation
    description.update(kind_details(shape, kind, address, frame, context, detail, slide_size))
    return description


def kind_details(shape, kind: str, address: str, frame: Frame, context: SlideContext, detail: bool, slide_size) -> dict:
    if kind == "group":
        return {"shapes": describe_shapes(shape.shapes, address, child_frame(frame, shape._element), context, detail, slide_size)}
    if kind == "table":
        return table_details(shape, detail)
    if kind == "chart":
        return {"chart": chart_details(shape.chart)}
    details = {}
    if shape._element.find(qn("p:txBody")) is not None:
        details.update(text_details(shape, context, detail))
    if kind == "picture" and detail:
        details["picture"] = picture_details(shape)
    if detail and kind in ("shape", "text"):
        details.update(paint_details(shape, context))
    return details


def text_details(shape, context: SlideContext, detail: bool) -> dict:
    text_frame = shape.text_frame
    text = frame_text(text_frame)
    paragraphs = text_frame.paragraphs
    first_run = next((run for paragraph in paragraphs for run in paragraph.runs), None)
    first_paragraph = first_run._r.getparent() if first_run is not None else paragraphs[0]._p
    properties = first_run._r.find(qn("a:rPr")) if first_run is not None else paragraphs[0]._p.find(qn("a:endParaRPr"))
    details = {"text": text}
    if text.strip():
        details["style"] = run_style(context, shape._element, first_paragraph, properties).compact(text)
    if detail:
        details["frame"] = frame_details(shape._element)
        details["paragraphs"] = [paragraph_details(paragraph, shape._element, context) for paragraph in paragraphs]
    return details


def frame_details(element) -> dict:
    body_properties = element.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
    autofit = next((AUTOFIT_NAMES[child.tag] for child in body_properties if child.tag in AUTOFIT_NAMES), "none")
    return {"autofit": autofit, "wrap": body_properties.get("wrap", "square") != "none", "anchor": body_properties.get("anchor", "t")}


def paragraph_details(paragraph, shape_element, context: SlideContext) -> dict:
    properties = paragraph._p.find(qn("a:pPr"))
    details = {"text": "".join(run.text for run in paragraph.runs), "level": paragraph.level}
    if properties is not None and properties.get("algn"):
        details["align"] = ALIGNMENT_NAMES.get(properties.get("algn"), properties.get("algn"))
    if properties is not None and properties.find(qn("a:buChar")) is not None:
        details["bullet"] = "bullet"
    elif properties is not None and properties.find(qn("a:buAutoNum")) is not None:
        details["bullet"] = "number"
    details["runs"] = [
        {"text": run.text, **run_style(context, shape_element, paragraph._p, run._r.find(qn("a:rPr"))).detailed(run.text)}
        for run in paragraph.runs
    ]
    return details


def table_details(shape, detail: bool) -> dict:
    table = shape.table
    details = {"rows": [[frame_text(cell.text_frame) for cell in row.cells] for row in table.rows]}
    if detail:
        details["merged"] = [
            {"row": row_index, "column": column_index, "rows": cell.span_height, "columns": cell.span_width}
            for row_index, row in enumerate(table.rows)
            for column_index, cell in enumerate(row.cells)
            if cell.is_merge_origin
        ]
    return details


def chart_details(chart) -> dict:
    plots = list(chart.plots)
    details = {"type": chart.chart_type.name.lower() if chart.chart_type is not None else "unknown"}
    if chart.has_title and chart.chart_title.has_text_frame:
        details["title"] = frame_text(chart.chart_title.text_frame)
    details["categories"] = list(plots[0].categories) if plots else []
    details["series"] = [{"name": series.name, "values": list(series.values)} for plot in plots for series in plot.series]
    return details


def picture_details(picture) -> dict:
    details = {"contentType": picture.image.content_type, "pixels": list(picture.image.size)}
    crop = {side: round(getattr(picture, f"crop_{side}"), 3) for side in ("left", "top", "right", "bottom")}
    if any(crop.values()):
        details["crop"] = crop
    return details


def paint_details(shape, context: SlideContext) -> dict:
    properties = shape._element.find(qn("p:spPr"))
    if properties is None:
        return {}
    details = {}
    fill = properties.find(qn("a:solidFill"))
    if fill is not None and len(fill):
        details["fill"] = resolve_color(context, fill[0])
    elif properties.find(qn("a:noFill")) is not None:
        details["fill"] = "none"
    line_fill = properties.find(f"{qn('a:ln')}/{qn('a:solidFill')}")
    if line_fill is not None and len(line_fill):
        details["line"] = resolve_color(context, line_fill[0])
    return details
