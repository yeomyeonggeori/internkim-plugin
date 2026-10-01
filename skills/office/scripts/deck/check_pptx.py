#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from deck_definitions import PPTX_NOT_RENDERED
from office_result import OfficeArgumentParser, Result, run_command
from pptx_layout_audit import audit_presentation, substitutions
from pptx_preview import preview_document
from pptx_slide_selection import select_slides


PREVIEW_NAME = "preview.html"
BOLD_WEIGHT = 700
REGULAR_WEIGHT = 400


def main() -> Result:
    arguments = parse_arguments()
    source_path = Path(arguments.presentation_path).expanduser()
    presentation = Presentation(str(source_path))
    slides = list(presentation.slides)
    numbers = select_slides(arguments.slides, len(slides))
    audit = audit_presentation(presentation, [(number, slides[number - 1]) for number in numbers])
    issues = list(audit.issues)
    details = {"checkedSlides": numbers, "fontsMeasuredWith": substitutions(audit.faces)}
    if not arguments.no_preview:
        directory = Path(arguments.output_directory).expanduser() if arguments.output_directory else source_path.with_name(f"{source_path.stem}-check")
        details.update(write_preview(presentation, numbers, directory))
        issues.append(PPTX_NOT_RENDERED.issue(f"the preview of {len(numbers)} slides was written as HTML to {directory / PREVIEW_NAME}, and no image of it was drawn", str(source_path)))
    return Result(summary=f"checked {len(numbers)} slides of {source_path}: {len(audit.issues)} layout issues", output_path=str(source_path), issues=tuple(issues), details=details)


def write_preview(presentation, numbers: list[int], directory: Path) -> dict:
    preview = preview_document(presentation, numbers)
    directory.mkdir(parents=True, exist_ok=True)
    preview_path = directory / PREVIEW_NAME
    preview_path.write_text(preview.html, encoding="utf-8")
    return {"preview": str(preview_path), "previewFonts": preview_fonts(preview.faces), "seen": False}


def preview_fonts(faces: frozenset) -> list[dict]:
    files = {(face.family, face.path, face.index, BOLD_WEIGHT if face.bold else REGULAR_WEIGHT) for _, face in faces}
    return [{"family": family, "path": path, "index": index, "weight": weight} for family, path, index, weight in sorted(files)]


def parse_arguments():
    parser = OfficeArgumentParser(description=(
        "Check an edited .pptx: text measured with the deck's fonts that overflows its box, shapes off the slide, overlapping text and stretched pictures, "
        "each with an operation deck apply accepts. Then write an HTML preview of the slides, drawn from the same geometry, styles and fonts, to look at."
    ))
    parser.add_argument("presentation_path")
    parser.add_argument("--slides", default="", help="slides to check, such as 2,4-6; default every slide")
    parser.add_argument("--output-directory", default="", help="where the preview goes, default <file name>-check beside the file")
    parser.add_argument("--no-preview", action="store_true", help="measure only, without writing the preview")
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
