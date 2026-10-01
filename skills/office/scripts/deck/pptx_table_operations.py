from __future__ import annotations

import copy

from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.table import _Cell

from core.office_operations import OPERATION_NOT_APPLICABLE, TARGET_NOT_FOUND, Change
from core.office_result import INVALID_VALUE, OfficeFailure
from deck.pptx_insert_operations import cell_text, chart_data, set_chart_title
from deck.pptx_targets import PptxEditing, ShapeTarget, require_kind, resolve_shape
from deck.pptx_text_operations import replace_text


MERGE_ATTRIBUTES = ("gridSpan", "rowSpan", "hMerge", "vMerge")
NON_CATEGORY_CHARTS = {
    XL_CHART_TYPE.XY_SCATTER, XL_CHART_TYPE.XY_SCATTER_LINES, XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
    XL_CHART_TYPE.XY_SCATTER_SMOOTH, XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS, XL_CHART_TYPE.BUBBLE, XL_CHART_TYPE.BUBBLE_THREE_D_EFFECT,
}


def table_element(target: ShapeTarget):
    return target.element.find(f"{qn('a:graphic')}/{qn('a:graphicData')}/{qn('a:tbl')}")


def table_rows(target: ShapeTarget) -> list:
    return table_element(target).findall(qn("a:tr"))


def grid_columns(target: ShapeTarget) -> list:
    return table_element(target).find(qn("a:tblGrid")).findall(qn("a:gridCol"))


def resolve_table(editing: PptxEditing, operation: dict, location: str) -> ShapeTarget:
    target = resolve_shape(editing, operation, location)
    require_kind(target, ("table",), location, f"{operation['op']} edits a table")
    return target


def require_index(index: int, count: int, noun: str, target: ShapeTarget, location: str, allow_end: bool = False) -> None:
    limit = count if allow_end else count - 1
    if index <= limit:
        return
    raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: {target.label} has no {noun} {index}", location, f"use a {noun} from 0 to {limit}"))


def require_unmerged(target: ShapeTarget, location: str) -> None:
    if any(cell.get(name) is not None for cell in table_element(target).iter(qn("a:tc")) for name in MERGE_ATTRIBUTES):
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.shape: {target.label} has merged cells, so rows and columns cannot be inserted or deleted safely", f"{location}.shape", "edit the cell text instead, or rebuild the table with add_table"))


def require_values_fit(values: list | None, count: int, location: str) -> None:
    if values is not None and len(values) > count:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.values: {len(values)} values for {count} cells", f"{location}.values", f"give at most {count} values"))


def fill_cells(cells: list, values: list | None, target: ShapeTarget) -> None:
    padded = list(values or []) + [""] * (len(cells) - len(values or []))
    for cell_element, value in zip(cells, padded):
        replace_text(_Cell(cell_element, target.shape).text_frame, None, cell_text(value))


def grow_frame(target: ShapeTarget, dimension: str, amount: int) -> None:
    extent = target.element.find(f"{qn('p:xfrm')}/{qn('a:ext')}")
    extent.set(dimension, str(max(1, int(extent.get(dimension)) + amount)))


def plan_set_table_cell(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_table(editing, operation, location)
    require_index(operation["row"], len(table_rows(target)), "row", target, f"{location}.row")
    require_index(operation["column"], len(grid_columns(target)), "column", target, f"{location}.column")
    cell = target.shape.table.cell(operation["row"], operation["column"])
    if cell.is_spanned:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: cell {operation['row']},{operation['column']} of {target.label} is covered by a merged cell", location, "write the text into the merged block's top-left cell"))

    def change() -> str:
        replace_text(cell.text_frame, None, operation["text"])
        editing.mark_edited(target.slide)
        return f"set cell {operation['row']},{operation['column']} of {target.label}"
    return change


def plan_insert_table_row(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_table(editing, operation, location)
    require_unmerged(target, location)
    rows = table_rows(target)
    at = operation["at"] if operation.get("at") is not None else len(rows)
    require_index(at, len(rows), "row", target, f"{location}.at", allow_end=True)
    require_values_fit(operation.get("values"), len(grid_columns(target)), location)

    def change() -> str:
        neighbor = rows[at - 1] if at > 0 else rows[0]
        clone = copy.deepcopy(neighbor)
        fill_cells(clone.findall(qn("a:tc")), operation.get("values"), target)
        if at < len(rows):
            rows[at].addprevious(clone)
        else:
            rows[-1].addnext(clone)
        grow_frame(target, "cy", int(clone.get("h")))
        editing.mark_edited(target.slide)
        return f"inserted row {at} into {target.label}"
    return change


def plan_delete_table_row(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_table(editing, operation, location)
    require_unmerged(target, location)
    rows = table_rows(target)
    require_index(operation["row"], len(rows), "row", target, f"{location}.row")
    require_more_than_one(len(rows), "row", target, location)

    def change() -> str:
        row = rows[operation["row"]]
        row.getparent().remove(row)
        grow_frame(target, "cy", -int(row.get("h")))
        editing.mark_edited(target.slide)
        return f"deleted row {operation['row']} of {target.label}"
    return change


def require_more_than_one(count: int, noun: str, target: ShapeTarget, location: str) -> None:
    if count > 1:
        return
    raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: {target.label} has only one {noun}", location, "delete the table with delete_shape instead"))


def plan_insert_table_column(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_table(editing, operation, location)
    require_unmerged(target, location)
    columns = grid_columns(target)
    at = operation["at"] if operation.get("at") is not None else len(columns)
    require_index(at, len(columns), "column", target, f"{location}.at", allow_end=True)
    require_values_fit(operation.get("values"), len(table_rows(target)), location)

    def change() -> str:
        source = at - 1 if at > 0 else 0
        new_column = copy.deepcopy(columns[source])
        place(columns, at, new_column)
        new_cells = []
        for row in table_rows(target):
            cells = row.findall(qn("a:tc"))
            new_cell = copy.deepcopy(cells[source])
            place(cells, at, new_cell)
            new_cells.append(new_cell)
        fill_cells(new_cells, operation.get("values"), target)
        grow_frame(target, "cx", int(new_column.get("w")))
        editing.mark_edited(target.slide)
        return f"inserted column {at} into {target.label}"
    return change


def place(siblings: list, at: int, element) -> None:
    if at < len(siblings):
        siblings[at].addprevious(element)
    else:
        siblings[-1].addnext(element)


def plan_delete_table_column(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_table(editing, operation, location)
    require_unmerged(target, location)
    columns = grid_columns(target)
    require_index(operation["column"], len(columns), "column", target, f"{location}.column")
    require_more_than_one(len(columns), "column", target, location)

    def change() -> str:
        column = columns[operation["column"]]
        column.getparent().remove(column)
        for row in table_rows(target):
            cell = row.findall(qn("a:tc"))[operation["column"]]
            row.remove(cell)
        grow_frame(target, "cx", -int(column.get("w")))
        editing.mark_edited(target.slide)
        return f"deleted column {operation['column']} of {target.label}"
    return change


def plan_merge_table_cells(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_table(editing, operation, location)
    first_row, first_column = operation["row"], operation["column"]
    last_row = first_row + (operation.get("rows") or 1) - 1
    last_column = first_column + (operation.get("columns") or 1) - 1
    require_index(last_row, len(table_rows(target)), "row", target, f"{location}.rows")
    require_index(last_column, len(grid_columns(target)), "column", target, f"{location}.columns")
    if (last_row, last_column) == (first_row, first_column):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: one cell is not a block to merge", location, "give rows or columns greater than 1"))
    table = target.shape.table
    block = [table.cell(row, column) for row in range(first_row, last_row + 1) for column in range(first_column, last_column + 1)]
    if any(cell.is_merge_origin or cell.is_spanned for cell in block):
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: the block overlaps cells that are already merged", location, "pick a block of unmerged cells"))

    def change() -> str:
        table.cell(first_row, first_column).merge(table.cell(last_row, last_column))
        editing.mark_edited(target.slide)
        return f"merged rows {first_row}-{last_row}, columns {first_column}-{last_column} of {target.label}"
    return change


def plan_set_chart_data(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_shape(editing, operation, location)
    require_kind(target, ("chart",), location, "set_chart_data edits a chart")
    chart = target.shape.chart
    if chart.chart_type in NON_CATEGORY_CHARTS:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.shape: {target.label} plots x-y points, which set_chart_data does not write", f"{location}.shape", "replace the chart with add_chart"))
    categories = operation.get("categories") or list(chart.plots[0].categories)
    data = chart_data(categories, operation["series"], location, existing_number_format(chart))

    def change() -> str:
        chart.replace_data(data)
        set_chart_title(chart, operation.get("title"))
        editing.mark_edited(target.slide)
        return f"replaced the data of {target.label} with {len(operation['series'])} series over {len(categories)} categories"
    return change


def existing_number_format(chart) -> str:
    format_code = chart._chartSpace.find(f".//{qn('c:val')}//{qn('c:formatCode')}")
    return format_code.text if format_code is not None and format_code.text else "General"


TABLE_AND_CHART_PLANNERS = {
    "set_table_cell": plan_set_table_cell,
    "insert_table_row": plan_insert_table_row,
    "delete_table_row": plan_delete_table_row,
    "insert_table_column": plan_insert_table_column,
    "delete_table_column": plan_delete_table_column,
    "merge_table_cells": plan_merge_table_cells,
    "set_chart_data": plan_set_chart_data,
}

