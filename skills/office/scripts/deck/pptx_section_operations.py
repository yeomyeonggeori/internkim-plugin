from __future__ import annotations

import uuid

from lxml import etree
from pptx.oxml.ns import qn

from office_operations import TARGET_NOT_FOUND, Change
from office_result import INVALID_VALUE, OfficeFailure
from office_schema import closest_suggestion
from pptx_sections import POWERPOINT_2010_NAMESPACE, SECTION_LIST_TAG, SECTION_SLIDE_TAG, SECTION_SLIDES_TAG, SECTION_TAG
from pptx_targets import PptxEditing, resolve_slide


SECTION_EXTENSION_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
DEFAULT_SECTION_NAME = "Default Section"


def section_list(editing: PptxEditing):
    return next(editing.presentation.part._element.iter(SECTION_LIST_TAG), None)


def sections_of(editing: PptxEditing) -> list:
    found = section_list(editing)
    return found.findall(SECTION_TAG) if found is not None else []


def section_slide_ids(section) -> list[str]:
    return [entry.get("id") for entry in section.iter(SECTION_SLIDE_TAG)]


def set_section_slide_ids(section, slide_ids: list[str]) -> None:
    slides = section.find(SECTION_SLIDES_TAG)
    if slides is None:
        slides = etree.SubElement(section, SECTION_SLIDES_TAG)
    for entry in list(slides):
        slides.remove(entry)
    for slide_id in slide_ids:
        etree.SubElement(slides, SECTION_SLIDE_TAG, id=slide_id)


def new_section(name: str, slide_ids: list[str]):
    section = etree.Element(SECTION_TAG, name=name, id="{" + str(uuid.uuid4()).upper() + "}", nsmap={"p14": POWERPOINT_2010_NAMESPACE})
    set_section_slide_ids(section, slide_ids)
    return section


def created_section_list(editing: PptxEditing):
    presentation_element = editing.presentation.part._element
    extensions = presentation_element.find(qn("p:extLst"))
    if extensions is None:
        extensions = etree.SubElement(presentation_element, qn("p:extLst"))
    extension = etree.SubElement(extensions, qn("p:ext"), uri=SECTION_EXTENSION_URI)
    listing = etree.SubElement(extension, SECTION_LIST_TAG, nsmap={"p14": POWERPOINT_2010_NAMESPACE})
    listing.append(new_section(DEFAULT_SECTION_NAME, [entry.get("id") for entry in editing.presentation.slides._sldIdLst]))
    return listing


def resolve_section(editing: PptxEditing, name: str, location: str):
    sections = sections_of(editing)
    match = next((section for section in sections if section.get("name") == name), None)
    if match is not None:
        return match
    names = [section.get("name") for section in sections]
    if not names:
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: the deck has no sections", location, "start one with add_section"))
    raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: the deck has no section {name!r}", location, closest_suggestion(name, names) or "use one of " + ", ".join(map(repr, names))))


def plan_add_section(editing: PptxEditing, operation: dict, location: str) -> Change:
    resolve_slide(editing, operation["slide"], f"{location}.slide")
    slide_id = editing.slide_id_elements[operation["slide"] - 1].get("id")

    def change() -> str:
        listing = section_list(editing)
        created = listing is None
        if created:
            listing = created_section_list(editing)
        sections = listing.findall(SECTION_TAG)
        owner = next((section for section in sections if slide_id in section_slide_ids(section)), sections[-1])
        members = section_slide_ids(owner)
        start = members.index(slide_id) if slide_id in members else len(members)
        set_section_slide_ids(owner, members[:start])
        owner.addnext(new_section(operation["name"], members[start:]))
        if created and start == 0:
            listing.remove(owner)
        return f"started section {operation['name']!r} at slide {operation['slide']}"
    return change


def plan_rename_section(editing: PptxEditing, operation: dict, location: str) -> Change:
    section = resolve_section(editing, operation["section"], f"{location}.section")

    def change() -> str:
        section.set("name", operation["name"])
        return f"renamed section {operation['section']!r} to {operation['name']!r}"
    return change


def plan_move_to_section(editing: PptxEditing, operation: dict, location: str) -> Change:
    section = resolve_section(editing, operation["section"], f"{location}.section")
    if len(set(operation["slides"])) != len(operation["slides"]):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.slides: a slide is listed twice", f"{location}.slides"))
    for index, number in enumerate(operation["slides"]):
        resolve_slide(editing, number, f"{location}.slides[{index}]")
    moved = [editing.slide_id_elements[number - 1] for number in operation["slides"]]

    def change() -> str:
        moved_ids = [element.get("id") for element in moved]
        for other in sections_of(editing):
            set_section_slide_ids(other, [slide_id for slide_id in section_slide_ids(other) if slide_id not in moved_ids])
        place_after_section(editing, section, moved)
        set_section_slide_ids(section, section_slide_ids(section) + moved_ids)
        editing.structure_changed = True
        return f"moved slides {', '.join(map(str, operation['slides']))} to the end of section {operation['section']!r}"
    return change


def place_after_section(editing: PptxEditing, section, moved: list) -> None:
    slide_list = editing.presentation.slides._sldIdLst
    for element in moved:
        slide_list.remove(element)
    by_id = {element.get("id"): element for element in slide_list}
    anchor = next((by_id[slide_id] for slide_id in reversed(section_slide_ids(section)) if slide_id in by_id), None)
    if anchor is None:
        anchor = last_slide_before(editing, section, by_id)
    for element in moved:
        if anchor is None:
            slide_list.insert(0, element)
        else:
            anchor.addnext(element)
        anchor = element


def last_slide_before(editing: PptxEditing, section, by_id: dict):
    sections = sections_of(editing)
    earlier = sections[:sections.index(section)]
    for previous in reversed(earlier):
        present = [by_id[slide_id] for slide_id in section_slide_ids(previous) if slide_id in by_id]
        if present:
            return present[-1]
    return None


def plan_remove_section(editing: PptxEditing, operation: dict, location: str) -> Change:
    section = resolve_section(editing, operation["section"], f"{location}.section")

    def change() -> str:
        sections = sections_of(editing)
        if len(sections) == 1:
            extension = section.getparent().getparent()
            extension.getparent().remove(extension)
            return f"removed section {operation['section']!r}, the deck's only one"
        position = sections.index(section)
        heir = sections[position - 1] if position > 0 else sections[1]
        merged = section_slide_ids(heir) + section_slide_ids(section) if position > 0 else section_slide_ids(section) + section_slide_ids(heir)
        set_section_slide_ids(heir, merged)
        section.getparent().remove(section)
        return f"removed section {operation['section']!r}; its slides joined {heir.get('name')!r}"
    return change


def describe_sections(presentation) -> list[dict]:
    listing = next(presentation.part._element.iter(SECTION_LIST_TAG), None)
    if listing is None:
        return []
    numbers = {element.get("id"): number for number, element in enumerate(presentation.slides._sldIdLst, start=1)}
    return [
        {"name": section.get("name"), "slides": [numbers[slide_id] for slide_id in section_slide_ids(section) if slide_id in numbers]}
        for section in listing.findall(SECTION_TAG)
    ]


SECTION_PLANNERS = {
    "add_section": plan_add_section,
    "rename_section": plan_rename_section,
    "move_to_section": plan_move_to_section,
    "remove_section": plan_remove_section,
}
