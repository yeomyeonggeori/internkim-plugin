from __future__ import annotations

import uuid

from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from core.office_operations import Change
from core.office_result import INVALID_VALUE, OfficeFailure
from deck.pptx_animation import remove_animations_of
from deck.pptx_element_operations import next_shape_identifier
from deck.pptx_shape_kinds import placeholder_of, placeholder_type
from deck.pptx_targets import PptxEditing, live_slides, resolve_slide
from deck.pptx_text import text_content


FOOTER_FIELDS = {"footer": "ftr", "slideNumber": "sldNum", "date": "dt"}
PLACEHOLDER_NAMES = {"ftr": "Footer", "sldNum": "Slide Number", "dt": "Date"}
FALLBACK_BOXES = {"dt": (0.05, 0.22, "l"), "ftr": (0.3, 0.4, "ctr"), "sldNum": (0.73, 0.22, "r")}
FALLBACK_TOP_SHARE = 0.93
FALLBACK_HEIGHT_SHARE = 0.045
FALLBACK_SIZE = 1200
FALLBACK_COLOR = "898989"


def plan_set_header_footer(editing: PptxEditing, operation: dict, location: str) -> Change:
    requested = {placeholder: operation[field] for field, placeholder in FOOTER_FIELDS.items() if operation.get(field) is not None}
    if not requested:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: set_header_footer changes nothing", location, "give footer, slideNumber or date"))
    numbered = [(operation["slide"], resolve_slide(editing, operation["slide"], f"{location}.slide"))] if operation.get("slide") is not None else numbered_live_slides(editing)

    def change() -> str:
        for number, slide in numbered:
            for placeholder, value in requested.items():
                replace_footer_placeholder(editing, slide, number, placeholder, value)
            editing.mark_edited(slide)
        scope = f"slide {operation['slide']}" if operation.get("slide") is not None else "every slide"
        return f"set {', '.join(field for field in FOOTER_FIELDS if operation.get(field) is not None)} on {scope}"
    return change


def numbered_live_slides(editing: PptxEditing) -> list[tuple[int, object]]:
    live = {id(slide) for slide in live_slides(editing)}
    return [(number, slide) for number, slide in enumerate(editing.slides, start=1) if id(slide) in live]


def replace_footer_placeholder(editing: PptxEditing, slide, number: int, placeholder: str, value) -> None:
    tree = slide.shapes._spTree
    existing = [shape for shape in tree if placeholder_type(shape) == placeholder]
    identifiers = {shape.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}").get("id") for shape in existing}
    for shape in existing:
        tree.remove(shape)
    remove_animations_of(slide._element, identifiers)
    if value in ("", False):
        return
    text = str(number) if placeholder == "sldNum" else value
    tree.append(footer_shape(editing, slide, placeholder, text))


def layout_placeholder(slide, placeholder: str):
    for owner in (slide.slide_layout, slide.slide_layout.slide_master):
        match = next((shape for shape in owner.shapes._spTree if placeholder_type(shape) == placeholder), None)
        if match is not None:
            return match
    return None


def footer_shape(editing: PptxEditing, slide, placeholder: str, text: str):
    inherited = layout_placeholder(slide, placeholder)
    identifier = next_shape_identifier(slide._element)
    index = placeholder_of(inherited).get("idx") if inherited is not None else None
    index_attribute = f' idx="{index}"' if index is not None else ""
    geometry, run_style = ("<p:spPr/>", "") if inherited is not None else fallback_geometry(editing, placeholder)
    return parse_xml(
        f'<p:sp {nsdecls("p", "a")}><p:nvSpPr><p:cNvPr id="{identifier}" name="{PLACEHOLDER_NAMES[placeholder]} {identifier}"/>'
        f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="{placeholder}"{index_attribute}/></p:nvPr></p:nvSpPr>'
        f'{geometry}<p:txBody><a:bodyPr/><a:lstStyle/><a:p>{paragraph_properties(inherited, placeholder)}{footer_run(placeholder, text, run_style)}</a:p></p:txBody></p:sp>'
    )


def fallback_geometry(editing: PptxEditing, placeholder: str) -> tuple[str, str]:
    width, height = editing.presentation.slide_width, editing.presentation.slide_height
    left_share, width_share, _ = FALLBACK_BOXES[placeholder]
    box = (round(width * left_share), round(height * FALLBACK_TOP_SHARE), round(width * width_share), round(height * FALLBACK_HEIGHT_SHARE))
    geometry = f'<p:spPr><a:xfrm><a:off x="{box[0]}" y="{box[1]}"/><a:ext cx="{box[2]}" cy="{box[3]}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
    run_style = f' sz="{FALLBACK_SIZE}"><a:solidFill><a:srgbClr val="{FALLBACK_COLOR}"/></a:solidFill></a:rPr>'
    return geometry, run_style


def paragraph_properties(inherited, placeholder: str) -> str:
    if inherited is not None:
        return ""
    return f'<a:pPr algn="{FALLBACK_BOXES[placeholder][2]}"/>'


def footer_run(placeholder: str, text: str, run_style: str) -> str:
    properties = f'<a:rPr lang="ko-KR"{run_style}' if run_style else '<a:rPr lang="ko-KR"/>'
    if placeholder == "sldNum":
        return f'<a:fld id="{{{str(uuid.uuid4()).upper()}}}" type="slidenum">{properties}<a:t>{text}</a:t></a:fld>'
    return f"<a:r>{properties}<a:t>{text_content(text)}</a:t></a:r>"


FOOTER_PLANNERS = {
    "set_header_footer": plan_set_header_footer,
}
