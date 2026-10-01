from __future__ import annotations

from dataclasses import dataclass
import html

from pptx.oxml.ns import qn

from charts.svg import chart_svg
from fonts.registry import resolved_face
from pptx_chart_look import chart_look, chart_model
from pptx_description import chart_details
from pptx_geometry import SLIDE_FRAME, Box, Frame, child_frame, local_box, rotation_degrees
from pptx_inheritance import SlideContext, slide_context
from pptx_preview_paint import PIXELS_PER_EMU, background_color, css, data_uri, element, fill_color, outline, pixels, svg_uri
from pptx_preview_shapes import geometry_html
from pptx_preview_table import table_html
from pptx_preview_text import TextPaint, body_html, body_layout
from pptx_shape_kinds import placeholder_of, shape_kind
from pptx_style import has_east_asian
from pptx_text_measure import font_face, grown_box, measure_text


DRAWABLE_IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/bmp", "image/svg+xml"}
CROP_SCALE = 100000
SLIDE_GAP_PIXELS = 16
UNDRAWN_OUTLINE = "1px dashed #9CA3AF"


@dataclass(frozen=True)
class Preview:
    html: str
    faces: frozenset


@dataclass(frozen=True)
class Canvas:
    context: SlideContext
    faces: set


def preview_document(presentation, numbers: list[int]) -> Preview:
    slides = list(presentation.slides)
    faces = set()
    sections = [slide_section(presentation, slides[number - 1], number, faces) for number in numbers]
    document = (
        '<!DOCTYPE html><html><head><meta charset="utf-8"><title>pptx preview</title></head>'
        f'<body style="margin:0px;background-color:#E5E7EB">{"".join(sections)}</body></html>'
    )
    return Preview(document, frozenset(faces))


def slide_section(presentation, slide, number: int, faces: set) -> str:
    canvas = Canvas(slide_context(presentation, slide), faces)
    layout, master = slide.slide_layout, slide.slide_layout.slide_master
    owners = [slide, layout, master]
    color, picture_fill, owner = background_color(canvas.context, [item._element for item in owners])
    layers = [background_picture_html(picture_fill, owners[owner].part) if picture_fill is not None else ""]
    if shows_master_shapes(slide._element) and shows_master_shapes(layout._element):
        layers.extend(decoration_html(master.shapes, canvas))
    if shows_master_shapes(slide._element):
        layers.extend(decoration_html(layout.shapes, canvas))
    layers.extend(shape_html(shape, SLIDE_FRAME, canvas) for shape in slide.shapes)
    section = {
        "position": "relative",
        "width": f"{pixels(presentation.slide_width)}px",
        "height": f"{pixels(presentation.slide_height)}px",
        "overflow": "hidden",
        "background-color": color or "#FFFFFF",
        "margin-bottom": f"{SLIDE_GAP_PIXELS}px",
    }
    return element("section", section, "".join(layers), f' data-slide="{number}"')


def shows_master_shapes(owner_element) -> bool:
    return owner_element.get("showMasterSp", "1") not in ("0", "false")


def decoration_html(shapes, canvas: Canvas) -> list[str]:
    return [shape_html(shape, SLIDE_FRAME, canvas) for shape in shapes if placeholder_of(shape._element) is None]


def background_picture_html(picture_fill, part) -> str:
    blip = picture_fill.find(qn("a:blip"))
    relationship_id = blip.get(qn("r:embed")) if blip is not None else None
    if relationship_id is None:
        return ""
    image = part.related_part(relationship_id)
    if image.content_type not in DRAWABLE_IMAGE_TYPES:
        return ""
    style = {"position": "absolute", "left": "0px", "top": "0px", "width": "100%", "height": "100%", "object-fit": "cover"}
    return f'<img src="{data_uri(image.content_type, image.blob)}" style="{html.escape(css(style))}"/>'


def placed(box: Box, element_tree, extra: dict | None = None) -> dict:
    rotation = rotation_degrees(element_tree)
    return {
        "position": "absolute",
        "left": f"{pixels(box.x)}px",
        "top": f"{pixels(box.y)}px",
        "width": f"{pixels(box.w)}px",
        "height": f"{pixels(box.h)}px",
        "transform": f"rotate({rotation:g}deg)" if rotation else None,
        **(extra or {}),
    }


def shape_html(shape, frame: Frame, canvas: Canvas) -> str:
    tree = shape._element
    kind = shape_kind(tree)
    if kind == "group":
        inner = child_frame(frame, tree)
        return "".join(shape_html(child, inner, canvas) for child in shape.shapes)
    box = frame.to_slide(local_box(tree, canvas.context))
    if kind in ("picture", "media"):
        return picture_html(shape, box, canvas)
    if kind == "table":
        return element("div", placed(box, tree), table_html(canvas.context, tree, pixels(box.w), pixels(box.h), canvas.faces))
    if kind == "chart":
        return chart_html(shape, box, canvas)
    if kind in ("text", "shape", "connector"):
        return autoshape_html(shape, box, canvas)
    return element("div", placed(box, tree, {"box-sizing": "border-box", "border": UNDRAWN_OUTLINE}))


def autoshape_html(shape, box: Box, canvas: Canvas) -> str:
    tree = shape._element
    body = tree.find(qn("p:txBody"))
    box = grown_box(tree, box, measure_text(canvas.context, tree, box))
    properties, style = tree.find(qn("p:spPr")), tree.find(qn("p:style"))
    fill, line = fill_color(canvas.context, properties, style), outline(canvas.context, properties, style)
    has_text = body is not None and any((node.text or "").strip() for node in body.iter(qn("a:t")))
    if not has_text and fill is None and line is None:
        return ""
    paint = geometry_html(properties, fill, line, pixels(box.w), pixels(box.h))
    text = ""
    if has_text:
        inherited = [inherited_tree.find(qn("p:txBody")) for _, inherited_tree in canvas.context.placeholder_chain(tree)]
        layout = body_layout(canvas.context, tree, body, [inherited_body for inherited_body in inherited if inherited_body is not None])
        text = body_html(TextPaint(canvas.context, tree, canvas.faces), body, layout)
    return element("div", placed(box, tree), paint + text)


def picture_html(shape, box: Box, canvas: Canvas) -> str:
    tree = shape._element
    blip = tree.find(f"{qn('p:blipFill')}/{qn('a:blip')}")
    if blip is None or blip.get(qn("r:embed")) is None or shape.image.content_type not in DRAWABLE_IMAGE_TYPES:
        return element("div", placed(box, tree, {"box-sizing": "border-box", "border": UNDRAWN_OUTLINE}))
    crop = tree.find(f"{qn('p:blipFill')}/{qn('a:srcRect')}")
    sides = {side: int(crop.get(side, "0")) / CROP_SCALE if crop is not None else 0.0 for side in ("l", "t", "r", "b")}
    width, height = pixels(box.w), pixels(box.h)
    inner_width = width / max(1 - sides["l"] - sides["r"], 0.01)
    inner_height = height / max(1 - sides["t"] - sides["b"], 0.01)
    image_style = {
        "position": "absolute",
        "left": f"{-sides['l'] * inner_width:.2f}px",
        "top": f"{-sides['t'] * inner_height:.2f}px",
        "width": f"{inner_width:.2f}px",
        "height": f"{inner_height:.2f}px",
    }
    image = f'<img src="{data_uri(shape.image.content_type, shape.image.blob)}" style="{html.escape(css(image_style))}"/>'
    line = outline(canvas.context, tree.find(qn("p:spPr")), tree.find(qn("p:style")))
    frame_style = {"overflow": "hidden", "box-sizing": "border-box", "border": f"{line.width}px solid {line.color}" if line else None}
    return element("div", placed(box, tree, frame_style), image)


def chart_html(shape, box: Box, canvas: Canvas) -> str:
    details = chart_details(shape.chart)
    labels = [str(category) for category in details["categories"]] + [entry["name"] for entry in details["series"]] + [details.get("title") or ""]
    korean = any(has_east_asian(label) for label in labels)
    own = shape.chart._chartSpace.find(f"{qn('c:txPr')}//{qn('a:defRPr')}/{qn('a:latin')}")
    minor = canvas.context.theme_element.find(f"{qn('a:themeElements')}/{qn('a:fontScheme')}/{qn('a:minorFont')}/{qn('a:latin')}")
    requested = next((typeface.get("typeface") for typeface in (own, minor) if typeface is not None), "sans-serif")
    face = font_face(requested, False, korean)
    canvas.faces.add((requested, face))
    width, height = box.w * PIXELS_PER_EMU, box.h * PIXELS_PER_EMU
    svg = chart_svg(chart_model(shape.chart, details), width, height, chart_look(shape.chart, canvas.context), resolved_face(requested).matched_family)
    return element("div", placed(box, shape._element), f'<img src="{svg_uri(svg)}" style="width:{width:.2f}px;height:{height:.2f}px"/>')
