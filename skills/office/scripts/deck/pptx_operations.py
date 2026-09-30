import copy
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph

from deck_definitions import OPERATIONS
from office_operations import OPERATION_NOT_APPLICABLE, TARGET_NOT_FOUND, Change, OperationSet
from office_result import INVALID_VALUE, OfficeFailure
from pptx_content import text_frames_in
from run_replacement import joined_text, replace_in_runs


PARAGRAPH_KEPT_TAGS = {qn("a:pPr"), qn("a:endParaRPr")}


@dataclass
class PptxEditing:
    presentation: Presentation
    slides: list
    slide_id_elements: list
    deleted: set = field(default_factory=set)
    touched: set = field(default_factory=set)


def load_editing(path: str) -> PptxEditing:
    presentation = Presentation(path)
    return PptxEditing(presentation, list(presentation.slides), list(presentation.slides._sldIdLst))


def save_editing(editing: PptxEditing, path: str) -> None:
    editing.presentation.save(path)


def resolve_slide(editing: PptxEditing, number: int, location: str):
    if number > len(editing.slides):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: slide {number} does not exist; the deck has {len(editing.slides)} slides", location))
    if number in editing.deleted:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: slide {number} is deleted by another operation in this batch", location))
    editing.touched.add(number)
    return editing.slides[number - 1]


def resolve_text_shape(editing: PptxEditing, operation: dict, location: str):
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")
    shapes = list(slide.shapes)
    if operation["shape"] >= len(shapes):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.shape: slide {operation['slide']} has {len(shapes)} shapes", f"{location}.shape"))
    shape = shapes[operation["shape"]]
    if not shape.has_text_frame:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.shape: shape {operation['shape']} of slide {operation['slide']} holds no text", f"{location}.shape"))
    return shape


def plan_set_text(editing: PptxEditing, operation: dict, location: str) -> Change:
    shape = resolve_text_shape(editing, operation, location)

    def change() -> str:
        set_frame_text(shape.text_frame, operation["text"])
        return f"set the text of shape {operation['shape']} on slide {operation['slide']}"
    return change


def set_frame_text(text_frame, text: str) -> None:
    paragraphs = text_frame.paragraphs
    if not paragraphs:
        text_frame.text = text
        return
    lines = text.split("\n")
    template = copy.deepcopy(paragraphs[0]._p)
    for paragraph in paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    set_paragraph_line(paragraphs[0], lines[0])
    anchor = paragraphs[0]._p
    for line in lines[1:]:
        clone = copy.deepcopy(template)
        anchor.addnext(clone)
        set_paragraph_line(_Paragraph(clone, paragraphs[0]._parent), line)
        anchor = clone


def set_paragraph_line(paragraph, line: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run().text = line
        return
    runs[0].text = line
    first_run = runs[0]._r
    for element in list(paragraph._p):
        if element is not first_run and element.tag not in PARAGRAPH_KEPT_TAGS:
            paragraph._p.remove(element)


def plan_find_replace(editing: PptxEditing, operation: dict, location: str) -> Change:
    frames = scoped_text_frames(editing, operation, location)
    runs_by_paragraph = [paragraph.runs for frame in frames for paragraph in frame.paragraphs]
    occurrences = sum(joined_text(runs).count(operation["find"]) for runs in runs_by_paragraph)
    if occurrences == 0:
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: {operation['find']!r} does not occur", location))

    def change() -> str:
        replaced = sum(replace_in_runs(runs, operation["find"], operation["replace"]) for runs in runs_by_paragraph)
        return f"replaced {replaced} occurrences of {operation['find']!r}"
    return change


def scoped_text_frames(editing: PptxEditing, operation: dict, location: str) -> list:
    if operation.get("slide") is not None:
        slides = [resolve_slide(editing, operation["slide"], f"{location}.slide")]
    else:
        slides = [slide for number, slide in enumerate(editing.slides, start=1) if number not in editing.deleted]
    return [frame for slide in slides for frame in text_frames_in(slide.shapes)]


def plan_set_notes(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")

    def change() -> str:
        slide.notes_slide.notes_text_frame.text = operation["text"]
        return f"set the notes of slide {operation['slide']}"
    return change


def plan_delete_slide(editing: PptxEditing, operation: dict, location: str) -> Change:
    number = operation["slide"]
    if number in editing.touched:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.slide: slide {number} is used by another operation in this batch", f"{location}.slide"))
    resolve_slide(editing, number, f"{location}.slide")
    editing.deleted.add(number)
    slide_id_element = editing.slide_id_elements[number - 1]

    def change() -> str:
        editing.presentation.part.drop_rel(slide_id_element.rId)
        slide_id_element.getparent().remove(slide_id_element)
        return f"deleted slide {number}"
    return change


def plan_reorder(editing: PptxEditing, operation: dict, location: str) -> Change:
    order = operation["order"]
    out_of_range = [number for number in order if number > len(editing.slides)]
    if out_of_range:
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.order: slide {out_of_range[0]} does not exist; the deck has {len(editing.slides)} slides", f"{location}.order"))
    if len(set(order)) != len(order):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.order: a slide is listed twice", f"{location}.order"))

    def change() -> str:
        remaining = [number for number in range(1, len(editing.slides) + 1) if number not in editing.deleted]
        if sorted(order) != remaining:
            raise OfficeFailure(INVALID_VALUE.issue(f"{location}.order: it must list exactly the slides that remain, {remaining}", f"{location}.order"))
        slide_id_list = editing.presentation.slides._sldIdLst
        for number in order:
            element = editing.slide_id_elements[number - 1]
            slide_id_list.remove(element)
            slide_id_list.append(element)
        return "reordered the slides to " + ", ".join(str(number) for number in order)
    return change


PPTX_OPERATIONS = OperationSet(OPERATIONS, {
    "set_text": plan_set_text,
    "find_replace": plan_find_replace,
    "set_notes": plan_set_notes,
    "delete_slide": plan_delete_slide,
    "reorder": plan_reorder,
})
