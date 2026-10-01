from __future__ import annotations

import copy

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.text.text import Font, _Paragraph
from pptx.util import Pt

from core.office_operations import TARGET_NOT_FOUND, Change
from core.office_result import INVALID_VALUE, OfficeFailure
from deck.pptx_content import text_frames_in
from deck.pptx_targets import PptxEditing, live_slides, require_text, resolve_paragraphs, resolve_shape, resolve_slide
from core.run_replacement import joined_text, replace_in_runs


PARAGRAPH_KEPT_TAGS = {qn("a:pPr"), qn("a:endParaRPr")}
BULLET_TAGS = (qn("a:buNone"), qn("a:buAutoNum"), qn("a:buChar"), qn("a:buBlip"))
BULLET_SUCCESSORS = ("a:tabLst", "a:defRPr", "a:extLst")
BULLET_CHARACTER = "•"
BULLET_INDENT_EMU = 285750
ALIGNMENTS = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
AUTOFIT = {"none": MSO_AUTO_SIZE.NONE, "shrink": MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE, "resize": MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT}
ANCHORS = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
RUN_STYLE_NAMES = ("font", "size", "bold", "italic", "underline", "color")
PARAGRAPH_STYLE_NAMES = ("align", "bullet", "level", "lineSpacing", "spaceBefore", "spaceAfter")


def plan_set_text(editing: PptxEditing, operation: dict, location: str) -> Change:
    target = resolve_shape(editing, operation, location)
    text_frame = require_text(target, location)
    resolve_paragraphs(text_frame, operation, target, location)

    def change() -> str:
        replace_text(text_frame, operation.get("paragraph"), operation["text"])
        editing.mark_edited(target.slide)
        scope = f"paragraph {operation['paragraph']} of " if operation.get("paragraph") is not None else ""
        return f"set the text of {scope}{target.label}"
    return change


def replace_text(text_frame, paragraph_index: int | None, text: str) -> None:
    paragraphs = list(text_frame.paragraphs)
    lines = text.split("\n")
    if paragraph_index is None:
        for paragraph in paragraphs[1:]:
            paragraph._p.getparent().remove(paragraph._p)
        replace_paragraph(paragraphs[0], lines)
        return
    replace_paragraph(paragraphs[paragraph_index], lines)


def replace_paragraph(paragraph, lines: list[str]) -> None:
    template = copy.deepcopy(paragraph._p)
    set_paragraph_line(paragraph, lines[0])
    anchor = paragraph._p
    for line in lines[1:]:
        clone = copy.deepcopy(template)
        anchor.addnext(clone)
        set_paragraph_line(_Paragraph(clone, paragraph._parent), line)
        anchor = clone


def set_paragraph_line(paragraph, line: str) -> None:
    runs = paragraph.runs
    if not runs:
        run = paragraph.add_run()
        inherit_end_formatting(paragraph._p, run._r)
        run.text = line
        return
    runs[0].text = line
    first_run = runs[0]._r
    for element in list(paragraph._p):
        if element is not first_run and element.tag not in PARAGRAPH_KEPT_TAGS:
            paragraph._p.remove(element)


def inherit_end_formatting(paragraph_element, run_element) -> None:
    end_properties = paragraph_element.find(qn("a:endParaRPr"))
    if end_properties is None:
        return
    run_properties = copy.deepcopy(end_properties)
    run_properties.tag = qn("a:rPr")
    run_element.insert(0, run_properties)


def plan_find_replace(editing: PptxEditing, operation: dict, location: str) -> Change:
    frames = scoped_text_frames(editing, operation, location)
    runs_by_paragraph = [paragraph.runs for frame in frames for paragraph in frame.paragraphs]
    occurrences = sum(joined_text(runs).count(operation["find"]) for runs in runs_by_paragraph)
    if occurrences == 0:
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: {operation['find']!r} does not occur", location, "copy the exact text from deck read; a match cannot span two paragraphs"))
    slides = scoped_slides(editing, operation, location)

    def change() -> str:
        replaced = sum(replace_in_runs(runs, operation["find"], operation["replace"]) for runs in runs_by_paragraph)
        for slide in slides:
            editing.mark_edited(slide)
        return f"replaced {replaced} occurrences of {operation['find']!r}"
    return change


def scoped_slides(editing: PptxEditing, operation: dict, location: str) -> list:
    if operation.get("slide") is not None:
        return [resolve_slide(editing, operation["slide"], f"{location}.slide")]
    if operation.get("shape") is not None:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.shape: a shape is only found within a slide", f"{location}.shape", "add the slide number, or leave out shape"))
    return live_slides(editing)


def scoped_text_frames(editing: PptxEditing, operation: dict, location: str) -> list:
    slides = scoped_slides(editing, operation, location)
    if operation.get("shape") is not None:
        target = resolve_shape(editing, operation, location)
        if target.kind in ("group", "table"):
            return text_frames_in([target.shape])
        return [require_text(target, location)]
    return [frame for slide in slides for frame in text_frames_in(slide.shapes)]


def plan_set_text_style(editing: PptxEditing, operation: dict, location: str) -> Change:
    require_any(operation, RUN_STYLE_NAMES, location)
    target = resolve_shape(editing, operation, location)
    paragraphs = resolve_paragraphs(require_text(target, location), operation, target, location)

    def change() -> str:
        for paragraph in paragraphs:
            for properties in character_properties(paragraph._p):
                apply_run_style(properties, operation)
        editing.mark_edited(target.slide)
        return f"restyled the text of {target.label}"
    return change


def require_any(operation: dict, names: tuple[str, ...], location: str) -> None:
    if any(operation.get(name) is not None for name in names):
        return
    raise OfficeFailure(INVALID_VALUE.issue(f"{location}: {operation['op']} changes nothing", location, f"give at least one of {', '.join(names)}"))


def character_properties(paragraph_element) -> list:
    properties = [run.get_or_add_rPr() for run in paragraph_element.iter(qn("a:r"))]
    properties.extend(field.get_or_add_rPr() for field in paragraph_element.iter(qn("a:fld")))
    end = paragraph_element.find(qn("a:endParaRPr"))
    if end is None:
        end = OxmlElement("a:endParaRPr")
        paragraph_element.append(end)
    properties.append(end)
    return properties


def apply_run_style(properties, operation: dict) -> None:
    font = Font(properties)
    if operation.get("size") is not None:
        font.size = Pt(operation["size"])
    for name in ("bold", "italic", "underline"):
        if operation.get(name) is not None:
            setattr(font, name, operation[name])
    if operation.get("color") is not None:
        font.color.rgb = RGBColor.from_string(operation["color"].lstrip("#").upper())
    if operation.get("font") is not None:
        set_typefaces(properties, operation["font"])


def set_typefaces(properties, typeface: str) -> None:
    Font(properties).name = typeface
    east_asian = properties.find(qn("a:ea"))
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        properties.find(qn("a:latin")).addnext(east_asian)
    east_asian.set("typeface", typeface)


def plan_set_paragraph(editing: PptxEditing, operation: dict, location: str) -> Change:
    require_any(operation, PARAGRAPH_STYLE_NAMES, location)
    target = resolve_shape(editing, operation, location)
    paragraphs = resolve_paragraphs(require_text(target, location), operation, target, location)

    def change() -> str:
        for paragraph in paragraphs:
            format_paragraph(paragraph, operation)
        editing.mark_edited(target.slide)
        return f"formatted the paragraphs of {target.label}"
    return change


def format_paragraph(paragraph, operation: dict) -> None:
    if operation.get("align") is not None:
        paragraph.alignment = ALIGNMENTS[operation["align"]]
    if operation.get("level") is not None:
        paragraph.level = operation["level"]
    if operation.get("lineSpacing") is not None:
        paragraph.line_spacing = float(operation["lineSpacing"])
    if operation.get("spaceBefore") is not None:
        paragraph.space_before = Pt(operation["spaceBefore"])
    if operation.get("spaceAfter") is not None:
        paragraph.space_after = Pt(operation["spaceAfter"])
    if operation.get("bullet") is not None:
        set_bullet(paragraph._p.get_or_add_pPr(), operation["bullet"])


def set_bullet(paragraph_properties, bullet: str) -> None:
    for existing in [child for child in paragraph_properties if child.tag in BULLET_TAGS]:
        paragraph_properties.remove(existing)
    if bullet == "none":
        paragraph_properties.insert_element_before(OxmlElement("a:buNone"), *BULLET_SUCCESSORS)
        return
    marker = OxmlElement("a:buChar" if bullet == "bullet" else "a:buAutoNum")
    marker.set(*(("char", BULLET_CHARACTER) if bullet == "bullet" else ("type", "arabicPeriod")))
    paragraph_properties.insert_element_before(marker, *BULLET_SUCCESSORS)
    if paragraph_properties.get("marL") is None:
        paragraph_properties.set("marL", str(BULLET_INDENT_EMU))
        paragraph_properties.set("indent", str(-BULLET_INDENT_EMU))


def plan_set_text_frame(editing: PptxEditing, operation: dict, location: str) -> Change:
    require_any(operation, ("autofit", "wrap", "anchor"), location)
    target = resolve_shape(editing, operation, location)
    text_frame = require_text(target, location)

    def change() -> str:
        if operation.get("autofit") is not None:
            text_frame.auto_size = AUTOFIT[operation["autofit"]]
        if operation.get("wrap") is not None:
            text_frame.word_wrap = operation["wrap"]
        if operation.get("anchor") is not None:
            text_frame.vertical_anchor = ANCHORS[operation["anchor"]]
        editing.mark_edited(target.slide)
        return f"set how the text of {target.label} fits"
    return change


def plan_set_notes(editing: PptxEditing, operation: dict, location: str) -> Change:
    slide = resolve_slide(editing, operation["slide"], f"{location}.slide")

    def change() -> str:
        slide.notes_slide.notes_text_frame.text = operation["text"]
        return f"set the notes of slide {operation['slide']}"
    return change


TEXT_PLANNERS = {
    "set_text": plan_set_text,
    "find_replace": plan_find_replace,
    "set_text_style": plan_set_text_style,
    "set_paragraph": plan_set_paragraph,
    "set_text_frame": plan_set_text_frame,
    "set_notes": plan_set_notes,
}
