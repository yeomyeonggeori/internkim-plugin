from __future__ import annotations


from fonts.docx_embedding import save_document
from core.office_result import MISSING_FIELD, OfficeFailure
from core.office_schema import require_valid
from paperwork.paperwork_definitions import CONTRACT_DOCUMENT, PAPERWORK_CONTENT_FIELDS, PAPERWORK_DOCUMENT
from paperwork.paperwork_design import COLOR_INK, FONT_KOREAN_DOCX, LINE_SPACING, SIZE_BODY, SIZE_CLAUSE_HEADING, SIZE_TITLE


SPEC_HINT = "read the form's spec at this skill's references/paperwork/<jurisdiction>/<form>.md and copy its Document JSON skeleton exactly"


def load_document(document: dict) -> dict:
    require_valid(PAPERWORK_DOCUMENT, document, "values")
    require_content(document)
    normalize_document(document)
    return document


def require_content(document):
    if any(field in document for field in PAPERWORK_CONTENT_FIELDS):
        return
    raise OfficeFailure(MISSING_FIELD.issue(f"values have no content: add at least one of {', '.join(PAPERWORK_CONTENT_FIELDS)}", "values", suggestion=SPEC_HINT))


def normalize_document(document):
    signature = document.get("signature")
    if isinstance(signature, str) and signature.strip():
        lines = [line.strip() for line in signature.splitlines() if line.strip()]
        if len(lines) >= 2:
            document["signature"] = {"date": lines[0], "line": " ".join(lines[1:])}
        else:
            document["signature"] = {"line": lines[0]}
    for name in ("lead", "notes"):
        lines = document.get(name)
        if isinstance(lines, str) and lines.strip():
            document[name] = [lines.strip()]


def load_contract_document(document: dict) -> dict:
    require_valid(CONTRACT_DOCUMENT, document, "values")
    return document


def generate_docx(document, output_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    word_document = Document()
    section = word_document.sections[0]
    margin_inches = (document.get("page") or {}).get("marginInches")
    margin = Inches(float(0.9 if margin_inches is None else margin_inches))
    section.top_margin = section.bottom_margin = margin
    section.left_margin = section.right_margin = margin
    font_name = text_or_default(document.get("fontName"), FONT_KOREAN_DOCX)
    font_size = float(SIZE_BODY if document.get("fontSize") is None else document["fontSize"])
    style = word_document.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(font_size)
    style.font.color.rgb = RGBColor(*COLOR_INK)
    style.paragraph_format.line_spacing = LINE_SPACING
    style.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font_name)
    title = text_or_default(document.get("title"), "").strip()
    if title:
        paragraph = word_document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(18)
        run = paragraph.add_run(title)
        run.bold = True
        run.font.size = Pt(SIZE_TITLE)
        run.font.color.rgb = RGBColor(*COLOR_INK)
    for block in document["blocks"]:
        append_docx_block(word_document, block)
    save_document(word_document, output_path)


def text_or_default(value, default):
    return default if value is None else str(value)


def append_docx_block(word_document, block):
    from docx.shared import Pt, RGBColor

    block_type = block["type"]
    if block_type == "heading":
        paragraph = word_document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(text_or_default(block.get("text"), ""))
        run.bold = True
        run.font.size = Pt(SIZE_CLAUSE_HEADING)
        run.font.color.rgb = RGBColor(*COLOR_INK)
    elif block_type == "paragraph":
        word_document.add_paragraph(text_or_default(block.get("text"), ""))
    elif block_type == "bullets":
        for item in block.get("items") or []:
            word_document.add_paragraph(str(item), style="List Bullet")
    elif block_type == "table":
        rows = block["rows"]
        table = word_document.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                if column_index < len(table.rows[row_index].cells):
                    table.rows[row_index].cells[column_index].text = "" if value is None else str(value)
