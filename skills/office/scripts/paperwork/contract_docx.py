from __future__ import annotations

from pathlib import Path

from fonts.docx_embedding import save_document
from paperwork.contract_plan import Block, ContractPlan
from paperwork.paperwork_design import (
    COLOR_INK,
    DOCX_PAGE_MARGIN_INCHES,
    FONT_KOREAN_DOCX,
    LINE_SPACING,
    SIZE_BODY,
    SIZE_CLAUSE_HEADING,
    SIZE_TITLE,
)

INK = COLOR_INK
BODY_FONT = FONT_KOREAN_DOCX


def styled_document():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(DOCX_PAGE_MARGIN_INCHES * 0.9)
    section.left_margin = section.right_margin = Inches(DOCX_PAGE_MARGIN_INCHES)
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(SIZE_BODY)
    normal.font.color.rgb = RGBColor(*INK)
    set_east_asia_font(normal.element, BODY_FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = LINE_SPACING
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return document


def set_east_asia_font(style_element, font_name):
    rPr = style_element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font_name)


def add_title(document, text):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(SIZE_TITLE)
    run.font.color.rgb = RGBColor(*INK)
    add_bottom_rule(paragraph)
    return paragraph


def add_bottom_rule(paragraph):
    from docx.oxml.ns import qn

    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pPr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single", qn("w:sz"): "12", qn("w:space"): "8", qn("w:color"): "".join(f"{channel:02X}" for channel in INK)})
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_centered(document, text):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def add_clause_heading(document, text):
    from docx.shared import Pt, RGBColor

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(SIZE_CLAUSE_HEADING)
    run.font.color.rgb = RGBColor(*INK)
    return paragraph


def add_body(document, text):
    return document.add_paragraph(text)


def add_numbered_item(document, text):
    paragraph = document.add_paragraph(text, style="List Number")
    numbering = document.part.numbering_part.element
    style_numbering = numbering.num_having_numId(paragraph.style.element.pPr.numPr.numId.val)
    restarted = numbering.add_num(style_numbering.abstractNumId.val)
    restarted.add_lvlOverride(ilvl=0).add_startOverride(1)
    paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = restarted.numId
    return paragraph


def add_signature_table(document, left_lines, right_lines):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    table = document.add_table(rows=1, cols=2)
    table.autofit = True
    for column_index, lines in ((0, left_lines), (1, right_lines)):
        cell = table.rows[0].cells[column_index]
        cell.text = ""
        for line_index, line in enumerate(lines):
            paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(line)
            if line_index == 0:
                run.bold = True
            paragraph.paragraph_format.space_after = Pt(2)
    return table


def write_contract(plan: ContractPlan, output_path: Path) -> None:
    document = styled_document()
    add_title(document, plan.title)
    add_blocks(document, plan.preamble)
    for heading, blocks in plan.articles:
        add_clause_heading(document, heading)
        add_blocks(document, blocks)
    add_blocks(document, plan.closing)
    if plan.signatures:
        add_signature_table(document, *plan.signatures)
    save_document(document, output_path)


def add_blocks(document, blocks: list[Block]) -> None:
    for block in blocks:
        if block.kind == "numbered":
            add_numbered_items(document, block.items)
        elif block.kind == "bullets":
            for item in block.items:
                add_body(document, f"- {item}")
        elif block.kind == "centered":
            add_centered(document, block.text)
        else:
            add_body(document, block.text)


def add_numbered_items(document, items: tuple[str, ...]) -> None:
    first = add_numbered_item(document, items[0])
    number_id = first._p.pPr.numPr.numId.val
    for item in items[1:]:
        paragraph = document.add_paragraph(item, style="List Number")
        paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = number_id

