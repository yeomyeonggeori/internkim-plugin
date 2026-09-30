#!/usr/bin/env python3
from __future__ import annotations

import sys
from pathlib import Path

from skill_runtime import ensure_requirements
from paperwork_design import (
    COLOR_INK,
    DOCX_PAGE_MARGIN_INCHES,
    FONT_KOREAN_DOCX,
    LINE_SPACING,
    SIZE_BODY,
    SIZE_CLAUSE_HEADING,
    SIZE_TITLE,
)

INK = COLOR_INK
BODY_FONT = FONT_KOREAN_DOCX


def styled_document():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(DOCX_PAGE_MARGIN_INCHES * 0.9)
    section.left_margin = section.right_margin = Inches(DOCX_PAGE_MARGIN_INCHES)
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(SIZE_BODY)
    normal.font.color.rgb = RGBColor(*INK)
    set_east_asia_font(normal.element, BODY_FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = LINE_SPACING
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return document


def set_east_asia_font(style_element, font_name):
    rPr = style_element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font_name)


def add_title(document, text):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(SIZE_TITLE)
    run.font.color.rgb = RGBColor(*INK)
    add_bottom_rule(paragraph)
    return paragraph


def add_bottom_rule(paragraph):
    from docx.oxml.ns import qn

    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pPr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single", qn("w:sz"): "12", qn("w:space"): "8", qn("w:color"): "1C2430"})
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_centered(document, text):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def add_clause_heading(document, text):
    from docx.shared import Pt, RGBColor

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(SIZE_CLAUSE_HEADING)
    run.font.color.rgb = RGBColor(*INK)
    return paragraph


def add_body(document, text):
    return document.add_paragraph(text)


def add_signature_table(document, left_lines, right_lines):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    table = document.add_table(rows=1, cols=2)
    table.autofit = True
    for column_index, lines in ((0, left_lines), (1, right_lines)):
        cell = table.rows[0].cells[column_index]
        cell.text = ""
        for line_index, line in enumerate(lines):
            paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(line)
            if line_index == 0:
                run.bold = True
            paragraph.paragraph_format.space_after = Pt(2)
    return table


def build_employment_contract(output_path):
    document = styled_document()
    add_title(document, "표 준 근 로 계 약 서")
    add_body(document, "{{ companyName }}(이하 \"사업주\"라 함)과(와) {{ employeeName }}(이하 \"근로자\"라 함)은 다음과 같이 근로계약을 체결한다.")
    add_clause_heading(document, "1. 근로개시일")
    add_body(document, "{% if endDate %}근로계약기간: {{ startDate }}부터 {{ endDate }}까지{% else %}{{ startDate }}부터 (기간의 정함이 없는 근로계약){% endif %}")
    add_clause_heading(document, "2. 근무장소")
    add_body(document, "{{ workplace }}")
    add_clause_heading(document, "3. 업무의 내용")
    add_body(document, "{{ duties }}")
    add_clause_heading(document, "4. 소정근로시간")
    add_body(document, "{{ workStartTime }}부터 {{ workEndTime }}까지 (휴게시간: {{ breakStart }} ~ {{ breakEnd }})")
    add_clause_heading(document, "5. 근무일 및 휴일")
    add_body(document, "매주 {{ workDays }} 근무, 주휴일 매주 {{ weeklyHoliday }}")
    add_clause_heading(document, "6. 임금")
    add_body(document, "- 월급: {{ monthlySalary }}")
    add_body(document, "- 상여금: {{ bonus }}")
    add_body(document, "- 기타급여(제수당 등): {{ otherAllowances }}")
    add_body(document, "- 임금지급일: {{ payday }} (휴일의 경우는 전일 지급)")
    add_body(document, "- 지급방법: {{ paymentMethod }}")
    add_clause_heading(document, "7. 연차유급휴가")
    add_body(document, "연차유급휴가는 근로기준법에서 정하는 바에 따라 부여함")
    add_clause_heading(document, "8. 사회보험 적용여부")
    add_body(document, "{{ insurances }}")
    add_clause_heading(document, "9. 근로계약서 교부")
    add_body(document, "사업주는 근로계약을 체결함과 동시에 본 계약서를 사본하여 근로자의 교부요구와 관계없이 근로자에게 교부함 (근로기준법 제17조 이행)")
    add_clause_heading(document, "10. 근로계약, 취업규칙 등의 성실한 이행의무")
    add_body(document, "사업주와 근로자는 각자가 근로계약, 취업규칙, 단체협약을 지키고 성실하게 이행하여야 함")
    add_clause_heading(document, "11. 기타")
    add_body(document, "이 계약에 정함이 없는 사항은 근로기준법령에 의함")
    add_body(document, "")
    add_centered(document, "{{ contractDate }}")
    add_body(document, "")
    add_signature_table(document,
        ["(사업주)", "사업체명: {{ companyName }} (전화: {{ companyPhone }})", "주소: {{ companyAddress }}", "대표자: {{ representative }} (서명)"],
        ["(근로자)", "주소: {{ employeeAddress }}", "연락처: {{ employeePhone }}", "성명: {{ employeeName }} (서명)"])
    document.save(str(output_path))


def build_service_agreement(output_path):
    document = styled_document()
    add_title(document, "용 역 계 약 서")
    add_body(document, "{{ clientName }}(이하 \"갑\"이라 한다)와 {{ providerName }}(이하 \"을\"이라 한다)은 {{ serviceName }}에 관하여 다음과 같이 용역계약을 체결한다.")
    add_clause_heading(document, "제1조 (계약의 목적)")
    add_body(document, "본 계약은 \"갑\"이 \"을\"에게 {{ serviceName }} 업무를 위탁하고 \"을\"이 이를 수행함에 필요한 제반사항과 \"갑\"과 \"을\" 간의 권리의무 및 협력사항을 정함으로써 상호 신뢰를 바탕으로 성공적인 용역 수행을 도모함에 그 목적이 있다.")
    add_clause_heading(document, "제2조 (용역의 대상 및 범위)")
    add_body(document, "\"을\"이 수행할 용역의 대상 및 범위는 다음 각 호와 같다.")
    add_body(document, "{%p for item in scopeItems %}")
    add_body(document, "{{ loop.index }}. {{ item }}")
    add_body(document, "{%p endfor %}")
    add_clause_heading(document, "제3조 (계약기간)")
    add_body(document, "본 계약의 용역 수행기간은 {{ startDate }}부터 {{ endDate }}까지로 한다.")
    add_clause_heading(document, "제4조 (계약금액 및 지급방법)")
    add_body(document, "① 본 용역의 계약금액은 금 {{ totalAmountKorean }}원整 (₩{{ totalAmount }}, 부가가치세 {{ vatNote }})으로 한다.")
    add_body(document, "② \"갑\"은 다음 각 호와 같이 \"을\"이 지정하는 계좌({{ bankAccount }})에 지급한다.")
    add_body(document, "{%p for payment in payments %}")
    add_body(document, "{{ loop.index }}. {{ payment }}")
    add_body(document, "{%p endfor %}")
    add_clause_heading(document, "제5조 (산출물 및 검수)")
    add_body(document, "① \"을\"은 용역기간 내에 다음 산출물을 \"갑\"에게 제출한다.")
    add_body(document, "{%p for deliverable in deliverables %}")
    add_body(document, "{{ loop.index }}. {{ deliverable }}")
    add_body(document, "{%p endfor %}")
    add_body(document, "② \"갑\"은 산출물을 제출받은 날로부터 14일 이내에 검수를 완료하고 그 결과를 서면으로 통지하며, 보완이 필요한 경우 \"을\"은 \"갑\"이 정한 기간 내에 보완하여 재제출한다.")
    add_clause_heading(document, "제6조 (계약당사자의 상호 의무)")
    add_body(document, "① \"갑\"과 \"을\"은 원활한 용역수행을 위하여 상호 협조하며, \"갑\"은 \"을\"이 용역수행에 필요한 제반자료나 업무요청을 하는 경우 적극 협조하여야 한다. ② \"을\"은 제2조에서 정한 용역의 대상 및 범위를 준수하여 성실히 용역을 수행한다.")
    add_clause_heading(document, "제7조 (보고 및 자료의 요청)")
    add_body(document, "\"갑\"은 \"을\"에게 용역수행의 경과나 그에 관련한 보고서 및 자료를 요청할 수 있으며, \"을\"은 요구를 받은 경우 지체 없이 이를 \"갑\"에게 제공하여야 한다.")
    add_clause_heading(document, "제8조 (지식재산권)")
    add_body(document, "① 본 용역의 수행 결과 발생하는 산출물에 대한 지식재산권은 잔금 지급 완료와 동시에 \"갑\"에게 귀속된다. ② 다만 \"을\"이 본 계약 이전부터 보유한 기술·노하우 및 범용 도구에 대한 권리는 \"을\"에게 유보된다.")
    add_clause_heading(document, "제9조 (비밀유지)")
    add_body(document, "\"갑\"과 \"을\"은 용역의 수행 및 본 계약의 이행과정에서 알게 된 상대방의 영업비밀을 상대방의 서면동의 없이 제3자에게 유출하거나 본 계약의 이행 이외의 목적으로 이용하여서는 아니 된다. 이 의무는 본 계약이 종료되거나 해제 또는 해지된 경우에도 3년간 존속한다.")
    add_clause_heading(document, "제10조 (지체상금)")
    add_body(document, "\"을\"이 약정 기한 내에 용역을 완성하지 못한 때에는 매 지체일수마다 계약금액의 1,000분의 {{ penaltyRate }}에 해당하는 지체상금을 \"갑\"에게 지급하거나 잔금에서 공제하며, 그 총액은 계약금액의 100분의 30을 초과하지 아니한다. 다만 불가항력 또는 \"갑\"의 귀책사유로 지체된 일수는 산입하지 아니한다.")
    add_clause_heading(document, "제11조 (하자보수)")
    add_body(document, "\"을\"은 검수 완료일로부터 {{ warrantyMonths }}개월간 산출물의 하자에 대하여 무상으로 보수할 책임을 지며, \"갑\"의 보수 요청을 받은 때에는 지체 없이 이에 응하여야 한다.")
    add_clause_heading(document, "제12조 (권리의무의 양도 금지)")
    add_body(document, "\"갑\" 또는 \"을\"은 상대방의 서면 동의 없이 본 계약상의 권리 또는 의무를 제3자에게 양도, 증여, 대물변제, 대여하거나 담보로 제공할 수 없다.")
    add_clause_heading(document, "제13조 (계약의 변경, 해제 및 해지)")
    add_body(document, "① 본 계약의 전부 또는 일부를 변경할 필요가 있는 경우 \"갑\"과 \"을\"은 서면합의로 변경할 수 있다. ② \"갑\" 또는 \"을\"은 상대방이 본 계약상의 의무를 위반하고 상당한 기간을 정한 시정 요구에도 이를 시정하지 아니한 경우 본 계약을 해제 또는 해지할 수 있으며, 그 통보는 서면으로 한다. ③ 해지 시 이미 수행된 기성부분에 대한 대가는 정산하여 지급한다.")
    add_clause_heading(document, "제14조 (손해배상)")
    add_body(document, "\"갑\" 또는 \"을\"이 본 계약을 위반하여 상대방에게 손해를 입힌 경우 그 손해를 배상하여야 한다. 다만 천재지변 등 불가항력으로 인한 손해에 대하여는 그러하지 아니하다.")
    add_clause_heading(document, "제15조 (분쟁의 해결)")
    add_body(document, "본 계약과 관련하여 분쟁이 발생한 경우 당사자의 상호 협의에 의한 해결을 모색하되, 합의가 이루어지지 아니한 경우에는 {{ jurisdiction }}을 합의관할로 하여 소송으로 해결한다.")
    add_body(document, "")
    add_body(document, "본 계약의 성립을 증명하기 위하여 계약서 2통을 작성하여 \"갑\"과 \"을\"이 서명 또는 기명날인한 후 각 1통씩 보관한다.")
    add_body(document, "")
    add_centered(document, "{{ contractDate }}")
    add_body(document, "")
    add_signature_table(document,
        ["(갑)", "상호: {{ clientName }}", "주소: {{ clientAddress }}", "대표자: {{ clientRepresentative }} (인)"],
        ["(을)", "상호: {{ providerName }}", "주소: {{ providerAddress }}", "대표자: {{ providerRepresentative }} (인)"])
    document.save(str(output_path))


def build_nda(output_path):
    document = styled_document()
    add_title(document, "비 밀 유 지 협 약 서")
    add_body(document, "{{ partyAName }}(이하 \"갑\"이라 한다)와 {{ partyBName }}(이하 \"을\"이라 한다)은 비밀정보의 제공과 관련하여 다음과 같이 비밀유지협약을 체결한다.")
    add_clause_heading(document, "제1조 (협약의 목적)")
    add_body(document, "본 협약은 \"갑\"과 \"을\"이 {{ purpose }}(이하 \"본 업무\"라 한다)와 관련하여 각자 상대방에게 제공하는 비밀정보를 비밀로 유지하고 보호하기 위하여 필요한 제반 사항을 규정함을 목적으로 한다.")
    add_clause_heading(document, "제2조 (비밀정보의 정의)")
    add_body(document, "① 본 협약에서 \"비밀정보\"라 함은 본 협약 체결 사실 자체 및 \"본 업무\"와 관련하여 어느 일방 당사자(이하 \"정보제공자\"라 한다)가 반대 당사자(이하 \"정보수령자\"라 한다)에게 제공하는 일체의 정보로서, 유·무형의 여부 및 그 기록 형태를 불문한다. ② 제1항의 비밀정보는 서면(전자문서를 포함하며, 이하 같음), 구두 혹은 기타 방법으로 제공되는 모든 노하우, 공정, 도면, 설계, 실험결과, 샘플, 사양, 데이터, 공식, 제법, 프로그램, 가격표, 거래명세서, 생산단가, 아이디어 등 모든 기술상 혹은 경영상의 정보와 그러한 정보가 수록된 물건 또는 장비 등을 모두 포함한다.")
    add_clause_heading(document, "제3조 (비밀의 표시)")
    add_body(document, "① 정보제공자가 정보수령자에게 서면 제출, 이메일 전송, 물품 인도 등 유형적 형태로 비밀정보를 제공하는 경우, 비밀임을 알리는 문구(\"비밀\" 또는 이와 유사한 표지)를 명확히 표시하여야 한다. ② 정보제공자가 정보수령자에게 구두나 영상 또는 당사자의 시설, 장비, 샘플 기타 품목들을 관찰·조사하게 하는 방법으로 비밀정보를 제공하는 경우에는, 그 즉시 정보수령자에게 해당 정보가 비밀정보에 속한다는 사실을 고지하고, 비밀정보 제공일로부터 30일 이내에 상대방에게 공개 범위, 공개 일자, 공개 장소 및 공개 대상자 등이 명시된 요약본을 서면 제출, 이메일 전송 등 유형적인 기록 형태로 제공하여야 한다. ③ 정보제공자가 비밀정보에 해당함에도 불구하고 제공 당시에 비밀정보임을 명확하게 표시하지 못하였거나 고지하지 못한 때에는, 정보제공자는 지체없이 정보수령자에 대하여 해당 정보가 비밀정보임을 고지함과 동시에 공개 범위, 공개 일자, 공개 장소 및 공개 대상자 등이 명시된 요약본을 서면 제출, 이메일 전송 등 유형적인 기록 형태로 제공하여야 하며, 이때부터 비밀정보로서 효력을 가진다.")
    add_clause_heading(document, "제4조 (비밀 유지 기간 등)")
    add_body(document, "① 본 협약은 본 협약 체결일로부터 {{ termYears }}년간 그 효력을 가진다. 다만 본 협약상 비밀유지의무는 협약기간의 만료 등의 사유로 본 협약이 종료된 이후에도 {{ survivalYears }}년간 효력을 가진다. ② 제1항에도 불구하고 본 협약에서 그 성질상 계속하여 효력을 유지하여야 하는 조항은 본 협약이 종료되거나 전항의 기간이 만료된 이후에도 계속하여 효력을 가진다.")
    add_clause_heading(document, "제5조 (정보의 사용용도 및 정보취급자 제한)")
    add_body(document, "① 정보수령자는 정보제공자의 사전 서면 승인이 없는 한 정보제공자의 비밀정보를 \"본 업무\"의 수행 또는 \"본 업무\"와 관련된 계약에서 정한 본래의 목적 및 용도로만 사용하여야 하며, \"본 업무\"와 관련하여 사용하는 경우에도 필요한 업무 수행의 범위를 초과하여 임의로 비밀정보를 복제, 수정, 저장, 변형 또는 분석할 수 없다. ② 정보수령자는 직접적·간접적으로 \"본 업무\"를 수행하는 임직원들에 한하여 정보제공자의 비밀정보를 취급할 수 있도록 필요한 조치를 취하여야 하며, 해당 임직원 각자에게 정보제공자의 비밀정보에 대한 비밀유지의무를 주지시켜야 한다. 이때 정보제공자는 정보수령자에게 해당 임직원으로부터 비밀유지서약서를 제출받는 등의 방법으로 해당 정보의 비밀성을 유지하기 위하여 필요한 조치를 요구할 수 있다. ③ 정보수령자가 \"본 업무\"의 수행을 위하여 정보제공자의 비밀정보를 \"본 업무\"를 수행하는 임직원들 이외의 제3자에게 제공하고자 할 때에는 사전에 정보제공자로부터 서면에 의한 동의를 얻어야 하며, 그 제3자와 사이에 해당 비밀정보의 유지 및 보호를 목적으로 하는 별도의 비밀유지협약을 체결한 이후에 그 제3자에게 해당 비밀정보를 제공하여야 한다.")
    add_clause_heading(document, "제6조 (비밀유지의무)")
    add_body(document, "① 정보수령자는 정보제공자의 사전 서면승낙 없이 비밀정보를 포함하여 본 협약의 내용, \"본 업무\"의 내용 등을 공표하거나 제3자에게 알려서는 아니 된다. 다만 객관적인 증거를 통하여 다음 각 호에 해당함이 입증되는 정보는 비밀정보가 아니거나 비밀유지의무가 없는 것으로 간주한다. 1. 비밀정보 제공 이전에 정보수령자가 이미 알고 있거나 알 수 있는 정보 2. 정보수령자의 고의 또는 과실에 의하지 않고 공지의 사실로 된 정보 3. 정보수령자가 정당하고 적법하게 제3자로부터 제공받은 정보 4. 정보수령자가 비밀정보를 이용하지 아니하고 독자적으로 개발하거나 알게 된 정보 5. 정보제공자가 비밀정보임을 고지하지 아니하고, 비밀정보에 속한다는 취지의 서면을 발송하지도 아니한 정보 6. 법원 기타 공공기관의 판결, 명령 또는 관련법령에 따른 공개의무에 따라서 공개한 정보 ② 정보수령자가 제1항 제6호에 따라 정보를 공개할 경우에는 사전에 정보제공자에게 그 사실을 서면으로 통지하고, 상대방으로 하여금 적절한 보호 및 대응조치를 할 수 있도록 하여야 한다. ③ 정보수령자는 비밀정보를 보호하고 관리하는 데에 필요한 모든 노력을 다하여야 한다. 다만 천재지변, 화재 등 불가항력 사유에 의해 비밀정보가 유출된 경우에는 유출에 대한 책임을 지지 아니한다.")
    add_clause_heading(document, "제7조 (손해배상, 위약벌)")
    add_body(document, "① 정보수령자가 본 협약을 위반한 경우, 정보수령자는 이로 인하여 정보제공자가 입은 모든 손해를 배상하는 것을 포함하여 법률상 배상 책임을 다하여야 한다.")
    add_body(document, "{% if penaltyAmount %}② 정보수령자가 본 협약을 위반한 경우, 정보수령자는 제1항의 손해배상과 별도로 정보제공자에게 위약벌로서 금 {{ penaltyAmount }}원 및 \"본 업무\" 관련 정보제공자·정보수령자 간 계약금액 중 큰 금원을 지급하여야 한다.{% endif %}")
    add_clause_heading(document, "제8조 (비밀정보의 반환 등)")
    add_body(document, "정보수령자는 협약기간의 만료 등의 사유로 본 협약이 종료된 경우, \"본 업무\"가 종료 또는 중단된 경우 또는 정보제공자의 요청이 있는 경우에는 지체없이 정보제공자의 비밀정보가 기재되어 있거나 이를 포함하고 있는 제반 자료, 장비, 서류, 샘플, 기타 유체물(복사본, 복사물, 모방물건, 모방장비 등을 포함)을 즉시 정보제공자에게 반환하거나, 정보제공자의 선택에 따라 이를 폐기하고 그 폐기를 증명하는 서류를 그때로부터 10일 내에 정보제공자에게 제공하여야 한다.")
    add_clause_heading(document, "제9조 (권리의 부존재 등)")
    add_body(document, "① 본 협약에 따라 제공되는 비밀정보에 관한 소유권, 지식재산권 등의 모든 권리는 정보제공자에 속하며, 비밀정보를 통하여 특허출원 등이 가능할 경우 특허 등을 출원할 권리는 정보제공자에게 있다. ② 본 협약은 어떠한 경우에도 정보수령자에게 비밀정보에 관한 어떠한 권리나 권리의 실시권 또는 사용권을 부여하는 것으로 해석되지 아니한다. ③ 본 협약은 어떠한 경우에도 당사자 간에 향후 어떠한 확정적인 협약의 체결, 제조물의 판매나 구입, 실시권의 허락 등을 암시하거나 이를 강제하지 아니하며, 기타 본 협약의 당사자가 비밀정보와 관련하여 다른 제3자와 어떠한 거래나 협약관계에 들어가는 것을 금지하거나 제한하지 아니한다. ④ 정보제공자는 비밀정보를 현 상태 그대로 제공하며, 비밀정보의 정확성 및 완전성이나 사업 목적에 대한 적합성 및 제3자의 권리 침해 여부에 대한 어떠한 보증도 하지 아니한다. ⑤ 정보제공자는 정보수령자가 비밀정보를 사용함에 따른 결과에 대하여 어떠한 책임도 지지 아니한다. ⑥ 각 당사자는 본 협약의 목적을 위하여 상대방의 시설을 방문하거나 이를 이용할 경우에는 상대방의 제반 규정 및 지시사항을 준수하여야 한다.")
    add_clause_heading(document, "제10조 (권리의무의 양도, 협약의 변경)")
    add_body(document, "① 각 당사자는 상대방의 사전 서면동의 없이 본 협약상의 권리의무를 각 당사자 이외의 제3자에게 양도하거나 이전할 수 없다. ② 본 협약의 수정이나 변경은 양 당사자의 정당한 대표자가 기명날인 또는 서명한 서면 합의로만 이루어질 수 있다.")
    add_clause_heading(document, "제11조 (협약의 분리가능성)")
    add_body(document, "본 협약 중 어느 규정이 법원에 의하여 위법, 무효 또는 집행 불가능하다고 선언될 경우에도, 이는 본 협약의 나머지 규정의 유효성에 영향을 미치지 아니한다.")
    add_clause_heading(document, "제12조 (분쟁의 해결)")
    add_body(document, "① 본 협약과 관련하여 분쟁이 발생한 경우 당사자의 상호 협의에 의한 해결을 모색하되, 분쟁에 관한 합의가 이루어지지 아니한 경우에는 중소기업기술 보호 지원에 관한 법률에 따른 중소기업기술분쟁조정·중재위원회의 조정에 따라 해결한다. ② 위 조정이 성립하지 아니한 경우 {{ jurisdiction }}을 제1심 관할법원으로 하여 소송을 통해 분쟁을 해결한다.")
    add_clause_heading(document, "제13조 (보칙)")
    add_body(document, "\"갑\"과 \"을\"은 본 협약의 성립을 증명하기 위하여 본 협약서 2부를 작성하여 각각 서명(또는 기명날인)한 후 각자 1부씩 보관한다.")
    add_body(document, "")
    add_centered(document, "{{ contractDate }}")
    add_body(document, "")
    add_signature_table(document,
        ["(갑)", "명칭: {{ partyAName }}", "주소: {{ partyAAddress }}", "대표자: {{ partyARepresentative }} (인)"],
        ["(을)", "명칭: {{ partyBName }}", "주소: {{ partyBAddress }}", "대표자: {{ partyBRepresentative }} (인)"])
    document.save(str(output_path))


def build_mou(output_path):
    document = styled_document()
    add_title(document, "업 무 협 약 서")
    add_body(document, "{{ orgAName }}과 {{ orgBName }}(이하 \"양 기관\"이라 한다)은 {{ purpose }}를 위하여 양 기관의 긴밀한 업무협력이 필요함을 깊이 인식하여 다음과 같이 협약을 체결한다.")
    add_clause_heading(document, "제1조 (목적)")
    add_body(document, "본 협약은 양 기관의 상호발전을 위하여 {{ orgAName }}와 {{ orgBName }} 간의 협력관계를 긴밀히 함에 그 목적이 있다.")
    add_clause_heading(document, "제2조 (협력분야)")
    add_body(document, "본 협약에 따른 협력분야는 다음과 같다.")
    add_body(document, "{%p for item in cooperationItems %}")
    add_body(document, "{{ loop.index }}. {{ item }}")
    add_body(document, "{%p endfor %}")
    add_body(document, "① {{ orgAName }}는 다음 각 호의 사항을 적극 이행한다.")
    add_body(document, "{%p for role in orgARoles %}")
    add_body(document, "{{ loop.index }}. {{ role }}")
    add_body(document, "{%p endfor %}")
    add_body(document, "② {{ orgBName }}는 다음 각 호의 사항을 적극 이행한다.")
    add_body(document, "{%p for role in orgBRoles %}")
    add_body(document, "{{ loop.index }}. {{ role }}")
    add_body(document, "{%p endfor %}")
    add_body(document, "③ 양 기관은 이 외에 기타 상호 협의에 따른 사항에 대하여 협력한다.")
    add_clause_heading(document, "제3조 (홍보)")
    add_body(document, "① 양 기관은 정보교환 또는 협력사항 등에 대한 보도자료를 작성·배포시 사전 통보한다. ② 양 기관은 상호 협력한 사항에 대한 보도자료를 배포할 경우 사전에 상대방 기관과 협의하여 작성·배포하고, \"양 기관 협력\" 문구를 반드시 명기한다.")
    add_clause_heading(document, "제4조 (실무협의회 구성 및 운영)")
    add_body(document, "① 본 협약서에 규정한 협력분야의 효율적 추진과 세부 업무의 상호 협의를 위하여 실무협의회를 구성한다. ② 실무협의회는 {{ orgAName }} 실무담당자와 {{ orgBName }} 실무담당자로 구성·운영하며, 양 기관간의 원활한 업무협조를 위해 필요한 경우 각 기관에 간사 1인을 추가로 둘 수 있다. ③ 실무협의회는 매분기 1회 정기 개최하며, 양 기관이 필요하다고 판단하는 경우 수시 개최할 수 있다.")
    add_clause_heading(document, "제5조 (비밀유지)")
    add_body(document, "① {{ orgAName }}와 {{ orgBName }}는 상호 업무협력을 수행함에 있어 취득한 정보에 대하여 비밀을 유지하여야 한다. ② {{ orgAName }}와 {{ orgBName }}의 비밀유지 의무는 본 협약이 종료된 이후에도 3년간 유지된다.")
    add_clause_heading(document, "제6조 (비용부담)")
    add_body(document, "업무협력을 위하여 소요되는 비용은 상호 협의를 통해 조정, 분담한다.")
    add_clause_heading(document, "제7조 (협의조정)")
    add_body(document, "본 협약서의 해석상 의견차가 있거나 추가 협의사항이 발생한 경우에는 제4조의 실무협의회를 통해 조정하며, 본 협약서에 언급하지 아니한 사항에 대해서는 양 기관이 별도로 협의하여 정한다.")
    add_clause_heading(document, "제8조 (협약의 효력)")
    add_body(document, "① 본 협약서의 효력은 협약 체결일로부터 발생하며, 유효기간은 {{ termYears }}년으로 한다. ② 본 협약서는 유효기간 만료 1개월 전까지 어느 일방으로부터 연장에 관한 서면 통보가 있을 경우 상호 합의하에 협약기간을 연장할 수 있다. ③ 양 기관은 서면 합의에 의하여 본 협약을 해지할 수 있으며, 어느 일방이 해지를 요청하는 경우 상대방과 사전 협의를 거쳐야 한다.")
    add_clause_heading(document, "제9조 (법적 구속력)")
    add_body(document, "이 협약은 법률적 구속력을 갖지 아니하며, 이 협약을 기반으로 파생되는 사업은 구체적인 조건과 기간을 명시하는 개별 계약에 의한다. 다만 제5조 비밀유지는 그러하지 아니하다.")
    add_body(document, "")
    add_body(document, "이 협약의 내용을 성실히 준행하고 협약을 증명하기 위하여 본 협약서 2부를 작성하여 양 기관의 대표자가 서명(날인)한 후 각 1부씩 보관한다.")
    add_body(document, "")
    add_centered(document, "{{ contractDate }}")
    add_body(document, "")
    add_signature_table(document,
        ["(기관A)", "기관명: {{ orgAName }}", "대표자: {{ orgARepresentative }} (서명 또는 인)"],
        ["(기관B)", "기관명: {{ orgBName }}", "대표자: {{ orgBRepresentative }} (서명 또는 인)"])
    document.save(str(output_path))


def build_offer_letter(output_path):
    document = styled_document()
    add_title(document, "채 용 제 안 서")
    add_body(document, "{{ candidateName }}님께, 안녕하세요. {{ companyName }}에서 함께 일하게 되어 기쁜 마음으로 아래와 같이 채용을 제안드립니다.")
    add_clause_heading(document, "포지션 요약")
    add_body(document, "직위: {{ position }}")
    add_body(document, "소속: {{ department }}")
    add_body(document, "근무지: {{ workplace }}")
    add_body(document, "입사예정일: {{ startDate }}")
    add_clause_heading(document, "보상")
    add_body(document, "연봉: {{ salary }}")
    add_body(document, "{% if equity %}스톡옵션: {{ equity }}{% endif %}")
    add_clause_heading(document, "복리후생")
    add_body(document, "{%p for benefit in benefits %}")
    add_body(document, "- {{ benefit }}")
    add_body(document, "{%p endfor %}")
    add_body(document, "{% if probationNote %}{{ probationNote }}{% endif %}")
    add_clause_heading(document, "제안 조건")
    add_body(document, "본 제안은 발송일로부터 {{ expiryDate }}까지 유효하며, 최종 입사는 배경확인 결과에 따라 확정됩니다.")
    add_body(document, "본 제안 내용에 대해 궁금한 점이 있으시면 언제든 회신 부탁드립니다.")
    add_body(document, "")
    add_body(document, "다시 한번 좋은 인연으로 함께하게 되기를 기대합니다.")
    add_body(document, "")
    add_centered(document, "{{ offerDate }}")
    add_body(document, "{{ companyName }}")
    add_body(document, "대표이사 {{ representative }}")
    document.save(str(output_path))


BUILDERS = {
    "employment-contract": build_employment_contract,
    "service-agreement": build_service_agreement,
    "nda": build_nda,
    "mou": build_mou,
    "offer-letter": build_offer_letter,
}


def main():
    if not ensure_requirements("office"):
        raise RuntimeError("paperwork dependencies are unavailable after bootstrap")
    templates_directory = Path(__file__).resolve().parents[2] / "assets" / "templates"
    templates_directory.mkdir(parents=True, exist_ok=True)
    requested = sys.argv[1:] or sorted(BUILDERS)
    for template_name in requested:
        builder = BUILDERS[template_name]
        output_path = templates_directory / f"{template_name}.docx"
        builder(output_path)
        print(output_path)


if __name__ == "__main__":
    main()
