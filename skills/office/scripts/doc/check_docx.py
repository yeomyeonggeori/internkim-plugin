#!/usr/bin/env python3
from __future__ import annotations

import re

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from doc_definitions import BROKEN_INTERNAL_REFERENCE, EAST_ASIA_FONT_MISSING, EAST_ASIA_LANGUAGE_NOT_KOREAN, STALE_TABLE_OF_CONTENTS, TRACKED_CHANGES_PRESENT
from docx_content_checks import chart_empty_issues, field_result_issues, heading_issues, missing_image_issues, unresolved_comment_issues
from docx_defaults import KOREAN_LANGUAGE
from docx_language import east_asia_font_issues, effective_east_asia_language
from docx_package import open_document
from docx_reference_operations import bookmark_names
from docx_revisions import collect_revisions, describe_pending
from docx_styles import run_styles
from docx_blocks import PARAGRAPH_TAG, body_block_elements, element_text, heading_level
from office_inputs import office_file
from office_result import Issue, OfficeArgumentParser, Result, run_command
from text_checks import PLACEHOLDER_LEFT, PLACEHOLDER_PATTERN, contains_korean


FIELD_REFERENCE_PATTERN = re.compile(r"^\s*(?:REF|PAGEREF|NOTEREF)\s+(\S+)", re.IGNORECASE)
MERGE_FIELD_PATTERN = re.compile(r"^\s*MERGEFIELD\s+(\S+)", re.IGNORECASE)
TABLE_OF_CONTENTS_PATTERN = re.compile(r"^\s*TOC\b", re.IGNORECASE)
CONTENTS_HEADING_DEPTH = 3


def main() -> Result:
    arguments = parse_arguments()
    document = open_document(arguments.document_path)
    elements = body_block_elements(document)
    issues = (
        placeholder_issues(elements)
        + part_placeholder_issues(document)
        + reference_issues(document, elements)
        + table_of_contents_issues(document, elements)
        + east_asia_font_issues(document)
        + east_asia_language_issues(document)
        + tracked_change_issues(document, elements)
        + missing_image_issues(document, elements)
        + chart_empty_issues(document, elements)
        + heading_issues(document, elements)
        + unresolved_comment_issues(document, elements)
        + field_result_issues(elements, fields_update_on_open(document))
    )
    return Result(summary=f"checked {arguments.document_path}: {len(issues)} issues", output_path=arguments.document_path, issues=tuple(issues))


def placeholder_issues(elements: list) -> list[Issue]:
    issues = []
    for index, element in enumerate(elements):
        issues.extend(placeholders_in(element, f"block {index}", index))
    return issues


def part_placeholder_issues(document) -> list[Issue]:
    issues = []
    for section_index, section in enumerate(document.sections):
        for part_name in ("header", "footer"):
            part = getattr(section, part_name)
            if part.is_linked_to_previous:
                continue
            issues.extend(placeholders_in(part._element, f"{part_name} of section {section_index}", None))
    return issues


def placeholders_in(element, location: str, block_index: int | None) -> list[Issue]:
    issues = [
        PLACEHOLDER_LEFT.issue(f"{location} still holds {placeholder}", location, suggestion=replace_suggestion(placeholder, block_index))
        for paragraph in element.iter(qn("w:p"))
        for placeholder in PLACEHOLDER_PATTERN.findall(element_text(paragraph))
    ]
    issues.extend(
        PLACEHOLDER_LEFT.issue(f"{location} holds merge field {field_name}", location, suggestion="replace the merge field with its value")
        for instruction in field_instructions(element)
        for field_name in MERGE_FIELD_PATTERN.findall(instruction)
    )
    return issues


def replace_suggestion(placeholder: str, block_index: int | None) -> dict:
    suggestion = {"op": "replace_text", "find": placeholder, "replace": "<value>"}
    if block_index is not None:
        suggestion["block"] = block_index
    return suggestion


def field_instructions(element) -> list[str]:
    simple = [field.get(qn("w:instr")) or "" for field in element.iter(qn("w:fldSimple"))]
    complex_fields = []
    current = None
    for node in element.iter(qn("w:fldChar"), qn("w:instrText")):
        if node.tag == qn("w:instrText"):
            if current is not None:
                current.append(node.text or "")
            continue
        character_type = node.get(qn("w:fldCharType"))
        if character_type == "begin":
            current = []
        elif character_type == "separate" and current is not None:
            complex_fields.append("".join(current))
            current = None
    return simple + complex_fields


def reference_issues(document, elements: list) -> list[Issue]:
    bookmarks = bookmark_names(document)
    issues = []
    for index, element in enumerate(elements):
        location = f"block {index}"
        for hyperlink in element.iter(qn("w:hyperlink")):
            anchor = hyperlink.get(qn("w:anchor"))
            if anchor and anchor not in bookmarks:
                issues.append(BROKEN_INTERNAL_REFERENCE.issue(f"{location} links to bookmark {anchor!r}, which does not exist", location))
        for instruction in field_instructions(element):
            for target in FIELD_REFERENCE_PATTERN.findall(instruction):
                if target not in bookmarks:
                    issues.append(BROKEN_INTERNAL_REFERENCE.issue(f"{location} cross-references bookmark {target!r}, which does not exist", location))
    return issues


def table_of_contents_issues(document, elements: list) -> list[Issue]:
    contents_blocks = [index for index, element in enumerate(elements) if any(TABLE_OF_CONTENTS_PATTERN.match(instruction) for instruction in field_instructions(element))]
    if not contents_blocks or fields_update_on_open(document):
        return []
    location = f"block {contents_blocks[0]}"
    contents_text = normalized(" ".join(element_text(elements[index]) for index in contents_blocks))
    headings = [normalized(element_text(element)) for element in elements if is_contents_heading(element, document)]
    missing = [heading for heading in headings if heading and heading not in contents_text]
    if not missing:
        return []
    return [STALE_TABLE_OF_CONTENTS.issue(f"the table of contents does not list {len(missing)} headings, such as {missing[0]!r}", location, suggestion={"op": "update_fields_on_open"})]


def fields_update_on_open(document) -> bool:
    update_fields = document.settings.element.find(qn("w:updateFields"))
    return update_fields is not None and update_fields.get(qn("w:val")) in ("true", "1", "on")


def is_contents_heading(element, document) -> bool:
    if element.tag != PARAGRAPH_TAG:
        return False
    level = heading_level(Paragraph(element, document._body))
    return level is not None and 1 <= level <= CONTENTS_HEADING_DEPTH


def normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def east_asia_language_issues(document) -> list[Issue]:
    runs = [run for run in document.element.body.iter(qn("w:r")) if contains_korean(element_text(run))]
    mistagged = [run for run in runs if effective_east_asia_language(run, document) != KOREAN_LANGUAGE]
    if not mistagged:
        return []
    return [EAST_ASIA_LANGUAGE_NOT_KOREAN.issue(f"{len(mistagged)} runs of Korean text have an East Asian language other than ko-KR", "document", suggestion={"op": "set_korean_language"})]


def tracked_change_issues(document, elements: list) -> list[Issue]:
    revisions = collect_revisions(document.element.body, elements)
    if not revisions:
        return []
    return [TRACKED_CHANGES_PRESENT.issue(f"the document holds tracked changes: {describe_pending(revisions)}", "document")]


def parse_arguments():
    parser = OfficeArgumentParser(description="Check a .docx for placeholders left, broken internal references, a stale table of contents, missing East Asian fonts, a wrong East Asian language, tracked changes, missing pictures, empty charts, empty or skipped headings, open comments and fields without a result. Issues suggest a doc apply operation where one fixes them.")
    parser.add_argument("document_path", type=office_file("docx"))
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
