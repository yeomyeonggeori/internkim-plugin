from __future__ import annotations

import re

from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from docx_blocks import heading_level
from docx_editing import DocxEditing, placement, resolve_paragraph
from docx_settings import insert_setting, request_field_update
from docx_parts import create_part, read_root, related_part, write_root
from docx_text import PARAGRAPH_TAG, live_runs, run_text, visible_text
from docx_tracking import mark_block_inserted, runs_between, split_runs_at
from office_operations import OPERATION_NOT_APPLICABLE, TARGET_NOT_FOUND, Change
from office_result import INVALID_VALUE, MISSING_FIELD, OfficeFailure
from office_theme import HYPERLINK_COLOR


BOOKMARK_NAME_PATTERN = re.compile(r"^[^\W\d_]\w{0,39}$")
CONTENTS_INDENT_INCHES = 0.25
NOTE_KINDS = {
    "footnote": ("footnotes", "/word/footnotes.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml", "FootnoteText", "FootnoteReference"),
    "endnote": ("endnotes", "/word/endnotes.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml", "EndnoteText", "EndnoteReference"),
}
NOTE_STYLE_NAMES = {"FootnoteText": "footnote text", "FootnoteReference": "footnote reference", "EndnoteText": "endnote text", "EndnoteReference": "endnote reference"}
CROSS_REFERENCE_INSTRUCTIONS = {"text": "REF {name} \\h", "page": "PAGEREF {name} \\h", "number": "REF {name} \\r \\h"}


def field_runs(instruction: str, result: str) -> list:
    runs = [field_character("begin"), instruction_run(instruction), field_character("separate")]
    result_run = OxmlElement("w:r")
    result_run.text = result
    return runs + [result_run, field_character("end")]


def field_character(character_type: str):
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:fldChar", attrs={qn("w:fldCharType"): character_type}))
    return run


def instruction_run(instruction: str):
    run = OxmlElement("w:r")
    text = OxmlElement("w:instrText", attrs={qn("xml:space"): "preserve"})
    text.text = f" {instruction} "
    run.append(text)
    return run


def plan_insert_table_of_contents(editing: DocxEditing, operation: dict, location: str) -> Change:
    levels = operation.get("levels") or 3
    place = placement(editing, operation, location)
    headings = [(level, visible_text(element)) for element in editing.elements if (level := contents_level(editing, element)) and level <= levels]

    def change() -> str:
        paragraphs = contents_paragraphs(editing, operation.get("title"), headings, levels)
        for paragraph in paragraphs:
            place(paragraph)
            if editing.tracking is not None:
                mark_block_inserted(paragraph, editing.tracking)
        request_field_update(editing.document)
        return f"inserted a table of contents listing {len(headings)} headings; Word refreshes its page numbers when the file opens"
    return change


def contents_level(editing: DocxEditing, element) -> int | None:
    if element.tag != PARAGRAPH_TAG:
        return None
    level = heading_level(Paragraph(element, editing.document._body))
    return level if level else None


def contents_paragraphs(editing: DocxEditing, title: str | None, headings: list, levels: int) -> list:
    paragraphs = []
    if title:
        style = "TOC Heading" if has_style(editing, "TOC Heading") else "Heading 1"
        paragraphs.append(editing.document.add_paragraph(title, style=style)._p)
    entries = headings or [(1, "")]
    for index, (level, text) in enumerate(entries):
        paragraph = editing.document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(CONTENTS_INDENT_INCHES * (level - 1))
        if index == 0:
            for run in field_runs(f'TOC \\o "1-{levels}" \\h \\z \\u', "")[:3]:
                paragraph._p.append(run)
        paragraph.add_run(text)
        if index == len(entries) - 1:
            paragraph._p.append(field_character("end"))
        paragraphs.append(paragraph._p)
    return paragraphs


def has_style(editing: DocxEditing, name: str) -> bool:
    return any(style.name == name for style in editing.document.styles)


def bookmark_names(document) -> set[str]:
    return {bookmark.get(qn("w:name")) for bookmark in document.element.body.iter(qn("w:bookmarkStart"))}


def plan_add_bookmark(editing: DocxEditing, operation: dict, location: str) -> Change:
    name = operation["name"]
    if not BOOKMARK_NAME_PATTERN.match(name):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.name: {name!r} is not a bookmark name", f"{location}.name", suggestion="start with a letter, then letters, digits or _, at most 40 characters"))
    if name in bookmark_names(editing.document):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.name: the document already has a bookmark {name!r}", f"{location}.name"))
    paragraph = resolve_paragraph(editing, operation["block"], f"{location}.block")
    identifier = str(max((int(mark.get(qn("w:id"))) for mark in editing.document.element.body.iter(qn("w:bookmarkStart"))), default=0) + 1)
    start = OxmlElement("w:bookmarkStart", attrs={qn("w:id"): identifier, qn("w:name"): name})
    properties = paragraph._p.find(qn("w:pPr"))
    if properties is None:
        paragraph._p.insert(0, start)
    else:
        properties.addnext(start)
    paragraph._p.append(OxmlElement("w:bookmarkEnd", attrs={qn("w:id"): identifier}))

    def change() -> str:
        return f"bookmarked block {operation['block']} as {name!r}"
    return change


def plan_insert_link(editing: DocxEditing, operation: dict, location: str) -> Change:
    if bool(operation.get("url")) == bool(operation.get("bookmark")):
        raise OfficeFailure(MISSING_FIELD.issue(f"{location}: give url or bookmark, not both", location))
    paragraph = resolve_paragraph(editing, operation["block"], f"{location}.block")
    runs = anchored_runs(paragraph._p, operation["find"], operation, location)
    if len({id(run.getparent()) for run in runs}) != 1:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.find: the text spans a tracked change or field, so it cannot become one link", f"{location}.find"))

    def change() -> str:
        hyperlink = OxmlElement("w:hyperlink")
        if operation.get("url"):
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(operation["url"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        else:
            hyperlink.set(qn("w:anchor"), operation["bookmark"])
        runs[0].addprevious(hyperlink)
        for run in runs:
            hyperlink.append(run)
            style_as_link(run)
        return f"linked {operation['find']!r} in block {operation['block']}"
    return change


def style_as_link(run) -> None:
    font = Run(run, None).font
    font.underline = True
    font.color.rgb = RGBColor.from_string(HYPERLINK_COLOR)


def anchored_runs(paragraph_element, find: str, operation: dict, location: str) -> list:
    text = "".join(run_text(run) for run in live_runs(paragraph_element))
    start = text.find(find)
    if start < 0:
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.find: {find!r} does not occur in block {operation['block']}", f"{location}.find", suggestion=f"copy the exact text from doc read: {text[:120]!r}"))
    split_runs_at(paragraph_element, (start, start + len(find)))
    return runs_between(paragraph_element, start, start + len(find))


def insertion_point(paragraph_element, find: str | None, operation: dict, location: str):
    if not find:
        return None
    return anchored_runs(paragraph_element, find, operation, location)[-1]


def place_runs(paragraph_element, anchor, runs: list) -> None:
    if anchor is None:
        for run in runs:
            paragraph_element.append(run)
        return
    for run in reversed(runs):
        anchor.addnext(run)


def plan_insert_cross_reference(editing: DocxEditing, operation: dict, location: str) -> Change:
    name = operation["bookmark"]
    targets = [bookmark for bookmark in editing.document.element.body.iter(qn("w:bookmarkStart")) if bookmark.get(qn("w:name")) == name]
    if not targets:
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.bookmark: the document has no bookmark {name!r}", f"{location}.bookmark", suggestion=f"add it with add_bookmark; existing: {', '.join(sorted(bookmark_names(editing.document))) or 'none'}"))
    paragraph = resolve_paragraph(editing, operation["block"], f"{location}.block")
    anchor = insertion_point(paragraph._p, operation.get("afterText"), operation, location)
    show = operation.get("show") or "text"

    def change() -> str:
        result = visible_text(targets[0].getparent()) if show == "text" else "?"
        place_runs(paragraph._p, anchor, field_runs(CROSS_REFERENCE_INSTRUCTIONS[show].format(name=name), result))
        request_field_update(editing.document)
        return f"inserted a cross-reference to {name!r} in block {operation['block']}"
    return change


def plan_insert_footnote(editing: DocxEditing, operation: dict, location: str) -> Change:
    return plan_insert_note(editing, operation, location, "footnote")


def plan_insert_endnote(editing: DocxEditing, operation: dict, location: str) -> Change:
    return plan_insert_note(editing, operation, location, "endnote")


def plan_insert_note(editing: DocxEditing, operation: dict, location: str, kind: str) -> Change:
    paragraph = resolve_paragraph(editing, operation["block"], f"{location}.block")
    anchor = insertion_point(paragraph._p, operation.get("afterText"), operation, location)

    def change() -> str:
        identifier = add_note(editing, kind, operation["text"])
        place_runs(paragraph._p, anchor, [note_reference_run(kind, identifier)])
        return f"added {kind} {identifier} to block {operation['block']}"
    return change


def add_note(editing: DocxEditing, kind: str, text: str) -> int:
    part_kind, part_name, content_type, text_style, reference_style = NOTE_KINDS[kind]
    ensure_note_styles(editing.document, text_style, reference_style)
    part = related_part(editing.document, getattr(RELATIONSHIP_TYPE, part_kind.upper())) or create_note_part(editing.document, kind)
    root = read_root(part)
    identifier = max((int(note.get(qn("w:id"))) for note in root.iter(qn(f"w:{kind}"))), default=0) + 1
    root.append(parse_xml(
        f'<w:{kind} {nsdecls("w")} w:id="{identifier}"><w:p><w:pPr><w:pStyle w:val="{text_style}"/></w:pPr>'
        f'<w:r><w:rPr><w:rStyle w:val="{reference_style}"/></w:rPr><w:{kind}Ref/></w:r>'
        f'<w:r><w:t xml:space="preserve"> {escape(text)}</w:t></w:r></w:p></w:{kind}>'
    ))
    write_root(part, root)
    return identifier


def escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def create_note_part(document, kind: str):
    part_kind, part_name, content_type, _, _ = NOTE_KINDS[kind]
    root = parse_xml(
        f'<w:{part_kind} {nsdecls("w")}>'
        f'<w:{kind} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{kind}>'
        f'<w:{kind} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{kind}>'
        f'</w:{part_kind}>'
    )
    declare_separators(document, kind)
    return create_part(document, getattr(RELATIONSHIP_TYPE, part_kind.upper()), part_name, content_type, root)


def declare_separators(document, kind: str) -> None:
    settings = document.settings.element
    if settings.find(qn(f"w:{kind}Pr")) is not None:
        return
    insert_setting(settings, parse_xml(f'<w:{kind}Pr {nsdecls("w")}><w:{kind} w:id="-1"/><w:{kind} w:id="0"/></w:{kind}Pr>'))


def ensure_note_styles(document, text_style: str, reference_style: str) -> None:
    styles = document.styles
    existing = {style.style_id for style in styles}
    if text_style not in existing:
        style = styles.add_style(NOTE_STYLE_NAMES[text_style], WD_STYLE_TYPE.PARAGRAPH)
        style.element.set(qn("w:styleId"), text_style)
        style.base_style = styles["Normal"]
        style.font.size = Pt(9)
        style.paragraph_format.space_after = Pt(0)
    if reference_style not in existing:
        style = styles.add_style(NOTE_STYLE_NAMES[reference_style], WD_STYLE_TYPE.CHARACTER)
        style.element.set(qn("w:styleId"), reference_style)
        style.font.superscript = True


def note_reference_run(kind: str, identifier: int):
    _, _, _, _, reference_style = NOTE_KINDS[kind]
    return parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rStyle w:val="{reference_style}"/></w:rPr><w:{kind}Reference w:id="{identifier}"/></w:r>')


def describe_notes(document) -> list[dict]:
    notes = []
    for kind, (part_kind, *_rest) in NOTE_KINDS.items():
        part = related_part(document, getattr(RELATIONSHIP_TYPE, part_kind.upper()))
        if part is None:
            continue
        for note in read_root(part).iter(qn(f"w:{kind}")):
            if note.get(qn("w:type")) is None:
                notes.append({"kind": kind, "id": int(note.get(qn("w:id"))), "text": visible_text(note).strip()})
    return notes
