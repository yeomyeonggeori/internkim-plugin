#!/usr/bin/env python3
from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from doc.doc_definitions import DOCUMENT_SPECIFICATION
from doc.docx_defaults import DOCUMENT_FONT, apply_korean_defaults, set_page, usable_width_inches
from fonts.docx_embedding import save_document
from doc.docx_tables import add_space_after_table, format_table
from doc.docx_lists import add_list_paragraph, start_list
from core.office_arguments import route_arguments
from core.office_result import INVALID_VALUE, Issue, OfficeFailure, Result, read_json_file, run_command
from core.office_schema import require_valid
from fonts.registry import BODY_SIZE_POINTS


DEFAULT_FONT_NAME = DOCUMENT_FONT
SIZED_STYLE_NAMES = ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]
LIST_BLOCK_TYPES = ("bullets", "numbered")


def main() -> Result:
    arguments = route_arguments("create", "docx")
    specification = read_specification(arguments.source)
    document = create_document(specification)
    output_path = Path(os.path.expanduser(arguments.output))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    save_document(document, output_path)
    return Result(summary=f"created {output_path}", output_path=str(output_path))


def read_specification(specification_path: str) -> dict:
    specification = read_json_file(specification_path)
    require_valid(DOCUMENT_SPECIFICATION, specification, "spec")
    require_rectangular_tables(specification["blocks"], "spec.blocks")
    return specification


def require_rectangular_tables(blocks: list[dict], location: str) -> None:
    problems = [
        problem
        for index, block in enumerate(blocks)
        if block["type"] == "table"
        for problem in table_width_problems(block["rows"], f"{location}[{index}].rows")
    ]
    if problems:
        raise OfficeFailure(*problems)


def table_width_problems(rows: list[list], location: str) -> list[Issue]:
    width = len(rows[0])
    if width == 0:
        return [INVALID_VALUE.issue(f"{location}[0]: table rows must contain at least one column", f"{location}[0]")]
    return [
        INVALID_VALUE.issue(f"{location}[{index}]: has {len(row)} cells but the header row has {width}", f"{location}[{index}]")
        for index, row in enumerate(rows)
        if len(row) != width
    ]


def create_document(specification: dict) -> Document:
    document = Document()
    set_document_page(document.sections[0], specification.get("page") or {})
    font_name = (specification.get("fontName") or DEFAULT_FONT_NAME).strip()
    set_text_sizes(document, float(specification.get("fontSize") or BODY_SIZE_POINTS))
    apply_korean_defaults(document, font_name)
    add_title(document, specification.get("title") or "")
    for block in specification["blocks"]:
        add_block(document, block)
    return document


def set_document_page(section, page: dict) -> None:
    margin_inches = page.get("marginInches")
    set_page(section, None if margin_inches is None else float(margin_inches), page.get("orientation") == "landscape")


def set_text_sizes(document: Document, font_size: float) -> None:
    for style_name in SIZED_STYLE_NAMES:
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        style.font.size = Pt(font_size if style_name == "Normal" else max(font_size + 1, 11))
        style.paragraph_format.line_spacing = 1.08
        style.paragraph_format.space_after = Pt(4)


def add_title(document: Document, title: str) -> None:
    if not title:
        return
    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_block(document: Document, block: dict) -> None:
    block_type = block["type"]
    if block_type == "heading":
        document.add_heading(block.get("text") or "", level=int(block.get("level") or 1))
    elif block_type == "paragraph":
        document.add_paragraph(block.get("text") or "")
    elif block_type in LIST_BLOCK_TYPES:
        add_list(document, block_type == "numbered", block.get("items") or [])
    elif block_type == "table":
        add_table(document, block)
    elif block_type == "pageBreak":
        document.add_page_break()


def add_list(document: Document, is_numbered: bool, items: list[str]) -> None:
    list_id = start_list(document, is_numbered)
    for item in items:
        add_list_paragraph(document, list_id, 0).add_run(item.strip())


def add_table(document: Document, block: dict) -> None:
    rows = block["rows"]
    width = len(rows[0])
    column_widths = read_column_widths(block, width, usable_width_inches(document.sections[-1]))
    table = document.add_table(rows=0, cols=width)
    table.style = block.get("style") or "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            fill_cell(cells[index], value, column_widths[index])
    format_table(table)
    add_space_after_table(document)


def fill_cell(cell, value: object, width_inches: float) -> None:
    cell.text = "" if value is None else str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell.width = Inches(width_inches)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)


def read_column_widths(block: dict, width: int, table_width: float) -> list[float]:
    column_widths = block.get("columnWidthsInches")
    if isinstance(column_widths, list) and len(column_widths) == width:
        return [float(value) for value in column_widths]
    return [table_width / width for _ in range(width)]


if __name__ == "__main__":
    raise SystemExit(run_command(main))
