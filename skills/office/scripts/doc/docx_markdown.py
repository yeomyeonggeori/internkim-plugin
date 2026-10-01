from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.image.exceptions import UnrecognizedImageError
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from doc_definitions import IMAGE_UNAVAILABLE
from docx_defaults import CODE_FONT, DOCUMENT_FONT, apply_korean_defaults, set_page
from docx_format_operations import ALIGNMENTS
from docx_tables import add_space_after_table, format_table
from docx_lists import add_list_paragraph, start_list
from docx_charts import add_chart_part, drawing_run, next_drawing_id, specification
from markdown_charts import Chart
from latex_math import OMML_NAMESPACE, LatexNotReadable, latex_omml
from markdown_blocks import Equation, Heading, Image, ListItem, Paragraph, Quote, Table, ThematicBreak, inline_segments, link_parts, local_image_problem, math_latex
from office_result import Issue


DEFAULT_DOCUMENT_FONT = DOCUMENT_FONT
DEFAULT_DOCUMENT_FONT_SIZE = 10.5
MAXIMUM_IMAGE_WIDTH = Inches(6)
LINK_COLOR = "0563C1"
CHART_HEIGHT_RATIO = 0.56
RULE_COLOR = "8C959F"
RULE_EIGHTHS_OF_A_POINT = "6"


def markdown_document(blocks: list, font_name: str, font_size: float, source_directory: Path) -> tuple[Document, list[Issue]]:
    document = Document()
    set_page(document.sections[0])
    set_base_font(document, font_name, font_size)
    issues = []
    list_ids: dict[bool, int] = {}
    for block in blocks:
        if not isinstance(block, ListItem):
            list_ids.clear()
        issues.extend(add_block(document, block, source_directory, list_ids))
    return document, issues


def set_base_font(document: Document, font_name: str, font_size: float) -> None:
    style = document.styles["Normal"]
    style.font.size = Pt(font_size)
    apply_korean_defaults(document, font_name)


def add_block(document: Document, block, source_directory: Path, list_ids: dict[bool, int]) -> list[Issue]:
    if isinstance(block, Heading):
        document.add_heading(block.text, level=block.level)
    elif isinstance(block, Table):
        add_table(document, block.rows, block.alignments)
    elif isinstance(block, ListItem):
        add_inline_runs(add_list_item(document, block, list_ids), block.text)
    elif isinstance(block, Quote):
        add_quote(document, block.text)
    elif isinstance(block, Image):
        return add_image(document, block, source_directory)
    elif isinstance(block, Chart):
        add_chart(document, block)
    elif isinstance(block, Paragraph):
        add_inline_runs(document.add_paragraph(), block.text)
    elif isinstance(block, ThematicBreak):
        add_rule(document)
    elif isinstance(block, Equation):
        add_display_equation(document, block)
    return []


def add_display_equation(document: Document, equation: Equation) -> None:
    try:
        formula = latex_omml(equation.latex, display=True)
    except LatexNotReadable:
        document.add_paragraph(equation.text)
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = ALIGNMENTS["center"]
    display = parse_xml(f'<m:oMathPara xmlns:m="{OMML_NAMESPACE}"/>')
    display.append(formula)
    paragraph._p.append(display)


def add_rule(document: Document) -> None:
    border = OxmlElement("w:bottom")
    for name, value in (("w:val", "single"), ("w:sz", RULE_EIGHTHS_OF_A_POINT), ("w:space", "1"), ("w:color", RULE_COLOR)):
        border.set(qn(name), value)
    borders = OxmlElement("w:pBdr")
    borders.append(border)
    document.add_paragraph()._p.get_or_add_pPr().append(borders)


def add_chart(document: Document, chart: Chart) -> None:
    section = document.sections[-1]
    width = int(section.page_width - section.left_margin - section.right_margin)
    relationship_id = add_chart_part(document, specification(chart.specification))
    drawing_id = next_drawing_id(document)
    document.add_paragraph()._p.append(drawing_run(relationship_id, width, int(width * CHART_HEIGHT_RATIO), drawing_id, f"Chart {drawing_id}"))


def add_list_item(document: Document, item: ListItem, list_ids: dict[bool, int]):
    if item.level == 0 and list_ids and item.is_numbered not in list_ids:
        list_ids.clear()
    if item.is_numbered not in list_ids:
        list_ids[item.is_numbered] = start_list(document, item.is_numbered)
    return add_list_paragraph(document, list_ids[item.is_numbered], item.level)


def add_table(document: Document, rows: list[list[str]], alignments: tuple[str, ...] = ()) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            paragraph = table.rows[row_index].cells[column_index].paragraphs[0]
            add_inline_runs(paragraph, row[column_index] if column_index < len(row) else "")
            alignment = alignments[column_index] if column_index < len(alignments) else ""
            if alignment:
                paragraph.alignment = ALIGNMENTS[alignment]
    format_table(table)
    add_space_after_table(document)


def add_quote(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(18)
    add_inline_runs(paragraph, text)
    for run in paragraph.runs:
        run.italic = True


def add_inline_runs(paragraph, text: str) -> None:
    for segment in inline_segments(text):
        link = link_parts(segment)
        latex = math_latex(segment)
        if latex is not None:
            add_inline_equation(paragraph, latex, segment)
        elif link:
            add_hyperlink(paragraph, *link)
        elif segment.startswith("**") and segment.endswith("**") and len(segment) > 4:
            paragraph.add_run(segment[2:-2]).bold = True
        elif segment.startswith("*") and segment.endswith("*") and len(segment) > 2:
            paragraph.add_run(segment[1:-1]).italic = True
        elif segment.startswith("`") and segment.endswith("`") and len(segment) > 2:
            run = paragraph.add_run(segment[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(segment)


def add_inline_equation(paragraph, latex: str, source: str) -> None:
    try:
        paragraph._p.append(latex_omml(latex))
    except LatexNotReadable:
        paragraph.add_run(source)


def add_hyperlink(paragraph, text: str, target: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    if target.startswith("#"):
        hyperlink.set(qn("w:anchor"), target[1:])
    else:
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor.from_string(LINK_COLOR)
    run.font.underline = True
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def add_image(document: Document, image: Image, source_directory: Path) -> list[Issue]:
    image_path = source_directory / image.source
    problem = local_image_problem(image.source, image_path)
    if problem:
        return [unavailable_image(document, image, problem)]
    try:
        picture = document.add_picture(str(image_path))
    except (UnrecognizedImageError, OSError) as error:
        return [unavailable_image(document, image, f"could not be read ({error})")]
    picture._inline.docPr.set("descr", image.alt)
    if picture.width > MAXIMUM_IMAGE_WIDTH:
        picture.height = int(picture.height * MAXIMUM_IMAGE_WIDTH / picture.width)
        picture.width = MAXIMUM_IMAGE_WIDTH
    return []


def unavailable_image(document: Document, image: Image, reason: str) -> Issue:
    document.add_paragraph().add_run(image.alt or image.source).italic = True
    return IMAGE_UNAVAILABLE.issue(f"image {image.source} {reason}; wrote its alt text instead", image.source)
