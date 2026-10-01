from __future__ import annotations

import copy
import os
import re

from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.image.exceptions import UnrecognizedImageError
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Twips
from docx.text.paragraph import Paragraph

from fonts.registry import OFFICE_KOREAN_FAMILY
from docx_drawing_operations import align_drawing, set_wrap
from docx_editing import DocxEditing, placement, resolve_block
from docx_format_operations import ALIGNMENTS
from docx_text import PARAGRAPH_TAG
from docx_tracking import mark_block_inserted
from office_operations import OPERATION_NOT_APPLICABLE, TARGET_NOT_FOUND, Change
from office_result import INPUT_NOT_FOUND, INVALID_VALUE, MISSING_FIELD, OfficeFailure
from page_sizes import PAPER_BY_NAME
from image_formats import PICTURE_FORMATS_TEXT


ORIENTATIONS = {"portrait": WD_ORIENT.PORTRAIT, "landscape": WD_ORIENT.LANDSCAPE}
SECTION_STARTS = {"nextPage": WD_SECTION_START.NEW_PAGE, "continuous": WD_SECTION_START.CONTINUOUS, "evenPage": WD_SECTION_START.EVEN_PAGE, "oddPage": WD_SECTION_START.ODD_PAGE}
MARGIN_FIELDS = {"marginTopInches": "top_margin", "marginBottomInches": "bottom_margin", "marginLeftInches": "left_margin", "marginRightInches": "right_margin"}
PAGE_SETUP_FIELDS = ("paper", "orientation", "marginInches", *MARGIN_FIELDS, "columns", "pageNumberStart")
PAGE_FIELD_PATTERN = re.compile(r"(\{PAGE\}|\{NUMPAGES\})")
PAGE_FIELD_INSTRUCTIONS = {"{PAGE}": "PAGE", "{NUMPAGES}": "NUMPAGES"}
WATERMARK_SHAPE_PREFIX = "PowerPlusWaterMarkObject"
PICTURE_WATERMARK_PREFIX = "WordPictureWatermark"
WASHOUT_ATTRIBUTES = ' gain="19661f" blacklevel="22938f"'
PICTURE_WATERMARK_TEMPLATE = """<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:w10="urn:schemas-microsoft-com:office:word"><w:pPr><w:pStyle w:val="Header"/></w:pPr><w:r><w:rPr><w:noProof/></w:rPr><w:pict><v:shapetype id="_x0000_t75" coordsize="21600,21600" o:spt="75" o:preferrelative="t" path="m@4@5l@4@11@9@11@9@5xe" filled="f" stroked="f"><v:stroke joinstyle="miter"/><v:formulas><v:f eqn="if lineDrawn pixelLineWidth 0"/><v:f eqn="sum @0 1 0"/><v:f eqn="sum 0 0 @1"/><v:f eqn="prod @2 1 2"/><v:f eqn="prod @3 21600 pixelWidth"/><v:f eqn="prod @3 21600 pixelHeight"/><v:f eqn="sum @0 0 1"/><v:f eqn="prod @6 1 2"/><v:f eqn="prod @7 21600 pixelWidth"/><v:f eqn="sum @8 21600 0"/><v:f eqn="prod @7 21600 pixelHeight"/><v:f eqn="sum @10 21600 0"/></v:formulas><v:path o:extrusionok="f" gradientshapeok="t" o:connecttype="rect"/><o:lock v:ext="edit" aspectratio="t"/></v:shapetype><v:shape id="{shape_id}" o:spid="_x0000_s{spid}" type="#_x0000_t75" style="position:absolute;margin-left:0;margin-top:0;width:{width}pt;height:{height}pt;z-index:-251656192;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin" o:allowincell="f"><v:imagedata r:id="{relationship_id}" o:title="{title}"{washout}/><w10:wrap anchorx="margin" anchory="margin"/></v:shape></w:pict></w:r></w:p>"""
WATERMARK_TEMPLATE = """<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:w10="urn:schemas-microsoft-com:office:word"><w:pPr><w:pStyle w:val="Header"/></w:pPr><w:r><w:rPr><w:noProof/></w:rPr><w:pict><v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e"><v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/><v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/><v:f eqn="if @0 @4 21600"/><v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/><v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas><v:path textpathok="t" o:connecttype="custom" o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/><v:textpath on="t" fitshape="t"/><v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/></v:handles><o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype><v:shape id="{shape_id}" o:spid="_x0000_s{spid}" type="#_x0000_t136" style="position:absolute;margin-left:0;margin-top:0;width:{width}pt;height:{height}pt;rotation:315;z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin" o:allowincell="f" fillcolor="#{color}" stroked="f"><v:fill opacity=".5"/><v:textpath style="font-family:&quot;{font}&quot;;font-size:1pt" string="{text}"/><w10:wrap anchorx="margin" anchory="margin"/></v:shape></w:pict></w:r></w:p>"""
WATERMARK_HEIGHT_POINTS = 120
WATERMARK_POINTS_PER_CHARACTER = 90
WATERMARK_MAXIMUM_WIDTH_POINTS = 480
DEFAULT_WATERMARK_COLOR = "C0C0C0"
SECTION_PROPERTIES_AFTER_PAGE_NUMBERING = tuple(f"w:{name}" for name in ("cols", "formProt", "vAlign", "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"))


def require_section(editing: DocxEditing, operation: dict, location: str) -> list:
    sections = list(editing.document.sections)
    if operation.get("section") is None:
        return sections
    if operation["section"] >= len(sections):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.section: the document has {len(sections)} sections", f"{location}.section", suggestion=f"use a section index from 0 to {len(sections) - 1}"))
    return [sections[operation["section"]]]


def plan_set_header(editing: DocxEditing, operation: dict, location: str) -> Change:
    return plan_header_or_footer(editing, operation, location, "header")


def plan_set_footer(editing: DocxEditing, operation: dict, location: str) -> Change:
    return plan_header_or_footer(editing, operation, location, "footer")


def plan_header_or_footer(editing: DocxEditing, operation: dict, location: str, part_name: str) -> Change:
    section = require_section(editing, {"section": operation.get("section") or 0}, location)[0]
    page = operation.get("page") or "default"

    def change() -> str:
        part = header_or_footer(editing, section, part_name, page)
        part.is_linked_to_previous = False
        for paragraph in part.paragraphs[1:]:
            paragraph._p.getparent().remove(paragraph._p)
        paragraph = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
        keep_watermarks(paragraph)
        write_with_page_fields(paragraph, operation["text"])
        if operation.get("align"):
            paragraph.alignment = ALIGNMENTS[operation["align"]]
        return f"set the {page} {part_name} of section {operation.get('section') or 0}"
    return change


def header_or_footer(editing: DocxEditing, section, part_name: str, page: str):
    if page == "first":
        section.different_first_page_header_footer = True
        return getattr(section, f"first_page_{part_name}")
    if page == "even":
        editing.document.settings.odd_and_even_pages_header_footer = True
        return getattr(section, f"even_page_{part_name}")
    return getattr(section, part_name)


def keep_watermarks(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr") and not is_watermark(child):
            paragraph._p.remove(child)


def write_with_page_fields(paragraph, text: str) -> None:
    for piece in PAGE_FIELD_PATTERN.split(text):
        if piece in PAGE_FIELD_INSTRUCTIONS:
            paragraph._p.append(page_field(PAGE_FIELD_INSTRUCTIONS[piece]))
        elif piece:
            paragraph.add_run(piece)


def page_field(instruction: str):
    field = OxmlElement("w:fldSimple", attrs={qn("w:instr"): f" {instruction} "})
    run = OxmlElement("w:r")
    run.text = "1"
    field.append(run)
    return field


def plan_set_page_setup(editing: DocxEditing, operation: dict, location: str) -> Change:
    if all(operation.get(name) is None for name in PAGE_SETUP_FIELDS):
        raise OfficeFailure(MISSING_FIELD.issue(f"{location}: give at least one of {', '.join(PAGE_SETUP_FIELDS)}", location))
    sections = require_section(editing, operation, location)

    def change() -> str:
        for section in sections:
            set_page(section, operation)
        return f"set up the page of {len(sections)} sections"
    return change


def set_page(section, operation: dict) -> None:
    orientation = operation.get("orientation")
    if operation.get("paper"):
        width, height = (Twips(length) for length in PAPER_BY_NAME[operation["paper"]].twips)
        section.page_width, section.page_height = width, height
        if (orientation or current_orientation(section)) == "landscape":
            section.page_width, section.page_height = height, width
    if orientation:
        turn(section, orientation)
    if operation.get("marginInches") is not None:
        for attribute in MARGIN_FIELDS.values():
            setattr(section, attribute, Inches(operation["marginInches"]))
    for field_name, attribute in MARGIN_FIELDS.items():
        if operation.get(field_name) is not None:
            setattr(section, attribute, Inches(operation[field_name]))
    if operation.get("columns"):
        set_columns(section, operation["columns"])
    if operation.get("pageNumberStart") is not None:
        set_page_number_start(section, operation["pageNumberStart"])


def current_orientation(section) -> str:
    return "landscape" if section.page_width and section.page_height and section.page_width > section.page_height else "portrait"


def turn(section, orientation: str) -> None:
    section.orientation = ORIENTATIONS[orientation]
    width, height = section.page_width, section.page_height
    if width is None or height is None:
        return
    if (orientation == "landscape") != (width > height):
        section.page_width, section.page_height = height, width


def set_columns(section, count: int) -> None:
    columns = section._sectPr.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        section._sectPr.append(columns)
    columns.set(qn("w:num"), str(count))
    columns.set(qn("w:space"), "425")


def set_page_number_start(section, start: int) -> None:
    numbering = section._sectPr.find(qn("w:pgNumType"))
    if numbering is None:
        numbering = OxmlElement("w:pgNumType")
        section._sectPr.insert_element_before(numbering, *SECTION_PROPERTIES_AFTER_PAGE_NUMBERING)
    numbering.set(qn("w:start"), str(start))


def plan_insert_section_break(editing: DocxEditing, operation: dict, location: str) -> Change:
    element = resolve_block(editing, operation["after"], f"{location}.after")
    if element.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.after: block {operation['after']} already ends a section", f"{location}.after"))
    start = SECTION_STARTS[operation.get("type") or "nextPage"]

    def change() -> str:
        carrier = element if element.tag == PARAGRAPH_TAG else empty_paragraph_after(element)
        following = governing_section_properties(editing, carrier)
        carrier.get_or_add_pPr().append(copy.deepcopy(following))
        section = next(section for section in editing.document.sections if section._sectPr is following)
        section.start_type = start
        if operation.get("orientation"):
            turn(section, operation["orientation"])
        return f"started a new section after block {operation['after']}"
    return change


def empty_paragraph_after(element):
    paragraph = OxmlElement("w:p")
    element.addnext(paragraph)
    return paragraph


def governing_section_properties(editing: DocxEditing, carrier):
    current = carrier.getnext()
    while current is not None:
        properties = current.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") if current.tag == PARAGRAPH_TAG else None
        if properties is not None:
            return properties
        if current.tag == qn("w:sectPr"):
            return current
        current = current.getnext()
    return editing.document.element.body.sectPr


def is_watermark(element) -> bool:
    return any((shape.get("id") or "").startswith((WATERMARK_SHAPE_PREFIX, PICTURE_WATERMARK_PREFIX)) for shape in element.iter("{urn:schemas-microsoft-com:vml}shape"))


def plan_set_watermark(editing: DocxEditing, operation: dict, location: str) -> Change:
    sections = require_section(editing, operation, location)
    image_path = require_watermark_source(operation, location)

    def change() -> str:
        for number, section in enumerate(sections, start=1):
            header = section.header
            if operation.get("section") is not None:
                header.is_linked_to_previous = False
            elif header.is_linked_to_previous and number > 1:
                continue
            header.is_linked_to_previous = False
            remove_watermarks(header._element)
            if image_path:
                header._element.append(picture_watermark_paragraph(header, section, operation, image_path, number, location))
            elif operation["text"]:
                header._element.append(watermark_paragraph(operation, number))
        if image_path:
            return f"set {os.path.basename(image_path)} as the watermark"
        return "removed the watermark" if not operation["text"] else f"set the watermark {operation['text']!r}"
    return change


def require_watermark_source(operation: dict, location: str) -> str | None:
    if (operation.get("text") is None) == (operation.get("image") is None):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: give exactly one of text and image", location, suggestion='text "" removes the watermark'))
    if operation.get("image") is None:
        return None
    return require_picture_file(operation["image"], f"{location}.image")


def require_picture_file(written_path: str, location: str) -> str:
    path = os.path.expanduser(written_path)
    if not os.path.isfile(path):
        raise OfficeFailure(INPUT_NOT_FOUND.issue(f"{location}: {written_path} does not exist", location, suggestion="pass an absolute path, or one relative to the directory doc apply runs in"))
    return path


def unreadable_picture(written_path: str, location: str) -> OfficeFailure:
    return OfficeFailure(INVALID_VALUE.issue(f"{location}: {written_path} is not a {PICTURE_FORMATS_TEXT} image", location))


def picture_watermark_paragraph(header, section, operation: dict, image_path: str, number: int, location: str):
    try:
        relationship_id, image = header.part.get_or_add_image(image_path)
    except UnrecognizedImageError as error:
        raise unreadable_picture(operation["image"], f"{location}.image") from error
    scale = operation["scale"] / 100 if operation.get("scale") else fitting_scale(section, image)
    return parse_xml(PICTURE_WATERMARK_TEMPLATE.format(
        shape_id=f"{PICTURE_WATERMARK_PREFIX}{number}",
        spid=2048 + number,
        width=round(Emu(image.width * scale).pt, 1),
        height=round(Emu(image.height * scale).pt, 1),
        relationship_id=relationship_id,
        title=xml_attribute(os.path.splitext(os.path.basename(image_path))[0]),
        washout=WASHOUT_ATTRIBUTES if operation.get("washout", True) else "",
    ))


def fitting_scale(section, image) -> float:
    text_width = section.page_width - section.left_margin - section.right_margin
    text_height = section.page_height - section.top_margin - section.bottom_margin
    return min(1, text_width / image.width, text_height / image.height)



def remove_watermarks(header_element) -> None:
    for paragraph in list(header_element.iter(PARAGRAPH_TAG)):
        for run in list(paragraph.iter(qn("w:r"))):
            if is_watermark(run):
                run.getparent().remove(run)


def watermark_paragraph(operation: dict, number: int):
    text = operation["text"]
    width = min(WATERMARK_MAXIMUM_WIDTH_POINTS, WATERMARK_POINTS_PER_CHARACTER * max(len(text), 2))
    return parse_xml(WATERMARK_TEMPLATE.format(
        shape_id=f"{WATERMARK_SHAPE_PREFIX}{number}",
        spid=2048 + number,
        width=width,
        height=WATERMARK_HEIGHT_POINTS,
        color=(operation.get("color") or DEFAULT_WATERMARK_COLOR).lstrip("#").upper(),
        font=OFFICE_KOREAN_FAMILY,
        text=xml_attribute(text),
    ))


def xml_attribute(text: str) -> str:
    return text.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;").replace(">", "&gt;")


def plan_insert_image(editing: DocxEditing, operation: dict, location: str) -> Change:
    path = require_picture_file(operation["path"], f"{location}.path")
    place = placement(editing, operation, location)

    def change() -> str:
        width = Inches(operation["widthInches"]) if operation.get("widthInches") else None
        height = Inches(operation["heightInches"]) if operation.get("heightInches") else None
        try:
            picture = editing.document.add_picture(path, width=width, height=height)
        except UnrecognizedImageError as error:
            raise unreadable_picture(operation["path"], f"{location}.path") from error
        fit_to_text_width(editing, picture, width, height)
        if operation.get("description"):
            picture._inline.docPr.set("descr", operation["description"])
        drawing = picture._inline.getparent()
        paragraph = drawing.getparent().getparent()
        if operation.get("wrap"):
            set_wrap(editing.document, drawing, operation["wrap"])
        if operation.get("align"):
            align_drawing(Paragraph(paragraph, None), drawing, operation["align"])
        place(paragraph)
        if editing.tracking is not None:
            mark_block_inserted(paragraph, editing.tracking)
        return f"inserted {os.path.basename(path)} at {Emu(picture.width).inches:.2f} x {Emu(picture.height).inches:.2f} inches"
    return change


def fit_to_text_width(editing: DocxEditing, picture, width, height) -> None:
    section = editing.document.sections[-1]
    text_width = section.page_width - section.left_margin - section.right_margin
    if width is not None or height is not None or picture.width <= text_width:
        return
    picture.height = int(picture.height * text_width / picture.width)
    picture.width = text_width
