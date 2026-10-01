from __future__ import annotations

import copy

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx_editing import DocxEditing, resolve_table
from docx_format_operations import ALIGNMENTS, require_any
from docx_text import PARAGRAPH_TAG, live_runs
from office_operations import OPERATION_NOT_APPLICABLE, TARGET_NOT_FOUND, Change
from office_result import INVALID_VALUE, OfficeFailure


VERTICAL_ALIGNMENTS = {"top": WD_CELL_VERTICAL_ALIGNMENT.TOP, "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER, "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}
CELL_PROPERTIES = ("fill", "bold", "align", "verticalAlign")
TABLE_CELL_SHADING_SUCCESSORS = ("noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange")


def require_unmerged(table, location: str) -> None:
    spans = [int(span.get(qn("w:val"), "1")) for span in table._tbl.iter(qn("w:gridSpan"))]
    if any(span > 1 for span in spans) or next(table._tbl.iter(qn("w:vMerge")), None) is not None:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: the table has merged cells, so its columns are not regular", location, suggestion="edit the cells with set_cell, or rebuild the table with insert_table"))


def column_count(table) -> int:
    return len(table._tbl.tblGrid.findall(qn("w:gridCol")))


def require_column(table, column: int, location: str) -> None:
    count = column_count(table)
    if column >= count:
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: the table has {count} columns", location, suggestion=f"use a column index from 0 to {count - 1}"))


def plan_insert_table_column(editing: DocxEditing, operation: dict, location: str) -> Change:
    table = resolve_table(editing, operation["block"], f"{location}.block")
    require_unmerged(table, location)
    after = column_to_follow(table, operation, location)
    cells = operation.get("cells") or []
    if len(cells) > len(table.rows):
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.cells: the table has {len(table.rows)} rows", f"{location}.cells"))

    def change() -> str:
        source = max(after, 0)
        total_width = grid_width(table)
        grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
        insert_copy(grid_columns[source], after)
        for row_index, row in enumerate(table._tbl.tr_lst):
            new_cell = insert_copy(row.tc_lst[source], after)
            value = cells[row_index] if row_index < len(cells) else ""
            set_cell_element_text(new_cell, "" if value is None else str(value))
        scale_widths(table, total_width)
        return f"inserted a column {'at the start' if after < 0 else f'after column {after}'} of block {operation['block']}"
    return change


def column_to_follow(table, operation: dict, location: str) -> int:
    if (operation.get("after") is None) == (operation.get("at") is None):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: give exactly one of after and at", location, suggestion='after takes a column index; at takes "start" or "end"'))
    if operation.get("at") == "start":
        return -1
    if operation.get("at") == "end":
        return column_count(table) - 1
    require_column(table, operation["after"], f"{location}.after")
    return operation["after"]


def insert_copy(element, after: int):
    duplicate = copy.deepcopy(element)
    if after < 0:
        element.addprevious(duplicate)
    else:
        element.addnext(duplicate)
    return duplicate


def set_cell_element_text(cell_element, text: str) -> None:
    paragraphs = cell_element.findall(PARAGRAPH_TAG)
    for paragraph in paragraphs[1:]:
        cell_element.remove(paragraph)
    first = paragraphs[0]
    runs = live_runs(first)
    for run in runs[1:]:
        run.getparent().remove(run)
    if runs:
        runs[0].text = text
        return
    run = OxmlElement("w:r")
    run.text = text
    first.append(run)


def plan_delete_table_column(editing: DocxEditing, operation: dict, location: str) -> Change:
    table = resolve_table(editing, operation["block"], f"{location}.block")
    require_unmerged(table, location)
    require_column(table, operation["column"], f"{location}.column")
    if column_count(table) == 1:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: the table's only column cannot be deleted; delete the block instead", location))

    def change() -> str:
        total_width = grid_width(table)
        grid_column = table._tbl.tblGrid.findall(qn("w:gridCol"))[operation["column"]]
        grid_column.getparent().remove(grid_column)
        for row in table._tbl.tr_lst:
            cell = row.tc_lst[operation["column"]]
            row.remove(cell)
        scale_widths(table, total_width)
        return f"deleted column {operation['column']} of block {operation['block']}"
    return change


def grid_width(table) -> int:
    return sum(int(column.get(qn("w:w"), "0")) for column in table._tbl.tblGrid.findall(qn("w:gridCol")))


def scale_widths(table, total_width: int) -> None:
    current = grid_width(table)
    if not total_width or not current:
        return
    factor = total_width / current
    for column in table._tbl.tblGrid.findall(qn("w:gridCol")):
        column.set(qn("w:w"), str(round(int(column.get(qn("w:w"), "0")) * factor)))
    for cell_width in table._tbl.iter(qn("w:tcW")):
        if cell_width.get(qn("w:type")) == "dxa":
            cell_width.set(qn("w:w"), str(round(int(cell_width.get(qn("w:w"), "0")) * factor)))


def cell_range(table, operation: dict, location: str) -> list:
    to_row = operation.get("toRow", operation["row"])
    to_column = operation.get("toColumn", operation["column"])
    if to_row is None:
        to_row = operation["row"]
    if to_column is None:
        to_column = operation["column"]
    if to_row < operation["row"] or to_column < operation["column"]:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: toRow and toColumn must not be before row and column", location))
    if to_row >= len(table.rows):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.toRow: the table has {len(table.rows)} rows", f"{location}.toRow", suggestion=f"use a row index from 0 to {len(table.rows) - 1}"))
    require_column(table, to_column, f"{location}.toColumn")
    return [(row, column) for row in range(operation["row"], to_row + 1) for column in range(operation["column"], to_column + 1)]


def plan_merge_cells(editing: DocxEditing, operation: dict, location: str) -> Change:
    table = resolve_table(editing, operation["block"], f"{location}.block")
    positions = cell_range(table, operation, location)
    if len(positions) < 2:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: give toRow or toColumn so the range covers more than one cell", location))
    first, last = positions[0], positions[-1]

    def change() -> str:
        table.cell(*first).merge(table.cell(*last))
        return f"merged cells ({first[0]}, {first[1]}) to ({last[0]}, {last[1]}) of block {operation['block']}"
    return change


def plan_format_cells(editing: DocxEditing, operation: dict, location: str) -> Change:
    require_any(operation, CELL_PROPERTIES, location)
    table = resolve_table(editing, operation["block"], f"{location}.block")
    positions = cell_range(table, operation, location)

    def change() -> str:
        cells = {id(cell._tc): cell for cell in (table.cell(row, column) for row, column in positions)}
        for cell in cells.values():
            format_cell(cell, operation)
        return f"formatted {len(cells)} cells of block {operation['block']}"
    return change


def format_cell(cell, operation: dict) -> None:
    if operation.get("fill"):
        set_cell_fill(cell, operation["fill"].lstrip("#").upper())
    if operation.get("verticalAlign"):
        cell.vertical_alignment = VERTICAL_ALIGNMENTS[operation["verticalAlign"]]
    for paragraph in cell.paragraphs:
        if operation.get("align"):
            paragraph.alignment = ALIGNMENTS[operation["align"]]
        if operation.get("bold") is not None:
            for run in paragraph.runs:
                run.bold = operation["bold"]


def set_cell_fill(cell, color: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        successor = next((child for child in cell_properties if child.tag in {qn(f"w:{name}") for name in TABLE_CELL_SHADING_SUCCESSORS}), None)
        if successor is None:
            cell_properties.append(shading)
        else:
            successor.addprevious(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
