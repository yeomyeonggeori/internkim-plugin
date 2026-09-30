#!/usr/bin/env python3
import argparse

from docx import Document

from create_docx import add_block, require_rectangular_tables
from doc_definitions import BLOCK_LIST
from documents_folder import resolve_document_path
from office_result import OfficeArgumentParser, Result, read_json_file, run_command
from office_schema import require_valid


def main() -> Result:
    arguments = parse_arguments()
    blocks = load_blocks(arguments.blocks) if arguments.blocks else []
    document_path = resolve_document_path(arguments.document_path, "docx")
    document = Document(document_path)
    for heading_text in arguments.heading:
        document.add_heading(heading_text, level=1)
    for paragraph_text in arguments.paragraph:
        document.add_paragraph(paragraph_text)
    for bullet_text in arguments.bullet:
        document.add_paragraph(bullet_text, style="List Bullet")
    for block in blocks:
        add_block(document, block)
    document.save(document_path)
    return Result(summary=f"appended to {document_path}", output_path=document_path)


def load_blocks(blocks_path: str) -> list[dict]:
    blocks = read_json_file(blocks_path)
    require_valid(BLOCK_LIST, blocks, "blocks")
    require_rectangular_tables(blocks, "blocks")
    return blocks


def parse_arguments() -> argparse.Namespace:
    parser = OfficeArgumentParser(description="Append content to an existing DOCX in place.")
    parser.add_argument("document_path", nargs="?", help="Path to the .docx; defaults to the newest .docx in ~/documents")
    parser.add_argument("--heading", action="append", default=[], metavar="TEXT", help="Append a level-1 heading")
    parser.add_argument("--paragraph", action="append", default=[], metavar="TEXT", help="Append a paragraph")
    parser.add_argument("--bullet", action="append", default=[], metavar="TEXT", help="Append a bullet item")
    parser.add_argument("--blocks", metavar="JSON_PATH", help="JSON file with an array of blocks; office guide doc describes them")
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
