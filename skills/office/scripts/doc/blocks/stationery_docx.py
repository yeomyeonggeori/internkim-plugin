from __future__ import annotations

from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from doc.blocks.stationery import LetterheadBlock, SignatureBlock
from doc.model.tables import TABLE_PROPERTIES_AFTER_BORDERS
from paperwork.paperwork_design import COLOR_MUTED, COLOR_RULE, SIZE_LETTERHEAD_DETAIL, SIZE_LETTERHEAD_NAME
from paperwork.stationery import Signer, hex_color

LOGO_HEIGHT_MILLIMETERS = 10
STAMP_HEIGHT_MILLIMETERS = 14
LETTERHEAD_RULE_EIGHTHS_OF_A_POINT = "11"
LETTERHEAD_RULE_SPACE_POINTS = "4"
LETTERHEAD_SPACE_AFTER_POINTS = 16
SIGNER_SIZE_POINTS = 12
SIGNATURE_SPACE_BEFORE_POINTS = {"center": 28, "left": 30}
PARAGRAPH_ALIGNMENTS = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}


def add_letterhead(document: Document, block: LetterheadBlock) -> None:
    letterhead = block.letterhead
    table = document.add_table(rows=1, cols=2)
    logo_cell, company_cell = table.rows[0].cells
    if letterhead.logo_path:
        logo_cell.paragraphs[0].add_run().add_picture(letterhead.logo_path, height=Mm(LOGO_HEIGHT_MILLIMETERS))
    company_lines = [(letterhead.name, SIZE_LETTERHEAD_NAME, True, None), *((line, SIZE_LETTERHEAD_DETAIL, False, COLOR_MUTED) for line in letterhead.details)]
    for index, (text, size, is_bold, color) in enumerate(company_lines):
        paragraph = company_cell.paragraphs[0] if index == 0 else company_cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        run.bold = is_bold
        if color:
            run.font.color.rgb = RGBColor(*color)
    rule_under(table)
    document.add_paragraph().paragraph_format.space_after = Pt(LETTERHEAD_SPACE_AFTER_POINTS)


def rule_under(table) -> None:
    border = OxmlElement("w:bottom")
    for name, value in (("w:val", "single"), ("w:sz", LETTERHEAD_RULE_EIGHTHS_OF_A_POINT), ("w:space", LETTERHEAD_RULE_SPACE_POINTS), ("w:color", hex_color(COLOR_RULE).lstrip("#"))):
        border.set(qn(name), value)
    borders = OxmlElement("w:tblBorders")
    borders.append(border)
    table._tbl.tblPr.insert_element_before(borders, *TABLE_PROPERTIES_AFTER_BORDERS)


def add_signature(document: Document, block: SignatureBlock) -> None:
    signer = block.signer
    alignment = PARAGRAPH_ALIGNMENTS[block.alignment]
    for index, line in enumerate(signer.lines):
        paragraph = document.add_paragraph()
        paragraph.alignment = alignment
        paragraph.paragraph_format.keep_with_next = index < len(signer.lines) - 1
        paragraph.paragraph_format.space_after = Pt(0)
        if index == 0:
            paragraph.paragraph_format.space_before = Pt(SIGNATURE_SPACE_BEFORE_POINTS[block.alignment])
        if index < len(signer.lines) - 1:
            paragraph.add_run(line)
            continue
        add_signer_line(paragraph, line, signer)


def add_signer_line(paragraph, line: str, signer: Signer) -> None:
    run = paragraph.add_run(line)
    if not signer.is_sealed:
        return
    run.bold = True
    run.font.size = Pt(SIGNER_SIZE_POINTS)
    if signer.stamp_path:
        paragraph.add_run(" ").add_picture(signer.stamp_path, height=Mm(STAMP_HEIGHT_MILLIMETERS))
    else:
        paragraph.add_run(f" {signer.seal_mark}").font.size = Pt(SIGNER_SIZE_POINTS)
