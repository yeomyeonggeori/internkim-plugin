#!/usr/bin/env python3
import argparse
import re

from docx import Document
from docx.oxml.ns import qn

from doc_definitions import (
    BODY_SIZE_UNUSUAL,
    DOCUMENT_EMPTY,
    DOCUMENT_SPARSE,
    LINE_SPACING_UNUSUAL,
    MARGIN_TOO_NARROW,
    MARGIN_TOO_WIDE,
    NO_STRUCTURED_TABLE,
    TABLE_DENSE_CELLS,
    TABLE_EMPTY_CELLS,
    TABLE_TOO_WIDE,
)
from office_result import Issue, OfficeArgumentParser, Result, run_command
from text_checks import korean_font_issues, text_presence_issues


DENSE_CELL_CHARACTERS = 90
MARGIN_SIDES = ["topMarginInches", "rightMarginInches", "bottomMarginInches", "leftMarginInches"]


def main() -> Result:
    arguments = parse_arguments()
    document = Document(arguments.document_path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    table_texts = [cell.text for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip()]
    visible_text = "\n".join(paragraphs + table_texts)
    tables = [table_metrics(table) for table in document.tables]
    typography = collect_typography(document, visible_text)
    issues = (
        paragraph_issues(paragraphs)
        + text_presence_issues(visible_text, arguments.required_text, arguments.forbidden_text)
        + korean_font_issues(visible_text, typography["fontNames"])
        + typography_issues(typography)
        + table_issues(tables)
    )
    details = {
        "paragraphCount": len(paragraphs),
        "headingCount": sum(1 for paragraph in document.paragraphs if paragraph.style and paragraph.style.name.startswith("Heading")),
        "tableCount": len(document.tables),
        "tables": tables,
        "firstParagraphs": paragraphs[:5],
        "visibleTextLength": len(visible_text),
        "typography": typography,
    }
    return Result(summary=f"checked {arguments.document_path}: {len(issues)} issues", output_path=arguments.document_path, issues=tuple(issues), details=details)


def table_metrics(table) -> dict:
    cells = [cell for row in table.rows for cell in row.cells]
    return {
        "rows": len(table.rows),
        "columns": len(table.columns),
        "emptyCellCount": sum(1 for cell in cells if not cell.text.strip()),
        "denseCellCount": sum(1 for cell in cells if len(cell.text.strip()) > DENSE_CELL_CHARACTERS),
    }


def collect_typography(document, visible_text: str) -> dict:
    normal_style = document.styles["Normal"] if "Normal" in document.styles else None
    paragraph_metrics = [metrics_for_paragraph(paragraph) for paragraph in document.paragraphs if paragraph.text.strip()]
    return {
        "hasKoreanText": bool(re.search(r"[가-힣]", visible_text)),
        "fontNames": sorted({font_name for font_name in collect_font_names(document) if font_name}),
        "fontSizePoints": sorted({size for size in collect_font_sizes(document, normal_style) if size is not None}),
        "normalFontSizePoints": point_value(normal_style.font.size) if normal_style is not None else None,
        "normalLineSpacing": line_spacing_value(normal_style.paragraph_format.line_spacing) if normal_style is not None else None,
        "paragraphLineSpacings": sorted({metric["lineSpacing"] for metric in paragraph_metrics if metric["lineSpacing"] is not None}),
        "sections": [section_metrics(section) for section in document.sections],
    }


def section_metrics(section) -> dict:
    return {
        "pageWidthInches": inch_value(section.page_width),
        "pageHeightInches": inch_value(section.page_height),
        "topMarginInches": inch_value(section.top_margin),
        "rightMarginInches": inch_value(section.right_margin),
        "bottomMarginInches": inch_value(section.bottom_margin),
        "leftMarginInches": inch_value(section.left_margin),
    }


def metrics_for_paragraph(paragraph) -> dict:
    return {
        "style": paragraph.style.name if paragraph.style else "",
        "lineSpacing": line_spacing_value(paragraph.paragraph_format.line_spacing),
        "spaceAfterPoints": point_value(paragraph.paragraph_format.space_after),
    }


def document_runs(document) -> list:
    body_runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
    cell_runs = [
        run
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
    ]
    return body_runs + cell_runs


def collect_font_names(document) -> list[str]:
    font_names = []
    for style in document.styles:
        font_names.extend(font_names_for_element(style.element))
        style_font = getattr(style, "font", None)
        if style_font is not None and style_font.name:
            font_names.append(style_font.name)
    font_names.extend(run.font.name for run in document_runs(document) if run.font.name)
    return font_names


def font_names_for_element(element) -> list[str]:
    run_properties = element.rPr
    if run_properties is None or run_properties.rFonts is None:
        return []
    values = [run_properties.rFonts.get(qn(attribute_name)) for attribute_name in ["w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"]]
    return [value for value in values if value]


def collect_font_sizes(document, normal_style) -> list[float | None]:
    sizes = []
    if normal_style is not None and getattr(normal_style, "font", None) is not None:
        sizes.append(point_value(normal_style.font.size))
    for paragraph in document.paragraphs:
        if paragraph.style is not None and getattr(paragraph.style, "font", None) is not None:
            sizes.append(point_value(paragraph.style.font.size))
    sizes.extend(point_value(run.font.size) for run in document_runs(document))
    return sizes


def point_value(value) -> float | None:
    if value is None:
        return None
    return round(value.pt, 2)


def inch_value(value) -> float | None:
    if value is None:
        return None
    return round(value.inches, 3)


def line_spacing_value(value) -> float | None:
    if value is None:
        return None
    if hasattr(value, "pt"):
        return round(value.pt, 2)
    return round(float(value), 2)


def paragraph_issues(paragraphs: list[str]) -> list[Issue]:
    if not paragraphs:
        return [DOCUMENT_EMPTY.issue("document has no visible paragraph text"), DOCUMENT_SPARSE.issue("document has very little visible text")]
    if len(paragraphs) < 3:
        return [DOCUMENT_SPARSE.issue("document has very little visible text")]
    return []


def typography_issues(typography: dict) -> list[Issue]:
    issues = []
    normal_size = typography["normalFontSizePoints"]
    if normal_size is not None and not 9.0 <= normal_size <= 12.5:
        issues.append(BODY_SIZE_UNUSUAL.issue(f"normal text size is {normal_size} pt; expected about 10-11 pt for business documents"))
    normal_spacing = typography["normalLineSpacing"]
    if normal_spacing is not None and not 1.0 <= normal_spacing <= 1.25:
        issues.append(LINE_SPACING_UNUSUAL.issue(f"normal line spacing is {normal_spacing}; expected about 1.05-1.2"))
    for index, section in enumerate(typography["sections"], start=1):
        issues.extend(margin_issues(section, f"section {index}"))
    return issues


def margin_issues(section: dict, location: str) -> list[Issue]:
    margins = [section[side] for side in MARGIN_SIDES if section[side] is not None]
    issues = []
    if any(value < 0.5 for value in margins):
        issues.append(MARGIN_TOO_NARROW.issue(f"{location} has a margin under 0.5 inches", location))
    if any(value > 1.25 for value in margins):
        issues.append(MARGIN_TOO_WIDE.issue(f"{location} has a margin above 1.25 inches", location))
    return issues


def table_issues(tables: list[dict]) -> list[Issue]:
    issues = [] if any(table["rows"] > 1 for table in tables) else [NO_STRUCTURED_TABLE.issue("document has no multi-row table for structured facts")]
    for index, table in enumerate(tables, start=1):
        location = f"table {index}"
        if table["columns"] > 5:
            issues.append(TABLE_TOO_WIDE.issue(f"{location} has more than 5 columns and may be too wide", location))
        if table["emptyCellCount"] > 0:
            issues.append(TABLE_EMPTY_CELLS.issue(f"{location} has {table['emptyCellCount']} empty cells", location))
        if table["denseCellCount"] > 0:
            issues.append(TABLE_DENSE_CELLS.issue(f"{location} has {table['denseCellCount']} dense cells that may need wrapping or shorter labels", location))
    return issues


def parse_arguments() -> argparse.Namespace:
    parser = OfficeArgumentParser(description="Validate and summarize a DOCX file.")
    parser.add_argument("document_path")
    parser.add_argument("--required-text", action="append", default=[])
    parser.add_argument("--forbidden-text", action="append", default=[])
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
