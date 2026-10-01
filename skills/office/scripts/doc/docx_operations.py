from __future__ import annotations

import copy

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from doc_definitions import OPERATIONS
from docx_language import make_east_asia_language_korean
from docx_format_operations import plan_define_style, plan_format_text, plan_set_paragraph_format
from docx_page_operations import plan_insert_image, plan_insert_section_break, plan_set_footer, plan_set_header, plan_set_page_setup, plan_set_watermark
from docx_reference_operations import (
    plan_add_bookmark, plan_insert_cross_reference, plan_insert_endnote, plan_insert_footnote, plan_insert_link, plan_insert_table_of_contents,
)
from docx_table_operations import plan_delete_table_column, plan_format_cells, plan_insert_table_column, plan_merge_cells
from docx_comments import plan_add_comment, plan_delete_comment, plan_reply_comment, plan_resolve_comment
from docx_blocks import PARAGRAPH_TAG, TABLE_TAG, paragraph_runs
from docx_settings import request_field_update
from docx_editing import DocxEditing, load_editing, placement, require_style, resolve_block, resolve_paragraph, resolve_table, save_editing
from docx_revision_operations import plan_accept_revisions, plan_reject_revisions
from docx_text import REMOVED_RUN_CONTAINER_TAGS
from docx_tracking import (
    mark_block_deleted, mark_block_inserted, mark_paragraph_deleted, mark_row, record_paragraph_property_change,
    snapshot_paragraph_properties, tracked_replace, tracked_set_text,
)
from office_operations import OPERATION_NOT_APPLICABLE, TARGET_NOT_FOUND, Change, OperationSet
from office_result import INVALID_VALUE, OfficeFailure
from run_replacement import joined_text, replace_in_runs


RUN_WRAPPER_TAGS = tuple(qn(f"w:{name}") for name in ("hyperlink", "ins", "moveTo", "smartTag", "customXml"))


def plan_replace_text(editing: DocxEditing, operation: dict, location: str) -> Change:
    paragraphs = scoped_paragraphs(editing, operation, location)
    occurrences = sum(joined_run_text(paragraph).count(operation["find"]) for paragraph in paragraphs)
    if occurrences == 0:
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}: {operation['find']!r} does not occur", location))

    def change() -> str:
        if editing.tracking is not None:
            replaced = sum(tracked_replace_in_paragraph(paragraph, operation["find"], operation["replace"], editing.tracking) for paragraph in paragraphs)
            return f"replaced {replaced} occurrences of {operation['find']!r} as tracked changes"
        replaced = sum(replace_in_paragraph(paragraph, operation["find"], operation["replace"]) for paragraph in paragraphs)
        return f"replaced {replaced} occurrences of {operation['find']!r}"
    return change


def scoped_paragraphs(editing: DocxEditing, operation: dict, location: str) -> list[Paragraph]:
    elements = [resolve_block(editing, operation["block"], f"{location}.block")] if operation.get("block") is not None else editing.elements
    return [Paragraph(paragraph_element, editing.document._body) for element in elements for paragraph_element in paragraph_elements_within(element)]


def paragraph_elements_within(element) -> list:
    if element.tag == PARAGRAPH_TAG:
        return [element]
    return list(element.iter(PARAGRAPH_TAG))


def joined_run_text(paragraph: Paragraph) -> str:
    return joined_text(paragraph_runs(paragraph))


def replace_in_paragraph(paragraph: Paragraph, find: str, replace: str) -> int:
    return replace_in_runs(paragraph_runs(paragraph), find, replace)


def tracked_replace_in_paragraph(paragraph: Paragraph, find: str, replace: str, tracking) -> int:
    text = joined_run_text(paragraph)
    positions = []
    position = text.find(find)
    while position >= 0:
        positions.append(position)
        position = text.find(find, position + len(find))
    for start in reversed(positions):
        tracked_replace(paragraph._p, start, start + len(find), replace, tracking)
    return len(positions)


def plan_set_text(editing: DocxEditing, operation: dict, location: str) -> Change:
    paragraph = resolve_paragraph(editing, operation["block"], f"{location}.block")

    def change() -> str:
        if editing.tracking is not None:
            tracked_set_text(paragraph._p, operation["text"], editing.tracking)
            return f"set the text of block {operation['block']} as tracked changes"
        set_paragraph_text(paragraph, operation["text"])
        return f"set the text of block {operation['block']}"
    return change


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    for removed in list(paragraph._p.iter(*REMOVED_RUN_CONTAINER_TAGS)):
        removed.getparent().remove(removed)
    runs = paragraph_runs(paragraph)
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)
    remove_empty_run_wrappers(paragraph._p)


def remove_empty_run_wrappers(paragraph_element) -> None:
    for wrapper in list(paragraph_element.iter(*RUN_WRAPPER_TAGS)):
        if wrapper.find(f".//{qn('w:r')}") is None and wrapper.getparent() is not None:
            wrapper.getparent().remove(wrapper)


def place_new_block(editing: DocxEditing, place, element) -> str:
    place(element)
    if editing.tracking is None:
        return ""
    mark_block_inserted(element, editing.tracking)
    return " as a tracked change"


def plan_insert_paragraph(editing: DocxEditing, operation: dict, location: str) -> Change:
    style = operation.get("style") or "Normal"
    require_style(editing, style, (WD_STYLE_TYPE.PARAGRAPH,), f"{location}.style")
    place = placement(editing, operation, location)

    def change() -> str:
        tracked = place_new_block(editing, place, editing.document.add_paragraph(operation["text"], style=style)._p)
        return f"inserted a {style} paragraph{tracked}"
    return change


def plan_insert_heading(editing: DocxEditing, operation: dict, location: str) -> Change:
    level = operation.get("level") or 1
    require_style(editing, f"Heading {level}", (WD_STYLE_TYPE.PARAGRAPH,), f"{location}.level")
    place = placement(editing, operation, location)

    def change() -> str:
        tracked = place_new_block(editing, place, editing.document.add_heading(operation["text"], level=level)._p)
        return f"inserted a level {level} heading{tracked}"
    return change


def plan_insert_list(editing: DocxEditing, operation: dict, location: str) -> Change:
    style = "List Number" if operation.get("numbered") else "List Bullet"
    require_style(editing, style, (WD_STYLE_TYPE.PARAGRAPH,), location)
    place = placement(editing, operation, location)

    def change() -> str:
        for item in operation["items"]:
            tracked = place_new_block(editing, place, editing.document.add_paragraph(item, style=style)._p)
        return f"inserted {len(operation['items'])} list items{tracked}"
    return change


def plan_insert_table(editing: DocxEditing, operation: dict, location: str) -> Change:
    rows = operation["rows"]
    width = len(rows[0])
    ragged = [index for index, row in enumerate(rows) if len(row) != width]
    if width == 0 or ragged:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.rows: every row needs the header row's {width} cells", f"{location}.rows"))
    style = operation.get("style") or "Table Grid"
    require_style(editing, style, (WD_STYLE_TYPE.TABLE,), f"{location}.style")
    place = placement(editing, operation, location)

    def change() -> str:
        table = editing.document.add_table(rows=len(rows), cols=width)
        table.style = style
        for row_cells, values in zip(table.rows, rows):
            for cell, value in zip(row_cells.cells, values):
                cell.text = "" if value is None else str(value)
        tracked = place_new_block(editing, place, table._tbl)
        return f"inserted a {len(rows)}x{width} table{tracked}"
    return change


def plan_insert_page_break(editing: DocxEditing, operation: dict, location: str) -> Change:
    place = placement(editing, operation, location)

    def change() -> str:
        paragraph = editing.document.add_paragraph()
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        return "inserted a page break" + place_new_block(editing, place, paragraph._p)
    return change


def plan_delete_block(editing: DocxEditing, operation: dict, location: str) -> Change:
    element = editing.elements[operation["block"]] if operation["block"] < len(editing.elements) else None
    if element is not None and id(element) in editing.touched:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: block {operation['block']} is used by another operation in this batch", location))
    element = resolve_block(editing, operation["block"], f"{location}.block")
    editing.deleted.add(id(element))

    def change() -> str:
        if editing.tracking is not None:
            mark_block_deleted(element, editing.tracking)
            return f"deleted block {operation['block']} as a tracked change"
        element.getparent().remove(element)
        return f"deleted block {operation['block']}"
    return change


def plan_set_style(editing: DocxEditing, operation: dict, location: str) -> Change:
    element = resolve_block(editing, operation["block"], f"{location}.block")
    is_table = element.tag == TABLE_TAG
    require_style(editing, operation["style"], (WD_STYLE_TYPE.TABLE,) if is_table else (WD_STYLE_TYPE.PARAGRAPH,), f"{location}.style")
    target = Table(element, editing.document._body) if is_table else Paragraph(element, editing.document._body)

    def change() -> str:
        old_properties = snapshot_paragraph_properties(element) if editing.tracking is not None and not is_table else None
        target.style = operation["style"]
        if old_properties is not None:
            record_paragraph_property_change(element, old_properties, editing.tracking)
        return f"set block {operation['block']} to style {operation['style']!r}"
    return change


def plan_set_cell(editing: DocxEditing, operation: dict, location: str) -> Change:
    table = resolve_table(editing, operation["block"], f"{location}.block")
    cell = table_cell(table, operation["row"], operation["column"], location)

    def change() -> str:
        replace_cell_text(cell, operation["text"], editing.tracking)
        return f"set cell ({operation['row']}, {operation['column']}) of block {operation['block']}"
    return change


def table_cell(table: Table, row: int, column: int, location: str):
    if row >= len(table.rows):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.row: the table has {len(table.rows)} rows", f"{location}.row"))
    cells = table.rows[row].cells
    if column >= len(cells):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.column: row {row} has {len(cells)} cells", f"{location}.column"))
    return cells[column]


def replace_cell_text(cell, text: str, tracking=None) -> None:
    paragraphs = cell.paragraphs
    if tracking is not None:
        tracked_set_text(paragraphs[0]._p, text, tracking)
        for paragraph in paragraphs[1:]:
            mark_paragraph_deleted(paragraph._p, tracking)
        return
    set_paragraph_text(paragraphs[0], text)
    for paragraph in paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)


def plan_insert_table_row(editing: DocxEditing, operation: dict, location: str) -> Change:
    table = resolve_table(editing, operation["block"], f"{location}.block")
    if operation["after"] >= len(table.rows):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.after: the table has {len(table.rows)} rows", f"{location}.after"))
    template_row = table.rows[operation["after"]]
    if len(operation["cells"]) > len(template_row.cells):
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.cells: the row has {len(template_row.cells)} cells", f"{location}.cells"))

    def change() -> str:
        new_row_element = copy.deepcopy(template_row._tr)
        template_row._tr.addnext(new_row_element)
        new_row = table.rows[operation["after"] + 1]
        for index, cell in enumerate(new_row.cells):
            value = operation["cells"][index] if index < len(operation["cells"]) else ""
            replace_cell_text(cell, "" if value is None else str(value))
        if editing.tracking is not None:
            mark_row(new_row_element, editing.tracking, inserted=True)
            return f"inserted a row after row {operation['after']} of block {operation['block']} as a tracked change"
        return f"inserted a row after row {operation['after']} of block {operation['block']}"
    return change


def plan_delete_table_row(editing: DocxEditing, operation: dict, location: str) -> Change:
    table = resolve_table(editing, operation["block"], f"{location}.block")
    if operation["row"] >= len(table.rows):
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.row: the table has {len(table.rows)} rows", f"{location}.row"))
    if len(table.rows) == 1:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: the table's only row cannot be deleted; delete the block instead", location))
    row_element = table.rows[operation["row"]]._tr

    def change() -> str:
        if editing.tracking is not None:
            mark_row(row_element, editing.tracking, inserted=False)
            return f"deleted row {operation['row']} of block {operation['block']} as a tracked change"
        row_element.getparent().remove(row_element)
        return f"deleted row {operation['row']} of block {operation['block']}"
    return change


def plan_set_east_asia_font(editing: DocxEditing, operation: dict, location: str) -> Change:
    def change() -> str:
        run_fonts = default_run_fonts(editing.document)
        run_fonts.set(qn("w:eastAsia"), operation["font"])
        run_fonts.attrib.pop(qn("w:eastAsiaTheme"), None)
        return f"set the default East Asian font to {operation['font']!r}"
    return change


def plan_set_korean_language(editing: DocxEditing, operation: dict, location: str) -> Change:
    def change() -> str:
        retagged = make_east_asia_language_korean(editing.document)
        return f"tagged the East Asian language as ko-KR; retagged {retagged} styles and runs that named another"
    return change


def default_run_fonts(document):
    styles = document.styles.element
    defaults = child_or_create(styles, "w:docDefaults", first=True)
    run_properties_default = child_or_create(defaults, "w:rPrDefault", first=True)
    run_properties = child_or_create(run_properties_default, "w:rPr", first=True)
    return child_or_create(run_properties, "w:rFonts", first=True)


def child_or_create(parent, tag: str, first: bool):
    child = parent.find(qn(tag))
    if child is not None:
        return child
    child = OxmlElement(tag)
    if first:
        parent.insert(0, child)
    else:
        parent.append(child)
    return child


def plan_update_fields_on_open(editing: DocxEditing, operation: dict, location: str) -> Change:
    def change() -> str:
        request_field_update(editing.document)
        return "asked Word to update fields when the file opens"
    return change


DOCX_OPERATIONS = OperationSet(OPERATIONS, {
    "replace_text": plan_replace_text,
    "set_text": plan_set_text,
    "insert_paragraph": plan_insert_paragraph,
    "insert_heading": plan_insert_heading,
    "insert_list": plan_insert_list,
    "insert_table": plan_insert_table,
    "insert_page_break": plan_insert_page_break,
    "delete_block": plan_delete_block,
    "set_style": plan_set_style,
    "set_cell": plan_set_cell,
    "insert_table_row": plan_insert_table_row,
    "delete_table_row": plan_delete_table_row,
    "set_header": plan_set_header,
    "set_footer": plan_set_footer,
    "set_page_setup": plan_set_page_setup,
    "insert_section_break": plan_insert_section_break,
    "set_watermark": plan_set_watermark,
    "insert_image": plan_insert_image,
    "insert_table_column": plan_insert_table_column,
    "delete_table_column": plan_delete_table_column,
    "merge_cells": plan_merge_cells,
    "format_cells": plan_format_cells,
    "format_text": plan_format_text,
    "set_paragraph_format": plan_set_paragraph_format,
    "define_style": plan_define_style,
    "insert_table_of_contents": plan_insert_table_of_contents,
    "add_bookmark": plan_add_bookmark,
    "insert_link": plan_insert_link,
    "insert_cross_reference": plan_insert_cross_reference,
    "insert_footnote": plan_insert_footnote,
    "insert_endnote": plan_insert_endnote,
    "add_comment": plan_add_comment,
    "reply_comment": plan_reply_comment,
    "resolve_comment": plan_resolve_comment,
    "delete_comment": plan_delete_comment,
    "accept_revisions": plan_accept_revisions,
    "reject_revisions": plan_reject_revisions,
    "set_east_asia_font": plan_set_east_asia_font,
    "set_korean_language": plan_set_korean_language,
    "update_fields_on_open": plan_update_fields_on_open,
})
