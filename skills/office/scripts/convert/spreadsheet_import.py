from __future__ import annotations

from dataclasses import dataclass, field
import datetime
from pathlib import Path
import zipfile

from lxml import etree
from openpyxl import Workbook
from python_calamine import CalamineWorkbook

from convert.convert_definitions import CONVERSION_APPROXIMATED
from sheet.formula_cache import cache_formula_values
from core.office_result import Issue
from core.excel_limits import fitting_sheet_name


OPEN_DOCUMENT_NAMESPACES = {"table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0"}
PACKAGE_RELATIONSHIP_NAMESPACE = "{http://schemas.openxmlformats.org/package/2006/relationships}"
BUNDLE_SHEET_RECORD = 0x009C
MERGE_CELL_RECORD = 0x00B0


@dataclass
class SheetData:
    name: str
    cells: dict[tuple[int, int], object] = field(default_factory=dict)
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)


def legacy_workbook_to_xlsx(input_path: Path, output_path: Path) -> list[Issue]:
    write_workbook(read_legacy_workbook(input_path), output_path)
    issues = list(cache_formula_values(str(output_path)))
    issues.append(CONVERSION_APPROXIMATED.issue("carried values, dates and merged cells, with each formula kept as its computed value; fonts, colors, borders, column widths and charts were not carried", input_path.name))
    return issues


def read_legacy_workbook(input_path: Path) -> list[SheetData]:
    book = CalamineWorkbook.from_path(str(input_path))
    merges = MERGE_READERS[input_path.suffix.lower()](input_path, book)
    return [sheet_data(name, book.get_sheet_by_name(name).to_python(skip_empty_area=False), merges.get(name, [])) for name in book.sheet_names]


def sheet_data(name: str, rows: list[list], merges: list) -> SheetData:
    data = SheetData(name, merges=merges)
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            if value != "" and value is not None:
                data.cells[(row_index, column_index)] = int(value) if isinstance(value, float) and value.is_integer() else value
    return data


def calamine_merges(input_path: Path, book: CalamineWorkbook) -> dict[str, list[tuple[int, int, int, int]]]:
    return {name: [(top + 1, left + 1, bottom + 1, right + 1) for (top, left), (bottom, right) in book.get_sheet_by_name(name).merged_cell_ranges or []] for name in book.sheet_names}


def ods_merges(input_path: Path, book: CalamineWorkbook) -> dict[str, list[tuple[int, int, int, int]]]:
    with zipfile.ZipFile(input_path) as archive:
        content = etree.fromstring(archive.read("content.xml"))
    return {table.get(open_document("table", "name")): ods_table_merges(table) for table in content.iterfind(".//table:table", OPEN_DOCUMENT_NAMESPACES)}


def ods_table_merges(table) -> list[tuple[int, int, int, int]]:
    merges, row_number = [], 1
    for row in table.iterfind(".//table:table-row", OPEN_DOCUMENT_NAMESPACES):
        merges.extend(ods_row_merges(row, row_number))
        row_number += ods_count(row, "number-rows-repeated")
    return merges


def ods_row_merges(row, row_number: int) -> list[tuple[int, int, int, int]]:
    merges, column = [], 1
    for cell in row:
        rows, columns = ods_count(cell, "number-rows-spanned"), ods_count(cell, "number-columns-spanned")
        if rows > 1 or columns > 1:
            merges.append((row_number, column, row_number + rows - 1, column + columns - 1))
        column += ods_count(cell, "number-columns-repeated")
    return merges


def ods_count(element, attribute: str) -> int:
    return int(element.get(open_document("table", attribute), "1"))


def xlsb_records(data: bytes):
    position = 0
    while position < len(data):
        kind, position = variable_integer(data, position, 2)
        size, position = variable_integer(data, position, 4)
        yield kind, data[position:position + size]
        position += size


def variable_integer(data: bytes, position: int, limit: int) -> tuple[int, int]:
    value = 0
    for shift in range(limit):
        byte = data[position]
        position += 1
        value |= (byte & 0x7F) << (7 * shift)
        if not byte & 0x80:
            break
    return value, position


def wide_text(payload: bytes, offset: int) -> tuple[str, int]:
    length = int.from_bytes(payload[offset:offset + 4], "little")
    end = offset + 4 + 2 * length
    return payload[offset + 4:end].decode("utf-16-le"), end


def xlsb_sheet_parts(archive: zipfile.ZipFile) -> dict[str, str]:
    relationships = etree.fromstring(archive.read("xl/_rels/workbook.bin.rels"))
    targets = {element.get("Id"): element.get("Target") for element in relationships.iter(f"{PACKAGE_RELATIONSHIP_NAMESPACE}Relationship")}
    parts = {}
    for kind, payload in xlsb_records(archive.read("xl/workbook.bin")):
        if kind == BUNDLE_SHEET_RECORD:
            identifier, end = wide_text(payload, 8)
            name, _ = wide_text(payload, end)
            parts[name] = "xl/" + targets.get(identifier, "").lstrip("/").removeprefix("xl/")
    return parts


def xlsb_merges(input_path: Path, book: CalamineWorkbook) -> dict[str, list[tuple[int, int, int, int]]]:
    with zipfile.ZipFile(input_path) as archive:
        names = set(archive.namelist())
        merges = {}
        for name, part in xlsb_sheet_parts(archive).items():
            if part not in names:
                continue
            records = [payload for kind, payload in xlsb_records(archive.read(part)) if kind == MERGE_CELL_RECORD]
            merges[name] = [merge_bounds(payload) for payload in records]
    return merges


def merge_bounds(payload: bytes) -> tuple[int, int, int, int]:
    first_row, last_row, first_column, last_column = (int.from_bytes(payload[offset:offset + 4], "little") for offset in (0, 4, 8, 12))
    return first_row + 1, first_column + 1, last_row + 1, last_column + 1


def write_workbook(sheets: list[SheetData], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        worksheet = workbook.create_sheet(fitting_sheet_name(sheet.name))
        for (row, column), value in sheet.cells.items():
            cell = worksheet.cell(row=row, column=column, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"
            if isinstance(value, (datetime.date, datetime.datetime)):
                cell.number_format = "yyyy-mm-dd" if not isinstance(value, datetime.datetime) else "yyyy-mm-dd hh:mm"
        for top, left, bottom, right in sheet.merges:
            worksheet.merge_cells(start_row=top, start_column=left, end_row=bottom, end_column=right)
    workbook.save(output_path)


def open_document(prefix: str, name: str) -> str:
    return f"{{{OPEN_DOCUMENT_NAMESPACES[prefix]}}}{name}"


MERGE_READERS = {".xls": calamine_merges, ".ods": ods_merges, ".xlsb": xlsb_merges}
