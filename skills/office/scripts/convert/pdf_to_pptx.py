from __future__ import annotations

from dataclasses import dataclass, field
import io
from pathlib import Path

import pdfplumber
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
import pypdfium2

from fonts.registry import OFFICE_KOREAN_FAMILY
from markdown_blocks import inline_segments, link_parts, strip_inline_markers
from office_inputs import unlocked_pdf_bytes
from pdf_to_blocks import MINIMUM_IMAGE_POINTS, Segment, crop, inside_any, page_image_bitmap, page_segments
from pdf_tables import page_tables


EMU_PER_POINT = 12700
SMALLEST_SLIDE_POINTS = 72
LARGEST_SLIDE_POINTS = 4032
LINE_X_TOLERANCE_POINTS = 4
LINE_GAP_FACTOR = 0.8
SIZE_TOLERANCE_POINTS = 0.6
BOX_SLACK_POINTS = 2
THIN_SHAPE_POINTS = 2.5
MOST_RULES = 40
UNDERLINE_REACH = 0.3
PLAIN_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
NO_TEXT_LAYER_REASON = "no text layer"


@dataclass
class SlideReport:
    page: int
    kind: str
    reason: str = ""
    text_boxes: int = 0
    tables: int = 0
    pictures: int = 0

    def to_json(self) -> dict:
        report = {"page": self.page, "slide": self.kind, "textBoxes": self.text_boxes, "tables": self.tables, "pictures": self.pictures}
        return report | ({"reason": self.reason} if self.reason else {})


@dataclass
class PdfSlides:
    reports: list[SlideReport] = field(default_factory=list)


def write_pdf_slides(input_path: Path, output_path: Path, password: str | None) -> PdfSlides:
    slides = PdfSlides()
    data = unlocked_pdf_bytes(str(input_path), password)
    rendered = pypdfium2.PdfDocument(data)
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            presentation = Presentation()
            presentation.slide_width, presentation.slide_height = slide_size(pdf.pages[0])
            for number, page in enumerate(pdf.pages, start=1):
                slides.reports.append(add_page_slide(presentation, page, rendered[number - 1], number))
            presentation.save(str(output_path))
    finally:
        rendered.close()
    return slides


def slide_size(page) -> tuple[Emu, Emu]:
    return tuple(Emu(round(min(max(float(points), SMALLEST_SLIDE_POINTS), LARGEST_SLIDE_POINTS) * EMU_PER_POINT)) for points in (page.width, page.height))


def add_page_slide(presentation, page, rendered_page, number: int) -> SlideReport:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    scale = presentation.slide_width / (float(page.width) * EMU_PER_POINT)
    image = page_image_bitmap(rendered_page)
    segments = page_segments(page)
    tables = page_tables(page) if segments else []
    reason = picture_reason(page, segments)
    if reason:
        add_picture(slide, image, 0, 0, presentation.slide_width, presentation.slide_height)
        slide.notes_slide.notes_text_frame.text = "\n".join(segment.text for segment in segments)
        return SlideReport(number, "picture", reason, pictures=1)
    report = SlideReport(number, "editable")
    table_boxes = [table.bbox for table in tables]
    rectangles = page_rectangles(page, table_boxes)
    for rectangle in [rectangle for rectangle in rectangles if rectangle.get("fill")]:
        add_rectangle(slide, rectangle, scale)
    for picture in page_pictures(page):
        add_picture(slide, crop(image, picture), *(points_to_emu(value, scale) for value in (picture["x0"], picture["top"], picture["x1"] - picture["x0"], picture["bottom"] - picture["top"])))
        report.pictures += 1
    for table in tables:
        add_table(slide, table, scale, table_fills(page, table))
    for rectangle in [rectangle for rectangle in rectangles if not rectangle.get("fill")]:
        add_rectangle(slide, rectangle, scale)
    for rule in page_rules(page, segments):
        add_rule(slide, rule, scale)
    content = [segment for segment in segments if not inside_any(segment, table_boxes)]
    for lines in text_blocks(content):
        add_text_box(slide, lines, scale, text_color(page.chars, lines[0]))
        report.text_boxes += 1
    report.tables = len(tables)
    return report


def picture_reason(page, segments: list[Segment]) -> str:
    if not segments:
        return NO_TEXT_LAYER_REASON
    if any(not is_thin(curve) for curve in page.curves):
        return "curved shapes or drawings"
    if not all(has_plain_colors(shape) for shape in [*page.rects, *page.lines]):
        return "pattern or gradient fills"
    if len(page_rules(page, segments)) > MOST_RULES:
        return "a drawing made of many lines"
    return ""


def has_plain_colors(shape: dict) -> bool:
    painted = [shape.get("non_stroking_color")] if shape.get("fill") else []
    painted += [shape.get("stroking_color")] if shape.get("stroke") or shape.get("object_type") == "line" else []
    return all(is_plain_color(color) for color in painted)


def is_plain_color(color) -> bool:
    if isinstance(color, (int, float)):
        return True
    return isinstance(color, (list, tuple)) and len(color) in (1, 3, 4) and all(isinstance(value, (int, float)) for value in color)


def page_rectangles(page, table_boxes: list) -> list[dict]:
    drawn = [rectangle for rectangle in page.rects if not is_thin(rectangle) and (rectangle.get("fill") or rectangle.get("stroke"))]
    return [rectangle for rectangle in drawn if not inside_box(rectangle, table_boxes)] + [
        {**rectangle, "fill": False} for rectangle in drawn if rectangle.get("stroke") and inside_box(rectangle, table_boxes)
    ]


def text_color(characters: list[dict], segment: Segment) -> str:
    inside = next((character for character in characters if segment.x0 <= (character["x0"] + character["x1"]) / 2 <= segment.x1 and segment.top <= (character["top"] + character["bottom"]) / 2 <= segment.bottom), None)
    return hex_color(inside.get("non_stroking_color") if inside is not None else 0)


def page_rules(page, segments: list[Segment]) -> list[dict]:
    return [shape for shape in [*page.lines, *page.rects] if is_thin(shape) and not underlines_text(shape, segments)]


def underlines_text(rule: dict, segments: list[Segment]) -> bool:
    return any(
        segment.x0 - 1 <= float(rule["x0"]) and float(rule["x1"]) <= segment.x1 + 1 and segment.bottom - segment.size * UNDERLINE_REACH <= float(rule["top"]) <= segment.bottom + segment.size * UNDERLINE_REACH
        for segment in segments
    )


def inside_box(shape: dict, boxes: list) -> bool:
    return any(x0 - 1 <= shape["x0"] and shape["x1"] <= x1 + 1 and top - 1 <= shape["top"] and shape["bottom"] <= bottom + 1 for x0, top, x1, bottom in boxes)


def is_thin(shape: dict) -> bool:
    return min(float(shape["x1"]) - float(shape["x0"]), float(shape["bottom"]) - float(shape["top"])) <= THIN_SHAPE_POINTS


def page_pictures(page) -> list[dict]:
    return [picture for picture in page.images if picture["x1"] - picture["x0"] >= MINIMUM_IMAGE_POINTS and picture["bottom"] - picture["top"] >= MINIMUM_IMAGE_POINTS]


def text_blocks(segments: list[Segment]) -> list[list[Segment]]:
    blocks: list[list[Segment]] = []
    for segment in sorted(segments, key=lambda segment: (round(segment.top), segment.x0)):
        block = next((block for block in reversed(blocks) if continues(block[-1], segment)), None)
        if block is None:
            blocks.append([segment])
        else:
            block.append(segment)
    return blocks


def continues(previous: Segment, segment: Segment) -> bool:
    same_column = abs(previous.x0 - segment.x0) <= LINE_X_TOLERANCE_POINTS
    same_size = abs(previous.size - segment.size) <= SIZE_TOLERANCE_POINTS
    gap = segment.top - previous.bottom
    return same_column and same_size and 0 <= gap <= previous.size * LINE_GAP_FACTOR


def points_to_emu(points: float, scale: float) -> Emu:
    return Emu(round(float(points) * EMU_PER_POINT * scale))


def add_text_box(slide, lines: list[Segment], scale: float, color: str) -> None:
    left, top = min(line.x0 for line in lines), lines[0].top
    width = max(line.x1 for line in lines) - left + BOX_SLACK_POINTS
    height = lines[-1].bottom - top
    box = slide.shapes.add_textbox(points_to_emu(left, scale), points_to_emu(top, scale), points_to_emu(width, scale), points_to_emu(height, scale))
    frame = box.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        for piece in inline_segments(line.markdown):
            add_run(paragraph, piece, line, scale, color)


def add_run(paragraph, piece: str, line: Segment, scale: float, color: str) -> None:
    link = link_parts(piece)
    marked = link[0] if link else piece
    run = paragraph.add_run()
    run.text = strip_inline_markers(marked)
    run.font.size = Pt(round(line.size * scale * 2) / 2)
    run.font.bold = line.bold or marked.startswith("**")
    run.font.color.rgb = RGBColor.from_string(color)
    if link:
        run.hyperlink.address = link[1]
        run.font.underline = True
    set_east_asian_typeface(run)


def set_east_asian_typeface(run) -> None:
    properties = run._r.get_or_add_rPr()
    east_asian = properties.find(qn("a:ea"))
    if east_asian is None:
        east_asian = properties.makeelement(qn("a:ea"), {})
        properties.append(east_asian)
    east_asian.set("typeface", OFFICE_KOREAN_FAMILY)


def table_fills(page, table) -> list[dict]:
    return [shape for shape in page.rects if shape.get("fill") and not is_thin(shape) and inside_box(shape, [table.bbox])]


def add_table(slide, table, scale: float, fills: list[dict]) -> None:
    x0, top, x1, bottom = table.bbox
    column_count = max(len(row) for row in table.rows)
    shape = slide.shapes.add_table(len(table.rows), column_count, points_to_emu(x0, scale), points_to_emu(top, scale), points_to_emu(x1 - x0, scale), points_to_emu(bottom - top, scale))
    shape.table._tbl.tblPr.find(qn("a:tableStyleId")).text = PLAIN_TABLE_STYLE
    shape.table.first_row = shape.table.horz_banding = False
    row_height, column_width = (bottom - top) / len(table.rows), (x1 - x0) / column_count
    size = Pt(max(8, round(min(12, row_height * 0.45) * scale)))
    for row_index, row in enumerate(table.rows):
        for column_index in range(column_count):
            cell = shape.table.cell(row_index, column_index)
            cell.text = row[column_index] if column_index < len(row) else ""
            center = (x0 + (column_index + 0.5) * column_width, top + (row_index + 0.5) * row_height)
            fill = next((rect for rect in fills if rect["x0"] <= center[0] <= rect["x1"] and rect["top"] <= center[1] <= rect["bottom"]), None)
            if fill is None:
                cell.fill.background()
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(hex_color(fill.get("non_stroking_color")))
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = size
                    set_east_asian_typeface(run)


def add_rectangle(slide, rectangle: dict, scale: float) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(points_to_emu(value, scale) for value in (rectangle["x0"], rectangle["top"], float(rectangle["x1"]) - float(rectangle["x0"]), float(rectangle["bottom"]) - float(rectangle["top"]))))
    shape.shadow.inherit = False
    if rectangle.get("fill"):
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(hex_color(rectangle.get("non_stroking_color")))
    else:
        shape.fill.background()
    if rectangle.get("stroke"):
        shape.line.color.rgb = RGBColor.from_string(hex_color(rectangle.get("stroking_color")))
        shape.line.width = Pt(max(float(rectangle.get("linewidth") or 0), 0.25) * scale)
    else:
        shape.line.fill.background()


def add_rule(slide, rule: dict, scale: float) -> None:
    horizontal = float(rule["x1"]) - float(rule["x0"]) >= float(rule["bottom"]) - float(rule["top"])
    middle = (float(rule["top"]) + float(rule["bottom"])) / 2 if horizontal else (float(rule["x0"]) + float(rule["x1"])) / 2
    start, end = ((rule["x0"], middle), (rule["x1"], middle)) if horizontal else ((middle, rule["top"]), (middle, rule["bottom"]))
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *(points_to_emu(value, scale) for value in (*start, *end)))
    thickness = float(rule.get("linewidth") or 0) if rule.get("object_type") == "line" else min(float(rule["x1"]) - float(rule["x0"]), float(rule["bottom"]) - float(rule["top"]))
    connector.line.width = Pt(max(thickness, 0.25) * scale)
    color = rule.get("stroking_color") if rule.get("object_type") == "line" else rule.get("non_stroking_color")
    connector.line.color.rgb = RGBColor.from_string(hex_color(color))


def hex_color(color) -> str:
    values = list(color) if isinstance(color, (list, tuple)) else [color] if isinstance(color, (int, float)) else [0]
    if len(values) == 4:
        cyan, magenta, yellow, black = values
        values = [(1 - cyan) * (1 - black), (1 - magenta) * (1 - black), (1 - yellow) * (1 - black)]
    elif len(values) != 3:
        values = [values[0]] * 3
    return "".join(f"{round(min(max(float(value), 0), 1) * 255):02X}" for value in values)


def add_picture(slide, image, left: Emu, top: Emu, width: Emu, height: Emu) -> None:
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    slide.shapes.add_picture(buffer, left, top, width, height)
