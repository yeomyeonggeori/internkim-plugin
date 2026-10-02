from __future__ import annotations

import csv
import datetime
from pathlib import Path
import re
import shutil
import tempfile

from openpyxl import load_workbook

from convert.convert_definitions import FORMULA_VALUE_MISSING
from sheet.create_xlsx import create_workbook
from sheet.formulas.cache import cache_formula_values
from core.office_result import INVALID_VALUE, Issue, OfficeFailure
from core.excel_limits import fitting_sheet_name


DELIMITERS = {"csv": ",", "tsv": "\t"}
FILE_NAME_UNSAFE = re.compile(r"[\\/:*?\"<>|\s]+")


def workbook_to_delimited(input_path: Path, output_path: Path, delimiter: str, sheet_name: str | None) -> tuple[list[str], list[Issue]]:
    formulas = load_workbook(str(input_path), data_only=False, read_only=False)
    values = load_workbook(str(input_path), data_only=True, read_only=False)
    if formulas_without_values(values, formulas):
        values = computed_values(input_path)
    names = [sheet_name] if sheet_name else values.sheetnames
    if sheet_name and sheet_name not in values.sheetnames:
        raise OfficeFailure(INVALID_VALUE.issue(f"--sheet {sheet_name!r} is not in the workbook", "--sheet", suggestion=f"use one of: {', '.join(values.sheetnames)}"))
    written, issues = [], []
    for name in names:
        path = output_path if len(names) == 1 else output_path.with_name(f"{output_path.stem}-{FILE_NAME_UNSAFE.sub('_', name)}{output_path.suffix}")
        issues.extend(write_sheet(values[name], formulas[name], path, delimiter))
        written.append(str(path))
    return written, issues


def formulas_without_values(values, formulas) -> bool:
    return any(
        value_cell.value is None and isinstance(formula_cell.value, str) and formula_cell.value.startswith("=")
        for name in formulas.sheetnames
        for value_row, formula_row in zip(values[name].iter_rows(), formulas[name].iter_rows())
        for value_cell, formula_cell in zip(value_row, formula_row)
    )


def computed_values(input_path: Path):
    with tempfile.TemporaryDirectory(prefix="office-convert-") as directory:
        computed = Path(directory) / "computed.xlsx"
        shutil.copyfile(input_path, computed)
        cache_formula_values(str(computed))
        return load_workbook(str(computed), data_only=True, read_only=False)


def write_sheet(value_sheet, formula_sheet, path: Path, delimiter: str) -> list[Issue]:
    missing = []
    rows = []
    for value_row, formula_row in zip(value_sheet.iter_rows(), formula_sheet.iter_rows()):
        row = []
        for value_cell, formula_cell in zip(value_row, formula_row):
            if value_cell.value is None and isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
                missing.append(f"{value_sheet.title}!{value_cell.coordinate}")
            row.append(text_value(value_cell.value))
        rows.append(row)
    while rows and not any(rows[-1]):
        rows.pop()
    with open(path, "w", newline="", encoding="utf-8-sig") as output:
        csv.writer(output, delimiter=delimiter).writerows(rows)
    if not missing:
        return []
    return [FORMULA_VALUE_MISSING.issue(f"{len(missing)} formula cells in {value_sheet.title!r} have no saved value, such as {missing[0]}", missing[0])]


def text_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time() else value.isoformat(sep=" ")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def delimited_to_workbook(input_path: Path, output_path: Path, delimiter: str) -> list[Issue]:
    specification = {"title": input_path.stem, "sheets": [{"title": fitting_sheet_name(input_path.stem), "csvPath": str(input_path), "delimiter": delimiter}]}
    workbook = create_workbook(specification)
    workbook.save(output_path)
    return list(cache_formula_values(str(output_path)))
