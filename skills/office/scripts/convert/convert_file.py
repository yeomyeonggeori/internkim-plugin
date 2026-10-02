#!/usr/bin/env python3
from __future__ import annotations

import base64
from dataclasses import dataclass, field, replace
import mimetypes
from pathlib import Path
import tempfile

from doc.blocks.writers import file_data_uri, html_document, markdown_text
from convert.convert_definitions import CONVERSION_APPROXIMATED, PAGE_WITHOUT_TEXT, TABLE_NOT_FOUND, UNSUPPORTED_CONVERSION
from core.office_arguments import route_arguments
from core.office_commands import KINDS, Conversion as ConversionRoute, command_text, conversion_sources, conversion_targets, find_conversion, find_route, normalized_extension
from docx.shared import Pt
from fonts.docx_embedding import save_document
from fonts.registry import BODY_SIZE_POINTS
from doc.blocks.docx import DEFAULT_DOCUMENT_FONT, markdown_document
from convert.docx_to_blocks import read_docx_blocks
from doc.export_document import export_pdf
from convert.html_to_blocks import read_html_blocks
from render.office_preview import PAGE_SELECTOR, Preview, write_preview
from doc.blocks.markdown import Image
from core.office_inputs import KINDS_BY_NAME, PDF, office_file, require_unlocked_pdf
from core.office_result import INVALID_VALUE, Issue, OfficeFailure, Result, run_command
from core.office_inputs import read_text_input, unlocked_pdf_bytes
from pdf.pdf_definitions import OCR_UNAVAILABLE, PAGE_READ_BY_OCR, page_reading_suggestion
from pdf.ocr.pdf_ocr import OcrUnavailable, pages_without_words, read_pages_by_ocr
from convert.pdf_to_blocks import read_pdf_blocks
from convert.pdf_workbook import read_pdf_tables, table_details, write_pdf_workbook
from convert.pdf_to_pptx import NO_TEXT_LAYER_REASON, write_pdf_slides
from pptx import Presentation
from powerpoint.preview.document import preview_document
from render.renderer import RENDER_FAILED, RENDERER_UNAVAILABLE, RenderFailed, RendererUnavailable, draw_preview
from doc.render_docx import docx_preview
from sheet.render_xlsx import xlsx_preview
from powerpoint.check_pptx import SLIDE_SELECTOR, preview_fonts
from convert.spreadsheet_import import legacy_workbook_to_xlsx
from convert.table_conversions import DELIMITERS, workbook_to_delimited


OCR_ROUTES = (("pdf", "docx"), ("pdf", "md"), ("pdf", "xlsx"))
MINIMUM_MARGIN_POINTS = 36
MAXIMUM_MARGIN_POINTS = 90


@dataclass
class Conversion:
    input_path: Path
    output_path: Path
    sheet: str | None
    password: str | None = None
    ocr: bool = False
    issues: list[Issue] = field(default_factory=list)
    details: dict = field(default_factory=dict)
    written: list[str] = field(default_factory=list)


def main() -> Result:
    arguments = route_arguments("convert")
    input_path = Path(arguments.input).expanduser()
    output_path = Path(arguments.output).expanduser()
    if not input_path.is_file():
        raise FileNotFoundError(2, "no such file", str(input_path))
    route = require_route(input_path, output_path)
    require_ocr_route(route, arguments.ocr)
    require_readable_source(arguments.input, route.source, arguments.password)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    conversion = Conversion(input_path, output_path, arguments.sheet, arguments.password, arguments.ocr)
    CONVERTERS[(route.source, route.target)](conversion)
    written = conversion.written or [str(output_path)]
    details = {"route": f"{route.source} -> {route.target}", "files": written, **conversion.details}
    return Result(summary=f"converted {input_path.name} to {', '.join(Path(path).name for path in written)}", output_path=written[0], issues=tuple(conversion.issues), details=details)


def require_readable_source(input_path: str, source: str, password: str | None) -> None:
    if source not in KINDS_BY_NAME:
        return
    office_file(source)(input_path)
    if KINDS_BY_NAME[source] == PDF:
        require_unlocked_pdf(input_path, password)


def require_ocr_route(route: ConversionRoute, ocr: bool) -> None:
    if ocr and (route.source, route.target) not in OCR_ROUTES:
        raise OfficeFailure(INVALID_VALUE.issue(f"--ocr reads scanned pages only for pdf to docx, md or xlsx, not {route.source} to {route.target}", "--ocr", suggestion="drop --ocr, or convert the PDF to .docx, .md or .xlsx"))


def require_route(input_path: Path, output_path: Path) -> ConversionRoute:
    source, target = normalized_extension(input_path.suffix), normalized_extension(output_path.suffix)
    route = find_conversion(source, target)
    if route is not None:
        return route
    raise OfficeFailure(UNSUPPORTED_CONVERSION.issue(f"no route converts .{source} to .{target}", f"{input_path.name} -> {output_path.name}", suggestion=route_suggestion(input_path, output_path, source)))


def route_suggestion(input_path: Path, output_path: Path, source: str) -> str:
    kind = next((kind.name for kind in KINDS if input_path.suffix.lower() in kind.extensions), None)
    creating = find_route("create", kind) if kind else None
    if creating is not None and output_path.suffix.lower() in creating.outputs:
        return command_text(["office", "create", str(output_path), str(input_path)])
    targets = conversion_targets(source)
    if targets:
        return f".{source} converts to {', '.join('.' + target for target in targets)}"
    return f"inputs this command reads: {', '.join('.' + candidate for candidate in conversion_sources())}"











def docx_to_pdf(conversion: Conversion) -> None:
    preview, fonts = docx_preview(conversion.input_path)
    report_simplified(conversion, preview)
    with tempfile.TemporaryDirectory(prefix="office-convert-") as directory:
        draw_pdf(conversion, write_preview(preview, Path(directory)), PAGE_SELECTOR, fonts)


def workbook_to_pdf(conversion: Conversion) -> None:
    conversion.issues.extend(draw_workbook_pdf(conversion, conversion.input_path, conversion.sheet))




def draw_workbook_pdf(conversion: Conversion, workbook_path: Path, sheet: str | None) -> list[Issue]:
    preview, fonts, _, print_issues = xlsx_preview(workbook_path, sheet)
    report_simplified(conversion, preview)
    with tempfile.TemporaryDirectory(prefix="office-convert-") as directory:
        draw_pdf(conversion, write_preview(preview, Path(directory)), PAGE_SELECTOR, fonts)
    return print_issues


def presentation_to_pdf(conversion: Conversion) -> None:
    presentation = Presentation(str(conversion.input_path))
    preview = preview_document(presentation, list(range(1, len(presentation.slides) + 1)))
    with tempfile.TemporaryDirectory(prefix="office-convert-") as directory:
        preview_path = Path(directory) / "preview.html"
        preview_path.write_text(preview.html, encoding="utf-8")
        draw_pdf(conversion, preview_path, SLIDE_SELECTOR, preview_fonts(preview.faces))


def draw_pdf(conversion: Conversion, preview_path: Path, page_selector: str, fonts: list[dict]) -> None:
    try:
        rendered = draw_preview(preview_path, page_selector, fonts, conversion.output_path.resolve(), draw_images=False)
    except RendererUnavailable as reason:
        raise OfficeFailure(RENDERER_UNAVAILABLE.issue(f"{conversion.output_path.name} was not written: {reason}", conversion.input_path.name))
    except RenderFailed as reason:
        raise OfficeFailure(RENDER_FAILED.issue(f"{conversion.output_path.name} was not written: {reason}", conversion.input_path.name))
    conversion.details["pageCount"] = len(rendered.page_sizes)


def report_simplified(conversion: Conversion, preview: Preview) -> None:
    if preview.approximations:
        listed = ", ".join(f"{count} {what}" for what, count in sorted(preview.approximations.items()))
        conversion.issues.append(CONVERSION_APPROXIMATED.issue(f"the PDF simplifies {listed}", conversion.input_path.name))


def docx_to_markdown(conversion: Conversion) -> None:
    media_name = f"{conversion.output_path.stem}-media"
    reading = read_docx_blocks(conversion.input_path, media_name)
    save_media(reading.media, conversion.output_path.parent / media_name, conversion)
    conversion.output_path.write_text(markdown_text(reading.blocks), encoding="utf-8")
    report_dropped(conversion, reading.dropped)


def docx_to_html(conversion: Conversion) -> None:
    reading = read_docx_blocks(conversion.input_path, "media")
    blocks = [with_data_source(block, reading.media) for block in reading.blocks]
    conversion.output_path.write_text(html_document(blocks, conversion.input_path.stem), encoding="utf-8")
    report_dropped(conversion, reading.dropped)


def html_to_docx(conversion: Conversion) -> None:
    with tempfile.TemporaryDirectory(prefix="office-convert-") as media_directory:
        write_docx(conversion, html_blocks(read_text(conversion.input_path), Path(media_directory)), conversion.input_path.parent)


def html_to_pdf(conversion: Conversion) -> None:
    with tempfile.TemporaryDirectory(prefix="office-convert-") as media_directory:
        blocks = html_blocks(read_text(conversion.input_path), Path(media_directory))
        conversion.issues.extend(export_pdf(blocks, conversion.output_path, conversion.input_path.parent, ""))


def html_blocks(html: str, media_directory: Path) -> list:
    return [decoded(block, media_directory) for block in read_html_blocks(html)]


def html_to_markdown(conversion: Conversion) -> None:
    conversion.output_path.write_text(markdown_text(read_html_blocks(read_text(conversion.input_path))), encoding="utf-8")


def pdf_to_docx(conversion: Conversion) -> None:
    with tempfile.TemporaryDirectory(prefix="office-convert-") as media_directory:
        reading = read_pdf_reading(conversion, Path(media_directory), Path(media_directory).name)
        document = write_docx(conversion, reading.blocks, Path(media_directory).parent, save=False)
        match_pdf_page(document, reading)
        save_document(document, conversion.output_path)


def pdf_to_markdown(conversion: Conversion) -> None:
    media_name = f"{conversion.output_path.stem}-media"
    media_directory = conversion.output_path.parent / media_name
    media_directory.mkdir(parents=True, exist_ok=True)
    reading = read_pdf_reading(conversion, media_directory, media_name)
    conversion.output_path.write_text(markdown_text(reading.blocks), encoding="utf-8")
    if not any(media_directory.iterdir()):
        media_directory.rmdir()


def pdf_to_presentation(conversion: Conversion) -> None:
    reports = write_pdf_slides(conversion.input_path, conversion.output_path, conversion.password).reports
    conversion.details["pages"] = [report.to_json() for report in reports]
    scanned = [report.page for report in reports if report.reason == NO_TEXT_LAYER_REASON]
    if scanned:
        conversion.issues.append(page_without_text_issue(conversion, scanned, "became pictures"))
    approximations = ["each line keeps its position and size in a text box, but not the PDF's own font"]
    drawn = [f"{report.page} ({report.reason})" for report in reports if report.kind == "picture" and report.page not in scanned]
    if drawn:
        approximations.append(f"pages {', '.join(drawn)} became one picture each, with their text in the slide notes")
    conversion.details["approximations"] = approximations
    conversion.issues.append(CONVERSION_APPROXIMATED.issue("; ".join(approximations), conversion.input_path.name))


def page_without_text_issue(conversion: Conversion, pages: list[int], outcome: str) -> Issue:
    listed = ",".join(map(str, pages))
    can_rerun_with_ocr = not conversion.ocr and route_of(conversion) in OCR_ROUTES
    rerun_command = f"office convert {conversion.input_path} {conversion.output_path}" if can_rerun_with_ocr else None
    return PAGE_WITHOUT_TEXT.issue(f"pages {listed} have no text layer and {outcome}", f"pages {listed}", suggestion=page_reading_suggestion(conversion.input_path, listed, rerun_command))


def route_of(conversion: Conversion) -> tuple[str, str]:
    return normalized_extension(conversion.input_path.suffix), normalized_extension(conversion.output_path.suffix)


def scanned_page_lines(conversion: Conversion) -> dict:
    if not conversion.ocr:
        return {}
    data = unlocked_pdf_bytes(str(conversion.input_path), conversion.password)
    scanned = pages_without_words(data)
    try:
        lines = read_pages_by_ocr(data, scanned)
    except OcrUnavailable as reason:
        conversion.issues.append(OCR_UNAVAILABLE.issue(str(reason), f"pages {','.join(map(str, scanned))}"))
        return {}
    read = [number for number, page_lines in lines.items() if page_lines]
    if read:
        conversion.issues.append(PAGE_READ_BY_OCR.issue(f"pages {','.join(map(str, read))} were read by OCR", f"pages {','.join(map(str, read))}"))
    return {number: page_lines for number, page_lines in lines.items() if page_lines}


def read_pdf_reading(conversion: Conversion, media_directory: Path, media_prefix: str):
    reading = read_pdf_blocks(conversion.input_path, media_directory, media_prefix, conversion.password, scanned_page_lines(conversion))
    conversion.details["pages"] = [report.to_json() for report in reading.reports]
    pictures = [report.page for report in reading.reports if not report.has_text]
    if pictures:
        conversion.issues.append(page_without_text_issue(conversion, pictures, "were kept as pictures"))
    approximations = ["text is reflowed into paragraphs, so line breaks and exact positions are not kept", "headings are inferred from font size and bold", "page breaks of the PDF are not kept"]
    if any(report.columns > 1 for report in reading.reports):
        approximations.append("two-column pages are read left column first, then right, into one column")
    if any(report.header_footer_lines for report in reading.reports):
        approximations.append("repeated header, footer and page number lines are dropped")
    conversion.details["approximations"] = approximations
    conversion.issues.append(CONVERSION_APPROXIMATED.issue("; ".join(approximations), conversion.input_path.name))
    return reading


def match_pdf_page(document, reading) -> None:
    if reading.page_size_points is None:
        return
    width, height = reading.page_size_points
    section = document.sections[0]
    section.page_width, section.page_height = Pt(width), Pt(height)
    left = reading.text_left if reading.text_left is not None else MAXIMUM_MARGIN_POINTS
    right = width - reading.text_right if reading.text_right is not None else MAXIMUM_MARGIN_POINTS
    section.left_margin = Pt(clamped_margin(left))
    section.right_margin = Pt(clamped_margin(right))


def clamped_margin(points: float) -> float:
    return min(max(points, MINIMUM_MARGIN_POINTS), MAXIMUM_MARGIN_POINTS)


def workbook_to_text(conversion: Conversion) -> None:
    delimiter = DELIMITERS[normalized_extension(conversion.output_path.suffix)]
    written, issues = workbook_to_delimited(conversion.input_path, conversion.output_path, delimiter, conversion.sheet)
    conversion.written = written
    conversion.issues.extend(issues)




def legacy_workbook(conversion: Conversion) -> None:
    conversion.issues.extend(legacy_workbook_to_xlsx(conversion.input_path, conversion.output_path))


def pdf_to_workbook(conversion: Conversion) -> None:
    found = read_pdf_tables(conversion.input_path, conversion.password, scanned_page_lines(conversion))
    if not found.tables:
        raise OfficeFailure(TABLE_NOT_FOUND.issue(f"{conversion.input_path.name} has no table on any of its pages", conversion.input_path.name))
    write_pdf_workbook(found.tables, conversion.output_path, conversion.input_path.stem)
    conversion.details["tables"] = [table_details(table) for table in found.tables]
    if found.pages_without_tables:
        listed = ", ".join(str(number) for number in found.pages_without_tables)
        conversion.issues.append(CONVERSION_APPROXIMATED.issue(f"page {listed} has no table and was left out" if len(found.pages_without_tables) == 1 else f"pages {listed} have no table and were left out", conversion.input_path.name))
    if found.pages_without_text:
        conversion.issues.append(page_without_text_issue(conversion, found.pages_without_text, "their tables were not read"))


def write_docx(conversion: Conversion, blocks: list, source_directory: Path, save: bool = True):
    document, issues = markdown_document(blocks, DEFAULT_DOCUMENT_FONT, BODY_SIZE_POINTS, source_directory)
    conversion.issues.extend(issues)
    if save:
        save_document(document, conversion.output_path)
    return document


def read_text(path: Path) -> str:
    return read_text_input(str(path))


def save_media(media: dict, directory: Path, conversion: Conversion) -> None:
    if not media:
        return
    directory.mkdir(parents=True, exist_ok=True)
    for name, blob in media.items():
        (directory / name).write_bytes(blob)
    conversion.details["mediaDirectory"] = str(directory)


def report_dropped(conversion: Conversion, dropped: dict) -> None:
    if dropped:
        listed = ", ".join(f"{count} {kind}" for kind, count in dropped.items())
        conversion.issues.append(CONVERSION_APPROXIMATED.issue(f"dropped {listed}; layout, fonts and colors are not carried", conversion.input_path.name))


def with_data_source(block, media: dict):
    if not isinstance(block, Image):
        return block
    name = Path(block.source).name
    return replace(block, source=file_data_uri(media[name], name)) if name in media else block


def decoded(block, media_directory: Path):
    if not isinstance(block, Image) or not block.source.startswith("data:") or ";base64," not in block.source:
        return block
    header, payload = block.source.split(",", 1)
    extension = mimetypes.guess_extension(header[5:].split(";", 1)[0]) or ".png"
    path = media_directory / f"image{len(list(media_directory.iterdir())) + 1}{extension}"
    path.write_bytes(base64.b64decode(payload))
    return replace(block, source=str(path))


CONVERTERS = {
    ("docx", "pdf"): docx_to_pdf,
    ("xlsx", "pdf"): workbook_to_pdf,
    ("pptx", "pdf"): presentation_to_pdf,
    ("docx", "md"): docx_to_markdown,
    ("docx", "html"): docx_to_html,
    ("html", "docx"): html_to_docx,
    ("html", "md"): html_to_markdown,
    ("html", "pdf"): html_to_pdf,
    ("pdf", "docx"): pdf_to_docx,
    ("pdf", "md"): pdf_to_markdown,
    ("pdf", "pptx"): pdf_to_presentation,
    ("xlsx", "csv"): workbook_to_text,
    ("xlsx", "tsv"): workbook_to_text,
    ("xls", "xlsx"): legacy_workbook,
    ("ods", "xlsx"): legacy_workbook,
    ("xlsb", "xlsx"): legacy_workbook,
    ("pdf", "xlsx"): pdf_to_workbook,
}



if __name__ == "__main__":
    raise SystemExit(run_command(main))
