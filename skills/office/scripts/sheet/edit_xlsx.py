#!/usr/bin/env python3
from __future__ import annotations

from openpyxl import load_workbook

from cell_values import typed_cell_value
from documents_folder import resolve_document_path
from office_result import OfficeArgumentParser, Result, read_json_file, run_command
from office_schema import require_valid
from sheet_definitions import ROWS


def main() -> Result:
    arguments = parse_arguments()
    json_rows = load_rows(arguments.rows) if arguments.rows else []
    workbook_path = resolve_document_path(arguments.workbook_path, "xlsx")
    workbook = load_workbook(workbook_path)
    worksheet = resolve_worksheet(workbook, arguments.sheet)
    for row_string in arguments.row:
        worksheet.append(parse_row(row_string))
    for row in json_rows:
        worksheet.append(row)
    workbook.save(workbook_path)
    return Result(summary=f"appended rows to {workbook_path}", output_path=workbook_path)


def load_rows(rows_path: str) -> list[list]:
    rows = read_json_file(rows_path)
    require_valid(ROWS, rows, "rows")
    return rows


def parse_row(row_string: str) -> list:
    return [typed_cell_value(cell.strip()) for cell in row_string.split(",")]


def resolve_worksheet(workbook, sheet_name: str | None):
    if sheet_name is None:
        return workbook.active
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return workbook.create_sheet(title=sheet_name)


def parse_arguments():
    parser = OfficeArgumentParser(description="Append rows to an existing XLSX workbook in place.")
    parser.add_argument("workbook_path", nargs="?", help="Path to the .xlsx; defaults to the newest .xlsx in ~/documents")
    parser.add_argument("--sheet", default=None, metavar="NAME", help="Sheet name (default: active sheet; created if missing)")
    parser.add_argument("--row", action="append", default=[], metavar="CELLS", help="Append one row; comma-separated cell values (repeatable)")
    parser.add_argument("--rows", metavar="JSON_PATH", help="JSON file with an array of row arrays")
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
