from __future__ import annotations

from dataclasses import dataclass, field
import datetime

from openpyxl.formula.tokenizer import Token, Tokenizer, TokenizerError
from openpyxl.utils import get_column_letter

from cell_values import typed_text
from excel_functions import EXCEL_FUNCTIONS
from formula_grammar import formula_problem
from formula_tree import function_name
from office_result import Issue
from sheet_definitions import VALUE_STORED_AS_TEXT
from sheet_styling import header_row_index

SHOWN_TEXT_LIMIT = 3


@dataclass(frozen=True)
class Reading:
    value: object
    number_format: str | None
    noun: str


@dataclass
class Run:
    sheet: str
    column: int
    first_row: int
    number_format: str | None
    noun: str
    texts: list[str] = field(default_factory=list)
    values: list[object] = field(default_factory=list)

    @property
    def last_row(self) -> int:
        return self.first_row + len(self.values) - 1

    def continues(self, column: int, row: int, reading: Reading) -> bool:
        return (self.column, self.last_row + 1, self.number_format, self.noun) == (column, row, reading.number_format, reading.noun)


def text_value_issues(workbook) -> list[Issue]:
    return [run_issue(run) for worksheet in workbook.worksheets for run in text_value_runs(worksheet)]


def text_value_runs(worksheet) -> list[Run]:
    header_row = header_row_index(worksheet)
    numeric_columns = columns_holding_numbers(worksheet, header_row)
    runs: list[Run] = []
    for column_cells in worksheet.iter_cols(min_row=header_row + 1):
        for cell in column_cells:
            reading = text_reading(cell, cell.column in numeric_columns)
            if reading is None:
                continue
            if not runs or not runs[-1].continues(cell.column, cell.row, reading):
                runs.append(Run(worksheet.title, cell.column, cell.row, reading.number_format, reading.noun))
            runs[-1].texts.append(cell.value)
            runs[-1].values.append(reading.value)
    return runs


def columns_holding_numbers(worksheet, header_row: int) -> set[int]:
    return {
        cell.column
        for row in worksheet.iter_rows(min_row=header_row + 1)
        for cell in row
        if cell.data_type == "f" or isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
    }


def text_reading(cell, in_numeric_column: bool) -> Reading | None:
    if cell.data_type != "s" or not isinstance(cell.value, str) or cell.quotePrefix:
        return None
    formula = formula_missing_equals(cell.value)
    if formula is not None:
        return Reading(formula, None, "a formula missing its =")
    value, number_format = typed_text(cell.value)
    if isinstance(value, datetime.date):
        return Reading(value.isoformat(), None, "a date")
    if not isinstance(value, (int, float)) or number_format is None and not in_numeric_column:
        return None
    return Reading(value, number_format, "a number")


def formula_missing_equals(text: str) -> str | None:
    formula = "=" + text.strip()
    try:
        tokens = Tokenizer(formula).items
    except (TokenizerError, IndexError):
        return None
    first = tokens[0] if tokens else None
    if first is None or first.type != Token.FUNC or function_name(first.value[:-1]) not in EXCEL_FUNCTIONS:
        return None
    return formula if formula_problem(formula) is None else None


def run_issue(run: Run) -> Issue:
    letter = get_column_letter(run.column)
    first, last = f"{letter}{run.first_row}", f"{letter}{run.last_row}"
    reference = first if first == last else f"{first}:{last}"
    location = f"{run.sheet}!{reference}"
    shown = ", ".join(repr(text) for text in run.texts[:SHOWN_TEXT_LIMIT]) + (" and more" if len(run.texts) > SHOWN_TEXT_LIMIT else "")
    verb = "holds the text" if len(run.texts) == 1 else "hold the texts"
    return VALUE_STORED_AS_TEXT.issue(f"{location} {verb} {shown}, which reads as {run.noun}", location, suggested_operations(run, first, reference))


def suggested_operations(run: Run, first: str, reference: str) -> list[dict]:
    if len(run.values) == 1:
        operations = [{"op": "set_cell", "sheet": run.sheet, "cell": first, "value": run.values[0]}]
    else:
        operations = [{"op": "set_range", "sheet": run.sheet, "cell": first, "values": [[value] for value in run.values]}]
    if run.number_format is not None:
        operations.append({"op": "format_range", "sheet": run.sheet, "range": reference, "numberFormat": run.number_format})
    return operations
