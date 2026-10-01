from __future__ import annotations

from copy import copy
import datetime
import re

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel

from office_operations import OPERATION_NOT_APPLICABLE, Change
from office_result import OfficeFailure
from workbook_access import column_index, parse_range, resolve_sheet, sheet_of
from workbook_snapshot import computed_values


NUMBER_RANK = 0
TEXT_RANK = 1
LOGICAL_RANK = 2
ERROR_RANK = 3
BLANK_RANK = 4


def sort_key(value: object) -> tuple:
    if value is None or value == "":
        return (BLANK_RANK, 0)
    if isinstance(value, bool):
        return (LOGICAL_RANK, int(value))
    if isinstance(value, (int, float)):
        return (NUMBER_RANK, value)
    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return (NUMBER_RANK, to_excel(value))
    text = str(value)
    if text.startswith("#"):
        return (ERROR_RANK, text)
    return (TEXT_RANK, text.casefold())


def ordered(rows: list, column: int, descending: bool) -> list:
    blank = [row for row in rows if sort_key(row[column]["value"])[0] == BLANK_RANK]
    filled = [row for row in rows if sort_key(row[column]["value"])[0] != BLANK_RANK]
    return sorted(filled, key=lambda row: sort_key(row[column]["value"]), reverse=descending) + blank


def plan_sort_range(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    min_row, min_column, max_row, max_column = parse_range(operation["range"], f"{location}.range")
    keys = [(column_index(operation["by"], f"{location}.by"), operation.get("descending", False))]
    if operation.get("thenBy"):
        keys.append((column_index(operation["thenBy"], f"{location}.thenBy"), operation.get("thenDescending", False)))
    for index, (column, _) in enumerate(keys):
        if not min_column <= column <= max_column:
            name = "by" if index == 0 else "thenBy"
            raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.{name}: column {get_column_letter(column)} is outside {operation['range'].upper()}", f"{location}.{name}"))
    first_row = min_row + 1 if operation.get("hasHeader", True) else min_row
    merged = [str(item) for item in worksheet.merged_cells.ranges if not (item.max_row < first_row or item.min_row > max_row or item.max_col < min_column or item.min_col > max_column)]
    if merged:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: sorting cannot move merged {', '.join(merged)}; unmerge first", f"{location}.range"))

    def change() -> str:
        values = computed_values(workbook, worksheet, (first_row, min_column, max_row, max_column))
        rows = [snapshot_row(worksheet, row, min_column, max_column, values) for row in range(first_row, max_row + 1)]
        for column, descending in reversed(keys):
            rows = ordered(rows, column - min_column, descending)
        for offset, row in enumerate(rows):
            place_row(worksheet, row, first_row + offset, min_column)
        return f"sorted {len(rows)} rows of {worksheet.title}!{operation['range'].upper()} by {operation['by'].upper()}"
    return change


def snapshot_row(worksheet, row: int, min_column: int, max_column: int, values: dict) -> list:
    snapshot = []
    for column in range(min_column, max_column + 1):
        cell = worksheet.cell(row=row, column=column)
        shown = values.get(cell.coordinate, cell.value) if cell.data_type == "f" else cell.value
        snapshot.append({"raw": cell.value, "value": shown, "style": copy(cell._style), "link": cell.hyperlink, "comment": cell.comment, "row": row, "coordinate": cell.coordinate})
    return snapshot


def place_row(worksheet, snapshot: list, row: int, min_column: int) -> None:
    for offset, item in enumerate(snapshot):
        cell = worksheet.cell(row=row, column=min_column + offset)
        raw = item["raw"]
        if isinstance(raw, str) and raw.startswith("=") and row != item["row"]:
            raw = Translator(raw, origin=item["coordinate"]).translate_formula(cell.coordinate)
        cell.value = raw
        cell._style = item["style"]
        cell.hyperlink = item["link"]
        cell.comment = copy(item["comment"]) if item["comment"] is not None else None


def plan_find_replace(workbook, operation: dict, location: str) -> Change:
    worksheets = [resolve_sheet(workbook, operation["sheet"], f"{location}.sheet")] if operation.get("sheet") else list(workbook.worksheets)
    scope = operation.get("in", "values")
    pattern = re.compile(("^" + re.escape(operation["find"]) + "$") if operation.get("wholeCell") else re.escape(operation["find"]), 0 if operation.get("matchCase") else re.IGNORECASE)

    def change() -> str:
        changed = 0
        for worksheet in worksheets:
            for cell in list(worksheet._cells.values()):
                if not isinstance(cell.value, str) or not wanted(cell, scope):
                    continue
                replaced = pattern.sub(lambda _: operation["replace"], cell.value)
                if replaced != cell.value:
                    cell.value = replaced
                    changed += 1
        return f"replaced {operation['find']!r} in {changed} cells"
    return change


def wanted(cell, scope: str) -> bool:
    is_formula = cell.data_type == "f"
    return scope == "all" or (scope == "formulas") == is_formula
