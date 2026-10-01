from __future__ import annotations

from copy import copy
from dataclasses import dataclass
import datetime
import re

from openpyxl.formula.translate import Translator
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import FilterColumn, Filters
from openpyxl.worksheet.formula import ArrayFormula

from core.office_operations import OPERATION_NOT_APPLICABLE, Change
from core.office_result import OfficeFailure
from sheet.workbook_access import column_index, parse_cell, parse_range, resolve_sheet, sheet_of
from sheet.workbook_snapshot import computed_values


TRAILING_NUMBER = re.compile(r"^(.*?)(\d+)$")


def range_label(bounds: tuple[int, int, int, int]) -> str:
    min_row, min_column, max_row, max_column = bounds
    return f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{max_row}"


def formula_of(cell) -> str | None:
    if isinstance(cell.value, ArrayFormula):
        return cell.value.text
    return cell.value if cell.data_type == "f" else None


def moved_formula(formula: str, origin: str, destination: str) -> str:
    return Translator(formula, origin=origin).translate_formula(destination)


def copy_cell(source, target, with_value: bool = True, with_style: bool = True) -> None:
    if with_value:
        formula = formula_of(source)
        target.value = moved_formula(formula, source.coordinate, target.coordinate) if formula else source.value
    if with_style:
        target._style = copy(source._style)


def series_value(value: object, steps: int, step: float) -> object:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value + step * steps
    if isinstance(value, datetime.datetime):
        return value + datetime.timedelta(days=step * steps)
    if isinstance(value, datetime.date):
        return value + datetime.timedelta(days=int(step * steps))
    match = TRAILING_NUMBER.match(value) if isinstance(value, str) and not value.startswith("=") else None
    if match:
        number = int(match.group(2)) + int(step * steps)
        return f"{match.group(1)}{str(number).zfill(len(match.group(2)))}"
    return None


def plan_fill_range(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    min_row, min_column, max_row, max_column = parse_range(operation["range"], f"{location}.range")
    downward = operation.get("direction", "down") == "down"
    if (max_row if downward else max_column) == (min_row if downward else min_column):
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: {operation['range'].upper()} has nothing to fill; it needs the first {'row' if downward else 'column'} and the cells after it", f"{location}.range"))
    series, step = operation.get("series", False), operation.get("step", 1)

    def change() -> str:
        lines = range(min_column, max_column + 1) if downward else range(min_row, max_row + 1)
        for line in lines:
            source = worksheet.cell(row=min_row, column=line) if downward else worksheet.cell(row=line, column=min_column)
            targets = range(min_row + 1, max_row + 1) if downward else range(min_column + 1, max_column + 1)
            for steps, position in enumerate(targets, start=1):
                target = worksheet.cell(row=position, column=line) if downward else worksheet.cell(row=line, column=position)
                copy_cell(source, target)
                stepped = series_value(source.value, steps, step) if series else None
                if stepped is not None:
                    target.value = stepped
        kind = "a series" if series else "copies"
        return f"filled {worksheet.title}!{operation['range'].upper()} {'down' if downward else 'right'} with {kind} of its first {'row' if downward else 'column'}"
    return change


def plan_copy_range(workbook, operation: dict, location: str) -> Change:
    source_sheet = sheet_of(workbook, operation, location)
    target_sheet = resolve_sheet(workbook, operation.get("toSheet") or source_sheet.title, f"{location}.toSheet")
    min_row, min_column, max_row, max_column = parse_range(operation["range"], f"{location}.range")
    top, left = parse_cell(operation["to"], f"{location}.to")
    paste = operation.get("paste", "all")
    row_offset, column_offset = top - min_row, left - min_column

    def change() -> str:
        values = computed_values(workbook, source_sheet, (min_row, min_column, max_row, max_column)) if paste == "values" else {}
        snapshot = [CellSnapshot.of(source_sheet.cell(row=row, column=column), values) for row in range(min_row, max_row + 1) for column in range(min_column, max_column + 1)]
        for source in snapshot:
            source.paste_into(target_sheet.cell(row=source.row + row_offset, column=source.column + column_offset), paste)
        copy_merges(source_sheet, target_sheet, (min_row, min_column, max_row, max_column), row_offset, column_offset)
        destination = range_label((top, left, max_row + row_offset, max_column + column_offset))
        return f"copied {paste} of {source_sheet.title}!{operation['range'].upper()} to {target_sheet.title}!{destination}"
    return change


@dataclass
class CellSnapshot:
    row: int
    column: int
    coordinate: str
    value: object
    formula: str | None
    computed: object
    style: object

    @classmethod
    def of(cls, cell, values: dict) -> "CellSnapshot":
        return cls(cell.row, cell.column, cell.coordinate, cell.value, formula_of(cell), values.get(cell.coordinate), copy(cell._style))

    def paste_into(self, target, paste: str) -> None:
        if paste == "values":
            target.value = self.computed if self.formula else self.value
            return
        if paste == "all":
            target.value = moved_formula(self.formula, self.coordinate, target.coordinate) if self.formula else self.value
        target._style = copy(self.style)


def copy_merges(source_sheet, target_sheet, bounds: tuple[int, int, int, int], row_offset: int, column_offset: int) -> None:
    min_row, min_column, max_row, max_column = bounds
    for merged in list(source_sheet.merged_cells.ranges):
        if merged.min_row >= min_row and merged.max_row <= max_row and merged.min_col >= min_column and merged.max_col <= max_column:
            target_sheet.merge_cells(start_row=merged.min_row + row_offset, start_column=merged.min_col + column_offset, end_row=merged.max_row + row_offset, end_column=merged.max_col + column_offset)


def plan_clear_range(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    bounds = parse_range(operation["range"], f"{location}.range")
    what = operation.get("what", "contents")

    def change() -> str:
        for row in worksheet.iter_rows(min_row=bounds[0], min_col=bounds[1], max_row=bounds[2], max_col=bounds[3]):
            for cell in row:
                clear_cell(cell, what)
        return f"cleared the {what} of {worksheet.title}!{operation['range'].upper()}"
    return change


def clear_cell(cell, what: str) -> None:
    if what in ("contents", "all"):
        cell.value = None
    if what in ("formats", "all"):
        cell._style = StyleArray()
    if what == "all":
        cell.hyperlink = None
        cell.comment = None


def plan_convert_to_values(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    bounds = parse_range(operation["range"], f"{location}.range")

    def change() -> str:
        values = computed_values(workbook, worksheet, bounds)
        formulas = [cell for row in worksheet.iter_rows(min_row=bounds[0], min_col=bounds[1], max_row=bounds[2], max_col=bounds[3]) for cell in row if formula_of(cell)]
        missing = [cell.coordinate for cell in formulas if cell.coordinate not in values]
        if missing:
            raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: {', '.join(missing[:10])} on {worksheet.title} have no computed value to keep; fix or remove those formulas first", f"{location}.range"))
        for cell in formulas:
            cell.value = values[cell.coordinate]
        return f"replaced {len(formulas)} formulas in {worksheet.title}!{operation['range'].upper()} with their values"
    return change


def plan_set_filter_criteria(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    column = column_index(operation["column"], f"{location}.column")
    if not worksheet.auto_filter.ref:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: {worksheet.title} has no filter; add one with set_auto_filter first", f"{location}.sheet"))
    bounds = parse_range(worksheet.auto_filter.ref, f"{location}.sheet")
    if not bounds[1] <= column <= bounds[3]:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.column: column {operation['column'].upper()} is outside the filter {worksheet.auto_filter.ref}", f"{location}.column"))

    def change() -> str:
        offset = column - bounds[1]
        kept = [item for item in worksheet.auto_filter.filterColumn if item.colId != offset]
        if operation.get("values"):
            kept.append(FilterColumn(colId=offset, filters=Filters(filter=list(operation["values"]))))
        worksheet.auto_filter.filterColumn = kept
        shown = show_matching_rows(workbook, worksheet, bounds)
        if not operation.get("values"):
            return f"cleared the filter on column {operation['column'].upper()} of {worksheet.title}; {shown} rows show"
        return f"filtered column {operation['column'].upper()} of {worksheet.title} to {', '.join(operation['values'])}; {shown} rows show"
    return change


def filter_text(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return "" if value is None else str(value).strip().casefold()


def show_matching_rows(workbook, worksheet, bounds: tuple[int, int, int, int]) -> int:
    min_row, min_column, max_row, max_column = bounds
    body = (min_row + 1, min_column, max_row, max_column)
    values = computed_values(workbook, worksheet, body)
    wanted = {item.colId: {filter_text(value) for value in item.filters.filter} for item in worksheet.auto_filter.filterColumn if item.filters is not None}
    shown = 0
    for row in range(min_row + 1, max_row + 1):
        cells = {offset: worksheet.cell(row=row, column=min_column + offset) for offset in wanted}
        visible = all(filter_text(values.get(cell.coordinate) if formula_of(cell) else cell.value) in wanted[offset] for offset, cell in cells.items())
        worksheet.row_dimensions[row].hidden = not visible
        shown += visible
    return shown
