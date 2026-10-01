from __future__ import annotations

from dataclasses import dataclass, field
import datetime
import inspect
import math
import os

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
from openpyxl.utils.formulas import FORMULAE

from formula_parser import Call, Literal, Postfix, Prefix, Reference, parse_formula
from formula_references import REFERENCE_ERROR, parse_end, reference_parts, unquote_sheet_name
from sheet_functions import FUNCTIONS, LAZY_FUNCTIONS, OPERATORS, POSTFIX_OPERATORS, PREFIX_OPERATORS
import sheet_semantics as semantics
from sheet_semantics import DIVIDE_BY_ZERO, INVALID_NUMBER, UNKNOWN_NAME, ExcelError, Item, Range


NUMBER = "n"
TEXT = "str"
BOOLEAN = "b"
ERROR = "e"
EXACT_INTEGER_LIMIT = 1e15
SIGNATURES = {name: inspect.signature(function) for name, function in FUNCTIONS.items()}


@dataclass(frozen=True)
class CachedValue:
    cell_type: str
    text: str

    @property
    def is_error(self) -> bool:
        return self.cell_type == ERROR


@dataclass
class Evaluation:
    values: dict = field(default_factory=dict)
    not_evaluated: list = field(default_factory=list)


def evaluate_workbook(path: str) -> Evaluation:
    workbook = load_workbook(path)
    evaluator = WorkbookEvaluator(workbook, allows_user_functions=os.path.splitext(path)[1].lower() == ".xlsm")
    evaluation = Evaluation()
    for (title, _, _), cell in evaluator.cells.items():
        if cell.data_type == "f":
            record(evaluation, (title, cell.coordinate), evaluator.computed_value(workbook[title], cell))
    return evaluation


def record(evaluation: Evaluation, key: tuple, value: CachedValue | None) -> None:
    if value is None:
        evaluation.not_evaluated.append(key)
    else:
        evaluation.values[key] = value


def is_unknown_function(name: str) -> bool:
    return name not in FORMULAE and not name.startswith("_XL")


class WorkbookEvaluator:
    def __init__(self, workbook, allows_user_functions: bool):
        self.workbook = workbook
        self.allows_user_functions = allows_user_functions
        self.worksheets = {worksheet.title.casefold(): worksheet for worksheet in workbook.worksheets}
        self.cells = {
            (worksheet.title, cell.row, cell.column): cell
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        }
        self.results = {}
        self.in_progress = set()

    def computed_value(self, worksheet, cell) -> CachedValue | None:
        try:
            return cached_value(self.formula_result(worksheet, cell))
        except (NotImplementedError, RecursionError):
            return None

    def formula_result(self, worksheet, cell) -> Item:
        key = (worksheet.title, cell.coordinate)
        if key in self.results:
            return self.known_result(key)
        if key in self.in_progress:
            raise NotImplementedError("a circular reference is not evaluated")
        self.in_progress.add(key)
        try:
            self.results[key] = self.compute(worksheet, cell)
        except NotImplementedError:
            self.results[key] = None
            raise
        finally:
            self.in_progress.discard(key)
        return self.results[key]

    def known_result(self, key: tuple) -> Item:
        if self.results[key] is None:
            raise NotImplementedError("a cell this formula reads is not evaluated")
        return self.results[key]

    def compute(self, worksheet, cell) -> Item:
        if not isinstance(cell.value, str):
            raise NotImplementedError("array and data table formulas are not evaluated")
        result = semantics.classify(self.evaluate(parse_formula(cell.value), worksheet))
        return Item(semantics.NUMBER, 0.0) if result.kind == semantics.BLANK else result

    def evaluate(self, node, worksheet):
        if isinstance(node, Literal):
            return node.item
        if isinstance(node, Reference):
            return self.reference(node.text, worksheet)
        if isinstance(node, Call):
            return self.call(node, worksheet)
        if isinstance(node, Prefix):
            return outcome(PREFIX_OPERATORS[node.operator], self.evaluate(node.operand, worksheet))
        if isinstance(node, Postfix):
            return outcome(POSTFIX_OPERATORS[node.operator], self.evaluate(node.operand, worksheet))
        return outcome(OPERATORS[node.operator], self.evaluate(node.left, worksheet), self.evaluate(node.right, worksheet))

    def call(self, node: Call, worksheet):
        if is_unknown_function(node.name) and not self.allows_user_functions:
            return Item(semantics.ERROR, ExcelError(UNKNOWN_NAME))
        name = node.name.removeprefix("_XLFN.")
        if name not in FUNCTIONS:
            raise NotImplementedError(f"{node.name} is not evaluated")
        if name in LAZY_FUNCTIONS:
            arguments = [lambda argument=argument: self.evaluate(argument, worksheet) for argument in node.arguments]
        else:
            arguments = [self.evaluate(argument, worksheet) for argument in node.arguments]
        try:
            SIGNATURES[name].bind(*arguments)
        except TypeError as error:
            raise NotImplementedError(f"{name} is given the wrong number of arguments") from error
        return outcome(FUNCTIONS[name], *arguments)

    def reference(self, text: str, worksheet):
        parts = reference_parts(text)
        if parts is None or len(parts) > 2:
            raise NotImplementedError(f"the reference {text} is not evaluated")
        if any(part.rest == REFERENCE_ERROR for part in parts):
            return Item(semantics.ERROR, ExcelError(REFERENCE_ERROR))
        if len(parts) == 1 and parts[0].prefix is None and not is_cell(parse_end(parts[0].rest)):
            return self.defined_name(text, worksheet)
        target = self.target_sheet(parts, worksheet)
        if target is None:
            return Item(semantics.ERROR, ExcelError(REFERENCE_ERROR))
        ends = [parse_end(part.rest) for part in parts]
        if any(end is None for end in ends):
            raise NotImplementedError(f"the reference {text} is not evaluated")
        rows, columns = self.spans(ends, target)
        return Range([[self.cell_item(target, row, column) for column in columns] for row in rows])

    def target_sheet(self, parts: list, worksheet):
        names = {unquote_sheet_name(part.prefix).casefold() for part in parts if part.prefix is not None}
        if len(names) > 1:
            raise NotImplementedError("a reference across sheets is not evaluated")
        if not names:
            return worksheet
        return self.worksheets.get(names.pop())

    def spans(self, ends: list, worksheet) -> tuple[range, range]:
        first, last = ends[0], ends[-1]
        if all(end.row is not None and end.column is not None for end in ends):
            return span(first.row, last.row), span(first.column, last.column)
        if len(ends) == 2 and all(end.row is None for end in ends):
            return span(1, worksheet.max_row), span(first.column, last.column)
        if len(ends) == 2 and all(end.column is None for end in ends):
            return span(first.row, last.row), span(1, worksheet.max_column)
        raise NotImplementedError("a partial row or column reference is not evaluated")

    def defined_name(self, name: str, worksheet):
        defined = self.find_defined_name(name, worksheet)
        if defined is None:
            return Item(semantics.ERROR, ExcelError(UNKNOWN_NAME))
        parts = reference_parts(defined.attr_text or "")
        if not parts or parts[0].prefix is None:
            raise NotImplementedError(f"the defined name {name} is not a reference to a sheet")
        return self.reference(defined.attr_text, worksheet)

    def find_defined_name(self, name: str, worksheet):
        for names in (worksheet.defined_names, self.workbook.defined_names):
            found = next((defined for key, defined in names.items() if key.casefold() == name.casefold()), None)
            if found is not None:
                return found
        return None

    def cell_item(self, worksheet, row: int, column: int) -> Item:
        cell = self.cells.get((worksheet.title, row, column))
        if cell is None:
            return Item(semantics.BLANK, None)
        if cell.data_type == "f":
            return self.formula_result(worksheet, cell)
        return stored_item(cell)


def is_cell(end) -> bool:
    return end is not None and end.row is not None and end.column is not None


def span(first: int, last: int) -> range:
    return range(min(first, last), max(first, last) + 1)


def stored_item(cell) -> Item:
    value = cell.value
    if cell.data_type == "e":
        return Item(semantics.ERROR, ExcelError(str(value)))
    if isinstance(value, bool):
        return Item(semantics.LOGICAL, value)
    if isinstance(value, (int, float)):
        return Item(semantics.NUMBER, float(value))
    if isinstance(value, str):
        return Item(semantics.TEXT, value)
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)):
        return Item(semantics.NUMBER, float(to_excel(value)))
    raise NotImplementedError(f"a {type(value).__name__} cell value is not evaluated")


def outcome(function, *arguments):
    try:
        return as_value(function(*arguments))
    except ExcelError as error:
        return Item(semantics.ERROR, error)
    except ZeroDivisionError:
        return Item(semantics.ERROR, ExcelError(DIVIDE_BY_ZERO))
    except OverflowError:
        return Item(semantics.ERROR, ExcelError(INVALID_NUMBER))


def as_value(result):
    if isinstance(result, (Item, Range)):
        return result
    if isinstance(result, ExcelError):
        return Item(semantics.ERROR, result)
    if isinstance(result, bool):
        return Item(semantics.LOGICAL, result)
    if isinstance(result, (int, float)):
        return number_item(float(result))
    if isinstance(result, str):
        return Item(semantics.TEXT, result)
    raise NotImplementedError(f"a {type(result).__name__} result is not evaluated")


def number_item(number: float) -> Item:
    if math.isnan(number) or math.isinf(number):
        return Item(semantics.ERROR, ExcelError(INVALID_NUMBER))
    return Item(semantics.NUMBER, number)


def cached_value(item: Item) -> CachedValue:
    if item.kind == semantics.ERROR:
        return CachedValue(ERROR, item.value.code)
    if item.kind == semantics.LOGICAL:
        return CachedValue(BOOLEAN, "1" if item.value else "0")
    if item.kind == semantics.TEXT:
        return CachedValue(TEXT, item.value)
    return number_value(item.value)


def number_value(number: float) -> CachedValue:
    if number.is_integer() and abs(number) < EXACT_INTEGER_LIMIT:
        return CachedValue(NUMBER, str(int(number)))
    return CachedValue(NUMBER, repr(number))
