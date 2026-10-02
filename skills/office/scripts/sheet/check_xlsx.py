#!/usr/bin/env python3
from __future__ import annotations

import math
from collections import defaultdict

from openpyxl.utils import range_boundaries

from sheet.formulas.cache import evaluation_issues
from sheet.formulas.names import name_issues, sheet_is_missing
from sheet.formulas.references import formula_references, referenced_sheet_names
from sheet.checks.number_display import displayed_number_width
from core.number_format import displayed
from core.office_arguments import route_arguments
from core.office_result import Issue, Result, run_command
from sheet.checks.chart_references import chart_reference_issues
from sheet.formulas.stale_values import stale_cached_value_issues
from sheet.sheet_definitions import BROKEN_DEFINED_NAME, FORMULA_ERROR, NUMBER_TOO_WIDE, PIVOT_VALUES_EMPTY
from core.text_checks import PLACEHOLDER_LEFT, PLACEHOLDER_PATTERN, text_presence_issues
from sheet.checks.text_values import text_value_issues
from sheet.checks.tables import table_findings
from sheet.workbook.access import open_workbook
from sheet.formulas.evaluation import evaluate_workbook
from sheet.preview.pages import sheet_print_width_issues


DEFAULT_COLUMN_WIDTH = 8.43
LISTED_CELL_LIMIT = 20
WIDTH_MARGIN = 2


def main() -> Result:
    arguments = route_arguments("check", "xlsx")
    workbook = open_workbook(arguments.file)
    evaluation = evaluate_workbook(arguments.file)
    table_issues, details = table_findings(workbook)
    presence = text_presence_issues(visible_text(workbook, evaluation), arguments.required_text, arguments.forbidden_text)
    issues = content_issues(workbook, arguments.file, evaluation) + table_issues + presence
    return Result(summary=f"checked {arguments.file}: {len(issues)} issues", output_path=arguments.file, issues=tuple(issues), details=details)


def content_issues(workbook, workbook_path: str, evaluation) -> list[Issue]:
    named = name_issues(workbook)
    return (
        stale_cached_value_issues(workbook_path, evaluation)
        + named
        + formula_error_issues(evaluation, {issue.location for issue in named})
        + defined_name_issues(workbook)
        + number_width_issues(workbook, evaluation)
        + text_value_issues(workbook)
        + chart_reference_issues(workbook)
        + empty_pivot_issues(workbook)
        + placeholder_issues(workbook)
        + evaluation_issues(evaluation)
        + sheet_print_width_issues(workbook)
    )


def visible_text(workbook, evaluation) -> str:
    return "\n".join(shown_text(cell, worksheet.title, evaluation) for worksheet in workbook.worksheets for row in worksheet.iter_rows() for cell in row if cell.value is not None)


def shown_text(cell, sheet: str, evaluation) -> str:
    if cell.data_type != "f":
        return displayed(cell.value, cell.number_format).text
    value = evaluation.values.get((sheet, cell.coordinate))
    if value is None:
        return ""
    number = number_shown_in(cell, sheet, evaluation)
    return displayed(number, cell.number_format).text if number is not None else value.text


def cell_label(sheet: str, coordinate: str) -> str:
    return f"{sheet}!{coordinate}"


def listed(labels: list[str]) -> str:
    hidden = len(labels) - LISTED_CELL_LIMIT
    return ", ".join(labels[:LISTED_CELL_LIMIT]) + (f" and {hidden} more" if hidden > 0 else "")


def formula_error_issues(evaluation, already_reported: set) -> list[Issue]:
    by_error = defaultdict(list)
    for (sheet, coordinate), value in evaluation.values.items():
        if value.is_error and cell_label(sheet, coordinate) not in already_reported:
            by_error[(sheet, value.text)].append(cell_label(sheet, coordinate))
    return [
        FORMULA_ERROR.issue(f"{len(labels)} formula cells on {sheet} compute {code}: {listed(sorted(labels))}", sorted(labels)[0])
        for (sheet, code), labels in sorted(by_error.items())
    ]


def defined_name_issues(workbook) -> list[Issue]:
    scoped = [(name, defined, "") for name, defined in workbook.defined_names.items()]
    for worksheet in workbook.worksheets:
        scoped.extend((name, defined, worksheet.title) for name, defined in worksheet.defined_names.items())
    return [
        BROKEN_DEFINED_NAME.issue(f"defined name {name} points at {defined.attr_text}", name)
        for name, defined, _ in scoped
        if defined_name_is_broken(workbook, defined.attr_text or "")
    ]


def defined_name_is_broken(workbook, reference_text: str) -> bool:
    references = formula_references("=" + reference_text)
    if "#REF!" in reference_text:
        return True
    return any(sheet_is_missing(workbook, name) for reference in references for name in referenced_sheet_names(reference))


def column_width(worksheet, column: int) -> float:
    for dimension in worksheet.column_dimensions.values():
        if (dimension.min or 1) <= column <= (dimension.max or dimension.min or 1) and dimension.width:
            return dimension.width
    return worksheet.sheet_format.defaultColWidth or DEFAULT_COLUMN_WIDTH


def is_hidden_column(worksheet, column: int) -> bool:
    return any(dimension.hidden and (dimension.min or 1) <= column <= (dimension.max or dimension.min or 1) for dimension in worksheet.column_dimensions.values())


def number_width_issues(workbook, evaluation) -> list[Issue]:
    issues = []
    for worksheet in workbook.worksheets:
        widest = widest_number_per_column(worksheet, evaluation)
        for column, (needed, coordinate) in sorted(widest.items()):
            issues.append(too_wide_issue(worksheet, column, needed, coordinate))
    return issues


def widest_number_per_column(worksheet, evaluation) -> dict:
    merged = {coordinate for merged_range in worksheet.merged_cells.ranges for row in worksheet[merged_range.coord] for coordinate in (cell.coordinate for cell in row)}
    widest = {}
    for row in worksheet.iter_rows():
        for cell in row:
            number = number_shown_in(cell, worksheet.title, evaluation)
            if number is None or cell.coordinate in merged or is_hidden_column(worksheet, cell.column):
                continue
            needed = displayed_number_width(number, cell.number_format)
            if needed > column_width(worksheet, cell.column) and needed > widest.get(cell.column, (0, ""))[0]:
                widest[cell.column] = (needed, cell.coordinate)
    return widest


def number_shown_in(cell, sheet: str, evaluation) -> float | None:
    if cell.data_type == "f":
        value = evaluation.values.get((sheet, cell.coordinate))
        return float(value.text) if value is not None and value.cell_type == "n" else None
    if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float)):
        return None
    return cell.value


def too_wide_issue(worksheet, column: int, needed: int, coordinate: str) -> Issue:
    letter = coordinate.rstrip("0123456789")
    width = math.ceil(needed) + WIDTH_MARGIN
    location = cell_label(worksheet.title, coordinate)
    widening = {"op": "set_column_width", "sheet": worksheet.title, "column": letter, "width": width}
    return NUMBER_TOO_WIDE.issue(f"{location} needs about {needed} characters but column {letter} holds {column_width(worksheet, column):g}", location, fix=[widening])


def empty_pivot_issues(workbook) -> list[Issue]:
    return [empty_pivot_issue(worksheet, pivot) for worksheet in workbook.worksheets for pivot in getattr(worksheet, "_pivots", []) if pivot_values_are_empty(worksheet, pivot)]


def pivot_value_cells(worksheet, pivot) -> list:
    min_column, min_row, max_column, max_row = range_boundaries(pivot.location.ref)
    first_row = min_row + (pivot.location.firstDataRow or 0)
    first_column = min_column + (pivot.location.firstDataCol or 0)
    return [cell for row in worksheet.iter_rows(min_row=first_row, max_row=max_row, min_col=first_column, max_col=max_column) for cell in row]


def pivot_values_are_empty(worksheet, pivot) -> bool:
    cells = pivot_value_cells(worksheet, pivot)
    return bool(cells) and all(cell.value in (None, "") for cell in cells)


def empty_pivot_issue(worksheet, pivot) -> Issue:
    location = cell_label(worksheet.title, pivot.location.ref)
    source = getattr(getattr(pivot.cache, "cacheSource", None), "worksheetSource", None)
    summarized = f" of {source.sheet}!{source.ref}" if source is not None and source.sheet else ""
    fields = ", ".join(repr(field.name) for field in pivot.dataFields if field.name) or "its values"
    return PIVOT_VALUES_EMPTY.issue(f"pivot table {pivot.name}{summarized} at {location} shows {fields} with every value cell empty; the summarized column holds no number", location)


def placeholder_issues(workbook) -> list[Issue]:
    issues = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "s" and isinstance(cell.value, str):
                    issues.extend(PLACEHOLDER_LEFT.issue(f"{cell_label(worksheet.title, cell.coordinate)} still holds {placeholder}", cell_label(worksheet.title, cell.coordinate)) for placeholder in PLACEHOLDER_PATTERN.findall(cell.value))
    return issues



if __name__ == "__main__":
    raise SystemExit(run_command(main))
