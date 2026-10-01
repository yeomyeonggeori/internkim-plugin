from __future__ import annotations

import datetime
import os
import re
from typing import Iterable
import warnings

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from office_inputs import holds_macros, package_stream
from office_operations import TARGET_NOT_FOUND
from office_result import INVALID_VALUE, OfficeFailure
from office_schema import closest_name, guess_text
from excel_functions import PARAMETER_PREFIX
from formula_tree import FUNCTION_PREFIXES, Call, parse_formula, render, tokens_in
from excel_limits import LAST_COLUMN, MAXIMUM_COLUMN, MAXIMUM_ROW, SHEET_LIMITS, column_number

warnings.filterwarnings("ignore", category=UserWarning, module=r"openpyxl\.")


def open_workbook(path: str, data_only: bool = False):
    if not os.path.exists(path):
        raise FileNotFoundError(2, "No such file or directory", path)
    return load_workbook(package_stream(path), data_only=data_only, keep_vba=holds_macros(path))


def resolve_sheet(workbook, sheet_name: str | None, location: str):
    if sheet_name is None:
        return workbook.worksheets[0]
    match = next((worksheet for worksheet in workbook.worksheets if worksheet.title == sheet_name), None)
    if match is None:
        raise OfficeFailure(missing_sheet_issue(workbook.sheetnames, sheet_name, location))
    return match


def missing_sheet_issue(names: list[str], sheet_name: str, location: str):
    nearest = closest_name(sheet_name, names)
    guess = guess_text(nearest)
    message = f"{location}: the workbook has no sheet named {sheet_name!r}{guess}; it has {', '.join(names)}"
    suggestion = f'use "sheet": "{nearest}"' if nearest else f"use one of the sheet names: {', '.join(names)}"
    return TARGET_NOT_FOUND.issue(message, location, suggestion)


CELL_GRAMMAR = re.compile(r"([A-Za-z]*)([0-9]*)")


def sheet_of(workbook, operation: dict, location: str):
    return resolve_sheet(workbook, operation.get("sheet"), f"{location}.sheet")


def column_index(text: str, location: str) -> int:
    letters = text.strip().upper()
    if not letters.isalpha() or not letters.isascii():
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: {text!r} is not a column letter such as C", location, f"write only the column letters, such as {''.join(character for character in letters if character.isascii() and character.isalpha()) or 'C'}"))
    if column_number(letters) > MAXIMUM_COLUMN:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: column {letters} is past {LAST_COLUMN}, the last column Excel has", location, f"use a column from A to {LAST_COLUMN}; {SHEET_LIMITS}"))
    return column_number(letters)


def field_name(location: str) -> str:
    return location.rsplit(".", 1)[-1].split("[", 1)[0]


def refuse_sheet_prefix(text: str, location: str) -> None:
    if "!" not in text:
        return
    sheet, _, cells = text.rpartition("!")
    sheet = sheet.strip("'")
    split = f"pass --sheet {sheet} {location} {cells}" if location.startswith("--") else f'put the sheet in its own field and the cells here: "sheet": "{sheet}", "{field_name(location)}": "{cells}"'
    raise OfficeFailure(INVALID_VALUE.issue(f"{location}: {text!r} names the sheet inside the address, and this field takes only the cells", location, split))


def cell_problems(text: str) -> list[str]:
    match = CELL_GRAMMAR.fullmatch(text)
    if match is None:
        return ["it is not column letters followed by a row number"]
    letters, digits = match.groups()
    problems = []
    if not letters:
        problems.append("the column letters are missing")
    elif column_number(letters) > MAXIMUM_COLUMN:
        problems.append(f"column {letters.upper()} is past {LAST_COLUMN}")
    if not digits:
        problems.append("the row number is missing")
    elif not 1 <= int(digits) <= MAXIMUM_ROW:
        problems.append(f"row {int(digits)} is outside 1 to {MAXIMUM_ROW}")
    return problems


def cell_example(text: str) -> str:
    match = CELL_GRAMMAR.fullmatch(text)
    letters, digits = match.groups() if match else ("", "")
    column = letters.upper() if letters and column_number(letters) <= MAXIMUM_COLUMN else "A"
    row = digits if digits and 1 <= int(digits) <= MAXIMUM_ROW else "1"
    return f"{column}{row}"


def refuse_cell(text: str, location: str, problems: list[str]) -> None:
    raise OfficeFailure(INVALID_VALUE.issue(
        f"{location}: in {text!r}, {' and '.join(problems)}; {SHEET_LIMITS}",
        location,
        f'use a cell such as "{cell_example(text)}"',
    ))


def parse_cell(text: str, location: str) -> tuple[int, int]:
    refuse_sheet_prefix(text, location)
    written = text.strip().replace("$", "")
    problems = cell_problems(written)
    if problems:
        refuse_cell(written, location, problems)
    letters, digits = CELL_GRAMMAR.fullmatch(written).groups()
    return int(digits), column_number(letters)


def parse_range(text: str, location: str) -> tuple[int, int, int, int]:
    refuse_sheet_prefix(text, location)
    corners = text.strip().replace("$", "").upper().split(":")
    if len(corners) > 2 or not all(corners):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: {text!r} is not a range such as A1:D10", location, "write two cells joined by a colon, such as A1:D10"))
    if len(corners) == 2 and (all(corner.isalpha() for corner in corners) or all(corner.isdigit() for corner in corners)):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: {text!r} must name both corners, such as A1:D10", location, f'add the row or column the corners lack, such as "{whole_line_example(corners)}"'))
    for corner in corners:
        problems = cell_problems(corner)
        if problems:
            raise OfficeFailure(INVALID_VALUE.issue(f"{location}: in the corner {corner} of {text!r}, {' and '.join(problems)}; {SHEET_LIMITS}", location, f'write each corner as a cell such as "{cell_example(corner)}"'))
    first_row, first_column = parse_cell(corners[0], location)
    last_row, last_column = parse_cell(corners[-1], location)
    return min(first_row, last_row), min(first_column, last_column), max(first_row, last_row), max(first_column, last_column)


def whole_line_example(corners: list[str]) -> str:
    if all(corner.isalpha() for corner in corners):
        return f"{corners[0]}1:{corners[-1]}100"
    return f"A{corners[0]}:Z{corners[-1]}"


def json_value(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def formula_text(cell) -> str | None:
    if cell.data_type != "f":
        return None
    stored = cell.value.text if isinstance(cell.value, ArrayFormula) else cell.value
    if not isinstance(stored, str):
        return None
    written = stored if stored.startswith("=") else "=" + stored
    nodes = parse_formula(written)
    if nodes is None:
        return written
    for token in tokens_in(nodes):
        if token.value.lower().startswith(PARAMETER_PREFIX):
            token.value = token.value[len(PARAMETER_PREFIX):]
    return "=" + render(nodes, without_storage_prefix)


def without_storage_prefix(call: Call) -> str | None:
    if not call.name.upper().startswith(FUNCTION_PREFIXES):
        return None
    return f"{call.function}(" + ",".join(render(argument, without_storage_prefix) for argument in call.arguments) + ")"


def cell_rows(worksheet, bounds: tuple[int, int, int, int]) -> Iterable:
    min_row, min_column, max_row, max_column = bounds
    return worksheet.iter_rows(min_row=min_row, min_col=min_column, max_row=max_row, max_col=max_column)
