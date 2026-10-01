from __future__ import annotations

from dataclasses import dataclass, field
import datetime
import re

from openpyxl.styles.fonts import DEFAULT_FONT
from openpyxl.utils import get_column_letter, range_boundaries

from core.number_format import Displayed, displayed
from render.office_preview import PageGeometry, Preview, escaped, page_section, pixels, positioned, style_attribute
from core.units import emu_to_pixels, inches_to_pixels, points_to_pixels
from fonts.registry import OFFICE_KOREAN_FAMILY
from fonts.preview import FontRegistry, FontRequest, css_font_family, draws_scripts_apart, script_font_family, script_runs
from sheet.sheet_formatting import STACKED_ROTATION, rotation_degrees
from sheet.sheet_objects import EXCEL_DEFAULT_FIT_PAGES
from sheet.xlsx_colors import css_color
from sheet.xlsx_conditional import ConditionalStyles
from sheet.xlsx_preview_charts import chart_html, chart_kind, chart_title, drawing_box, image_html, is_whole
from core.page_sizes import DEFAULT_PAPER, PAPER_BY_SPREADSHEET_CODE


DEFAULT_COLUMN_CHARACTERS = 8.43
DEFAULT_ROW_POINTS = 15
MAXIMUM_DIGIT_PIXELS = 7
CELL_PADDING_PIXELS = 2
INDENT_PIXELS = 9
GRIDLINE = "1px solid #d0d0d0"
BORDER_STYLES = {
    "thin": ("1px", "solid"), "hair": ("1px", "solid"), "medium": ("2px", "solid"), "thick": ("3px", "solid"), "double": ("3px", "double"),
    "dashed": ("1px", "dashed"), "dotted": ("1px", "dotted"), "mediumDashed": ("2px", "dashed"), "dashDot": ("1px", "dashed"),
    "mediumDashDot": ("2px", "dashed"), "dashDotDot": ("1px", "dotted"), "mediumDashDotDot": ("2px", "dotted"), "slantDashDot": ("2px", "dashed"),
}
HORIZONTAL = {"left": "flex-start", "center": "center", "right": "flex-end", "centerContinuous": "center", "justify": "flex-start", "fill": "flex-start", "distributed": "center"}
VERTICAL = {"top": "flex-start", "center": "center", "bottom": "flex-end", "justify": "flex-start", "distributed": "center"}
HEADER_CODES = (("&P", "{page}"), ("&N", "{pages}"), ("&A", "{sheet}"), ("&F", "{file}"), ("&D", "{date}"), ("&T", ""), ("&B", ""), ("&I", ""), ("&U", ""), ("&&", "&"))
HEADER_FORMATTING = re.compile(r'&K[0-9A-Fa-f]{6}|&"[^"]*"|&\d+|&[SXYEOHG]')


@dataclass
class SheetFrame:
    worksheet: object
    values: object
    columns: list[int]
    rows: list[int]
    widths: dict[int, float]
    heights: dict[int, float]
    merges: dict[tuple[int, int], tuple[int, int]]
    covered: set
    title_rows: list[int] = field(default_factory=list)


class SheetPreviewer:
    def __init__(self, workbook, values, file_name: str, palette: tuple, fonts: FontRegistry, preview: Preview):
        self.workbook = workbook
        self.values = values
        self.file_name = file_name
        self.palette = palette
        self.fonts = fonts
        self.preview = preview
        self.page_painted = False
        self.page_drawings: list[dict] = []
        self.page_contents: list[dict] = []

    def sheet_pages(self, worksheet, sheet_values) -> list[tuple[PageGeometry, list[int], list[int], float, SheetFrame]]:
        frame = self.frame(worksheet, sheet_values)
        if frame is None:
            return []
        geometry, scale = self.page_geometry(worksheet, frame)
        column_chunks = chunks(frame.columns, [frame.widths[column] * scale for column in frame.columns], geometry.content_width, set())
        title_height = sum(frame.heights[row] for row in frame.title_rows) * scale
        breaks = {brk.id for brk in worksheet.row_breaks.brk} if worksheet.row_breaks else set()
        body_rows = [row for row in frame.rows if row not in frame.title_rows]
        row_chunks = chunks(body_rows, [frame.heights[row] * scale for row in body_rows], geometry.height - geometry.margin_top - geometry.margin_bottom - title_height, breaks)
        if len(column_chunks) > 1:
            self.preview.approximate(f"{worksheet.title}: columns printed across {len(column_chunks)} pages wide")
        return [(geometry, columns, frame.title_rows + rows, scale, frame) for columns in column_chunks for rows in row_chunks]

    def frame(self, worksheet, sheet_values) -> SheetFrame | None:
        bounds = print_bounds(worksheet, sheet_values)
        if bounds is None:
            return None
        min_column, min_row, max_column, max_row = bounds
        column_settings = column_dimension_map(worksheet)
        columns = [column for column in range(min_column, max_column + 1) if not column_settings.get(column, {}).get("hidden")]
        rows = [row for row in range(min_row, max_row + 1) if not row_hidden(worksheet, row)]
        widths = {column: column_pixels(column_settings.get(column, {}).get("width") or default_column_width(worksheet)) for column in columns}
        merges, covered = merge_maps(worksheet)
        # Excel's row AutoFit ignores merged cells (Microsoft support, "Change the column width or row height in Excel")
        heights = {row: self.row_height(worksheet, sheet_values, row, [column for column in columns if (row, column) not in covered], widths) for row in rows}
        title_rows = print_title_rows(worksheet, rows)
        return SheetFrame(worksheet, sheet_values, columns, rows, widths, heights, merges, covered, title_rows)

    def row_height(self, worksheet, sheet_values, row: int, columns: list[int], widths: dict) -> float:
        dimension = worksheet.row_dimensions.get(row) if hasattr(worksheet.row_dimensions, "get") else None
        if dimension is not None and dimension.height:
            return points_to_pixels(dimension.height)
        default = points_to_pixels(worksheet.sheet_format.defaultRowHeight or DEFAULT_ROW_POINTS)
        wrapped = [column for column in columns if worksheet.cell(row=row, column=column).alignment.wrap_text and sheet_values.cell(row=row, column=column).value is not None]
        if not wrapped:
            return default
        return max(default, *(self.wrapped_height(worksheet, sheet_values, row, column, widths[column]) for column in wrapped))

    def wrapped_height(self, worksheet, sheet_values, row: int, column: int, width: float) -> float:
        cell = worksheet.cell(row=row, column=column)
        text = displayed(sheet_values.cell(row=row, column=column).value, cell.number_format).text
        request = font_request(cell.font)
        lines = 0
        for paragraph in text.split("\n"):
            lines += max(1, int(self.fonts.width(request, paragraph) // max(width - 2 * CELL_PADDING_PIXELS, 1)) + 1)
        return lines * self.fonts.line_height(request) + 2

    def page_geometry(self, worksheet, frame: SheetFrame) -> tuple[PageGeometry, float]:
        setup = worksheet.page_setup
        width_inches, height_inches = PAPER_BY_SPREADSHEET_CODE.get(int(setup.paperSize or DEFAULT_PAPER.spreadsheet_code), DEFAULT_PAPER).inches
        if setup.orientation == "landscape":
            width_inches, height_inches = height_inches, width_inches
        margins = worksheet.page_margins
        geometry = PageGeometry(
            width=inches_to_pixels(width_inches),
            height=inches_to_pixels(height_inches),
            margin_top=inches_to_pixels(margins.top if margins.top is not None else 0.75),
            margin_right=inches_to_pixels(margins.right if margins.right is not None else 0.7),
            margin_bottom=inches_to_pixels(margins.bottom if margins.bottom is not None else 0.75),
            margin_left=inches_to_pixels(margins.left if margins.left is not None else 0.7),
            header_distance=inches_to_pixels(margins.header if margins.header is not None else 0.3),
            footer_distance=inches_to_pixels(margins.footer if margins.footer is not None else 0.3),
        )
        return geometry, print_scale(worksheet, frame, geometry)

    def page_html(self, page_number: int, page_count: int, geometry: PageGeometry, columns: list[int], rows: list[int], scale: float, frame: SheetFrame) -> str:
        grid = self.grid_html(frame, columns, rows, scale)
        if not self.page_painted:
            self.preview.blank_pages.append(f"page {page_number} ({frame.worksheet.title})")
        self.page_contents.append(page_content(page_number, frame, columns, rows, self.page_drawings))
        grid_width = sum(frame.widths[column] for column in columns) * scale
        left = geometry.margin_left + max(0.0, (geometry.content_width - grid_width) / 2) if frame.worksheet.print_options.horizontalCentered else geometry.margin_left
        parts = [positioned(left, geometry.margin_top, grid_width, grid)]
        substitutions = {"page": page_number, "pages": page_count, "sheet": frame.worksheet.title, "file": self.file_name, "date": datetime.date.today().isoformat()}
        parts.extend(self.header_footer_html(frame.worksheet, geometry, substitutions))
        return page_section(page_number, geometry, "".join(parts))

    def header_footer_html(self, worksheet, geometry: PageGeometry, substitutions: dict) -> list[str]:
        parts = []
        for block, top in ((worksheet.oddHeader, geometry.header_distance), (worksheet.oddFooter, geometry.height - geometry.footer_distance - 16)):
            pieces = [(getattr(block, side).text if getattr(block, side) is not None else None, alignment) for side, alignment in (("left", "left"), ("center", "center"), ("right", "right"))]
            for text, alignment in pieces:
                if text:
                    content = escaped(header_text(text).format(**substitutions))
                    parts.append(positioned(geometry.margin_left, top, geometry.content_width, content, {"text-align": alignment, "font-size": "12px", "font-family": css_font_family(DEFAULT_FONT.name, OFFICE_KOREAN_FAMILY), "white-space": "pre"}))
        return parts

    def grid_html(self, frame: SheetFrame, columns: list[int], rows: list[int], scale: float) -> str:
        column_index = {column: index for index, column in enumerate(columns)}
        row_index = {row: index for index, row in enumerate(rows)}
        conditional = ConditionalStyles(frame.worksheet, frame.values, self.palette)
        backgrounds, texts = [], []
        for row in rows:
            for column in columns:
                if (row, column) in frame.covered and (row, column) not in frame.merges:
                    continue
                background, text = self.cell_html(frame, row, column, columns, rows, column_index, row_index, scale, conditional)
                backgrounds.append(background)
                texts.append(text)
        overlays = self.drawings_html(frame, columns, rows, scale)
        self.page_painted = bool(overlays) or any(backgrounds) or any(texts)
        declarations = {
            "position": "relative",
            "display": "grid",
            "grid-template-columns": " ".join(pixels(frame.widths[column] * scale) for column in columns),
            "grid-template-rows": " ".join(pixels(frame.heights[row] * scale) for row in rows),
            "width": pixels(sum(frame.widths[column] for column in columns) * scale),
        }
        return f"<div{style_attribute(declarations)}>{''.join(backgrounds)}{''.join(texts)}{overlays}</div>"

    def cell_html(self, frame: SheetFrame, row: int, column: int, columns: list[int], rows: list[int], column_index: dict, row_index: dict, scale: float, conditional: ConditionalStyles) -> tuple[str, str]:
        worksheet = frame.worksheet
        cell = worksheet.cell(row=row, column=column)
        value = frame.values.cell(row=row, column=column).value
        last_row, last_column = frame.merges.get((row, column), (row, column))
        row_span = sum(1 for candidate in rows if row <= candidate <= last_row)
        column_span = sum(1 for candidate in columns if column <= candidate <= last_column)
        placement = {"grid-column": f"{column_index[column] + 1} / span {max(column_span, 1)}", "grid-row": f"{row_index[row] + 1} / span {max(row_span, 1)}"}
        extra = conditional.style_for(row, column, value)
        background = {**placement, "box-sizing": "border-box", "background": extra.get("background") or fill_color(cell, self.palette), **self.border_css(frame, row, column, last_row, last_column, scale)}
        bar = extra.get("bar")
        bar_html = f'<div{style_attribute({"width": f"{bar[0]:.1f}%", "height": "70%", "background": bar[1], "margin-top": "2px"})}></div>' if bar else ""
        text = displayed(value, cell.number_format)
        width = sum(frame.widths[candidate] for candidate in columns if column <= candidate <= last_column) * scale
        painted = bar_html or any(name.startswith("border") or (name == "background" and value_set) for name, value_set in background.items())
        background_html = f"<div{style_attribute(background)}>{bar_html}</div>" if painted else ""
        return background_html, self.text_html(cell, text, value, placement, width, scale, extra, frame, row, column, columns)

    def text_html(self, cell, text: Displayed, value, placement: dict, width: float, scale: float, extra: dict, frame: SheetFrame, row: int, column: int, columns: list[int]) -> str:
        if not text.text:
            return ""
        font = cell.font
        request = font_request(font, scale)
        alignment = cell.alignment
        horizontal = alignment.horizontal or ("right" if text.is_number else "center" if isinstance(value, bool) else "left")
        content = stacked(text.text) if alignment.textRotation == STACKED_ROTATION else text.text
        self.fonts.use(request, content)
        wraps = bool(alignment.wrap_text)
        if text.is_number and not wraps and self.fonts.width(request, content) > width - 2 * CELL_PADDING_PIXELS * scale:
            content = "#" * max(1, int((width - 2 * CELL_PADDING_PIXELS * scale) // max(self.fonts.width(request, "#"), 1)))
        span = placement
        if not wraps and not text.is_number and horizontal == "left" and (row, column) not in frame.merges:
            span = {**placement, "grid-column": f"{placement['grid-column'].split(' / ')[0]} / span {spill_span(frame, row, column, columns)}"}
        declarations = {
            **span,
            "display": "flex",
            "justify-content": HORIZONTAL.get(horizontal, "flex-start"),
            "align-items": "center" if rotation_style(alignment.textRotation) else VERTICAL.get(alignment.vertical or "bottom", "flex-end"),
            "padding": f"0 {pixels(CELL_PADDING_PIXELS * scale)}",
            "padding-left": pixels((CELL_PADDING_PIXELS + INDENT_PIXELS * (alignment.indent or 0)) * scale) if alignment.indent else None,
            "box-sizing": "border-box",
            "overflow": "hidden",
            "white-space": "pre-wrap" if wraps else "pre",
            "font-family": css_font_family(request.latin, request.east_asia),
            "font-size": pixels(points_to_pixels(request.size)),
            "font-weight": "700" if font.b else None,
            "font-style": "italic" if font.i else None,
            "color": extra.get("color") or text.color or css_color(font.color, self.palette),
            "line-height": pixels(self.fonts.line_height(request)),
            "text-align": horizontal if horizontal in ("left", "center", "right") else None,
        }
        span = {"text-decoration": text_decoration(font), **rotation_style(alignment.textRotation)}
        return f"<div{style_attribute(declarations)}><span{style_attribute(span)}>{script_spans(request, content)}</span></div>"

    def border_css(self, frame: SheetFrame, row: int, column: int, last_row: int, last_column: int, scale: float) -> dict:
        worksheet = frame.worksheet
        first = worksheet.cell(row=row, column=column).border
        last = worksheet.cell(row=last_row, column=last_column).border
        sides = {"top": first.top, "left": first.left, "bottom": last.bottom, "right": last.right}
        declarations = {}
        gridlines = bool(frame.worksheet.print_options.gridLines)
        for side, border in sides.items():
            value = border_value(border, self.palette)
            if value is None and side in ("left", "top"):
                neighbor = neighbor_border(worksheet, row, column, side)
                if neighbor is not None:
                    continue
            if value is None and gridlines and side in ("right", "bottom"):
                value = GRIDLINE
            if value is not None:
                declarations[f"border-{side}"] = value
        return declarations

    def drawings_html(self, frame: SheetFrame, columns: list[int], rows: list[int], scale: float) -> str:
        body_rows = [row for row in rows if row not in frame.title_rows]
        page_width = sum(frame.widths[column] for column in columns) * scale
        body_height = sum(frame.heights[row] for row in body_rows) * scale
        title_height = sum(frame.heights[row] for row in rows if row in frame.title_rows) * scale
        self.page_drawings = []
        html = []
        if getattr(frame.worksheet, "_charts", []):
            self.fonts.width(FontRequest(OFFICE_KOREAN_FAMILY, OFFICE_KOREAN_FAMILY, 10), "가")
        for chart in getattr(frame.worksheet, "_charts", []):
            box = drawing_box(chart.anchor, frame, columns, body_rows, scale)
            if box is not None:
                html.append(chart_html(chart, box, self.values, self.palette, self.preview))
                self.page_drawings.append({"chart": chart_title(chart) or None, "type": chart_kind(chart), "shown": "whole" if is_whole(box, page_width, body_height) else "part, cut at the page edge"})
        for image in getattr(frame.worksheet, "_images", []):
            box = drawing_box(image.anchor, frame, columns, body_rows, scale, image)
            if box is not None:
                html.append(image_html(image, box))
                self.page_drawings.append({"image": True, "shown": "whole" if is_whole(box, page_width, body_height) else "part, cut at the page edge"})
        if not html:
            return ""
        clip = {"position": "absolute", "left": "0", "top": pixels(title_height), "width": pixels(page_width), "height": pixels(body_height), "overflow": "hidden"}
        return f"<div{style_attribute(clip)}>{''.join(html)}</div>"


def page_content(page_number: int, frame: SheetFrame, columns: list[int], rows: list[int], drawings: list[dict]) -> dict:
    body_rows = [row for row in rows if row not in frame.title_rows] or rows
    cells = f"{get_column_letter(columns[0])}{body_rows[0]}:{get_column_letter(columns[-1])}{body_rows[-1]}"
    content = {"page": page_number, "sheet": frame.worksheet.title, "cells": cells}
    titles = [row for row in rows if row in frame.title_rows]
    if titles:
        content["repeatedRows"] = f"{titles[0]}:{titles[-1]}"
    if drawings:
        content["drawings"] = drawings
    return content


def text_decoration(font) -> str | None:
    return " ".join(name for name, flag in (("underline", font.u and font.u != "none"), ("line-through", font.strike)) if flag) or None


def stacked(text: str) -> str:
    return "\n".join(text)


def rotation_style(rotation: int | None) -> dict:
    if not rotation or rotation == STACKED_ROTATION:
        return {}
    return {"display": "inline-block", "transform": f"rotate({-rotation_degrees(rotation)}deg)"}


def font_request(font, scale: float = 1.0) -> FontRequest:
    return FontRequest(font.name or DEFAULT_FONT.name, OFFICE_KOREAN_FAMILY, (font.sz or DEFAULT_FONT.sz) * scale, bool(font.b))


def fill_color(cell, palette: tuple) -> str | None:
    fill = cell.fill
    if getattr(fill, "patternType", None) != "solid":
        return None
    return css_color(fill.fgColor, palette)


def border_value(side, palette: tuple) -> str | None:
    if side is None or not side.style:
        return None
    width, style = BORDER_STYLES.get(side.style, ("1px", "solid"))
    return f"{width} {style} {css_color(side.color, palette) or '#000000'}"


def neighbor_border(worksheet, row: int, column: int, side: str):
    if side == "left" and column > 1:
        border = worksheet.cell(row=row, column=column - 1).border.right
    elif side == "top" and row > 1:
        border = worksheet.cell(row=row - 1, column=column).border.bottom
    else:
        return None
    return border if border is not None and border.style else None


def spill_span(frame: SheetFrame, row: int, column: int, columns: list[int]) -> int:
    span = 1
    for candidate in columns[columns.index(column) + 1:]:
        if frame.values.cell(row=row, column=candidate).value not in (None, "") or (row, candidate) in frame.covered:
            break
        span += 1
    return span


def print_bounds(worksheet, sheet_values) -> tuple[int, int, int, int] | None:
    area = worksheet.print_area
    if area:
        first = area[0] if isinstance(area, (list, tuple)) else str(area).split(",")[0]
        return range_boundaries(first.split("!")[-1].replace("$", "").replace("'", ""))
    used = [(cell.row, cell.column) for row in sheet_values.iter_rows() for cell in row if cell.value not in (None, "")]
    used += [(merged.max_row, merged.max_col) for merged in worksheet.merged_cells.ranges]
    used += drawing_corners(worksheet)
    if not used:
        return None
    return min(column for _, column in used), min(row for row, _ in used), max(column for _, column in used), max(row for row, _ in used)


def drawing_spans(worksheet) -> list[tuple[int, int, int, int]]:
    spans = []
    for drawing in [*getattr(worksheet, "_charts", []), *getattr(worksheet, "_images", [])]:
        start = getattr(drawing.anchor, "_from", None)
        if start is None:
            continue
        end = getattr(drawing.anchor, "to", None)
        extent = getattr(drawing.anchor, "ext", None)
        if end is not None:
            spans.append((start.row + 1, start.col + 1, end.row + 1, end.col + 1))
        elif extent is not None:
            rows = int(emu_to_pixels(extent.height) // points_to_pixels(DEFAULT_ROW_POINTS)) + 1
            columns = int(emu_to_pixels(extent.width) // column_pixels(DEFAULT_COLUMN_CHARACTERS)) + 1
            spans.append((start.row + 1, start.col + 1, start.row + rows, start.col + columns))
    return spans


def drawing_corners(worksheet) -> list[tuple[int, int]]:
    return [corner for top, left, bottom, right in drawing_spans(worksheet) for corner in ((top, left), (bottom, right))]


def column_dimension_map(worksheet) -> dict[int, dict]:
    settings = {}
    for dimension in worksheet.column_dimensions.values():
        start = dimension.min or 0
        end = dimension.max or start
        for column in range(start, end + 1):
            settings[column] = {"width": dimension.width if dimension.customWidth or dimension.width else None, "hidden": dimension.hidden}
    return settings


def row_hidden(worksheet, row: int) -> bool:
    dimension = worksheet.row_dimensions.get(row) if hasattr(worksheet.row_dimensions, "get") else None
    return bool(dimension is not None and dimension.hidden)


def default_column_width(worksheet) -> float:
    return worksheet.sheet_format.defaultColWidth or DEFAULT_COLUMN_CHARACTERS


def column_pixels(characters: float) -> float:
    return int(((256 * characters + int(128 / MAXIMUM_DIGIT_PIXELS)) / 256) * MAXIMUM_DIGIT_PIXELS)


def merge_maps(worksheet) -> tuple[dict, set]:
    merges, covered = {}, set()
    for merged in worksheet.merged_cells.ranges:
        merges[(merged.min_row, merged.min_col)] = (merged.max_row, merged.max_col)
        covered.update((row, column) for row in range(merged.min_row, merged.max_row + 1) for column in range(merged.min_col, merged.max_col + 1))
    return merges, covered


def print_title_rows(worksheet, rows: list[int]) -> list[int]:
    titles = worksheet.print_title_rows
    if not titles:
        return []
    start, _, end = titles.replace("$", "").partition(":")
    return [row for row in rows if int(start) <= row <= int(end or start)]


def print_scale(worksheet, frame: SheetFrame, geometry: PageGeometry) -> float:
    setup = worksheet.page_setup
    properties = worksheet.sheet_properties.pageSetUpPr
    if properties is not None and properties.fitToPage:
        total_width = sum(frame.widths.values())
        total_height = sum(frame.heights.values())
        width_pages = EXCEL_DEFAULT_FIT_PAGES if setup.fitToWidth is None else int(setup.fitToWidth)
        height_pages = EXCEL_DEFAULT_FIT_PAGES if setup.fitToHeight is None else int(setup.fitToHeight)
        scales = [1.0]
        if width_pages:
            scales.append(geometry.content_width * width_pages / total_width)
        if height_pages:
            scales.append((geometry.height - geometry.margin_top - geometry.margin_bottom) * height_pages / total_height)
        return min(scales)
    return (setup.scale or 100) / 100


def chunks(items: list[int], sizes: list[float], limit: float, breaks_after: set) -> list[list[int]]:
    groups, current, used = [], [], 0.0
    for item, size in zip(items, sizes):
        if current and used + size > limit + 0.5:
            groups.append(current)
            current, used = [], 0.0
        current.append(item)
        used += size
        if item in breaks_after:
            groups.append(current)
            current, used = [], 0.0
    if current:
        groups.append(current)
    return groups or [[]]


def header_text(text: str) -> str:
    cleaned = HEADER_FORMATTING.sub("", text.replace("{", "{{").replace("}", "}}"))
    for code, replacement in HEADER_CODES:
        cleaned = cleaned.replace(code, replacement)
    return cleaned


def script_spans(request: FontRequest, text: str) -> str:
    if not draws_scripts_apart(request):
        return escaped(text)
    return "".join(f'<span{style_attribute({"font-family": script_font_family(request, piece)})}>{escaped(piece)}</span>' for piece in script_runs(text))
