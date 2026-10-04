from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import get_column_letter, range_boundaries

from core.office_result import OfficeFailure
from core.source_snapshot import read_source
from sheet.workbook_declaration import COMPILED_CELLS


SHOWN_CHANGES = 5


@dataclass(frozen=True)
class CompiledCells:
    workbook_path: str
    declaration: str
    ranges: tuple[tuple[str, str], ...]

    def values(self, workbook) -> dict:
        values = {}
        for sheet_name, cell_range in self.ranges:
            if sheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet_name]
            first_column, first_row, last_column, last_row = range_boundaries(cell_range)
            for row in range(first_row, last_row + 1):
                for column in range(first_column, last_column + 1):
                    values[(sheet_name, row, column)] = worksheet.cell(row, column).value
        return values

    def require_unchanged(self, before: dict, workbook) -> None:
        missing = sorted({sheet_name for sheet_name, _ in self.ranges if sheet_name not in workbook.sheetnames})
        changed = [f"{sheet_name}!{get_column_letter(column)}{row}" for (sheet_name, row, column), value in before.items() if sheet_name in workbook.sheetnames and workbook[sheet_name].cell(row, column).value != value]
        if not missing and not changed:
            return
        touched = [f"sheet {name}" for name in missing] + changed
        shown = ", ".join(touched[:SHOWN_CHANGES]) + (f" and {len(touched) - SHOWN_CHANGES} more" if len(touched) > SHOWN_CHANGES else "")
        raise OfficeFailure(COMPILED_CELLS.issue(
            f"{shown} of {self.workbook_path} {'is' if len(touched) == 1 else 'are'} compiled from {self.declaration}; nothing was written",
            touched[0],
            f"change {self.declaration} and run office create {self.workbook_path} {self.declaration} again; a null input stays blank for the person, and apply may still format compiled cells or write outside them",
        ))


def compiled_cells_beside(workbook_path: str) -> CompiledCells | None:
    source = read_source(workbook_path)
    if not source.get("compiled"):
        return None
    ranges = tuple((entry["sheet"], entry["range"]) for entry in source["compiled"])
    return CompiledCells(workbook_path, str(source.get("declaration", "")), ranges)
