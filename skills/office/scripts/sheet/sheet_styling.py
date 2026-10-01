from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from display_width import display_width


HEADER_FILL_COLOR = "DCEAF7"
HEADING_FILL_COLOR = "EAF3F8"
BORDER_COLOR = "CBD5E1"
HEADING_FONT_SIZE = 14
MINIMUM_TABLE_ROWS = 2
MINIMUM_TABLE_COLUMNS = 2
MINIMUM_DATA_ROWS = 10
MINIMUM_COLUMN_WIDTH = 10
MAXIMUM_COLUMN_WIDTH = 48
TITLE_ROWS_ABOVE_HEADER = 2


def is_filled(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def non_blank_count(values) -> int:
    return sum(1 for value in values if is_filled(value))


def header_row_index(worksheet, bounds: tuple[int, int, int, int] | None = None) -> int:
    min_row, min_column, max_row, max_column = bounds or (1, 1, worksheet.max_row, worksheet.max_column)
    last_candidate = min(max_row, min_row + TITLE_ROWS_ABOVE_HEADER)
    rows = worksheet.iter_rows(min_row=min_row, max_row=last_candidate, min_col=min_column, max_col=max_column, values_only=True)
    for row_number, values in enumerate(rows, start=min_row):
        if non_blank_count(values) >= MINIMUM_TABLE_COLUMNS:
            return row_number
    return min_row


def data_bounds(worksheet) -> tuple[int, int, int, int] | None:
    cells = [cell for cell in worksheet._cells.values() if is_filled(cell.value)]
    if not cells:
        return None
    rows = [cell.row for cell in cells]
    columns = [cell.column for cell in cells]
    return min(rows), min(columns), max(rows), max(columns)


def touches(outer: tuple[int, int, int, int], block: tuple[int, int, int, int]) -> bool:
    first_row, first_column, last_row, last_column = outer
    return block[0] <= last_row + 1 and block[2] >= first_row - 1 and block[1] <= last_column + 1 and block[3] >= first_column - 1


def default_border() -> Border:
    side = Side(style="thin", color=BORDER_COLOR)
    return Border(left=side, right=side, top=side, bottom=side)


def style_data_cell(cell) -> None:
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = default_border()
    if isinstance(cell.value, bool):
        return
    if isinstance(cell.value, int):
        cell.number_format = "#,##0"
    if isinstance(cell.value, float):
        cell.number_format = "#,##0.00"


def style_header_cell(cell) -> None:
    style_data_cell(cell)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor=HEADER_FILL_COLOR)


def style_heading_row(worksheet) -> None:
    for cell in worksheet[1]:
        cell.font = Font(bold=True, size=HEADING_FONT_SIZE)
        cell.fill = PatternFill("solid", fgColor=HEADING_FILL_COLOR)
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        cell.border = Border()


def style_cell(cell, header_row: int) -> None:
    if cell.row == header_row:
        style_header_cell(cell)
    else:
        style_data_cell(cell)


def style_table(worksheet, has_heading: bool) -> None:
    header_row = 2 if has_heading else 1
    for row in worksheet.iter_rows():
        for cell in row:
            style_cell(cell, header_row)
    if has_heading:
        style_heading_row(worksheet)
    fit_column_widths(worksheet, first_row=header_row)


def merged_bounds(existing: tuple[int, int, int, int] | None, block: tuple[int, int, int, int]) -> tuple[int, int, int, int] | None:
    if existing is None:
        return block
    if not touches(existing, block):
        return None
    return min(existing[0], block[0]), min(existing[1], block[1]), max(existing[2], block[2]), max(existing[3], block[3])


def style_written_cells(worksheet, cells: list, existing: tuple[int, int, int, int] | None) -> None:
    cells = [cell for cell in cells if is_filled(cell.value)]
    if not cells:
        return
    block = (min(cell.row for cell in cells), min(cell.column for cell in cells), max(cell.row for cell in cells), max(cell.column for cell in cells))
    region = merged_bounds(existing, block)
    if region is None or region[2] - region[0] + 1 < MINIMUM_TABLE_ROWS or region[3] - region[1] + 1 < MINIMUM_TABLE_COLUMNS:
        return
    header_row = header_row_index(worksheet)
    for cell in cells:
        if not cell.has_style:
            style_cell(cell, header_row)


def fit_column_widths(worksheet, first_row: int) -> None:
    for column_cells in worksheet.iter_cols(min_row=first_row):
        content_width = max(display_width(cell.value) for cell in column_cells)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(content_width + 2, MINIMUM_COLUMN_WIDTH), MAXIMUM_COLUMN_WIDTH)
