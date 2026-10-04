from __future__ import annotations

import csv
from dataclasses import dataclass, field
from decimal import Decimal
import io
import os
from pathlib import Path as FilePath

from openpyxl.utils import get_column_letter

from core.office_inputs import read_text_input
from core.office_result import OfficeFailure
from core.office_schema import require_valid
from schemas.expression import Binary, Call, Negate, Number, Path, parse_expression, referenced_paths
from schemas.typed_values import parse_number
from sheet.workbook_declaration import DECLARATION_INVALID, WORKBOOK_DECLARATION


NUMBER_TYPES = ("amount", "quantity", "percent")
CURRENCY_FORMATS = {"USD": '"$"#,##0', "EUR": '"€"#,##0', "GBP": '"£"#,##0', "JPY": '"¥"#,##0'}
TOTAL_WORDS = {"ko": "합계", "en": "Total"}
PARTIAL_MARK = "*"
PARTIAL_NOTES = {"ko": "* 주어진 값만 더했습니다. 주어지지 않은 값: {labels}", "en": "* Only the values given are summed. Not given: {labels}"}
COVERAGE_NOTES = {"ko": "범위가 다릅니다: {labels}", "en": "Unequal coverage: {labels}"}
COVERAGE_LABELS = {"ko": "{member} ({dimension} {most}개 중 {covered}개)", "en": "{member} ({covered} of {most} {dimension})"}
INPUT_FILL = "FFF2CC"
CHANGE_FUNCTIONS = ("change", "percentChange")


@dataclass(frozen=True)
class Column:
    name: str
    type: str
    role: str
    unit: str
    letter: str


@dataclass
class SourceTable:
    name: str
    columns: list
    rows: list

    def column(self, name: str, location: str) -> Column:
        found = next((column for column in self.columns if column.name == name), None)
        if found is None:
            raise invalid(f"{location}: table {self.name!r} has no column {name!r}; it has {', '.join(column.name for column in self.columns)}", location)
        return found

    def range(self, name: str) -> str:
        letter = self.column(name, name).letter
        return f"{quoted(self.name)}!${letter}$2:${letter}${len(self.rows) + 1}"

    def index(self, name: str) -> int:
        return next(index for index, column in enumerate(self.columns) if column.name == name)

    def members(self, dimension: str) -> list:
        return distinct(row[self.index(dimension)] for row in self.rows if row[self.index(dimension)] is not None)

    def members_where(self, dimension: str, filter_dimension: str, filter_member) -> list:
        return distinct(row[self.index(dimension)] for row in self.rows if row[self.index(filter_dimension)] == filter_member and row[self.index(dimension)] is not None)

    def combinations(self, dimensions: list[str]) -> list[tuple]:
        return distinct(tuple(row[self.index(name)] for name in dimensions) for row in self.rows)


@dataclass
class Placed:
    sheet: str
    title_row: int
    header_row: int
    body_rows: list
    total_row: int | None
    label_count: int
    base_columns: list
    total_column: int | None
    added: list = field(default_factory=list)
    cells: dict = field(default_factory=dict)
    formats: dict = field(default_factory=dict)
    partial: set = field(default_factory=set)
    note_rows: list = field(default_factory=list)
    width: int = 0


def invalid(message: str, location: str, suggestion: str | None = None) -> OfficeFailure:
    return OfficeFailure(DECLARATION_INVALID.issue(message, location, suggestion))


def quoted(sheet_name: str) -> str:
    return "'" + sheet_name.replace("'", "''") + "'"


def distinct(values) -> list:
    seen, ordered = set(), []
    for value in values:
        if value not in seen:
            seen.add(value)
            ordered.append(value)
    return ordered


def plain_number(number: Decimal) -> int | float:
    return int(number) if number == number.to_integral_value() else float(number)


def declared_column(entry: dict, index: int) -> Column:
    column_type = entry.get("type", "text")
    role = entry.get("role") or ("measure" if column_type in NUMBER_TYPES else "dimension")
    return Column(entry["name"], column_type, role, entry.get("unit", ""), get_column_letter(index + 1))


def typed_cell(column: Column, value: object, location: str):
    if value is None or isinstance(value, str) and not value.strip():
        return None
    if column.type in NUMBER_TYPES:
        return plain_number(parse_number(value, location))
    if isinstance(value, str) and value.strip().lstrip("-").isdigit():
        return int(value.strip())
    return value.strip() if isinstance(value, str) else value


def delimited_records(csv_path: str) -> list[list[str]]:
    text = read_text_input(csv_path)
    first_line = text.split("\n", 1)[0]
    delimiter = "\t" if csv_path.lower().endswith(".tsv") or "\t" in first_line and "," not in first_line else ","
    return list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))[1:]


def source_rows(entry: dict, columns: list, location: str) -> list:
    records = delimited_records(entry["csvPath"]) if entry.get("csvPath") else entry.get("rows") or []
    rows = []
    for row_index, record in enumerate(records):
        if len(record) != len(columns):
            raise invalid(f"{location}.rows[{row_index}]: {len(record)} cells for {len(columns)} columns", f"{location}.rows[{row_index}]")
        rows.append([typed_cell(column, value, f"{location}.rows[{row_index}][{index}]") for index, (column, value) in enumerate(zip(columns, record))])
    return rows


def source_table(entry: dict, location: str) -> SourceTable:
    columns = [declared_column(column, index) for index, column in enumerate(entry["columns"])]
    return SourceTable(entry["name"], columns, source_rows(entry, columns, location))


def marked_partial(number_format_text: str) -> str:
    return f'{number_format_text}"{PARTIAL_MARK}"'


def missing_labels(table: SourceTable, measures: list) -> list[str]:
    dimensions = [column for column in table.columns if column.role == "dimension"]
    labels = []
    for row in table.rows:
        for measure in measures:
            if row[table.index(measure.name)] is None:
                labels.append(" ".join([*(str(row[table.index(dimension.name)]) for dimension in dimensions if row[table.index(dimension.name)] is not None), measure.name]))
    return labels


def number_format(value_type: str, unit: str, has_fraction: bool = False) -> str:
    if value_type == "percent":
        return "0.0%"
    if value_type == "date":
        return "yyyy-mm-dd"
    if value_type == "amount" and unit.upper() in CURRENCY_FORMATS:
        return CURRENCY_FORMATS[unit.upper()] + (".00" if has_fraction else "")
    if value_type in ("amount", "quantity"):
        return "#,##0.00" if has_fraction else "#,##0"
    return "General"


def header_text(column: Column) -> str:
    if column.type in ("amount", "quantity") and column.unit and column.unit.casefold() not in column.name.casefold():
        return f"{column.name} ({column.unit})"
    return column.name


class Compiler:
    def __init__(self, declaration: dict):
        self.declaration = declaration
        self.language = declaration.get("language", "ko")
        self.tables = [source_table(entry, f"tables[{index}]") for index, entry in enumerate(declaration["tables"])]
        self.next_row: dict[str, int] = {}
        self.views: list[Placed] = []
        self.blanks: list[dict] = []
        self.chart_blocks: list[tuple[str, str]] = []

    def table(self, name: str | None, location: str) -> SourceTable:
        if name is None:
            return self.tables[0]
        found = next((table for table in self.tables if table.name == name), None)
        if found is None:
            raise invalid(f"{location}: no table {name!r}", location)
        return found

    def dimension(self, table: SourceTable, name: str, location: str) -> Column:
        column = table.column(name, location)
        if column.role != "dimension":
            raise invalid(f"{location}: {name!r} is a measure, and this place takes a dimension", location)
        return column

    def measure(self, table: SourceTable, name: str, location: str) -> Column:
        column = table.column(name, location)
        if column.role != "measure":
            raise invalid(f"{location}: {name!r} is a dimension, and this place takes a measure", location)
        return column

    def compile_view(self, entry: dict, location: str) -> Placed:
        table = self.table(entry.get("table"), f"{location}.table")
        row_dimensions = [self.dimension(table, name, f"{location}.rows") for name in entry["rows"]]
        column_dimension = self.dimension(table, entry["columns"], f"{location}.columns") if entry.get("columns") else None
        if column_dimension is not None:
            if not entry.get("measure"):
                raise invalid(f"{location}.measure: a view with columns names the one measure its cells show", f"{location}.measure")
            measures = [self.measure(table, entry["measure"], f"{location}.measure")]
            base = [("member", member) for member in table.members(column_dimension.name)]
        else:
            if not entry.get("measures"):
                raise invalid(f"{location}.measures: a view without columns names its measures", f"{location}.measures")
            measures = [self.measure(table, name, f"{location}.measures") for name in entry["measures"]]
            base = [("measure", measure.name) for measure in measures]
        added_entries = entry.get("add") or []
        shares = [share_dimension(added, f"{location}.add[{index}]") for index, added in enumerate(added_entries)]
        column_name = column_dimension.name if column_dimension is not None else None
        needs_totals = bool(entry.get("totals")) or any(share not in (None, column_name) for share in shares)
        needs_total_column = bool(entry.get("totalColumn")) or column_name is not None and column_name in shares
        sheet = entry["sheet"]
        title_row = self.next_row.get(sheet, 1)
        header_row = title_row + 1
        combinations = table.combinations([dimension.name for dimension in row_dimensions])
        body_rows = [header_row + 1 + index for index in range(len(combinations))]
        total_row = body_rows[-1] + 1 if needs_totals and body_rows else None
        label_count = len(row_dimensions)
        total_column = label_count + len(base) + 1 if needs_total_column and column_dimension is not None else None
        placed = Placed(sheet, title_row, header_row, body_rows, total_row, label_count, base, total_column)
        context = ViewContext(self, table, placed, row_dimensions, column_dimension, measures, combinations)
        context.write()
        for index, added in enumerate(added_entries):
            context.add_column(added, f"{location}.add[{index}]")
        placed.width = max(placed.cells, key=lambda key: key[1])[1] if placed.cells else 1
        last_row = total_row or (body_rows[-1] if body_rows else header_row)
        for note in self.view_notes(table, placed, measures, compared_dimensions(column_dimension, added_entries, location)):
            last_row += 1
            placed.note_rows.append(last_row)
            placed.cells[(last_row, 1)] = note
        self.next_row[sheet] = last_row + 3
        return placed

    def view_notes(self, table: SourceTable, placed: Placed, measures: list, compared: list[str]) -> list[str]:
        notes = []
        if placed.partial:
            notes.append(PARTIAL_NOTES.get(self.language, PARTIAL_NOTES["en"]).format(labels=", ".join(missing_labels(table, measures))))
        uneven = coverage_labels(table, compared, COVERAGE_LABELS.get(self.language, COVERAGE_LABELS["en"]))
        if uneven:
            notes.append(COVERAGE_NOTES.get(self.language, COVERAGE_NOTES["en"]).format(labels=", ".join(uneven)))
        return notes


class ViewContext:
    def __init__(self, compiler: Compiler, table: SourceTable, placed: Placed, row_dimensions: list, column_dimension: Column | None, measures: list, combinations: list):
        self.compiler = compiler
        self.table = table
        self.placed = placed
        self.row_dimensions = row_dimensions
        self.column_dimension = column_dimension
        self.measures = measures
        self.combinations = combinations

    def cell(self, row: int, column: int) -> str:
        return f"{get_column_letter(column)}{row}"

    def put(self, row: int, column: int, value, number_format: str | None = None) -> None:
        self.placed.cells[(row, column)] = value
        if number_format:
            self.placed.formats[(row, column)] = number_format

    def base_column_index(self, position: int) -> int:
        return self.placed.label_count + position + 1

    def write(self) -> None:
        placed = self.placed
        self.put(placed.title_row, 1, None)
        for index, dimension in enumerate(self.row_dimensions):
            self.put(placed.header_row, index + 1, dimension.name)
        for position, (kind, value) in enumerate(placed.base_columns):
            header = value if kind == "member" else header_text(self.table.column(value, value))
            self.put(placed.header_row, self.base_column_index(position), header)
        if placed.total_column:
            self.put(placed.header_row, placed.total_column, TOTAL_WORDS.get(self.compiler.language, "Total"))
        for row, combination in zip(placed.body_rows, self.combinations):
            for index, member in enumerate(combination):
                self.put(row, index + 1, member)
            self.write_row(row, self.row_criteria(row))
        if placed.total_row:
            self.put(placed.total_row, 1, TOTAL_WORDS.get(self.compiler.language, "Total"))
            self.write_row(placed.total_row, [])

    def members_at(self, row: int, column: int | None) -> dict:
        members = {}
        if row != self.placed.total_row and row in self.placed.body_rows:
            combination = self.combinations[self.placed.body_rows.index(row)]
            members.update({dimension.name: member for dimension, member in zip(self.row_dimensions, combination)})
        if column is not None and self.base_kind(column) == "member":
            members[self.column_dimension.name] = self.placed.cells[(self.placed.header_row, column)]
        return members

    def coverage(self, measure: Column, members: dict) -> tuple[int, int]:
        matching = [row for row in self.table.rows if all(row[self.table.index(name)] == member for name, member in members.items())]
        return len(matching), sum(1 for row in matching if row[self.table.index(measure.name)] is not None)

    def is_partial(self, measure: Column, members: dict) -> bool:
        rows, values = self.coverage(measure, members)
        return 0 < values < rows

    def put_marked(self, row: int, column: int, formula: str, number_format_text: str, is_partial: bool) -> None:
        self.put(row, column, formula, marked_partial(number_format_text) if is_partial else number_format_text)
        if is_partial:
            self.placed.partial.add((row, column))

    def row_criteria(self, row: int) -> list[tuple[str, str]]:
        return [(self.table.range(dimension.name), f"${get_column_letter(index + 1)}{row}") for index, dimension in enumerate(self.row_dimensions)]

    def member_criteria(self, column: int) -> list[tuple[str, str]]:
        return [(self.table.range(self.column_dimension.name), f"{get_column_letter(column)}${self.placed.header_row}")]

    def criteria_for(self, row: int, column: int) -> list[tuple[str, str]]:
        row_part = self.row_criteria(row) if row != self.placed.total_row else []
        kind = self.base_kind(column)
        column_part = self.member_criteria(column) if kind == "member" else []
        return row_part + column_part

    def base_kind(self, column: int) -> str:
        position = column - self.placed.label_count - 1
        if 0 <= position < len(self.placed.base_columns):
            return self.placed.base_columns[position][0]
        return "total"

    def measure_at(self, column: int) -> Column:
        position = column - self.placed.label_count - 1
        if 0 <= position < len(self.placed.base_columns) and self.placed.base_columns[position][0] == "measure":
            return self.table.column(self.placed.base_columns[position][1], "measure")
        return self.measures[0]

    def write_row(self, row: int, row_criteria: list) -> None:
        for position in range(len(self.placed.base_columns)):
            column = self.base_column_index(position)
            measure = self.measure_at(column)
            partial = self.is_partial(measure, self.members_at(row, column))
            self.put_marked(row, column, aggregate_formula(self.table.range(measure.name), self.criteria_for(row, column)), number_format(measure.type, measure.unit), partial)
        if self.placed.total_column:
            measure = self.measures[0]
            partial = self.is_partial(measure, self.members_at(row, None))
            self.put_marked(row, self.placed.total_column, aggregate_formula(self.table.range(measure.name), row_criteria), number_format(measure.type, measure.unit), partial)

    def next_free_column(self) -> int:
        return max(column for _, column in self.placed.cells) + 1

    def add_column(self, entry: dict, location: str) -> None:
        expression = parse_expression(entry["expression"], f"{location}.expression")
        if isinstance(expression, Call) and expression.function in CHANGE_FUNCTIONS + ("share",):
            self.add_function_columns(entry, expression, location)
            return
        if self.column_dimension is not None:
            raise invalid(f"{location}.expression: a view with columns adds only percentChange, change or share; arithmetic over measures belongs in a view with measures", f"{location}.expression")
        self.add_arithmetic_column(entry, expression, location)

    def function_dimension(self, expression: Call, location: str) -> str:
        if len(expression.arguments) != 2 or not all(isinstance(argument, Path) for argument in expression.arguments):
            raise invalid(f"{location}.expression: {expression.function} takes a measure and a dimension", f"{location}.expression")
        self.compiler.measure(self.table, expression.arguments[0].name, f"{location}.expression")
        return self.compiler.dimension(self.table, expression.arguments[1].name, f"{location}.expression").name

    def add_function_columns(self, entry: dict, expression: Call, location: str) -> None:
        dimension = self.function_dimension(expression, location)
        measure_name = expression.arguments[0].name
        result_type = "percent" if expression.function != "change" else self.table.column(measure_name, measure_name).type
        result_unit = self.table.column(measure_name, measure_name).unit
        targets = self.member_columns() if self.column_dimension is not None else [self.base_column_index(position) for position, (_, name) in enumerate(self.placed.base_columns) if name == measure_name]
        if self.column_dimension is not None and dimension == self.column_dimension.name and expression.function in CHANGE_FUNCTIONS:
            self.add_change_across_columns(entry, expression.function, targets, result_type, result_unit)
        elif self.column_dimension is not None and dimension == self.column_dimension.name:
            self.add_share_of_row_total(entry, targets)
        elif dimension in [dimension.name for dimension in self.row_dimensions] and expression.function == "share":
            self.add_share_of_column_total(entry, targets)
        elif dimension in [dimension.name for dimension in self.row_dimensions]:
            self.add_change_down_rows(entry, expression.function, dimension, targets, result_type, result_unit)
        else:
            raise invalid(f"{location}.expression: {dimension!r} is neither this view's columns nor one of its rows", f"{location}.expression")

    def member_columns(self) -> list[int]:
        return [self.base_column_index(position) for position, (kind, _) in enumerate(self.placed.base_columns) if kind == "member"]

    def header_for(self, entry: dict, member, several: bool) -> str:
        return f"{entry['name']} {member}" if several else entry["name"]

    def add_change_across_columns(self, entry: dict, function: str, targets: list, result_type: str, result_unit: str) -> None:
        pairs = list(zip(targets, targets[1:]))
        for previous, current in pairs:
            column = self.next_free_column()
            self.put(self.placed.header_row, column, self.header_for(entry, self.placed.cells[(self.placed.header_row, current)], len(pairs) > 1))
            for row in self.rows_with_total():
                row_part = self.row_criteria(row) if row != self.placed.total_row else []
                measure = self.measures[0]
                comparable = complete_and_equal(self.table, measure, row_part + self.member_criteria(current), row_part + self.member_criteria(previous))
                self.put(row, column, change_formula(function, self.cell(row, current), self.cell(row, previous), comparable), number_format(result_type, result_unit))

    def add_share_of_row_total(self, entry: dict, targets: list) -> None:
        for target in targets:
            column = self.next_free_column()
            self.put(self.placed.header_row, column, self.header_for(entry, self.placed.cells[(self.placed.header_row, target)], len(targets) > 1))
            for row in self.rows_with_total():
                partial = (row, target) in self.placed.partial or (row, self.placed.total_column) in self.placed.partial
                self.put_marked(row, column, share_formula(self.cell(row, target), self.cell(row, self.placed.total_column)), "0.0%", partial)

    def add_share_of_column_total(self, entry: dict, targets: list) -> None:
        for target in targets:
            column = self.next_free_column()
            member = self.placed.cells[(self.placed.header_row, target)]
            self.put(self.placed.header_row, column, self.header_for(entry, member, len(targets) > 1))
            for row in self.rows_with_total():
                partial = (row, target) in self.placed.partial or (self.placed.total_row, target) in self.placed.partial
                self.put_marked(row, column, share_formula(self.cell(row, target), self.cell(self.placed.total_row, target)), "0.0%", partial)

    def add_change_down_rows(self, entry: dict, function: str, dimension: str, targets: list, result_type: str, result_unit: str) -> None:
        dimension_index = [item.name for item in self.row_dimensions].index(dimension)
        members = self.table.members(dimension)
        for target in targets:
            column = self.next_free_column()
            self.put(self.placed.header_row, column, entry["name"])
            for row, combination in zip(self.placed.body_rows, self.combinations):
                previous_row = self.previous_row(combination, dimension_index, members)
                if previous_row is None:
                    self.put(row, column, None)
                    continue
                measure = self.measure_at(target)
                column_part = self.member_criteria(target) if self.base_kind(target) == "member" else []
                comparable = complete_and_equal(self.table, measure, self.row_criteria(row) + column_part, self.row_criteria(previous_row) + column_part)
                self.put(row, column, change_formula(function, self.cell(row, target), self.cell(previous_row, target), comparable), number_format(result_type, result_unit))

    def previous_row(self, combination: tuple, dimension_index: int, members: list) -> int | None:
        position = members.index(combination[dimension_index])
        if position == 0:
            return None
        wanted = combination[:dimension_index] + (members[position - 1],) + combination[dimension_index + 1:]
        return next((row for row, other in zip(self.placed.body_rows, self.combinations) if other == wanted), None)

    def add_arithmetic_column(self, entry: dict, expression, location: str) -> None:
        names = referenced_paths(expression)
        columns = {}
        for name in names:
            self.compiler.measure(self.table, name, f"{location}.expression")
            position = next((index for index, (_, measure) in enumerate(self.placed.base_columns) if measure == name), None)
            if position is None:
                raise invalid(f"{location}.expression: {name!r} is not one of this view's measures", f"{location}.expression")
            columns[name] = self.base_column_index(position)
        result_type = entry.get("type") or inferred_type(expression, self.table)
        has_fraction = result_type == "amount" and isinstance(expression, Binary) and expression.operator == "/"
        unit = self.table.column(names[0], names[0]).unit if names else ""
        column = self.next_free_column()
        self.put(self.placed.header_row, column, entry["name"])
        for row in self.rows_with_total():
            references = {name: self.cell(row, index) for name, index in columns.items()}
            criteria = self.row_criteria(row) if row != self.placed.total_row else []
            counts = [values_count(self.table.range(name), criteria) for name in columns]
            partial = any((row, index) in self.placed.partial for index in columns.values())
            self.put_marked(row, column, arithmetic_formula(expression, references, counts), number_format(result_type, unit, has_fraction), partial)

    def rows_with_total(self) -> list[int]:
        return self.placed.body_rows + ([self.placed.total_row] if self.placed.total_row else [])


def compared_dimensions(column_dimension: Column | None, added_entries: list, location: str) -> list[str]:
    compared = [column_dimension.name] if column_dimension is not None else []
    for index, entry in enumerate(added_entries):
        expression = parse_expression(entry["expression"], f"{location}.add[{index}].expression")
        if isinstance(expression, Call) and expression.function in CHANGE_FUNCTIONS and len(expression.arguments) == 2 and isinstance(expression.arguments[1], Path):
            compared.append(expression.arguments[1].name)
    return distinct(compared)


def coverage_labels(table: SourceTable, compared: list[str], template: str) -> list[str]:
    dimensions = [column.name for column in table.columns if column.role == "dimension"]
    labels = []
    for compared_name in compared:
        for other_name in dimensions:
            if other_name == compared_name:
                continue
            covered = {member: len(table.members_where(other_name, compared_name, member)) for member in table.members(compared_name)}
            most = max(covered.values(), default=0)
            labels += [template.format(member=member, dimension=other_name, covered=count, most=most) for member, count in covered.items() if count < most]
    return labels


def share_dimension(entry: dict, location: str) -> str | None:
    expression = parse_expression(entry["expression"], f"{location}.expression")
    if isinstance(expression, Call) and expression.function == "share" and len(expression.arguments) == 2 and isinstance(expression.arguments[1], Path):
        return expression.arguments[1].name
    return None


def criteria_text(criteria: list[tuple[str, str]]) -> str:
    return ",".join(f"{range_text},{value}" for range_text, value in criteria)


def count_formula(criteria: list, table: SourceTable) -> str:
    if not criteria:
        return f"ROWS({table.range(table.columns[0].name)})"
    return f"COUNTIFS({criteria_text(criteria)})"


def values_count(measure_range: str, criteria: list) -> str:
    if not criteria:
        return f"COUNT({measure_range})"
    return f'COUNTIFS({measure_range},"<>",{criteria_text(criteria)})'


def aggregate_formula(measure_range: str, criteria: list) -> str:
    if not criteria:
        return f'=IF(COUNT({measure_range})=0,"",SUM({measure_range}))'
    return f'=IF({values_count(measure_range, criteria)}=0,"",SUMIFS({measure_range},{criteria_text(criteria)}))'


def complete_and_equal(table: SourceTable, measure: Column, current: list, previous: list) -> str:
    measure_range = table.range(measure.name)
    current_rows, previous_rows = count_formula(current, table), count_formula(previous, table)
    return f"{current_rows}={previous_rows},{values_count(measure_range, current)}={current_rows},{values_count(measure_range, previous)}={previous_rows}"


def change_formula(function: str, current: str, previous: str, comparable: str) -> str:
    computed = f'IF({previous}=0,"",{current}/{previous}-1)' if function == "percentChange" else f"{current}-{previous}"
    return f'=IF(AND({comparable},ISNUMBER({current}),ISNUMBER({previous})),{computed},"")'


def share_formula(part: str, whole: str) -> str:
    return f'=IF(AND(ISNUMBER({part}),ISNUMBER({whole})),IF({whole}=0,"",{part}/{whole}),"")'


def excel_text(expression, references: dict) -> str:
    if isinstance(expression, Number):
        return str(expression.value)
    if isinstance(expression, Path):
        return references[expression.name]
    if isinstance(expression, Negate):
        return f"-{excel_text(expression.operand, references)}"
    if isinstance(expression, Binary):
        return f"({excel_text(expression.left, references)}{expression.operator}{excel_text(expression.right, references)})"
    raise invalid(f"{expression.function}() does not belong in arithmetic over measures", expression.function)


def denominators(expression) -> list:
    if isinstance(expression, Binary):
        own = [expression.right] if expression.operator == "/" else []
        return own + denominators(expression.left) + denominators(expression.right)
    if isinstance(expression, Negate):
        return denominators(expression.operand)
    return []


def arithmetic_formula(expression, references: dict, counts: list[str]) -> str:
    cells = list(dict.fromkeys(references.values()))
    body = excel_text(expression, references)
    zero_checks = [f"{excel_text(denominator, references)}=0" for denominator in denominators(expression)]
    if zero_checks:
        body = f'IF(OR({",".join(zero_checks)}),"",{body})'
    if not cells:
        return f"={body}"
    missing = [f"COUNT({','.join(cells)})<{len(cells)}"]
    if len(counts) > 1:
        missing.append(f"NOT(AND({','.join(f'{count}={counts[0]}' for count in counts[1:])}))")
    return f'=IF(OR({",".join(missing)}),"",{body})'


def inferred_type(expression, table: SourceTable) -> str:
    if isinstance(expression, Binary) and expression.operator == "/":
        numerator_types = {table.column(name, name).type for name in referenced_paths(expression.left)}
        denominator_types = {table.column(name, name).type for name in referenced_paths(expression.right)}
        return "percent" if numerator_types == denominator_types else next(iter(numerator_types), "amount")
    types = {table.column(name, name).type for name in referenced_paths(expression)}
    return next(iter(types), "amount")


TABLE_EXTENSIONS = (".csv", ".tsv")


def require_attached_tables_read(declaration: dict, attachments: tuple) -> None:
    attached = [attachment for attachment in attachments if FilePath(str(attachment.get("path", ""))).suffix.lower() in TABLE_EXTENSIONS]
    read_paths = [entry["csvPath"] for entry in declaration["tables"] if entry.get("csvPath")]
    unread = [attachment for attachment in attached if not any(is_same_file(attachment["path"], read_path) for read_path in read_paths)]
    typed = next((index for index, entry in enumerate(declaration["tables"]) if not entry.get("csvPath")), None)
    if not unread or typed is None:
        return
    location = f"declaration.tables[{typed}].rows"
    raise invalid(f"{location}: the request attached {', '.join(attachment['path'] for attachment in unread)}; an attached table is read as it is, never typed", location, f'set "csvPath": "{unread[0]["path"]}" on that table in place of its rows')


def is_same_file(first: str, second: str) -> bool:
    first_path, second_path = FilePath(first).expanduser(), FilePath(second).expanduser()
    if first_path.exists() and second_path.exists():
        return os.path.samefile(first_path, second_path)
    return first_path.resolve() == second_path.resolve()


def compile_declaration(declaration: object, attachments: tuple = ()) -> Compiler:
    require_valid(WORKBOOK_DECLARATION, declaration, "declaration")
    require_attached_tables_read(declaration, attachments)
    compiler = Compiler(declaration)
    for index, entry in enumerate(declaration.get("views") or []):
        compiler.views.append(compiler.compile_view(entry, f"declaration.views[{index}]"))
    return compiler
