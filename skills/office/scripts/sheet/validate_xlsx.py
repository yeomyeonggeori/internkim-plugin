#!/usr/bin/env python3
from __future__ import annotations

from openpyxl.utils.cell import range_boundaries

from office_inputs import office_file
from office_result import Issue, OfficeArgumentParser, Result, run_command
from sheet_definitions import AUTO_FILTER_MISSING, BLANK_HEADER_CELLS, HEADER_NOT_FROZEN
from sheet_styling import MINIMUM_DATA_ROWS, MINIMUM_TABLE_COLUMNS, header_row_index, is_filled, non_blank_count
from workbook_access import open_workbook


FORMULA_CELL_LIMIT = 50


def main() -> Result:
    arguments = parse_arguments()
    workbook = open_workbook(arguments.workbook_path)
    sheet_summaries = [summarize_sheet(worksheet) for worksheet in workbook.worksheets]
    issues = [issue for summary in sheet_summaries for issue in sheet_issues(summary)]
    formula_cells = [
        f"{worksheet.title}!{cell.coordinate}"
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if is_formula(cell.value)
    ]
    details = {
        "sheetCount": len(workbook.worksheets),
        "sheets": sheet_summaries,
        "formulaCells": formula_cells[:FORMULA_CELL_LIMIT],
        "formulaCellCount": len(formula_cells),
    }
    return Result(summary=f"checked {arguments.workbook_path}: {len(issues)} issues", output_path=arguments.workbook_path, issues=tuple(issues), details=details)


def is_inside(cell, bounds: tuple[int, int, int, int]) -> bool:
    min_column, min_row, max_column, max_row = bounds
    return min_column <= cell.column <= max_column and min_row <= cell.row <= max_row


def holds_only_pivots(worksheet) -> bool:
    pivots = [range_boundaries(pivot.location.ref) for pivot in getattr(worksheet, "_pivots", [])]
    filled = [cell for cell in worksheet._cells.values() if is_filled(cell.value)]
    return bool(pivots) and all(any(is_inside(cell, bounds) for bounds in pivots) for cell in filled)


def summarize_sheet(worksheet) -> dict:
    header_row = header_row_index(worksheet)
    is_table_sheet = worksheet.max_row >= header_row and not holds_only_pivots(worksheet)
    header_values = [cell.value for cell in worksheet[header_row]] if is_table_sheet else []
    data_rows = sum(1 for row in worksheet.iter_rows(min_row=header_row + 1) if non_blank_count(cell.value for cell in row) > 0)
    return {
        "title": worksheet.title,
        "rows": worksheet.max_row,
        "columns": worksheet.max_column,
        "freezePanes": str(worksheet.freeze_panes) if worksheet.freeze_panes else None,
        "autoFilter": bool(worksheet.auto_filter.ref),
        "dataRows": data_rows,
        "isDataTable": non_blank_count(header_values) >= MINIMUM_TABLE_COLUMNS and data_rows >= MINIMUM_DATA_ROWS,
        "blankHeaderCount": sum(1 for value in header_values if value is None or str(value).strip() == ""),
    }


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def sheet_issues(summary: dict) -> list[Issue]:
    title = summary["title"]
    issues = []
    if summary["isDataTable"] and not summary["freezePanes"]:
        issues.append(HEADER_NOT_FROZEN.issue(f"{title}: header row is not frozen", title))
    if summary["isDataTable"] and not summary["autoFilter"]:
        issues.append(AUTO_FILTER_MISSING.issue(f"{title}: auto filter is missing", title))
    if summary["blankHeaderCount"] > 0:
        issues.append(BLANK_HEADER_CELLS.issue(f"{title}: {summary['blankHeaderCount']} blank header cells", title))
    return issues


def parse_arguments():
    parser = OfficeArgumentParser()
    parser.add_argument("workbook_path", type=office_file("xlsx"))
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
