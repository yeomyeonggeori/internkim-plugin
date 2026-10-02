#!/usr/bin/env python3
from __future__ import annotations

import csv
import io
import os
from pathlib import Path
from types import SimpleNamespace

from openpyxl.utils import get_column_letter

from sheet.workbook.cell_values import typed_cell_value
from sheet.formulas.functions import written_value
from core.office_inputs import read_text_input
from core.office_operations import apply_batch, save_atomically
from core.office_arguments import route_arguments
from core.office_result import INVALID_VALUE, OfficeFailure, Result, read_json_file, run_command
from core.office_schema import closest_name, require_valid
from sheet.sheet_definitions import WORKBOOK_SPECIFICATION
from sheet.operations.operation_set import SHEET_OPERATIONS, SheetEditing, save_editing
from sheet.operations.styling import style_table
from sheet.operations.sheets import validate_sheet_name
from sheet.operations.written_cells import require_writable_rows
from core.excel_limits import MAXIMUM_COLUMN, fitting_sheet_name
from sheet.workbook.access import column_index


def optional_text(value):
    return (value or "").strip()


def read_specification(specification_path):
    specification = read_json_file(specification_path)
    require_valid(WORKBOOK_SPECIFICATION, specification, "spec")
    require_sheet_titles(specification, "spec")
    return specification


def require_sheet_titles(specification, location):
    taken = SimpleNamespace(sheetnames=[])
    for index, sheet_specification in enumerate(specification["sheets"]):
        title = sheet_specification["title"].strip()
        validate_sheet_name(taken, title, f"{location}.sheets[{index}].title")
        taken.sheetnames.append(title)


def create_workbook(specification):
    from openpyxl import Workbook

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    workbook_title = optional_text(specification.get("title"))
    if workbook_title:
        workbook.properties.title = workbook_title

    for index, sheet_specification in enumerate(specification["sheets"]):
        worksheet = add_sheet(workbook, sheet_specification, get_column_letter, f"spec.sheets[{index}]")
        apply_default_formatting(worksheet, sheet_specification, f"spec.sheets[{index}]")

    return workbook


def add_sheet(workbook, sheet_specification, get_column_letter, location):
    title = sheet_specification["title"].strip()
    worksheet = workbook.create_sheet(title=fitting_sheet_name(title))
    rows = read_rows(sheet_specification)
    require_writable_rows(rows, sheet_specification["csvPath"].strip() if sheet_specification.get("csvPath") else f"{location}.rows")
    heading = optional_text(sheet_specification.get("heading"))
    if heading:
        worksheet.append([heading])
    for row in rows:
        worksheet.append([written_value(value) for value in row])
    default_freeze_panes = "A3" if heading else "A2"
    freeze_panes = sheet_specification.get("freezePanes", default_freeze_panes)
    header_row = header_row_index(sheet_specification)
    if isinstance(freeze_panes, str) and freeze_panes.strip() and worksheet.max_row >= header_row:
        worksheet.freeze_panes = freeze_panes.strip()
    if sheet_specification.get("autoFilter", True) and worksheet.max_row > 0 and worksheet.max_column > 0:
        worksheet.auto_filter.ref = auto_filter_reference(worksheet, header_row, get_column_letter)
    return worksheet


def header_row_index(sheet_specification):
    return 2 if optional_text(sheet_specification.get("heading")) else 1


def auto_filter_reference(worksheet, header_row, get_column_letter):
    if worksheet.max_row < header_row:
        return worksheet.dimensions
    return f"A{header_row}:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"


def read_rows(sheet_specification):
    if sheet_specification.get("csvPath"):
        return read_delimited_rows(sheet_specification)
    return sheet_specification.get("rows") or []


def read_delimited_rows(sheet_specification):
    csv_path = sheet_specification["csvPath"].strip()
    delimiter = sheet_specification.get("delimiter") or ","
    if delimiter == "\\t":
        delimiter = "\t"
    text = read_text_input(csv_path)
    return [[typed_cell_value(value) for value in row] for row in csv.reader(io.StringIO(text, newline=""), delimiter=line_delimiter(text, delimiter))]


def line_delimiter(text, delimiter):
    first_line = text.split("\n", 1)[0]
    return "\t" if delimiter == "," and "," not in first_line and "\t" in first_line else delimiter


def apply_default_formatting(worksheet, sheet_specification, location):
    if worksheet.max_row == 0:
        return
    style_table(worksheet, bool(optional_text(sheet_specification.get("heading"))))
    headers = header_letters(worksheet, header_row_index(sheet_specification))
    for column_letter, width in by_column_letter(sheet_specification, "columnWidths", headers, location).items():
        worksheet.column_dimensions[column_letter].width = max(float(width), 4.0)
    for column_letter, number_format in by_column_letter(sheet_specification, "numberFormats", headers, location).items():
        for cell in worksheet[column_letter]:
            cell.number_format = number_format


def header_letters(worksheet, header_row):
    return {str(cell.value): cell.column_letter for cell in worksheet[header_row] if cell.value not in (None, "")}


def by_column_letter(sheet_specification, field, headers, location):
    return {keyed_column_letter(key, headers, f"{location}.{field}.{key}"): value for key, value in (sheet_specification.get(field) or {}).items()}


def keyed_column_letter(key, headers, location):
    if key.strip().isascii() and key.strip().isalpha():
        return get_column_letter(column_index(key, location))
    meant = meant_column(key.strip(), headers)
    suggestion = f'use "{meant[0]}", {meant[1]}' if meant else "use a column letter such as \"C\""
    raise OfficeFailure(INVALID_VALUE.issue(f"{location}: {key!r} is not a column letter", location, suggestion))


def meant_column(key, headers):
    header = key if key in headers else closest_name(key, list(headers))
    if header is not None:
        return headers[header], f"the column headed {header!r}"
    if key.isdigit() and 1 <= int(key) <= MAXIMUM_COLUMN:
        return get_column_letter(int(key)), f"column number {key}"
    return None



def main():
    arguments = route_arguments("create", "xlsx")
    specification = read_specification(arguments.source)
    workbook = create_workbook(specification)
    output_path = Path(os.path.expanduser(arguments.output))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    editing = SheetEditing(workbook, None)
    changes = apply_batch(SHEET_OPERATIONS, editing, specification.get("operations") or [])
    issues = save_atomically(lambda temporary_path: save_editing(editing, temporary_path), str(output_path))
    details = {"changes": changes} if changes else None
    return Result(summary=f"created {output_path}", output_path=str(output_path), issues=tuple(issues), details=details)


if __name__ == "__main__":
    raise SystemExit(run_command(main))
