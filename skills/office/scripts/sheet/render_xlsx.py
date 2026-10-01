#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path
import shutil
import sys
import tempfile

SCRIPTS_PATH = Path(__file__).resolve().parents[1]
sys.path[1:1] = [str(SCRIPTS_PATH / "deck")]

from openpyxl import load_workbook  # noqa: E402

from formula_cache import cache_formula_values  # noqa: E402
from office_preview import PAGES_NOT_RENDERED, Preview, approximation_issues, write_preview  # noqa: E402
from office_result import INVALID_VALUE, OfficeArgumentParser, OfficeFailure, Result, run_command  # noqa: E402
from preview_fonts import FontRegistry  # noqa: E402
from xlsx_colors import theme_palette  # noqa: E402
from xlsx_preview import SheetPreviewer  # noqa: E402


def workbook_values(source_path: Path):
    with tempfile.TemporaryDirectory(prefix="office-preview-") as directory:
        computed = Path(directory) / "computed.xlsx"
        shutil.copyfile(source_path, computed)
        cache_formula_values(str(computed))
        return load_workbook(str(computed), data_only=True)


def xlsx_preview(source_path: Path, sheet_name: str | None) -> tuple[Preview, list[dict]]:
    workbook = load_workbook(str(source_path))
    values = workbook_values(source_path)
    names = [sheet_name] if sheet_name else [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    if sheet_name and sheet_name not in workbook.sheetnames:
        raise OfficeFailure(INVALID_VALUE.issue(f"--sheet {sheet_name!r} is not in the workbook", "--sheet", suggestion=f"use one of: {', '.join(workbook.sheetnames)}"))
    fonts = FontRegistry()
    preview = Preview(title=source_path.stem)
    previewer = SheetPreviewer(workbook, values, source_path.name, theme_palette(workbook.loaded_theme), fonts, preview)
    planned = [page for name in names for page in previewer.sheet_pages(workbook[name], values[name])]
    preview.pages = [previewer.page_html(number, len(planned), *page) for number, page in enumerate(planned, start=1)]
    return preview, fonts.preview_fonts()


def main() -> Result:
    arguments = parse_arguments()
    source_path = Path(arguments.source_path).expanduser()
    if not source_path.is_file():
        raise FileNotFoundError(2, "no such file", str(source_path))
    output_directory = Path(arguments.output_directory).expanduser() if arguments.output_directory else source_path.with_name(f"{source_path.stem}-preview")
    preview, preview_fonts = xlsx_preview(source_path, arguments.sheet)
    preview_path = write_preview(preview, output_directory)
    issues = [*approximation_issues(preview, source_path.name), PAGES_NOT_RENDERED.issue(f"wrote {preview_path.name} with {len(preview.pages)} pages; page images come from the shared renderer", str(preview_path))]
    details = {"preview": str(preview_path), "pageCount": len(preview.pages), "previewFonts": preview_fonts, "approximations": preview.approximations}
    return Result(summary=f"laid out {source_path.name} as {len(preview.pages)} printed pages in {preview_path}", output_path=str(preview_path), issues=tuple(issues), details=details)


def parse_arguments():
    parser = OfficeArgumentParser(description="Lay out each visible sheet as printed pages of preview HTML: print area, page setup and scaling, column widths, row heights, merges, number formats as displayed, fonts, fills, borders, conditional colors, charts and pictures.")
    parser.add_argument("source_path", help="the .xlsx to lay out")
    parser.add_argument("--sheet", help="lay out only this sheet")
    parser.add_argument("--output-directory", help="where preview.html goes; default <name>-preview beside the file")
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
