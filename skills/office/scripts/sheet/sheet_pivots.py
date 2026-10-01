from __future__ import annotations

from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from office_operations import OPERATION_NOT_APPLICABLE, Change
from office_result import INVALID_VALUE, OfficeFailure
from cell_values import typed_date
from display_width import display_width
from office_schema import closest_suggestion, did_you_mean
from pivot_formula import parse_formula
from pivot_layout import PivotAxis, PivotGrid, PivotModel, PivotValue, TopFilter, as_datetime, binned_field, build_model, grouped_field, numbers, pivot_grid, plain_field
from pivot_parts import PivotPlacement, add_pivot_parts
from workbook_access import parse_cell, parse_range, sheet_of
from workbook_snapshot import cell_values


FUNCTION_CAPTIONS = {"sum": "Sum", "count": "Count", "average": "Average", "max": "Max", "min": "Min"}
DEFAULT_NUMBER_FORMAT = "#,##0"
DEFAULT_PERCENT_FORMAT = "0.0%"
DEFAULT_TOTAL_LABEL = "Grand Total"
DEFAULT_TARGET_SHEET = "Pivot"
DEFAULT_TARGET_CELL = "A3"
DATE_FORMAT = "yyyy-mm-dd"
ALL_ITEMS_LABEL = "(All)"
HEADER_FILL = "DCEAF7"
RULE_COLOR = "94A3B8"
MINIMUM_COLUMN_WIDTH = 14
LABEL_MARGIN = 2


def refusal(message: str, location: str) -> OfficeFailure:
    return OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: {message}", location))


def field_index(headers: list, name: str, location: str) -> int:
    if name in headers:
        return headers.index(name)
    raise OfficeFailure(INVALID_VALUE.issue(f"{location}: {name!r} is not a header of the source{did_you_mean(name, headers)}; it has {', '.join(headers)}", location, closest_suggestion(name, headers)))


def named_list(operation: dict, key: str) -> list[str]:
    written = operation.get(key)
    if written is None:
        return []
    return [written] if isinstance(written, str) else list(written)


def field_indexes(headers: list, operation: dict, key: str, location: str) -> list[int]:
    suffix = "" if isinstance(operation.get(key), str) else "[{}]"
    return [field_index(headers, name, f"{location}.{key}{suffix.format(position)}") for position, name in enumerate(named_list(operation, key))]


def require_distinct_roles(headers: list, roles: dict[str, list[int]], location: str) -> None:
    seen: dict[int, str] = {}
    for role, indexes in roles.items():
        for index in indexes:
            if index in seen:
                raise refusal(f"{headers[index]!r} is named in {seen[index]} and in {role}; a header takes one place in a pivot", f"{location}.{role}")
            seen[index] = role


def source_headers(worksheet, bounds: tuple, location: str) -> list[str]:
    if bounds[2] <= bounds[0]:
        raise refusal("a pivot needs a header row and at least one data row", f"{location}.range")
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=bounds[0], max_row=bounds[0], min_col=bounds[1], max_col=bounds[3]))]
    if any(header is None or str(header).strip() == "" for header in headers) or len({str(header) for header in headers}) != len(headers):
        raise refusal(f"every header cell of the source needs its own name; row {bounds[0]} holds {headers}", f"{location}.range")
    return [str(header) for header in headers]


def value_entries(operation: dict) -> list[dict]:
    return [{"field": entry} if isinstance(entry, str) else entry for entry in operation["values"]]


def plan_value(headers: list, operation: dict, entry: dict, location: str) -> PivotValue:
    function = entry.get("function", operation.get("function", "sum"))
    show_as = entry.get("showAs", "value")
    number_format = entry.get("numberFormat") or (DEFAULT_PERCENT_FORMAT if show_as != "value" else operation.get("numberFormat", DEFAULT_NUMBER_FORMAT))
    if "formula" not in entry:
        index = field_index(headers, entry["field"], f"{location}.field")
        caption = entry.get("label") or f"{FUNCTION_CAPTIONS[function]} of {headers[index]}"
        return PivotValue(caption, index, function, show_as, number_format)
    if entry["field"] in headers:
        raise refusal(f"a calculated value needs a new name, and {entry['field']!r} is already a header of the source", f"{location}.field")
    if function != "sum":
        raise refusal(f"a calculated value adds up its fields first and then applies the formula, so its function is sum, not {function}", f"{location}.function")
    formula = parse_formula(entry["formula"], headers, f"{location}.formula")
    caption = entry.get("label") or f"Sum of {entry['field']}"
    return PivotValue(caption, None, "sum", show_as, number_format, formula, entry["field"])


def plan_values(headers: list, operation: dict, location: str) -> list[PivotValue]:
    values = [plan_value(headers, operation, entry, f"{location}.values[{position}]") for position, entry in enumerate(value_entries(operation))]
    names = [value.calculated_name.casefold() for value in values if value.calculated_name]
    if len(names) != len(set(names)):
        raise refusal("two calculated values share a name", f"{location}.values")
    return values


def date_groups(headers: list, operation: dict, axis_indexes: list[int], location: str) -> dict[int, str]:
    groups = {}
    for name, unit in (operation.get("groupDates") or {}).items():
        index = field_index(headers, name, f"{location}.groupDates.{name}")
        if index not in axis_indexes:
            raise refusal(f"{name!r} is grouped by {unit}, so it must also be named in row or column", f"{location}.groupDates.{name}")
        groups[index] = unit
    return groups


def number_groups(headers: list, operation: dict, axis_indexes: list[int], date_indexes: dict, location: str) -> dict[int, dict]:
    groups = {}
    for name, group in (operation.get("groupNumbers") or {}).items():
        group_location = f"{location}.groupNumbers.{name}"
        index = field_index(headers, name, group_location)
        if index not in axis_indexes:
            raise refusal(f"{name!r} is grouped into bins, so it must also be named in row or column", group_location)
        if index in date_indexes:
            raise refusal(f"{name!r} is grouped by date in groupDates, so it cannot also be grouped into bins", group_location)
        if group["step"] <= 0:
            raise refusal(f"step must be more than 0, and it is {group['step']}", f"{group_location}.step")
        groups[index] = group
    return groups


def require_numbers(records: list, index: int, header: str, first_row: int, location: str) -> None:
    for offset, record in enumerate(records):
        if not numbers([record[index]]):
            shown = "nothing" if record[index] in (None, "") else repr(record[index])
            raise refusal(f"grouping {header!r} into bins needs a number in every row, and row {first_row + offset} holds {shown}", f"{location}.groupNumbers.{header}")


def top_filter(headers: list, operation: dict, axis_indexes: list[int], location: str) -> TopFilter | None:
    top = operation.get("top")
    if top is None:
        return None
    index = field_index(headers, top["field"], f"{location}.top.field")
    if index not in axis_indexes:
        raise refusal(f"{top['field']!r} is filtered by its top items, so it must also be named in row or column", f"{location}.top.field")
    return TopFilter(index, top["count"], top.get("bottom", False))


def stored_date(value: object) -> object:
    return typed_date(value) if isinstance(value, str) else value


def store_text_dates(worksheet, records: list, index: int, header: str, first_row: int, column: int, location: str) -> int:
    dates = [stored_date(record[index]) for record in records]
    for offset, value in enumerate(dates):
        if as_datetime(value) is None:
            shown = "nothing" if value in (None, "") else repr(value)
            raise refusal(f"grouping {header!r} by date needs a date in every row, and row {first_row + offset} holds {shown}; write dates as YYYY-MM-DD such as 2024-01-31", f"{location}.groupDates.{header}")
    converted = 0
    for offset, (record, value) in enumerate(zip(records, dates)):
        if value is record[index]:
            continue
        cell = worksheet.cell(row=first_row + offset, column=column, value=value)
        cell.number_format = DATE_FORMAT
        record[index] = value
        converted += 1
    return converted


def date_notes(converted: dict[str, int]) -> str:
    stored = [f"{count} YYYY-MM-DD texts of {header!r}" for header, count in converted.items() if count]
    return f"; stored {' and '.join(stored)} as dates so they group" if stored else ""


def pivot_axis(records: list, indexes: list[int], groups: dict[int, str], bins: dict[int, dict]) -> PivotAxis:
    return PivotAxis([axis_field(records, index, groups, bins) for index in indexes])


def axis_field(records: list, index: int, groups: dict[int, str], bins: dict[int, dict]):
    if index in groups:
        return grouped_field(index, records, groups[index])
    if index in bins:
        return binned_field(index, records, bins[index]["step"], bins[index].get("start"))
    return plain_field(index, records)


def pivot_names(workbook) -> set:
    return {pivot.name.casefold() for worksheet in workbook.worksheets for pivot in getattr(worksheet, "_pivots", [])}


def pivot_name(editing, operation: dict, location: str) -> str:
    taken = pivot_names(editing.workbook) | editing.reserved_names
    name = operation.get("name") or next(f"PivotTable{number}" for number in range(1, 10000) if f"pivottable{number}" not in taken)
    if name.casefold() in taken:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.name: the workbook already has a pivot table named {name!r}", f"{location}.name"))
    return name


def plan_add_pivot_table(editing, operation: dict, location: str) -> Change:
    workbook = editing.workbook
    worksheet = sheet_of(workbook, operation, location)
    bounds = parse_range(operation["range"], f"{location}.range")
    headers = source_headers(worksheet, bounds, location)
    row_indexes = field_indexes(headers, operation, "row", location)
    column_indexes = field_indexes(headers, operation, "column", location)
    page_indexes = field_indexes(headers, operation, "filters", location)
    require_distinct_roles(headers, {"row": row_indexes, "column": column_indexes, "filters": page_indexes}, location)
    values = plan_values(headers, operation, location)
    if column_indexes and len(values) != 1:
        raise refusal("a pivot with a column field summarizes exactly one value", f"{location}.values")
    groups = date_groups(headers, operation, row_indexes + column_indexes, location)
    bins = number_groups(headers, operation, row_indexes + column_indexes, groups, location)
    top = top_filter(headers, operation, row_indexes + column_indexes, location)
    target_row, target_column = parse_cell(operation.get("targetCell", DEFAULT_TARGET_CELL), f"{location}.targetCell")
    name = pivot_name(editing, operation, location)

    def change() -> str:
        records = cell_values(workbook, worksheet, (bounds[0] + 1, bounds[1], bounds[2], bounds[3]))
        converted = {headers[index]: store_text_dates(worksheet, records, index, headers[index], bounds[0] + 1, bounds[1] + index, location) for index in groups}
        for index in bins:
            require_numbers(records, index, headers[index], bounds[0] + 1, location)
        pages = [plain_field(index, records) for index in page_indexes]
        rows, columns = pivot_axis(records, row_indexes, groups, bins), pivot_axis(records, column_indexes, groups, bins)
        model = build_model(PivotModel(headers, records, rows, columns, pages, values, operation.get("totalLabel", DEFAULT_TOTAL_LABEL), top))
        target = target_sheet(workbook, operation.get("targetSheet", DEFAULT_TARGET_SHEET))
        grid = pivot_grid(model)
        table_row = write_page_fields(target, model, target_row, target_column, location)
        reference = write_grid(target, grid, table_row, target_column, location)
        placement = PivotPlacement(name, worksheet.title, operation["range"].replace("$", "").upper(), reference, grid.header_rows)
        editing.reserved_names.add(name.casefold())
        editing.package_patches.append(lambda package: add_pivot_parts(package, target.title, model, placement))
        return f"added pivot table {name} of {worksheet.title}!{placement.source_reference} on {target.title}!{reference}{date_notes(converted)}"
    return change


def target_sheet(workbook, name: str):
    match = next((worksheet for worksheet in workbook.worksheets if worksheet.title == name), None)
    return match if match is not None else workbook.create_sheet(title=name)


def require_empty(worksheet, top: int, left: int, height: int, width: int, location: str) -> None:
    occupied = [cell.coordinate for row in worksheet.iter_rows(min_row=top, max_row=top + height - 1, min_col=left, max_col=left + width - 1) for cell in row if cell.value is not None]
    if occupied:
        area = f"{get_column_letter(left)}{top}:{get_column_letter(left + width - 1)}{top + height - 1}"
        raise refusal(f"the pivot needs {area} empty, but {worksheet.title} has values in {', '.join(occupied[:5])}", f"{location}.targetCell")


def write_page_fields(worksheet, model: PivotModel, top: int, left: int, location: str) -> int:
    if not model.pages:
        return top
    require_empty(worksheet, top, left, len(model.pages) + 1, 2, location)
    for offset, pivot_field in enumerate(model.pages):
        worksheet.cell(row=top + offset, column=left, value=model.headers[pivot_field.index]).font = Font(bold=True)
        worksheet.cell(row=top + offset, column=left + 1, value=ALL_ITEMS_LABEL)
    return top + len(model.pages) + 1


def write_grid(worksheet, grid: PivotGrid, top: int, left: int, location: str) -> str:
    width = max(len(row) for row in grid.rows)
    require_empty(worksheet, top, left, len(grid.rows), width, location)
    for offset, values in enumerate(grid.rows):
        for column_offset, value in enumerate(values):
            cell = worksheet.cell(row=top + offset, column=left + column_offset, value=value)
            style_pivot_cell(cell, grid.kinds[offset], grid.number_formats.get(column_offset))
    for column_offset in range(width):
        letter = get_column_letter(left + column_offset)
        widest_label = max((display_width(row[column_offset]) for row in grid.rows if column_offset < len(row) and isinstance(row[column_offset], str)), default=0)
        worksheet.column_dimensions[letter].width = max(worksheet.column_dimensions[letter].width or 0, MINIMUM_COLUMN_WIDTH, widest_label + LABEL_MARGIN)
    return f"{get_column_letter(left)}{top}:{get_column_letter(left + width - 1)}{top + len(grid.rows) - 1}"


def style_pivot_cell(cell, kind: str, number_format: str | None) -> None:
    if kind == "header":
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        return
    if kind in ("subtotal", "total"):
        cell.font = Font(bold=True)
    if kind == "total":
        cell.border = Border(top=Side(style="thin", color=RULE_COLOR))
    if number_format:
        cell.number_format = number_format
