from __future__ import annotations

import csv
import io

from openpyxl.utils import get_column_letter

from sheet.workbook.cell_values import typed_cell_value
from sheet.formulas.functions import written_value
from core.office_inputs import read_text_input
from sheet.operations.styling import style_table
from sheet.operations.written_cells import require_writable_rows
from core.excel_limits import fitting_sheet_name


HEADER_ROW = 1


def create_workbook(specification):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook_title = (specification.get("title") or "").strip()
    if workbook_title:
        workbook.properties.title = workbook_title
    for index, sheet_specification in enumerate(specification["sheets"]):
        worksheet = add_sheet(workbook, sheet_specification, f"spec.sheets[{index}]")
        if worksheet.max_row:
            style_table(worksheet, False)
    return workbook


def add_sheet(workbook, sheet_specification, location):
    worksheet = workbook.create_sheet(title=fitting_sheet_name(sheet_specification["title"].strip()))
    rows = read_rows(sheet_specification)
    require_writable_rows(rows, sheet_specification["csvPath"].strip() if sheet_specification.get("csvPath") else f"{location}.rows")
    for row in rows:
        worksheet.append([written_value(value) for value in row])
    freeze_panes = sheet_specification.get("freezePanes", f"A{HEADER_ROW + 1}")
    if isinstance(freeze_panes, str) and freeze_panes.strip() and worksheet.max_row >= HEADER_ROW:
        worksheet.freeze_panes = freeze_panes.strip()
    if sheet_specification.get("autoFilter", True) and worksheet.max_row > 0 and worksheet.max_column > 0:
        worksheet.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    return worksheet


def read_rows(sheet_specification):
    if sheet_specification.get("csvPath"):
        return read_delimited_rows(sheet_specification)
    return sheet_specification.get("rows") or []


def read_delimited_rows(sheet_specification):
    csv_path = sheet_specification["csvPath"].strip()
    delimiter = sheet_specification.get("delimiter") or ","
    if delimiter == "\\t":
        delimiter = "\t"
    text = read_text_input(csv_path)
    return [[typed_cell_value(value) for value in row] for row in csv.reader(io.StringIO(text, newline=""), delimiter=line_delimiter(text, delimiter))]


def line_delimiter(text, delimiter):
    first_line = text.split("\n", 1)[0]
    return "\t" if delimiter == "," and "," not in first_line and "\t" in first_line else delimiter
