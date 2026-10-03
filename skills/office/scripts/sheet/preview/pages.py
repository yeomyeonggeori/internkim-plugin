from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field, replace
import datetime
import re

from openpyxl.styles.fonts import DEFAULT_FONT
from openpyxl.utils import get_column_letter, range_boundaries

from core.number_format import Displayed, displayed
from render.office_preview import PageGeometry, Preview, escaped, page_section, pixels, positioned, style_attribute
from core.units import emu_to_pixels, inches_to_pixels, millimetres_to_pixels, points_to_pixels
from fonts.registry import OFFICE_KOREAN_FAMILY
from fonts.preview import FontRegistry, FontRequest, css_font_family, draws_scripts_apart, script_font_family, script_runs
from sheet.operations.formatting import STACKED_ROTATION, rotation_degrees
from core.office_result import Issue
from sheet.sheet_definitions import READABLE_PRINT_POINTS, SHEET_PRINTS_SMALL, SHEET_PRINTS_WIDE
from sheet.operations.objects import EXCEL_DEFAULT_FIT_PAGES, is_fitted_to_pages
from sheet.operations.charts import DEFAULT_ANCHOR_GAP, anchor_cell
from sheet.preview.colors import css_color, theme_palette
from sheet.preview.conditional import ConditionalStyles
from sheet.preview.charts import VISIBLE_PIXELS, chart_html, chart_kind, chart_title, drawing_box, image_html, is_whole
from core.page_sizes import DEFAULT_PAPER, PAPER_BY_SPREADSHEET_CODE


DEFAULT_COLUMN_CHARACTERS = 8.43
DEFAULT_ROW_POINTS = 15
FIT_SEARCH_STEPS = 30
MAXIMUM_DIGIT_PIXELS = 7
CELL_PADDING_PIXELS = 2
INDENT_PIXELS = 9
GRIDLINE = "1px solid #d0d0d0"
BORDER_STYLES = {
    "thin": ("1px", "solid"), "hair": ("1px", "solid"), "medium": ("2px", "solid"), "thick": ("3px", "solid"), "double": ("3px", "double"),
    "dashed": ("1px", "dashed"), "dotted": ("1px", "dotted"), "mediumDashed": ("2px", "dashed"), "dashDot": ("1px", "dashed"),
    "mediumDashDot": ("2px", "dashed"), "dashDotDot": ("1px", "dotted"), "mediumDashDotDot": ("2px", "dotted"), "slantDashDot": ("2px", "dashed"),
}
HORIZONTAL = {"left": "flex-start", "center": "center", "right": "flex-end", "centerContinuous": "center", "justify": "flex-start", "fill": "flex-start", "distributed": "center"}
VERTICAL = {"top": "flex-start", "center": "center", "bottom": "flex-end", "justify": "flex-start", "distributed": "center"}
HEADER_CODES = (("&P", "{page}"), ("&N", "{pages}"), ("&A", "{sheet}"), ("&F", "{file}"), ("&D", "{date}"), ("&T", ""), ("&B", ""), ("&I", ""), ("&U", ""), ("&&", "&"))
HEADER_FORMATTING = re.compile(r'&K[0-9A-Fa-f]{6}|&"[^"]*"|&\d+|&[SXYEOHG]')


@dataclass
class SheetFrame:
    worksheet: object
    values: object
    columns: list[int]
    rows: list[int]
    widths: dict[int, float]
    heights: dict[int, float]
    merges: dict[tuple[int, int], tuple[int, int]]
    covered: set
    title_rows: list[int] = field(default_factory=list)


@dataclass(frozen=True)
class PrintLayout:
    frame: SheetFrame
    geometry: PageGeometry
    scale: float
    column_chunks: list[list[int]]


@dataclass(frozen=True)
class PrintExtent:
    column_widths: list[float]
    height: float


@dataclass(frozen=True)
class ChartMove:
    chart: int
    anchor: str


@dataclass(frozen=True)
class PageFit:
    orientation: str
    width_pages: int


@dataclass(frozen=True)
class PrintPlan:
    moves: list[ChartMove]
    fit: PageFit | None
    scale: float
    body_points: float


class SheetPreviewer:
    def __init__(self, workbook, values, file_name: str, palette: tuple, fonts: FontRegistry, preview: Preview):
        self.workbook = workbook
        self.values = values
        self.file_name = file_name
        self.palette = palette
        self.fonts = fonts
        self.preview = preview
        self.page_painted = False
        self.page_drawings: list[dict] = []
        self.page_contents: list[dict] = []
        self.print_issues: list[Issue] = []

    def sheet_pages(self, worksheet, sheet_values) -> list[tuple[PageGeometry, list[int], list[int], float, SheetFrame]]:
        layout = self.print_layout(worksheet, sheet_values)
        if layout is None:
            return []
        frame, geometry, scale = layout.frame, layout.geometry, layout.scale
        title_height = sum(frame.heights[row] for row in frame.title_rows) * scale
        breaks = {brk.id for brk in worksheet.row_breaks.brk} if worksheet.row_breaks else set()
        body_rows = [row for row in frame.rows if row not in frame.title_rows]
        row_chunks = chunks(body_rows, [frame.heights[row] * scale for row in body_rows], geometry.content_height - title_height, breaks)
        self.print_issues.extend(print_layout_issues(worksheet, layout))
        return [(geometry, columns, frame.title_rows + rows, scale, frame) for columns in layout.column_chunks for rows in row_chunks]

    def print_layout(self, worksheet, sheet_values) -> PrintLayout | None:
        frame = self.frame(worksheet, sheet_values)
        if frame is None:
            return None
        geometry = page_geometry(worksheet)
        scale = print_scale(worksheet, frame_extent(frame), geometry)
        column_chunks = chunks(frame.columns, [frame.widths[column] * scale for column in frame.columns], geometry.content_width, set())
        return PrintLayout(frame, geometry, scale, column_chunks)

    def frame(self, worksheet, sheet_values) -> SheetFrame | None:
        bounds = print_bounds(worksheet, sheet_values)
        if bounds is None:
            return None
        min_column, min_row, max_column, max_row = bounds
        column_settings = column_dimension_map(worksheet)
        columns = [column for column in range(min_column, max_column + 1) if not column_settings.get(column, {}).get("hidden")]
        rows = [row for row in range(min_row, max_row + 1) if not row_hidden(worksheet, row)]
        widths = {column: column_pixels(column_settings.get(column, {}).get("width") or default_column_width(worksheet)) for column in columns}
        merges, covered = merge_maps(worksheet)
        # Excel's row AutoFit ignores merged cells (Microsoft support, "Change the column width or row height in Excel")
        heights = {row: self.row_height(worksheet, sheet_values, row, [column for column in columns if (row, column) not in covered], widths) for row in rows}
        title_rows = print_title_rows(worksheet, rows)
        return SheetFrame(worksheet, sheet_values, columns, rows, widths, heights, merges, covered, title_rows)

    def row_height(self, worksheet, sheet_values, row: int, columns: list[int], widths: dict) -> float:
        dimension = worksheet.row_dimensions.get(row) if hasattr(worksheet.row_dimensions, "get") else None
        if dimension is not None and dimension.height:
            return points_to_pixels(dimension.height)
        default = default_row_pixels(worksheet)
        wrapped = [column for column in columns if worksheet.cell(row=row, column=column).alignment.wrap_text and sheet_values.cell(row=row, column=column).value is not None]
        if not wrapped:
            return default
        return max(default, *(self.wrapped_height(worksheet, sheet_values, row, column, widths[column]) for column in wrapped))

    def wrapped_height(self, worksheet, sheet_values, row: int, column: int, width: float) -> float:
        cell = worksheet.cell(row=row, column=column)
        text = displayed(sheet_values.cell(row=row, column=column).value, cell.number_format).text
        request = font_request(cell.font)
        lines = 0
        for paragraph in text.split("\n"):
            lines += max(1, int(self.fonts.width(request, paragraph) // max(width - 2 * CELL_PADDING_PIXELS, 1)) + 1)
        return lines * self.fonts.line_height(request) + 2

    def page_html(self, page_number: int, page_count: int, geometry: PageGeometry, columns: list[int], rows: list[int], scale: float, frame: SheetFrame) -> str:
        grid = self.grid_html(frame, columns, rows, scale)
        if not self.page_painted:
            self.preview.blank_pages.append(f"page {page_number} ({frame.worksheet.title})")
        self.page_contents.append(page_content(page_number, frame, columns, rows, self.page_drawings))
        grid_width = sum(frame.widths[column] for column in columns) * scale
        left = geometry.margin_left + max(0.0, (geometry.content_width - grid_width) / 2) if frame.worksheet.print_options.horizontalCentered else geometry.margin_left
        parts = [positioned(left, geometry.margin_top, grid_width, grid)]
        substitutions = {"page": page_number, "pages": page_count, "sheet": frame.worksheet.title, "file": self.file_name, "date": datetime.date.today().isoformat()}
        parts.extend(self.header_footer_html(frame.worksheet, geometry, substitutions))
        return page_section(page_number, geometry, "".join(parts))

    def header_footer_html(self, worksheet, geometry: PageGeometry, substitutions: dict) -> list[str]:
        parts = []
        for block, top in ((worksheet.oddHeader, geometry.header_distance), (worksheet.oddFooter, geometry.height - geometry.footer_distance - 16)):
            pieces = [(getattr(block, side).text if getattr(block, side) is not None else None, alignment) for side, alignment in (("left", "left"), ("center", "center"), ("right", "right"))]
            for text, alignment in pieces:
                if text:
                    content = escaped(header_text(text).format(**substitutions))
                    parts.append(positioned(geometry.margin_left, top, geometry.content_width, content, {"text-align": alignment, "font-size": "12px", "font-family": css_font_family(DEFAULT_FONT.name, OFFICE_KOREAN_FAMILY), "white-space": "pre"}))
        return parts

    def grid_html(self, frame: SheetFrame, columns: list[int], rows: list[int], scale: float) -> str:
        column_index = {column: index for index, column in enumerate(columns)}
        row_index = {row: index for index, row in enumerate(rows)}
        conditional = ConditionalStyles(frame.worksheet, frame.values, self.palette)
        backgrounds, texts = [], []
        for row in rows:
            for column in columns:
                if (row, column) in frame.covered and (row, column) not in frame.merges:
                    continue
                background, text = self.cell_html(frame, row, column, columns, rows, column_index, row_index, scale, conditional)
                backgrounds.append(background)
                texts.append(text)
        overlays = self.drawings_html(frame, columns, rows, scale)
        self.page_painted = bool(overlays) or any(backgrounds) or any(texts)
        declarations = {
            "position": "relative",
            "display": "grid",
            "grid-template-columns": " ".join(pixels(frame.widths[column] * scale) for column in columns),
            "grid-template-rows": " ".join(pixels(frame.heights[row] * scale) for row in rows),
            "width": pixels(sum(frame.widths[column] for column in columns) * scale),
        }
        return f"<div{style_attribute(declarations)}>{''.join(backgrounds)}{''.join(texts)}{overlays}</div>"

    def cell_html(self, frame: SheetFrame, row: int, column: int, columns: list[int], rows: list[int], column_index: dict, row_index: dict, scale: float, conditional: ConditionalStyles) -> tuple[str, str]:
        worksheet = frame.worksheet
        cell = worksheet.cell(row=row, column=column)
        value = frame.values.cell(row=row, column=column).value
        last_row, last_column = frame.merges.get((row, column), (row, column))
        row_span = sum(1 for candidate in rows if row <= candidate <= last_row)
        column_span = sum(1 for candidate in columns if column <= candidate <= last_column)
        placement = {"grid-column": f"{column_index[column] + 1} / span {max(column_span, 1)}", "grid-row": f"{row_index[row] + 1} / span {max(row_span, 1)}"}
        extra = conditional.style_for(row, column, value)
        background = {**placement, "box-sizing": "border-box", "background": extra.get("background") or fill_color(cell, self.palette), **self.border_css(frame, row, column, last_row, last_column, scale)}
        bar = extra.get("bar")
        bar_html = f'<div{style_attribute({"width": f"{bar[0]:.1f}%", "height": "70%", "background": bar[1], "margin-top": "2px"})}></div>' if bar else ""
        text = displayed(value, cell.number_format)
        width = sum(frame.widths[candidate] for candidate in columns if column <= candidate <= last_column) * scale
        painted = bar_html or any(name.startswith("border") or (name == "background" and value_set) for name, value_set in background.items())
        background_html = f"<div{style_attribute(background)}>{bar_html}</div>" if painted else ""
        return background_html, self.text_html(cell, text, value, placement, width, scale, extra, frame, row, column, columns)

    def text_html(self, cell, text: Displayed, value, placement: dict, width: float, scale: float, extra: dict, frame: SheetFrame, row: int, column: int, columns: list[int]) -> str:
        if not text.text:
            return ""
        font = cell.font
        request = font_request(font, scale)
        alignment = cell.alignment
        horizontal = alignment.horizontal or ("right" if text.is_number else "center" if isinstance(value, bool) else "left")
        content = stacked(text.text) if alignment.textRotation == STACKED_ROTATION else text.text
        self.fonts.use(request, content)
        wraps = bool(alignment.wrap_text)
        if text.is_number and not wraps and self.fonts.width(request, content) > width - 2 * CELL_PADDING_PIXELS * scale:
            content = "#" * max(1, int((width - 2 * CELL_PADDING_PIXELS * scale) // max(self.fonts.width(request, "#"), 1)))
        span = placement
        if not wraps and not text.is_number and horizontal == "left" and (row, column) not in frame.merges:
            span = {**placement, "grid-column": f"{placement['grid-column'].split(' / ')[0]} / span {spill_span(frame, row, column, columns)}"}
        declarations = {
            **span,
            "display": "flex",
            "justify-content": HORIZONTAL.get(horizontal, "flex-start"),
            "align-items": "center" if rotation_style(alignment.textRotation) else VERTICAL.get(alignment.vertical or "bottom", "flex-end"),
            "padding": f"0 {pixels(CELL_PADDING_PIXELS * scale)}",
            "padding-left": pixels((CELL_PADDING_PIXELS + INDENT_PIXELS * (alignment.indent or 0)) * scale) if alignment.indent else None,
            "box-sizing": "border-box",
            "overflow": "hidden",
            "white-space": "pre-wrap" if wraps else "pre",
            "font-family": css_font_family(request.latin, request.east_asia),
            "font-size": pixels(points_to_pixels(request.size)),
            "font-weight": "700" if font.b else None,
            "font-style": "italic" if font.i else None,
            "color": extra.get("color") or text.color or css_color(font.color, self.palette),
            "line-height": pixels(self.fonts.line_height(request)),
            "text-align": horizontal if horizontal in ("left", "center", "right") else None,
        }
        span = {"text-decoration": text_decoration(font), **rotation_style(alignment.textRotation)}
        return f"<div{style_attribute(declarations)}><span{style_attribute(span)}>{script_spans(request, content)}</span></div>"

    def border_css(self, frame: SheetFrame, row: int, column: int, last_row: int, last_column: int, scale: float) -> dict:
        worksheet = frame.worksheet
        first = worksheet.cell(row=row, column=column).border
        last = worksheet.cell(row=last_row, column=last_column).border
        sides = {"top": first.top, "left": first.left, "bottom": last.bottom, "right": last.right}
        declarations = {}
        gridlines = bool(frame.worksheet.print_options.gridLines)
        for side, border in sides.items():
            value = border_value(border, self.palette)
            if value is None and side in ("left", "top"):
                neighbor = neighbor_border(worksheet, row, column, side)
                if neighbor is not None:
                    continue
            if value is None and gridlines and side in ("right", "bottom"):
                value = GRIDLINE
            if value is not None:
                declarations[f"border-{side}"] = value
        return declarations

    def drawings_html(self, frame: SheetFrame, columns: list[int], rows: list[int], scale: float) -> str:
        body_rows = [row for row in rows if row not in frame.title_rows]
        page_width = sum(frame.widths[column] for column in columns) * scale
        body_height = sum(frame.heights[row] for row in body_rows) * scale
        title_height = sum(frame.heights[row] for row in rows if row in frame.title_rows) * scale
        self.page_drawings = []
        html = []
        if getattr(frame.worksheet, "_charts", []):
            self.fonts.width(FontRequest(OFFICE_KOREAN_FAMILY, OFFICE_KOREAN_FAMILY, 10), "가")
        for chart in getattr(frame.worksheet, "_charts", []):
            box = drawing_box(chart.anchor, frame, columns, body_rows, scale)
            if box is not None:
                html.append(chart_html(chart, box, self.values, self.palette, self.preview))
                self.page_drawings.append({"chart": chart_title(chart) or None, "type": chart_kind(chart), "shown": "whole" if is_whole(box, page_width, body_height) else "part, cut at the page edge"})
        for image in getattr(frame.worksheet, "_images", []):
            box = drawing_box(image.anchor, frame, columns, body_rows, scale, image)
            if box is not None:
                html.append(image_html(image, box))
                self.page_drawings.append({"image": True, "shown": "whole" if is_whole(box, page_width, body_height) else "part, cut at the page edge"})
        if not html:
            return ""
        clip = {"position": "absolute", "left": "0", "top": pixels(title_height), "width": pixels(page_width), "height": pixels(body_height), "overflow": "hidden"}
        return f"<div{style_attribute(clip)}>{''.join(html)}</div>"


def page_geometry(worksheet) -> PageGeometry:
    setup = worksheet.page_setup
    width_inches, height_inches = PAPER_BY_SPREADSHEET_CODE.get(int(setup.paperSize or DEFAULT_PAPER.spreadsheet_code), DEFAULT_PAPER).inches
    if page_orientation(worksheet) == "landscape":
        width_inches, height_inches = height_inches, width_inches
    margins = worksheet.page_margins
    return PageGeometry(
        width=inches_to_pixels(width_inches),
        height=inches_to_pixels(height_inches),
        margin_top=inches_to_pixels(margins.top if margins.top is not None else 0.75),
        margin_right=inches_to_pixels(margins.right if margins.right is not None else 0.7),
        margin_bottom=inches_to_pixels(margins.bottom if margins.bottom is not None else 0.75),
        margin_left=inches_to_pixels(margins.left if margins.left is not None else 0.7),
        header_distance=inches_to_pixels(margins.header if margins.header is not None else 0.3),
        footer_distance=inches_to_pixels(margins.footer if margins.footer is not None else 0.3),
    )


def page_orientation(worksheet) -> str:
    return "landscape" if worksheet.page_setup.orientation == "landscape" else "portrait"


def page_content(page_number: int, frame: SheetFrame, columns: list[int], rows: list[int], drawings: list[dict]) -> dict:
    body_rows = [row for row in rows if row not in frame.title_rows] or rows
    cells = f"{get_column_letter(columns[0])}{body_rows[0]}:{get_column_letter(columns[-1])}{body_rows[-1]}"
    content = {"page": page_number, "sheet": frame.worksheet.title, "cells": cells}
    titles = [row for row in rows if row in frame.title_rows]
    if titles:
        content["repeatedRows"] = f"{titles[0]}:{titles[-1]}"
    if drawings:
        content["drawings"] = drawings
    return content


def text_decoration(font) -> str | None:
    return " ".join(name for name, flag in (("underline", font.u and font.u != "none"), ("line-through", font.strike)) if flag) or None


def stacked(text: str) -> str:
    return "\n".join(text)


def rotation_style(rotation: int | None) -> dict:
    if not rotation or rotation == STACKED_ROTATION:
        return {}
    return {"display": "inline-block", "transform": f"rotate({-rotation_degrees(rotation)}deg)"}


def font_request(font, scale: float = 1.0) -> FontRequest:
    return FontRequest(font.name or DEFAULT_FONT.name, OFFICE_KOREAN_FAMILY, (font.sz or DEFAULT_FONT.sz) * scale, bool(font.b))


def fill_color(cell, palette: tuple) -> str | None:
    fill = cell.fill
    if getattr(fill, "patternType", None) != "solid":
        return None
    return css_color(fill.fgColor, palette)


def border_value(side, palette: tuple) -> str | None:
    if side is None or not side.style:
        return None
    width, style = BORDER_STYLES.get(side.style, ("1px", "solid"))
    return f"{width} {style} {css_color(side.color, palette) or '#000000'}"


def neighbor_border(worksheet, row: int, column: int, side: str):
    if side == "left" and column > 1:
        border = worksheet.cell(row=row, column=column - 1).border.right
    elif side == "top" and row > 1:
        border = worksheet.cell(row=row - 1, column=column).border.bottom
    else:
        return None
    return border if border is not None and border.style else None


def spill_span(frame: SheetFrame, row: int, column: int, columns: list[int]) -> int:
    span = 1
    for candidate in columns[columns.index(column) + 1:]:
        if frame.values.cell(row=row, column=candidate).value not in (None, "") or (row, candidate) in frame.covered:
            break
        span += 1
    return span


def print_bounds(worksheet, sheet_values) -> tuple[int, int, int, int] | None:
    area = worksheet.print_area
    if area:
        first = area[0] if isinstance(area, (list, tuple)) else str(area).split(",")[0]
        return range_boundaries(first.split("!")[-1].replace("$", "").replace("'", ""))
    return used_bounds(worksheet, sheet_values, [*sheet_charts(worksheet), *getattr(worksheet, "_images", [])])


def used_bounds(worksheet, sheet_values, drawings: list) -> tuple[int, int, int, int] | None:
    used = [(cell.row, cell.column) for row in sheet_values.iter_rows() for cell in row if cell.value not in (None, "")]
    used += [(merged.max_row, merged.max_col) for merged in worksheet.merged_cells.ranges]
    used += [corner for top, left, bottom, right in drawing_spans(worksheet, drawings) for corner in ((top, left), (bottom, right))]
    if not used:
        return None
    return min(column for _, column in used), min(row for row, _ in used), max(column for _, column in used), max(row for row, _ in used)


def sheet_charts(worksheet) -> list:
    return list(getattr(worksheet, "_charts", []))


def drawing_spans(worksheet, drawings: list) -> list[tuple[int, int, int, int]]:
    spans = [drawing_span(worksheet, drawing) for drawing in drawings]
    return [span for span in spans if span is not None]


def drawing_span(worksheet, drawing) -> tuple[int, int, int, int] | None:
    start = getattr(drawing.anchor, "_from", None)
    end = getattr(drawing.anchor, "to", None)
    extent = getattr(drawing.anchor, "ext", None)
    if start is not None and end is not None:
        return start.row + 1, start.col + 1, max(start.row + 1, last_cell_reached(end.row, end.rowOff)), max(start.col + 1, last_cell_reached(end.col, end.colOff))
    if start is not None and extent is not None:
        columns = cells_spanned(start.col + 1, emu_to_pixels(extent.width + (start.colOff or 0)), sheet_column_pixels(worksheet))
        rows = cells_spanned(start.row + 1, emu_to_pixels(extent.height + (start.rowOff or 0)), lambda row: sheet_row_pixels(worksheet, row))
        return start.row + 1, start.col + 1, start.row + rows, start.col + columns
    return None


def last_cell_reached(marker_index: int, offset_emu: int | None) -> int:
    return marker_index + 1 if emu_to_pixels(offset_emu or 0) >= VISIBLE_PIXELS else marker_index


def cells_spanned(first: int, length: float, pixels_of) -> int:
    count, covered = 0, 0.0
    while covered < length - VISIBLE_PIXELS:
        covered += max(pixels_of(first + count), 1)
        count += 1
    return max(count, 1)


def sheet_column_pixels(worksheet):
    settings = column_dimension_map(worksheet)
    default = default_column_width(worksheet)
    return lambda column: 0 if settings.get(column, {}).get("hidden") else column_pixels(settings.get(column, {}).get("width") or default)


def sheet_row_pixels(worksheet, row: int) -> float:
    if row_hidden(worksheet, row):
        return 0
    dimension = worksheet.row_dimensions.get(row) if hasattr(worksheet.row_dimensions, "get") else None
    if dimension is not None and dimension.height:
        return points_to_pixels(dimension.height)
    return default_row_pixels(worksheet)


def column_dimension_map(worksheet) -> dict[int, dict]:
    settings = {}
    for dimension in worksheet.column_dimensions.values():
        start = dimension.min or 0
        end = dimension.max or start
        for column in range(start, end + 1):
            settings[column] = {"width": dimension.width if dimension.customWidth or dimension.width else None, "hidden": dimension.hidden}
    return settings


def row_hidden(worksheet, row: int) -> bool:
    dimension = worksheet.row_dimensions.get(row) if hasattr(worksheet.row_dimensions, "get") else None
    return bool(dimension is not None and dimension.hidden)


def default_row_pixels(worksheet) -> float:
    return points_to_pixels(worksheet.sheet_format.defaultRowHeight or DEFAULT_ROW_POINTS)


def default_column_width(worksheet) -> float:
    return worksheet.sheet_format.defaultColWidth or DEFAULT_COLUMN_CHARACTERS


def column_pixels(characters: float) -> float:
    return int(((256 * characters + int(128 / MAXIMUM_DIGIT_PIXELS)) / 256) * MAXIMUM_DIGIT_PIXELS)


def merge_maps(worksheet) -> tuple[dict, set]:
    merges, covered = {}, set()
    for merged in worksheet.merged_cells.ranges:
        merges[(merged.min_row, merged.min_col)] = (merged.max_row, merged.max_col)
        covered.update((row, column) for row in range(merged.min_row, merged.max_row + 1) for column in range(merged.min_col, merged.max_col + 1))
    return merges, covered


def print_title_rows(worksheet, rows: list[int]) -> list[int]:
    titles = worksheet.print_title_rows
    if not titles:
        return []
    start, _, end = titles.replace("$", "").partition(":")
    return [row for row in rows if int(start) <= row <= int(end or start)]


def sheet_print_issues(workbook) -> list[Issue]:
    previewer = SheetPreviewer(workbook, workbook, "", theme_palette(workbook.loaded_theme), FontRegistry(), Preview(title=""))
    layouts = [(worksheet, previewer.print_layout(worksheet, worksheet)) for worksheet in workbook.worksheets if worksheet.sheet_state == "visible"]
    return [issue for worksheet, layout in layouts if layout is not None for issue in print_layout_issues(worksheet, layout)]


def print_layout_issues(worksheet, layout: PrintLayout) -> list[Issue]:
    return [*print_width_issues(worksheet, layout), *print_size_issues(worksheet, layout)]


def print_width_issues(worksheet, layout: PrintLayout) -> list[Issue]:
    pages_wide = len(layout.column_chunks)
    if pages_wide <= chosen_width_pages(worksheet):
        return []
    first_page = layout.column_chunks[0]
    rest = [column for chunk in layout.column_chunks[1:] for column in chunk]
    message = (
        f"{worksheet.title} prints {pages_wide} pages wide: columns {get_column_letter(first_page[0])}-{get_column_letter(first_page[-1])} fill the first page "
        f"and {get_column_letter(rest[0])}-{get_column_letter(rest[-1])} print on {'a page' if pages_wide == 2 else 'pages'} of their own, so each printed row is cut apart"
    )
    plan = readable_print_plan(worksheet, layout, body_points(layout.frame))
    return [print_layout_issue(SHEET_PRINTS_WIDE, message, worksheet, plan)]


def print_size_issues(worksheet, layout: PrintLayout) -> list[Issue]:
    points = body_points(layout.frame)
    if layout.scale >= readable_scale(points):
        return []
    plan = readable_print_plan(worksheet, layout, points)
    message = (
        f"{worksheet.title} prints at {percent(layout.scale)} of full size, so its {format_points(points)} pt body text prints at {format_points(points * layout.scale)} pt, "
        f"under the {READABLE_PRINT_POINTS} pt that stays readable on paper: {shrink_cause(worksheet, layout, plan)}"
    )
    return [print_layout_issue(SHEET_PRINTS_SMALL, message, worksheet, plan)]


def print_layout_issue(kind, message: str, worksheet, plan: PrintPlan) -> Issue:
    return kind.issue(message, worksheet.title, plan_suggestion(worksheet, plan), plan_operations(worksheet, plan))


def readable_scale(points: float) -> float:
    return min(1.0, READABLE_PRINT_POINTS / points)


def body_points(frame: SheetFrame) -> float:
    rows, columns = set(frame.rows), set(frame.columns)
    sizes = Counter(
        cell.font.sz or DEFAULT_FONT.sz
        for row in frame.values.iter_rows() for cell in row
        if cell.value not in (None, "") and cell.row in rows and cell.column in columns
    )
    return float(sizes.most_common(1)[0][0] if sizes else DEFAULT_FONT.sz)


def readable_print_plan(worksheet, layout: PrintLayout, points: float) -> PrintPlan:
    moves, extent = beside_chart_moves(worksheet, layout.frame)
    minimum = readable_scale(points)
    current = print_scale(worksheet, extent, layout.geometry)
    if current >= minimum and pages_across(extent.column_widths, current, layout.geometry.content_width) <= chosen_width_pages(worksheet):
        return PrintPlan(moves, None, current, points)
    fit, scale = readable_fit(worksheet, extent, layout.geometry, minimum)
    return PrintPlan(moves, fit, scale, points)


def beside_chart_moves(worksheet, frame: SheetFrame) -> tuple[list[ChartMove], PrintExtent]:
    charts = sheet_charts(worksheet)
    images = list(getattr(worksheet, "_images", []))
    content = used_bounds(worksheet, frame.values, images)
    beside = [] if worksheet.print_area else [index for index, chart in enumerate(charts) if is_beside(drawing_span(worksheet, chart), content)]
    if not beside:
        return [], frame_extent(frame)
    kept = used_bounds(worksheet, frame.values, [*images, *(chart for index, chart in enumerate(charts) if index not in beside)])
    sizes = [chart_pixels(charts[index], frame) for index in beside]
    row_pixels = default_row_pixels(worksheet)
    return stacked_moves(beside, sizes, kept, row_pixels), moved_extent(worksheet, frame, kept, sizes, row_pixels)


def is_beside(span: tuple[int, int, int, int] | None, content: tuple[int, int, int, int] | None) -> bool:
    if span is None or content is None:
        return False
    top, _, _, right = span
    _, _, last_column, last_row = content
    return right > last_column and top <= last_row


def chart_pixels(chart, frame: SheetFrame) -> tuple[float, float]:
    box = drawing_box(chart.anchor, frame, frame.columns, frame.rows, 1.0) if hasattr(chart.anchor, "_from") else None
    if box is not None:
        return box.width, box.height
    return millimetres_to_pixels(chart.width * 10), millimetres_to_pixels(chart.height * 10)


def stacked_moves(indexes: list[int], sizes: list[tuple[float, float]], kept: tuple[int, int, int, int], row_pixels: float) -> list[ChartMove]:
    first_column, _, _, last_row = kept
    moves, row = [], last_row + DEFAULT_ANCHOR_GAP
    for index, (_, height) in zip(indexes, sizes):
        moves.append(ChartMove(index, f"{get_column_letter(first_column)}{row}"))
        row += int(height // row_pixels) + DEFAULT_ANCHOR_GAP
    return moves


def moved_extent(worksheet, frame: SheetFrame, kept: tuple[int, int, int, int], sizes: list[tuple[float, float]], row_pixels: float) -> PrintExtent:
    first_column, _, last_column, last_row = kept
    widths = [frame.widths[column] for column in frame.columns if first_column <= column <= last_column]
    widest = max(width for width, _ in sizes)
    default_width = max(column_pixels(default_column_width(worksheet)), 1)
    column = last_column + 1
    while sum(widths) < widest:
        widths.append(frame.widths.get(column, default_width))
        column += 1
    kept_height = sum(frame.heights[row] for row in frame.rows if row <= last_row)
    moved_height = sum(height + (DEFAULT_ANCHOR_GAP - 1) * row_pixels for _, height in sizes)
    return PrintExtent(widths, kept_height + moved_height)


def frame_extent(frame: SheetFrame) -> PrintExtent:
    return PrintExtent([frame.widths[column] for column in frame.columns], sum(frame.heights.values()))


def readable_fit(worksheet, extent: PrintExtent, geometry: PageGeometry, minimum: float) -> tuple[PageFit, float]:
    current = page_orientation(worksheet)
    orientations = list(dict.fromkeys([current, "landscape"]))
    widest_pages = pages_across(extent.column_widths, 1.0, geometry.content_width)
    for pages in range(1, widest_pages + 1):
        for orientation in orientations:
            fit = PageFit(orientation, pages)
            scale = fitted_scale(extent, oriented(geometry, orientation, current), pages, 0)
            if scale >= minimum:
                return fit, scale
    return fit, scale


def oriented(geometry: PageGeometry, orientation: str, current: str) -> PageGeometry:
    return geometry if orientation == current else replace(geometry, width=geometry.height, height=geometry.width)


def shrink_cause(worksheet, layout: PrintLayout, plan: PrintPlan) -> str:
    columns = f"columns {get_column_letter(layout.frame.columns[0])}-{get_column_letter(layout.frame.columns[-1])}"
    if plan.moves:
        charts = sheet_charts(worksheet)
        named = ", ".join(f"chart {move.chart} at {anchor_cell(charts[move.chart])}" for move in plan.moves)
        return f"{named}, beside the data, widens the printed page to {columns}"
    if is_fitted_to_pages(worksheet):
        return f"its page setup fits {columns} onto {fitted_pages_text(worksheet)}"
    return f"its page setup prints at a scale of {worksheet.page_setup.scale}%"


def fitted_pages_text(worksheet) -> str:
    width, height = fitted_page_count(worksheet.page_setup.fitToWidth), fitted_page_count(worksheet.page_setup.fitToHeight)
    parts = [f"{pages_text(width)} wide" if width else "", f"{pages_text(height)} tall" if height else ""]
    return " and ".join(part for part in parts if part)


def plan_suggestion(worksheet, plan: PrintPlan) -> str:
    steps = [f'move chart {move.chart} under the data with edit_chart "anchor": "{move.anchor}"' for move in plan.moves]
    if plan.fit is not None:
        turned = " in landscape" if plan.fit.orientation != page_orientation(worksheet) else ""
        steps.append(f"fit the columns onto {pages_text(plan.fit.width_pages)} wide{turned}")
    return f"apply the operations in fix: they {' and '.join(steps)}, so the sheet prints at {percent(plan.scale)} and its body text at {format_points(plan.body_points * plan.scale)} pt"


def plan_operations(worksheet, plan: PrintPlan) -> list[dict]:
    moves = [{"op": "edit_chart", "sheet": worksheet.title, "chart": move.chart, "anchor": move.anchor} for move in plan.moves]
    return moves + ([page_fit_operation(worksheet, plan.fit)] if plan.fit is not None else [])


def page_fit_operation(worksheet, fit: PageFit) -> dict:
    operation = {"op": "set_page_setup", "sheet": worksheet.title}
    if fit.orientation != page_orientation(worksheet):
        operation["orientation"] = fit.orientation
    operation["fitToWidth"] = fit.width_pages
    if is_fitted_to_pages(worksheet) and fitted_page_count(worksheet.page_setup.fitToHeight):
        operation["fitToHeight"] = 0
    return operation


def pages_text(count: int) -> str:
    return "1 page" if count == 1 else f"{count} pages"


def percent(scale: float) -> str:
    return f"{scale:.0%}"


def format_points(points: float) -> str:
    return f"{points:.1f}".removesuffix(".0")


def chosen_width_pages(worksheet) -> int:
    if not is_fitted_to_pages(worksheet) or worksheet.page_setup.fitToWidth is None:
        return 1
    return max(1, int(worksheet.page_setup.fitToWidth))


def print_scale(worksheet, extent: PrintExtent, geometry: PageGeometry) -> float:
    if not is_fitted_to_pages(worksheet):
        return (worksheet.page_setup.scale or 100) / 100
    return fitted_scale(extent, geometry, fitted_page_count(worksheet.page_setup.fitToWidth), fitted_page_count(worksheet.page_setup.fitToHeight))


def fitted_page_count(pages) -> int:
    return EXCEL_DEFAULT_FIT_PAGES if pages is None else int(pages)


def fitted_scale(extent: PrintExtent, geometry: PageGeometry, width_pages: int, height_pages: int) -> float:
    scales = [1.0]
    if width_pages:
        scales.append(width_fit_scale(extent.column_widths, geometry.content_width, width_pages))
    if height_pages and extent.height:
        scales.append(geometry.content_height * height_pages / extent.height)
    return min(scales)


def width_fit_scale(widths: list[float], content_width: float, pages: int) -> float:
    total = sum(widths)
    if not total or pages_across(widths, 1.0, content_width) <= pages:
        return 1.0
    fitting, failing = 0.0, min(1.0, content_width * pages / total)
    if pages_across(widths, failing, content_width) <= pages:
        return failing
    for _ in range(FIT_SEARCH_STEPS):
        middle = (fitting + failing) / 2
        fitting, failing = (middle, failing) if pages_across(widths, middle, content_width) <= pages else (fitting, middle)
    return fitting


def pages_across(widths: list[float], scale: float, content_width: float) -> int:
    return len(chunks(list(range(len(widths))), [width * scale for width in widths], content_width, set()))


def chunks(items: list[int], sizes: list[float], limit: float, breaks_after: set) -> list[list[int]]:
    groups, current, used = [], [], 0.0
    for item, size in zip(items, sizes):
        if current and used + size > limit + 0.5:
            groups.append(current)
            current, used = [], 0.0
        current.append(item)
        used += size
        if item in breaks_after:
            groups.append(current)
            current, used = [], 0.0
    if current:
        groups.append(current)
    return groups or [[]]


def header_text(text: str) -> str:
    cleaned = HEADER_FORMATTING.sub("", text.replace("{", "{{").replace("}", "}}"))
    for code, replacement in HEADER_CODES:
        cleaned = cleaned.replace(code, replacement)
    return cleaned


def script_spans(request: FontRequest, text: str) -> str:
    if not draws_scripts_apart(request):
        return escaped(text)
    return "".join(f'<span{style_attribute({"font-family": script_font_family(request, piece)})}>{escaped(piece)}</span>' for piece in script_runs(text))
