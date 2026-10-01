from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter

from excel_functions import written_value
from formula_cache import cell_labels, evaluation_issues, listed, save_workbook_with_values
from formula_references import COLUMN_AXIS, ROW_AXIS, Shift
from office_operations import Change, OperationSet
from office_result import INVALID_VALUE, OfficeFailure
from sheet_charts import plan_add_chart, plan_delete_chart, plan_edit_chart
from sheet_data import plan_find_replace, plan_sort_range
from sheet_definitions import CIRCULAR_REFERENCE, CONTENT_DROPPED, CONTENT_WOULD_BE_LOST, OPERATIONS
from sheet_formatting import plan_format_range, plan_hide_columns, plan_hide_rows, plan_hide_sheet, plan_merge_cells, plan_set_row_height, plan_unmerge_cells
from sheet_objects import plan_add_table, plan_set_comment, plan_set_hyperlink, plan_set_page_setup
from sheet_pivots import plan_add_pivot_table
from sheet_rules import plan_add_conditional_format, plan_add_data_validation, plan_clear_conditional_formats, plan_clear_data_validations
from sheet_drawings import plan_add_image, plan_add_shape
from sheet_ranges import plan_clear_range, plan_convert_to_values, plan_copy_range, plan_fill_range, plan_set_filter_criteria
from sheet_sparklines import plan_add_sparklines
from sheet_workbook import plan_add_defined_name, plan_delete_defined_name, plan_delete_sheet, plan_duplicate_sheet, plan_move_sheet, plan_protect_sheet, validate_sheet_name
from sheet_styling import data_bounds, style_written_cells
from workbook_access import column_index, open_workbook, parse_cell, parse_range, resolve_sheet, sheet_of
from workbook_fidelity import EditRecord, preserve_source_content
from workbook_package import Package, read_package, write_package
from workbook_structure import isolate_column, rename_sheet_references, shift_workbook
from workbook_values import evaluate_workbook
from written_cells import require_writable


@dataclass
class SheetEditing:
    workbook: object
    source: Package | None
    allows_loss: bool = False
    source_path: str | None = None
    record: EditRecord = field(default_factory=EditRecord)
    package_patches: list = field(default_factory=list)
    reserved_names: set = field(default_factory=set)
    warnings: list = field(default_factory=list)


def load_editing(path: str, allows_loss: bool = False) -> SheetEditing:
    workbook = open_workbook(path)
    return SheetEditing(workbook, read_package(path), allows_loss, path)


def save_editing(editing: SheetEditing, path: str) -> list:
    evaluation = save_workbook_with_values(editing.workbook, path)
    refuse_new_circular_references(editing, evaluation)
    issues = [issue for issue in evaluation_issues(evaluation) if issue.kind is not CIRCULAR_REFERENCE]
    output = read_package(path)
    for patch in editing.package_patches:
        patch(output)
    lost = preserve_source_content(editing.source, output, editing.record) if editing.source is not None else []
    if lost and not editing.allows_loss:
        raise OfficeFailure(CONTENT_WOULD_BE_LOST.issue(f"saving would drop what the editor cannot carry: {', '.join(lost)}", lost[0]))
    write_package(output, path)
    return issues + editing.warnings + ([CONTENT_DROPPED.issue(f"dropped what the editor cannot carry: {', '.join(lost)}", lost[0])] if lost else [])


def refuse_new_circular_references(editing: SheetEditing, evaluation) -> None:
    if not evaluation.circular or editing.workbook.calculation.iterate:
        return
    existing = evaluate_workbook(editing.source_path).circular if editing.source_path else []
    if len(evaluation.circular) <= len(existing):
        return
    cells = cell_labels([key for key in evaluation.circular if key not in existing] or evaluation.circular)
    raise OfficeFailure(CIRCULAR_REFERENCE.issue(f"{len(cells)} formula cells would read their own value: {listed(cells)}; nothing was written", cells[0]))


def on_workbook(planner):
    return lambda editing, operation, location: planner(editing.workbook, operation, location)


def store_value(cell, value, value_type) -> None:
    if value_type == "text" and isinstance(value, str):
        cell.value = value
        cell.data_type = "s"
        cell.quotePrefix = True
        return
    cell.value = written_value(value)


def plan_set_cell(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    row, column = parse_cell(operation["cell"], f"{location}.cell")
    require_writable(operation.get("value"), f"{location}.value", keeps_text=operation.get("type") == "text")

    def change() -> str:
        existing = data_bounds(worksheet)
        cell = worksheet.cell(row=row, column=column)
        store_value(cell, operation.get("value"), operation.get("type"))
        style_written_cells(worksheet, [cell], existing)
        return f"set {worksheet.title}!{operation['cell'].upper()}"
    return change


def plan_set_range(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    first_row, first_column = parse_cell(operation["cell"], f"{location}.cell")
    for row_offset, values in enumerate(operation["values"]):
        for column_offset, value in enumerate(values):
            require_writable(value, f"{location}.values[{row_offset}][{column_offset}]", keeps_text=operation.get("type") == "text")

    def change() -> str:
        existing = data_bounds(worksheet)
        written = []
        for row_offset, values in enumerate(operation["values"]):
            for column_offset, value in enumerate(values):
                cell = worksheet.cell(row=first_row + row_offset, column=first_column + column_offset)
                store_value(cell, value, operation.get("type"))
                written.append(cell)
        style_written_cells(worksheet, written, existing)
        rows = len(operation["values"])
        return f"wrote {rows} rows from {worksheet.title}!{operation['cell'].upper()}"
    return change


def plan_set_column_width(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    index = column_index(operation["column"], f"{location}.column")

    def change() -> str:
        dimension = isolate_column(worksheet, index) if has_group_over(worksheet, index) else worksheet.column_dimensions[get_column_letter(index)]
        dimension.width = operation["width"]
        return f"set column {get_column_letter(index)} of {worksheet.title} to width {operation['width']}"
    return change


def has_group_over(worksheet, index: int) -> bool:
    return any((dimension.min or 1) <= index <= (dimension.max or dimension.min or 1) for dimension in worksheet.column_dimensions.values())


def plan_freeze_panes(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    cell = (operation.get("cell") or "").strip()
    if cell:
        parse_cell(cell, f"{location}.cell")

    def change() -> str:
        worksheet.freeze_panes = cell.replace("$", "").upper() or None
        return f"froze panes at {cell.upper()} on {worksheet.title}" if cell else f"unfroze panes on {worksheet.title}"
    return change


def plan_set_auto_filter(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    range_text = (operation.get("range") or "").strip()
    if range_text:
        parse_range(range_text, f"{location}.range")

    def change() -> str:
        worksheet.auto_filter = AutoFilter(ref=range_text.replace("$", "").upper() or None)
        return f"set the filter of {worksheet.title} to {range_text.upper()}" if range_text else f"removed the filter of {worksheet.title}"
    return change


def plan_add_sheet(workbook, operation: dict, location: str) -> Change:
    validate_sheet_name(workbook, operation["name"], f"{location}.name")

    def change() -> str:
        workbook.create_sheet(title=operation["name"], index=operation.get("index"))
        return f"added sheet {operation['name']}"
    return change


def plan_rename_sheet(editing: SheetEditing, operation: dict, location: str) -> Change:
    workbook = editing.workbook
    worksheet = resolve_sheet(workbook, operation["sheet"], f"{location}.sheet")
    validate_sheet_name(workbook, operation["name"], f"{location}.name", renaming=worksheet.title)

    def change() -> str:
        old_name = worksheet.title
        rewritten = rename_sheet_references(workbook, old_name, operation["name"])
        worksheet.title = operation["name"]
        editing.record.renamed(old_name, operation["name"])
        return f"renamed sheet {old_name} to {operation['name']}; rewrote {rewritten} formulas"
    return change


def plan_structure(axis: str, sign: int, noun: str, verb: str):
    def plan(editing: SheetEditing, operation: dict, location: str) -> Change:
        workbook = editing.workbook
        worksheet = sheet_of(workbook, operation, location)
        at = operation["at"] if axis == ROW_AXIS else column_index(operation["at"], f"{location}.at")
        count = operation.get("count", 1)
        shift = Shift(worksheet.title, axis, at, sign * count)

        def change() -> str:
            rewritten = shift_workbook(workbook, worksheet, shift)
            editing.record.shifted(shift)
            return f"{verb} {count} {noun} {'before' if sign > 0 else 'from'} {axis} {operation['at']} of {worksheet.title}; rewrote {rewritten} formulas"
        return change
    return plan


def plan_recalculate(workbook, operation: dict, location: str) -> Change:
    def change() -> str:
        return "stored freshly computed values for every formula the workbook can compute"
    return change


SHEET_OPERATIONS = OperationSet(OPERATIONS, {
    "set_cell": on_workbook(plan_set_cell),
    "set_range": on_workbook(plan_set_range),
    "format_range": on_workbook(plan_format_range),
    "set_column_width": on_workbook(plan_set_column_width),
    "freeze_panes": on_workbook(plan_freeze_panes),
    "set_auto_filter": on_workbook(plan_set_auto_filter),
    "add_sheet": on_workbook(plan_add_sheet),
    "rename_sheet": plan_rename_sheet,
    "insert_rows": plan_structure(ROW_AXIS, 1, "rows", "inserted"),
    "delete_rows": plan_structure(ROW_AXIS, -1, "rows", "deleted"),
    "insert_columns": plan_structure(COLUMN_AXIS, 1, "columns", "inserted"),
    "delete_columns": plan_structure(COLUMN_AXIS, -1, "columns", "deleted"),
    "recalculate": on_workbook(plan_recalculate),
    "add_chart": plan_add_chart,
    "edit_chart": plan_edit_chart,
    "delete_chart": on_workbook(plan_delete_chart),
    "merge_cells": on_workbook(plan_merge_cells),
    "unmerge_cells": on_workbook(plan_unmerge_cells),
    "set_row_height": on_workbook(plan_set_row_height),
    "hide_rows": on_workbook(plan_hide_rows),
    "hide_columns": on_workbook(plan_hide_columns),
    "hide_sheet": on_workbook(plan_hide_sheet),
    "sort_range": on_workbook(plan_sort_range),
    "find_replace": on_workbook(plan_find_replace),
    "add_conditional_format": on_workbook(plan_add_conditional_format),
    "clear_conditional_formats": on_workbook(plan_clear_conditional_formats),
    "add_data_validation": on_workbook(plan_add_data_validation),
    "clear_data_validations": on_workbook(plan_clear_data_validations),
    "add_table": on_workbook(plan_add_table),
    "set_hyperlink": on_workbook(plan_set_hyperlink),
    "set_comment": on_workbook(plan_set_comment),
    "set_page_setup": on_workbook(plan_set_page_setup),
    "add_sparklines": plan_add_sparklines,
    "delete_sheet": plan_delete_sheet,
    "duplicate_sheet": on_workbook(plan_duplicate_sheet),
    "move_sheet": on_workbook(plan_move_sheet),
    "add_defined_name": on_workbook(plan_add_defined_name),
    "delete_defined_name": on_workbook(plan_delete_defined_name),
    "fill_range": on_workbook(plan_fill_range),
    "copy_range": on_workbook(plan_copy_range),
    "clear_range": on_workbook(plan_clear_range),
    "convert_to_values": on_workbook(plan_convert_to_values),
    "set_filter_criteria": on_workbook(plan_set_filter_criteria),
    "protect_sheet": on_workbook(plan_protect_sheet),
    "add_image": on_workbook(plan_add_image),
    "add_shape": plan_add_shape,
    "add_pivot_table": plan_add_pivot_table,
}, sequential=True)
