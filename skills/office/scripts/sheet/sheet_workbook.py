from __future__ import annotations

from copy import copy
import re

from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.protection import SheetProtection

from office_operations import OPERATION_NOT_APPLICABLE, TARGET_NOT_FOUND, Change
from office_result import INVALID_VALUE, MISSING_FIELD, OfficeFailure
from office_schema import closest_name
from formula_references import deleted_sheet_reference, quote_sheet_name
from workbook_access import parse_range, resolve_sheet, sheet_of
from workbook_structure import rewrite_chart_references, rewrite_defined_names, rewrite_formulas


DEFINED_NAME_PATTERN = re.compile(r"^[^\W\d][\w.]{0,254}$")
CELL_LIKE_NAME = re.compile(r"^([A-Za-z]{1,3}\d+|[Rr]\d*[Cc]\d*)$")
COPY_SUFFIX = " ({number})"
MAXIMUM_SHEET_NAME_LENGTH = 31
FORBIDDEN_SHEET_NAME_CHARACTERS = set("[]:*?/\\")


def validate_sheet_name(workbook, name: str, location: str, renaming: str | None = None) -> None:
    field = location.rsplit(".", 1)[-1]
    if len(name) > MAXIMUM_SHEET_NAME_LENGTH:
        shortened = name[:MAXIMUM_SHEET_NAME_LENGTH].rstrip()
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: Excel sheet names hold at most {MAXIMUM_SHEET_NAME_LENGTH} characters, and {name!r} has {len(name)}", location, f'use a name of {MAXIMUM_SHEET_NAME_LENGTH} characters or fewer, such as "{field}": "{shortened}"'))
    if FORBIDDEN_SHEET_NAME_CHARACTERS & set(name) or name.startswith("'") or name.endswith("'"):
        cleaned = "".join(character for character in name if character not in FORBIDDEN_SHEET_NAME_CHARACTERS).strip("'") or "Sheet"
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: a sheet name cannot hold any of [ ] : * ? / \\ or start or end with an apostrophe", location, f'use "{field}": "{cleaned}"'))
    taken = [title for title in workbook.sheetnames if title.casefold() == name.casefold() and title != renaming]
    if taken:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: the workbook already has a sheet named {taken[0]!r}", location, f'pick a name no sheet has, such as "{field}": "{copy_title(workbook, name)}"'))


def visible_sheets(workbook) -> list:
    return [worksheet for worksheet in workbook.worksheets if worksheet.sheet_state == "visible"]


def plan_delete_sheet(editing, operation: dict, location: str) -> Change:
    workbook = editing.workbook
    worksheet = resolve_sheet(workbook, operation["sheet"], f"{location}.sheet")
    if visible_sheets(workbook) == [worksheet]:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}: {worksheet.title} is the only visible sheet, and a workbook keeps one", f"{location}.sheet"))

    def change() -> str:
        title = worksheet.title
        rewrite = lambda reference, host: deleted_sheet_reference(reference, title)
        workbook.remove(worksheet)
        broken = rewrite_formulas(workbook, rewrite)
        rewrite_defined_names(workbook, rewrite)
        rewrite_chart_references(workbook, rewrite)
        editing.record.deleted(title)
        return f"deleted sheet {title}" + (f"; {broken} formulas that read it now show #REF!" if broken else "")
    return change


def copy_title(workbook, title: str) -> str:
    for number in range(2, 1000):
        suffix = COPY_SUFFIX.format(number=number)
        candidate = title[:MAXIMUM_SHEET_NAME_LENGTH - len(suffix)] + suffix
        if candidate.casefold() not in {name.casefold() for name in workbook.sheetnames}:
            return candidate
    raise OfficeFailure(INVALID_VALUE.issue(f"no free copy name for {title}", title))


def plan_duplicate_sheet(workbook, operation: dict, location: str) -> Change:
    worksheet = resolve_sheet(workbook, operation["sheet"], f"{location}.sheet")
    name = operation.get("name") or copy_title(workbook, worksheet.title)
    validate_sheet_name(workbook, name, f"{location}.name")

    def change() -> str:
        duplicate = workbook.copy_worksheet(worksheet)
        duplicate.title = name
        copy_sheet_rules(worksheet, duplicate)
        workbook.move_sheet(duplicate, workbook.index(worksheet) + 1 - workbook.index(duplicate))
        left_behind = left_behind_objects(worksheet)
        return f"copied sheet {worksheet.title} to {name}" + (f"; {', '.join(left_behind)} stay on {worksheet.title}" if left_behind else "")
    return change


def copy_sheet_rules(source, duplicate) -> None:
    duplicate.freeze_panes = source.freeze_panes
    duplicate.auto_filter.ref = source.auto_filter.ref
    for formatting in source.conditional_formatting:
        for rule in formatting.rules:
            duplicate.conditional_formatting.add(str(formatting.sqref), copy(rule))
    for validation in source.data_validations.dataValidation:
        duplicate.add_data_validation(copy(validation))
    if source.print_title_rows:
        duplicate.print_title_rows = source.print_title_rows


def left_behind_objects(worksheet) -> list[str]:
    kinds = (("charts", worksheet._charts), ("images", worksheet._images), ("tables", worksheet.tables), ("pivot tables", getattr(worksheet, "_pivots", [])))
    return [name for name, objects in kinds if objects]


def plan_move_sheet(workbook, operation: dict, location: str) -> Change:
    worksheet = resolve_sheet(workbook, operation["sheet"], f"{location}.sheet")
    last = len(workbook.worksheets) - 1
    if operation["index"] > last:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.index: the workbook has {last + 1} sheets, numbered from 0 to {last}", f"{location}.index"))

    def change() -> str:
        workbook.move_sheet(worksheet, operation["index"] - workbook.index(worksheet))
        workbook.active = workbook.index(visible_sheets(workbook)[0])
        return f"moved sheet {worksheet.title} to position {operation['index']}"
    return change


def validate_defined_name(name: str, location: str) -> None:
    if not DEFINED_NAME_PATTERN.match(name) or CELL_LIKE_NAME.match(name) or name.upper() in ("TRUE", "FALSE"):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: {name!r} cannot be a name; start with a letter or _, use letters, digits, _ and ., and do not write a cell address such as A1", location))


def defined_name_target(workbook, operation: dict, location: str) -> str:
    if operation.get("range") and operation.get("value") is not None:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}: give range or value, not both", f"{location}.value"))
    if operation.get("value") is not None:
        value = operation["value"]
        return f'"{value}"' if isinstance(value, str) else str(value)
    if not operation.get("range"):
        raise OfficeFailure(MISSING_FIELD.issue(f"{location}.range: a name needs the range it stands for, or a value", f"{location}.range"))
    worksheet = sheet_of(workbook, operation, location)
    parse_range(operation["range"], f"{location}.range")
    absolute = ":".join(absolute_cell(part) for part in operation["range"].replace("$", "").upper().split(":"))
    return f"{quote_sheet_name(worksheet.title)}!{absolute}"


def absolute_cell(coordinate: str) -> str:
    letters = coordinate.rstrip("0123456789")
    return f"${letters}${coordinate[len(letters):]}"


def name_scope(workbook, operation: dict, location: str):
    return sheet_of(workbook, operation, location) if operation.get("local") else None


def plan_add_defined_name(workbook, operation: dict, location: str) -> Change:
    validate_defined_name(operation["name"], f"{location}.name")
    target = defined_name_target(workbook, operation, location)
    scope = name_scope(workbook, operation, location)
    names = scope.defined_names if scope is not None else workbook.defined_names

    def change() -> str:
        replaced = operation["name"] in names
        names[operation["name"]] = DefinedName(operation["name"], attr_text=target)
        return f"{'replaced' if replaced else 'added'} the name {operation['name']} for {target}"
    return change


def plan_delete_defined_name(workbook, operation: dict, location: str) -> Change:
    scopes = [("workbook", workbook.defined_names)] + [(worksheet.title, worksheet.defined_names) for worksheet in workbook.worksheets]
    found = [(scope, names) for scope, names in scopes if operation["name"] in names]
    if not found:
        known = sorted({name for _, names in scopes for name in names})
        nearest = closest_name(operation["name"], known)
        guess = f" (did you mean {nearest!r}?)" if nearest else ""
        raise OfficeFailure(TARGET_NOT_FOUND.issue(f"{location}.name: the workbook has no name {operation['name']!r}{guess}; it has {', '.join(known) or 'none'}", f"{location}.name"))

    def change() -> str:
        for _, names in found:
            del names[operation["name"]]
        return f"deleted the name {operation['name']}"
    return change


def plan_protect_sheet(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    protected = operation.get("protected", True)

    def change() -> str:
        worksheet.protection = SheetProtection(sheet=protected)
        if protected and operation.get("password"):
            worksheet.protection.password = operation["password"]
        return f"{'protected' if protected else 'unprotected'} {worksheet.title}"
    return change
