from __future__ import annotations

from dataclasses import dataclass, replace

from openpyxl.formula.tokenizer import Token
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

from sheet.formulas.references import MAXIMUM_COLUMN, MAXIMUM_ROW, formula_references, is_bare_name, parse_end, reference_parts, unquote_sheet_name
from sheet.formulas.tree import Call, calls_in, meaningful, parse_formula, render, text_literal


VOLATILE_REFERENCE_FUNCTIONS = frozenset(("INDIRECT", "OFFSET"))
PARAMETER_PREFIX = "_XLPM."


@dataclass(frozen=True)
class Region:
    sheet: str | None
    min_row: int
    min_column: int
    max_row: int
    max_column: int

    def contains(self, sheet: str, row: int, column: int) -> bool:
        if self.sheet is not None and self.sheet != sheet:
            return False
        return self.min_row <= row <= self.max_row and self.min_column <= column <= self.max_column


EVERYWHERE = Region(None, 1, 1, MAXIMUM_ROW, MAXIMUM_COLUMN)


class DependencyReader:
    def __init__(self, workbook):
        self.workbook = workbook
        self.tables = {table.name.casefold(): (worksheet.title, table.ref) for worksheet in workbook.worksheets for table in worksheet.tables.values()}
        self.names = {name.casefold(): (defined.attr_text or "", None) for name, defined in workbook.defined_names.items()}
        for worksheet in workbook.worksheets:
            self.names.update({(worksheet.title.casefold(), name.casefold()): (defined.attr_text or "", worksheet.title) for name, defined in worksheet.defined_names.items()})

    def regions(self, formula: str, host_sheet: str, depth: int = 0) -> list[Region]:
        nodes = parse_formula(formula)
        if nodes is None or depth > 8:
            return [EVERYWHERE]
        computed = [self.computed_reference_regions(call, host_sheet, depth) for call in calls_in(nodes) if call.function in VOLATILE_REFERENCE_FUNCTIONS]
        if any(regions is None for regions in computed):
            return [EVERYWHERE]
        written = [region for reference in formula_references(formula) for region in self.reference_regions(reference, host_sheet, depth)]
        return written + [region for regions in computed for region in regions]

    def computed_reference_regions(self, call: Call, host_sheet: str, depth: int) -> list[Region] | None:
        if call.function == "INDIRECT":
            return self.indirect_regions(call, host_sheet, depth)
        return offset_regions(call, host_sheet)

    def indirect_regions(self, call: Call, host_sheet: str, depth: int) -> list[Region] | None:
        reference = text_literal(call.arguments[0]) if call.arguments else None
        if reference is None or len(call.arguments) > 2 or len(call.arguments) == 2 and not is_true_literal(call.arguments[1]):
            return None
        return self.reference_regions(reference, host_sheet, depth) or None

    def reference_regions(self, reference: str, host_sheet: str, depth: int) -> list[Region]:
        if reference.upper().startswith(PARAMETER_PREFIX):
            return []
        if "[" in reference and "!" not in reference.split("[", 1)[0]:
            return self.table_regions(reference.split("[", 1)[0], host_sheet)
        if is_bare_name(reference):
            return self.name_regions(reference, host_sheet, depth)
        region = cell_region(reference, host_sheet)
        return [region] if region is not None else []

    def table_regions(self, table_name: str, host_sheet: str) -> list[Region]:
        found = self.tables.get(table_name.casefold())
        if found is None:
            return [Region(host_sheet.casefold(), 1, 1, MAXIMUM_ROW, MAXIMUM_COLUMN)]
        return [cell_region(found[1], found[0])]

    def name_regions(self, name: str, host_sheet: str, depth: int) -> list[Region]:
        found = self.names.get((host_sheet.casefold(), name.casefold())) or self.names.get(name.casefold())
        if found is None:
            return []
        return self.regions("=" + found[0], found[1] or host_sheet, depth + 1)


def cell_region(reference: str, host_sheet: str) -> Region | None:
    parts = reference_parts(reference)
    if parts is None or len(parts) > 2:
        return None
    sheet = unquote_sheet_name(parts[0].prefix) if parts[0].prefix is not None else host_sheet
    ends = [parse_end(part.rest) for part in parts]
    if any(end is None for end in ends):
        return None
    first, last = ends[0], ends[-1]
    return Region(
        sheet.casefold(),
        first.row if first.row is not None else 1,
        first.column if first.column is not None else 1,
        last.row if last.row is not None else MAXIMUM_ROW,
        last.column if last.column is not None else MAXIMUM_COLUMN,
    )


def is_true_literal(nodes: list) -> bool:
    significant = meaningful(nodes)
    return len(significant) == 1 and isinstance(significant[0], Token) and significant[0].value.upper() in ("TRUE", "1")


def integer_literal(nodes: list) -> int | None:
    text = render(meaningful(nodes))
    try:
        return int(text)
    except ValueError:
        return None


def offset_regions(call: Call, host_sheet: str) -> list[Region] | None:
    if not 3 <= len(call.arguments) <= 5:
        return None
    base = cell_region(render(meaningful(call.arguments[0])), host_sheet)
    numbers = [integer_literal(argument) for argument in call.arguments[1:]]
    if base is None or any(number is None for number in numbers):
        return None
    rows, columns, *size = numbers
    height = size[0] if size else base.max_row - base.min_row + 1
    width = size[1] if len(size) > 1 else base.max_column - base.min_column + 1
    top, left = base.min_row + rows, base.min_column + columns
    return [replace(base, min_row=top, min_column=left, max_row=top + height - 1, max_column=left + width - 1)]


def cell_position(coordinate: str) -> tuple[int, int]:
    letters, row = coordinate_from_string(coordinate)
    return row, column_index_from_string(letters)


def propagate(roots: set, formulas: dict, reader: DependencyReader) -> set:
    if not roots:
        return set(roots)
    unevaluated = set(roots)
    regions = {key: reader.regions(formula, key[0]) for key, formula in formulas.items() if key not in unevaluated}
    positions = {key: (key[0].casefold(), *cell_position(key[1])) for key in formulas}
    changed = True
    while changed:
        changed = False
        root_positions = [positions[key] for key in unevaluated if key in positions]
        for key, key_regions in list(regions.items()):
            if any(region.contains(*position) for region in key_regions for position in root_positions):
                unevaluated.add(key)
                del regions[key]
                changed = True
    return unevaluated
