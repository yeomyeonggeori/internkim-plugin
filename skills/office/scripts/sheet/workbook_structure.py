from __future__ import annotations

from copy import copy
from typing import Callable

from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.dimensions import ColumnDimension

from formula_references import COLUMN_AXIS, REFERENCE_ERROR, ROW_AXIS, Shift, rename_sheet_reference, rewrite_formula, shift_reference, shift_single, shift_span
from office_operations import OPERATION_NOT_APPLICABLE
from office_result import OfficeFailure


Rewrite = Callable[[str, str], str]


def shift_workbook(workbook, worksheet, shift: Shift) -> int:
    rewrite = lambda reference, host: shift_reference(reference, host, shift)
    rewritten = rewrite_formulas(workbook, rewrite)
    rewrite_defined_names(workbook, rewrite)
    rewrite_chart_references(workbook, rewrite)
    shift_sheet_structure(worksheet, shift)
    return rewritten


def rename_sheet_references(workbook, old_name: str, new_name: str) -> int:
    rewrite = lambda reference, host: rename_sheet_reference(reference, old_name, new_name)
    rewritten = rewrite_formulas(workbook, rewrite)
    rewrite_defined_names(workbook, rewrite)
    rewrite_chart_references(workbook, rewrite)
    return rewritten


def rewrite_formulas(workbook, rewrite: Rewrite) -> int:
    rewritten = 0
    for worksheet in workbook.worksheets:
        for cell in formula_cells(worksheet):
            updated = rewrite_formula_value(cell.value, worksheet.title, rewrite)
            if updated is not cell.value:
                cell.value = updated
                rewritten += 1
    return rewritten


def formula_cells(worksheet) -> list:
    return [cell for cell in worksheet._cells.values() if cell.data_type == "f"]


def rewrite_formula_value(value, host_sheet: str, rewrite: Rewrite):
    if isinstance(value, str):
        updated = rewrite_formula(value, lambda reference: rewrite(reference, host_sheet))
        return value if updated == value else updated
    updated_text = rewrite_formula(value.text, lambda reference: rewrite(reference, host_sheet))
    if updated_text == value.text:
        return value
    return type(value)(ref=value.ref, text=updated_text)


def rewrite_defined_names(workbook, rewrite: Rewrite) -> None:
    for defined in workbook.defined_names.values():
        defined.attr_text = rewrite_reference_text(defined.attr_text, "", rewrite)
    for worksheet in workbook.worksheets:
        for defined in worksheet.defined_names.values():
            defined.attr_text = rewrite_reference_text(defined.attr_text, worksheet.title, rewrite)


def rewrite_reference_text(text: str | None, host_sheet: str, rewrite: Rewrite) -> str | None:
    if not text:
        return text
    return rewrite_formula("=" + text, lambda reference: rewrite(reference, host_sheet))[1:]


def chart_series(chart) -> list:
    return [series for sub_chart in chart._charts for series in sub_chart.series]


def series_references(series) -> list:
    sources = [getattr(series, name, None) for name in ("tx", "cat", "val", "xVal", "yVal", "bubbleSize")]
    holders = [getattr(source, kind, None) for source in sources if source is not None for kind in ("numRef", "strRef", "multiLvlStrRef")]
    return [holder for holder in holders if holder is not None]


def rewrite_chart_references(workbook, rewrite: Rewrite) -> None:
    for worksheet in workbook.worksheets:
        for chart in worksheet._charts:
            for series in chart_series(chart):
                for holder in series_references(series):
                    holder.f = rewrite_reference_text(holder.f, worksheet.title, rewrite)


def shift_sheet_structure(worksheet, shift: Shift) -> None:
    guard_tables(worksheet, shift)
    move_cells(worksheet, shift)
    shift_merged_ranges(worksheet, shift)
    shift_filter(worksheet, shift)
    shift_tables(worksheet, shift)
    shift_validations(worksheet, shift)
    shift_conditional_formats(worksheet, shift)
    shift_dimensions(worksheet, shift)
    shift_chart_anchors(worksheet, shift)
    refresh_hyperlinks(worksheet)


def move_cells(worksheet, shift: Shift) -> None:
    operations = {
        (ROW_AXIS, True): worksheet.insert_rows,
        (ROW_AXIS, False): worksheet.delete_rows,
        (COLUMN_AXIS, True): worksheet.insert_cols,
        (COLUMN_AXIS, False): worksheet.delete_cols,
    }
    operations[(shift.axis, shift.is_insert)](shift.at, shift.removed)


def shifted_range_text(range_text: str, shift: Shift, sheet_title: str) -> str | None:
    shifted = shift_reference(range_text, sheet_title, shift)
    return None if shifted == REFERENCE_ERROR else shifted


def shift_merged_ranges(worksheet, shift: Shift) -> None:
    shifted = [shifted_range_text(str(merged), shift, shift.sheet) for merged in worksheet.merged_cells.ranges]
    worksheet.merged_cells = MultiCellRange()
    for text in shifted:
        if text is not None and ":" in text and text.split(":")[0] != text.split(":")[1]:
            worksheet.merged_cells.add(CellRange(text))


def shift_filter(worksheet, shift: Shift) -> None:
    if not worksheet.auto_filter.ref:
        return
    worksheet.auto_filter.ref = shifted_range_text(worksheet.auto_filter.ref, shift, shift.sheet)


def guard_tables(worksheet, shift: Shift) -> None:
    for table in worksheet.tables.values():
        if table_is_broken_by(table, shift):
            raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"table {table.name} on {worksheet.title} would gain or lose a column or lose its header row, which sheet apply does not edit; change rows and columns outside the table, or rebuild the table", worksheet.title))


def table_is_broken_by(table, shift: Shift) -> bool:
    before = CellRange(table.ref)
    after_text = shifted_range_text(table.ref, shift, shift.sheet)
    if after_text is None:
        return True
    after = CellRange(after_text)
    if shift.axis == COLUMN_AXIS:
        return after.max_col - after.min_col != before.max_col - before.min_col
    return not shift.is_insert and shift.at <= before.min_row <= shift.last_removed


def shift_tables(worksheet, shift: Shift) -> None:
    for table in worksheet.tables.values():
        table.ref = shifted_range_text(table.ref, shift, shift.sheet)
        if table.autoFilter is not None and table.autoFilter.ref:
            table.autoFilter.ref = table.ref


def shift_validations(worksheet, shift: Shift) -> None:
    kept = []
    for validation in worksheet.data_validations.dataValidation:
        shifted = [shifted_range_text(str(cell_range), shift, shift.sheet) for cell_range in validation.sqref.ranges]
        texts = [text for text in shifted if text is not None]
        if texts:
            validation.sqref = MultiCellRange(" ".join(texts))
            kept.append(validation)
    worksheet.data_validations.dataValidation = kept


def shift_conditional_formats(worksheet, shift: Shift) -> None:
    rebuilt = ConditionalFormattingList()
    for formatting in worksheet.conditional_formatting:
        shifted = [shifted_range_text(str(cell_range), shift, shift.sheet) for cell_range in formatting.sqref.ranges]
        texts = [text for text in shifted if text is not None]
        for rule in formatting.rules if texts else ():
            rebuilt.add(" ".join(texts), rule)
    worksheet.conditional_formatting = rebuilt


def shift_dimensions(worksheet, shift: Shift) -> None:
    if shift.axis == ROW_AXIS:
        shift_row_dimensions(worksheet, shift)
        return
    shift_column_dimensions(worksheet, shift)


def shift_row_dimensions(worksheet, shift: Shift) -> None:
    moved = {}
    for index, dimension in list(worksheet.row_dimensions.items()):
        new_index = shift_single(index, shift)
        if new_index is not None:
            dimension.index = new_index
            moved[new_index] = dimension
    worksheet.row_dimensions.clear()
    worksheet.row_dimensions.update(moved)


def shift_column_dimensions(worksheet, shift: Shift) -> None:
    moved = {}
    for dimension in list(worksheet.column_dimensions.values()):
        span = shift_span(dimension.min or 1, dimension.max or dimension.min or 1, shift)
        if span is not None:
            dimension.index = get_column_letter(span[0])
            dimension.min, dimension.max = span
            moved[dimension.index] = dimension
    worksheet.column_dimensions.clear()
    worksheet.column_dimensions.update(moved)


def isolate_column(worksheet, index: int) -> ColumnDimension:
    for dimension in list(worksheet.column_dimensions.values()):
        first, last = dimension.min or 1, dimension.max or dimension.min or 1
        if first <= index <= last and first != last:
            del worksheet.column_dimensions[dimension.index]
            for piece_first, piece_last in ((first, index - 1), (index, index), (index + 1, last)):
                if piece_first <= piece_last:
                    duplicate = column_dimension_copy(worksheet, dimension, piece_first, piece_last)
                    worksheet.column_dimensions[duplicate.index] = duplicate
    return worksheet.column_dimensions[get_column_letter(index)]


def column_dimension_copy(worksheet, dimension, first: int, last: int) -> ColumnDimension:
    duplicate = ColumnDimension(worksheet, index=get_column_letter(first), width=dimension.width, bestFit=dimension.bestFit, hidden=dimension.hidden, outlineLevel=dimension.outlineLevel, collapsed=dimension.collapsed, customWidth=dimension.customWidth, min=first, max=last)
    duplicate._style = copy(dimension._style)
    return duplicate


def shift_chart_anchors(worksheet, shift: Shift) -> None:
    attribute = "row" if shift.axis == ROW_AXIS else "col"
    for chart in worksheet._charts:
        for marker in anchor_markers(chart.anchor):
            moved = shift_single(getattr(marker, attribute) + 1, shift)
            setattr(marker, attribute, (moved if moved is not None else shift.at) - 1)


def anchor_markers(anchor) -> list:
    if isinstance(anchor, str):
        return []
    return [marker for marker in (getattr(anchor, "_from", None), getattr(anchor, "to", None)) if marker is not None]


def refresh_hyperlinks(worksheet) -> None:
    for cell in worksheet._cells.values():
        if cell.hyperlink is not None:
            cell.hyperlink.ref = cell.coordinate
