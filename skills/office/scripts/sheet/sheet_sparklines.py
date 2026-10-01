from __future__ import annotations

from lxml import etree
from openpyxl.utils import get_column_letter

from sheet.formula_references import quote_sheet_name
from core.office_operations import OPERATION_NOT_APPLICABLE, Change
from core.office_result import OfficeFailure
from sheet.sheet_formatting import require_colors
from sheet.workbook_access import parse_range, sheet_of
from sheet.workbook_fidelity import EditRecord, merge_extension, rewrite_extension_references
from sheet.workbook_package import Package, main_tag, worksheet_parts


SPARKLINE_URI = "{05C60535-1F16-4fd2-B633-F4F36F0B64E0}"
X14_NAMESPACE = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
EXCEL_NAMESPACE = "http://schemas.microsoft.com/office/excel/2006/main"
DEFAULT_COLOR = "2563EB"
NEGATIVE_COLOR = "EF4444"
HIGH_COLOR = "10B981"
LOW_COLOR = "EF4444"
AXIS_COLOR = "64748B"
TYPES = {"line": None, "column": "column", "win_loss": "stacked"}


def sparkline_cells(data: tuple[int, int, int, int], target: tuple[int, int, int, int], location: str) -> list[tuple[str, str]]:
    data_rows = data[2] - data[0] + 1
    target_cells = (target[2] - target[0] + 1) * (target[3] - target[1] + 1)
    if target[1] != target[3] and target[0] != target[2] or target_cells != data_rows:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.target: give one target cell per data row, {data_rows} in one column or row; the target has {target_cells}", f"{location}.target"))
    targets = [f"{get_column_letter(column)}{row}" for row in range(target[0], target[2] + 1) for column in range(target[1], target[3] + 1)]
    sources = [f"{get_column_letter(data[1])}{row}:{get_column_letter(data[3])}{row}" for row in range(data[0], data[2] + 1)]
    return list(zip(sources, targets))


def plan_add_sparklines(editing, operation: dict, location: str) -> Change:
    worksheet = sheet_of(editing.workbook, operation, location)
    data = parse_range(operation["range"], f"{location}.range")
    target = parse_range(operation["target"], f"{location}.target")
    cells = sparkline_cells(data, target, location)
    require_colors(operation, location)

    def change() -> str:
        recorded = len(editing.record.steps)
        title = worksheet.title
        editing.package_patches.append(lambda package: add_sparkline_group(package, worksheet.title, group_element(title, cells, operation), EditRecord(editing.record.steps[recorded:]), title))
        return f"added {len(cells)} sparklines to {worksheet.title}!{operation['target'].upper()}"
    return change


def color(parent, name: str, value: str) -> None:
    etree.SubElement(parent, f"{{{X14_NAMESPACE}}}{name}", rgb=f"FF{value.upper()}")


def group_element(title: str, cells: list[tuple[str, str]], operation: dict):
    group = etree.Element(f"{{{X14_NAMESPACE}}}sparklineGroup", nsmap={"x14": X14_NAMESPACE, "xm": EXCEL_NAMESPACE})
    group.set("displayEmptyCellsAs", "gap")
    if TYPES[operation.get("type", "line")]:
        group.set("type", TYPES[operation.get("type", "line")])
    if operation.get("markers"):
        group.set("markers", "1")
    if operation.get("highLow"):
        group.set("high", "1")
        group.set("low", "1")
    series = operation.get("color", DEFAULT_COLOR)
    for name, value in (("colorSeries", series), ("colorNegative", NEGATIVE_COLOR), ("colorAxis", AXIS_COLOR), ("colorMarkers", series), ("colorFirst", series), ("colorLast", series), ("colorHigh", HIGH_COLOR), ("colorLow", LOW_COLOR)):
        color(group, name, value)
    sparklines = etree.SubElement(group, f"{{{X14_NAMESPACE}}}sparklines")
    for source, target in cells:
        sparkline = etree.SubElement(sparklines, f"{{{X14_NAMESPACE}}}sparkline")
        etree.SubElement(sparkline, f"{{{EXCEL_NAMESPACE}}}f").text = f"{quote_sheet_name(title)}!{source}"
        etree.SubElement(sparkline, f"{{{EXCEL_NAMESPACE}}}sqref").text = target
    return group


def add_sparkline_group(package: Package, title: str, group, later_steps: EditRecord, original_title: str) -> None:
    part = next(part for part, name in worksheet_parts(package).items() if name == title)
    root = package.xml(part)
    extension = etree.Element(main_tag("ext"), uri=SPARKLINE_URI, nsmap={"x14": X14_NAMESPACE})
    groups = etree.SubElement(extension, f"{{{X14_NAMESPACE}}}sparklineGroups", nsmap={"xm": EXCEL_NAMESPACE})
    groups.append(group)
    rewrite_extension_references(extension, original_title, later_steps)
    extensions = root.find(main_tag("extLst"))
    if extensions is None:
        extensions = etree.SubElement(root, main_tag("extLst"))
    existing = next((item for item in extensions if item.get("uri") == SPARKLINE_URI), None)
    merge_extension(extensions, existing, extension)
    package.set_xml(part, root)
