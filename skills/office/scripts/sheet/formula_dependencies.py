from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

from formula_references import MAXIMUM_COLUMN, MAXIMUM_ROW, formula_references, is_bare_name, parse_end, reference_parts, unquote_sheet_name
from formula_tree import calls_in, parse_formula


VOLATILE_REFERENCE_FUNCTIONS = frozenset(("INDIRECT", "OFFSET"))
PARAMETER_PREFIX = "_XLPM."


@dataclass(frozen=True)
class Region:
    sheet: str | None
    min_row: int
    min_column: int
    max_row: int
    max_column: int

    def contains(self, sheet: str, row: int, column: int) -> bool:
        if self.sheet is not None and self.sheet != sheet:
            return False
        return self.min_row <= row <= self.max_row and self.min_column <= column <= self.max_column


EVERYWHERE = Region(None, 1, 1, MAXIMUM_ROW, MAXIMUM_COLUMN)


class DependencyReader:
    def __init__(self, workbook):
        self.workbook = workbook
        self.tables = {table.name.casefold(): (worksheet.title, table.ref) for worksheet in workbook.worksheets for table in worksheet.tables.values()}
        self.names = {name.casefold(): (defined.attr_text or "", None) for name, defined in workbook.defined_names.items()}
        for worksheet in workbook.worksheets:
            self.names.update({(worksheet.title.casefold(), name.casefold()): (defined.attr_text or "", worksheet.title) for name, defined in worksheet.defined_names.items()})

    def regions(self, formula: str, host_sheet: str, depth: int = 0) -> list[Region]:
        nodes = parse_formula(formula)
        if nodes is None or depth > 8:
            return [EVERYWHERE]
        if any(call.function in VOLATILE_REFERENCE_FUNCTIONS for call in calls_in(nodes)):
            return [EVERYWHERE]
        return [region for reference in formula_references(formula) for region in self.reference_regions(reference, host_sheet, depth)]

    def reference_regions(self, reference: str, host_sheet: str, depth: int) -> list[Region]:
        if reference.upper().startswith(PARAMETER_PREFIX):
            return []
        if "[" in reference and "!" not in reference.split("[", 1)[0]:
            return self.table_regions(reference.split("[", 1)[0], host_sheet)
        if is_bare_name(reference):
            return self.name_regions(reference, host_sheet, depth)
        region = cell_region(reference, host_sheet)
        return [region] if region is not None else []

    def table_regions(self, table_name: str, host_sheet: str) -> list[Region]:
        found = self.tables.get(table_name.casefold())
        if found is None:
            return [Region(host_sheet.casefold(), 1, 1, MAXIMUM_ROW, MAXIMUM_COLUMN)]
        return [cell_region(found[1], found[0])]

    def name_regions(self, name: str, host_sheet: str, depth: int) -> list[Region]:
        found = self.names.get((host_sheet.casefold(), name.casefold())) or self.names.get(name.casefold())
        if found is None:
            return []
        return self.regions("=" + found[0], found[1] or host_sheet, depth + 1)


def cell_region(reference: str, host_sheet: str) -> Region | None:
    parts = reference_parts(reference)
    if parts is None or len(parts) > 2:
        return None
    sheet = unquote_sheet_name(parts[0].prefix) if parts[0].prefix is not None else host_sheet
    ends = [parse_end(part.rest) for part in parts]
    if any(end is None for end in ends):
        return None
    first, last = ends[0], ends[-1]
    return Region(
        sheet.casefold(),
        first.row if first.row is not None else 1,
        first.column if first.column is not None else 1,
        last.row if last.row is not None else MAXIMUM_ROW,
        last.column if last.column is not None else MAXIMUM_COLUMN,
    )


def cell_position(coordinate: str) -> tuple[int, int]:
    letters, row = coordinate_from_string(coordinate)
    return row, column_index_from_string(letters)


def propagate(roots: set, formulas: dict, reader: DependencyReader) -> set:
    if not roots:
        return set(roots)
    unevaluated = set(roots)
    regions = {key: reader.regions(formula, key[0]) for key, formula in formulas.items() if key not in unevaluated}
    positions = {key: (key[0].casefold(), *cell_position(key[1])) for key in formulas}
    changed = True
    while changed:
        changed = False
        root_positions = [positions[key] for key in unevaluated if key in positions]
        for key, key_regions in list(regions.items()):
            if any(region.contains(*position) for region in key_regions for position in root_positions):
                unevaluated.add(key)
                del regions[key]
                changed = True
    return unevaluated
