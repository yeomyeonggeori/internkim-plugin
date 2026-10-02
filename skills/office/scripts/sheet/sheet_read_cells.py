from __future__ import annotations

import datetime

from openpyxl.styles.fonts import DEFAULT_FONT
from openpyxl.utils import get_column_letter

from core.office_result import INVALID_VALUE, OfficeFailure
from sheet.sheet_styling import header_row_index
from sheet.workbook_access import cell_rows, column_index, formula_text, json_value
from sheet.workbook_values import EXCEL_ERROR_CODES


GENERAL = "General"
STAT_DIGITS = 10


def selected_columns(text: str | None, bounds: tuple[int, int, int, int]) -> list[int]:
    width = bounds[3] - bounds[1] + 1
    if not text:
        return list(range(width))
    offsets = []
    for piece in text.split(","):
        first, _, last = piece.strip().partition(":")
        start, end = column_index(first, "--columns"), column_index(last or first, "--columns")
        if start > end:
            raise OfficeFailure(INVALID_VALUE.issue(f"--columns: {piece.strip()!r} runs from right to left, so it names no column", "--columns", f"write it left to right: {get_column_letter(end)}:{get_column_letter(start)}"))
        offsets.extend(range(start - bounds[1], end - bounds[1] + 1))
    outside = [offset for offset in offsets if not 0 <= offset < width]
    if outside:
        raise OfficeFailure(INVALID_VALUE.issue(f"--columns: {text!r} reaches outside the range's columns {get_column_letter(bounds[1])} to {get_column_letter(bounds[3])}", "--columns"))
    return offsets


def kind_of(value: object, is_formula: bool) -> str:
    if is_formula and value is None:
        return "uncomputed"
    if value is None or value == "":
        return "empty"
    if isinstance(value, bool):
        return "logical"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return "date"
    if isinstance(value, str) and value in EXCEL_ERROR_CODES:
        return "error"
    return "text"


def paired_cells(worksheet, values_worksheet, bounds: tuple[int, int, int, int], columns: list[int]):
    for formula_row, value_row in zip(cell_rows(worksheet, bounds), cell_rows(values_worksheet, bounds)):
        formula_row, value_row = list(formula_row), list(value_row)
        for column in columns:
            yield formula_row[column], value_row[column].value


def cells_where(worksheet, values_worksheet, bounds: tuple[int, int, int, int], columns: list[int], where: str) -> list[dict]:
    found = []
    for cell, value in paired_cells(worksheet, values_worksheet, bounds, columns):
        is_formula = cell.data_type == "f"
        if where == "formula" and not is_formula or where != "formula" and kind_of(value, is_formula) != where:
            continue
        entry = {"cell": cell.coordinate}
        if where != "empty":
            entry["value"] = json_value(value)
        if is_formula:
            entry["formula"] = formula_text(cell)
        found.append(entry)
    return found


def column_stats(worksheet, values_worksheet, bounds: tuple[int, int, int, int], columns: list[int]) -> tuple[int, list[dict]]:
    header = header_row_index(values_worksheet, bounds)
    header_bounds = (header, bounds[1], header, bounds[3])
    headers = [cell.value for cell in list(cell_rows(values_worksheet, header_bounds))[0]]
    body = (header + 1, bounds[1], bounds[2], bounds[3])
    collected = {column: {"types": {}, "numbers": [], "formulas": 0} for column in columns}
    if body[0] <= body[2]:
        for formula_row, value_row in zip(cell_rows(worksheet, body), cell_rows(values_worksheet, body)):
            formula_row, value_row = list(formula_row), list(value_row)
            for column in columns:
                tally(collected[column], formula_row[column], value_row[column].value)
    return header, [summary(bounds[1] + column, headers[column], collected[column]) for column in columns]


def tally(column: dict, cell, value: object) -> None:
    kind = kind_of(value, cell.data_type == "f")
    column["types"][kind] = column["types"].get(kind, 0) + 1
    if cell.data_type == "f":
        column["formulas"] += 1
    if kind == "number":
        column["numbers"].append(value)


def summary(column_number: int, header: object, column: dict) -> dict:
    result = {"column": get_column_letter(column_number), "header": json_value(header), "types": column["types"]}
    if column["formulas"]:
        result["formulas"] = column["formulas"]
    numbers = column["numbers"]
    if numbers:
        result.update({"count": len(numbers), "min": min(numbers), "max": max(numbers), "sum": round(sum(numbers), STAT_DIGITS), "mean": round(sum(numbers) / len(numbers), STAT_DIGITS)})
    return result


def color_text(color) -> str | None:
    if color is None:
        return None
    if color.type == "rgb" and isinstance(color.rgb, str):
        return None if color.rgb in ("00000000", "FF000000") else color.rgb[-6:]
    if color.type == "theme":
        return f"theme {color.theme}"
    return None


def font_text(font) -> str | None:
    parts = [word for word, present in (("bold", font.b), ("italic", font.i), ("underline", font.u)) if present]
    if font.sz and float(font.sz) != DEFAULT_FONT.sz:
        parts.append(f"{float(font.sz):g}pt")
    if font.name and font.name != DEFAULT_FONT.name:
        parts.append(font.name)
    color = color_text(font.color)
    if color:
        parts.append(color)
    return " ".join(parts) or None


def border_text(border) -> str | None:
    sides = {side: getattr(border, side).style for side in ("left", "right", "top", "bottom") if getattr(border, side) is not None and getattr(border, side).style}
    if not sides:
        return None
    styles = set(sides.values())
    if len(sides) == 4 and len(styles) == 1:
        return f"{styles.pop()} all"
    return " ".join(f"{side} {style}" for side, style in sides.items())


def alignment_text(alignment) -> str | None:
    parts = [value for value in (alignment.horizontal, alignment.vertical and f"vertical {alignment.vertical}") if value]
    if alignment.wrap_text:
        parts.append("wrap")
    if alignment.indent:
        parts.append(f"indent {int(alignment.indent)}")
    return " ".join(parts) or None


def cell_format(cell) -> dict:
    fill = color_text(cell.fill.fgColor) if cell.fill is not None and cell.fill.fill_type == "solid" else None
    properties = {
        "numberFormat": cell.number_format if cell.number_format != GENERAL else None,
        "font": font_text(cell.font),
        "fill": fill,
        "border": border_text(cell.border),
        "alignment": alignment_text(cell.alignment),
    }
    return {name: value for name, value in properties.items() if value}


def formats_in(worksheet, bounds: tuple[int, int, int, int], columns: list[int]) -> dict:
    cells = []
    for row in cell_rows(worksheet, bounds):
        row = list(row)
        for column in columns:
            properties = cell_format(row[column])
            if properties:
                cells.append({"cell": row[column].coordinate, **properties})
    letters = [get_column_letter(bounds[1] + column) for column in columns]
    widths = {letter: worksheet.column_dimensions[letter].width for letter in letters if letter in worksheet.column_dimensions and worksheet.column_dimensions[letter].width}
    heights = {str(row): worksheet.row_dimensions[row].height for row in range(bounds[0], bounds[2] + 1) if row in worksheet.row_dimensions and worksheet.row_dimensions[row].height}
    return {"cells": cells, "columnWidths": widths, "rowHeights": heights}
