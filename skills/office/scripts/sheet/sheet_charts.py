from __future__ import annotations

from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, RadarChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from office_operations import OPERATION_NOT_APPLICABLE, Change
from office_result import MISSING_FIELD, OfficeFailure
from sheet_definitions import CHART_COLUMN_LEFT_OUT
from text_values import conversion_operations, holds_number, holds_or_reads_as_number
from workbook_access import parse_cell, parse_range, sheet_of


DEFAULT_ANCHOR_GAP = 2
SERIES_PALETTE = ("2563EB", "F59E0B", "10B981", "EF4444", "8B5CF6", "14B8A6", "EC4899", "64748B")
LINE_WIDTH_EMU = 28575
DEFAULT_WIDTH = 16
DEFAULT_HEIGHT = 8
DEFAULT_LEGEND = "bottom"
SECONDARY_AXIS_ID = 200
BAR_GAP_WIDTH = 80
DOUGHNUT_HOLE = 55
LEGEND_POSITIONS = {"bottom": "b", "right": "r", "top": "t"}
ROUND_CHARTS = ("pie", "doughnut")
ROUND_PLOTS = ("pieChart", "doughnutChart", "pie3DChart", "ofPieChart")
LINE_PLOTS = ("lineChart", "line3DChart", "radarChart")
SCATTER_PLOT = "scatterChart"
GENERAL = "General"


def require_block(worksheet, bounds: tuple[int, int, int, int], location: str) -> None:
    min_row, min_column, max_row, max_column = bounds
    if max_row <= min_row or max_column <= min_column:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: a chart needs a header row, a category column and at least one row and one series column", f"{location}.range", block_suggestion(worksheet, bounds)))


def block_suggestion(worksheet, bounds: tuple[int, int, int, int]) -> str:
    min_row, min_column, max_row, max_column = bounds
    last_row = max_row + 1 if max_row == min_row else max_row
    first_column, last_column = min_column, max_column
    if max_column == min_column:
        series_columns = number_columns_right_of(worksheet, (min_row, min_column, last_row, max_column))
        if series_columns:
            last_column = series_columns[-1]
        elif min_column > 1:
            first_column = min_column - 1
    if (first_column, last_column, last_row) == (min_column, max_column, max_row):
        return "point range at a block whose first column holds the categories and whose other columns hold numbers"
    widened = f"{get_column_letter(first_column)}{min_row}:{get_column_letter(last_column)}{last_row}"
    return f"include the header row, the category column and the number columns beside it, such as \"range\": \"{widened}\""


def number_columns_right_of(worksheet, bounds: tuple[int, int, int, int]) -> list[int]:
    min_row, _, max_row, max_column = bounds
    found = []
    for column in range(max_column + 1, worksheet.max_column + 1):
        filled = [cell for cell in (worksheet.cell(row=row, column=column) for row in range(min_row + 1, max_row + 1)) if cell.value not in (None, "")]
        if not filled:
            break
        if any(holds_or_reads_as_number(cell) for cell in filled):
            found.append(column)
    return found


def number_columns(worksheet, bounds: tuple[int, int, int, int]) -> list[int]:
    min_row, min_column, max_row, max_column = bounds
    columns = worksheet.iter_cols(min_row=min_row + 1, max_row=max_row, min_col=min_column + 1, max_col=max_column)
    return [cells[0].column for cells in columns if any(holds_number(cell) for cell in cells)]


def column_label(worksheet, bounds: tuple[int, int, int, int], column: int) -> str:
    letter = get_column_letter(column)
    header = worksheet.cell(row=bounds[0], column=column).value
    named = f" ({header})" if header not in (None, "") else ""
    return f"{letter}{bounds[0] + 1}:{letter}{bounds[2]}{named}"


def drawn_columns(worksheet, bounds: tuple[int, int, int, int], operation: dict, location: str) -> list[int]:
    require_block(worksheet, bounds, location)
    columns = number_columns(worksheet, bounds)
    if not columns:
        first, last = get_column_letter(bounds[1] + 1), get_column_letter(bounds[3])
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(
            f"{location}.range: no series column {first}{bounds[0] + 1}:{last}{bounds[2]} of {worksheet.title} holds a number, so the chart would be empty",
            f"{location}.range",
            "write the numbers first, or point range at a block whose first column holds the categories and whose other columns hold numbers",
        ))
    if operation["type"] == "combo" and len(columns) < 2:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: a combo chart needs two number columns or more, bars first and the line last", f"{location}.range"))
    return columns


def column_labels(worksheet, bounds: tuple[int, int, int, int], columns: list[int]) -> str:
    return ", ".join(column_label(worksheet, bounds, column) for column in columns)


def left_out_issue(worksheet, bounds: tuple[int, int, int, int], columns: list[int], location: str, chart_index: int):
    left_out = [column for column in range(bounds[1] + 1, bounds[3] + 1) if column not in columns]
    if not left_out:
        return None
    drawing = f"it draws {column_labels(worksheet, bounds, columns)} over the categories in {column_label(worksheet, bounds, bounds[1])}"
    conversions = {column: operations for column in left_out if (operations := conversion_operations(worksheet, column, bounds[0] + 1, bounds[2]))}
    if not conversions:
        return CHART_COLUMN_LEFT_OUT.issue(f"{location}.range: {worksheet.title}!{column_labels(worksheet, bounds, left_out)} holds no number, so the chart leaves it out; {drawing}", f"{location}.range", text_column_suggestion(bounds, columns, left_out))
    texts = [column for column in left_out if column not in conversions]
    text_note = f", and {column_labels(worksheet, bounds, texts)} holds no number" if texts else ""
    rebuilt = {"op": "edit_chart", "sheet": worksheet.title, "chart": chart_index, "range": range_text(bounds)}
    return CHART_COLUMN_LEFT_OUT.issue(
        f"{location}.range: {worksheet.title}!{column_labels(worksheet, bounds, list(conversions))} holds numbers stored as text{text_note}, so the chart leaves them out; {drawing}",
        f"{location}.range",
        [*(operation for operations in conversions.values() for operation in operations), rebuilt],
    )


def text_column_suggestion(bounds: tuple[int, int, int, int], columns: list[int], left_out: list[int]) -> str | None:
    min_row, min_column, max_row, _ = bounds
    if min(left_out) > max(columns):
        trimmed = range_text((min_row, min_column, max_row, max(columns)))
        return f"leave the text column out of the chart with \"range\": \"{trimmed}\", or leave it as it is"
    if max(left_out) < min(columns):
        relabelled = range_text((min_row, max(left_out), max_row, bounds[3]))
        return f"to label the categories with that text column instead, start the range at it: \"range\": \"{relabelled}\"; or leave it as it is"
    return None


def range_text(bounds: tuple[int, int, int, int]) -> str:
    min_row, min_column, max_row, max_column = bounds
    return f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{max_row}"


def build_chart(worksheet, operation: dict, location: str, chart_index: int):
    bounds = parse_range(operation["range"], f"{location}.range")
    columns = drawn_columns(worksheet, bounds, operation, location)
    if operation["type"] == "scatter":
        chart = scatter_chart(worksheet, bounds, columns)
    elif operation["type"] == "combo":
        chart = combo_chart(worksheet, bounds, columns, operation)
    else:
        chart = category_chart(worksheet, bounds, columns, operation)
    chart.width = operation.get("width", DEFAULT_WIDTH)
    chart.height = operation.get("height", DEFAULT_HEIGHT)
    apply_labels(chart, {"legend": DEFAULT_LEGEND, **operation})
    paint_series(chart, operation.get("colors") or [])
    format_value_axes(chart, worksheet, bounds, columns)
    show_axes(chart)
    return chart, left_out_issue(worksheet, bounds, columns, location, chart_index)


def add_columns(chart, worksheet, bounds, columns: list[int]) -> None:
    for column in columns:
        chart.add_data(Reference(worksheet, min_col=column, max_col=column, min_row=bounds[0], max_row=bounds[2]), titles_from_data=True)


def categories(worksheet, bounds) -> Reference:
    return Reference(worksheet, min_col=bounds[1], min_row=bounds[0] + 1, max_row=bounds[2])


def new_chart(chart_type: str, operation: dict):
    if chart_type == "bar":
        chart = BarChart()
        chart.type = "bar" if operation.get("horizontal") else "col"
        chart.gapWidth = BAR_GAP_WIDTH
    elif chart_type == "line":
        chart = LineChart()
    elif chart_type == "area":
        chart = AreaChart()
    elif chart_type == "radar":
        chart = RadarChart()
        chart.type = "marker"
    elif chart_type == "doughnut":
        chart = DoughnutChart()
        chart.holeSize = DOUGHNUT_HOLE
    else:
        chart = PieChart()
    if operation.get("stacked") and chart_type in ("bar", "line", "area"):
        chart.grouping = "stacked"
        if chart_type == "bar":
            chart.overlap = 100
    return chart


def category_chart(worksheet, bounds, columns: list[int], operation: dict):
    chart_type = operation["type"]
    chart = new_chart(chart_type, operation)
    add_columns(chart, worksheet, bounds, columns[:1] if chart_type in ROUND_CHARTS else columns)
    chart.set_categories(categories(worksheet, bounds))
    return chart


def combo_chart(worksheet, bounds, columns: list[int], operation: dict):
    line_count = min(operation.get("lineSeries", 1), len(columns) - 1)
    bars = new_chart("bar", operation)
    add_columns(bars, worksheet, bounds, columns[:-line_count])
    bars.set_categories(categories(worksheet, bounds))
    line = LineChart()
    add_columns(line, worksheet, bounds, columns[-line_count:])
    line.set_categories(categories(worksheet, bounds))
    if operation.get("secondaryAxis", True):
        line.y_axis.axId = SECONDARY_AXIS_ID
        line.y_axis.crosses = "max"
        line.y_axis.majorGridlines = None
    bars += line
    return bars


def scatter_chart(worksheet, bounds, columns: list[int]):
    chart = ScatterChart()
    x_values = Reference(worksheet, min_col=bounds[1], min_row=bounds[0] + 1, max_row=bounds[2])
    for column in columns:
        values = Reference(worksheet, min_col=column, min_row=bounds[0], max_row=bounds[2])
        chart.series.append(Series(values, x_values, title_from_data=True))
    return chart


def solid(color: str) -> GraphicalProperties:
    properties = GraphicalProperties(solidFill=color)
    properties.line.solidFill = color
    return properties


def series_color(colors: list[str], index: int) -> str:
    if index < len(colors):
        return colors[index].lstrip("#").upper()
    return SERIES_PALETTE[index % len(SERIES_PALETTE)]


def style_line(series, color: str) -> None:
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = LINE_WIDTH_EMU
    series.smooth = False
    series.marker = Marker(symbol="circle", size=6)
    series.marker.graphicalProperties = solid(color)


def style_marker(series, color: str) -> None:
    series.marker = Marker(symbol="circle", size=7)
    series.marker.graphicalProperties = solid(color)
    series.graphicalProperties.line.noFill = True


def style_fill(series, color: str) -> None:
    series.graphicalProperties.solidFill = color
    series.graphicalProperties.line.solidFill = color


def point_count(series) -> int:
    reference = series.val.numRef.f if series.val is not None and series.val.numRef is not None else None
    if not reference:
        return 0
    min_column, min_row, max_column, max_row = range_boundaries(reference.rpartition("!")[2].replace("$", ""))
    return (max_row - min_row + 1) * (max_column - min_column + 1)


def paint_series(chart, colors: list[str]) -> None:
    index = 0
    for plot in chart._charts:
        for series in plot.series:
            paint_one(plot.tagname, series, colors, index)
            index += 1


def paint_one(plot_kind: str, series, colors: list[str], index: int) -> None:
    if plot_kind in ROUND_PLOTS:
        series.dPt = [point_with_color(point, series_color(colors, point)) for point in range(point_count(series))]
    elif plot_kind in LINE_PLOTS:
        style_line(series, series_color(colors, index))
    elif plot_kind == SCATTER_PLOT:
        style_marker(series, series_color(colors, index))
    else:
        style_fill(series, series_color(colors, index))


def point_with_color(index: int, color: str) -> DataPoint:
    point = DataPoint(idx=index)
    point.graphicalProperties.solidFill = color
    return point


def format_value_axes(chart, worksheet, bounds, columns: list[int]) -> None:
    for sub_chart in chart._charts:
        y_axis = getattr(sub_chart, "y_axis", None)
        if y_axis is None:
            continue
        column = columns[-1] if sub_chart is not chart else columns[0]
        number_format = worksheet.cell(row=bounds[0] + 1, column=column).number_format
        if number_format != GENERAL:
            y_axis.number_format = number_format
            y_axis.numFmt.sourceLinked = False


def apply_labels(chart, operation: dict) -> None:
    if operation.get("title") is not None:
        chart.title = operation["title"] or None
    if operation.get("xTitle") is not None and hasattr(chart, "x_axis"):
        chart.x_axis.title = operation["xTitle"] or None
    if operation.get("yTitle") is not None and hasattr(chart, "y_axis"):
        chart.y_axis.title = operation["yTitle"] or None
    if operation.get("legend") == "none":
        chart.legend = None
    elif operation.get("legend") is not None:
        chart.legend = chart.legend or Legend()
        chart.legend.position = LEGEND_POSITIONS[operation["legend"]]
    if operation.get("dataLabels") is not None:
        chart.dataLabels = value_labels() if operation["dataLabels"] else None


def value_labels() -> DataLabelList:
    return DataLabelList(showVal=True, showLegendKey=False, showCatName=False, showSerName=False, showPercent=False, showBubbleSize=False)


def show_axes(chart) -> None:
    for sub_chart in chart._charts:
        for axis_name in ("x_axis", "y_axis"):
            axis = getattr(sub_chart, axis_name, None)
            if axis is not None:
                axis.delete = False


def chart_anchor(operation: dict, location: str) -> str:
    if operation.get("anchor"):
        row, column = parse_cell(operation["anchor"], f"{location}.anchor")
        return f"{get_column_letter(column)}{row}"
    min_row, _, _, max_column = parse_range(operation["range"], f"{location}.range")
    return f"{get_column_letter(max_column + DEFAULT_ANCHOR_GAP)}{min_row}"


def plan_add_chart(editing, operation: dict, location: str) -> Change:
    worksheet = sheet_of(editing.workbook, operation, location)
    chart, left_out = build_chart(worksheet, operation, location, len(worksheet._charts))
    anchor = chart_anchor(operation, location)

    def change() -> str:
        worksheet.add_chart(chart, anchor)
        if left_out is not None:
            editing.warnings.append(left_out)
        return f"added a {operation['type']} chart of {worksheet.title}!{operation['range'].upper()} at {anchor}"
    return change


def existing_chart(worksheet, operation: dict, location: str):
    charts = worksheet._charts
    if operation["chart"] >= len(charts):
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.chart: {worksheet.title} has {len(charts)} charts, numbered from 0", f"{location}.chart"))
    return charts[operation["chart"]]


def anchor_cell(chart) -> str:
    if isinstance(chart.anchor, str):
        return chart.anchor
    marker = chart.anchor._from
    return f"{get_column_letter(marker.col + 1)}{marker.row + 1}"


def plan_edit_chart(editing, operation: dict, location: str) -> Change:
    worksheet = sheet_of(editing.workbook, operation, location)
    chart = existing_chart(worksheet, operation, location)
    if operation.get("type") and not operation.get("range"):
        raise OfficeFailure(MISSING_FIELD.issue(f"{location}.range: changing the chart type needs the data range", f"{location}.range"))
    replacement, left_out = build_chart(worksheet, {"type": "bar", **operation}, location, operation["chart"]) if operation.get("range") else (None, None)

    def change() -> str:
        index = operation["chart"]
        anchor = chart_anchor(operation, location) if operation.get("anchor") else anchor_cell(chart)
        if replacement is not None:
            if operation.get("title") is None and chart.title is not None:
                replacement.title = chart.title
            replacement.anchor = anchor
            worksheet._charts[index] = replacement
            if left_out is not None:
                editing.warnings.append(left_out)
            return f"rebuilt chart {index} of {worksheet.title}"
        apply_labels(chart, operation)
        if operation.get("colors"):
            paint_series(chart, operation["colors"])
        if operation.get("anchor") or "width" in operation or "height" in operation:
            chart.width = operation.get("width", chart.width)
            chart.height = operation.get("height", chart.height)
            chart.anchor = anchor
        return f"edited chart {index} of {worksheet.title}"
    return change


def plan_delete_chart(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    existing_chart(worksheet, operation, location)

    def change() -> str:
        del worksheet._charts[operation["chart"]]
        return f"deleted chart {operation['chart']} of {worksheet.title}"
    return change
