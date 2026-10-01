from __future__ import annotations

import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from office_operations import OPERATION_NOT_APPLICABLE, Change
from office_result import INVALID_VALUE, OfficeFailure
from workbook_access import cell_rows, column_index, parse_range, resolve_sheet, sheet_of
from workbook_structure import isolate_column


COLOR_PATTERN = re.compile(r"^[0-9A-Fa-f]{6}$")
COLOR_FIELDS = ("fill", "fontColor", "borderColor", "color")
FONT_FIELDS = ("bold", "italic", "underline", "fontSize", "fontName", "fontColor")
ALIGNMENT_FIELDS = ("alignment", "verticalAlignment", "indent", "wrapText")
DEFAULT_BORDER_COLOR = "94A3B8"
BORDER_EDGES = ("left", "right", "top", "bottom")


def require_colors(operation: dict, location: str) -> None:
    for name in COLOR_FIELDS:
        if isinstance(operation.get(name), str) and not COLOR_PATTERN.match(operation[name]):
            raise OfficeFailure(INVALID_VALUE.issue(f"{location}.{name}: {operation[name]!r} is not six hex digits such as DCEAF7", f"{location}.{name}"))


def plan_format_range(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    bounds = parse_range(operation["range"], f"{location}.range")
    require_colors(operation, location)

    def change() -> str:
        rows = [list(row) for row in cell_rows(worksheet, bounds)]
        for row in rows:
            for cell in row:
                apply_format(cell, operation)
        if "border" in operation:
            apply_border(rows, operation)
        return f"formatted {sum(len(row) for row in rows)} cells in {worksheet.title}!{operation['range'].upper()}"
    return change


def apply_format(cell, operation: dict) -> None:
    if "numberFormat" in operation:
        cell.number_format = operation["numberFormat"]
    if any(name in operation for name in FONT_FIELDS):
        cell.font = changed_font(cell.font, operation)
    if "fill" in operation:
        cell.fill = PatternFill("solid", fgColor=operation["fill"].upper())
    if any(name in operation for name in ALIGNMENT_FIELDS):
        cell.alignment = changed_alignment(cell.alignment, operation)


def changed_font(font, operation: dict) -> Font:
    return Font(
        name=operation.get("fontName", font.name),
        size=operation.get("fontSize", font.sz),
        bold=operation.get("bold", font.b),
        italic=operation.get("italic", font.i),
        underline=("single" if operation["underline"] else None) if "underline" in operation else font.u,
        strike=font.strike,
        color=operation["fontColor"].upper() if "fontColor" in operation else font.color,
        vertAlign=font.vertAlign,
        family=font.family,
        charset=font.charset,
        scheme=None if "fontName" in operation else font.scheme,
    )


def changed_alignment(alignment, operation: dict) -> Alignment:
    return Alignment(
        horizontal=operation.get("alignment", alignment.horizontal),
        vertical=operation.get("verticalAlignment", alignment.vertical),
        wrap_text=operation.get("wrapText", alignment.wrap_text),
        shrink_to_fit=alignment.shrink_to_fit,
        indent=operation.get("indent", alignment.indent),
        text_rotation=alignment.text_rotation,
    )


def apply_border(rows: list[list], operation: dict) -> None:
    style = None if operation["border"] == "none" else operation.get("borderStyle", "thin")
    side = Side(style=style, color=operation.get("borderColor", DEFAULT_BORDER_COLOR).upper() if style else None)
    last_row, last_column = len(rows) - 1, len(rows[0]) - 1
    for row_index, row in enumerate(rows):
        for column_position, cell in enumerate(row):
            edges = border_edges(operation["border"], row_index, column_position, last_row, last_column)
            sides = {edge: side if edge in edges else getattr(cell.border, edge) for edge in BORDER_EDGES}
            cell.border = Border(**sides)


def border_edges(border: str, row: int, column: int, last_row: int, last_column: int) -> set:
    if border in ("all", "none"):
        return set(BORDER_EDGES)
    if border == "top":
        return {"top"} if row == 0 else set()
    if border == "bottom":
        return {"bottom"} if row == last_row else set()
    edges = set()
    edges.update({"top"} if row == 0 else set())
    edges.update({"bottom"} if row == last_row else set())
    edges.update({"left"} if column == 0 else set())
    edges.update({"right"} if column == last_column else set())
    return edges


def plan_merge_cells(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    min_row, min_column, max_row, max_column = parse_range(operation["range"], f"{location}.range")
    reference = f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{max_row}"
    if min_row == max_row and min_column == max_column:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: merging needs at least two cells", f"{location}.range"))
    overlapping = [str(merged) for merged in worksheet.merged_cells.ranges if not (merged.max_row < min_row or merged.min_row > max_row or merged.max_col < min_column or merged.min_col > max_column)]
    if overlapping:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: {reference} overlaps merged {', '.join(overlapping)}; unmerge it first", f"{location}.range"))

    def change() -> str:
        worksheet.merge_cells(reference)
        return f"merged {worksheet.title}!{reference}"
    return change


def plan_unmerge_cells(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    min_row, min_column, max_row, max_column = parse_range(operation["range"], f"{location}.range")
    inside = [str(merged) for merged in worksheet.merged_cells.ranges if merged.min_row >= min_row and merged.max_row <= max_row and merged.min_col >= min_column and merged.max_col <= max_column]
    if not inside:
        merged = ", ".join(str(item) for item in worksheet.merged_cells.ranges) or "none"
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: no merged range lies inside {operation['range'].upper()}; the sheet has {merged}", f"{location}.range"))

    def change() -> str:
        for reference in inside:
            worksheet.unmerge_cells(reference)
        return f"unmerged {', '.join(inside)} on {worksheet.title}"
    return change


def plan_set_row_height(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    first, count = operation["row"], operation.get("count", 1)

    def change() -> str:
        for row in range(first, first + count):
            worksheet.row_dimensions[row].height = operation["height"]
        return f"set {count} rows from {first} of {worksheet.title} to height {operation['height']}"
    return change


def plan_hide_rows(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    first, count, hidden = operation["at"], operation.get("count", 1), operation.get("hidden", True)

    def change() -> str:
        for row in range(first, first + count):
            worksheet.row_dimensions[row].hidden = hidden
        return f"{'hid' if hidden else 'showed'} {count} rows from {first} of {worksheet.title}"
    return change


def plan_hide_columns(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    first, count, hidden = column_index(operation["at"], f"{location}.at"), operation.get("count", 1), operation.get("hidden", True)

    def change() -> str:
        for index in range(first, first + count):
            isolate_column(worksheet, index).hidden = hidden
        return f"{'hid' if hidden else 'showed'} {count} columns from {get_column_letter(first)} of {worksheet.title}"
    return change


def plan_hide_sheet(workbook, operation: dict, location: str) -> Change:
    worksheet = resolve_sheet(workbook, operation["sheet"], f"{location}.sheet")
    hidden = operation.get("hidden", True)
    visible_others = [other for other in workbook.worksheets if other is not worksheet and other.sheet_state == "visible"]
    if hidden and not visible_others:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.sheet: {worksheet.title} is the only visible sheet, and a workbook needs one", f"{location}.sheet"))

    def change() -> str:
        worksheet.sheet_state = "hidden" if hidden else "visible"
        if hidden and workbook.active is worksheet:
            workbook.active = workbook.worksheets.index(visible_others[0])
        return f"{'hid' if hidden else 'showed'} sheet {worksheet.title}"
    return change
