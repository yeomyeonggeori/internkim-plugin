from __future__ import annotations

import re

from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.pagebreak import Break, ColBreak, RowBreak
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.office_operations import OPERATION_NOT_APPLICABLE, Change
from core.office_result import INVALID_VALUE, OfficeFailure
from core.office_theme import HYPERLINK_COLOR
from core.page_sizes import PAPER_BY_NAME
from sheet.workbook_access import column_index, parse_cell, parse_range, sheet_of


TABLE_NAME = re.compile(r"^[A-Za-z_\u0080-￿][A-Za-z0-9_.\u0080-￿]*$")
DEFAULT_TABLE_STYLE = "TableStyleMedium2"
EXTERNAL_LINK = re.compile(r"^(https?://|mailto:|file:)", re.IGNORECASE)
COMMENT_WIDTH = 240
COMMENT_HEIGHT = 90
# ECMA-376 Part 1, 18.3.1.63 pageSetup: fitToWidth and fitToHeight default to 1, and 0 leaves that direction free
EXCEL_DEFAULT_FIT_PAGES = 1
MARGINS = {
    "normal": {"left": 0.7, "right": 0.7, "top": 0.75, "bottom": 0.75, "header": 0.3, "footer": 0.3},
    "narrow": {"left": 0.25, "right": 0.25, "top": 0.75, "bottom": 0.75, "header": 0.3, "footer": 0.3},
    "wide": {"left": 1.0, "right": 1.0, "top": 1.0, "bottom": 1.0, "header": 0.5, "footer": 0.5},
}
PAGE_NUMBER_FOOTER = "&P / &N"
HEADER_FOOTER_CODES = {"{page}": "&P", "{pages}": "&N", "{date}": "&D", "{sheet}": "&A", "{file}": "&F"}


def table_names(workbook) -> set:
    return {table.casefold() for worksheet in workbook.worksheets for table in worksheet.tables}


def range_text(bounds: tuple[int, int, int, int]) -> str:
    min_row, min_column, max_row, max_column = bounds
    return f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{max_row}"


def ranges_overlap(first: str, second: str) -> bool:
    one, two = CellRange(first), CellRange(second)
    return not (one.max_row < two.min_row or one.min_row > two.max_row or one.max_col < two.min_col or one.min_col > two.max_col)


def table_name(workbook, operation: dict, location: str) -> str:
    taken = table_names(workbook)
    name = operation.get("name") or next(f"Table{number}" for number in range(1, 10000) if f"table{number}" not in taken)
    if not TABLE_NAME.match(name) or re.match(r"^[A-Za-z]{1,3}[0-9]+$", name):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.name: {name!r} must start with a letter or underscore, hold no spaces, and not look like a cell address", f"{location}.name"))
    if name.casefold() in taken:
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.name: the workbook already has a table named {name!r}", f"{location}.name"))
    return name


def require_table_header(worksheet, bounds: tuple[int, int, int, int], location: str) -> None:
    min_row, min_column, max_row, max_column = bounds
    if max_row <= min_row:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: a table needs a header row and at least one data row", f"{location}.range"))
    headers = [worksheet.cell(row=min_row, column=column).value for column in range(min_column, max_column + 1)]
    names = [str(header).strip() for header in headers if header is not None and str(header).strip()]
    if len(names) != len(headers) or len({name.casefold() for name in names}) != len(names):
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: every header cell of a table needs its own name; row {min_row} holds {headers}", f"{location}.range"))


def plan_add_table(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    bounds = parse_range(operation["range"], f"{location}.range")
    reference = range_text(bounds)
    require_table_header(worksheet, bounds, location)
    clashing = [table.name for table in worksheet.tables.values() if ranges_overlap(table.ref, reference)]
    if clashing:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.range: {reference} overlaps table {clashing[0]}", f"{location}.range"))
    name = table_name(workbook, operation, location)

    def change() -> str:
        for column in range(bounds[1], bounds[3] + 1):
            header = worksheet.cell(row=bounds[0], column=column)
            header.value = str(header.value).strip()
        let_table_style_show(worksheet, bounds)
        if worksheet.auto_filter.ref and ranges_overlap(worksheet.auto_filter.ref, reference):
            worksheet.auto_filter.ref = None
        style = TableStyleInfo(name=operation.get("style", DEFAULT_TABLE_STYLE), showRowStripes=operation.get("bandedRows", True), showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
        worksheet.add_table(Table(displayName=name, ref=reference, tableStyleInfo=style))
        return f"added table {name} over {worksheet.title}!{reference}"
    return change


def let_table_style_show(worksheet, bounds: tuple[int, int, int, int]) -> None:
    for row in worksheet.iter_rows(min_row=bounds[0], max_row=bounds[2], min_col=bounds[1], max_col=bounds[3]):
        for cell in row:
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border()
            if cell.row == bounds[0]:
                cell.font = Font(name=cell.font.name, size=cell.font.sz)


def plan_set_hyperlink(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    row, column = parse_cell(operation["cell"], f"{location}.cell")
    url = operation["url"].strip()
    if not url.startswith("#") and not EXTERNAL_LINK.match(url):
        raise OfficeFailure(INVALID_VALUE.issue(f"{location}.url: {url!r} must start with https://, http://, mailto: or # for a place in the workbook", f"{location}.url"))

    def change() -> str:
        cell = worksheet.cell(row=row, column=column)
        if url.startswith("#"):
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=url[1:], tooltip=operation.get("tooltip"))
        else:
            cell.hyperlink = Hyperlink(ref=cell.coordinate, target=url, tooltip=operation.get("tooltip"))
        if operation.get("text") or cell.value is None:
            cell.value = operation.get("text") or url.lstrip("#")
        cell.font = Font(name=cell.font.name, size=cell.font.sz, bold=cell.font.b, color=HYPERLINK_COLOR, underline="single")
        return f"linked {worksheet.title}!{cell.coordinate} to {url}"
    return change


def plan_set_comment(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    row, column = parse_cell(operation["cell"], f"{location}.cell")

    def change() -> str:
        cell = worksheet.cell(row=row, column=column)
        if not operation["text"]:
            cell.comment = None
            return f"removed the note on {worksheet.title}!{cell.coordinate}"
        author = operation.get("author") or ""
        cell.comment = Comment(operation["text"], author, width=COMMENT_WIDTH, height=COMMENT_HEIGHT)
        return f"set the note on {worksheet.title}!{cell.coordinate}"
    return change


def plan_set_page_setup(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    for name in ("printTitleRows", "printArea"):
        if operation.get(name):
            validate_print_reference(operation[name], name, location)
    for letter in operation.get("pageBreakColumns", []):
        column_index(letter, f"{location}.pageBreakColumns")

    def change() -> str:
        apply_page_setup(worksheet, operation)
        return f"set the page setup of {worksheet.title}"
    return change


def validate_print_reference(text: str, name: str, location: str) -> None:
    if name == "printTitleRows":
        if not re.match(r"^\$?\d+:\$?\d+$", text.strip()):
            raise OfficeFailure(INVALID_VALUE.issue(f"{location}.{name}: {text!r} is not a row span such as 1:1", f"{location}.{name}"))
        return
    parse_range(text, f"{location}.{name}")


def apply_page_setup(worksheet, operation: dict) -> None:
    if "orientation" in operation:
        worksheet.page_setup.orientation = operation["orientation"]
    if "paperSize" in operation:
        worksheet.page_setup.paperSize = PAPER_BY_NAME[operation["paperSize"]].spreadsheet_code
    if "fitToWidth" in operation or "fitToHeight" in operation:
        fit_to_pages(worksheet, fitted_pages(worksheet, operation, "fitToWidth"), fitted_pages(worksheet, operation, "fitToHeight"))
    if "printGridlines" in operation:
        worksheet.print_options.gridLines = operation["printGridlines"]
    if "printTitleRows" in operation:
        worksheet.print_title_rows = operation["printTitleRows"].replace("$", "") or None
    if "printArea" in operation:
        worksheet.print_area = operation["printArea"].replace("$", "").upper() or None
    if "margins" in operation:
        for side, inches in MARGINS[operation["margins"]].items():
            setattr(worksheet.page_margins, side, inches)
    if "centerHorizontally" in operation:
        worksheet.print_options.horizontalCentered = operation["centerHorizontally"]
    if "pageNumbers" in operation:
        worksheet.oddFooter.center.text = PAGE_NUMBER_FOOTER if operation["pageNumbers"] else None
    if "scale" in operation:
        fit_to_pages(worksheet, 0, 0)
        worksheet.page_setup.scale = operation["scale"]
    if "header" in operation:
        worksheet.oddHeader.center.text = header_footer_text(operation["header"])
    if "footer" in operation:
        worksheet.oddFooter.center.text = header_footer_text(operation["footer"])
    if "pageBreakRows" in operation:
        worksheet.row_breaks = RowBreak(brk=[Break(id=row) for row in operation["pageBreakRows"]])
    if "pageBreakColumns" in operation:
        worksheet.col_breaks = ColBreak(brk=[Break(id=column_index(letter, "pageBreakColumns")) for letter in operation["pageBreakColumns"]])


def header_footer_text(text: str) -> str | None:
    if not text:
        return None
    escaped = text.replace("&", "&&")
    for placeholder, code in HEADER_FOOTER_CODES.items():
        escaped = escaped.replace(placeholder, code)
    return escaped


def fitted_pages(worksheet, operation: dict, name: str) -> int:
    if name in operation:
        return int(operation[name])
    if not is_fitted_to_pages(worksheet):
        return 0
    current = getattr(worksheet.page_setup, name)
    return EXCEL_DEFAULT_FIT_PAGES if current is None else int(current)


def is_fitted_to_pages(worksheet) -> bool:
    properties = worksheet.sheet_properties.pageSetUpPr
    return bool(properties is not None and properties.fitToPage)


def fit_to_pages(worksheet, width_pages: int, height_pages: int) -> None:
    if worksheet.sheet_properties.pageSetUpPr is None:
        worksheet.sheet_properties.pageSetUpPr = PageSetupProperties()
    enabled = width_pages > 0 or height_pages > 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = enabled
    worksheet.page_setup.fitToWidth = width_pages if enabled else None
    worksheet.page_setup.fitToHeight = height_pages if enabled else None


