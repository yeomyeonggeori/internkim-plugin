from __future__ import annotations

from openpyxl.utils.cell import range_boundaries

from core.office_result import Issue
from sheet.sheet_definitions import AUTO_FILTER_MISSING, BLANK_HEADER_CELLS, HEADER_NOT_FROZEN
from sheet.operations.styling import MINIMUM_DATA_ROWS, MINIMUM_TABLE_COLUMNS, header_row_index, is_filled, non_blank_count


FORMULA_CELL_LIMIT = 50


def table_findings(workbook) -> tuple[list[Issue], dict]:
    sheet_summaries = [summarize_sheet(worksheet) for worksheet in workbook.worksheets]
    issues = [issue for summary in sheet_summaries for issue in sheet_issues(summary)]
    formula_cells = [
        f"{worksheet.title}!{cell.coordinate}"
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if is_formula(cell.value)
    ]
    details = {
        "sheetCount": len(workbook.worksheets),
        "sheets": sheet_summaries,
        "formulaCells": formula_cells[:FORMULA_CELL_LIMIT],
        "formulaCellCount": len(formula_cells),
    }
    return issues, details


def is_inside(cell, bounds: tuple[int, int, int, int]) -> bool:
    min_column, min_row, max_column, max_row = bounds
    return min_column <= cell.column <= max_column and min_row <= cell.row <= max_row


def holds_only_pivots(worksheet) -> bool:
    pivots = [range_boundaries(pivot.location.ref) for pivot in getattr(worksheet, "_pivots", [])]
    filled = [cell for cell in worksheet._cells.values() if is_filled(cell.value)]
    return bool(pivots) and all(any(is_inside(cell, bounds) for bounds in pivots) for cell in filled)


def summarize_sheet(worksheet) -> dict:
    header_row = header_row_index(worksheet)
    is_table_sheet = worksheet.max_row >= header_row and not holds_only_pivots(worksheet)
    hidden = hidden_columns(worksheet)
    header_values = filled_span([cell.value for cell in worksheet[header_row] if cell.column not in hidden]) if is_table_sheet else []
    row_counts = [non_blank_count(cell.value for cell in row) for row in worksheet.iter_rows(min_row=header_row + 1)]
    data_rows = sum(1 for count in row_counts if count > 0)
    is_one_block = not any(count == 0 for count in trimmed(row_counts))
    return {
        "title": worksheet.title,
        "rows": worksheet.max_row,
        "columns": worksheet.max_column,
        "freezePanes": str(worksheet.freeze_panes) if worksheet.freeze_panes else None,
        "autoFilter": bool(worksheet.auto_filter.ref),
        "dataRows": data_rows,
        "isDataTable": is_one_block and non_blank_count(header_values) >= MINIMUM_TABLE_COLUMNS and data_rows >= MINIMUM_DATA_ROWS,
        "blankHeaderCount": sum(1 for value in header_values if value is None or str(value).strip() == ""),
    }


def hidden_columns(worksheet) -> set[int]:
    return {column for dimension in worksheet.column_dimensions.values() if dimension.hidden for column in range(dimension.min or 1, (dimension.max or dimension.min or 1) + 1)}


def filled_span(values: list) -> list:
    filled = [index for index, value in enumerate(values) if is_filled(value)]
    return values[filled[0]:filled[-1] + 1] if filled else []


def trimmed(counts: list[int]) -> list[int]:
    filled = [index for index, count in enumerate(counts) if count > 0]
    return counts[filled[0]:filled[-1] + 1] if filled else []


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def sheet_issues(summary: dict) -> list[Issue]:
    title = summary["title"]
    issues = []
    if summary["isDataTable"] and not summary["freezePanes"]:
        issues.append(HEADER_NOT_FROZEN.issue(f"{title}: header row is not frozen", title))
    if summary["isDataTable"] and not summary["autoFilter"]:
        issues.append(AUTO_FILTER_MISSING.issue(f"{title}: auto filter is missing", title))
    if summary["blankHeaderCount"] > 0:
        issues.append(BLANK_HEADER_CELLS.issue(f"{title}: {summary['blankHeaderCount']} blank header cells", title))
    return issues

