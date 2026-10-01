from __future__ import annotations

import datetime

from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule, Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.datavalidation import DataValidation

from excel_functions import stored_formula
from office_operations import OPERATION_NOT_APPLICABLE, Change
from office_result import MISSING_FIELD, OfficeFailure
from sheet_formatting import style_color
from workbook_access import parse_range, sheet_of
from written_cells import require_writable


HIGHLIGHT_FILL = "FFC7CE"
HIGHLIGHT_FONT = "9C0006"
DATA_BAR_COLOR = "638EC6"
COMPARISON_OPERATORS = {
    "greater_than": "greaterThan",
    "less_than": "lessThan",
    "between": "between",
    "not_between": "notBetween",
    "equal": "equal",
    "not_equal": "notEqual",
    "greater_or_equal": "greaterThanOrEqual",
    "less_or_equal": "lessThanOrEqual",
}
COLOR_SCALES = {
    "red_yellow_green": ("F8696B", "FFEB84", "63BE7B"),
    "green_yellow_red": ("63BE7B", "FFEB84", "F8696B"),
    "white_green": ("FFFFFF", "63BE7B"),
    "white_red": ("FFFFFF", "F8696B"),
    "white_blue": ("FFFFFF", "5A8AC6"),
}
ICON_SETS = {
    "3_traffic_lights": "3TrafficLights1",
    "3_arrows": "3Arrows",
    "3_flags": "3Flags",
    "3_symbols": "3Symbols",
    "4_arrows": "4Arrows",
    "4_rating": "4Rating",
    "5_arrows": "5Arrows",
    "5_rating": "5Rating",
}
# the formulas Excel itself writes for these rules
TEXT_RULES = {
    "contains_text": ("containsText", "containsText", "NOT(ISERROR(SEARCH({text},{cell})))"),
    "not_contains_text": ("notContainsText", "notContains", "ISERROR(SEARCH({text},{cell}))"),
    "begins_with": ("beginsWith", "beginsWith", "LEFT({cell},LEN({text}))={text}"),
    "ends_with": ("endsWith", "endsWith", "RIGHT({cell},LEN({text}))={text}"),
}
BLANK_RULES = {
    "blank": ("containsBlanks", "LEN(TRIM({cell}))=0"),
    "not_blank": ("notContainsBlanks", "LEN(TRIM({cell}))>0"),
}
MIDDLE_PERCENTILE = 50
VALIDATION_TYPES = {"list": "list", "whole": "whole", "decimal": "decimal", "date": "date", "text_length": "textLength", "custom": "custom"}
LIST_LENGTH_LIMIT = 255


def require(operation: dict, name: str, location: str, reason: str) -> None:
    if operation.get(name) is None:
        raise OfficeFailure(MISSING_FIELD.issue(f"{location}.{name}: {reason}", f"{location}.{name}"))


def operand(value: object) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return repr(value)
    text = str(value)
    if text.startswith("="):
        return stored_formula(text)[1:]
    return '"' + text.replace('"', '""') + '"'


def date_operand(value: object) -> str:
    if isinstance(value, str) and len(value) == 10 and value[4] == "-":
        day = datetime.date.fromisoformat(value)
        return f"DATE({day.year},{day.month},{day.day})"
    return operand(value)


def highlight_style(operation: dict) -> DifferentialStyle:
    return DifferentialStyle(
        font=Font(color=style_color(operation.get("fontColor", HIGHLIGHT_FONT)), bold=operation.get("bold")),
        fill=PatternFill(bgColor=style_color(operation.get("fill", HIGHLIGHT_FILL)), fill_type="solid"),
    )


def top_left(range_text: str) -> str:
    return range_text.split(":")[0].replace("$", "").upper()


def conditional_rule(operation: dict, location: str):
    rule = operation["rule"]
    if rule in COMPARISON_OPERATORS:
        return comparison_rule(operation, location)
    if rule in TEXT_RULES:
        return text_rule(operation, location)
    if rule in BLANK_RULES:
        kind, formula = BLANK_RULES[rule]
        return Rule(type=kind, dxf=highlight_style(operation), formula=[formula.format(cell=top_left(operation["range"]))])
    if rule in ("duplicate", "unique"):
        return Rule(type="duplicateValues" if rule == "duplicate" else "uniqueValues", dxf=highlight_style(operation))
    if rule in ("top", "bottom"):
        return Rule(type="top10", rank=operation.get("rank", 10), percent=operation.get("percent"), bottom=rule == "bottom" or None, dxf=highlight_style(operation))
    if rule in ("above_average", "below_average"):
        return Rule(type="aboveAverage", aboveAverage=rule == "above_average", dxf=highlight_style(operation))
    if rule == "formula":
        require(operation, "formula", location, "a formula rule needs the formula")
        return FormulaRule(formula=[rule_formula(operation["formula"], f"{location}.formula")], font=highlight_style(operation).font, fill=highlight_style(operation).fill)
    return visual_rule(operation)


def text_rule(operation: dict, location: str) -> Rule:
    require(operation, "value", location, f"{operation['rule']} needs the text to look for")
    kind, operator, formula = TEXT_RULES[operation["rule"]]
    text = str(operation["value"])
    written = formula.format(text=operand(text), cell=top_left(operation["range"]))
    return Rule(type=kind, operator=operator, text=text, dxf=highlight_style(operation), formula=[written])


def ensure_equals(formula: str) -> str:
    return formula if formula.startswith("=") else "=" + formula


def rule_formula(text: str, location: str) -> str:
    formula = ensure_equals(text)
    require_writable(formula, location)
    return stored_formula(formula)[1:]


def comparison_rule(operation: dict, location: str):
    require(operation, "value", location, f"{operation['rule']} needs the value to compare with")
    values = [operand(operation["value"])]
    if operation["rule"] in ("between", "not_between"):
        require(operation, "value2", location, "between needs value2, the upper bound")
        values.append(operand(operation["value2"]))
    style = highlight_style(operation)
    return CellIsRule(operator=COMPARISON_OPERATORS[operation["rule"]], formula=values, font=style.font, fill=style.fill)


def visual_rule(operation: dict):
    if operation["rule"] == "color_scale":
        return color_scale_rule(operation)
    if operation["rule"] == "data_bar":
        return DataBarRule(start_type="min", end_type="max", color=style_color(operation.get("fill", DATA_BAR_COLOR)))
    icons = operation.get("icons", "3_traffic_lights")
    count = int(icons[0])
    return IconSetRule(ICON_SETS[icons], "percent", [round(100 * step / count) for step in range(count)])


def color_scale_rule(operation: dict):
    low, middle, high = scale_colors(operation)
    start = threshold(operation, "minValue", "min", None)
    end = threshold(operation, "maxValue", "max", None)
    if middle is None:
        return ColorScaleRule(start_type=start[0], start_value=start[1], start_color=low, end_type=end[0], end_value=end[1], end_color=high)
    mid = threshold(operation, "midValue", "percentile", MIDDLE_PERCENTILE)
    return ColorScaleRule(start_type=start[0], start_value=start[1], start_color=low, mid_type=mid[0], mid_value=mid[1], mid_color=middle, end_type=end[0], end_value=end[1], end_color=high)


def scale_colors(operation: dict) -> tuple:
    preset = COLOR_SCALES[operation.get("scale", "red_yellow_green")]
    if not any(name in operation for name in ("minColor", "midColor", "maxColor")):
        middle = style_color(preset[1]) if len(preset) == 3 else None
        return style_color(preset[0]), middle, style_color(preset[-1])
    middle = style_color(operation["midColor"]) if "midColor" in operation else None
    return style_color(operation.get("minColor", preset[0])), middle, style_color(operation.get("maxColor", preset[-1]))


def threshold(operation: dict, name: str, default_type: str, default_value: float | None) -> tuple:
    return ("num", operation[name]) if name in operation else (default_type, default_value)


def plan_add_conditional_format(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    parse_range(operation["range"], f"{location}.range")
    rule = conditional_rule(operation, location)

    def change() -> str:
        worksheet.conditional_formatting.add(operation["range"].replace("$", "").upper(), rule)
        return f"added a {operation['rule']} conditional format to {worksheet.title}!{operation['range'].upper()}"
    return change


def overlaps(first: CellRange, second: CellRange) -> bool:
    return not (first.max_row < second.min_row or first.min_row > second.max_row or first.max_col < second.min_col or first.min_col > second.max_col)


def plan_clear_conditional_formats(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    target = CellRange(operation["range"].replace("$", "").upper()) if operation.get("range") else None
    if target is not None:
        parse_range(operation["range"], f"{location}.range")

    def change() -> str:
        kept = ConditionalFormattingList()
        removed = 0
        for formatting in worksheet.conditional_formatting:
            if target is None or any(overlaps(cell_range, target) for cell_range in formatting.sqref.ranges):
                removed += len(formatting.rules)
                continue
            for rule in formatting.rules:
                kept.add(str(formatting.sqref), rule)
        worksheet.conditional_formatting = kept
        return f"removed {removed} conditional formats from {worksheet.title}"
    return change


def validation_from(operation: dict, location: str) -> DataValidation:
    kind = operation["type"]
    validation = DataValidation(
        type=VALIDATION_TYPES[kind],
        allow_blank=operation.get("allowBlank", True),
        showErrorMessage=True,
        showInputMessage=bool(operation.get("prompt")),
        prompt=operation.get("prompt") or None,
        error=operation.get("error") or None,
    )
    if kind == "list":
        validation.formula1 = list_source(operation, location)
    elif kind == "custom":
        require(operation, "formula", location, "a custom validation needs the formula")
        validation.formula1 = rule_formula(operation["formula"], f"{location}.formula")
    else:
        bound_validation(validation, operation, location)
    return validation


def list_source(operation: dict, location: str) -> str:
    if operation.get("source"):
        return rule_formula(operation["source"], f"{location}.source")
    require(operation, "values", location, "a list needs values, or source naming a range of choices")
    if any("," in value for value in operation["values"]):
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.values: a choice cannot hold a comma; put the choices in cells and pass source", f"{location}.values"))
    joined = ",".join(operation["values"])
    if len(joined) > LIST_LENGTH_LIMIT:
        raise OfficeFailure(OPERATION_NOT_APPLICABLE.issue(f"{location}.values: the choices take {len(joined)} characters and Excel allows {LIST_LENGTH_LIMIT}; put them in cells and pass source", f"{location}.values"))
    return f'"{joined}"'


def bound_validation(validation: DataValidation, operation: dict, location: str) -> None:
    bounds = [operation.get(name) for name in ("minimum", "maximum")]
    if all(bound is None for bound in bounds):
        raise OfficeFailure(MISSING_FIELD.issue(f"{location}.minimum: {operation['type']} needs a minimum, a maximum or both", f"{location}.minimum"))
    default = "between" if None not in bounds else ("greater_or_equal" if bounds[0] is not None else "less_or_equal")
    validation.operator = COMPARISON_OPERATORS[operation.get("operator", default)]
    present = [bound for bound in bounds if bound is not None]
    convert = date_operand if operation["type"] == "date" else operand
    validation.formula1 = convert(present[0])
    if validation.operator in ("between", "notBetween"):
        if len(present) < 2:
            raise OfficeFailure(MISSING_FIELD.issue(f"{location}.maximum: between needs both minimum and maximum", f"{location}.maximum"))
        validation.formula2 = convert(present[1])


def plan_add_data_validation(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    parse_range(operation["range"], f"{location}.range")
    validation = validation_from(operation, location)

    def change() -> str:
        validation.add(operation["range"].replace("$", "").upper())
        worksheet.add_data_validation(validation)
        return f"added a {operation['type']} validation to {worksheet.title}!{operation['range'].upper()}"
    return change


def plan_clear_data_validations(workbook, operation: dict, location: str) -> Change:
    worksheet = sheet_of(workbook, operation, location)
    parse_range(operation["range"], f"{location}.range")
    target = CellRange(operation["range"].replace("$", "").upper())

    def change() -> str:
        kept = [validation for validation in worksheet.data_validations.dataValidation if not any(overlaps(cell_range, target) for cell_range in validation.sqref.ranges)]
        removed = len(worksheet.data_validations.dataValidation) - len(kept)
        worksheet.data_validations.dataValidation = kept
        return f"removed {removed} data validations from {worksheet.title}"
    return change

