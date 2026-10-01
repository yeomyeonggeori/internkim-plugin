#!/usr/bin/env python3
from __future__ import annotations

import os
import tempfile

from lxml import etree
from openpyxl.utils import get_column_letter

from formula_cache import cache_formula_values
from office_operations import apply_batch, save_atomically
from office_inputs import office_file
from office_result import OfficeArgumentParser, Result, read_json_file, run_command
from office_schema import require_valid
from formula_references import ROW_AXIS, parse_end, rebuild_reference, reference_parts, same_sheet, unquote_sheet_name
from sheet_operations import SHEET_OPERATIONS, load_editing, save_editing
from workbook_structure import rewrite_chart_references, rewrite_defined_names, rewrite_formulas
from template_merge import MERGE_VALUES, MISSING, PLACEHOLDER, MergeReport, fill_text_nodes, repeated_list_name, whole_placeholder, write_package
from workbook_package import MAIN_NAMESPACE, read_package, worksheet_parts


SHARED_STRINGS_PART = "xl/sharedStrings.xml"
HEADER_FOOTER_TAGS = ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter")
PRESERVE_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def main() -> Result:
    arguments = parse_arguments()
    values = read_json_file(arguments.values_path)
    require_valid(MERGE_VALUES, values, "values")
    output_path = os.path.expanduser(arguments.output_path)
    report = MergeReport(values)
    with tempfile.TemporaryDirectory() as directory:
        template_path, list_names = expanded_template(os.path.expanduser(arguments.template_path), values, os.path.join(directory, "template.xlsx"))
        for list_name in list_names:
            report.mark_used(list_name)
        parts = merged_parts(template_path, report)
        report.require_complete()
        issues = save_atomically(lambda path: write_filled_workbook(template_path, parts, path), output_path)
    return Result(summary=f"filled {report.filled} placeholders into {output_path}", output_path=output_path, issues=(*issues, *report.unused_issues()), details={"placeholders": sorted(report.placeholder_names)})


def expanded_template(template_path: str, values: dict, expanded_path: str) -> tuple[str, set[str]]:
    editing = load_editing(template_path)
    list_rows = [(worksheet, row_number, list_name) for worksheet in editing.workbook.worksheets for row_number, list_name in rows_naming_lists(worksheet, values)]
    if not list_rows:
        return template_path, set()
    for worksheet, row_number, list_name in sorted(list_rows, key=lambda entry: (entry[0].title, -entry[1])):
        expand_list_row(editing, worksheet, row_number, list_name, len(values[list_name]))
    save_editing(editing, expanded_path)
    return expanded_path, {list_name for _, _, list_name in list_rows}


def expand_list_row(editing, worksheet, row_number: int, list_name: str, item_count: int) -> None:
    if item_count == 0:
        clear_row_placeholders(worksheet, row_number)
        return
    if item_count > 1:
        apply_batch(SHEET_OPERATIONS, editing, row_copy_operations(worksheet, row_number, item_count))
        extend_ranges_over_copies(editing.workbook, worksheet.title, row_number, item_count - 1)
    for index in range(item_count):
        index_row_placeholders(worksheet, row_number + index, list_name, index)


def clear_row_placeholders(worksheet, row_number: int) -> None:
    for cell in worksheet[row_number]:
        if isinstance(cell.value, str) and PLACEHOLDER.search(cell.value):
            cell.value = None


def rows_naming_lists(worksheet, values: dict) -> list[tuple[int, str]]:
    rows = []
    for row in worksheet.iter_rows():
        text = " ".join(cell.value for cell in row if isinstance(cell.value, str))
        list_name = repeated_list_name(text, values)
        if list_name is not None:
            rows.append((row[0].row, list_name))
    return rows


def row_copy_operations(worksheet, row_number: int, item_count: int) -> list[dict]:
    last_column = get_column_letter(worksheet.max_column)
    copies = [{"op": "copy_range", "sheet": worksheet.title, "range": f"A{row_number}:{last_column}{row_number}", "to": f"A{row_number + offset}"} for offset in range(1, item_count)]
    return [{"op": "insert_rows", "sheet": worksheet.title, "at": row_number + 1, "count": item_count - 1}, *copies]


def extend_ranges_over_copies(workbook, sheet_title: str, row_number: int, added_rows: int) -> None:
    if added_rows <= 0:
        return
    rewrite = lambda reference, host_sheet: extended_reference(reference, host_sheet, sheet_title, row_number, added_rows)
    rewrite_formulas(workbook, rewrite)
    rewrite_defined_names(workbook, rewrite)
    rewrite_chart_references(workbook, rewrite)


def extended_reference(reference: str, host_sheet: str, sheet_title: str, row_number: int, added_rows: int) -> str:
    parts = reference_parts(reference)
    if parts is None or len(parts) != 2:
        return reference
    target_sheet = unquote_sheet_name(parts[0].prefix) if parts[0].prefix is not None else host_sheet
    ends = [parse_end(part.rest) for part in parts]
    if not same_sheet(target_sheet, sheet_title) or any(end is None or end.row is None for end in ends):
        return reference
    if not ends[0].row <= row_number == ends[1].row:
        return reference
    return rebuild_reference(parts, [ends[0], ends[1].moved(ROW_AXIS, row_number + added_rows)])


def index_row_placeholders(worksheet, row_number: int, list_name: str, index: int) -> None:
    for cell in worksheet[row_number]:
        if isinstance(cell.value, str):
            cell.value = PLACEHOLDER.sub(lambda match: indexed_placeholder(match, list_name, index), cell.value)


def indexed_placeholder(match, list_name: str, index: int) -> str:
    name, _, rest = match.group(1).partition(".")
    if name != list_name or not rest or rest.split(".", 1)[0].isdigit():
        return match.group(0)
    return f"{{{{ {list_name}.{index}.{rest} }}}}"


def merged_parts(template_path: str, report: MergeReport) -> dict[str, bytes]:
    package = read_package(template_path)
    sheet_names = worksheet_parts(package)
    shared = package.xml(SHARED_STRINGS_PART) if SHARED_STRINGS_PART in package.entries else None
    sheets = {part: package.xml(part) for part in sheet_names}
    shared_items = list(shared.iterchildren(qualified("si"))) if shared is not None else []
    filled_items: set[int] = set()
    for part, worksheet in sheets.items():
        fill_worksheet(worksheet, sheet_names[part], shared_items, filled_items, report)
    rewritten = {part: serialize(worksheet) for part, worksheet in sheets.items()}
    if filled_items:
        rewritten[SHARED_STRINGS_PART] = serialize(shared)
    return rewritten


def fill_worksheet(worksheet, sheet_name: str, shared_items: list, filled_items: set[int], report: MergeReport) -> None:
    for cell in worksheet.iter(qualified("c")):
        if cell.find(qualified("f")) is not None:
            continue
        fill_cell(cell, f"{sheet_name}!{cell.get('r')}", shared_items, filled_items, report)
    for tag in HEADER_FOOTER_TAGS:
        for element in worksheet.iter(qualified(tag)):
            fill_text_nodes([element], report, f"{sheet_name} {tag}")


def fill_cell(cell, location: str, shared_items: list, filled_items: set[int], report: MergeReport) -> None:
    item = string_item(cell, shared_items)
    if item is None:
        return
    nodes = text_nodes(item)
    if write_typed_value(cell, nodes, location, report):
        return
    index = int(cell.findtext(qualified("v"))) if cell.get("t") == "s" else None
    if index in filled_items:
        return
    if fill_text_nodes(nodes, report, location):
        preserve_spaces(nodes)
        if index is not None:
            filled_items.add(index)


def string_item(cell, shared_items: list):
    if cell.get("t") == "inlineStr":
        return cell.find(qualified("is"))
    if cell.get("t") != "s":
        return None
    index = int(cell.findtext(qualified("v")))
    return shared_items[index] if index < len(shared_items) else None


def text_nodes(item) -> list:
    return [*item.findall(qualified("t")), *item.findall(f"{qualified('r')}/{qualified('t')}")]


def write_typed_value(cell, nodes: list, location: str, report: MergeReport) -> bool:
    path = whole_placeholder("".join(node.text or "" for node in nodes))
    if path is None:
        return False
    value = report.resolve(path, location)
    if value is MISSING or isinstance(value, str) or value is None:
        if value is not MISSING:
            write_inline_text(cell, value or "")
        return True
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)
    if isinstance(value, bool):
        cell.set("t", "b")
    etree.SubElement(cell, qualified("v")).text = str(int(value)) if isinstance(value, bool) else repr(value)
    return True


def write_inline_text(cell, text: str) -> None:
    for child in list(cell):
        cell.remove(child)
    cell.set("t", "inlineStr")
    item = etree.SubElement(cell, qualified("is"))
    node = etree.SubElement(item, qualified("t"))
    node.text = text
    preserve_spaces([node])


def preserve_spaces(nodes: list) -> None:
    for node in nodes:
        node.set(PRESERVE_SPACE, "preserve")


def write_filled_workbook(template_path: str, parts: dict[str, bytes], output_path: str):
    write_package(template_path, parts, output_path)
    return cache_formula_values(output_path)


def qualified(tag: str) -> str:
    return f"{{{MAIN_NAMESPACE}}}{tag}"


def serialize(element) -> bytes:
    return etree.tostring(element, xml_declaration=True, encoding="UTF-8", standalone=True)


def parse_arguments():
    parser = OfficeArgumentParser()
    parser.add_argument("template_path", type=office_file("xlsx"), help="the .xlsx template")
    parser.add_argument("values_path", help="JSON object mapping each placeholder name to its value")
    parser.add_argument("output_path", help="the filled .xlsx to write")
    return parser.parse_args()


if __name__ == "__main__":
    raise SystemExit(run_command(main))
