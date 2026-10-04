from __future__ import annotations

from math import ceil

from openpyxl.utils import get_column_letter

from sheet.workbook.display_width import character_width, display_width


MINIMUM_WIDTH = 10
MAXIMUM_WIDTH = 40
CELL_PADDING = 2
LINE_POINTS = 15
HEADER_LINES = 2


def header_floor(text: str) -> int:
    longest_word = max((display_width(word) for word in str(text).split()), default=0)
    return max(longest_word, ceil(display_width(text) / HEADER_LINES))


def wrapped_line_count(text: str, capacity: int) -> int:
    lines = 0
    for paragraph in str(text).split("\n"):
        lines += wrapped_paragraph_lines(paragraph, max(capacity, 1))
    return max(lines, 1)


def wrapped_paragraph_lines(paragraph: str, capacity: int) -> int:
    lines, used = 1, 0
    for word in paragraph.split(" "):
        for piece in pieces_within(word, capacity):
            size = display_width(piece)
            separator = 1 if used else 0
            if used and used + separator + size > capacity:
                lines, used, separator = lines + 1, 0, 0
            used += separator + size
    return lines


def pieces_within(word: str, capacity: int) -> list[str]:
    pieces, current, size = [], "", 0
    for character in word:
        width = character_width(character)
        if current and size + width > capacity:
            pieces.append(current)
            current, size = "", 0
        current, size = current + character, size + width
    return pieces + [current]


def column_width(sheet, column: int) -> float:
    return sheet.column_dimensions[get_column_letter(column)].width or MINIMUM_WIDTH


def fit_columns(sheet, overflowing_rows: set, header_rows: frozenset = frozenset()) -> None:
    widths: dict[int, float] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None or isinstance(cell.value, str) and cell.value.startswith("=") or cell.row in overflowing_rows:
                continue
            content = header_floor(cell.value) if cell.row in header_rows else display_width(str(cell.value))
            widths[cell.column] = max(widths.get(cell.column, 0), content + CELL_PADDING)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width, MINIMUM_WIDTH), MAXIMUM_WIDTH)


def fit_row_height(sheet, row: int, spans: dict[int, int]) -> None:
    needed = 1
    for cell in sheet[row]:
        if cell.value is None or isinstance(cell.value, str) and cell.value.startswith("="):
            continue
        capacity = sum(column_width(sheet, cell.column + offset) for offset in range(spans.get(cell.column, 1))) - CELL_PADDING
        needed = max(needed, wrapped_line_count(cell.value, int(capacity)))
    if needed > 1:
        sheet.row_dimensions[row].height = LINE_POINTS * needed + 3
