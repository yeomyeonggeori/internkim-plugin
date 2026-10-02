#!/usr/bin/env python3
from __future__ import annotations

from openpyxl.utils import get_column_letter

from core.office_arguments import route_arguments
from core.office_result import Result, run_command
from sheet.checks.chart_references import describe_charts
from sheet.sheet_definitions import READ_ROW_LIMIT
from sheet.read_cells import cells_where, column_stats, formats_in, selected_columns
from sheet.workbook.access import cell_rows, formula_text, json_value, open_workbook, parse_range, resolve_sheet
from sheet.workbook.package import read_package, worksheet_parts


SPARKLINE_TAG = "{http://schemas.microsoft.com/office/spreadsheetml/2009/9/main}sparkline"


def main() -> Result:
    arguments = parse_arguments()
    formulas_workbook = open_workbook(arguments.file)
    values_workbook = open_workbook(arguments.file, data_only=True)
    worksheet = resolve_sheet(formulas_workbook, arguments.sheet, "--sheet")
    bounds = requested_bounds(worksheet, arguments.range)
    sparklines = sparkline_counts(arguments.file)
    details = {
        "sheets": [describe_sheet(sheet, sparklines.get(sheet.title, 0)) for sheet in formulas_workbook.worksheets],
        "definedNames": describe_defined_names(formulas_workbook),
        "range": describe_range(worksheet, values_workbook[worksheet.title], bounds, arguments),
    }
    return Result(summary=f"read {details['range']['range']} of {worksheet.title} from {arguments.file}", output_path=arguments.file, details=details)


def requested_bounds(worksheet, range_text: str | None) -> tuple[int, int, int, int]:
    if range_text:
        return parse_range(range_text, "--range")
    return 1, 1, max(worksheet.max_row, 1), max(worksheet.max_column, 1)


def limited_bounds(bounds: tuple[int, int, int, int], limit: int) -> tuple[tuple[int, int, int, int], bool]:
    min_row, min_column, max_row, max_column = bounds
    last_row = min(max_row, min_row + limit - 1)
    return (min_row, min_column, last_row, max_column), last_row < max_row


def sparkline_counts(path: str) -> dict:
    package = read_package(path)
    return {title: sum(1 for _ in package.xml(part).iter(SPARKLINE_TAG)) for part, title in worksheet_parts(package).items()}


def describe_sheet(worksheet, sparklines: int) -> dict:
    description = {
        "name": worksheet.title,
        "visible": worksheet.sheet_state == "visible",
        "dimensions": worksheet.dimensions,
        "rows": worksheet.max_row,
        "columns": worksheet.max_column,
        "frozenPanes": worksheet.freeze_panes,
        "autoFilter": worksheet.auto_filter.ref or None,
        "tables": [{"name": table.name, "range": table.ref} for table in worksheet.tables.values()],
        "charts": describe_charts(worksheet),
        "mergedCells": [str(merged) for merged in worksheet.merged_cells.ranges],
    }
    features = feature_counts(worksheet, sparklines)
    if features:
        description["features"] = features
    return description


def feature_counts(worksheet, sparklines: int) -> dict:
    counts = {
        "conditionalFormats": sum(len(formatting.rules) for formatting in worksheet.conditional_formatting),
        "dataValidations": len(worksheet.data_validations.dataValidation),
        "hyperlinks": sum(1 for cell in worksheet._cells.values() if cell.hyperlink is not None),
        "comments": sum(1 for cell in worksheet._cells.values() if cell.comment is not None),
        "pivotTables": len(getattr(worksheet, "_pivots", [])),
        "images": len(getattr(worksheet, "_images", [])),
        "sparklines": sparklines,
    }
    return {name: count for name, count in counts.items() if count}


def describe_defined_names(workbook) -> list[dict]:
    names = [{"name": name, "value": defined.attr_text, "scope": None} for name, defined in workbook.defined_names.items()]
    for worksheet in workbook.worksheets:
        names.extend({"name": name, "value": defined.attr_text, "scope": worksheet.title} for name, defined in worksheet.defined_names.items())
    return names


def range_label(bounds: tuple[int, int, int, int]) -> str:
    min_row, min_column, max_row, max_column = bounds
    return f"{get_column_letter(min_column)}{min_row}:{get_column_letter(max_column)}{max_row}"


def describe_range(worksheet, values_worksheet, bounds: tuple[int, int, int, int], arguments) -> dict:
    columns = selected_columns(arguments.columns, bounds)
    description = {"sheet": worksheet.title, "range": range_label(bounds)}
    if arguments.stats:
        description["headerRow"], description["stats"] = column_stats(worksheet, values_worksheet, bounds, columns)
        return description
    shown_bounds, truncated = limited_bounds(bounds, arguments.limit)
    description.update({"range": range_label(shown_bounds), "truncated": truncated})
    if arguments.where:
        description["cells"] = cells_where(worksheet, values_worksheet, shown_bounds, columns, arguments.where)
    elif arguments.formats:
        description.update(formats_in(worksheet, shown_bounds, columns))
    else:
        description.update(values_and_formulas(worksheet, values_worksheet, shown_bounds, columns))
    return description


def values_and_formulas(worksheet, values_worksheet, bounds: tuple[int, int, int, int], columns: list[int]) -> dict:
    values = [[json_value(row[column].value) for column in columns] for row in grid(values_worksheet, bounds)]
    formulas = [[formula_text(row[column]) for column in columns] for row in grid(worksheet, bounds)]
    shown = {"values": values}
    if any(formula is not None for row in formulas for formula in row):
        shown["formulas"] = formulas
    if len(columns) != bounds[3] - bounds[1] + 1:
        shown["columns"] = [get_column_letter(bounds[1] + column) for column in columns]
    return shown


def grid(worksheet, bounds: tuple[int, int, int, int]) -> list:
    return [list(row) for row in cell_rows(worksheet, bounds)]


def parse_arguments():
    return route_arguments("read", "xlsx", limit=READ_ROW_LIMIT)


if __name__ == "__main__":
    raise SystemExit(run_command(main))
