from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from sheet.column_fit import fit_columns, fit_row_height
from sheet.compile_workbook import INPUT_FILL, Compiler, Placed, compile_declaration, header_text, invalid, number_format
from sheet.operations.charts import DEFAULT_WIDTH as DEFAULT_CHART_WIDTH
from sheet.operations.styling import BORDER_COLOR, HEADER_FILL_COLOR


CHART_OPERATIONS = {"line": {"type": "line"}, "column": {"type": "bar"}, "bar": {"type": "bar", "horizontal": True}, "pie": {"type": "pie"}}
CHART_COLUMNS_GAP = 20
CENTIMETERS_PER_WIDTH_UNIT = 0.19
DEFAULT_COLUMN_WIDTH = 8.43
CHART_HEIGHT_ROWS = 17


def thin_border() -> Border:
    side = Side(style="thin", color=BORDER_COLOR)
    return Border(left=side, right=side, top=side, bottom=side)


def style_header(cell) -> None:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor=HEADER_FILL_COLOR)
    cell.border = thin_border()
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_tables(workbook: Workbook, compiler: Compiler) -> None:
    for table in compiler.tables:
        sheet = workbook.create_sheet(table.name)
        for index, column in enumerate(table.columns, start=1):
            sheet.cell(1, index, header_text(column))
            style_header(sheet.cell(1, index))
        for row_index, row in enumerate(table.rows, start=2):
            for index, (column, value) in enumerate(zip(table.columns, row), start=1):
                cell = sheet.cell(row_index, index, value)
                cell.border = thin_border()
                if column.role == "measure":
                    cell.number_format = number_format(column.type, column.unit, isinstance(value, float))
                if value is None:
                    cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
                    compiler.blanks.append({"field": f"{table.name}!{cell.coordinate}", "label": blank_label(table, row, column)})
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(table.columns))}{len(table.rows) + 1}"
        fit_columns(sheet, set())


def blank_label(table, row: list, column) -> str:
    dimensions = [str(value) for other, value in zip(table.columns, row) if other.role == "dimension" and value is not None]
    return " ".join([*dimensions, column.name])


def write_view(workbook: Workbook, placed: Placed, title: str) -> None:
    sheet = workbook[placed.sheet] if placed.sheet in workbook.sheetnames else workbook.create_sheet(placed.sheet)
    sheet.cell(placed.title_row, 1, title).font = Font(bold=True, size=12)
    for (row, column), value in placed.cells.items():
        if (row, column) == (placed.title_row, 1):
            continue
        cell = sheet.cell(row, column, value)
        if row in (placed.header_row, placed.group_row):
            style_header(cell)
            continue
        if row in placed.note_rows:
            cell.font = Font(italic=True, size=9)
            continue
        cell.border = thin_border()
        if (row, column) in placed.formats:
            cell.number_format = placed.formats[(row, column)]
        if row == placed.total_row or column == placed.total_column:
            cell.font = Font(bold=True)
    for group in placed.groups:
        merge_group_header(sheet, placed, group)


def merge_group_header(sheet, placed: Placed, group: dict) -> None:
    for column in range(group["first"], group["last"] + 1):
        style_header(sheet.cell(placed.group_row, column))
    sheet.merge_cells(start_row=placed.group_row, start_column=group["first"], end_row=placed.group_row, end_column=group["last"])


def compiled_ranges(compiler: Compiler) -> list[dict]:
    ranges = [{"sheet": table.name, "range": f"A1:{get_column_letter(len(table.columns))}{len(table.rows) + 1}"} for table in compiler.tables]
    for placed in compiler.views:
        last_row = max(row for row, _ in placed.cells)
        ranges.append({"sheet": placed.sheet, "range": f"A{placed.title_row}:{get_column_letter(placed.width)}{last_row}"})
    return ranges + [{"sheet": sheet, "range": block} for sheet, block in compiler.chart_blocks]


def chart_operations(workbook: Workbook, compiler: Compiler, titles: list[str]) -> list[dict]:
    operations = []
    placed_by_title = dict(zip(titles, compiler.views))
    sheet_widths: dict[str, int] = {}
    for placed in compiler.views:
        sheet_widths[placed.sheet] = max(sheet_widths.get(placed.sheet, 0), placed.width)
    helper_columns = {sheet: width + CHART_COLUMNS_GAP for sheet, width in sheet_widths.items()}
    for index, chart in enumerate(compiler.declaration.get("charts") or []):
        placed = placed_by_title.get(chart["view"])
        location = f"declaration.charts[{index}].view"
        if placed is None:
            raise invalid(f"{location}: no view titled {chart['view']!r}; the views are {', '.join(titles)}", location)
        first_column = helper_columns[placed.sheet]
        helper_columns[placed.sheet] += len(placed.base_columns) + 2
        data_range = chart_data_block(workbook[placed.sheet], placed, first_column)
        compiler.chart_blocks.append((placed.sheet, data_range))
        anchor_row = compiler.next_row.get(placed.sheet, 1)
        start_chart_on_new_page(workbook[placed.sheet], anchor_row)
        compiler.next_row[placed.sheet] = anchor_row + CHART_HEIGHT_ROWS + 1
        operation = {"op": "add_chart", "sheet": placed.sheet, "range": data_range, "anchor": f"A{anchor_row}", **CHART_OPERATIONS[chart["type"]]}
        if chart.get("title"):
            operation["title"] = chart["title"]
        operations.append(operation)
        chart_end = max(sheet_widths[placed.sheet], columns_spanned(workbook[placed.sheet], DEFAULT_CHART_WIDTH))
        fit_print_width(workbook[placed.sheet], chart_end, compiler.next_row[placed.sheet])
    return operations


def start_chart_on_new_page(sheet, anchor_row: int) -> None:
    sheet.row_breaks.append(Break(id=anchor_row - 1))


def fit_print_width(sheet, last_column: int, last_row: int) -> None:
    sheet.print_area = f"A1:{get_column_letter(last_column)}{last_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def columns_spanned(sheet, centimeters: float) -> int:
    column, covered = 0, 0.0
    while covered < centimeters:
        column += 1
        covered += (sheet.column_dimensions[get_column_letter(column)].width or DEFAULT_COLUMN_WIDTH) * CENTIMETERS_PER_WIDTH_UNIT
    return column


def chart_data_block(sheet, placed: Placed, first_column: int) -> str:
    width = 1 + len(placed.base_columns)
    for offset in range(width):
        sheet.column_dimensions[get_column_letter(first_column + offset)].hidden = True
    for source_row in [placed.header_row, *placed.body_rows]:
        sheet.cell(source_row, first_column, joined_labels(placed, source_row))
        for offset in range(1, width):
            sheet.cell(source_row, first_column + offset, charted_value(placed, source_row, placed.label_count + offset))
    return f"{get_column_letter(first_column)}{placed.header_row}:{get_column_letter(first_column + width - 1)}{placed.body_rows[-1]}"


def charted_value(placed: Placed, row: int, column: int) -> str:
    source = f"{get_column_letter(column)}{row}"
    if row == placed.header_row:
        return f"={source}"
    if (row, column) in placed.partial:
        return "=NA()"
    return f"=IF(ISNUMBER({source}),{source},NA())"


def joined_labels(placed: Placed, row: int) -> str:
    labels = [f"{get_column_letter(column)}{row}" for column in range(1, placed.label_count + 1)]
    if row == placed.header_row:
        return f"={labels[-1]}"
    return "=" + '&" "&'.join(labels)


def show_hidden_chart_data(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        for chart in sheet._charts:
            chart.visible_cells_only = False
            chart.display_blanks = "gap"


def shown_views(path: str, compiler: Compiler, titles: list[str]) -> list[dict]:
    from openpyxl import load_workbook
    from core.number_format import displayed

    workbook = load_workbook(path, data_only=True)
    return [shown_view(workbook[placed.sheet], placed, title, displayed) for placed, title in zip(compiler.views, titles)]


def shown_view(sheet, placed: Placed, title: str, displayed) -> dict:
    columns = sorted({column for row, column in placed.cells if row == placed.header_row})
    rows = [*([placed.group_row] if placed.group_row else []), placed.header_row, *placed.body_rows, *([placed.total_row] if placed.total_row else [])]
    shown = [[displayed(sheet.cell(row, column).value, sheet.cell(row, column).number_format).text for column in columns] for row in rows]
    view = {"title": title, "sheet": placed.sheet, "rows": shown}
    if placed.note_rows:
        view["notes"] = [placed.cells[(row, 1)] for row in placed.note_rows]
    return view


def fit_header_heights(sheet, placed: Placed) -> None:
    fit_row_height(sheet, placed.header_row, {})
    if placed.group_row:
        fit_row_height(sheet, placed.group_row, {group["first"]: group["last"] - group["first"] + 1 for group in placed.groups})


def declared_workbook(declaration: dict, attachments: tuple = ()) -> tuple[Workbook, list[dict], Compiler]:
    compiler = compile_declaration(declaration, attachments)
    workbook = Workbook()
    workbook.remove(workbook.active)
    if declaration.get("title"):
        workbook.properties.title = declaration["title"]
    write_tables(workbook, compiler)
    titles = [view["title"] for view in declaration.get("views") or []]
    for placed, title in zip(compiler.views, titles):
        write_view(workbook, placed, title)
    for sheet_name in dict.fromkeys(placed.sheet for placed in compiler.views):
        placed_here = [placed for placed in compiler.views if placed.sheet == sheet_name]
        overflowing_rows = {row for placed in placed_here for row in (placed.title_row, placed.group_row, *placed.note_rows) if row}
        header_rows = frozenset(placed.header_row for placed in placed_here)
        fit_columns(workbook[sheet_name], overflowing_rows, header_rows)
        for placed in placed_here:
            fit_header_heights(workbook[sheet_name], placed)
    return workbook, chart_operations(workbook, compiler, titles), compiler
