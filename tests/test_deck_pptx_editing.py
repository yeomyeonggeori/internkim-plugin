import base64
import difflib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

import lxml.html
from PIL import Image
from pptx import Presentation

from deck_fixture import OFFICE_ENTRY, SCRIPTS_PATH
from doc_fixture import write_json
from render_fixture import assert_pages_drawn, can_render
from pptx_edit_fixture import CUSTOM_PART_NAME, UNKNOWN_EXTENSION_URI, build_korean_deck, part_contents, sample_photo


sys.path.insert(0, str(SCRIPTS_PATH))
sys.path.insert(0, str(SCRIPTS_PATH / "deck"))

from pptx_text_measure import font_face  # noqa: E402


def run_office(arguments, directory):
    completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), *arguments], capture_output=True, text=True, cwd=directory)
    return json.loads(completed.stdout)


def codes(envelope):
    return [issue["code"] for issue in envelope["issues"]]


class KoreanDeckFixture(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._shared = tempfile.TemporaryDirectory()
        cls.deck_bytes_path = Path(cls._shared.name) / "deck.pptx"
        build_korean_deck(cls.deck_bytes_path)

    @classmethod
    def tearDownClass(cls):
        cls._shared.cleanup()

    def setUp(self):
        self._temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self._temporary_directory.cleanup)
        self.directory = Path(self._temporary_directory.name)
        shutil.copy(self.deck_bytes_path, self.directory / "deck.pptx")
        (self.directory / "photo.png").write_bytes(sample_photo(400, 400))

    def apply(self, operations, *options):
        write_json(self.directory / "ops.json", operations)
        return run_office(["deck", "apply", "deck.pptx", "ops.json", *options], self.directory)

    def read(self, *options, name="deck.pptx"):
        return run_office(["deck", "read", name, *options], self.directory)["details"]

    def slide(self, number, *options):
        return self.read("--slides", str(number), *options)["slides"][0]

    def presentation(self, name="deck.pptx"):
        return Presentation(str(self.directory / name))


class ReadTest(KoreanDeckFixture):
    def test_each_shape_has_index_id_kind_box_and_effective_style(self):
        details = self.read()
        self.assertEqual(details["slideSize"]["w"], 12192000)
        title, card, label, body, picture = self.slide(2)["shapes"]
        self.assertEqual((title["kind"], title["placeholder"]), ("text", "title"))
        self.assertEqual((card["kind"], picture["kind"]), ("shape", "picture"))
        self.assertEqual(label["box"], {"x": 762000, "y": 1981200, "w": 3048000, "h": 457200})
        self.assertEqual(label["percent"]["x"], 6.2)
        self.assertEqual(label["style"]["size"], 24.0)
        self.assertTrue(label["style"]["bold"])
        self.assertEqual(label["style"]["color"], "#FFFFFF")
        self.assertEqual(body["text"], "신규 고객 42곳 확보\n재구매율 68%로 상승")

    def test_detail_names_where_each_inherited_value_comes_from(self):
        title = self.slide(2, "--detail")["shapes"][0]
        run = title["paragraphs"][0]["runs"][0]
        self.assertEqual(run["size"], {"value": 44.0, "from": "master"})
        self.assertEqual(run["font"]["from"], "theme")
        label_run = self.slide(2, "--detail")["shapes"][2]["paragraphs"][0]["runs"][0]
        self.assertEqual(label_run["size"], {"value": 24.0, "from": "run"})

    def test_groups_tables_charts_and_notes_are_described(self):
        table_slide = self.slide(4, "--detail")
        group = table_slide["shapes"][2]
        self.assertEqual([child["index"] for child in group["shapes"]], ["2.0", "2.1"])
        self.assertEqual(group["shapes"][1]["text"], "담당 박예시")
        self.assertEqual(table_slide["shapes"][1]["rows"][1], ["플랫폼", "50억", "58억"])
        self.assertEqual(table_slide["animated"], [str(group["id"])])
        chart = self.slide(3)["shapes"][1]["chart"]
        self.assertEqual(chart["categories"], ["1분기", "2분기", "3분기"])
        self.assertEqual(chart["series"][0], {"name": "매출", "values": [96.0, 110.0, 128.0]})
        self.assertEqual(self.slide(1)["notes"], "인사 후 3분기 요약으로 시작합니다.")


class TextOperationTest(KoreanDeckFixture):
    def test_set_text_on_one_paragraph_keeps_its_run_formatting(self):
        envelope = self.apply([{"op": "set_text", "slide": 2, "shape": 3, "paragraph": 1, "text": "재구매율 71%로 상승"}])
        self.assertEqual(envelope["status"], "ok", envelope["issues"])
        body = self.presentation().slides[1].shapes[3].text_frame
        self.assertEqual([paragraph.text for paragraph in body.paragraphs], ["신규 고객 42곳 확보", "재구매율 71%로 상승"])
        self.assertEqual(body.paragraphs[1].runs[0].font.size.pt, 20)

    def test_style_paragraph_and_frame_operations_write_their_properties(self):
        self.apply([
            {"op": "set_text_style", "slide": 2, "shape": 3, "paragraph": 0, "bold": True, "color": "1F4E79", "font": "Nanum Gothic", "size": 22},
            {"op": "set_paragraph", "slide": 2, "shape": 3, "bullet": "bullet", "align": "right", "spaceAfter": 6},
            {"op": "set_text_frame", "slide": 2, "shape": 3, "autofit": "shrink", "anchor": "middle"},
        ])
        body = self.presentation().slides[1].shapes[3].text_frame
        first_run = body.paragraphs[0].runs[0]
        self.assertTrue(first_run.font.bold)
        self.assertEqual((str(first_run.font.color.rgb), first_run.font.size.pt, first_run.font.name), ("1F4E79", 22, "Nanum Gothic"))
        self.assertIsNone(body.paragraphs[1].runs[0].font.bold)
        self.assertEqual(body.paragraphs[1]._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar").get("char"), "•")
        self.assertIn("normAutofit", body._bodyPr.xml)

    def test_find_replace_reaches_tables_groups_and_one_shape(self):
        self.apply([{"op": "find_replace", "find": "억", "replace": " 억원", "slide": 4}, {"op": "find_replace", "find": "박예시", "replace": "최견본"}])
        rows = self.slide(4)["shapes"][1]["rows"]
        self.assertEqual(rows[1], ["플랫폼", "50 억원", "58 억원"])
        self.assertEqual(self.slide(4)["shapes"][2]["shapes"][1]["text"], "담당 최견본")
        scoped = self.apply([{"op": "find_replace", "find": "분기", "replace": "Q", "slide": 3, "shape": 0}])
        self.assertEqual(scoped["status"], "ok")
        self.assertEqual(self.slide(3)["shapes"][0]["text"], "Q별 매출 추이")


class ElementOperationTest(KoreanDeckFixture):
    def test_set_transform_moves_a_group_child_in_slide_coordinates(self):
        self.apply([{"op": "set_transform", "slide": 4, "shape": "2.1", "x": 8534400, "w": 2743200}])
        child = self.slide(4)["shapes"][2]["shapes"][1]
        self.assertEqual((child["box"]["x"], child["box"]["w"]), (8534400, 2743200))

    def test_deleting_an_animated_shape_removes_its_animation(self):
        envelope = self.apply([{"op": "delete_shape", "slide": 4, "shape": 2}])
        self.assertIn("1 animations", envelope["details"]["changes"][0]["change"])
        slide = self.presentation().slides[3]
        self.assertNotIn("spTgt", slide._element.xml)
        self.assertEqual(len(slide.shapes), 2)

    def test_a_duplicated_chart_gets_its_own_chart_part(self):
        self.apply([{"op": "duplicate_shape", "slide": 3, "shape": 1, "y": 3600000}, {"op": "set_chart_data", "slide": 3, "shape": 1, "series": [{"name": "매출", "values": [1, 2, 3]}]}])
        charts = [shape.chart for shape in self.presentation().slides[2].shapes if shape.has_chart]
        self.assertEqual(len({chart.part.partname for chart in charts}), 2)
        self.assertEqual([list(chart.plots[0].series[0].values) for chart in charts], [[1.0, 2.0, 3.0], [96.0, 110.0, 128.0]])

    def test_z_order_fill_line_and_picture_replacement(self):
        self.apply([
            {"op": "set_z_order", "slide": 2, "shape": 1, "to": "back"},
            {"op": "set_fill", "slide": 2, "shape": 1, "color": "#1F4E79"},
            {"op": "set_line", "slide": 2, "shape": 1, "color": "none"},
            {"op": "replace_picture", "slide": 2, "shape": 4, "image": "photo.png"},
        ])
        shapes = self.presentation().slides[1].shapes
        self.assertEqual(shapes[0].name, "Rounded Rectangle 2")
        self.assertEqual(str(shapes[0].fill.fore_color.rgb), "1F4E79")
        picture = shapes[4]
        self.assertEqual(picture.image.size, (400, 400))
        self.assertAlmostEqual(picture.crop_top, (1 - 2743200 / 4572000) / 2, places=3)
        self.assertEqual((picture.crop_left, picture.width), (0.0, 4572000))

    def test_a_picture_for_an_image_that_is_not_one_is_refused(self):
        (self.directory / "notes.txt").write_text("not an image")
        self.assertEqual(codes(self.apply([{"op": "add_picture", "slide": 1, "image": "notes.txt", "x": 0, "y": 0}])), ["PICTURE_UNREADABLE"])


class InsertOperationTest(KoreanDeckFixture):
    def test_inserted_shapes_table_and_chart_carry_their_content(self):
        envelope = self.apply([
            {"op": "add_text_box", "slide": 5, "x": 609600, "y": 5800000, "w": 4000000, "h": 400000, "text": "출처: 내부 집계", "size": 12, "color": "666666"},
            {"op": "add_shape", "slide": 5, "kind": "rounded_rectangle", "x": 9000000, "y": 5600000, "w": 2400000, "h": 600000, "fill": "C00000", "text": "결론"},
            {"op": "add_picture", "slide": 5, "image": "photo.png", "x": 100000, "y": 100000, "w": 900000},
            {"op": "add_table", "slide": 5, "x": 609600, "y": 3600000, "w": 6000000, "rows": [["항목", "값"], ["매출", 128]]},
            {"op": "add_chart", "slide": 5, "type": "pie", "x": 7000000, "y": 2400000, "w": 4000000, "h": 2600000, "categories": ["플랫폼", "컨설팅"], "series": [{"name": "비중", "values": [58, 37]}], "title": "사업부 비중"},
        ])
        self.assertNotEqual(envelope["status"], "error", envelope["issues"])
        shapes = self.slide(5)["shapes"]
        self.assertEqual([shape["kind"] for shape in shapes[2:]], ["text", "shape", "picture", "table", "chart"])
        self.assertEqual(shapes[2]["style"]["size"], 12.0)
        self.assertEqual(shapes[4]["box"]["h"], 900000)
        self.assertEqual(shapes[5]["rows"], [["항목", "값"], ["매출", "128"]])
        self.assertEqual(shapes[6]["chart"]["title"], "사업부 비중")

    def test_series_must_match_the_categories(self):
        envelope = self.apply([{"op": "add_chart", "slide": 1, "type": "column", "x": 0, "y": 0, "w": 100, "h": 100, "categories": ["가", "나"], "series": [{"name": "값", "values": [1]}]}])
        self.assertEqual(envelope["issues"][0]["location"], "ops[0].series[0].values")


class TableAndChartOperationTest(KoreanDeckFixture):
    def test_a_new_row_copies_its_neighbor_and_grows_the_frame(self):
        self.apply([{"op": "insert_table_row", "slide": 4, "shape": 1, "values": ["교육", "10억", "12억"]}, {"op": "set_table_cell", "slide": 4, "shape": 1, "row": 0, "column": 0, "text": "부문"}])
        table = self.slide(4)["shapes"][1]
        self.assertEqual(table["rows"][0][0], "부문")
        self.assertEqual(table["rows"][3], ["교육", "10억", "12억"])
        self.assertEqual(table["box"]["h"], 1371600 + 457200)

    def test_columns_are_inserted_and_deleted_and_cells_merged(self):
        self.apply([{"op": "insert_table_column", "slide": 4, "shape": 1, "at": 1, "values": ["비중", "61%", "39%"]}, {"op": "delete_table_column", "slide": 4, "shape": 1, "column": 2}])
        self.assertEqual(self.slide(4)["shapes"][1]["rows"][0], ["사업부", "비중", "실적"])
        self.apply([{"op": "merge_table_cells", "slide": 4, "shape": 1, "row": 1, "column": 1, "rows": 2}])
        merged = self.slide(4, "--detail")["shapes"][1]["merged"]
        self.assertEqual(merged, [{"row": 1, "column": 1, "rows": 2, "columns": 1}])
        refused = self.apply([{"op": "insert_table_row", "slide": 4, "shape": 1}])
        self.assertEqual(codes(refused), ["OPERATION_NOT_APPLICABLE"])

    def test_chart_data_replaces_values_and_keeps_the_chart_type(self):
        self.apply([{"op": "set_chart_data", "slide": 3, "shape": 1, "categories": ["2분기", "3분기"], "series": [{"name": "매출", "values": [110, 131]}], "title": "매출"}])
        chart = self.slide(3)["shapes"][1]["chart"]
        self.assertEqual((chart["type"], chart["title"], chart["categories"]), ("column_clustered", "매출", ["2분기", "3분기"]))
        self.assertEqual(chart["series"], [{"name": "매출", "values": [110.0, 131.0]}])


class SlideOperationTest(KoreanDeckFixture):
    def test_a_slide_added_from_a_layout_is_filled_and_placed(self):
        self.apply([{"op": "add_slide", "layout": "Title and Content", "after": 1, "title": "요약", "body": "매출 성장\n비용 절감"}])
        details = self.read()
        self.assertEqual(details["slideCount"], 6)
        added = details["slides"][1]
        self.assertEqual((added["layout"], added["shapes"][0]["text"], added["shapes"][1]["text"]), ("Title and Content", "요약", "매출 성장\n비용 절감"))

    def test_an_unknown_layout_suggests_the_closest_name(self):
        envelope = self.apply([{"op": "add_slide", "layout": "Title Onl"}])
        self.assertEqual(envelope["issues"][0]["suggestion"], "did you mean 'Title Only'?")

    def test_a_duplicated_slide_keeps_notes_chart_animation_and_section(self):
        self.apply([{"op": "duplicate_slide", "slide": 3}, {"op": "duplicate_slide", "slide": 4}])
        presentation = self.presentation()
        copy_of_chart, copy_of_table = presentation.slides[3], presentation.slides[5]
        self.assertTrue(any(shape.has_chart for shape in copy_of_chart.shapes))
        self.assertNotEqual(copy_of_chart.shapes[1].chart.part.partname, presentation.slides[2].shapes[1].chart.part.partname)
        self.assertIn("spTgt", copy_of_table._element.xml)
        notes = self.read()["slides"]
        self.assertEqual(notes[3]["notes"], notes[2]["notes"])
        section_slides = [len(section) for section in presentation.part._element.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}sldIdLst")]
        self.assertEqual(section_slides, [2, 5])

    def test_a_reorder_after_an_added_slide_keeps_the_added_slide_in_place(self):
        self.apply([{"op": "add_slide", "layout": "Blank", "after": 2}, {"op": "reorder", "order": [5, 4, 3, 2, 1]}])
        layouts = [slide["layout"] for slide in self.read()["slides"]]
        self.assertEqual(layouts, ["Title Slide", "Title Only", "Blank", "Title Only", "Title Only", "Title Slide"])

    def test_hide_background_layout_and_delete(self):
        self.apply([
            {"op": "set_slide_hidden", "slide": 2, "hidden": True},
            {"op": "set_background", "slide": 1, "color": "F3F6FA"},
            {"op": "set_layout", "slide": 5, "layout": "Title Only"},
            {"op": "delete_slide", "slide": 3},
        ])
        details = self.read()
        self.assertTrue(details["slides"][1]["hidden"])
        self.assertEqual(details["slides"][3]["layout"], "Title Only")
        self.assertEqual(details["slideCount"], 4)
        self.assertEqual(str(self.presentation().slides[0].background.fill.fore_color.rgb), "F3F6FA")
        section_slides = [len(section) for section in self.presentation().part._element.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}sldIdLst")]
        self.assertEqual(section_slides, [2, 2])


class DeckOperationTest(KoreanDeckFixture):
    def test_theme_colors_and_fonts_change_what_slides_inherit(self):
        self.apply([{"op": "set_theme", "colors": {"accent1": "2E75B6"}, "koreanFont": "Nanum Gothic", "headingFont": "Arial"}])
        title = self.slide(2, "--detail")["shapes"][0]["paragraphs"][0]["runs"][0]
        self.assertEqual(title["font"], {"value": "Nanum Gothic", "from": "theme"})
        self.assertEqual(self.slide(2, "--detail")["shapes"][1]["fill"], "#2E75B6")

    def test_an_unknown_theme_slot_is_refused_with_the_closest_slot(self):
        envelope = self.apply([{"op": "set_theme", "colors": {"accent7": "2E75B6"}}])
        self.assertEqual(codes(envelope), ["INVALID_VALUE"])
        self.assertIn("accent", envelope["issues"][0]["suggestion"])

    def test_slide_size_scales_the_shapes(self):
        self.apply([{"op": "set_slide_size", "width": 9144000, "height": 6858000}])
        details = self.read()
        self.assertEqual(details["slideSize"]["w"], 9144000)
        self.assertEqual(details["slides"][1]["shapes"][2]["box"]["x"], 571500)


class ValidationTest(KoreanDeckFixture):
    def test_an_unknown_operation_or_field_says_which_was_meant(self):
        envelope = self.apply([{"op": "set_txt", "slide": 1, "shape": 0, "text": "a"}, {"op": "set_text", "slide": 1, "shape": 0, "txt": "a"}])
        self.assertEqual([issue["suggestion"] for issue in envelope["issues"][:2]], ['use "op": "set_text"', "rename the field to 'text'"])

    def test_a_missing_shape_lists_the_shapes_the_slide_has(self):
        envelope = self.apply([{"op": "set_text", "slide": 3, "shape": 7, "text": "a"}])
        self.assertEqual(codes(envelope), ["TARGET_NOT_FOUND"])
        self.assertEqual(envelope["issues"][0]["suggestion"], "use one of 0 'Title 1' (text), 1 'Chart 2' (chart)")
        missing_slide = self.apply([{"op": "set_notes", "slide": 9, "text": "a"}])
        self.assertEqual(missing_slide["issues"][0]["suggestion"], "use a slide number from 1 to 5")

    def test_text_operations_refuse_shapes_without_text(self):
        envelope = self.apply([{"op": "set_text", "slide": 3, "shape": 1, "text": "a"}])
        self.assertEqual(codes(envelope), ["OPERATION_NOT_APPLICABLE"])

    def test_one_failing_operation_writes_nothing(self):
        original = (self.directory / "deck.pptx").read_bytes()
        envelope = self.apply([{"op": "set_notes", "slide": 1, "text": "바뀜"}, {"op": "delete_table_row", "slide": 4, "shape": 1, "row": 9}])
        self.assertEqual(codes(envelope), ["TARGET_NOT_FOUND"])
        self.assertEqual((self.directory / "deck.pptx").read_bytes(), original)


class FidelityTest(KoreanDeckFixture):
    def test_a_text_edit_rewrites_only_its_slide_and_only_the_edited_bytes(self):
        self.apply([{"op": "set_text", "slide": 2, "shape": 3, "paragraph": 1, "text": "재구매율 68%로 상승, 이탈률 4%"}], "--output", "edited.pptx")
        before, after = part_contents(self.directory / "deck.pptx"), part_contents(self.directory / "edited.pptx")
        self.assertEqual(sorted(before), sorted(after))
        self.assertEqual([name for name in before if before[name] != after[name]], ["ppt/slides/slide2.xml"])
        matcher = difflib.SequenceMatcher(None, before["ppt/slides/slide2.xml"], after["ppt/slides/slide2.xml"], autojunk=False)
        changed = [(tag, after["ppt/slides/slide2.xml"][j1:j2].decode()) for tag, _, _, j1, j2 in matcher.get_opcodes() if tag != "equal"]
        self.assertEqual(changed, [("insert", ", 이탈률 4%")])

    def test_structural_edits_keep_animation_extensions_and_unknown_parts(self):
        self.apply([{"op": "duplicate_slide", "slide": 4}, {"op": "set_chart_data", "slide": 3, "shape": 1, "series": [{"name": "매출", "values": [1, 2, 3]}]}, {"op": "insert_table_row", "slide": 4, "shape": 1}], "--output", "edited.pptx")
        before, after = part_contents(self.directory / "deck.pptx"), part_contents(self.directory / "edited.pptx")
        unchanged = [name for name in before if name.startswith(("ppt/slideLayouts/", "ppt/slideMasters/", "ppt/theme/", "ppt/media/", "ppt/notesSlides/", "docProps/")) or name in ("ppt/slides/slide5.xml", CUSTOM_PART_NAME.lstrip("/"))]
        self.assertEqual([name for name in unchanged if before[name] != after[name]], [])
        self.assertIn(UNKNOWN_EXTENSION_URI.encode(), after["ppt/slides/slide5.xml"])
        self.assertIn(b"p:timing", after["ppt/slides/slide4.xml"])


class LayoutAuditTest(KoreanDeckFixture):
    def test_growing_text_is_reported_with_an_operation_that_clears_it(self):
        long_text = "3분기 매출은 128억 원으로 전년 동기 대비 23% 성장했고, 신규 고객 42곳과 재구매율 68%가 함께 성장을 이끌었습니다.\n다음 분기에는 컨설팅 사업부의 수주 회복과 플랫폼 사업부의 가격 정책 조정을 함께 추진해 같은 흐름을 유지합니다.\n세부 실행 계획은 다음 장에서 사업부별로 설명합니다."
        envelope = self.apply([{"op": "set_text", "slide": 2, "shape": 3, "text": long_text}, {"op": "set_text_frame", "slide": 2, "shape": 3, "autofit": "none"}])
        overflow = [issue for issue in envelope["issues"] if issue["code"] == "CONTENT_OVERFLOW" and issue["location"] == "slide 2 shape 3"]
        self.assertEqual(len(overflow), 1)
        fix = overflow[0]["fix"]
        self.assertEqual([(operation["slide"], operation["shape"]) for operation in fix], [(2, 3)])
        fixed = self.apply(fix)
        self.assertNotIn("slide 2 shape 3", [issue["location"] for issue in fixed["issues"] if issue["code"] == "CONTENT_OVERFLOW"])

    def test_off_slide_overlap_and_stretched_pictures_are_found(self):
        envelope = self.apply([
            {"op": "set_transform", "slide": 3, "shape": 1, "y": 3600000},
            {"op": "set_transform", "slide": 2, "shape": 3, "x": 609600, "y": 1828800},
            {"op": "set_transform", "slide": 2, "shape": 4, "w": 6000000},
        ])
        by_code = {issue["code"]: issue for issue in envelope["issues"]}
        self.assertEqual(by_code["OUT_OF_FRAME"]["fix"], [{"op": "set_transform", "slide": 3, "shape": 1, "y": 6858000 - 4572000}])
        self.assertEqual([operation["op"] for operation in by_code["TEXT_OVERLAP"]["fix"]], ["set_transform"])
        self.assertEqual([operation["op"] for operation in by_code["IMAGE_DISTORTED"]["fix"]], ["set_transform"])

    def test_text_that_grows_past_its_card_is_found(self):
        envelope = self.apply([
            {"op": "set_text", "slide": 2, "shape": 2, "text": "3분기 매출 128억 원, 전년 동기 대비 23% 성장하며 분기 최고치를 다시 경신"},
            {"op": "set_text_frame", "slide": 2, "shape": 2, "wrap": True},
        ])
        spill = [issue for issue in envelope["issues"] if "grows with its text past shape 1" in issue["message"]]
        self.assertEqual([operation["shape"] for operation in spill[0]["fix"]], [1])

    def test_text_that_grows_past_the_slide_bottom_gets_an_operation_not_prose(self):
        long_text = "이 문장은 상자에 비해 훨씬 길어서 슬라이드 아래로 넘칠 것입니다. " * 30
        envelope = self.apply([{"op": "set_text", "slide": 2, "shape": 3, "text": long_text}, {"op": "set_text_frame", "slide": 2, "shape": 3, "autofit": "resize", "wrap": True}])
        off_slide = [issue for issue in envelope["issues"] if issue["code"] == "OUT_OF_FRAME" and issue["location"] == "slide 2 shape 3"]
        self.assertEqual(len(off_slide), 1)
        fix = off_slide[0]["fix"]
        self.assertEqual(len(fix), 1)
        fixed = self.apply(fix)
        self.assertNotIn("slide 2 shape 3", [issue["location"] for issue in fixed["issues"] if issue["code"] == "OUT_OF_FRAME"])
        self.assertTrue(all(issue["fix"] for issue in fixed["issues"] if issue["code"] in ("OUT_OF_FRAME", "CONTENT_OVERFLOW")))

    def test_each_suggested_operation_clears_its_issue_or_the_issue_says_no_single_operation_can(self):
        moderate = "3분기 매출은 128억 원으로 전년 동기 대비 23% 성장했고, 신규 고객 42곳과 재구매율 68%가 함께 성장을 이끌었습니다.\n" * 3
        endless = "이 문장은 상자에 비해 훨씬 길어서 슬라이드 아래로 넘칠 것입니다. " * 30
        scenarios = {"moderate": (moderate, "none"), "endless": (endless, "none"), "endless growing": (endless, "resize")}
        for name, (text, autofit) in scenarios.items():
            with self.subTest(name):
                shutil.copy(self.deck_bytes_path, self.directory / "deck.pptx")
                envelope = self.apply([{"op": "set_text", "slide": 2, "shape": 3, "text": text}, {"op": "set_text_frame", "slide": 2, "shape": 3, "autofit": autofit, "wrap": True}])
                found = [issue for issue in envelope["issues"] if issue["location"] == "slide 2 shape 3" and issue["code"] in ("CONTENT_OVERFLOW", "OUT_OF_FRAME")]
                self.assertTrue(found)
                for issue in found:
                    if not issue["fix"]:
                        self.assertIn("no single operation", issue["suggestion"])
                        continue
                    shutil.copy(self.directory / "deck.pptx", self.directory / "before.pptx")
                    after = self.apply(issue["fix"])
                    remaining = [(other["code"], other["location"]) for other in after["issues"]]
                    self.assertNotIn((issue["code"], issue["location"]), remaining, issue["fix"])
                    shutil.copy(self.directory / "before.pptx", self.directory / "deck.pptx")

    def test_check_without_a_preview_reports_the_same_findings(self):
        self.apply([{"op": "set_transform", "slide": 3, "shape": 1, "y": 3600000}])
        envelope = run_office(["deck", "check", "deck.pptx", "--no-preview", "--slides", "3"], self.directory)
        self.assertEqual(codes(envelope), ["OUT_OF_FRAME"])
        self.assertEqual(envelope["details"]["checkedSlides"], [3])


class PreviewTest(KoreanDeckFixture):
    def check_preview(self, *operations):
        if operations:
            self.apply(list(operations))
        envelope = run_office(["deck", "check", "deck.pptx"], self.directory)
        document = lxml.html.fromstring((self.directory / envelope["details"]["preview"]).read_text(encoding="utf-8"))
        return envelope, document

    def section(self, document, number):
        return document.xpath(f"//section[@data-slide='{number}']")[0]

    def styled(self, node):
        return dict(part.split(":", 1) for part in node.get("style").split(";"))

    def test_every_slide_is_a_page_at_the_slide_size(self):
        envelope, document = self.check_preview({"op": "set_slide_hidden", "slide": 2, "hidden": True}, {"op": "set_background", "slide": 1, "color": "F3F6FA"})
        sections = document.xpath("//section")
        self.assertEqual([section.get("data-slide") for section in sections], ["1", "2", "3", "4", "5"])
        self.assertEqual({(self.styled(section)["width"], self.styled(section)["height"]) for section in sections}, {("1280.0px", "720.0px")})
        self.assertEqual(self.styled(sections[0])["background-color"], "#F3F6FA")
        self.assertTrue(all(Path(font["path"]).exists() for font in envelope["details"]["previewFonts"]))

    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_every_slide_is_drawn_and_seen(self):
        envelope, _ = self.check_preview()
        self.assertEqual(codes(envelope), [])
        assert_pages_drawn(self, envelope["details"], self.directory, 5, (1280, 720))

    def test_text_keeps_its_box_style_and_korean_face(self):
        _, document = self.check_preview()
        label = next(span for span in self.section(document, 2).iter("span") if span.text == "매출")
        style = self.styled(label)
        self.assertEqual((style["font-size"], style["font-weight"], style["color"]), ("32.0px", "700", "#FFFFFF"))
        self.assertEqual(style["font-family"], f"'{font_face('맑은 고딕', True, True).family}'")
        box = self.styled(label.getparent().getparent().getparent())
        self.assertEqual((box["left"], box["top"], box["width"]), ("80.0px", "208.0px", "320.0px"))

    def test_pictures_are_cropped_and_group_children_placed_on_the_slide(self):
        _, document = self.check_preview()
        frame = next(node for node in self.section(document, 2).iter("div") if self.styled(node).get("overflow") == "hidden")
        image = self.styled(frame[0])
        self.assertEqual((self.styled(frame)["width"], image["width"], image["left"]), ("480.0px", "600.00px", "-60.00px"))
        member = next(span for span in self.section(document, 4).iter("span") if span.text == "담당")
        self.assertEqual(self.styled(member.getparent().getparent().getparent())["left"], f"{8382000 * 96 / 914400:.1f}px")

    def test_tables_and_charts_are_drawn_from_their_data(self):
        _, document = self.check_preview()
        cells = [node for node in self.section(document, 4).iter("div") if self.styled(node).get("border") == "1px solid #FFFFFF"]
        self.assertEqual(len(cells), 9)
        self.assertEqual(self.styled(cells[0])["background-color"], "#4F81BD")
        chart = next(image for image in self.section(document, 3).iter("img"))
        svg = base64.b64decode(chart.get("src").split(",", 1)[1]).decode("utf-8")
        bars, legend_swatches = 6, 2
        self.assertEqual(svg.count("<rect"), bars + legend_swatches)
        self.assertIn("3분기", svg)

    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_a_chart_page_stays_clear_around_the_bars(self):
        envelope, _ = self.check_preview()
        page = Image.open(self.directory / envelope["details"]["pages"][2]).convert("RGB")
        scale = page.width / Presentation(str(self.directory / "deck.pptx")).slide_width
        corner = (round((609600 + 60000) * scale), round((1524000 + 60000) * scale))
        self.assertGreater(min(page.getpixel(corner)), 200)

    def test_the_preview_uses_only_inline_css_a_renderer_without_selectors_draws(self):
        _, document = self.check_preview()
        markup = lxml.html.tostring(document, encoding="unicode")
        for unsupported in (":has(", "::before", "::after", "counter(", "<style", "<script"):
            self.assertNotIn(unsupported, markup)


if __name__ == "__main__":
    unittest.main()
