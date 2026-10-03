from pathlib import Path
import subprocess
import sys
import tempfile
import textwrap
import unittest

from doc_fixture import OFFICE_ENTRY, run_office


FACTS = ["Example Holdings Co., Ltd.", "Seoul, Korea", "1,234,567", "예시물산(주), 서울 지점", "October 12, 2026"]

MARKDOWN = """# Supply note

Seller: Example Holdings Co., Ltd., registered in Seoul, Korea.

Buyer: 예시물산(주), 서울 지점, order dated October 12, 2026.

| Item | Amount |
| --- | --- |
| Total | 1,234,567 |
"""

SLIDES = """
from pptx import Presentation
from pptx.util import Inches

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text_frame
box.text = "Seller: Example Holdings Co., Ltd., registered in Seoul, Korea."
box.add_paragraph().text = "Buyer: 예시물산(주), 서울 지점, order dated October 12, 2026. Total 1,234,567"
presentation.save("note.pptx")
"""

CSV = 'Party,Place,Total,Branch,Date\n"Example Holdings Co., Ltd.","Seoul, Korea","1,234,567","예시물산(주), 서울 지점","October 12, 2026"\n'


class RequiredTextValueTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)
        (self.directory / "note.md").write_text(MARKDOWN, encoding="utf-8")
        (self.directory / "note.csv").write_text(CSV, encoding="utf-8")
        (self.directory / "slides.py").write_text(textwrap.dedent(SLIDES), encoding="utf-8")
        run_office(["create", "note.docx", "note.md"], self.directory)
        run_office(["create", "note.pdf", "note.md"], self.directory)
        run_office(["create", "note.xlsx", "note.csv"], self.directory)
        subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "slides.py"], check=True, cwd=self.directory)

    def tearDown(self):
        self.temporary_directory.cleanup()

    def missing(self, name: str, values: list[str]) -> list[str]:
        arguments = ["check", name, "--no-preview"] if name.endswith(".pptx") else ["check", name]
        for value in values:
            arguments += ["--required-text", value]
        envelope = run_office(arguments, self.directory)
        return [issue["location"] for issue in envelope["issues"] if issue["code"] == "REQUIRED_TEXT_MISSING"]

    def test_a_value_is_matched_whole_with_its_commas_in_every_format(self):
        for name in ("note.docx", "note.pdf", "note.xlsx", "note.pptx"):
            with self.subTest(file=name):
                self.assertEqual(self.missing(name, FACTS), [])

    def test_facts_joined_into_one_value_are_one_value_and_the_issue_says_to_pass_each_on_its_own(self):
        joined = ",".join(FACTS[:2])
        envelope = run_office(["check", "note.docx", "--required-text", joined], self.directory)
        issue = next(issue for issue in envelope["issues"] if issue["code"] == "REQUIRED_TEXT_MISSING")
        self.assertEqual(issue["location"], joined)
        self.assertIn("its own --required-text", issue["suggestion"])


if __name__ == "__main__":
    unittest.main()
