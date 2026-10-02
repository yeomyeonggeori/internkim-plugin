import zipfile

SPARKLINE_URI = "{05C60535-1F16-4fd2-B633-F4F36F0B64E0}"
CONDITIONAL_FORMATTING_URI = "{78C0D931-6437-407d-A8EE-F0AAD7539E65}"
DATA_VALIDATION_URI = "{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}"
SLICER_LIST_URI = "{A8765BA9-456A-4dab-B4F3-ACF838C121DE}"
WORKBOOK_SLICER_CACHES_URI = "{BBE1A952-AA13-448e-AADC-164F8A28A991}"
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM = "http://schemas.microsoft.com/office/excel/2006/main"
DATA_BAR_ID = "{00000000-000E-0000-0000-000001000000}"

BUILD = """
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.worksheet.table import Table

workbook = Workbook()
data = workbook.active
data.title = "Data"
data.append(["item", "q1", "q2", "q3", "total"])
for index in range(5):
    data.append([f"r{index}", index + 1, index * 2 + 1, index * 3 + 1, f"=SUM(B{index + 2}:D{index + 2})"])
data.add_table(Table(displayName="Items", ref="A1:E6"))
data.conditional_formatting.add("B2:B6", DataBarRule(start_type="min", end_type="max", color="638EC6"))
chart = BarChart()
chart.add_data(Reference(data, min_col=2, min_row=1, max_row=6), titles_from_data=True)
data.add_chart(chart, "H2")
lists = workbook.create_sheet("Lists")
for value in ["r0", "r1", "r2"]:
    lists.append([value])
workbook.save("{path}")
"""

SHEET_EXTENSIONS = (
    f'<extLst>'
    f'<ext uri="{CONDITIONAL_FORMATTING_URI}" xmlns:x14="{X14}"><x14:conditionalFormattings><x14:conditionalFormatting xmlns:xm="{XM}">'
    f'<x14:cfRule type="dataBar" id="{DATA_BAR_ID}"><x14:dataBar minLength="0" maxLength="100" gradient="0" negativeBarColorSameAsPositive="1"><x14:cfvo type="autoMin"/><x14:cfvo type="autoMax"/></x14:dataBar></x14:cfRule>'
    f'<xm:sqref>B2:B6</xm:sqref></x14:conditionalFormatting></x14:conditionalFormattings></ext>'
    f'<ext uri="{DATA_VALIDATION_URI}" xmlns:x14="{X14}"><x14:dataValidations count="1" xmlns:xm="{XM}">'
    f'<x14:dataValidation type="list" allowBlank="1" showInputMessage="1" showErrorMessage="1"><x14:formula1><xm:f>Lists!$A$1:$A$3</xm:f></x14:formula1><xm:sqref>A8:A12</xm:sqref></x14:dataValidation>'
    f'</x14:dataValidations></ext>'
    f'<ext uri="{SPARKLINE_URI}" xmlns:x14="{X14}"><x14:sparklineGroups xmlns:xm="{XM}"><x14:sparklineGroup displayEmptyCellsAs="gap">'
    f'<x14:colorSeries rgb="FF376092"/><x14:sparklines>'
    f'<x14:sparkline><xm:f>Data!B2:D2</xm:f><xm:sqref>F2</xm:sqref></x14:sparkline>'
    f'<x14:sparkline><xm:f>Data!B6:D6</xm:f><xm:sqref>F6</xm:sqref></x14:sparkline>'
    f'</x14:sparklines></x14:sparklineGroup></x14:sparklineGroups></ext>'
    f'<ext uri="{SLICER_LIST_URI}" xmlns:x14="{X14}"><x14:slicerList><x14:slicer r:id="rIdSlicer"/></x14:slicerList></ext>'
    f'</extLst>'
)
SHAPE_ANCHOR = (
    '<xdr:twoCellAnchor xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>8</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
    '<xdr:to><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>11</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="9" name="Note box"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>'
    '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1828800" cy="571500"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
    '<xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>keep this note</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>'
)
SLICER = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<slicers xmlns="{X14}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><slicer name="item" cache="Slicer_item" caption="item" rowHeight="241300"/></slicers>'
)
SLICER_CACHE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<slicerCacheDefinition xmlns="{X14}" name="Slicer_item" sourceName="item"><extLst/></slicerCacheDefinition>'
)
WORKBOOK_EXTENSIONS = (
    f'<extLst><ext uri="{WORKBOOK_SLICER_CACHES_URI}" xmlns:x14="{X14}"><x14:slicerCaches><x14:slicerCache r:id="rIdSlicerCache"/></x14:slicerCaches></ext></extLst>'
)
CUSTOM_XML = b'<?xml version="1.0" encoding="UTF-8" standalone="no"?><note xmlns="urn:example">keep me</note>'
CUSTOM_XML_PROPERTIES = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="no"?>'
    b'<ds:datastoreItem ds:itemID="{11111111-2222-3333-4444-555555555555}" xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml"><ds:schemaRefs/></ds:datastoreItem>'
)
RELATIONSHIPS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_RELATIONSHIPS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def relationships_xml(items):
    body = "".join(f'<Relationship Id="{identifier}" Type="{kind}" Target="{target}"/>' for identifier, kind, target in items)
    return f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{RELATIONSHIPS}">{body}</Relationships>'.encode()


def add_relationship(xml: bytes, identifier: str, kind: str, target: str) -> bytes:
    return xml.replace(b"</Relationships>", f'<Relationship Id="{identifier}" Type="{kind}" Target="{target}"/></Relationships>'.encode())


def enrich(path):
    with zipfile.ZipFile(path) as source:
        entries = {info.filename: source.read(info.filename) for info in source.infolist()}
    sheet = entries["xl/worksheets/sheet1.xml"].decode()
    if 'xmlns:r=' not in sheet.split(">", 2)[1]:
        sheet = sheet.replace("<worksheet ", f'<worksheet xmlns:r="{OFFICE_RELATIONSHIPS}" ', 1)
    sheet = sheet.replace("</dataBar></cfRule>", f'</dataBar><extLst><ext uri="{{B025F937-C7B1-47D3-B67F-A62EFF666E3E}}" xmlns:x14="{X14}"><x14:id>{DATA_BAR_ID}</x14:id></ext></extLst></cfRule>')
    sheet = sheet.replace("<pageMargins", '<phoneticPr fontId="1" type="noConversion"/><pageMargins').replace("<drawing ", '<ignoredErrors><ignoredError sqref="E2:E6" formula="1"/></ignoredErrors><drawing ')
    entries["xl/worksheets/sheet1.xml"] = sheet.replace("</worksheet>", SHEET_EXTENSIONS + "</worksheet>").encode()
    entries["xl/worksheets/_rels/sheet1.xml.rels"] = add_relationship(entries["xl/worksheets/_rels/sheet1.xml.rels"], "rIdSlicer", "http://schemas.microsoft.com/office/2007/relationships/slicer", "../slicers/slicer1.xml")
    drawing = entries["xl/drawings/drawing1.xml"].decode()
    entries["xl/drawings/drawing1.xml"] = drawing.replace("</wsDr>", SHAPE_ANCHOR + "</wsDr>").encode()
    workbook = entries["xl/workbook.xml"].decode()
    if 'xmlns:r=' not in workbook.split(">", 2)[1]:
        workbook = workbook.replace("<workbook ", f'<workbook xmlns:r="{OFFICE_RELATIONSHIPS}" ', 1)
    entries["xl/workbook.xml"] = workbook.replace("</workbook>", WORKBOOK_EXTENSIONS + "</workbook>").encode()
    workbook_relationships = add_relationship(entries["xl/_rels/workbook.xml.rels"], "rIdSlicerCache", "http://schemas.microsoft.com/office/2007/relationships/slicerCache", "slicerCaches/slicerCache1.xml")
    entries["xl/_rels/workbook.xml.rels"] = add_relationship(workbook_relationships, "rIdCustom", f"{OFFICE_RELATIONSHIPS}/customXml", "../customXml/item1.xml")
    entries["xl/slicers/slicer1.xml"] = SLICER.encode()
    entries["xl/slicerCaches/slicerCache1.xml"] = SLICER_CACHE.encode()
    entries["customXml/item1.xml"] = CUSTOM_XML
    entries["customXml/itemProps1.xml"] = CUSTOM_XML_PROPERTIES
    entries["customXml/_rels/item1.xml.rels"] = relationships_xml([("rId1", f"{OFFICE_RELATIONSHIPS}/customXmlProps", "itemProps1.xml")])
    content_types = entries["[Content_Types].xml"].decode()
    overrides = (
        '<Override PartName="/xl/slicers/slicer1.xml" ContentType="application/vnd.ms-excel.slicer+xml"/>'
        '<Override PartName="/xl/slicerCaches/slicerCache1.xml" ContentType="application/vnd.ms-excel.slicerCache+xml"/>'
        '<Override PartName="/customXml/itemProps1.xml" ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/>'
    )
    entries["[Content_Types].xml"] = content_types.replace("</Types>", overrides + "</Types>").encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in entries.items():
            target.writestr(name, content)


def build_rich_workbook(directory, run_office_python, name="rich.xlsx"):
    run_office_python(BUILD.replace("{path}", name), directory)
    enrich(directory / name)
    return directory / name
