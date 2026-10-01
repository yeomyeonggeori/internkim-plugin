import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from doc_fixture import OFFICE_ENTRY, SCRIPTS_PATH, block_texts, run_office, run_office_python
from pdf_fixture import newsletter_pdf_code
from pptx_edit_fixture import build_korean_deck
from render_fixture import can_render, pdf_page_count
from report_fixture import CHART_IMAGE, REPORT_MARKDOWN


FONT_DIRECTORY = SCRIPTS_PATH.parent / "assets" / "fonts" / "paperlogy"
WORKBOOK = """
from openpyxl import Workbook
book = Workbook()
sales = book.active
sales.title = "매출"
for row in (["월", "매출"], ["1월", 1200], ["2월", 1500], ["합계", "=SUM(B2:B3)"]):
    sales.append(row)
book.create_sheet("담당자").append(["이름", "부서"])
book.save("실적.xlsx")
"""


def convert(source, target, working_directory, *flags):
    return run_office(["convert", source, target, *flags], working_directory)


class ConversionFixture(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)
        (self.directory / "보고서.md").write_text(REPORT_MARKDOWN, encoding="utf-8")
        run_office_python(CHART_IMAGE, self.directory)

    def tearDown(self):
        self.temporary_directory.cleanup()


class RouteTableTest(unittest.TestCase):
    def test_every_declared_route_has_a_converter(self):
        script = SCRIPTS_PATH / "convert" / "convert_file.py"
        check = "import runpy, sys; sys.path.insert(0, sys.argv[2]); namespace = runpy.run_path(sys.argv[1], run_name='routes'); print(sorted(namespace['CONVERTERS']) == sorted((route.source, route.target) for route in namespace['ROUTES']))"
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "-c", check, str(script), str(script.parent)], capture_output=True, text=True, check=True)
        self.assertEqual(completed.stdout.strip(), "True")

    def test_an_unsupported_pair_names_the_outputs_the_input_has(self):
        with tempfile.TemporaryDirectory() as directory:
            Path(directory, "보고서.md").write_text("# 제목\n", encoding="utf-8")
            envelope = convert("보고서.md", "보고서.xlsx", directory)
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["UNSUPPORTED_CONVERSION"])
        self.assertEqual(envelope["issues"][0]["suggestion"], ".md converts to .docx, .html, .pdf")


class DocumentRouteTest(ConversionFixture):
    def test_markdown_survives_docx_and_back(self):
        self.assertEqual(convert("보고서.md", "보고서.docx", self.directory)["status"], "ok")
        envelope = convert("보고서.docx", "왕복.md", self.directory)
        self.assertEqual(envelope["status"], "ok", envelope["issues"])
        markdown = (self.directory / "왕복.md").read_text(encoding="utf-8")
        for expected in ("# 2026년 3분기 영업 실적 보고서", "**42억 3,000만 원**", "[사내 대시보드](https://dashboard.example.com)", "| 매출 | 3,820 | 4,230 |", "   - 대형 고객 3곳 다년 계약 전환", "1. 연말 프로모션 집행", "![분기별 매출](왕복-media/image1.png)"):
            self.assertIn(expected, markdown)
        self.assertTrue((self.directory / "왕복-media" / "image1.png").exists())

    def test_html_carries_the_structure_into_docx_and_back(self):
        convert("보고서.md", "보고서.html", self.directory)
        page = (self.directory / "보고서.html").read_text(encoding="utf-8")
        self.assertIn("<strong>42억 3,000만 원</strong>", page)
        self.assertIn('src="data:image/png;base64,', page)
        self.assertEqual(convert("보고서.html", "웹.docx", self.directory)["status"], "ok")
        blocks = block_texts(self.directory, "웹.docx")
        self.assertIn(("heading", "2. 주요 지표"), blocks)
        self.assertIn(("table", [["구분", "2분기", "3분기"], ["매출", "3,820", "4,230"], ["영업이익", "290", "343"]]), blocks)
        self.assertIn(("listItem", "대형 고객 3곳 다년 계약 전환"), blocks)
        convert("웹.docx", "웹.html", self.directory)
        self.assertIn("<ol><li>수도권 영업 인력 2명 충원</li>", (self.directory / "웹.html").read_text(encoding="utf-8"))


class PdfSourceRouteTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)
        run_office_python(newsletter_pdf_code(FONT_DIRECTORY), self.directory)

    def tearDown(self):
        self.temporary_directory.cleanup()

    def test_a_text_pdf_becomes_an_editable_document(self):
        envelope = convert("newsletter.pdf", "newsletter.docx", self.directory)
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["PAGE_WITHOUT_TEXT", "CONVERSION_APPROXIMATED"])
        self.assertEqual([(page["columns"], page["headerFooterLinesDropped"], page["hasText"]) for page in envelope["details"]["pages"]], [(2, 2, True), (1, 2, True), (1, 0, False)])
        blocks = block_texts(self.directory, "newsletter.docx")
        self.assertEqual(blocks[:3], [("heading", "2026년 하반기 교육 안내"), ("heading", "1. 교육 운영 방향"), ("paragraph", "올해 하반기 사내 교육은 직무 역량과 리더십 두 축으로 운영한다. 교육 일정은 부서별 수요 조사를 반영해 확정했다.")])
        texts = [text for _, text in blocks]
        self.assertLess(texts.index("직무 교육은 매월 둘째 주 화요일에 열리며 사전 신청자에 한해 수료증을 발급한다."), texts.index("리더십 과정은 팀장급을 대상으로 하며 외부 강사를 초빙한다. 과정별 정원은 20명이다."))
        self.assertIn(("table", [["과정", "일정", "정원"], ["직무 기초", "9월 9일", "30명"], ["리더십", "10월 14일", "20명"]]), blocks)
        self.assertIn(("listItem", "취소는 교육 3일 전까지 가능하다."), blocks)
        self.assertNotIn("주식회사 예시상사 사내 소식", texts)
        self.assertTrue(read_has_picture(self.directory, "newsletter.docx"))


    def test_each_page_becomes_a_slide_of_text_boxes_or_one_picture(self):
        run_office_python("""
            from fpdf import FPDF
            pdf = FPDF(format="A4")
            pdf.add_font("Korean", "", "{font}")
            pdf.add_page()
            pdf.set_font("Korean", "", 12)
            pdf.set_xy(20, 20)
            pdf.cell(100, 8, "원형 도식 설명")
            pdf.ellipse(60, 60, 80, 80, style="F")
            pdf.output("circle.pdf")
        """.replace("{font}", str(FONT_DIRECTORY / "Paperlogy-4Regular.ttf")), self.directory)
        envelope = convert("newsletter.pdf", "newsletter.pptx", self.directory)
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["PAGE_WITHOUT_TEXT", "CONVERSION_APPROXIMATED"])
        self.assertIn("pdf render newsletter.pdf --pages 3 --scale 2", envelope["issues"][0]["suggestion"])
        self.assertEqual([(page["slide"], page["tables"]) for page in envelope["details"]["pages"]], [("editable", 0), ("editable", 1), ("picture", 0)])
        slides = slide_texts(self.directory, "newsletter.pptx")
        self.assertIn("2026년 하반기 교육 안내", slides[0])
        self.assertIn("리더십 과정은 팀장급을 대상으로 하며 외부 강사를", slides[0])
        self.assertIn("10월 14일", slides[1])
        drawn = convert("circle.pdf", "circle.pptx", self.directory)
        self.assertEqual([(page["slide"], page.get("reason")) for page in drawn["details"]["pages"]], [("picture", "curved shapes or drawings")])
        self.assertIn("pages 1 (curved shapes or drawings) became one picture each", drawn["issues"][-1]["message"])


def slide_texts(directory, name):
    code = (
        "import json\n"
        "from pptx import Presentation\n"
        f"presentation = Presentation({name!r})\n"
        "print(json.dumps([' '.join(cell.text for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells) + ' ' + ' '.join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame) for slide in presentation.slides], ensure_ascii=False))\n"
    )
    completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "-c", code], capture_output=True, text=True, cwd=directory, check=True)
    return json.loads(completed.stdout)

def read_has_picture(directory, name):
    return any(block.get("picture") for block in run_office(["doc", "read", name], directory)["details"]["blocks"])


class TableRouteTest(unittest.TestCase):
    def test_each_sheet_becomes_a_csv_with_computed_formulas_and_csv_becomes_a_styled_workbook(self):
        with tempfile.TemporaryDirectory() as name:
            directory = Path(name)
            run_office_python(WORKBOOK, directory)
            envelope = convert("실적.xlsx", "실적.csv", directory)
            self.assertEqual(envelope["details"]["files"], ["실적-매출.csv", "실적-담당자.csv"])
            self.assertEqual((directory / "실적-매출.csv").read_text(encoding="utf-8-sig").splitlines(), ["월,매출", "1월,1200", "2월,1500", "합계,2700"])
            self.assertEqual(convert("실적.xlsx", "매출.tsv", directory, "--sheet", "매출")["details"]["files"], ["매출.tsv"])
            self.assertEqual(convert("실적-매출.csv", "다시.xlsx", directory)["status"], "ok")
            sheet = run_office(["sheet", "read", "다시.xlsx"], directory)["details"]
            self.assertEqual(sheet["range"]["values"][3], ["합계", 2700])
            self.assertEqual(sheet["sheets"][0]["frozenPanes"], "A2")

    def test_an_ods_sheet_becomes_a_workbook_with_values_dates_and_merges(self):
        with tempfile.TemporaryDirectory() as name:
            directory = Path(name)
            write_ods(directory / "예산.ods")
            envelope = convert("예산.ods", "예산.xlsx", directory)
            self.assertEqual([issue["code"] for issue in envelope["issues"]], ["CONVERSION_APPROXIMATED"])
            sheet = run_office(["sheet", "read", "예산.xlsx"], directory)["details"]
            self.assertEqual(sheet["sheets"][0]["name"], "예산")
            self.assertEqual(sheet["sheets"][0]["mergedCells"], ["A1:B1"])
            self.assertEqual(sheet["range"]["values"], [["2026년 예산", None], ["인건비", 1200], ["인건비", 1200], ["마감", "2026-10-01T00:00:00"], ["=수식 아님", 2400]])


DOCUMENT = """
from docx import Document
document = Document()
document.add_heading("분기 보고서", 0)
document.add_paragraph("3분기 매출은 128억 원입니다.")
document.add_page_break()
document.add_paragraph("다음 분기 계획")
document.save("보고서.docx")
"""


@unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
class PdfRouteTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)

    def tearDown(self):
        self.temporary_directory.cleanup()

    def assert_pdf(self, name, page_count, required_text):
        self.assertEqual(pdf_page_count(self.directory / name), page_count)
        validation = run_office(["pdf", "validate", name, "--required-text", required_text], self.directory)
        self.assertEqual(validation["status"], "ok", validation["issues"])

    def test_a_document_becomes_a_pdf_page_for_each_laid_out_page(self):
        run_office_python(DOCUMENT, self.directory)
        envelope = convert("보고서.docx", "보고서.pdf", self.directory)
        self.assertEqual(envelope["status"], "ok", envelope["issues"])
        self.assert_pdf("보고서.pdf", envelope["details"]["pageCount"], "128억 원")
        self.assertEqual(envelope["details"]["pageCount"], 2)

    def test_a_workbook_becomes_its_printed_pages(self):
        run_office_python(WORKBOOK, self.directory)
        envelope = convert("실적.xlsx", "실적.pdf", self.directory, "--sheet", "매출")
        self.assertEqual(envelope["details"]["pageCount"], 1)
        self.assert_pdf("실적.pdf", 1, "2700")

    def test_a_presentation_becomes_one_page_per_slide(self):
        build_korean_deck(self.directory / "deck.pptx")
        envelope = convert("deck.pptx", "deck.pdf", self.directory)
        self.assertEqual(envelope["status"], "ok", envelope["issues"])
        self.assert_pdf("deck.pdf", 5, "분기별 매출 추이")


class LegacyWorkbookTest(unittest.TestCase):
    def test_xls_cells_keep_their_types(self):
        sys.path[0:0] = [str(SCRIPTS_PATH), str(SCRIPTS_PATH / "convert"), str(SCRIPTS_PATH / "sheet")]
        import xlrd
        from spreadsheet_import import xls_value

        class Cell:
            def __init__(self, ctype, value):
                self.ctype, self.value = ctype, value

        class Book:
            datemode = 0

        values = [xls_value(Book(), Cell(ctype, value)) for ctype, value in ((xlrd.XL_CELL_NUMBER, 3.0), (xlrd.XL_CELL_NUMBER, 2.5), (xlrd.XL_CELL_DATE, 46296.0), (xlrd.XL_CELL_BOOLEAN, 1), (xlrd.XL_CELL_TEXT, "영업"), (xlrd.XL_CELL_EMPTY, ""))]
        self.assertEqual([str(value) for value in values], ["3", "2.5", "2026-10-01", "True", "영업", "None"])


ODS_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">
<office:body><office:spreadsheet><table:table table:name="예산">
<table:table-row><table:table-cell table:number-columns-spanned="2" office:value-type="string"><text:p>2026년 예산</text:p></table:table-cell><table:covered-table-cell/></table:table-row>
<table:table-row table:number-rows-repeated="2"><table:table-cell office:value-type="string"><text:p>인건비</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="1200"/></table:table-row>
<table:table-row><table:table-cell office:value-type="string"><text:p>마감</text:p></table:table-cell><table:table-cell office:value-type="date" office:date-value="2026-10-01"/></table:table-row>
<table:table-row><table:table-cell office:value-type="string"><text:p>=수식 아님</text:p></table:table-cell><table:table-cell table:formula="of:=SUM([.B2:.B3])" office:value-type="float" office:value="2400"/><table:table-cell table:number-columns-repeated="1020"/></table:table-row>
<table:table-row table:number-rows-repeated="1048570"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>
</table:table></office:spreadsheet></office:body></office:document-content>"""


def write_ods(path):
    import zipfile
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        archive.writestr("content.xml", ODS_CONTENT)


if __name__ == "__main__":
    unittest.main()
