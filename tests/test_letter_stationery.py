import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import docx
import pypdfium2
from PIL import Image

from doc_fixture import OFFICE_ENTRY
from render_fixture import can_render
from task_context_fixture import environment_with_context, write_context_at

from paperwork.jurisdictions import jurisdiction_codes
from schemas.schema_document import bundled_schema_names, load_schema


KOREAN_PROFILE = {"name": "샘플테크 주식회사", "representative": "최견본", "representativeTitle": "대표이사", "address": "서울특별시 예시구 견본로 1", "registrationNumber": "000-00-00000"}
ENGLISH_PROFILE = {"name": "Sample Tech Inc.", "representative": "Choi Gyeonbon", "representativeTitle": "CEO", "address": "1 Sample-ro, Example-gu, Seoul"}
KOREAN_LETTER = {"language": "ko", "recipient": "견본상사 대표이사", "title": "사무실 이전 안내", "sections": [{"heading": None, "blocks": [{"type": "paragraph", "text": "당사 사무실이 11월 1일부로 이전합니다."}]}], "enclosures": ["약도 1부"]}
ENGLISH_LETTER = {"language": "en", "recipient": "Example Imports LLC", "title": "Notice of office relocation", "salutation": "Dear Ms. Park,", "sections": [{"heading": None, "blocks": [{"type": "paragraph", "text": "We are moving our office on November 1."}]}]}


class SchemaPlacementTest(unittest.TestCase):
    def test_every_bundled_schema_sits_in_a_jurisdiction_folder(self):
        for name in bundled_schema_names():
            with self.subTest(name):
                self.assertIn(name.split("/")[0], jurisdiction_codes())
                self.assertEqual(load_schema(name).name, name)

    def test_a_bare_name_one_schema_answers_to_resolves_to_it(self):
        self.assertEqual(load_schema("letter").name, "intl/letter")
        self.assertEqual(load_schema("report").name, "intl/report")


@unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
class LetterTest(unittest.TestCase):
    def setUp(self):
        self.directory = Path(self.enterContext(tempfile.TemporaryDirectory()))

    def merged(self, language: str, profile: dict, values: dict, output: str) -> dict:
        (self.directory / "values.json").write_text(json.dumps(values, ensure_ascii=False), encoding="utf-8")
        context = write_context_at(self.directory / "context.json", {"requester": {"name": "이샘플"}, "today": "2026-10-23", "company": {language: profile}})
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "merge", "intl/letter", "values.json", output], capture_output=True, text=True, cwd=self.directory, env=environment_with_context(context))
        result = json.loads(completed.stdout)
        self.assertEqual(result["status"], "ok", result)
        return result

    def pdf_text(self, name: str) -> str:
        document = pypdfium2.PdfDocument(str(self.directory / name))
        try:
            return "\n".join(page.get_textpage().get_text_range() for page in document)
        finally:
            document.close()

    def docx_text(self, name: str) -> str:
        document = docx.Document(str(self.directory / name))
        cells = [paragraph.text for table in document.tables for cell in table._cells for paragraph in cell.paragraphs]
        return "\n".join([*cells, *(paragraph.text for paragraph in document.paragraphs)])

    def test_a_korean_letter_is_a_gongmun_on_the_company_letterhead(self):
        self.merged("ko", KOREAN_PROFILE, KOREAN_LETTER, "letter.pdf")
        text = self.pdf_text("letter.pdf")
        for expected in ("샘플테크 주식회사", "사업자등록번호 000-00-00000", "서울특별시 예시구 견본로 1", "수신", "제목", "붙임", "1. 약도 1부", "끝.", "2026년 10월 23일", "대표이사 최견본", "(인)"):
            with self.subTest(expected):
                self.assertIn(expected, text)

    def test_an_english_letter_signs_off_without_a_seal(self):
        self.merged("en", ENGLISH_PROFILE, ENGLISH_LETTER, "letter.pdf")
        text = self.pdf_text("letter.pdf")
        for expected in ("Sample Tech Inc.", "October 23, 2026", "Example Imports LLC", "Notice of office relocation", "Dear Ms. Park,", "Sincerely,", "Choi Gyeonbon", "CEO, Sample Tech Inc."):
            with self.subTest(expected):
                self.assertIn(expected, text)
        self.assertNotIn("(인)", text)
        self.assertNotIn("끝.", text)

    def test_a_letterhead_leaves_out_what_the_profile_lacks_instead_of_drawing_a_blank(self):
        result = self.merged("ko", KOREAN_PROFILE, KOREAN_LETTER, "letter.pdf")
        self.assertEqual(result["details"]["blanks"], [])
        self.assertNotIn("전화", self.pdf_text("letter.pdf"))

    def test_word_letters_carry_the_same_stationery(self):
        stamp = self.directory / "stamp.png"
        Image.new("RGBA", (60, 60), (200, 30, 30, 255)).save(stamp)
        self.merged("ko", {**KOREAN_PROFILE, "stampPath": str(stamp)}, KOREAN_LETTER, "letter.docx")
        text = self.docx_text("letter.docx")
        for expected in ("샘플테크 주식회사", "사업자등록번호 000-00-00000", "수신", "끝.", "대표이사 최견본"):
            with self.subTest(expected):
                self.assertIn(expected, text)
        self.assertEqual(len(docx.Document(str(self.directory / "letter.docx")).inline_shapes), 1)


if __name__ == "__main__":
    unittest.main()
