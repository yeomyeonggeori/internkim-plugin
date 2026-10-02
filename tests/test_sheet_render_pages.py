import re
import unittest
import zipfile

from openpyxl import load_workbook

from sheet_fixture import WorkbookFixture, run_office

MONTHS = [["month", "sales", "margin"], ["Jan", 120, 0.21], ["Feb", 150, 0.24], ["Mar", 90, 0.18], ["Apr", 170, 0.27]]


class RenderedPagesTest(WorkbookFixture):
    def render(self):
        envelope = run_office(["render", "book.xlsx"], self.directory)
        self.assertNotEqual(envelope["status"], "error", envelope)
        return envelope, (self.directory / "book-preview" / "preview.html").read_text(encoding="utf-8")

    def chart_svgs(self, preview):
        return re.findall(r"<svg.*?</svg>", preview, flags=re.DOTALL)

    def test_a_combo_chart_draws_its_bars_and_its_line_on_a_second_axis(self):
        self.create_workbook([{"title": "Sales", "rows": MONTHS}])
        self.apply([{"op": "add_chart", "type": "combo", "range": "A1:C5", "anchor": "E2", "title": "Sales and margin"}])
        _, preview = self.render()
        svg = self.chart_svgs(preview)[0]
        self.assertEqual(svg.count("<rect x=") - 1 - 2, 4)
        self.assertEqual(svg.count("<polyline"), 1)
        self.assertRegex(svg, r'text-anchor="start" dominant-baseline="middle">0.3<')
        self.assertIn(">margin<", svg)

    def test_a_pie_is_drawn_in_the_slice_colors_it_was_given(self):
        self.create_workbook([{"title": "Sales", "rows": MONTHS}])
        self.apply([{"op": "add_chart", "type": "pie", "range": "A1:B5", "anchor": "E2", "colors": ["111111", "222222", "333333", "444444"]}])
        _, preview = self.render()
        fills = re.findall(r'fill="(#[0-9a-f]{6})"', self.chart_svgs(preview)[0])
        self.assertEqual([color for color in ("#111111", "#222222", "#333333", "#444444") if color in fills], ["#111111", "#222222", "#333333", "#444444"])
        self.apply([{"op": "add_chart", "type": "bar", "range": "A1:C5", "anchor": "E20", "colors": ["AA0000", "00AA00"]}])
        _, preview = self.render()
        bar_fills = set(re.findall(r'fill="(#[0-9a-f]{6})"', self.chart_svgs(preview)[1]))
        self.assertTrue({"#aa0000", "#00aa00"} <= bar_fills, bar_fills)

    def test_horizontal_stacked_bars_stay_horizontal_and_stacked(self):
        self.create_workbook([{"title": "Sales", "rows": [["team", "won", "lost"], ["North", 4, 2], ["South", 3, 5]]}])
        self.apply([{"op": "add_chart", "type": "bar", "range": "A1:C3", "anchor": "E2", "horizontal": True, "stacked": True}])
        _, preview = self.render()
        bars = re.findall(r'<rect x="([\d.]+)" y="([\d.]+)" width="([\d.]+)" height="([\d.]+)" fill="#', self.chart_svgs(preview)[0])
        north_won, north_lost = bars[1], bars[3]
        self.assertEqual(north_won[1], north_lost[1])
        self.assertAlmostEqual(float(north_won[0]) + float(north_won[2]), float(north_lost[0]), delta=0.2)
        self.assertGreater(float(north_won[2]), float(north_won[3]))

    def test_a_chart_across_a_page_break_draws_its_rest_on_the_next_page_and_says_so(self):
        self.create_workbook([{"title": "Sales", "rows": [["day", "sales"]] + [[f"d{index}", index] for index in range(1, 61)]}])
        self.apply([{"op": "add_chart", "type": "line", "range": "A1:B13", "anchor": "D40", "title": "Daily sales", "width": 10}])
        envelope, preview = self.render()
        self.assertNotIn("BLANK_PAGE", [issue["code"] for issue in envelope["issues"]])
        contents = envelope["details"]["pageContents"]
        self.assertEqual([(page["page"], page["sheet"]) for page in contents], [(1, "Sales"), (2, "Sales")])
        self.assertTrue(contents[0]["cells"].startswith("A1:") and contents[1]["cells"].startswith("A"))
        self.assertEqual([page["drawings"] for page in contents], [[{"chart": "Daily sales", "type": "line", "shown": "part, cut at the page edge"}]] * 2)
        self.assertEqual(len(self.chart_svgs(preview)), 2)

    def test_a_long_heading_stays_on_one_line_and_a_merged_title_does_not_grow_its_row(self):
        heading = "판매 실적 원본 데이터 (판매실적_2026.csv · 2026-04~2026-09)"
        self.create_workbook([{"title": "Sales", "heading": heading, "rows": MONTHS}])
        sheet = load_workbook(self.directory / "book.xlsx")["Sales"]
        self.assertFalse(sheet["A1"].alignment.wrap_text)
        self.assertLess(sheet.column_dimensions["A"].width, 20)
        self.apply([{"op": "merge_cells", "range": "A1:C1"}, {"op": "format_range", "range": "A1", "wrapText": True}])
        _, preview = self.render()
        first_row = re.search(r"grid-template-rows:([\d.]+)px", preview).group(1)
        self.assertLess(float(first_row), 30)

    def test_a_page_with_nothing_on_it_is_reported_with_its_number(self):
        self.create_workbook([{"title": "Sales", "rows": MONTHS}])
        self.apply([{"op": "set_cell", "cell": "A200", "value": "note"}])
        envelope, _ = self.render()
        issue = next(issue for issue in envelope["issues"] if issue["code"] == "BLANK_PAGE")
        self.assertIn("3 of 5 printed pages show nothing: page 2 (Sales), page 3 (Sales), page 4 (Sales)", issue["message"])
        self.assertIn("print area", issue["suggestion"])

    def test_a_full_sheet_reports_no_blank_page(self):
        self.create_workbook([{"title": "Sales", "rows": MONTHS}])
        envelope, _ = self.render()
        self.assertNotIn("BLANK_PAGE", [issue["code"] for issue in envelope["issues"]])

    def test_a_sheet_printing_pages_wide_is_a_print_defect_with_a_fit_to_width_fix(self):
        header = [f"항목{index}" for index in range(1, 13)]
        self.create_workbook([{"title": "넓은표", "rows": [header, list(range(1, 13))], "columnWidths": {letter: 10 for letter in "ABCDEFGHIJKL"}}])
        rendered, _ = self.render()
        checked = run_office(["check", "book.xlsx"], self.directory)
        for envelope in (rendered, checked):
            issue = next(issue for issue in envelope["issues"] if issue["code"] == "SHEET_PRINTS_WIDE")
            self.assertEqual((issue["location"], issue["fix"]), ("넓은표", [{"op": "set_page_setup", "sheet": "넓은표", "fitToWidth": 1}]))
        self.assertNotIn("PREVIEW_APPROXIMATED", [issue["code"] for issue in rendered["issues"]])
        self.assertEqual(self.apply(issue["fix"])["status"], "ok")
        self.assertNotIn("SHEET_PRINTS_WIDE", [issue["code"] for issue in run_office(["check", "book.xlsx"], self.directory)["issues"]])
        self.assertEqual(self.render()[0]["details"]["pageCount"], 1)

    def test_a_width_the_author_spread_over_pages_is_left_alone(self):
        header = [f"항목{index}" for index in range(1, 13)]
        self.create_workbook([{"title": "넓은표", "rows": [header, list(range(1, 13))], "columnWidths": {letter: 18 for letter in "ABCDEFGHIJKL"}}])
        self.apply([{"op": "set_page_setup", "fitToWidth": 2, "fitToHeight": 0}])
        self.assertNotIn("SHEET_PRINTS_WIDE", [issue["code"] for issue in self.render()[0]["issues"]])

    def test_a_chart_cut_at_the_page_edge_is_offered_a_place_under_the_tables(self):
        self.create_workbook([{"title": "Sales", "rows": MONTHS}])
        self.apply([{"op": "add_chart", "type": "line", "range": "A1:B5", "anchor": "F2", "width": 20}])
        issue = next(issue for issue in run_office(["check", "book.xlsx"], self.directory)["issues"] if issue["code"] == "SHEET_PRINTS_WIDE")
        self.assertIn('edit_chart "anchor": "A7"', issue["suggestion"])

    def test_a_chart_beside_the_data_that_shrinks_the_fitted_page_is_moved_under_the_data(self):
        rows = [["지역", "목표", "실적", "달성률", "고객수", "비고"]] + [[f"지역{index}", 100 + index, 90 + index, 0.9, index, "-"] for index in range(1, 21)]
        self.create_workbook([{"title": "요약", "rows": rows}])
        self.apply([{"op": "add_chart", "type": "line", "range": "A1:C21", "anchor": "G2", "width": 20}, {"op": "set_page_setup", "fitToWidth": 1, "fitToHeight": 0}])
        rendered, checked = self.render()[0], run_office(["check", "book.xlsx"], self.directory)
        for envelope in (rendered, checked):
            issue = next(issue for issue in envelope["issues"] if issue["code"] == "SHEET_PRINTS_SMALL")
            self.assertIn("11 pt body text prints at", issue["message"])
            self.assertIn("8 pt", issue["message"])
            self.assertIn("chart 0", issue["message"])
            self.assertEqual(issue["fix"], [{"op": "edit_chart", "sheet": "요약", "chart": 0, "anchor": "A23"}])
        self.assertEqual(self.apply(issue["fix"])["status"], "ok")
        print_codes = {"SHEET_PRINTS_SMALL", "SHEET_PRINTS_WIDE"}
        self.assertEqual([issue["code"] for issue in run_office(["check", "book.xlsx"], self.directory)["issues"] if issue["code"] in print_codes], [])
        self.assertEqual([issue["code"] for issue in self.render()[0]["issues"] if issue["code"] in print_codes], [])

    def test_a_table_too_wide_to_read_on_one_page_turns_landscape_or_spreads_over_pages(self):
        cases = (
            (10, {"op": "set_page_setup", "sheet": "넓은표", "orientation": "landscape", "fitToWidth": 1}),
            (12, {"op": "set_page_setup", "sheet": "넓은표", "fitToWidth": 2}),
        )
        for column_count, page_setup in cases:
            with self.subTest(columns=column_count):
                letters = "ABCDEFGHIJKL"[:column_count]
                header = [f"항목{index}" for index in range(1, column_count + 1)]
                self.create_workbook([{"title": "넓은표", "rows": [header, list(range(1, column_count + 1))], "columnWidths": {letter: 18 for letter in letters}}])
                wide = next(issue for issue in run_office(["check", "book.xlsx"], self.directory)["issues"] if issue["code"] == "SHEET_PRINTS_WIDE")
                self.assertEqual(wide["fix"], [page_setup])
                self.assertEqual(self.apply([{"op": "set_page_setup", "fitToWidth": 1}])["status"], "ok")
                small = next(issue for issue in run_office(["check", "book.xlsx"], self.directory)["issues"] if issue["code"] == "SHEET_PRINTS_SMALL")
                self.assertEqual(small["fix"], [page_setup])
                self.assertEqual(self.apply(small["fix"])["status"], "ok")
                remaining = [issue["code"] for issue in run_office(["check", "book.xlsx"], self.directory)["issues"]]
                self.assertNotIn("SHEET_PRINTS_SMALL", remaining)
                self.assertNotIn("SHEET_PRINTS_WIDE", remaining)

    def test_a_fit_that_keeps_the_text_readable_and_a_low_print_scale(self):
        header = [f"항목{index}" for index in range(1, 10)]
        self.create_workbook([{"title": "표", "rows": [header, list(range(1, 10))], "columnWidths": {letter: 12 for letter in "ABCDEFGHI"}}])
        self.apply([{"op": "set_page_setup", "fitToWidth": 1}])
        self.assertNotIn("SHEET_PRINTS_SMALL", [issue["code"] for issue in run_office(["check", "book.xlsx"], self.directory)["issues"]])
        self.apply([{"op": "set_page_setup", "scale": 50}])
        issue = next(issue for issue in run_office(["check", "book.xlsx"], self.directory)["issues"] if issue["code"] == "SHEET_PRINTS_SMALL")
        self.assertIn("prints at 50% of full size", issue["message"])
        self.assertEqual(self.apply(issue["fix"])["status"], "ok")
        self.assertNotIn("SHEET_PRINTS_SMALL", [issue["code"] for issue in run_office(["check", "book.xlsx"], self.directory)["issues"]])

    def test_fit_to_one_page_tall_and_printed_gridlines_shape_the_pages(self):
        self.create_workbook([{"title": "일지", "rows": [["일자", "건수"]] + [[f"{day}일", day] for day in range(1, 151)]}])
        self.assertGreater(self.render()[0]["details"]["pageCount"], 1)
        self.assertEqual(self.apply([{"op": "set_page_setup", "fitToHeight": 1, "fitToWidth": 0, "printGridlines": True}])["status"], "ok")
        sheet = load_workbook(self.directory / "book.xlsx")["일지"]
        self.assertEqual((sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight, sheet.print_options.gridLines), (0, 1, True))
        self.apply([{"op": "set_cell", "sheet": "일지", "cell": "D2", "value": "메모"}])
        envelope, preview = self.render()
        self.assertEqual(envelope["details"]["pageCount"], 1)
        self.assertIn("border-right:1px solid #d0d0d0", preview.replace(": ", ":"))


SHARES = [["분기", "국내", "해외"], ["1분기", 30, 10], ["2분기", 20, 20], ["3분기", 45, 5]]


class PercentChartTest(WorkbookFixture):
    def chart_xml(self, number):
        with zipfile.ZipFile(self.directory / "book.xlsx") as archive:
            return archive.read(f"xl/charts/chart{number}.xml").decode()

    def test_percent_stacked_columns_and_percent_slices_are_written_and_drawn(self):
        self.create_workbook([{"title": "매출", "rows": SHARES}])
        envelope = self.apply([
            {"op": "add_chart", "type": "bar", "range": "A1:C4", "stacked": "percent", "dataLabels": "value", "anchor": "E2"},
            {"op": "add_chart", "type": "pie", "range": "A1:B4", "dataLabels": "category_percent", "anchor": "E20"},
            {"op": "add_chart", "type": "area", "range": "A1:C4", "stacked": "percent", "anchor": "E38"},
        ])
        self.assertEqual(envelope["status"], "ok", envelope)
        columns = self.chart_xml(1)
        self.assertIn('<grouping val="percentStacked"/>', columns)
        self.assertIn('formatCode="0%"', columns)
        self.assertIn('<showVal val="1"/>', columns)
        self.assertIn('<showCatName val="1"/><showSerName val="0"/><showPercent val="1"/>', self.chart_xml(2))
        self.assertIn('<grouping val="percentStacked"/>', self.chart_xml(3))
        run_office(["render", "book.xlsx"], self.directory)
        svgs = re.findall(r"<svg.*?</svg>", (self.directory / "book-preview" / "preview.html").read_text(encoding="utf-8"), flags=re.DOTALL)
        heights = [float(height) for height in re.findall(r'<rect x="[\d.]+" y="[\d.]+" width="[\d.]+" height="([\d.]+)" fill="#', svgs[0])[1:]]
        self.assertAlmostEqual(heights[0] + heights[3], heights[1] + heights[4], delta=0.2)
        self.assertIn(">100%<", svgs[0])
        self.assertIn(">30<", svgs[0])
        self.assertIn(">1분기 32%<", svgs[1])
        self.assertIn(">100%<", svgs[2])

    def test_percent_labels_on_a_chart_without_slices_are_refused(self):
        self.create_workbook([{"title": "매출", "rows": SHARES}])
        issue = self.apply([{"op": "add_chart", "type": "bar", "range": "A1:C4", "dataLabels": "percent"}])["issues"][0]
        self.assertEqual((issue["code"], issue["location"]), ("OPERATION_NOT_APPLICABLE", "ops[0].dataLabels"))
        self.assertIn('"stacked": "percent"', issue["suggestion"])



class StyleRuleTest(WorkbookFixture):
    def test_strikethrough_rotation_theme_colors_and_text_rules_are_written_and_drawn(self):
        self.create_workbook([{"title": "현황", "rows": [["이름", "상태", "점수"], ["이샘플", "완료", 90], ["박예시", "진행 중", 40], ["최견본", None, 70]]}])
        envelope = self.apply([
            {"op": "format_range", "sheet": "현황", "range": "A2", "strikethrough": True, "textRotation": 45},
            {"op": "format_range", "sheet": "현황", "range": "A3", "textRotation": -30},
            {"op": "format_range", "sheet": "현황", "range": "B1", "textRotation": "vertical"},
            {"op": "format_range", "sheet": "현황", "range": "A1:C1", "fill": "accent1+40%", "fontColor": "dk2"},
            {"op": "add_conditional_format", "sheet": "현황", "range": "B2:B4", "rule": "begins_with", "value": "완", "fill": "C6EFCE"},
            {"op": "add_conditional_format", "sheet": "현황", "range": "B2:B4", "rule": "ends_with", "value": "중", "fill": "FFEB9C"},
            {"op": "add_conditional_format", "sheet": "현황", "range": "B2:B4", "rule": "blank", "fill": "D9D9D9"},
            {"op": "add_conditional_format", "sheet": "현황", "range": "C2:C4", "rule": "color_scale", "minColor": "F8696B", "midColor": "FFFFFF", "maxColor": "63BE7B", "minValue": 0, "midValue": 50, "maxValue": 100},
        ])
        self.assertEqual(envelope["status"], "ok", envelope)
        sheet = load_workbook(self.directory / "book.xlsx")["현황"]
        self.assertEqual((sheet["A2"].font.strike, sheet["A2"].alignment.textRotation, sheet["A3"].alignment.textRotation, sheet["B1"].alignment.textRotation), (True, 45, 120, 255))
        self.assertEqual((sheet["A1"].fill.fgColor.theme, sheet["A1"].fill.fgColor.tint, sheet["A1"].font.color.theme), (4, 0.4, 3))
        rules = [rule for formatting in sheet.conditional_formatting for rule in formatting.rules]
        self.assertEqual([rule.type for rule in rules], ["beginsWith", "endsWith", "containsBlanks", "colorScale"])
        self.assertEqual([(cfvo.type, cfvo.val) for cfvo in rules[3].colorScale.cfvo], [("num", 0.0), ("num", 50.0), ("num", 100.0)])
        run_office(["render", "book.xlsx"], self.directory)
        compact = (self.directory / "book-preview" / "preview.html").read_text(encoding="utf-8").replace(": ", ":")
        for drawn in ("text-decoration:line-through;display:inline-block;transform:rotate(-45deg)", "rotate(30deg)", "background:#95b3d7", "background:#c6efce", "background:#ffeb9c", "background:#d9d9d9", "background:#fee1e1"):
            with self.subTest(drawn=drawn):
                self.assertIn(drawn, compact)

    def test_a_color_name_or_a_theme_alias_is_answered_with_the_exact_value(self):
        self.create_workbook([{"title": "현황", "rows": [["이름", "점수"], ["이샘플", 90]]}])
        issues = self.apply([{"op": "format_range", "range": "A1", "fill": "red", "fontColor": "background1+10%"}])["issues"]
        self.assertEqual([(issue["location"], issue["suggestion"]) for issue in issues], [("ops[0].fontColor", 'use "lt1+10%"'), ("ops[0].fill", 'use "FF0000" for red')])


if __name__ == "__main__":
    unittest.main()
