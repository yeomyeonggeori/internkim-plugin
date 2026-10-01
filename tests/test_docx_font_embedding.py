import io
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from fontTools.ttLib import TTFont
from lxml import etree

OFFICE_SCRIPTS_PATH = Path(__file__).resolve().parents[1] / "skills" / "office" / "scripts"
sys.path.insert(0, str(OFFICE_SCRIPTS_PATH))

from fonts.registry import SANS_BODY, default_family  # noqa: E402
from doc.docx_defaults import apply_korean_defaults  # noqa: E402
from fonts.docx_embedding import embed_named_fonts, obfuscated, save_document  # noqa: E402

WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
RELATIONSHIP_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
BODY_FAMILY = default_family(SANS_BODY).name
TEXT = "갑과 을은 다음과 같이 용역 계약을 체결한다. 뷁 똠양꿍 ①"


def written_document(path: Path, font_name: str) -> None:
    document = Document()
    apply_korean_defaults(document, font_name)
    document.add_heading("용역 계약서", 1)
    document.add_paragraph(TEXT)
    save_document(document, path)


def archive_parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def font_entry(parts: dict[str, bytes], name: str):
    table = etree.fromstring(parts["word/fontTable.xml"])
    return next((entry for entry in table.iter(f"{WORD}font") if entry.get(f"{WORD}name") == name), None)


def embedded_font(parts: dict[str, bytes], embed) -> TTFont:
    relationships = etree.fromstring(parts["word/_rels/fontTable.xml.rels"])
    target = next(relationship.get("Target") for relationship in relationships if relationship.get("Id") == embed.get(RELATIONSHIP_ID))
    return TTFont(io.BytesIO(obfuscated(parts[f"word/{target}"], embed.get(f"{WORD}fontKey"))))


class DocxFontEmbeddingTest(unittest.TestCase):
    def setUp(self):
        self.directory = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: __import__("shutil").rmtree(self.directory, ignore_errors=True))

    def test_a_docx_carries_the_regular_and_bold_bundled_faces_it_names(self):
        path = self.directory / "계약서.docx"
        written_document(path, BODY_FAMILY)
        parts = archive_parts(path)
        entry = font_entry(parts, BODY_FAMILY)
        styles = {}
        for element in ("embedRegular", "embedBold"):
            font = embedded_font(parts, entry.find(f"{WORD}{element}"))
            self.assertEqual(font["name"].getDebugName(1), BODY_FAMILY)
            self.assertEqual([character for character in TEXT if ord(character) not in font.getBestCmap()], [])
            styles[element] = font["name"].getDebugName(2)
        self.assertEqual(styles, {"embedRegular": "Regular", "embedBold": "Bold"})
        self.assertIsNotNone(etree.fromstring(parts["word/settings.xml"]).find(f"{WORD}embedTrueTypeFonts"))
        self.assertIn(b'Extension="odttf"', parts["[Content_Types].xml"])

    def test_saving_an_unchanged_document_again_keeps_its_fonts_and_font_table(self):
        path = self.directory / "계약서.docx"
        written_document(path, BODY_FAMILY)
        first = archive_parts(path)
        save_document(Document(path), path)
        again = archive_parts(path)
        for name in [name for name in first if name.endswith(".odttf") or "fontTable" in name]:
            self.assertEqual(again[name], first[name], name)
        before = path.read_bytes()
        embed_named_fonts(path)
        self.assertEqual(path.read_bytes(), before)

    def test_a_docx_naming_only_fonts_the_skill_does_not_ship_carries_none(self):
        path = self.directory / "계약서.docx"
        written_document(path, "맑은 고딕")
        parts = archive_parts(path)
        self.assertFalse([name for name in parts if name.endswith(".odttf")])
        self.assertIsNone(etree.fromstring(parts["word/settings.xml"]).find(f"{WORD}embedTrueTypeFonts"))


if __name__ == "__main__":
    unittest.main()
