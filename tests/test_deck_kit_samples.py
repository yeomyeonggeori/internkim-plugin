import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unittest


TESTS_PATH = Path(__file__).resolve().parent
SCRIPTS_PATH = TESTS_PATH.parent / "skills" / "office" / "scripts"
OFFICE_ENTRY = SCRIPTS_PATH / "office"
SAMPLE_DECKS_PATH = TESTS_PATH / "fixtures" / "deck-kit"
sys.path.insert(0, str(SCRIPTS_PATH))
sys.path.insert(0, str(TESTS_PATH))

from deck.check_deck import CheckRequest, check_deck  # noqa: E402
from png_fixture import write_png  # noqa: E402
from render_fixture import can_render, pdf_page_count  # noqa: E402


LAYOUT_DEFECT_CODES = {"CONTENT_OVERFLOW", "TEXT_OVERLAP", "OUT_OF_FRAME", "VERTICAL_DEAD_ZONE"}
PERCENT_DONUT_DECK = """<body data-theme="corporate">
<section data-layout="cover"><h1>제품별 매출 비중</h1><p class="lead">주식회사 예시랩 · 2026년 3분기</p></section>
<section data-layout="chart"><h2>매출의 절반 이상이 클라우드입니다</h2>
<figure data-chart="donut" data-labels="클라우드, 온프레미스, 컨설팅" data-values="52, 31, 17" data-unit="%"><figcaption>제품별 매출 비중</figcaption></figure></section>
<section data-layout="closing"><h2>클라우드 비중을 더 키웁니다</h2></section>
</body>
"""
BRAND_TOKEN_DECK = """<head><style>:root { --accent: #E4002B; }</style></head><body data-theme="corporate">
<section data-layout="cover"><h1>브랜드 색으로 그립니다</h1><p class="lead">주식회사 예시랩</p></section>
<section data-layout="chart"><h2>매출이 늘었습니다</h2><figure data-chart="column" data-labels="1Q, 2Q" data-values="96, 128" data-unit="억"><figcaption>분기 매출</figcaption></figure></section>
</body>
"""
GENERATED_PHOTO_SIZE = (960, 640)


def sample_deck_paths() -> list[Path]:
    return sorted(path for path in SAMPLE_DECKS_PATH.iterdir() if (path / "slides.html").exists())


def slide_count(deck_path: Path) -> int:
    return (deck_path / "slides.html").read_text(encoding="utf-8").count("<section")


def referenced_images(deck_path: Path) -> list[str]:
    return re.findall(r'src="(images/[^"]+\.png)"', (deck_path / "slides.html").read_text(encoding="utf-8"))


def legend_amounts(deck_path: Path) -> list[str]:
    source = (deck_path / "slides.html").read_text(encoding="utf-8")
    round_charts = re.findall(r'data-chart="(?:donut|pie)"[^>]*data-values="([^"]+)"[^>]*data-unit="([^"]*)"', source)
    return [f"{int(value):,}{unit}" for values, unit in round_charts for value in values.split(",")]


def write_gradient_photo(path: Path) -> None:
    width, height = GENERATED_PHOTO_SIZE
    rows = [[(40 + 120 * column // width, 70 + 90 * row // height, 110, 255) for column in range(width)] for row in range(height)]
    path.parent.mkdir(parents=True, exist_ok=True)
    write_png(path, width, height, rows)


def copy_sample_deck(sample_path: Path, directory: Path) -> Path:
    deck_path = directory / sample_path.name
    shutil.copytree(sample_path, deck_path)
    for image in referenced_images(deck_path):
        write_gradient_photo(deck_path / image)
    return deck_path


class SampleDeckCheckTest(unittest.TestCase):
    def test_every_sample_deck_passes_the_check_at_its_slide_count(self):
        for sample_path in sample_deck_paths():
            with self.subTest(sample_path.name), tempfile.TemporaryDirectory() as directory:
                deck_path = copy_sample_deck(sample_path, Path(directory))
                result = check_deck(CheckRequest(deck_path / "slides.html", slide_count(deck_path), ()))
                self.assertEqual([issue.kind.code for issue in result.issues if issue.kind.severity == "error"], [])


class SampleDeckBuildTest(unittest.TestCase):
    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_every_sample_deck_builds_an_acceptable_pdf_without_layout_defects(self):
        for sample_path in sample_deck_paths():
            with self.subTest(sample_path.name), tempfile.TemporaryDirectory() as directory:
                deck_path = copy_sample_deck(sample_path, Path(directory))
                count = slide_count(deck_path)
                completed = subprocess.run(
                    [sys.executable, str(OFFICE_ENTRY), "deck", "build", "--format", "all", "--slide-count", str(count)],
                    capture_output=True,
                    text=True,
                    cwd=deck_path,
                )
                envelope = json.loads(completed.stdout)
                if envelope["details"]["review"]["renderSource"] != "layout":
                    self.skipTest("the renderer could not run on this host")
                self.assertEqual({issue["code"] for issue in envelope["issues"]} & LAYOUT_DEFECT_CODES, set())
                self.assertTrue(envelope["details"]["acceptance"]["acceptable"], envelope["summary"])
                self.assertEqual(pdf_page_count(str(deck_path / "build" / f"{deck_path.name}.pdf")), count)
                self.assert_review_measured_every_page(deck_path, count)
                self.assert_pptx_keeps_the_layout(deck_path / "build" / f"{deck_path.name}.pptx")
                self.assert_kit_tables_are_native_tables(deck_path)

    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_a_donut_of_percentages_lists_each_share_once(self):
        with tempfile.TemporaryDirectory() as directory:
            deck_path = Path(directory) / "shares"
            deck_path.mkdir()
            (deck_path / "slides.html").write_text(PERCENT_DONUT_DECK, encoding="utf-8")
            subprocess.run([sys.executable, str(OFFICE_ENTRY), "deck", "build", "--format", "pptx"], capture_output=True, text=True, cwd=deck_path)
            pptx_path = deck_path / "build" / "shares.pptx"
            read = json.loads(subprocess.run([sys.executable, str(OFFICE_ENTRY), "deck", "read", str(pptx_path)], capture_output=True, text=True).stdout)
            texts = [shape.get("text", "") for slide in read["details"]["slides"] for shape in slide["shapes"]]
            self.assertEqual(texts.count("31%"), 1, texts)
            self.assert_pptx_keeps_the_layout(pptx_path)

    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_root_tokens_win_over_the_theme_the_body_names(self):
        with tempfile.TemporaryDirectory() as directory:
            deck_path = Path(directory) / "brand"
            deck_path.mkdir()
            (deck_path / "slides.html").write_text(BRAND_TOKEN_DECK, encoding="utf-8")
            subprocess.run([sys.executable, str(OFFICE_ENTRY), "deck", "build", "--format", "pptx"], capture_output=True, text=True, cwd=deck_path)
            layout = json.loads((deck_path / "build" / "review" / "pptx-layers" / "layout.json").read_text(encoding="utf-8"))
        chart = layout["slides"][1]["charts"][0]
        self.assertEqual(chart["colors"]["series"], ["rgb(228, 0, 43)"])
        self.assertEqual(chart["colors"]["background"], "rgb(255, 255, 255)")

    def assert_pptx_keeps_the_layout(self, pptx_path: Path):
        check = json.loads(subprocess.run([sys.executable, str(OFFICE_ENTRY), "deck", "check", str(pptx_path)], capture_output=True, text=True).stdout)
        self.assertEqual([issue["message"] for issue in check["issues"]], [])
        read = json.loads(subprocess.run([sys.executable, str(OFFICE_ENTRY), "deck", "read", str(pptx_path)], capture_output=True, text=True).stdout)
        texts = [shape.get("text", "") for slide in read["details"]["slides"] for shape in slide["shapes"]]
        for amount in legend_amounts(pptx_path.parent.parent):
            self.assertIn(amount, texts)
        self.assert_pdf_and_pptx_agree_on_text_sizes(pptx_path.with_suffix(".pdf"), read)

    def assert_kit_tables_are_native_tables(self, deck_path: Path):
        from pptx import Presentation

        source_rows = [table.count("<tr") for table in re.findall(r"<table>.*?</table>", (deck_path / "slides.html").read_text(encoding="utf-8"), re.DOTALL)]
        presentation = Presentation(str(deck_path / "build" / f"{deck_path.name}.pptx"))
        tables = [(slide, shape.table) for slide in presentation.slides for shape in slide.shapes if shape.has_table]
        self.assertEqual([len(table.rows) for _, table in tables], source_rows)
        for slide, table in tables:
            cell_texts = {cell.text for row in table.rows for cell in row.cells} - {""}
            loose_texts = {shape.text_frame.text for shape in slide.shapes if shape.has_text_frame}
            self.assertEqual(cell_texts & loose_texts, set())

    def assert_pdf_and_pptx_agree_on_text_sizes(self, pdf_path: Path, read: dict):
        import pdfplumber

        size = read["details"]["slideSize"]
        with pdfplumber.open(pdf_path) as pdf:
            for slide, page in zip(read["details"]["slides"], pdf.pages):
                points_per_emu = page.width / size["w"]
                words = page.extract_words(extra_attrs=["size"])
                for shape in slide["shapes"]:
                    if not shape.get("text") or not shape.get("style"):
                        continue
                    box = shape["box"]
                    inside = [word["size"] for word in words if box["x"] * points_per_emu <= (word["x0"] + word["x1"]) / 2 <= (box["x"] + box["w"]) * points_per_emu and box["y"] * points_per_emu <= (word["top"] + word["bottom"]) / 2 <= (box["y"] + box["h"]) * points_per_emu]
                    if inside:
                        drawn_points = max(inside) / (points_per_emu * size["emuPerPoint"])
                        self.assertAlmostEqual(drawn_points, shape["style"]["size"], delta=shape["style"]["size"] * 0.04, msg=f"slide {slide['slide']}: {shape['text'][:20]}")

    def assert_review_measured_every_page(self, deck_path: Path, count: int):
        review = json.loads((deck_path / "build" / "review" / "slide-review.json").read_text(encoding="utf-8"))
        self.assertTrue(all(slide["contentBounds"] for slide in review["slides"]))
        self.assertEqual(len(review["contactSheets"]), -(-count // 4))
        self.assertTrue(all((deck_path / "build" / "review" / sheet["filename"]).is_file() for sheet in review["contactSheets"]))


if __name__ == "__main__":
    unittest.main()
