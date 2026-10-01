import subprocess
import sys
import unittest

from sheet_fixture import OFFICE_ENTRY, WorkbookFixture, run_office, run_office_python, write_json

FIXTURE_WORKBOOK = """
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

workbook = Workbook()
sales = workbook.active
sales.title = "Sales"
for row in [["item", "amount"], ["A", 10], ["B", 20], ["C", 30]]:
    sales.append(row)
sales["B5"] = "=SUM(B2:B4)"
sales["D1"] = "=B5*2"
sales["D2"] = "=1/0"
sales["A1"].font = Font(bold=True, size=14)
sales["B2"].number_format = "#,##0"
sales["B2"].fill = PatternFill("solid", fgColor="FFEEAA")
sales["B2"].border = Border(*(Side(style="thin") for _ in range(4)))
sales.column_dimensions["A"].width = 22
sales["A2"].hyperlink = "https://example.com"
sales["A3"].comment = Comment("checked", "이샘플")
sales.conditional_formatting.add("B2:B4", CellIsRule(operator="greaterThan", formula=["15"], fill=PatternFill("solid", fgColor="FFC7CE")))
validation = DataValidation(type="list", formula1='"A,B,C"')
validation.add("A2:A4")
sales.add_data_validation(validation)
sales.freeze_panes = "A2"
sales.auto_filter.ref = "A1:B4"
sales.merge_cells("D3:E3")
sales.add_table(Table(displayName="SalesTable", ref="A1:B4"))
chart = BarChart()
chart.add_data(Reference(sales, min_col=2, min_row=1, max_row=4), titles_from_data=True)
sales.add_chart(chart, "G2")
workbook.create_sheet("Notes")["A1"] = "memo"
workbook.defined_names["Total"] = DefinedName("Total", attr_text="Sales!$B$5")
workbook.save("fixture.xlsx")
"""


class ReadTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        run_office_python(FIXTURE_WORKBOOK, self.directory)
        write_json(self.directory / "ops.json", [{"op": "recalculate"}])
        run_office(["sheet", "apply", "fixture.xlsx", "ops.json"], self.directory)

    def read(self, *arguments):
        envelope = run_office(["sheet", "read", "fixture.xlsx", *arguments], self.directory)
        self.assertEqual(envelope["status"], "ok", envelope)
        return envelope["details"]

    def test_each_sheet_reports_its_dimensions_and_features(self):
        sales = self.read()["sheets"][0]
        self.assertEqual(sales["name"], "Sales")
        self.assertEqual(sales["dimensions"], "A1:E5")
        self.assertEqual(sales["frozenPanes"], "A2")
        self.assertEqual(sales["autoFilter"], "A1:B4")
        self.assertEqual(sales["mergedCells"], ["D3:E3"])
        self.assertEqual(sales["tables"], [{"name": "SalesTable", "range": "A1:B4"}])
        self.assertEqual(sales["charts"], [{"chart": 0, "type": "bar", "title": None, "anchor": "G2", "data": ["'Sales'!$B$2:$B$4"]}])
        self.assertEqual(sales["features"], {"conditionalFormats": 1, "dataValidations": 1, "hyperlinks": 1, "comments": 1})

    def test_defined_names_are_listed_with_what_they_point_at(self):
        self.assertEqual(self.read()["definedNames"], [{"name": "Total", "value": "Sales!$B$5", "scope": None}])

    def test_a_range_returns_values_and_formulas_side_by_side(self):
        selected = self.read("--range", "A4:B5")["range"]
        self.assertEqual(selected["formulas"], [[None, None], [None, "=SUM(B2:B4)"]])
        self.assertEqual(selected["values"][0], ["C", 30])

    def test_where_formula_lists_only_formula_cells(self):
        cells = self.read("--where", "formula")["range"]["cells"]
        self.assertEqual([(cell["cell"], cell["formula"]) for cell in cells], [("D1", "=B5*2"), ("D2", "=1/0"), ("B5", "=SUM(B2:B4)")])

    def test_where_names_the_cells_of_one_kind(self):
        def cells_where(kind):
            return [cell["cell"] for cell in self.read("--range", "A1:D2", "--where", kind)["range"]["cells"]]
        self.assertEqual(cells_where("error"), ["D2"])
        self.assertEqual(cells_where("number"), ["D1", "B2"])
        self.assertEqual(cells_where("text"), ["A1", "B1", "A2"])
        self.assertEqual(cells_where("empty"), ["C1", "C2"])

    def test_stats_summarize_each_column_under_its_header(self):
        stats = self.read("--range", "A1:B5", "--stats")["range"]["stats"]
        self.assertEqual(stats[0], {"column": "A", "header": "item", "types": {"text": 3, "empty": 1}})
        self.assertEqual(stats[1], {"column": "B", "header": "amount", "types": {"number": 4}, "formulas": 1, "count": 4, "min": 10, "max": 60, "sum": 120, "mean": 30})

    def test_stats_skip_a_title_row_above_the_header(self):
        self.create_workbook([{"title": "S", "heading": "2026 실적", "rows": [["담당", "1월", "2월"], ["이샘플", 10, 20], ["박예시", 30, 40]]}])
        selected = run_office(["sheet", "read", "book.xlsx", "--stats"], self.directory)["details"]["range"]
        self.assertEqual(selected["headerRow"], 2)
        self.assertEqual([(column["header"], column.get("sum")) for column in selected["stats"]], [("담당", None), ("1월", 40), ("2월", 60)])

    def test_stats_count_a_formula_without_a_stored_value_apart(self):
        run_office_python(FIXTURE_WORKBOOK, self.directory)
        stats = self.read("--range", "B1:B5", "--stats")["range"]["stats"]
        self.assertEqual(stats[0]["types"], {"number": 3, "uncomputed": 1})

    def test_cols_narrow_the_range_to_the_named_columns(self):
        selected = self.read("--range", "A1:D3", "--cols", "A,C:D")["range"]
        self.assertEqual(selected["columns"], ["A", "C", "D"])
        self.assertEqual(selected["values"][0], ["item", None, 120])

    def test_formats_show_what_is_styled(self):
        formats = self.read("--range", "A1:B2", "--formats")["range"]
        by_cell = {cell["cell"]: cell for cell in formats["cells"]}
        self.assertEqual(by_cell["A1"]["font"], "bold 14pt")
        self.assertEqual({key: by_cell["B2"][key] for key in ("numberFormat", "fill", "border")}, {"numberFormat": "#,##0", "fill": "FFEEAA", "border": "thin all"})
        self.assertEqual(formats["columnWidths"], {"A": 22})

    def test_rows_beyond_the_limit_are_left_out_and_flagged(self):
        selected = self.read("--max-rows", "2")["range"]
        self.assertEqual((len(selected["values"]), selected["truncated"]), (2, True))

    def test_another_sheet_is_chosen_by_name(self):
        self.assertEqual(self.read("--sheet", "Notes")["range"]["values"], [["memo"]])

    def test_an_unknown_sheet_is_refused_with_the_names_that_exist(self):
        envelope = run_office(["sheet", "read", "fixture.xlsx", "--sheet", "Missing"], self.directory)
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["TARGET_NOT_FOUND"])
        self.assertIn("Sales, Notes", envelope["issues"][0]["message"])


class LibraryWarningTest(WorkbookFixture):
    def test_reading_a_workbook_with_sparklines_prints_only_the_result(self):
        self.create_workbook([{"title": "S", "rows": [["name", "1월", "2월", "3월", "trend"], ["a", 1, 2, 3, None], ["b", 3, 2, 1, None]]}])
        self.assertEqual(self.apply([{"op": "add_sparklines", "range": "B2:D3", "target": "E2:E3"}])["status"], "ok")
        for command in (["sheet", "read"], ["sheet", "check"], ["sheet", "validate"]):
            completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), *command, "book.xlsx"], capture_output=True, text=True, cwd=self.directory)
            self.assertEqual(completed.stderr, "", command)


if __name__ == "__main__":
    unittest.main()
