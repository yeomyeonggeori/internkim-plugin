from __future__ import annotations

import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

import openpyxl
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn

from png_fixture import read_png, write_png
from render_fixture import can_render
from free_deck_fixture import build_pptx, run_office_json, write_free_deck


SCRIPTS_PATH = Path(__file__).resolve().parents[1] / "skills" / "office" / "scripts"
sys.path.insert(0, str(SCRIPTS_PATH))

from deck.deck_kit import chart_types  # noqa: E402
from deck.pptx_export.editable import read_text_layers, write_editable_pptx  # noqa: E402
from deck.pptx_export.charts import NATIVE_CHART_TYPES  # noqa: E402
from powerpoint.preview.chart_look import chart_look, chart_model  # noqa: E402
from powerpoint.description import chart_details  # noqa: E402
from powerpoint.model.inheritance import slide_context  # noqa: E402
from charts.svg import chart_svg  # noqa: E402


ACCENT = "rgb(26, 86, 219)"
MUTED = "rgb(200, 210, 224)"
SECOND = "rgb(141, 176, 238)"
TEXT = {"fontFamily": "Paperlogy", "fontWeight": 600, "sizePx": 19, "color": "rgb(47, 61, 82)"}
STRONG = {"fontFamily": "Paperlogy", "fontWeight": 700, "sizePx": 21, "color": "rgb(12, 26, 48)"}
EXPECTED_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked100": XL_CHART_TYPE.COLUMN_STACKED_100,
    "area": XL_CHART_TYPE.AREA_STACKED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "donut": XL_CHART_TYPE.DOUGHNUT,
    "pie": XL_CHART_TYPE.PIE,
}


def chart_layout(kind: str, series: list[dict], **overrides) -> dict:
    labels = ["1분기", "2분기", "3분기"]
    is_round = kind in ("donut", "pie")
    series_colors = [ACCENT, SECOND][: len(series)]
    points = [[ACCENT, SECOND, MUTED]] if is_round else [[color] * len(labels) for color in series_colors]
    layout = {
        "type": kind, "labels": labels, "series": series, "units": ["억"] * len(series), "decimals": [0] * len(series), "startsAtZero": kind != "line",
        "valueRange": {"minimum": 50, "maximum": 150} if kind == "line" else None, "secondaryRange": None, "gapWidth": 100,
        "box": {"left": 100, "top": 200, "right": 1100, "bottom": 800},
        "colors": {"series": series_colors, "points": points, "grid": "rgb(213, 221, 232)", "background": "rgb(255, 255, 255)"},
        "text": {"base": TEXT, "category": TEXT, "legend": TEXT, "share": STRONG, "axisTitle": TEXT},
        "pointLabels": [] if is_round else [{"series": 0, "point": point, "position": "", "text": STRONG} for point in range(len(labels))],
    }
    return layout | overrides


def write_layers(review_path: Path, charts: list[dict]) -> None:
    layers_path = review_path / "pptx-layers"
    layers_path.mkdir(parents=True)
    slides = [{"width": 1600, "height": 900, "visibleText": "", "pictureTexts": [], "blocks": [], "shapes": [], "boxesKeptAsPicture": 0, "charts": [chart]} for chart in charts]
    (layers_path / "layout.json").write_text(json.dumps({"language": "ko", "slides": slides}), encoding="utf-8")
    for number in range(1, len(charts) + 1):
        write_png(layers_path / f"background.{number:03}.png", 32, 18, [[(255, 255, 255, 255)] * 32 for _ in range(18)])


def chart_frame(slide):
    return next(shape for shape in slide.shapes if shape.has_chart)


class NativeChartPackageTest(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)

    def write(self, charts: list[dict]):
        write_layers(self.directory / "review", charts)
        written = write_editable_pptx(read_text_layers(self.directory / "review", len(charts)), [""] * len(charts), self.directory / "deck.pptx")
        return written, Presentation(str(self.directory / "deck.pptx"))

    def test_every_kit_chart_type_has_a_native_chart_type(self):
        self.assertEqual(set(chart_types()), set(NATIVE_CHART_TYPES))

    def test_each_kit_chart_becomes_a_native_chart_with_its_categories_and_series(self):
        one = [{"name": "", "values": [96, 104.5, 128]}]
        two = [{"name": "프리미엄", "values": [38, 47, 61]}, {"name": "일반", "values": [52, 50, 49]}]
        charts = [chart_layout(kind, one if kind in ("donut", "pie", "column", "bar") else two) for kind in EXPECTED_TYPES]
        written, presentation = self.write(charts)
        self.assertEqual(written.chart_count, len(EXPECTED_TYPES))
        for slide, layout in zip(presentation.slides, charts):
            chart = chart_frame(slide).chart
            with self.subTest(layout["type"]):
                self.assertEqual(chart.chart_type, EXPECTED_TYPES[layout["type"]])
                self.assertEqual(list(chart.plots[0].categories), layout["labels"])
                self.assertEqual([list(series.values) for series in chart.plots[0].series], [[float(value) for value in series["values"]] for series in layout["series"]])

    def test_the_chart_sits_in_the_box_the_layout_measured(self):
        _, presentation = self.write([chart_layout("column", [{"name": "", "values": [1, 2, 3]}])])
        frame = chart_frame(presentation.slides[0])
        emu_per_pixel = 12192000 / 1600
        self.assertEqual((frame.left, frame.top, frame.width, frame.height), tuple(round(value * emu_per_pixel) for value in (100, 200, 1000, 600)))

    def test_the_embedded_workbook_holds_the_same_data(self):
        _, presentation = self.write([chart_layout("stacked", [{"name": "프리미엄", "values": [38, 47, 61]}, {"name": "일반", "values": [52, 50, 49.5]}])])
        workbook = openpyxl.load_workbook(io.BytesIO(chart_frame(presentation.slides[0]).chart.part.chart_workbook.xlsx_part.blob))
        rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
        self.assertEqual(rows, [[None, "프리미엄", "일반"], ["1분기", 38, 52], ["2분기", 47, 50], ["3분기", 61, 49.5]])

    def test_series_and_highlight_colors_and_the_kits_labels_are_kept(self):
        highlighted = chart_layout("column", [{"name": "", "values": [96, 104, 128]}], colors={"series": [ACCENT], "points": [[MUTED, MUTED, ACCENT]], "grid": "rgb(213, 221, 232)", "background": "rgb(255, 255, 255)"})
        _, presentation = self.write([highlighted])
        chart = chart_frame(presentation.slides[0]).chart
        series = chart.plots[0].series[0]._element
        self.assertEqual(series.find(f"{qn('c:spPr')}/{qn('a:solidFill')}/{qn('a:srgbClr')}").get("val"), "1A56DB")
        recolored = {int(point.find(qn("c:idx")).get("val")): point.find(f".//{qn('a:srgbClr')}").get("val") for point in series.findall(qn("c:dPt"))}
        self.assertEqual(recolored, {0: "C8D2E0", 1: "C8D2E0"})
        labels = series.findall(f"{qn('c:dLbls')}/{qn('c:dLbl')}")
        self.assertEqual([label.find(qn("c:numFmt")).get("formatCode") for label in labels], ['#,##0"억"'] * 3)
        self.assertEqual({label.find(qn("c:dLblPos")).get("val") for label in labels}, {"outEnd"})
        self.assertFalse(chart.value_axis.visible)

    def test_a_donut_is_only_its_ring_so_the_kit_legend_and_center_stay_text(self):
        _, presentation = self.write([chart_layout("donut", [{"name": "", "values": [42, 30, 28]}])])
        chart = chart_frame(presentation.slides[0]).chart
        self.assertEqual(chart._chartSpace.find(f".//{qn('c:holeSize')}").get("val"), "60")
        colors = [point.find(f".//{qn('a:solidFill')}/{qn('a:srgbClr')}").get("val") for point in chart.plots[0].series[0]._element.findall(qn("c:dPt"))]
        self.assertEqual(colors, ["1A56DB", "8DB0EE", "C8D2E0"])
        self.assertFalse(chart.has_legend)
        self.assertFalse(any(shape.has_text_frame for shape in presentation.slides[0].shapes))

    def test_the_preview_draws_the_charts_own_colors_and_only_its_labels(self):
        highlighted = chart_layout("column", [{"name": "", "values": [96, 104, 128]}], colors={"series": [ACCENT], "points": [[MUTED, MUTED, ACCENT]], "grid": "rgb(213, 221, 232)", "background": "rgb(255, 255, 255)"})
        highlighted["pointLabels"] = [{"series": 0, "point": 2, "position": "", "text": STRONG}]
        _, presentation = self.write([highlighted])
        slide = presentation.slides[0]
        chart = chart_frame(slide).chart
        svg = chart_svg(chart_model(chart, chart_details(chart)), 600, 400, chart_look(chart, slide_context(presentation, slide)), "Paperlogy")
        self.assertEqual((svg.count('fill="#1A56DB"'), svg.count('fill="#C8D2E0"')), (1, 2))
        self.assertIn(">128억<", svg)
        self.assertNotIn(">96억<", svg)
        self.assertEqual(svg.count("<line"), 1, "the category axis is drawn and the hidden value axis leaves no gridlines")

    def test_a_combo_chart_draws_its_last_series_as_a_line_on_a_second_axis(self):
        revenue, margin = {"name": "매출", "values": [96, 104, 128]}, {"name": "이익률", "values": [11.2, 12.5, 14.2]}
        combo = chart_layout("combo", [revenue, margin], units=["억", "%"], decimals=[0, 1], valueRange={"minimum": 0, "maximum": 213}, secondaryRange={"minimum": 0, "maximum": 16, "step": 4})
        _, presentation = self.write([combo])
        chart = chart_frame(presentation.slides[0]).chart
        self.assertEqual([plot.__class__.__name__ for plot in chart.plots], ["BarPlot", "LinePlot"])
        self.assertEqual([[series.name for series in plot.series] for plot in chart.plots], [["매출"], ["이익률"]])
        axes = chart._chartSpace.findall(f".//{qn('c:valAx')}")
        self.assertEqual([axis.find(f"{qn('c:scaling')}/{qn('c:max')}").get("val") for axis in axes], ["213", "16"])
        self.assertEqual(axes[1].find(qn("c:numFmt")).get("formatCode"), '#,##0"%"', "ticks every whole step need no decimals")

    def test_a_scatter_chart_plots_the_first_series_across_and_names_each_point(self):
        stores = ["강남점", "판교점", "부산점"]
        scatter = chart_layout("scatter", [{"name": "매출", "values": [120, 95, 70]}, {"name": "이익률", "values": [14.5, 12.1, 9.8]}], labels=stores, units=["억", "%"], decimals=[0, 1],
                               valueRange={"minimum": 50, "maximum": 140}, secondaryRange={"minimum": 8, "maximum": 16},
                               colors={"series": [ACCENT, SECOND], "points": [[ACCENT, MUTED, MUTED]], "grid": "rgb(213, 221, 232)", "background": "rgb(255, 255, 255)"})
        _, presentation = self.write([scatter])
        slide = presentation.slides[0]
        chart = chart_frame(slide).chart
        self.assertEqual(chart.chart_type, XL_CHART_TYPE.XY_SCATTER)
        details = chart_details(chart)
        self.assertEqual((details["series"][0]["x"], details["series"][0]["values"]), ([120.0, 95.0, 70.0], [14.5, 12.1, 9.8]))
        names = [label.findtext(f".//{qn('a:t')}") for label in chart._chartSpace.iter(qn("c:dLbl"))]
        self.assertEqual(names, stores)
        self.assertFalse(chart.has_legend)
        svg = chart_svg(chart_model(chart, details), 600, 400, chart_look(chart, slide_context(presentation, slide)), "Paperlogy")
        self.assertTrue(all(f">{store}<" in svg for store in stores))
        self.assertEqual(svg.count('fill="#C8D2E0"'), 2)

    def test_the_chart_fonts_are_embedded(self):
        written, _ = self.write([chart_layout("bar", [{"name": "", "values": [1, 2, 3]}])])
        self.assertEqual(written.embedded_typefaces, ("Paperlogy 6 SemiBold", "Paperlogy 7 Bold"))


def figure(kind: str, labels: str, series: str, unit: str, extra: str = "", caption: str = "") -> str:
    attribute = "data-series" if ";" in series else "data-values"
    return f'<figure data-chart="{kind}" data-labels="{labels}" {attribute}="{series}" data-unit="{unit}" {extra}><figcaption>{caption}</figcaption></figure>'


NEW_CHART_SECTIONS = [
    "<h2>매출과 이익률이 함께 올랐습니다</h2>" + figure("combo", "1Q, 2Q, 3Q, 4Q", "매출: 96, 104, 113, 128; 영업이익률: 11.2, 12.5, 13.1, 14.2", "억, %", caption="분기 매출과 영업이익률"),
    "<h2>클라우드 매출이 해마다 커졌습니다</h2>" + figure("area", "2022, 2023, 2024, 2025, 2026", "클라우드: 20, 32, 45, 60, 78; 온프레미스: 60, 58, 55, 50, 44", "억", caption="사업별 매출, 단위 억 원"),
    "<h2>매출이 큰 매장일수록 이익률도 높습니다</h2>" + figure("scatter", "강남점, 판교점, 부산점, 대구점, 광주점", "매출: 120, 95, 70, 52, 40; 이익률: 14.5, 12.1, 9.8, 8.2, 6.5", "억, %", 'data-highlight="강남점"', "매장별 연 매출과 이익률"),
    "<h2>프리미엄 비중이 분기마다 늘었습니다</h2>" + figure("stacked100", "1Q, 2Q, 3Q, 4Q", "프리미엄: 38, 45, 52, 58; 일반: 52, 48, 44, 40; 기타: 10, 7, 4, 2", "%", caption="제품군별 매출 비중"),
]

PROPOSAL_SECTIONS = [
    "<h2>직원이 가장 많이 쓰는 시간은 재고 확인입니다</h2>" + figure("bar", "재고 확인, 발주 작성, 매출 정산, 직원 일정, 세금 서류", "9.5, 6.2, 4.8, 3.1, 2.4", "시간", caption="주간 업무 시간"),
    "<h2>도입 매장의 지표가 비교 매장보다 앞섭니다</h2>" + figure("line", "1월, 2월, 3월, 4월", "도입 매장: 80, 86, 93, 101; 비교 매장: 79, 81, 82, 84", "점", caption="월별 지표"),
    "<h2>도입 의향은 절반을 넘었습니다</h2>" + figure("donut", "있음, 검토 중, 없음", "52, 30, 18", "%", caption="점주 응답"),
]


class BuiltNativeChartTest(unittest.TestCase):
    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_combo_area_scatter_and_percent_charts_build_clean_into_native_charts(self):
        with tempfile.TemporaryDirectory() as directory:
            deck_path = write_free_deck(Path(directory) / "charts", NEW_CHART_SECTIONS)
            built = build_pptx(deck_path)
            pptx_path = Path(built["details"]["outputs"]["pptx"])
            read = run_office_json(["read", str(pptx_path)], deck_path)
            checked = run_office_json(["check", str(pptx_path)], deck_path)
        self.assertTrue(built["details"]["acceptance"]["acceptable"], built["summary"])
        charts = [shape["chart"] for slide in read["details"]["slides"] for shape in slide["shapes"] if shape["kind"] == "chart"]
        self.assertEqual([chart["type"] for chart in charts], ["column_clustered+line_markers", "area_stacked", "xy_scatter", "column_stacked_100"])
        self.assertEqual(charts[2]["series"][0]["x"], [120.0, 95.0, 70.0, 52.0, 40.0])
        self.assertEqual([issue["message"] for issue in checked["issues"]], [])
        self.assertTrue(checked["details"]["seen"])

    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_a_built_deck_carries_native_charts_read_back_and_previewed(self):
        with tempfile.TemporaryDirectory() as directory:
            deck_path = write_free_deck(Path(directory) / "proposal", PROPOSAL_SECTIONS)
            built = build_pptx(deck_path)
            pptx_path = Path(built["details"]["outputs"]["pptx"])
            read = run_office_json(["read", str(pptx_path)], deck_path)
            checked = run_office_json(["check", str(pptx_path)], deck_path)
            layout = json.loads((deck_path / "build" / "review" / "pptx-layers" / "layout.json").read_text(encoding="utf-8"))
            backgrounds = {number: read_png(deck_path / "build" / "review" / "pptx-layers" / f"background.{number:03}.png") for number, slide in enumerate(layout["slides"], start=1) if slide.get("charts")}
            preview = (deck_path / checked["details"]["preview"]).read_text(encoding="utf-8")
        self.assertEqual(built["details"]["pptx"]["charts"], 3)
        charts = [shape["chart"] for slide in read["details"]["slides"] for shape in slide["shapes"] if shape["kind"] == "chart"]
        self.assertEqual([chart["type"] for chart in charts], ["bar_clustered", "line_markers", "doughnut"])
        self.assertEqual(charts[0]["categories"], ["재고 확인", "발주 작성", "매출 정산", "직원 일정", "세금 서류"])
        self.assertEqual(charts[1]["series"][1]["name"], "비교 매장")
        self.assert_chart_left_the_background(layout, backgrounds)
        self.assertNotIn("PPTX_NOT_RENDERED", {issue["code"] for issue in checked["issues"]})
        self.assertIn("svg", preview)

    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_a_unit_that_is_a_word_stands_apart_from_its_number_and_a_symbol_does_not(self):
        sections = [
            "<h2>Paying cafes grew every quarter</h2>" + figure("column", "Q1, Q2, Q3", "3100, 3500, 4200", "cafés"),
            "<h2>Churn fell every quarter</h2>" + figure("column", "Q1, Q2, Q3", "4.1, 3.6, 3.1", "%"),
        ]
        with tempfile.TemporaryDirectory() as directory:
            deck_path = write_free_deck(Path(directory) / "units", sections)
            built = build_pptx(deck_path)
            presentation = Presentation(built["details"]["outputs"]["pptx"])
            formats = [{label.find(qn("c:numFmt")).get("formatCode") for label in chart_frame(slide).chart.plots[0].series[0]._element.iter(qn("c:dLbl"))} for slide in presentation.slides]
        self.assertEqual(formats, [{'#,##0" cafés"'}, {'#,##0.0"%"'}])

    def assert_chart_left_the_background(self, layout: dict, backgrounds: dict) -> None:
        for number, background in backgrounds.items():
            chart = layout["slides"][number - 1]["charts"][0]
            box = chart["box"]
            scale = background["width"] / layout["slides"][number - 1]["width"]
            columns = range(round(box["left"] * scale), round(box["right"] * scale), 7)
            rows = range(round(box["top"] * scale), round(box["bottom"] * scale), 7)
            colors = {background["rows"][row][column][:3] for row in rows for column in columns}
            self.assertEqual(len(colors), 1, f"slide {number} still draws its chart in the background")


if __name__ == "__main__":
    unittest.main()
