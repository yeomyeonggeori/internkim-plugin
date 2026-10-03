import json
from pathlib import Path
import subprocess
import sys
import tempfile
import textwrap
import unittest

from doc_fixture import OFFICE_ENTRY, block_texts, run_office, write_json


FORM_BUILDER = """
import json, sys
from docx import Document

shape = json.loads(sys.argv[1])
document = Document()
document.add_paragraph(shape["title"])
header = shape["header"]
item_rows = shape["itemRows"]
summaries = shape["summaries"]
table = document.add_table(rows=1 + item_rows + len(summaries), cols=len(header))
table.style = "Table Grid"
for cell, text in zip(table.rows[0].cells, header):
    cell.text = text
for number in range(1, item_rows + 1):
    table.rows[number].cells[0].text = str(number)
label_span = shape["labelSpan"]
for offset, label in enumerate(summaries):
    row = table.rows[1 + item_rows + offset]
    merged = row.cells[0].merge(row.cells[label_span - 1]) if label_span > 1 else row.cells[0]
    merged.text = label
document.save(shape["name"])
"""


def build_form(directory: Path, shape: dict) -> None:
    script = directory / "build_form.py"
    script.write_text(textwrap.dedent(FORM_BUILDER), encoding="utf-8")
    subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", str(script), json.dumps(shape, ensure_ascii=False)], check=True, cwd=directory)


FORM_SHAPES = [
    {"name": "purchase.docx", "title": "구매 요청", "header": ["No", "품명", "수량", "단가", "금액"], "itemRows": 10, "summaries": ["소계", "부가세", "합계"], "labelSpan": 4},
    {"name": "invoice.docx", "title": "Invoice", "header": ["#", "Description", "Hours", "Rate", "Amount"], "itemRows": 3, "summaries": ["Subtotal", "Tax", "Total due"], "labelSpan": 1},
    {"name": "settlement.docx", "title": "출장비 정산", "header": ["순번", "일자", "내역", "금액"], "itemRows": 1, "summaries": ["합계"], "labelSpan": 3},
    {"name": "long.docx", "title": "Expense report", "header": ["No", "Date", "Vendor", "Category", "Note", "Amount"], "itemRows": 25, "summaries": ["Net", "VAT", "Grand total"], "labelSpan": 5},
]


class LabelAddressedCellTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)

    def tearDown(self):
        self.temporary_directory.cleanup()

    def table_rows(self, name: str) -> list[list[str]]:
        return next(cells for kind, cells in block_texts(self.directory, name) if kind == "table")

    def test_a_summary_row_receives_the_value_its_label_names_whatever_the_number_of_item_rows(self):
        for shape in FORM_SHAPES:
            with self.subTest(form=shape["name"]):
                build_form(self.directory, shape)
                amount_column = shape["header"][-1]
                operations = [
                    {"op": "set_cell", "block": 1, "rowLabel": label, "columnLabel": amount_column, "text": f"value-{offset}"}
                    for offset, label in enumerate(shape["summaries"])
                ]
                operations.append({"op": "set_cell", "block": 1, "rowLabel": "1", "columnLabel": shape["header"][1], "text": "first item"})
                write_json(self.directory / "ops.json", operations)
                envelope = run_office(["apply", shape["name"], "ops.json"], self.directory)
                self.assertEqual(envelope["status"], "ok", envelope)
                rows = self.table_rows(shape["name"])
                summary_rows = rows[1 + shape["itemRows"]:]
                self.assertEqual([row[-1] for row in summary_rows], [f"value-{offset}" for offset in range(len(shape["summaries"]))])
                self.assertEqual(rows[1][1], "first item")
                self.assertTrue(all(row[-1] == "" for row in rows[1:1 + shape["itemRows"]]))

    def test_a_label_written_in_several_rows_is_refused_and_names_them(self):
        build_form(self.directory, FORM_SHAPES[1])
        write_json(self.directory / "ops.json", [{"op": "set_cell", "block": 1, "row": 2, "column": 1, "text": "Total due"}])
        run_office(["apply", "invoice.docx", "ops.json"], self.directory)
        write_json(self.directory / "ops.json", [{"op": "set_cell", "block": 1, "rowLabel": "Total due", "column": 4, "text": "x"}])
        envelope = run_office(["apply", "invoice.docx", "ops.json"], self.directory)
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["TARGET_NOT_FOUND"])
        self.assertIn("rows 2 and 6", envelope["issues"][0]["message"])

    def test_a_label_no_cell_holds_is_refused_and_lists_the_labels_there_are(self):
        build_form(self.directory, FORM_SHAPES[0])
        write_json(self.directory / "ops.json", [{"op": "set_cell", "block": 1, "rowLabel": "총계", "column": 4, "text": "x"}])
        envelope = run_office(["apply", "purchase.docx", "ops.json"], self.directory)
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["TARGET_NOT_FOUND"])
        self.assertEqual(envelope["issues"][0]["location"], "ops[0].rowLabel")
        self.assertIn("합계", envelope["issues"][0]["suggestion"])

    def test_a_cell_is_named_by_exactly_one_index_or_label_on_each_axis(self):
        build_form(self.directory, FORM_SHAPES[2])
        for operation, location in [
            ({"op": "set_cell", "block": 1, "row": 2, "rowLabel": "합계", "column": 3, "text": "x"}, "ops[0]"),
            ({"op": "set_cell", "block": 1, "column": 3, "text": "x"}, "ops[0]"),
            ({"op": "set_cell", "block": 1, "row": 2, "text": "x"}, "ops[0]"),
        ]:
            with self.subTest(operation=operation):
                write_json(self.directory / "ops.json", [operation])
                envelope = run_office(["apply", "settlement.docx", "ops.json"], self.directory)
                self.assertEqual(envelope["status"], "error")
                self.assertEqual(envelope["issues"][0]["location"], location)


if __name__ == "__main__":
    unittest.main()


SLIDE_TABLE_BUILDER = """
from pptx import Presentation
from pptx.util import Inches

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
rows = [["Quarter", "Region", "Revenue"], ["Q1", "North", ""], ["Q2", "North", ""], ["Q3", "South", ""], ["Q4", "South", ""], ["Year total", "", ""]]
table = slide.shapes.add_table(len(rows), 3, Inches(1), Inches(1), Inches(6), Inches(3)).table
for row, values in zip(table.rows, rows):
    for cell, value in zip(row.cells, values):
        cell.text = value
presentation.save("slides.pptx")
"""


class SlideTableLabelTest(unittest.TestCase):
    def test_a_slide_table_cell_is_named_by_its_row_and_column_labels(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            (directory / "build.py").write_text(textwrap.dedent(SLIDE_TABLE_BUILDER), encoding="utf-8")
            subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "build.py"], check=True, cwd=directory)
            write_json(directory / "ops.json", [
                {"op": "set_table_cell", "slide": 1, "shape": 0, "rowLabel": "Year total", "columnLabel": "Revenue", "text": "4,200"},
                {"op": "set_table_cell", "slide": 1, "shape": 0, "rowLabel": "Q3", "column": 2, "text": "1,100"},
            ])
            envelope = run_office(["apply", "slides.pptx", "ops.json"], directory)
            self.assertEqual(envelope["status"], "ok", envelope)
            write_json(directory / "ops.json", [{"op": "set_table_cell", "slide": 1, "shape": 0, "rowLabel": "North", "column": 2, "text": "x"}])
            refused = run_office(["apply", "slides.pptx", "ops.json"], directory)
            self.assertEqual([issue["code"] for issue in refused["issues"]], ["TARGET_NOT_FOUND"])
            rows = next(shape for shape in run_office(["read", "slides.pptx"], directory)["details"]["slides"][0]["shapes"] if "rows" in shape)["rows"]
            self.assertEqual([row[2] for row in rows], ["Revenue", "", "", "1,100", "", "4,200"])
