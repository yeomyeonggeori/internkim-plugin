import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from doc_fixture import ContractFixture, run_office, write_json


class DecorationFixture(ContractFixture):
    def apply(self, operations, *flags):
        write_json(self.directory / "ops.json", operations)
        return run_office(["doc", "apply", "contract.docx", "ops.json", *flags], self.directory)

    def document(self):
        return Document(str(self.directory / "contract.docx"))


class BaselineTest(DecorationFixture):
    def test_found_text_is_raised_or_lowered_and_can_be_put_back(self):
        envelope = self.apply([{"op": "format_text", "block": 5, "find": "10,000,000", "baseline": "superscript"}, {"op": "format_text", "block": 7, "find": "30일", "baseline": "subscript"}])
        self.assertEqual(envelope["status"], "ok", envelope)
        paragraphs = self.document().paragraphs
        raised = [run for run in paragraphs[5].runs if run.font.superscript]
        lowered = [run for run in paragraphs[7].runs if run.font.subscript]
        self.assertEqual(([run.text for run in raised], [run.text for run in lowered]), (["10,000,000"], ["30일"]))
        self.apply([{"op": "format_text", "block": 5, "find": "10,000,000", "baseline": "none"}])
        self.assertFalse(any(run.font.superscript for run in self.document().paragraphs[5].runs))


class ParagraphDecorationTest(DecorationFixture):
    def properties(self, index):
        return self.document().paragraphs[index]._p.pPr

    def test_shading_borders_and_right_indent_are_written_in_schema_order(self):
        envelope = self.apply([{"op": "set_paragraph_format", "block": 7, "align": "justify", "indentRightInches": 0.5, "shadingFill": "F2F2F2", "borders": ["top", "bottom"], "borderColor": "1F4E79"}])
        self.assertEqual(envelope["status"], "ok", envelope)
        properties = self.properties(7)
        self.assertEqual(properties.find(qn("w:shd")).get(qn("w:fill")), "F2F2F2")
        lines = properties.find(qn("w:pBdr"))
        self.assertEqual([(line.tag.split("}")[1], line.get(qn("w:color"))) for line in lines], [("top", "1F4E79"), ("bottom", "1F4E79")])
        self.assertEqual(self.document().paragraphs[7].paragraph_format.right_indent, Inches(0.5))
        tags = [child.tag.split("}")[1] for child in properties]
        self.assertLess(tags.index("pBdr"), tags.index("shd"))
        self.assertLess(tags.index("shd"), tags.index("jc"))

    def test_none_and_an_empty_list_take_the_decoration_away(self):
        self.apply([{"op": "set_paragraph_format", "block": 7, "shadingFill": "F2F2F2", "borders": ["left"]}])
        self.apply([{"op": "set_paragraph_format", "block": 7, "shadingFill": "none", "borders": []}])
        properties = self.properties(7)
        self.assertIsNone(properties.find(qn("w:shd")))
        self.assertIsNone(properties.find(qn("w:pBdr")))

    def test_a_color_alone_recolors_the_lines_already_there(self):
        self.apply([{"op": "set_paragraph_format", "block": 7, "borders": ["bottom"]}])
        self.apply([{"op": "set_paragraph_format", "block": 7, "borderColor": "C00000"}])
        self.assertEqual(self.properties(7).find(qn("w:pBdr")).find(qn("w:bottom")).get(qn("w:color")), "C00000")

    def test_tracked_decoration_records_the_old_properties(self):
        self.apply([{"op": "set_paragraph_format", "block": 7, "shadingFill": "FFF2CC"}], "--track")
        self.assertIsNotNone(self.properties(7).find(qn("w:pPrChange")))


class StyleDecorationTest(DecorationFixture):
    def test_a_style_carries_baseline_shading_borders_and_right_indent(self):
        envelope = self.apply([
            {"op": "define_style", "name": "Callout", "basedOn": "Normal", "shadingFill": "DEEAF6", "borders": ["left"], "borderColor": "2E75B6", "indentRightInches": 0.25},
            {"op": "define_style", "name": "Footnote Mark", "type": "character", "baseline": "superscript"},
        ])
        self.assertEqual(envelope["status"], "ok", envelope)
        styles = self.document().styles
        callout = styles["Callout"].element.pPr
        self.assertEqual(callout.find(qn("w:shd")).get(qn("w:fill")), "DEEAF6")
        self.assertEqual(callout.find(qn("w:pBdr")).find(qn("w:left")).get(qn("w:color")), "2E75B6")
        self.assertEqual(styles["Callout"].paragraph_format.right_indent, Inches(0.25))
        self.assertTrue(styles["Footnote Mark"].font.superscript)


if __name__ == "__main__":
    unittest.main()
