import re
import sys
import unittest
import zipfile

from openpyxl import load_workbook

from sheet_fixture import SCRIPTS_PATH, WorkbookFixture

sys.path.insert(0, str(SCRIPTS_PATH))

from sheet.sheet_definitions import CHART_COLUMN_LEFT_OUT  # noqa: E402

SALES = [["region", "product", "qty", "amount"]] + [
    [region, product, quantity, quantity * 1000]
    for region, product, quantity in [
        ("Seoul", "A", 3), ("Busan", "A", 5), ("Seoul", "B", 2), ("Daegu", "A", 7), ("Busan", "B", 1), ("Seoul", "A", 4),
    ]
]


def archive_text(path, part):
    with zipfile.ZipFile(path) as archive:
        return archive.read(part).decode()


def archive_names(path):
    with zipfile.ZipFile(path) as archive:
        return archive.namelist()


class OperationFixture(WorkbookFixture):
    def setUp(self):
        super().setUp()
        self.create_workbook([{"title": "Sales", "rows": SALES}, {"title": "Notes", "rows": [["memo"], ["first"]]}])

    def edit(self, operations, *extra):
        envelope = self.apply(operations, *extra)
        self.assertNotEqual(envelope["status"], "error", envelope)
        return envelope

    def refused(self, operations):
        envelope = self.apply(operations)
        self.assertEqual(envelope["status"], "error", envelope)
        return envelope["issues"][0]

    def workbook(self):
        return load_workbook(self.directory / "book.xlsx")

    def sheet_xml(self, number=1):
        return archive_text(self.directory / "book.xlsx", f"xl/worksheets/sheet{number}.xml")


class FormattingTest(OperationFixture):
    def test_fonts_alignment_and_borders(self):
        self.edit([{"op": "format_range", "range": "A1:D2", "italic": True, "underline": True, "fontSize": 13, "fontName": "Malgun Gothic", "verticalAlignment": "center", "indent": 1, "border": "outline", "borderStyle": "medium", "borderColor": "1F2937"}])
        sheet = self.workbook()["Sales"]
        first, last = sheet["A1"], sheet["D2"]
        self.assertEqual((first.font.i, first.font.u, first.font.sz, first.font.name), (True, "single", 13, "Malgun Gothic"))
        self.assertEqual((first.alignment.vertical, first.alignment.indent), ("center", 1))
        self.assertEqual((first.border.top.style, first.border.left.style, first.border.bottom.style), ("medium", "medium", "thin"))
        self.assertEqual((last.border.bottom.style, last.border.right.style, last.border.bottom.color.rgb), ("medium", "medium", "001F2937"))

    def test_merge_unmerge_row_height_and_hiding(self):
        self.edit([
            {"op": "merge_cells", "range": "F1:G1"},
            {"op": "set_row_height", "row": 2, "height": 24, "count": 2},
            {"op": "hide_rows", "at": 5},
            {"op": "hide_columns", "at": "C", "count": 2},
            {"op": "hide_sheet", "sheet": "Notes"},
        ])
        workbook = self.workbook()
        sheet = workbook["Sales"]
        self.assertEqual([str(merged) for merged in sheet.merged_cells.ranges], ["F1:G1"])
        self.assertEqual((sheet.row_dimensions[3].height, sheet.row_dimensions[5].hidden), (24, True))
        self.assertTrue(sheet.column_dimensions["D"].hidden)
        self.assertEqual(workbook["Notes"].sheet_state, "hidden")
        self.edit([{"op": "unmerge_cells", "range": "A1:H1"}, {"op": "hide_rows", "at": 5, "hidden": False}])
        sheet = self.workbook()["Sales"]
        self.assertEqual((list(sheet.merged_cells.ranges), sheet.row_dimensions[5].hidden), ([], False))

    def test_the_last_visible_sheet_cannot_be_hidden(self):
        issue = self.refused([{"op": "hide_sheet", "sheet": "Notes"}, {"op": "hide_sheet", "sheet": "Sales"}])
        self.assertEqual(issue["code"], "OPERATION_NOT_APPLICABLE")


class RuleTest(OperationFixture):
    def test_every_conditional_format_rule_is_written(self):
        rules = [
            {"rule": "greater_than", "value": 3}, {"rule": "between", "value": 2, "value2": 4}, {"rule": "contains_text", "value": "Seo"},
            {"rule": "duplicate"}, {"rule": "top", "rank": 2}, {"rule": "below_average"}, {"rule": "formula", "formula": "=$C2>4", "fill": "C6EFCE", "fontColor": "006100"},
            {"rule": "color_scale"}, {"rule": "data_bar"}, {"rule": "icon_set", "icons": "3_arrows"},
        ]
        self.edit([{"op": "add_conditional_format", "range": "C2:C7", **rule} for rule in rules])
        sheet = self.sheet_xml()
        for kind in ('type="cellIs"', 'operator="between"', 'type="containsText"', 'type="duplicateValues"', 'type="top10"', 'aboveAverage="0"', 'type="expression"', "<colorScale>", "<dataBar>", 'iconSet="3Arrows"'):
            self.assertIn(kind, sheet)
        self.assertIn("<formula>$C2&gt;4</formula>", sheet)
        self.edit([{"op": "clear_conditional_formats", "range": "C5"}])
        self.assertNotIn("conditionalFormatting", self.sheet_xml())

    def test_data_validation_kinds(self):
        self.edit([
            {"op": "add_data_validation", "range": "A2:A50", "type": "list", "values": ["Seoul", "Busan"], "prompt": "pick a region"},
            {"op": "add_data_validation", "range": "B2:B50", "type": "list", "source": "=Notes!$A$1:$A$2"},
            {"op": "add_data_validation", "range": "C2:C50", "type": "whole", "minimum": 0, "maximum": 100, "error": "0 to 100"},
            {"op": "add_data_validation", "range": "E2:E50", "type": "date", "minimum": "2026-01-01"},
            {"op": "add_data_validation", "range": "F2:F50", "type": "custom", "formula": "=LEN(F2)<=10"},
        ])
        sheet = self.sheet_xml()
        self.assertIn("<formula1>\"Seoul,Busan\"</formula1>", sheet)
        self.assertIn("<formula1>Notes!$A$1:$A$2</formula1>", sheet)
        self.assertIn('operator="between"', sheet)
        self.assertIn("<formula1>DATE(2026,1,1)</formula1>", sheet)
        self.assertIn('operator="greaterThanOrEqual"', sheet)
        self.edit([{"op": "clear_data_validations", "range": "A1:B2"}])
        self.assertNotIn("Seoul,Busan", self.sheet_xml())

    def test_a_list_choice_with_a_comma_is_refused(self):
        issue = self.refused([{"op": "add_data_validation", "range": "A2", "type": "list", "values": ["a,b"]}])
        self.assertIn("source", issue["message"])


class ObjectTest(OperationFixture):
    def test_a_table_takes_its_style_and_replaces_the_sheet_filter(self):
        self.edit([{"op": "add_table", "range": "A1:D7", "name": "Sales_Table"}])
        sheet = self.workbook()["Sales"]
        table = sheet.tables["Sales_Table"]
        self.assertEqual((table.ref, table.tableStyleInfo.name, sheet.auto_filter.ref), ("A1:D7", "TableStyleMedium2", None))
        self.assertIsNone(sheet["A1"].fill.fill_type)

    def test_a_table_needs_distinct_headers(self):
        issue = self.refused([{"op": "set_cell", "cell": "C1", "value": "Region"}, {"op": "add_table", "range": "A1:D7"}])
        self.assertEqual(issue["code"], "OPERATION_NOT_APPLICABLE")

    def test_links_notes_and_page_setup(self):
        self.edit([
            {"op": "set_hyperlink", "cell": "F1", "url": "https://example.com", "text": "site"},
            {"op": "set_hyperlink", "cell": "F2", "url": "#Notes!A1"},
            {"op": "set_comment", "cell": "D2", "text": "checked", "author": "이샘플"},
            {"op": "set_page_setup", "orientation": "landscape", "paperSize": "A4", "fitToWidth": True, "printTitleRows": "1:1", "margins": "narrow", "pageNumbers": True},
        ])
        sheet = self.workbook()["Sales"]
        self.assertEqual((sheet["F1"].value, sheet["F1"].hyperlink.target), ("site", "https://example.com"))
        self.assertEqual(sheet["F2"].hyperlink.location, "Notes!A1")
        self.assertEqual((sheet["D2"].comment.text, sheet["D2"].comment.author), ("checked", "이샘플"))
        self.assertEqual((sheet.page_setup.orientation, sheet.page_setup.paperSize, sheet.page_setup.fitToWidth), ("landscape", 9, 1))
        self.assertEqual((sheet.print_title_rows, sheet.page_margins.left, sheet.oddFooter.center.text), ("$1:$1", 0.25, "&P / &N"))
        self.edit([{"op": "set_comment", "cell": "D2", "text": ""}])
        self.assertIsNone(self.workbook()["Sales"]["D2"].comment)


class DataTest(OperationFixture):
    def test_sort_moves_whole_rows_and_their_formulas(self):
        self.edit([{"op": "set_range", "cell": "E1", "values": [["double"]] + [[f"=D{row}*2"] for row in range(2, 8)]}])
        self.edit([{"op": "sort_range", "range": "A1:E7", "by": "D", "descending": True}])
        sheet = self.workbook()["Sales"]
        self.assertEqual([sheet[f"A{row}"].value for row in range(1, 8)], ["region", "Daegu", "Busan", "Seoul", "Seoul", "Seoul", "Busan"])
        self.assertEqual((sheet["E2"].value, sheet["D2"].value), ("=D2*2", 7000))

    def test_sort_by_a_formula_column_uses_its_values_and_breaks_ties(self):
        self.edit([{"op": "set_range", "cell": "E1", "values": [["key"]] + [[f"=-C{row}"] for row in range(2, 8)]}])
        self.edit([{"op": "sort_range", "range": "A1:E7", "by": "B", "thenBy": "E"}])
        sheet = self.workbook()["Sales"]
        self.assertEqual([sheet[f"C{row}"].value for row in range(2, 8)], [7, 5, 4, 3, 2, 1])

    def test_find_replace_in_values_or_formulas(self):
        self.edit([{"op": "set_cell", "cell": "F1", "value": '=COUNTIF(A2:A7,"Seoul")'}])
        self.edit([{"op": "find_replace", "find": "seoul", "replace": "서울"}])
        sheet = self.workbook()["Sales"]
        self.assertEqual((sheet["A2"].value, sheet["F1"].value), ("서울", '=COUNTIF(A2:A7,"Seoul")'))
        self.edit([{"op": "find_replace", "sheet": "Sales", "find": '"Seoul"', "replace": '"서울"', "in": "formulas"}])
        self.assertEqual(self.cells()["F1"]["value"], "3")


class ChartTest(OperationFixture):
    def chart_xml(self, number):
        return archive_text(self.directory / "book.xlsx", f"xl/charts/chart{number}.xml")

    def test_every_chart_kind(self):
        kinds = ["area", "doughnut", "scatter", "radar", "combo"]
        self.edit([{"op": "add_chart", "type": kind, "range": "B1:D4" if kind == "combo" else "A1:C4", "anchor": f"H{index * 20 + 1}"} for index, kind in enumerate(kinds)])
        for number, element in enumerate(("<areaChart>", "<doughnutChart>", "<scatterChart>", "<radarChart>", "<barChart>"), 1):
            self.assertIn(element, self.chart_xml(number))
        combo = self.chart_xml(5)
        self.assertIn("<lineChart>", combo)
        self.assertEqual(combo.count("<valAx>"), 2)
        self.assertIn('<crosses val="max"/>', combo)

    def test_edit_and_delete_a_chart(self):
        self.edit([{"op": "add_chart", "type": "bar", "range": "A1:C4", "title": "Before"}, {"op": "add_chart", "type": "line", "range": "A1:C4", "anchor": "H30"}])
        self.edit([{"op": "edit_chart", "chart": 0, "title": "After", "yTitle": "Amount", "legend": "none"}, {"op": "delete_chart", "chart": 1}])
        charts = self.workbook()["Sales"]._charts
        self.assertEqual(len(charts), 1)
        self.assertIn("After", self.chart_xml(1))
        self.assertIn("Amount", self.chart_xml(1))
        self.edit([{"op": "edit_chart", "chart": 0, "type": "line", "range": "A1:D7"}])
        self.assertIn("<lineChart>", self.chart_xml(1))
        self.assertIn("After", self.chart_xml(1))

    def test_data_labels_show_the_value_alone(self):
        self.edit([{"op": "add_chart", "type": "bar", "range": "A1:C4", "dataLabels": True}])
        labels = self.chart_xml(1).split("<dLbls>")[1].split("</dLbls>")[0]
        self.assertEqual(labels, '<showLegendKey val="0"/><showVal val="1"/><showCatName val="0"/><showSerName val="0"/><showPercent val="0"/><showBubbleSize val="0"/>')

    def test_a_text_column_inside_the_range_is_left_out_with_a_warning(self):
        envelope = self.edit([{"op": "add_chart", "type": "combo", "range": "A1:D4"}])
        self.assertEqual([(issue["code"], issue["location"]) for issue in envelope["issues"]], [("CHART_COLUMN_LEFT_OUT", "ops[0].range")])
        self.assertIn("B2:B4 (product) holds no number", envelope["issues"][0]["message"])
        chart = self.workbook()["Sales"]._charts[0]
        self.assertEqual([series.val.numRef.f for plot in chart._charts for series in plot.series], ["'Sales'!$C$2:$C$4", "'Sales'!$D$2:$D$4"])

    def test_a_number_column_stored_as_text_is_converted_and_drawn_by_the_suggestion(self):
        self.edit([{"op": "set_range", "cell": "E1", "values": [["share"], ["12"], ["9"], ["7"], ["15"], ["6"], ["4"]]}])
        envelope = self.edit([{"op": "add_chart", "type": "bar", "range": "A1:E7"}])
        issue = envelope["issues"][0]
        self.assertIn("E2:E7 (share) holds numbers stored as text, and B2:B7 (product) holds no number", issue["message"])
        self.assertEqual(issue["fix"], [
            {"op": "set_range", "sheet": "Sales", "cell": "E2", "values": [[12], [9], [7], [15], [6], [4]]},
            {"op": "edit_chart", "sheet": "Sales", "chart": 0, "range": "A1:E7"},
        ])
        self.edit(issue["fix"])
        chart = self.workbook()["Sales"]._charts[0]
        self.assertIn("'Sales'!$E$2:$E$7", [series.val.numRef.f for series in chart.series])

    def test_a_trailing_text_column_is_left_out_of_the_suggested_range(self):
        self.edit([{"op": "set_range", "cell": "E1", "values": [["memo"], ["가"], ["나"]]}])
        issue = self.edit([{"op": "add_chart", "type": "bar", "range": "C1:E7"}])["issues"][0]
        self.assertIn('"range": "C1:D7"', issue["suggestion"])

    def test_text_columns_on_both_sides_of_the_numbers_get_a_suggestion_of_their_own(self):
        self.edit([{"op": "set_range", "cell": "E1", "values": [["memo"], ["가"], ["나"]]}])
        issue = self.edit([{"op": "add_chart", "type": "bar", "range": "A1:E7"}])["issues"][0]
        self.assertEqual((issue["code"], issue["fix"]), ("CHART_COLUMN_LEFT_OUT", []))
        self.assertNotEqual(issue["suggestion"], CHART_COLUMN_LEFT_OUT.suggestion)

    def test_a_range_that_starts_under_its_header_is_refused_with_the_range_that_holds_it(self):
        original = (self.directory / "book.xlsx").read_bytes()
        issue = self.refused([{"op": "add_chart", "type": "bar", "range": "B2:D4"}])
        self.assertEqual((issue["code"], issue["location"]), ("OPERATION_NOT_APPLICABLE", "ops[0].range"))
        self.assertIn('"range": "B1:D4"', issue["suggestion"])
        self.assertEqual((self.directory / "book.xlsx").read_bytes(), original)

    def test_one_category_column_is_widened_to_the_number_columns_beside_it(self):
        issue = self.refused([{"op": "add_chart", "type": "bar", "range": "A1:A7"}])
        self.assertIn('"range": "A1:D7"', issue["suggestion"])

    def test_a_range_without_a_number_is_refused_before_anything_is_written(self):
        original = (self.directory / "book.xlsx").read_bytes()
        issue = self.refused([{"op": "add_chart", "type": "bar", "range": "F1:H5"}])
        self.assertEqual((issue["code"], issue["location"]), ("OPERATION_NOT_APPLICABLE", "ops[0].range"))
        self.assertIn("would be empty", issue["message"])
        issue = self.refused([{"op": "add_chart", "type": "bar", "range": "D1:D7"}])
        self.assertIn('"range": "C1:D7"', issue["suggestion"])
        self.assertEqual((self.directory / "book.xlsx").read_bytes(), original)

    def test_each_series_and_each_slice_takes_the_color_it_is_given(self):
        self.edit([
            {"op": "add_chart", "type": "bar", "range": "B1:D4", "colors": ["1F4E79", "#f59e0b"]},
            {"op": "add_chart", "type": "pie", "range": "B1:C4", "anchor": "H30", "colors": ["111111", "222222"]},
        ])
        bars = self.chart_xml(1)
        self.assertLess(bars.index('<a:srgbClr val="1F4E79"/>'), bars.index('<a:srgbClr val="F59E0B"/>'))
        pie = self.chart_xml(2)
        self.assertEqual(re.findall(r'<dPt><idx val="(\d)"/><spPr><a:solidFill[^>]*><a:srgbClr val="(\w+)"/>', pie), [("0", "111111"), ("1", "222222"), ("2", "10B981")])
        self.edit([{"op": "edit_chart", "chart": 0, "colors": ["AA0000"]}])
        self.assertIn('<a:srgbClr val="AA0000"/>', self.chart_xml(1))
        issue = self.refused([{"op": "edit_chart", "chart": 0, "colors": ["red"]}])
        self.assertEqual((issue["code"], issue["location"]), ("INVALID_VALUE", "ops[0].colors[0]"))

    def test_names_from_other_vocabularies_are_answered_with_the_exact_word(self):
        cases = [
            ({"op": "add_chart", "type": "column", "range": "A1:C4"}, "ops[0].type", "use 'bar'"),
            ({"op": "add_chart", "type": "bar", "range": "A1:C4", "colors": ["red"]}, "ops[0].colors[0]", 'use "FF0000" for red'),
            ({"op": "add_pivot_table", "range": "A1:D7", "row": "region", "values": ["amount"], "function": "avg"}, "ops[0].function", "use 'average'"),
        ]
        for operation, location, suggestion in cases:
            with self.subTest(location=location):
                issue = self.refused([operation])
                self.assertEqual((issue["location"], issue["suggestion"]), (location, suggestion))

    def test_a_chart_index_beyond_the_sheet_is_refused(self):
        issue = self.refused([{"op": "delete_chart", "chart": 0}])
        self.assertIn("has 0 charts", issue["message"])

    def test_a_chart_index_beyond_the_sheet_names_the_indexes_it_has(self):
        self.edit([{"op": "add_chart", "type": "bar", "range": "A1:C4"}])
        issue = self.refused([{"op": "edit_chart", "chart": 3, "title": "After"}])
        self.assertIn("use chart 0", issue["suggestion"])
        self.edit([{"op": "delete_chart", "chart": 0}])
        issue = self.refused([{"op": "edit_chart", "chart": 3, "title": "After"}])
        self.assertIn("add_chart", issue["suggestion"])


class SparklineTest(OperationFixture):
    def test_sparklines_are_written_and_survive_another_edit(self):
        self.edit([{"op": "add_sparklines", "range": "C2:D7", "target": "F2:F7", "type": "column", "highLow": True}])
        sheet = self.sheet_xml()
        self.assertIn("{05C60535-1F16-4fd2-B633-F4F36F0B64E0}", sheet)
        self.assertIn('type="column"', sheet)
        self.assertIn("<xm:f>Sales!C7:D7</xm:f><xm:sqref>F7</xm:sqref>", sheet)
        self.edit([{"op": "insert_rows", "at": 2}])
        self.assertIn("<xm:f>Sales!C8:D8</xm:f><xm:sqref>F8</xm:sqref>", self.sheet_xml())

    def test_a_target_that_does_not_match_the_rows_is_refused(self):
        issue = self.refused([{"op": "add_sparklines", "range": "C2:D7", "target": "F2:F4"}])
        self.assertIn("one target cell per data row", issue["message"])


class PivotTest(OperationFixture):
    def test_a_pivot_writes_its_parts_and_its_numbers(self):
        envelope = self.edit([{"op": "add_pivot_table", "range": "A1:D7", "row": "region", "values": ["qty", "amount"], "targetSheet": "Pivot"}])
        names = archive_names(self.directory / "book.xlsx")
        for part in ("xl/pivotTables/pivotTable1.xml", "xl/pivotCache/pivotCacheDefinition1.xml", "xl/pivotCache/pivotCacheRecords1.xml"):
            self.assertIn(part, names)
        self.assertIn("<pivotCaches>", archive_text(self.directory / "book.xlsx", "xl/workbook.xml"))
        pivot = self.workbook()["Pivot"]
        grid = [[cell.value for cell in row] for row in pivot.iter_rows(min_row=3, max_row=7, max_col=3)]
        self.assertEqual(grid, [["region", "Sum of qty", "Sum of amount"], ["Busan", 6, 6000], ["Daegu", 7, 7000], ["Seoul", 9, 9000], ["Grand Total", 22, 22000]])
        self.assertIn("added pivot table PivotTable1", envelope["details"]["changes"][0]["change"])

    def test_a_crosstab_pivot_and_a_second_edit_keep_the_pivot(self):
        self.edit([{"op": "add_pivot_table", "range": "A1:D7", "row": "region", "column": "product", "values": [{"field": "amount", "label": "건수"}], "function": "count", "totalLabel": "합계"}])
        pivot = self.workbook()["Pivot"]
        grid = [[cell.value for cell in row] for row in pivot.iter_rows(min_row=3, max_row=8, max_col=4)]
        self.assertEqual(grid[0][:2], ["건수", "product"])
        self.assertEqual(grid[1], ["region", "A", "B", "합계"])
        self.assertEqual(grid[4], ["Seoul", 2, 1, 3])
        self.assertEqual(grid[5], ["합계", 4, 2, 6])
        self.edit([{"op": "set_cell", "sheet": "Notes", "cell": "A3", "value": "later"}])
        self.assertIn("xl/pivotTables/pivotTable1.xml", archive_names(self.directory / "book.xlsx"))
        self.assertIn('subtotal="count"', archive_text(self.directory / "book.xlsx", "xl/pivotTables/pivotTable1.xml"))

    def test_two_pivots_get_two_names_and_an_unknown_header_is_named(self):
        self.edit([
            {"op": "add_pivot_table", "range": "A1:D7", "row": "region", "values": ["qty"]},
            {"op": "add_pivot_table", "range": "A1:D7", "row": "product", "values": ["qty"], "targetCell": "A12"},
        ])
        names = sorted(text.split('name="')[1].split('"')[0] for text in (archive_text(self.directory / "book.xlsx", f"xl/pivotTables/pivotTable{number}.xml") for number in (1, 2)))
        self.assertEqual(names, ["PivotTable1", "PivotTable2"])
        issue = self.refused([{"op": "add_pivot_table", "range": "A1:D7", "row": "regoin", "values": ["qty"], "targetSheet": "P2"}])
        self.assertEqual(issue["suggestion"], "use 'region'")


class SpecOperationsTest(WorkbookFixture):
    def test_sheet_create_runs_operations_after_building_the_sheets(self):
        envelope = self.create_workbook([{"title": "Sales", "rows": SALES}])
        self.assertEqual(envelope["status"], "ok")
        from sheet_fixture import write_json, run_office
        write_json(self.directory / "spec.json", {"sheets": [{"title": "Sales", "rows": SALES}], "operations": [{"op": "add_chart", "type": "bar", "range": "A1:C4"}, {"op": "add_conditional_format", "range": "C2:C7", "rule": "data_bar"}]})
        envelope = run_office(["create", "spec.xlsx", "spec.json"], self.directory)
        self.assertEqual([change["op"] for change in envelope["details"]["changes"]], ["add_chart", "add_conditional_format"])
        self.assertEqual(len(load_workbook(self.directory / "spec.xlsx")["Sales"]._charts), 1)


if __name__ == "__main__":
    unittest.main()
