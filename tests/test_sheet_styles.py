import unittest

from openpyxl import load_workbook

from sheet_fixture import WorkbookFixture, run_office, run_office_python, write_json


STYLED_WORKBOOK = """
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

workbook = Workbook()
sheet = workbook.active
sheet.title = "Sales"
side = Side(style="thin", color="CBD5E1")
for row in [["item", "amount", "share"], ["A", 1000, 0.25], ["B", 2000, 0.5], ["C", 3000, 0.75]]:
    sheet.append(row)
for row in sheet["A3:C3"]:
    for cell in row:
        cell.font = Font(bold=True, color="FF0000")
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
        cell.border = Border(left=side, right=side, top=side, bottom=side)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
sheet["B3"].number_format = "#,##0"
sheet["C3"].number_format = "0.0%"
for row in range(1, 5):
    sheet.cell(row=row, column=2).fill = PatternFill("solid", fgColor="DDEBF7")
    sheet.cell(row=row, column=2).number_format = "#,##0.00"
workbook.save("fixture.xlsx")
"""


def style_of(cell):
    return (cell.number_format, cell.font.b, cell.font.color.rgb if cell.font.color else None, cell.fill.fgColor.rgb, cell.border.left.style, cell.alignment.horizontal, cell.alignment.wrap_text)


class InsertedCellsTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        run_office_python(STYLED_WORKBOOK, self.directory)

    def insert(self, operation):
        envelope = self.apply([operation], name="fixture.xlsx")
        self.assertEqual(envelope["status"], "ok", envelope)
        return load_workbook(self.directory / "fixture.xlsx")["Sales"]

    def test_inserted_rows_copy_the_style_of_the_row_above(self):
        sheet = self.insert({"op": "insert_rows", "sheet": "Sales", "at": 4, "count": 2})
        for row in (4, 5):
            for column in "ABC":
                self.assertEqual(style_of(sheet[f"{column}{row}"]), style_of(sheet[f"{column}3"]), f"{column}{row}")
        self.assertEqual(sheet["C5"].number_format, "0.0%")

    def test_inserted_columns_copy_the_style_of_the_column_to_the_left(self):
        sheet = self.insert({"op": "insert_columns", "sheet": "Sales", "at": "C", "count": 2})
        for row in range(1, 5):
            for column in "CD":
                self.assertEqual(style_of(sheet[f"{column}{row}"]), style_of(sheet[f"B{row}"]), f"{column}{row}")
        self.assertEqual(sheet["D3"].number_format, "#,##0.00")
        self.assertEqual(sheet["E3"].number_format, "0.0%")

    def test_a_row_inserted_at_the_top_has_nothing_above_to_copy(self):
        sheet = self.insert({"op": "insert_rows", "sheet": "Sales", "at": 1})
        self.assertFalse(sheet["A1"].has_style)


def rule_of(cell):
    return (cell.border.left.style, cell.border.top.style, cell.alignment.vertical, cell.alignment.wrap_text)


DEFAULT_RULE = ("thin", "thin", "top", True)


class DefaultTableStyleTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        (self.directory / "sales.csv").write_text("item,amount\nA,1000\nB,2000\n", encoding="utf-8")
        write_json(self.directory / "spec.json", {"sheets": [
            {"title": "Sales", "csvPath": "sales.csv"},
            {"title": "Calculated", "rows": [["item", "double"], ["A", "=2*5"], ["B", 7]]},
        ]})
        envelope = run_office(["create", "book.xlsx", "spec.json"], self.directory)
        self.assertEqual(envelope["status"], "ok", envelope)

    def apply(self, operations):
        envelope = super().apply(operations)
        self.assertEqual(envelope["status"], "ok", envelope)
        return load_workbook(self.directory / "book.xlsx")

    def test_created_formula_cells_follow_the_same_rule_as_the_rest(self):
        sheet = load_workbook(self.directory / "book.xlsx")["Calculated"]
        self.assertEqual({rule_of(cell) for row in sheet.iter_rows() for cell in row}, {DEFAULT_RULE})

    def test_cells_written_beside_a_table_take_its_default_style(self):
        sheet = self.apply([
            {"op": "set_cell", "sheet": "Sales", "cell": "C1", "value": "share"},
            {"op": "set_cell", "sheet": "Sales", "cell": "C2", "value": "=B2/B4"},
            {"op": "set_range", "sheet": "Sales", "cell": "A4", "values": [["total", 3000, "=SUM(C2:C3)"]]},
        ])["Sales"]
        for address in ("C1", "C2", "A4", "B4", "C4"):
            self.assertEqual(rule_of(sheet[address]), DEFAULT_RULE, address)
        self.assertEqual((sheet["C1"].font.b, sheet["C1"].fill.fgColor.rgb), (sheet["A1"].font.b, sheet["A1"].fill.fgColor.rgb))
        self.assertEqual(sheet["B4"].number_format, "#,##0")

    def test_a_block_written_to_a_new_sheet_becomes_a_table(self):
        self.apply([{"op": "add_sheet", "name": "Summary"}, {"op": "set_range", "sheet": "Summary", "cell": "A1", "values": [["item", "value"], ["total", 3000]]}])
        sheet = load_workbook(self.directory / "book.xlsx")["Summary"]
        self.assertEqual({rule_of(cell) for row in sheet.iter_rows() for cell in row}, {DEFAULT_RULE})
        self.assertTrue(sheet["A1"].font.b)
        self.assertFalse(sheet["A2"].font.b)

    def test_cells_far_from_any_table_a_lone_cell_and_styled_cells_are_left_alone(self):
        workbook = self.apply([
            {"op": "format_range", "sheet": "Sales", "range": "B3", "fill": "FFFF00"},
            {"op": "set_cell", "sheet": "Sales", "cell": "B3", "value": 5},
            {"op": "set_cell", "sheet": "Sales", "cell": "H20", "value": "note"},
            {"op": "add_sheet", "name": "Notes"},
            {"op": "set_cell", "sheet": "Notes", "cell": "A1", "value": "memo"},
            {"op": "set_cell", "sheet": "Sales", "cell": "A2", "value": None},
        ])
        self.assertEqual(workbook["Sales"]["B3"].fill.fgColor.rgb, "00FFFF00")
        self.assertFalse(workbook["Sales"]["H20"].has_style)
        self.assertFalse(workbook["Notes"]["A1"].has_style)


class ChartColorTest(WorkbookFixture):
    def charts(self):
        self.create_workbook([{"title": "Sales", "rows": [["month", "Seoul", "Busan"], ["Jan", 10, 5], ["Feb", 12, 6], ["Mar", 9, 7]]}])
        operations = [{"op": "add_chart", "sheet": "Sales", "type": kind, "range": "A1:C4", "anchor": anchor} for kind, anchor in (("bar", "F2"), ("line", "F20"), ("pie", "F40"))]
        self.assertEqual(self.apply(operations)["status"], "ok")
        return load_workbook(self.directory / "book.xlsx")["Sales"]._charts

    def test_every_series_has_its_own_explicit_color(self):
        bar, line, _ = self.charts()
        bar_colors = [series.graphicalProperties.solidFill.srgbClr for series in bar.series]
        line_colors = [series.graphicalProperties.line.solidFill.srgbClr for series in line.series]
        self.assertEqual(len(set(bar_colors)), 2)
        self.assertEqual(bar_colors, line_colors)
        self.assertTrue(all(len(color) == 6 for color in bar_colors))

    def test_every_pie_slice_has_its_own_explicit_color(self):
        pie = self.charts()[2]
        colors = [point.graphicalProperties.solidFill.srgbClr for point in pie.series[0].dPt]
        self.assertEqual(len(colors), 3)
        self.assertEqual(len(set(colors)), 3)


if __name__ == "__main__":
    unittest.main()
