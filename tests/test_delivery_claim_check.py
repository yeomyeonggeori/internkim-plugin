import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from host_fixture import FakeHost
from task_context_fixture import environment_with_context, write_task_context

SCRIPTS_PATH = Path(__file__).resolve().parents[1] / "skills" / "office" / "scripts"
OFFICE_ENTRY = SCRIPTS_PATH / "office"
sys.path.insert(0, str(SCRIPTS_PATH))

from delivery import claim_kinds  # noqa: E402
from delivery.claim_kinds import COMPACT_PROFILE, TODAY_PROFILE, AttachmentText, Claim, Sources, judge  # noqa: E402
from delivery.claim_rewrites import recompute, treat  # noqa: E402
from host import script_host  # noqa: E402

REQUEST = "사무실 이전 안내문 만들어 줘. 10월 20일 새 사무실로 이전합니다."
INVENTED = "이전 기간에는 전화 응대가 어렵습니다."
COMPANY = {"name": "주식회사 예시", "representative": "최견본", "representativeTitle": "대표이사"}
NOTICE = {
    "language": "ko",
    "recipient": "고객 여러분",
    "title": "사무실 이전 안내",
    "sections": [{"heading": "이전 안내", "blocks": [{"type": "paragraph", "text": f"10월 20일 새 사무실로 이전합니다. {INVENTED}"}]}],
}


def kind_answers(kinds: dict, default: str = "source"):
    def answer(body: dict) -> dict:
        answers = {}
        for name in body["questions"]:
            text = body["state"]["claims"][name]["text"]
            kind = kinds.get(text, default)
            answers[name] = {"type": "choice", "choice": kind, "probabilities": {kind: 0.9, "source": 0.1} if kind != "source" else {"source": 1.0}}
        return {"answers": answers, "modelName": "jev", "usage": {"costUSD": 0.001}}
    return answer


def recompute_answers(wrong: set):
    def answer(body: dict) -> dict:
        if "checks" in json.dumps(body["schema"]):
            units = json.loads(body["prompt"].split("\nUnits: ", 1)[1])
            return {"answer": {"checks": [{"key": key, "working": "w", "isWrong": unit["text"] in wrong} for key, unit in units.items()]}, "usage": {"costUSD": 0.002}}
        return {"answer": {"text": ""}, "usage": {"costUSD": 0.002}}
    return answer


def rewrites(texts: dict):
    def answer(body: dict) -> dict:
        unit = body["prompt"].split("\nUnit: ", 1)[1].split("\n", 1)[0]
        return {"answer": {"text": texts.get(unit, "")}, "usage": {"costUSD": 0.002}}
    return answer


def claim(text: str, path: str = "", at: str = "") -> Claim:
    return Claim(path=path or text, text=text, at=at)


class JudgeTest(unittest.TestCase):
    def test_a_claim_copied_from_the_sources_is_not_asked(self):
        with FakeHost(decide=kind_answers({})) as host:
            judgment = judge(Sources(request=(REQUEST,)), [claim("10월 20일 새 사무실로 이전합니다."), claim(INVENTED)])
        self.assertTrue(judgment.verdicts[0].is_copied)
        asked = host.requests_to("decide")[0]["state"]["claims"]
        self.assertEqual([unit["text"] for unit in asked.values()], [INVENTED])

    def test_a_number_found_only_inside_a_longer_number_is_asked(self):
        with FakeHost(decide=kind_answers({})) as host:
            judge(Sources(request=("2026년 10월 7일까지 접수합니다.",)), [claim("7"), claim("20"), claim("2026")])
        asked = host.requests_to("decide")[0]["state"]["claims"]
        self.assertEqual([unit["text"] for unit in asked.values()], ["20"])

    def test_only_a_claim_at_or_above_the_threshold_is_a_defect(self):
        def answer(body):
            return {"answers": {"claim0": {"choice": "source", "probabilities": {"source": 0.51, "claim": 0.49}}, "claim1": {"choice": "claim", "probabilities": {"claim": 0.5, "source": 0.5}}}}
        with FakeHost(decide=answer):
            judgment = judge(Sources(), [claim("a"), claim("b")])
        self.assertEqual([verdict.defect for verdict in judgment.verdicts], ["", "claim"])

    def test_every_question_asks_the_same_kinds_and_the_state_holds_their_meaning(self):
        with FakeHost(decide=kind_answers({})) as host:
            judge(Sources(request=("요청",)), [claim("a"), claim("b")])
        body = host.requests_to("decide")[0]
        self.assertEqual(body["state"]["kinds"], COMPACT_PROFILE.kinds)
        for name, question in body["questions"].items():
            self.assertEqual(set(question["criteria"]), set(COMPACT_PROFILE.kinds))
            self.assertIn(f"claims.{name}", question["instructions"])
            self.assertIn("do not add a requirement the sources do not state", question["instructions"])

    def test_many_claims_are_split_into_bounded_calls(self):
        with FakeHost(decide=kind_answers({})) as host:
            judgment = judge(Sources(), [claim(f"unit {index}") for index in range(claim_kinds.CLAIMS_PER_CALL + 5)])
        self.assertEqual([len(body["questions"]) for body in host.requests_to("decide")], [claim_kinds.CLAIMS_PER_CALL, 5])
        self.assertEqual(judgment.calls, 2)

    def test_an_attachment_is_shown_up_to_its_bound(self):
        with FakeHost(decide=kind_answers({})) as host:
            judge(Sources(attachments=(AttachmentText("a.txt", "가" * 9000),)), [claim("x")])
        shown = host.requests_to("decide")[0]["state"]["attachments"][0]["text"]
        self.assertEqual(len(shown), claim_kinds.ATTACHMENT_MAXIMUM_CHARACTERS)

    def test_a_missing_answer_is_an_error(self):
        with FakeHost(decide=lambda body: {"answers": {}}):
            with self.assertRaises(script_host.HostFailure):
                judge(Sources(), [claim("x")])

    def test_each_defect_kind_is_routed_to_its_treatment(self):
        kinds = {"c": "claim", "m": "mistake", "e": "error", "h": "hollow"}
        with FakeHost(decide=kind_answers(kinds)):
            judgment = judge(Sources(), [claim(text) for text in kinds])
        self.assertEqual({verdict.claim.text: verdict.treatment for verdict in judgment.verdicts}, {"c": "blank", "m": "blank", "e": "blank", "h": "rewrite"})

    def test_a_defect_split_across_kinds_is_flagged_as_the_strongest(self):
        verdict = COMPACT_PROFILE.verdict_for(claim("x"), "source", {"source": 0.4, "claim": 0.25, "mistake": 0.35})
        self.assertEqual(verdict.defect, "mistake")

    def test_the_today_profile_knows_only_the_old_kinds(self):
        self.assertEqual(set(TODAY_PROFILE.kinds), {"source", "derived", "expression", "claim"})
        self.assertTrue(set(COMPACT_PROFILE.treatments) <= set(COMPACT_PROFILE.kinds))

    def test_withdrawn_values_are_shown_as_removed_and_only_when_there_are_some(self):
        with FakeHost(decide=kind_answers({})) as host:
            judge(Sources(), [claim("x")])
            judge(Sources(removed=(claim("빠진 값"),)), [claim("y")])
        first, second = host.requests_to("decide")
        self.assertNotIn("removed", first["state"])
        self.assertEqual(second["state"]["removed"], [{"text": "빠진 값"}])


class RewriteTest(unittest.TestCase):
    def test_recompute_flags_only_the_derived_units_the_writer_got_wrong(self):
        kinds = {"합계 300": "derived", "합계 200": "derived", "인사": "expression"}
        with FakeHost(decide=kind_answers(kinds), generate=recompute_answers({"합계 300"})):
            judgment = recompute(Sources(), judge(Sources(), [claim(text) for text in kinds]))
        self.assertEqual({verdict.claim.text: verdict.defect for verdict in judgment.verdicts}, {"합계 300": "error", "합계 200": "", "인사": ""})

    def test_recompute_asks_nothing_when_no_unit_is_derived(self):
        with FakeHost(decide=kind_answers({}), generate=recompute_answers(set())) as host:
            recompute(Sources(), judge(Sources(), [claim("x")]))
        self.assertEqual(host.requests_to("generate"), [])

    def test_recompute_fails_loudly_on_an_answer_that_is_not_the_closed_shape(self):
        with FakeHost(decide=kind_answers({"x": "derived"}), generate=lambda body: {"answer": {"verdict": "ok"}}):
            with self.assertRaises(script_host.HostFailure):
                recompute(Sources(), judge(Sources(), [claim("x")]))

    def test_treat_blanks_what_cannot_be_rewritten_and_rewrites_hollow(self):
        kinds = {"invented": "claim", "최고의 서비스": "hollow"}
        with FakeHost(decide=kind_answers(kinds), generate=rewrites({"최고의 서비스": "10월 20일 이전"})):
            outcome = treat(Sources(), judge(Sources(), [claim(text) for text in kinds]))
        self.assertEqual([verdict.claim.text for verdict in outcome.blank], ["invented"])
        self.assertEqual([replaced.text for replaced in outcome.replaced], ["10월 20일 이전"])

    def test_treat_keeps_hollow_that_is_still_hollow_after_one_rewrite(self):
        kinds = {"최고의 서비스": "hollow", "여전히 최고": "hollow"}
        with FakeHost(decide=kind_answers(kinds), generate=rewrites({"최고의 서비스": "여전히 최고"})):
            outcome = treat(Sources(), judge(Sources(), [claim("최고의 서비스")]))
        self.assertEqual([verdict.claim.text for verdict in outcome.kept], ["최고의 서비스"])
        self.assertEqual(outcome.blank + outcome.replaced + outcome.removed, [])

    def test_a_free_sentence_whose_rewrite_is_empty_is_removed_and_a_required_slot_kept(self):
        kinds = {"빈말": "hollow"}
        free = Claim(path="sections[0].blocks[0].text#1", text="빈말", is_free=True)
        slot = Claim(path="title", text="빈말")
        with FakeHost(decide=kind_answers(kinds), generate=rewrites({})):
            outcome = treat(Sources(), judge(Sources(), [free, slot]))
        self.assertEqual([verdict.claim.path for verdict in outcome.removed], [free.path])
        self.assertEqual([verdict.claim.path for verdict in outcome.kept], [slot.path])

    def test_an_unreadable_rewrite_keeps_that_unit_and_never_fails_the_document(self):
        with FakeHost(decide=kind_answers({"빈말": "hollow"}), generate=lambda body: {"answer": {"other": 1}}):
            outcome = treat(Sources(), judge(Sources(), [claim("빈말")]))
        self.assertEqual([verdict.claim.text for verdict in outcome.kept], ["빈말"])

    def test_treat_asks_nothing_when_nothing_is_flagged(self):
        with FakeHost(decide=kind_answers({}), generate=rewrites({})) as host:
            outcome = treat(Sources(), judge(Sources(), [claim("x")]))
        self.assertTrue(outcome.changes_nothing())
        self.assertEqual(host.requests_to("generate"), [])


def deliver(directory: Path, name: str, context_path) -> dict:
    completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "delivery-check", name], capture_output=True, text=True, cwd=directory, env=environment_with_context(context_path))
    return json.loads(completed.stdout)


class DeliveredMadeFileTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        (self.directory / "notice.json").write_text(json.dumps(NOTICE, ensure_ascii=False), encoding="utf-8")

    def facts(self, **overrides) -> dict:
        return {"requester": {"name": "이샘플", "email": "sample@example.com"}, "today": "2026-10-04", "company": {"ko": COMPANY}, "request": [REQUEST]} | overrides

    def merge(self, facts: dict | None = None) -> dict:
        context_path = write_task_context(self.directory, facts or self.facts())
        subprocess.run([sys.executable, str(OFFICE_ENTRY), "merge", "letter", "notice.json", "notice.pdf"], capture_output=True, text=True, cwd=self.directory, env=environment_with_context(context_path), check=True)
        return deliver(self.directory, "notice.pdf", context_path)

    def snapshot(self) -> dict:
        return json.loads((self.directory / "notice.pdf.source.json").read_text(encoding="utf-8"))

    def metadata(self) -> dict:
        return json.loads((self.directory / "notice.pdf.meta.json").read_text(encoding="utf-8"))

    def test_an_unsupported_claim_is_blanked_by_remaking_the_file_and_the_reply_is_told_where_but_never_what_it_said(self):
        with FakeHost(decide=kind_answers({INVENTED: "claim"})):
            result = self.merge()
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "blanked")
        given_text = self.snapshot()["given"]["sections"][0]["blocks"][0]["text"]
        self.assertNotIn(INVENTED, given_text or "")
        self.assertIn("sections[0].blocks[0].text#1", [blank["field"] for blank in self.snapshot()["blanks"]])
        note = self.metadata()["notes"][0]
        self.assertIn("left blank, for the reply to offer to complete", note)
        self.assertNotIn(INVENTED, note)
        self.assertIn("never repeats", note)
        self.assertIn("nothing the person gave supports it", note)

    def test_a_supported_file_is_left_as_it_is_with_no_note(self):
        with FakeHost(decide=kind_answers({})):
            result = self.merge()
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "supported")
        self.assertEqual(self.metadata()["notes"], [])
        self.assertEqual(self.metadata()["holds"]["schema"], "letter")

    def test_a_value_the_writer_left_blank_is_named_without_saying_the_file_lost_it(self):
        notice = NOTICE | {"recipient": None, "sections": [{"heading": "이전 안내", "blocks": [{"type": "paragraph", "text": "10월 20일 새 사무실로 이전합니다."}]}]}
        (self.directory / "notice.json").write_text(json.dumps(notice, ensure_ascii=False), encoding="utf-8")
        with FakeHost(decide=kind_answers({})):
            self.merge()
        note = self.metadata()["notes"][0]
        self.assertIn("left blank, for the reply to offer to complete", note)
        self.assertNotIn("no longer holds", note)

    def test_a_note_says_only_the_taken_out_places_no_longer_hold_what_they_said(self):
        notice = NOTICE | {"recipient": None}
        (self.directory / "notice.json").write_text(json.dumps(notice, ensure_ascii=False), encoding="utf-8")
        with FakeHost(decide=kind_answers({INVENTED: "claim"})):
            self.merge()
        note = self.metadata()["notes"][0]
        never_had, _, taken_out = note.partition("no longer holds")
        self.assertIn("수신", never_had)
        self.assertNotIn("이전 안내", never_had)

    def test_a_mistake_is_blanked_and_the_note_asks_the_person_to_confirm(self):
        with FakeHost(decide=kind_answers({INVENTED: "mistake"})):
            self.merge()
        self.assertIn("ask them to confirm the right value", self.metadata()["notes"][0])

    def test_a_claim_is_not_blanked_when_an_attachment_could_not_be_read(self):
        facts = self.facts(attachments=[{"name": "plan.hwp", "path": str(self.directory / "plan.hwp"), "text": "", "current": True}])
        with FakeHost(decide=kind_answers({INVENTED: "claim"})):
            result = self.merge(facts)
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "not_enforced_unread_attachment")
        self.assertIn(INVENTED, self.snapshot()["given"]["sections"][0]["blocks"][0]["text"])
        self.assertIn("an attachment was not read", self.metadata()["notes"][0])

    def test_a_hollow_sentence_is_replaced_by_its_clean_rewrite_when_the_file_is_remade(self):
        with FakeHost(decide=kind_answers({INVENTED: "hollow"}), generate=rewrites({INVENTED: "10월 20일부터 새 사무실에서 응대합니다."})):
            result = self.merge()
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "rewritten")
        self.assertIn("10월 20일부터 새 사무실에서 응대합니다.", self.snapshot()["given"]["sections"][0]["blocks"][0]["text"])

    def test_a_hollow_sentence_with_an_empty_rewrite_is_dropped_because_a_sentence_may_go(self):
        with FakeHost(decide=kind_answers({INVENTED: "hollow"}), generate=rewrites({})):
            self.merge()
        self.assertNotIn(INVENTED, self.snapshot()["given"]["sections"][0]["blocks"][0]["text"] or "")

    def test_a_host_without_a_generation_model_still_blanks_and_never_rewrites(self):
        with FakeHost(decide=kind_answers({INVENTED: "hollow", "고객 여러분": "claim"})):
            result = self.merge()
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "blanked")
        self.assertIn(INVENTED, self.snapshot()["given"]["sections"][0]["blocks"][0]["text"])

    def test_a_statement_left_beside_blanked_values_is_judged_against_what_was_removed(self):
        notice = NOTICE | {"sections": [{"heading": "이전 안내", "blocks": [{"type": "paragraph", "text": f"{INVENTED} 위 사항을 양해 바랍니다."}]}]}
        (self.directory / "notice.json").write_text(json.dumps(notice, ensure_ascii=False), encoding="utf-8")
        removed_seen = []

        def answer(body):
            removed_seen.append(body["state"].get("removed"))
            return kind_answers({INVENTED: "claim", "위 사항을 양해 바랍니다.": "claim" if body["state"].get("removed") else "expression"})(body)
        with FakeHost(decide=answer):
            result = self.merge()
        self.assertEqual(removed_seen[-1], [{"at": result["details"]["claimCheck"][0]["blanked"][0], "text": INVENTED}])
        self.assertEqual(len(result["details"]["claimCheck"][0]["blanked"]), 2)

    def test_a_file_made_off_any_host_gets_no_check_and_no_metadata(self):
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "merge", "letter", "notice.json", "notice.pdf"], capture_output=True, text=True, cwd=self.directory, env=environment_with_context(None))
        result = json.loads(completed.stdout)
        self.assertNotIn("claimCheck", result.get("details") or {})
        self.assertFalse((self.directory / "notice.pdf.meta.json").exists())



FIGURE = "재택근무로 생산성이 13% 향상됩니다."
REPORT = {"language": "ko", "title": "재택근무 안내", "sections": [{"heading": "기대 효과", "blocks": [{"type": "fields", "fields": [{"label": "생산성 향상", "type": "percent", "value": 13}]}]}]}


class DeliveredFileTextTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.context_path = write_task_context(self.directory, {"requester": {"name": "이샘플", "email": "sample@example.com"}, "today": "2026-10-04", "company": {"ko": COMPANY}, "request": [REQUEST]})

    def office(self, *words: str) -> dict:
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), *words], capture_output=True, text=True, cwd=self.directory, env=environment_with_context(self.context_path))
        return json.loads(completed.stdout)

    def text_of(self, name: str) -> str:
        return json.dumps(self.office("read", name)["details"], ensure_ascii=False)

    def write_by_script(self, name: str, paragraphs: list[str]) -> None:
        from docx import Document

        document = Document()
        document.add_heading("이전 안내", level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        document.save(self.directory / name)

    def test_a_file_a_script_wrote_is_judged_by_its_own_text_and_the_invented_figure_is_blanked(self):
        self.write_by_script("memo.docx", ["10월 20일 새 사무실로 이전합니다.", f"이전 후에도 업무는 이어집니다. {FIGURE}", INVENTED])
        with FakeHost(decide=kind_answers({FIGURE: "claim", INVENTED: "claim"})):
            result = deliver(self.directory, "memo.docx", self.context_path)
        text = self.text_of("memo.docx")
        self.assertNotIn("13%", text)
        self.assertIn("10월 20일 새 사무실로 이전합니다.", text)
        self.assertIn("이전 후에도 업무는 이어집니다.", text)
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "blanked")
        notes = json.loads((self.directory / "memo.docx.meta.json").read_text(encoding="utf-8"))["notes"]
        self.assertNotIn("13%", notes[0])
        self.assertNotIn(INVENTED, notes[0])

    def test_a_merged_file_rewritten_afterwards_is_judged_by_what_it_now_says(self):
        (self.directory / "notice.json").write_text(json.dumps(NOTICE | {"sections": [{"heading": "이전 안내", "blocks": [{"type": "paragraph", "text": "10월 20일 새 사무실로 이전합니다."}]}]}, ensure_ascii=False), encoding="utf-8")
        self.office("merge", "letter", "notice.json", "notice.docx")
        from docx import Document

        document = Document(self.directory / "notice.docx")
        document.add_paragraph(FIGURE)
        document.save(self.directory / "notice.docx")
        with FakeHost(decide=kind_answers({FIGURE: "claim"})):
            deliver(self.directory, "notice.docx", self.context_path)
        self.assertNotIn("13%", self.text_of("notice.docx"))

    def test_an_invented_percent_in_a_typed_field_is_blanked_when_the_file_is_made_again(self):
        (self.directory / "report.json").write_text(json.dumps(REPORT, ensure_ascii=False), encoding="utf-8")
        self.office("merge", "report", "report.json", "report.docx")
        with FakeHost(decide=kind_answers({"13%": "claim"})):
            result = deliver(self.directory, "report.docx", self.context_path)
        self.assertNotIn("13%", self.text_of("report.docx"))
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "blanked")

    def blocks_of(self, name: str) -> list[tuple[str, str]]:
        return [(block["kind"], block.get("text") or "") for block in self.office("read", name)["details"]["blocks"]]

    def write_memo_with_invented_sections(self) -> None:
        from docx import Document

        document = Document()
        document.add_heading("사무실 이전 안내", level=1)
        document.add_paragraph("발신: 경영지원팀")
        document.add_paragraph("10월 20일 새 사무실로 이전합니다.")
        document.add_heading("기대 효과", level=2)
        document.add_paragraph("출근 시간이 20% 줄어듭니다.", style="List Bullet")
        document.add_paragraph("업무 만족도가 15% 오릅니다.", style="List Bullet")
        document.add_heading("유의사항", level=2)
        document.add_paragraph("이전 기간에는 전화 응대가 어렵습니다.")
        document.sections[0].header.paragraphs[0].text = "주식회사 예시 대외비 등급 A"
        document.save(self.directory / "memo.docx")

    def test_a_blanked_line_keeps_its_label_a_blanked_sentence_goes_and_an_emptied_section_keeps_one_line_to_fill(self):
        self.write_memo_with_invented_sections()
        flagged = {"발신: 경영지원팀": "claim", "출근 시간이 20% 줄어듭니다.": "claim", "업무 만족도가 15% 오릅니다.": "claim", "유의사항": "claim", INVENTED: "claim"}
        with FakeHost(decide=kind_answers(flagged)):
            deliver(self.directory, "memo.docx", self.context_path)
        self.assertEqual(self.blocks_of("memo.docx"), [
            ("heading", "사무실 이전 안내"),
            ("paragraph", "발신: __________"),
            ("paragraph", "10월 20일 새 사무실로 이전합니다."),
            ("heading", "기대 효과"),
            ("listItem", "__________"),
        ])

    def test_an_invented_header_is_judged_and_blanked(self):
        self.write_memo_with_invented_sections()
        with FakeHost(decide=kind_answers({"주식회사 예시 대외비 등급 A": "claim"})) as host:
            deliver(self.directory, "memo.docx", self.context_path)
        asked = [unit["text"] for body in host.requests_to("decide") for unit in body["state"]["claims"].values()]
        self.assertIn("주식회사 예시 대외비 등급 A", asked)
        self.assertEqual(self.office("read", "memo.docx")["details"]["sections"][0]["header"], "")

    def test_a_chart_holding_an_invented_value_is_taken_out(self):
        self.write_by_script("memo.docx", ["10월 20일 새 사무실로 이전합니다."])
        operations = self.directory / "chart.json"
        operations.write_text(json.dumps([{"op": "insert_chart", "at": "end", "type": "column", "categories": ["9월", "10월"], "series": [{"name": "방문객", "values": [120, 340]}]}]), encoding="utf-8")
        self.office("apply", "memo.docx", str(operations))
        with FakeHost(decide=kind_answers({"10월 방문객 340": "claim"})):
            result = deliver(self.directory, "memo.docx", self.context_path)
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "blanked")
        self.assertNotIn("charts", self.office("read", "memo.docx")["details"])

    def test_a_pdf_converted_from_a_document_is_blanked_through_that_document_and_converted_again(self):
        self.write_by_script("memo.docx", ["10월 20일 새 사무실로 이전합니다.", INVENTED])
        (self.directory / "out").mkdir()
        self.office("convert", "memo.docx", "out/memo.pdf")
        with FakeHost(decide=kind_answers({INVENTED: "claim"})):
            result = deliver(self.directory / "out", "memo.pdf", self.context_path)
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "blanked")
        self.assertNotIn(INVENTED, self.text_of("out/memo.pdf"))
        self.assertIn("10월 20일 새 사무실로 이전합니다.", self.text_of("out/memo.pdf"))

    def judged_state(self, host) -> dict:
        return host.requests_to("decide")[0]["state"]

    def test_the_judge_reads_the_company_profile_the_task_recorded(self):
        self.write_by_script("memo.docx", ["발신: 주식회사 예시 총무팀"])
        with FakeHost(decide=kind_answers({})) as host:
            deliver(self.directory, "memo.docx", self.context_path)
        self.assertEqual(self.judged_state(host)["runtimeFacts"]["company"]["ko"]["name"], COMPANY["name"])

    def test_the_judge_asks_for_the_company_profile_when_the_task_never_read_it(self):
        context_path = write_task_context(self.directory / "unread", {"requester": {"name": "이샘플", "email": "sample@example.com"}, "today": "2026-10-04", "request": [REQUEST]})
        profile_path = self.directory / "company-profile.json"
        profile_path.write_text(json.dumps(COMPANY | {"logoImage": "logo.png", "phone": ""}, ensure_ascii=False), encoding="utf-8")
        self.write_by_script("memo.docx", ["발신: 주식회사 예시"])
        answer = {"result": {}, "files": [{"name": "company-profile.json", "path": str(profile_path)}]}
        with FakeHost(decide=kind_answers({}), tools={"company_info_get": lambda body: answer}) as host:
            deliver(self.directory, "memo.docx", context_path)
        self.assertEqual(self.judged_state(host)["runtimeFacts"]["company"], {"ko": COMPANY})

    def test_an_addressing_line_is_judged_by_whom_it_names_and_is_a_mistake_only_against_a_named_party(self):
        self.write_by_script("memo.docx", ["수신: 사내 전체"])
        with FakeHost(decide=kind_answers({})) as host:
            deliver(self.directory, "memo.docx", self.context_path)
        state = self.judged_state(host)
        self.assertIn("addressing line", state["kinds"]["source"])
        self.assertIn("sources name a different sender or recipient", next(iter(host.requests_to("decide")[0]["questions"].values()))["instructions"])

    def test_a_supported_script_written_file_is_delivered_as_it_is(self):
        self.write_by_script("memo.docx", ["10월 20일 새 사무실로 이전합니다."])
        before = (self.directory / "memo.docx").read_bytes()
        with FakeHost(decide=kind_answers({})):
            result = deliver(self.directory, "memo.docx", self.context_path)
        self.assertEqual(result["details"]["claimCheck"][0]["outcome"], "supported")
        self.assertEqual((self.directory / "memo.docx").read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
