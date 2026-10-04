import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from render_fixture import can_render


SCRIPTS_PATH = Path(__file__).resolve().parents[1] / "skills" / "office" / "scripts"
OFFICE_ENTRY = SCRIPTS_PATH / "office"
sys.path.insert(0, str(SCRIPTS_PATH))

from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from charts.combo import combo_chart_space  # noqa: E402
from powerpoint.chart_audit import chart_issues  # noqa: E402


C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
YEARS = ("2023", "2024", "2025", "2026")
REVENUE_AND_LOSS = [("매출", (8.2, 19.5, 36.8, 61.0), False), ("영업이익", (-12.4, -7.1, -1.9, 4.3), True)]


def set_limits(axis, minimum: float, maximum: float) -> None:
    scaling = axis.find(f"{C}scaling")
    for tag in ("min", "max"):
        for existing in scaling.findall(f"{C}{tag}"):
            scaling.remove(existing)
    for tag, value in (("max", maximum), ("min", minimum)):
        scaling.append(scaling.makeelement(f"{C}{tag}", {"val": repr(value)}))


def codes(issues) -> list[str]:
    return [issue.kind.code for issue in issues]


def column_presentation(path: Path, values: tuple[float, ...], maximum: float | None) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = list(YEARS[: len(values)])
    data.add_series("매출", list(values))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(8), Inches(5), data).chart
    if maximum is not None:
        chart.value_axis.maximum_scale = maximum
    presentation.save(str(path))


class ChartAxisAuditTest(unittest.TestCase):
    def test_a_combo_whose_axes_share_their_zero_holds_every_point(self):
        root = combo_chart_space(YEARS, REVENUE_AND_LOSS, True)
        self.assertEqual(chart_issues(root, "slide 9", "chart 0"), [])

    def test_two_axes_that_draw_zero_at_different_heights_are_misaligned(self):
        root = combo_chart_space(YEARS, REVENUE_AND_LOSS, True)
        primary, secondary = root.findall(f".//{C}valAx")
        set_limits(primary, 0, 101.66666666666667)
        set_limits(secondary, -59.71666666666667, 9.86666666666666)
        issues = chart_issues(root, "slide 9", "chart 0")
        self.assertEqual(codes(issues), ["CHART_ZERO_MISALIGNED"])
        self.assertIn("영업이익", issues[0].message)
        self.assertEqual(issues[0].location, "slide 9")

    def test_a_value_beyond_its_axis_limit_is_named(self):
        root = combo_chart_space(YEARS, REVENUE_AND_LOSS, False)
        set_limits(root.find(f".//{C}valAx"), 0, 80)
        issues = chart_issues(root, "slide 9", "chart 0")
        self.assertEqual(codes(issues), ["CHART_POINT_OUTSIDE_AXIS"])
        self.assertIn("-12.4, -7.1, -1.9", issues[0].message)

    def test_the_check_of_a_pptx_reports_a_column_its_axis_cuts_off(self):
        with tempfile.TemporaryDirectory() as directory:
            broken, whole = Path(directory) / "broken.pptx", Path(directory) / "whole.pptx"
            column_presentation(broken, (12, 30, 55), 40)
            column_presentation(whole, (12, 30, 55), None)
            reported = {name: self.check(path) for name, path in (("broken", broken), ("whole", whole))}
        self.assertIn(("CHART_POINT_OUTSIDE_AXIS", "slide 1"), reported["broken"])
        self.assertNotIn("CHART_POINT_OUTSIDE_AXIS", {code for code, _ in reported["whole"]})

    def check(self, path: Path) -> set[tuple[str, str]]:
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "check", str(path), "--no-preview"], capture_output=True, text=True)
        return {(issue["code"], issue["location"]) for issue in json.loads(completed.stdout)["issues"]}

    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_a_built_deck_with_a_loss_making_combo_writes_axes_that_hold_it(self):
        source = (
            '<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>재무</title></head><body data-theme="corporate">'
            '<section data-layout="chart"><h2>영업이익은 2026년에 흑자로 돌아섭니다</h2>'
            '<figure data-chart="combo" data-labels="2023, 2024, 2025, 2026" data-series="매출: 820, 1950, 3680, 6100; 영업이익: -12.4, -7.1, -1.9, 4.3" data-unit="백만원, 억원">'
            "<figcaption>연도별 매출과 영업이익 · 자료: 재무팀</figcaption></figure><aside class=\"notes\">재무</aside></section></body></html>"
        )
        with tempfile.TemporaryDirectory() as directory:
            (Path(directory) / "slides.html").write_text(source, encoding="utf-8")
            completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "create", "build/deck.pptx", "slides.html"], capture_output=True, text=True, cwd=directory)
            envelope = json.loads(completed.stdout)
            presentation = Presentation(str(Path(directory) / "build" / "deck.pptx"))
        chart_space = next(shape for shape in presentation.slides[0].shapes if shape.has_chart).chart._chartSpace
        self.assertEqual(len(chart_space.findall(f".//{C}valAx")), 2)
        self.assertEqual(codes(chart_issues(chart_space, "slide 1", "chart")), [])
        self.assertEqual({issue["code"] for issue in envelope["issues"]} & {"CHART_POINT_OUTSIDE_AXIS", "CHART_ZERO_MISALIGNED"}, set())


if __name__ == "__main__":
    unittest.main()
