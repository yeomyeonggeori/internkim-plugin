import unittest
import zipfile

from sheet_fixture import WorkbookFixture, run_office, run_office_python, stored_cells, write_json


TEMPLATE = """
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
workbook = Workbook()
sheet = workbook.active
sheet.title = "견적서"
sheet["A1"] = "견적서: {{ customer.name }} 귀하"
sheet["A2"] = CellRichText("담당: {{ man", TextBlock(InlineFont(b=True), "ager }}"))
sheet["B3"] = "{{ amount }}"
sheet["B4"] = "=B3*1.1"
sheet["B5"] = "{{ issued }}"
sheet["A6"] = "{{ items.0.name }}"
sheet.oddHeader.center.text = "{{ customer.name }}"
workbook.save("quote.xlsx")
"""


class SheetMergeTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        run_office_python(TEMPLATE, self.directory)

    def merge(self, values):
        write_json(self.directory / "values.json", values)
        return run_office(["sheet", "merge", "quote.xlsx", "values.json", "filled.xlsx"], self.directory)

    def test_a_whole_cell_placeholder_takes_the_value_type_and_formulas_recompute(self):
        envelope = self.merge({"customer": {"name": "박예시"}, "manager": "최견본", "amount": 1200000, "issued": "2026-10-01", "items": [{"name": "연간 유지보수"}]})
        self.assertEqual(envelope["status"], "ok", envelope)
        cells = stored_cells(self.directory / "filled.xlsx")
        self.assertEqual((cells["B3"]["type"], cells["B3"]["value"]), (None, "1200000"))
        self.assertEqual(cells["B4"]["value"], "1320000")
        self.assertEqual(cells["B5"]["type"], "inlineStr")
        read = run_office(["sheet", "read", "filled.xlsx", "--range", "A1:A6"], self.directory)
        texts = [value for row in read["details"]["range"]["values"] for value in row if value]
        self.assertEqual(texts, ["견적서: 박예시 귀하", "담당: 최견본", "연간 유지보수"])
        with zipfile.ZipFile(self.directory / "filled.xlsx") as archive:
            self.assertIn("박예시".encode(), archive.read("xl/worksheets/sheet1.xml"))

    def test_a_missing_value_or_an_object_in_one_cell_writes_nothing(self):
        envelope = self.merge({"customer": {"name": "박예시"}, "manager": "최견본", "amount": 1, "issued": {"date": "2026-10-01"}})
        self.assertEqual(sorted((issue["code"], issue["location"]) for issue in envelope["issues"]), [("LIST_NEEDS_A_ROW", "견적서!B5"), ("UNRESOLVED_PLACEHOLDER", "견적서!A6")])
        self.assertFalse((self.directory / "filled.xlsx").exists())

    def test_untouched_parts_are_kept_byte_for_byte(self):
        self.merge({"customer": {"name": "박예시"}, "manager": "최견본", "amount": 5, "issued": "오늘", "items": [{"name": "품목"}]})
        with zipfile.ZipFile(self.directory / "quote.xlsx") as before, zipfile.ZipFile(self.directory / "filled.xlsx") as after:
            unchanged = [name for name in before.namelist() if name not in ("xl/worksheets/sheet1.xml", "xl/sharedStrings.xml") and before.read(name) != after.read(name)]
        self.assertEqual(unchanged, [])


LIST_TEMPLATE = """
from openpyxl import Workbook
workbook = Workbook()
sheet = workbook.active
sheet.title = "견적"
sheet["A1"], sheet["B1"], sheet["C1"], sheet["D1"] = "품목", "수량", "단가", "금액"
sheet["A2"], sheet["B2"], sheet["C2"], sheet["D2"] = "{{ items.name }}", "{{ items.qty }}", "{{ items.price }}", "=B2*C2"
sheet["A3"], sheet["D3"] = "합계", "=SUM(D2:D2)"
workbook.save("list.xlsx")
"""


class SheetListMergeTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        run_office_python(LIST_TEMPLATE, self.directory)

    def merge(self, items):
        write_json(self.directory / "values.json", {"items": items})
        return run_office(["sheet", "merge", "list.xlsx", "values.json", "filled.xlsx"], self.directory)

    def formulas_and_values(self, cell_range):
        read = run_office(["sheet", "read", "filled.xlsx", "--range", cell_range, "--where", "formula"], self.directory)
        return {cell["cell"]: (cell["formula"], cell["value"]) for cell in read["details"]["range"]["cells"]}

    def test_a_row_naming_a_list_repeats_per_item_and_the_total_grows_with_it(self):
        envelope = self.merge([{"name": "노트북", "qty": 2, "price": 1200000}, {"name": "모니터", "qty": 3, "price": 300000}, {"name": "키보드", "qty": 5, "price": 89000}])
        self.assertEqual(envelope["status"], "ok", envelope)
        self.assertEqual(self.formulas_and_values("D2:D5"), {
            "D2": ("=B2*C2", 2400000), "D3": ("=B3*C3", 900000), "D4": ("=B4*C4", 445000), "D5": ("=SUM(D2:D4)", 3745000),
        })

    def test_an_empty_list_leaves_the_row_blank_and_the_total_whole(self):
        self.assertEqual(self.merge([])["status"], "ok")
        self.assertEqual(self.formulas_and_values("D2:D3"), {"D2": ("=B2*C2", 0), "D3": ("=SUM(D2:D2)", 0)})


    def test_a_row_naming_a_list_of_values_repeats_with_each_value(self):
        run_office_python("""
from openpyxl import Workbook
workbook = Workbook()
sheet = workbook.active
sheet["A1"], sheet["A2"] = "{{ regions }}", "끝"
workbook.save("list.xlsx")
""", self.directory)
        write_json(self.directory / "values.json", {"regions": ["서울", "부산"]})
        self.assertEqual(run_office(["sheet", "merge", "list.xlsx", "values.json", "filled.xlsx"], self.directory)["status"], "ok")
        read = run_office(["sheet", "read", "filled.xlsx", "--range", "A1:A3"], self.directory)
        self.assertEqual(read["details"]["range"]["values"], [["서울"], ["부산"], ["끝"]])


if __name__ == "__main__":
    unittest.main()
