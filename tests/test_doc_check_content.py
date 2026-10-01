import re
import zipfile

from doc_fixture import DocumentFixture, run_office, run_office_python, write_json


CONTENT_CODES = {"MISSING_IMAGE", "CHART_EMPTY", "EMPTY_HEADING", "HEADING_SKIP", "UNRESOLVED_COMMENTS", "FIELD_NOT_EVALUATED"}


def without_member(path, prefix):
    with zipfile.ZipFile(path) as archive:
        members = {member.filename: archive.read(member) for member in archive.infolist()}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            if not name.startswith(prefix):
                archive.writestr(name, data)


def emptied_chart_values(path):
    with zipfile.ZipFile(path) as archive:
        members = {member.filename: archive.read(member) for member in archive.infolist()}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            if name.startswith("word/charts/chart"):
                data = re.sub(rb"(<c:val>.*?</c:val>)", lambda match: re.sub(rb"<c:pt idx=\"\d+\"><c:v>[^<]*</c:v></c:pt>", b"", match.group(1)), data, flags=re.S)
            archive.writestr(name, data)


class ContentCheckTest(DocumentFixture):
    def build_defects(self):
        run_office_python("""
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from PIL import Image
            Image.new("RGB", (40, 30), "navy").save("picture.png")
            document = Document("fixture.docx")
            document.add_heading("", 2)
            document.add_heading("세부 내역", 3)
            date = document.add_paragraph("작성일: ")
            date._p.append(OxmlElement("w:fldSimple", attrs={qn("w:instr"): " DATE "}))
            document.add_picture("picture.png")
            document.save("fixture.docx")
        """, self.directory)
        write_json(self.directory / "setup.json", [
            {"op": "insert_chart", "at": "end", "type": "column", "categories": ["1분기", "2분기"], "series": [{"name": "매출", "values": [3, 4]}]},
            {"op": "add_comment", "block": 1, "text": "금액 확인 필요", "author": "박예시"},
        ])
        self.assertEqual(run_office(["doc", "apply", "fixture.docx", "setup.json"], self.directory)["status"], "ok")
        emptied_chart_values(self.directory / "fixture.docx")
        without_member(self.directory / "fixture.docx", "word/media/")

    def test_each_content_finding_names_the_operation_that_clears_it(self):
        self.build_defects()
        envelope = run_office(["doc", "check", "fixture.docx"], self.directory)
        findings = {issue["code"]: issue for issue in envelope["issues"] if issue["code"] in CONTENT_CODES}
        self.assertEqual(set(findings), CONTENT_CODES)
        self.assertEqual(findings["HEADING_SKIP"]["fix"][0]["style"], "Heading 2")
        chart_fix = findings["CHART_EMPTY"]["fix"][0]
        self.assertEqual(chart_fix["categories"], ["1분기", "2분기"])
        chart_fix["series"] = [{"name": "매출", "values": [5, 6]}]
        fixes = [operation for code in ("MISSING_IMAGE", "EMPTY_HEADING", "HEADING_SKIP", "UNRESOLVED_COMMENTS", "FIELD_NOT_EVALUATED") for operation in findings[code]["fix"]] + [chart_fix]
        write_json(self.directory / "fixes.json", fixes)
        self.assertEqual(run_office(["doc", "apply", "fixture.docx", "fixes.json"], self.directory)["status"], "ok")
        remaining = {issue["code"] for issue in run_office(["doc", "check", "fixture.docx"], self.directory)["issues"]}
        self.assertFalse(remaining & CONTENT_CODES, remaining)

    def test_a_document_this_skill_writes_raises_none_of_them(self):
        write_json(self.directory / "edits.json", [
            {"op": "insert_heading", "at": "end", "text": "부록", "level": 2},
            {"op": "add_bookmark", "block": 0, "name": "overview"},
            {"op": "insert_cross_reference", "block": 1, "bookmark": "overview", "show": "page"},
            {"op": "set_footer", "text": "{PAGE} / {NUMPAGES}"},
        ])
        self.assertEqual(run_office(["doc", "apply", "fixture.docx", "edits.json"], self.directory)["status"], "ok")
        codes = {issue["code"] for issue in run_office(["doc", "check", "fixture.docx"], self.directory)["issues"]}
        self.assertFalse(codes & CONTENT_CODES, codes)
