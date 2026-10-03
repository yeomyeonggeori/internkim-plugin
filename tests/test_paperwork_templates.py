import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest

from doc_fixture import OFFICE_ENTRY, SCRIPTS_PATH, run_office


from paperwork.contract_template import Text, placeholders
from paperwork.forms import Form, form_slugs
from paperwork.jurisdictions import JURISDICTIONS


SPECIFICATIONS_PATH = SCRIPTS_PATH.parent / "references" / "paperwork" / "kr"
ENGLISH_SPECIFICATIONS_PATH = SPECIFICATIONS_PATH.parent / "intl"
SKELETON_PATTERN = re.compile(r"## Context JSON skeleton\s+```json\n(.*?)\n```", re.DOTALL)
DOCUMENT_SKELETON_PATTERN = re.compile(r"## Document JSON skeleton\s+```json\n(.*?)\n```", re.DOTALL)
CLAUSE_LINE = re.compile(r"^- `(\w+)`", re.MULTILINE)
TAX_RATE_MENTION = re.compile(r"(?:부가세|VAT|Tax) ?\((\d+)%\)|(?:세액은 공급가액의|tax as) (\d+)%")
PROFILE_PLACEHOLDER = re.compile(r"\{ \.\.\.[^}]*\.\.\. \}")
DOCUMENT_TEXT = """
import json, sys
from docx import Document
document = Document(sys.argv[1])
paragraphs = [[paragraph.text, paragraph._p.pPr.numPr.numId.val if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None else None] for paragraph in document.paragraphs]
cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
print(json.dumps({"paragraphs": paragraphs, "cells": cells}, ensure_ascii=False))
"""


def contract_templates():
    forms = [Form(jurisdiction, slug) for jurisdiction in JURISDICTIONS for slug in form_slugs(jurisdiction.code)]
    return {form.name: form.contract_template for form in forms if form.contract_template is not None}


def sample_value(term, index):
    if term.type == "list":
        return [f"{term.name} 항목 하나", f"{term.name} 항목 둘"]
    if term.type == "number":
        return str(index + 7)
    if term.type == "amount":
        return f"{index + 1},000,000"
    if term.type == "choice":
        return term.options[-1]
    return f"{term.name} 값"


def complete_values(template):
    return {term.name: sample_value(term, index) for index, term in enumerate(template.terms) if not term.is_derived}


def template_texts(template):
    pieces = [*template.preamble, *template.closing, *[piece for article in template.articles for piece in article.pieces]]
    return [piece for piece in pieces], [line for column in template.signatures for line in column]


def stated_terms(template):
    pieces, signature_lines = template_texts(template)
    names = set()
    for piece in [*pieces, *(Text(line) for line in signature_lines)]:
        if not isinstance(piece, Text):
            names.add(piece.term)
        elif not piece.when:
            names.update(name for is_article, name in placeholders(piece.text) if not is_article)
    return names


def specification_text(form_name):
    return (SPECIFICATIONS_PATH.parent / f"{form_name}.md").read_text(encoding="utf-8")


def documented_fields(form_name):
    match = SKELETON_PATTERN.search(specification_text(form_name))
    return json.loads(match.group(1)) if match else None


def documented_clauses(form_name):
    specification = specification_text(form_name)
    section = specification.split("## Clauses", 1)[1].split("\n## ", 1)[0] if "## Clauses" in specification else ""
    return CLAUSE_LINE.findall(section)


def codes(envelope):
    return [(issue["code"], issue["location"]) for issue in envelope["issues"]]


def document_skeleton_fields(path):
    match = DOCUMENT_SKELETON_PATTERN.search(path.read_text(encoding="utf-8"))
    return field_paths(json.loads(PROFILE_PLACEHOLDER.sub("{}", match.group(1)))) if match else None


def field_paths(value, prefix=""):
    if isinstance(value, dict):
        return {path for key, child in value.items() for path in {prefix + key} | field_paths(child, f"{prefix}{key}.")}
    if isinstance(value, list):
        return {path for item in value for path in field_paths(item, f"{prefix}[].")}
    return set()


class ContractTemplateShapeTest(unittest.TestCase):
    def test_the_bundled_contracts_are_the_templates_the_catalog_names(self):
        self.assertEqual(sorted(contract_templates()), ["kr/employment-contract", "kr/mou", "kr/nda", "kr/offer-letter", "kr/service-agreement"])

    def test_every_placeholder_names_a_term_and_every_reference_a_clause(self):
        for name, template in contract_templates().items():
            pieces, signature_lines = template_texts(template)
            texts = [piece.text for piece in pieces if isinstance(piece, Text)] + signature_lines
            for is_article, placeholder in (found for text in texts for found in placeholders(text)):
                known = template.article_keys if is_article else [term.name for term in template.terms]
                self.assertIn(placeholder, known, name)

    def test_every_term_a_template_defines_is_printed_somewhere(self):
        for name, template in contract_templates().items():
            pieces, _ = template_texts(template)
            conditional = {piece.when for piece in pieces if isinstance(piece, Text) and piece.when}
            self.assertEqual({term.name for term in template.terms}, stated_terms(template) | conditional, name)

    def test_the_spec_skeleton_gives_exactly_the_terms_a_caller_writes(self):
        for name, template in contract_templates().items():
            documented = documented_fields(name)
            self.assertIsNotNone(documented, f"{name}.md has no context skeleton")
            self.assertEqual(set(documented) - {"form"}, {term.name for term in template.terms if not term.is_derived}, name)
            self.assertEqual(sorted(field for field, value in documented.items() if isinstance(value, list)), sorted(term.name for term in template.terms if term.type == "list"), name)

    def test_the_spec_lists_exactly_the_template_clauses_in_order(self):
        for name, template in contract_templates().items():
            self.assertEqual(documented_clauses(name), template.article_keys, name)

    def test_the_guide_prints_each_template_terms_and_clauses_from_the_template(self):
        guide = subprocess.run([sys.executable, str(OFFICE_ENTRY), "guide", "form"], capture_output=True, text=True, check=True).stdout
        for template in contract_templates().values():
            self.assertIn(f"    clauses: {', '.join(template.article_keys)}", guide)
            for term in template.terms:
                self.assertIn(f"    {term.describe()}", guide)
        for code in ("MISSING_TERM", "UNKNOWN_TERM", "TERM_NOT_STATED", "DUPLICATE_CLAUSE"):
            self.assertIn(code, guide)


class SpecificationTest(unittest.TestCase):
    def test_the_skill_routes_every_catalog_form_to_merge(self):
        skill = (SCRIPTS_PATH.parent / "SKILL.md").read_text(encoding="utf-8")
        route = " ".join(line for line in skill.splitlines() if "`merge " in line and line.startswith("|")).lower()
        for slug in sorted({slug for jurisdiction in JURISDICTIONS for slug in form_slugs(jurisdiction.code)}):
            self.assertIn(slug.replace("-", " "), route, slug)

    def test_english_and_korean_specs_describe_the_same_document_fields(self):
        for korean in sorted(SPECIFICATIONS_PATH.glob("*.md")):
            english = ENGLISH_SPECIFICATIONS_PATH / korean.name
            self.assertTrue(english.exists(), korean.name)
            korean_fields = document_skeleton_fields(korean)
            if korean_fields is not None:
                self.assertEqual(document_skeleton_fields(english), korean_fields, korean.name)

    def test_every_tax_rate_a_spec_states_is_the_rate_its_jurisdiction_checks(self):
        for jurisdiction in JURISDICTIONS:
            if jurisdiction.tax_rate_percent is None:
                continue
            for specification in sorted((SPECIFICATIONS_PATH.parent / jurisdiction.code).glob("*.md")):
                for match in TAX_RATE_MENTION.finditer(specification.read_text(encoding="utf-8")):
                    self.assertEqual(int(match.group(1) or match.group(2)), jurisdiction.tax_rate_percent, f"{jurisdiction.code}/{specification.name}: {match.group(0)}")

    def test_every_spec_skeleton_names_its_own_form(self):
        for specification in sorted(SPECIFICATIONS_PATH.parent.glob("*/*.md")):
            match = DOCUMENT_SKELETON_PATTERN.search(specification.read_text(encoding="utf-8")) or SKELETON_PATTERN.search(specification.read_text(encoding="utf-8"))
            named = re.search(r'"form": "([^"]+)"', match.group(1))
            self.assertEqual(named.group(1), f"{specification.parent.name}/{specification.stem}")


class ContractRun(unittest.TestCase):
    def setUp(self):
        self.directory = Path(self.enterContext(tempfile.TemporaryDirectory()))

    def run_form(self, verb, form_name, values):
        (self.directory / "values.json").write_text(json.dumps({"form": form_name, **values}, ensure_ascii=False), encoding="utf-8")
        arguments = ["merge", form_name, "values.json", "out.docx"] if verb == "merge" else ["check", "values.json"]
        return run_office(arguments, self.directory)

    def merged_text(self, form_name, values):
        envelope = self.run_form("merge", form_name, values)
        self.assertEqual(envelope["status"], "ok", envelope)
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "-c", DOCUMENT_TEXT, str(self.directory / "out.docx")], capture_output=True, text=True, check=True)
        return envelope, json.loads(completed.stdout)

    def paragraphs(self, text):
        return [paragraph for paragraph, _ in text["paragraphs"]]

    def headings(self, text):
        return [paragraph for paragraph in self.paragraphs(text) if re.fullmatch(r"제\d+조 \(.+\)", paragraph)]


class EveryTermIsTypedTest(ContractRun):
    def test_a_complete_contract_prints_every_value_it_is_given_and_nothing_left_unfilled(self):
        for name, template in contract_templates().items():
            values = complete_values(template)
            envelope, text = self.merged_text(name, values)
            whole = "\n".join(self.paragraphs(text) + text["cells"])
            self.assertNotIn("{{", whole, name)
            for term_name, value in values.items():
                for item in value if isinstance(value, list) else [value]:
                    self.assertIn(item, whole, f"{name}: {term_name}")
            derived = {term.name for term in template.terms if term.is_derived}
            self.assertEqual(set(envelope["details"]["terms"]) - derived, set(values), name)

    def test_leaving_out_any_term_a_printed_clause_states_is_refused_never_defaulted(self):
        for name, template in contract_templates().items():
            for term in template.terms:
                if term.is_derived or term.is_optional or term.name not in stated_terms(template):
                    continue
                values = complete_values(template)
                values.pop(term.name)
                envelope = self.run_form("merge", name, values)
                self.assertEqual(codes(envelope), [("MISSING_TERM", f"values.{term.name}")], f"{name}: {term.name}")
                self.assertFalse((self.directory / "out.docx").exists())

    def test_an_optional_party_detail_may_be_blank(self):
        values = complete_values(contract_templates()["kr/nda"]) | {"partyAAddress": "", "partyBRepresentative": ""}
        self.assertEqual(self.run_form("check", "kr/nda", values)["issues"], [])

    def test_a_value_naming_no_term_is_refused_with_the_term_it_resembles(self):
        values = complete_values(contract_templates()["kr/service-agreement"])
        values["latePenaltyCap"] = values.pop("latePenaltyCapPercent")
        envelope = self.run_form("merge", "kr/service-agreement", values)
        self.assertEqual(codes(envelope), [("UNKNOWN_TERM", "values.latePenaltyCap"), ("MISSING_TERM", "values.latePenaltyCapPercent")])
        self.assertIn("latePenaltyCapPercent", envelope["issues"][0]["suggestion"])

    def test_a_term_the_template_does_not_have_points_to_writing_it_as_a_clause(self):
        values = complete_values(contract_templates()["kr/mou"]) | {"governingLanguage": "한국어"}
        issue = self.run_form("check", "kr/mou", values)["issues"][0]
        self.assertEqual(issue["code"], "UNKNOWN_TERM")
        self.assertIn("addedClauses", issue["suggestion"])

    def test_a_number_term_holds_the_number_alone_and_a_choice_one_of_its_options(self):
        values = complete_values(contract_templates()["kr/service-agreement"]) | {"latePenaltyPerMille": "1.5/1000", "vatNote": "면제"}
        self.assertEqual(sorted(codes(self.run_form("check", "kr/service-agreement", values))), [("INVALID_TERM", "values.latePenaltyPerMille"), ("INVALID_TERM", "values.vatNote")])
        values |= {"latePenaltyPerMille": 1.5, "vatNote": "포함"}
        self.assertEqual(self.run_form("check", "kr/service-agreement", values)["issues"], [])

    def test_an_amount_in_words_is_written_from_its_amount_and_never_given(self):
        values = complete_values(contract_templates()["kr/service-agreement"]) | {"totalAmount": "50,000,000"}
        _, text = self.merged_text("kr/service-agreement", values)
        self.assertIn("① 본 용역의 계약금액은 금 오천만원整 (₩50,000,000, 부가가치세 포함)으로 한다.", self.paragraphs(text))
        given = self.run_form("check", "kr/service-agreement", values | {"totalAmountKorean": "오천만"})
        self.assertEqual(codes(given), [("INVALID_TERM", "values.totalAmountKorean")])


class ClauseTest(ContractRun):
    def test_a_replacing_clause_takes_the_template_clause_place_number_and_heading(self):
        template = contract_templates()["kr/nda"]
        values = complete_values(template) | {"clauses": {"returnOfInformation": {"paragraphs": ["정보수령자는 요청을 받은 날로부터 {{ returnDays }}일 이내에 반환하거나 폐기하고 확인서를 제출한다."]}}}
        envelope, text = self.merged_text("kr/nda", values)
        paragraphs = self.paragraphs(text)
        position = paragraphs.index("제8조 (비밀정보의 반환 등)")
        self.assertEqual(paragraphs[position + 1], f"정보수령자는 요청을 받은 날로부터 {values['returnDays']}일 이내에 반환하거나 폐기하고 확인서를 제출한다.")
        self.assertEqual(paragraphs[position + 2], "제9조 (권리의 부존재 등)")
        self.assertEqual(len(self.headings(text)), len(template.articles))
        self.assertIn({"number": "제8조", "key": "returnOfInformation", "heading": "비밀정보의 반환 등", "source": "replaced"}, envelope["details"]["clauses"])

    def test_a_replacing_clause_that_drops_a_given_term_is_refused(self):
        values = complete_values(contract_templates()["kr/service-agreement"]) | {"clauses": {"confidentiality": {"paragraphs": ["양 당사자는 계약 종료 후 5년간 비밀을 유지한다."]}}}
        envelope = self.run_form("merge", "kr/service-agreement", values)
        self.assertEqual(codes(envelope), [("TERM_NOT_STATED", "values.confidentialitySurvivalYears")])
        self.assertIn("clauses.confidentiality replaced", envelope["issues"][0]["message"])

    def test_a_replaced_clause_whose_term_is_left_out_no_longer_needs_it(self):
        values = complete_values(contract_templates()["kr/service-agreement"])
        values.pop("latePenaltyCapPercent")
        values["clauses"] = {"latePenalty": {"paragraphs": ["지체상금은 매 지체일수마다 계약금액의 1,000분의 {{ latePenaltyPerMille }}로 하며 상한을 두지 아니한다."]}}
        self.assertEqual(self.run_form("check", "kr/service-agreement", values)["issues"], [])

    def test_an_added_clause_cannot_repeat_a_clause_the_contract_already_has(self):
        values = complete_values(contract_templates()["kr/service-agreement"])
        envelope = self.run_form("merge", "kr/service-agreement", values | {"addedClauses": [{"key": "confidentiality", "heading": "비밀유지", "paragraphs": ["비밀을 유지한다."]}]})
        self.assertEqual(codes(envelope), [("DUPLICATE_CLAUSE", "addedClauses[0].key")])
        self.assertIn("clauses.confidentiality", envelope["issues"][0]["suggestion"])
        twice = [{"key": "audit", "heading": "감사", "paragraphs": ["감사한다."]}] * 2
        self.assertEqual(codes(self.run_form("check", "kr/service-agreement", values | {"addedClauses": twice})), [("DUPLICATE_CLAUSE", "addedClauses[1].key")])

    def test_an_added_clause_goes_after_the_clause_it_names_and_references_follow_the_numbers(self):
        added = {"key": "dataSharing", "heading": "자료 공유", "paragraphs": ["양 기관은 {{ article:cooperation }}의 협력을 위하여 자료를 공유한다."], "after": "purpose"}
        envelope, text = self.merged_text("kr/mou", complete_values(contract_templates()["kr/mou"]) | {"addedClauses": [added]})
        self.assertEqual(self.headings(text)[1:3], ["제2조 (자료 공유)", "제3조 (협력분야)"])
        paragraphs = self.paragraphs(text)
        self.assertIn("양 기관은 제3조의 협력을 위하여 자료를 공유한다.", paragraphs)
        self.assertTrue(any("제5조의 실무협의회" in paragraph for paragraph in paragraphs))
        self.assertEqual(envelope["details"]["clauses"][1]["source"], "added")

    def test_a_removed_clause_takes_its_terms_and_its_number_with_it(self):
        template = contract_templates()["kr/service-agreement"]
        values = complete_values(template)
        values.pop("warrantyMonths")
        _, text = self.merged_text("kr/service-agreement", values | {"removedClauses": ["warranty"]})
        self.assertEqual(len(self.headings(text)), len(template.articles) - 1)
        given = self.run_form("check", "kr/service-agreement", complete_values(template) | {"removedClauses": ["warranty"]})
        self.assertEqual(codes(given), [("TERM_NOT_STATED", "values.warrantyMonths")])
        self.assertIn("removedClauses removes warranty", given["issues"][0]["message"])
        values.pop("scopeItems")
        referenced = self.run_form("check", "kr/service-agreement", values | {"removedClauses": ["warranty", "scope"]})
        self.assertEqual(codes(referenced), [("UNKNOWN_CLAUSE", "template.duties")])

    def test_an_unknown_clause_key_is_refused_with_the_key_it_resembles(self):
        values = complete_values(contract_templates()["kr/nda"]) | {"clauses": {"returnInformation": {"paragraphs": ["반환한다."]}}}
        envelope = self.run_form("check", "kr/nda", values)
        self.assertEqual(codes(envelope), [("UNKNOWN_CLAUSE", "clauses.returnInformation")])
        self.assertEqual(envelope["issues"][0]["suggestion"], "use returnOfInformation")

    def test_a_clause_text_naming_an_unknown_term_is_refused(self):
        values = complete_values(contract_templates()["kr/nda"]) | {"addedClauses": [{"key": "injunction", "heading": "가처분", "paragraphs": ["{{ injunctionCourt }}에 가처분을 신청할 수 있다."]}]}
        self.assertEqual(codes(self.run_form("check", "kr/nda", values)), [("UNKNOWN_TERM", "addedClauses[0]")])

    def test_check_reports_what_merge_would_refuse_without_writing(self):
        values = complete_values(contract_templates()["kr/nda"])
        values.pop("oralConfirmationDays")
        self.assertEqual(codes(self.run_form("check", "kr/nda", values)), [("MISSING_TERM", "values.oralConfirmationDays")])
        self.assertFalse((self.directory / "out.docx").exists())


class LayoutTest(ContractRun):
    def test_each_list_item_is_a_numbered_paragraph_and_each_list_restarts_at_one(self):
        values = complete_values(contract_templates()["kr/service-agreement"]) | {"scopeItems": ["설계", "구축", "운영"], "payments": ["선급금 30%", "잔금 70%"], "deliverables": ["결과 보고서"]}
        _, text = self.merged_text("kr/service-agreement", values)
        numbered = {}
        for paragraph, number in text["paragraphs"]:
            if number is not None:
                numbered.setdefault(number, []).append(paragraph)
        self.assertEqual(list(numbered.values()), [["설계", "구축", "운영"], ["선급금 30%", "잔금 70%"], ["결과 보고서"]])

    def test_a_blank_conditional_term_leaves_its_paragraph_out(self):
        values = complete_values(contract_templates()["kr/offer-letter"]) | {"benefits": ["식대 지원", "자기계발비"], "equity": "1,000주", "probationNote": ""}
        _, text = self.merged_text("kr/offer-letter", values)
        paragraphs = self.paragraphs(text)
        self.assertIn("스톡옵션: 1,000주", paragraphs)
        self.assertEqual([paragraph for paragraph in paragraphs if paragraph.startswith("- ")], ["- 식대 지원", "- 자기계발비"])
        _, with_note = self.merged_text("kr/offer-letter", values | {"probationNote": "수습 3개월"})
        self.assertEqual(len(paragraphs), len(with_note["paragraphs"]) - 1)

    def test_an_employment_contract_without_an_end_date_has_no_contract_period(self):
        values = complete_values(contract_templates()["kr/employment-contract"]) | {"startDate": "2026-11-01", "endDate": ""}
        _, open_ended = self.merged_text("kr/employment-contract", values)
        self.assertIn("근로개시일: 2026-11-01", self.paragraphs(open_ended))
        self.assertFalse(any(paragraph.startswith("근로계약기간") for paragraph in self.paragraphs(open_ended)))
        _, fixed_term = self.merged_text("kr/employment-contract", values | {"endDate": "2027-10-31"})
        self.assertIn("근로계약기간: 2026-11-01부터 2027-10-31까지", self.paragraphs(fixed_term))


if __name__ == "__main__":
    unittest.main()
