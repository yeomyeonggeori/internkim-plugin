import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader

from render_fixture import can_render

SCRIPTS_PATH = Path(__file__).resolve().parents[1] / "skills" / "office" / "scripts"
OFFICE_ENTRY = SCRIPTS_PATH / "office"
sys.path.insert(0, str(SCRIPTS_PATH))

from staged_deck_fixture import write_staged_deck  # noqa: E402
from fonts.registry import FAMILIES, resolved_face  # noqa: E402

KOREAN_TEXT = "일금 일백만원整 다람쥐 헌 쳇바퀴에 타고파"
DECK_PAGE = f'<h1>{KOREAN_TEXT}</h1><p style="font-size: 32px">본문 문장입니다.</p>'
DECK_STYLE = "section { justify-content: center; font-family: system-ui, sans-serif; } h1 { font-size: 72px; }"


def fontconfig_without_fonts(directory: Path) -> Path:
    empty = directory / "no-fonts"
    empty.mkdir()
    configuration = directory / "fonts.conf"
    configuration.write_text(f'<?xml version="1.0"?>\n<fontconfig><dir>{empty}</dir><cachedir>{directory / "fontconfig-cache"}</cachedir></fontconfig>\n', encoding="utf-8")
    return configuration


def embedded_font_names(pdf_path: Path) -> set[str]:
    names = set()
    for page in PdfReader(str(pdf_path)).pages:
        fonts = page["/Resources"].get("/Font", {})
        for key in fonts:
            font = fonts[key].get_object()
            descendants = font.get("/DescendantFonts")
            face = descendants[0].get_object() if descendants else font
            names.add(str(face["/BaseFont"]).lstrip("/").split("+", 1)[-1])
    return names


def pdf_text(pdf_path: Path) -> str:
    return "".join(page.extract_text() for page in PdfReader(str(pdf_path)).pages)


@unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
class HostWithoutFontsTest(unittest.TestCase):
    def setUp(self):
        self.directory = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: shutil.rmtree(self.directory, ignore_errors=True))
        self.environment = {**os.environ, "FONTCONFIG_FILE": str(fontconfig_without_fonts(self.directory))}

    def run_office(self, *arguments) -> dict:
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), *arguments], capture_output=True, text=True, cwd=self.directory, env=self.environment)
        envelope = json.loads(completed.stdout)
        self.assertNotEqual(envelope["status"], "error", envelope)
        return envelope

    def assert_bundled_pdf(self, pdf_path: Path, text: str) -> None:
        names = embedded_font_names(pdf_path)
        self.assertTrue(names)
        self.assertEqual({name for name in names if resolved_face(name).is_substitute}, set())
        self.assertIn(text, pdf_text(pdf_path).replace("\n", " "))

    def assert_bundled_preview(self, envelope: dict) -> None:
        paths = {Path(font["path"]) for font in envelope["details"]["previewFonts"]}
        self.assertTrue(paths)
        self.assertLessEqual(paths, {family.path(face) for family in FAMILIES for face in family.faces})
        self.assertTrue(envelope["details"]["pages"])

    def test_a_deck_draws_korean_with_bundled_fonts(self):
        write_staged_deck(self.directory, [DECK_PAGE], DECK_STYLE)
        envelope = self.run_office("create", "build/deck.pdf", ".")
        self.assert_bundled_pdf(self.directory / envelope["outputPath"], KOREAN_TEXT)

    def test_a_document_draws_korean_with_bundled_fonts(self):
        document = Document()
        document.add_paragraph(KOREAN_TEXT)
        document.save(self.directory / "문서.docx")
        envelope = self.run_office("render", "문서.docx")
        self.assert_bundled_preview(envelope)
        self.assert_bundled_pdf(self.directory / envelope["details"]["pdf"], KOREAN_TEXT)

    def test_a_workbook_draws_korean_with_bundled_fonts(self):
        workbook = Workbook()
        workbook.active["A1"] = KOREAN_TEXT
        workbook.save(self.directory / "표.xlsx")
        envelope = self.run_office("render", "표.xlsx")
        self.assert_bundled_preview(envelope)
        self.assert_bundled_pdf(self.directory / envelope["details"]["pdf"], KOREAN_TEXT)

    def test_pdf_create_draws_korean_with_bundled_fonts(self):
        (self.directory / "spec.json").write_text(json.dumps({"title": KOREAN_TEXT, "sections": [{"paragraphs": ["본문 문장입니다."]}]}, ensure_ascii=False), encoding="utf-8")
        envelope = self.run_office("create", "보고서.pdf", "spec.json")
        self.assert_bundled_pdf(self.directory / "보고서.pdf", KOREAN_TEXT)
        self.assertEqual(envelope["status"], "ok")


if __name__ == "__main__":
    unittest.main()
