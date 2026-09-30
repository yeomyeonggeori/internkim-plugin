import json
import subprocess
import sys
import textwrap
import unittest
import zipfile

from deck_fixture import DeckFixture, OFFICE_ENTRY, SCRIPTS_PATH
from doc_fixture import run_office, run_office_python, write_json


FIXTURE_DECK = """
from pptx import Presentation
presentation = Presentation()
title_layout, content_layout = presentation.slide_layouts[0], presentation.slide_layouts[1]
first = presentation.slides.add_slide(title_layout)
first.shapes.title.text = "2024 매출 보고"
first.placeholders[1].text = "여명거리 {{ quarter }}"
second = presentation.slides.add_slide(content_layout)
second.shapes.title.text = "핵심 지표"
body = second.placeholders[1].text_frame
body.text = "매출 100"
body.add_paragraph().text = "비용 50"
second.notes_slide.notes_text_frame.text = "둘째 슬라이드 노트"
third = presentation.slides.add_slide(content_layout)
third.shapes.title.text = "다음 단계"
presentation.save("fixture.pptx")
"""

INSPECT_DECK = """
import json, sys
from lxml import etree
from pptx import Presentation

presentation = Presentation(sys.argv[1])
def canonical(part):
    return etree.tostring(etree.fromstring(part.blob), method="c14n").decode()
print(json.dumps({
    "titles": [slide.shapes.title.text for slide in presentation.slides],
    "layouts": [slide.slide_layout.name for slide in presentation.slides],
    "masters": len(presentation.slide_masters),
    "layoutNames": [layout.name for layout in presentation.slide_layouts],
    "layoutXml": [canonical(layout.part) for layout in presentation.slide_layouts],
    "masterXml": [canonical(master.part) for master in presentation.slide_masters],
    "notes": [slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else None for slide in presentation.slides],
}, ensure_ascii=False))
"""


class DeckEditFixture(DeckFixture):
    def setUp(self):
        super().setUp()
        run_office_python(FIXTURE_DECK, self.directory)

    def inspect(self, name="fixture.pptx"):
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "-c", INSPECT_DECK, name], capture_output=True, text=True, check=True, cwd=self.directory)
        return json.loads(completed.stdout)

    def apply(self, operations, *options):
        write_json(self.directory / "ops.json", operations)
        return run_office(["deck", "apply", "fixture.pptx", "ops.json", *options], self.directory)


class ReadTest(DeckEditFixture):
    def test_slides_are_numbered_with_layout_shapes_and_notes(self):
        envelope = run_office(["deck", "read", "fixture.pptx"], self.directory)
        self.assertEqual(envelope["status"], "ok")
        slides = envelope["details"]["slides"]
        self.assertEqual([slide["slide"] for slide in slides], [1, 2, 3])
        self.assertEqual([slide["layout"] for slide in slides], ["Title Slide", "Title and Content", "Title and Content"])
        self.assertEqual(slides[0]["shapes"][0]["text"], "2024 매출 보고")
        self.assertEqual(slides[1]["shapes"][1]["text"], "매출 100\n비용 50")
        self.assertEqual([slide["notes"] for slide in slides], ["", "둘째 슬라이드 노트", ""])


class ApplyTest(DeckEditFixture):
    def test_a_batch_edits_text_notes_order_and_slides_against_the_numbers_read_before_it(self):
        envelope = self.apply([
            {"op": "set_text", "slide": 3, "shape": 0, "text": "다음 단계 요약"},
            {"op": "find_replace", "find": "{{ quarter }}", "replace": "3분기"},
            {"op": "set_notes", "slide": 1, "text": "첫 노트\n둘째 줄"},
            {"op": "delete_slide", "slide": 2},
            {"op": "reorder", "order": [3, 1]},
        ])
        self.assertEqual(envelope["status"], "ok")
        edited = self.inspect()
        self.assertEqual(edited["titles"], ["다음 단계 요약", "2024 매출 보고"])
        self.assertEqual(edited["notes"], [None, "첫 노트\n둘째 줄"])
        read = run_office(["deck", "read", "fixture.pptx"], self.directory)
        self.assertEqual(read["details"]["slides"][1]["shapes"][1]["text"], "여명거리 3분기")

    def test_editing_keeps_every_layout_and_master_and_each_slides_layout(self):
        before = self.inspect()
        self.apply([{"op": "set_text", "slide": 1, "shape": 0, "text": "바뀐 제목"}, {"op": "delete_slide", "slide": 3}])
        after = self.inspect()
        self.assertEqual(after["masters"], before["masters"])
        self.assertEqual(after["layoutNames"], before["layoutNames"])
        self.assertEqual(after["layoutXml"], before["layoutXml"])
        self.assertEqual(after["masterXml"], before["masterXml"])
        self.assertEqual(after["layouts"], before["layouts"][:2])

    def test_a_replacement_keeps_the_formatting_of_the_run_it_starts_in(self):
        self.apply([{"op": "set_text", "slide": 1, "shape": 1, "text": "첫 줄\n둘째 줄"}])
        read = run_office(["deck", "read", "fixture.pptx"], self.directory)
        self.assertEqual(read["details"]["slides"][0]["shapes"][1]["text"], "첫 줄\n둘째 줄")

    def test_one_bad_operation_leaves_the_file_untouched(self):
        original = (self.directory / "fixture.pptx").read_bytes()
        envelope = self.apply([
            {"op": "set_notes", "slide": 1, "text": "바뀜"},
            {"op": "set_text", "slide": 9, "shape": 0, "text": "없음"},
        ])
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["TARGET_NOT_FOUND"])
        self.assertEqual(envelope["issues"][0]["location"], "ops[1].slide")
        self.assertEqual((self.directory / "fixture.pptx").read_bytes(), original)

    def test_a_dry_run_reports_changes_and_writes_nothing(self):
        original = (self.directory / "fixture.pptx").read_bytes()
        envelope = self.apply([{"op": "delete_slide", "slide": 1}], "--dry-run")
        self.assertTrue(envelope["details"]["dryRun"])
        self.assertEqual(len(envelope["details"]["changes"]), 1)
        self.assertEqual((self.directory / "fixture.pptx").read_bytes(), original)

    def test_output_writes_a_copy_and_leaves_the_source(self):
        original = (self.directory / "fixture.pptx").read_bytes()
        self.apply([{"op": "delete_slide", "slide": 1}], "--output", "edited.pptx")
        self.assertEqual((self.directory / "fixture.pptx").read_bytes(), original)
        self.assertEqual(self.inspect("edited.pptx")["titles"], ["핵심 지표", "다음 단계"])

    def test_a_deleted_slide_cannot_be_used_and_a_used_slide_cannot_be_deleted(self):
        used_then_deleted = self.apply([{"op": "set_notes", "slide": 2, "text": "바뀜"}, {"op": "delete_slide", "slide": 2}])
        deleted_then_used = self.apply([{"op": "delete_slide", "slide": 2}, {"op": "set_notes", "slide": 2, "text": "바뀜"}])
        self.assertEqual([issue["code"] for issue in used_then_deleted["issues"]], ["OPERATION_NOT_APPLICABLE"])
        self.assertEqual([issue["code"] for issue in deleted_then_used["issues"]], ["OPERATION_NOT_APPLICABLE"])

    def test_a_reorder_must_list_exactly_the_slides_that_remain(self):
        envelope = self.apply([{"op": "reorder", "order": [3, 1]}])
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["INVALID_VALUE"])

    def test_text_cannot_be_set_on_a_shape_without_a_text_frame(self):
        envelope = self.apply([{"op": "set_text", "slide": 1, "shape": 7, "text": "없음"}])
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["TARGET_NOT_FOUND"])

    def test_untouched_parts_keep_their_content_when_only_notes_change(self):
        self.apply([{"op": "set_notes", "slide": 1, "text": "노트"}], "--output", "edited.pptx")
        with zipfile.ZipFile(self.directory / "fixture.pptx") as source, zipfile.ZipFile(self.directory / "edited.pptx") as edited:
            changed = [name for name in source.namelist() if name.startswith("ppt/slideLayouts/") and source.read(name) != edited.read(name)]
        self.assertEqual(changed, [])


class GuideTest(unittest.TestCase):
    def test_the_guide_lists_every_operation_and_its_fields(self):
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "guide", "deck"], capture_output=True, text=True, check=True)
        for operation in ("set_text", "find_replace", "set_notes", "delete_slide", "reorder"):
            self.assertIn(f'op "{operation}"', completed.stdout)
        self.assertIn("deck apply <file.pptx> <ops.json>", completed.stdout)
        self.assertIn("TARGET_NOT_FOUND", completed.stdout)

    def test_every_declared_operation_has_a_planner(self):
        code = """
        import json, sys
        sys.path.insert(0, sys.argv[1])
        from pptx_operations import PPTX_OPERATIONS as operations
        print(json.dumps(sorted(operations.planners) == sorted(record.name for record in operations.shape.records)))
        """
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "-c", textwrap.dedent(code), str(SCRIPTS_PATH / "deck")], capture_output=True, text=True, check=True)
        self.assertEqual(completed.stdout.strip(), "true")
