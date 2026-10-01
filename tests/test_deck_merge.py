import json
from pathlib import Path
import subprocess
import sys
import zipfile

from deck_fixture import OFFICE_ENTRY, DeckFixture, run_office_python_json


TEMPLATE = """
from pptx import Presentation
from pptx.util import Inches
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[5])
title = slide.shapes.title.text_frame.paragraphs[0]
title.add_run().text = "{{ project.na"
title.add_run().text = "me }} 분기 보고"
title.runs[1].font.bold = True
table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.2)).table
for column, (header, cell) in enumerate((("품목", "{{ items.name }}"), ("금액", "{{ items.amount }}"))):
    table.cell(0, column).text = header
    table.cell(1, column).text = cell
table.cell(2, 0).text = "합계"
table.cell(2, 1).text = "{{ total }}"
slide.notes_slide.notes_text_frame.text = "발표자: {{ presenter }}"
presentation.save("report.pptx")
"""

READ = """
import json
from pptx import Presentation
slide = Presentation("filled.pptx").slides[0]
table = next(shape for shape in slide.shapes if shape.has_table)
print(json.dumps({
    "title": slide.shapes.title.text_frame.text,
    "rows": [[cell.text for cell in row.cells] for row in table.table.rows],
    "height": table.height,
    "notes": slide.notes_slide.notes_text_frame.text,
}, ensure_ascii=False))
"""


class DeckMergeTest(DeckFixture):
    def setUp(self):
        super().setUp()
        run_office_python_json(TEMPLATE + "\nprint('{}')", self.directory)

    def merge(self, values):
        (self.directory / "values.json").write_text(json.dumps(values, ensure_ascii=False), encoding="utf-8")
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "deck", "merge", "report.pptx", "values.json", "filled.pptx"], capture_output=True, text=True, cwd=self.directory)
        return json.loads(completed.stdout)

    def test_a_split_title_fills_and_a_table_row_repeats_per_item(self):
        envelope = self.merge({"project": {"name": "샘플 프로젝트"}, "items": [{"name": "설계", "amount": "1,200만"}, {"name": "구축", "amount": "3,400만"}], "total": "4,600만", "presenter": "이샘플"})
        self.assertEqual(envelope["status"], "ok", envelope)
        filled = run_office_python_json(READ, self.directory)
        self.assertEqual(filled["title"], "샘플 프로젝트 분기 보고")
        self.assertEqual(filled["rows"], [["품목", "금액"], ["설계", "1,200만"], ["구축", "3,400만"], ["합계", "4,600만"]])
        self.assertEqual(filled["notes"], "발표자: 이샘플")
        self.assertEqual(filled["height"], 1463040)

    def test_a_missing_value_writes_nothing_and_untouched_parts_stay(self):
        envelope = self.merge({"project": {"name": "샘플"}, "items": [], "total": "0"})
        self.assertEqual([(issue["code"], issue["location"]) for issue in envelope["issues"]], [("UNRESOLVED_PLACEHOLDER", "notes of slide 1")])
        self.assertFalse(Path(self.directory / "filled.pptx").exists())
        self.merge({"project": {"name": "샘플"}, "items": [], "total": "0", "presenter": "박예시"})
        with zipfile.ZipFile(self.directory / "report.pptx") as before, zipfile.ZipFile(self.directory / "filled.pptx") as after:
            changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
        self.assertEqual(changed, ["ppt/slides/slide1.xml", "ppt/notesSlides/notesSlide1.xml"])
