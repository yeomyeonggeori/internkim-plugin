from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn

from doc_fixture import run_office


def korean_document(path: Path, east_asia_font: str | None) -> None:
    document = Document()
    run = document.add_paragraph().add_run("주식회사 예시상사 연간 유지보수 계약")
    if east_asia_font:
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia_font)
    else:
        for fonts in document.styles.element.iter(qn("w:rFonts")):
            for attribute in ("w:eastAsia", "w:eastAsiaTheme"):
                fonts.attrib.pop(qn(attribute), None)
    document.save(path)


def font_codes(envelope) -> set[str]:
    return {issue["code"] for issue in envelope["issues"]} & {"EAST_ASIA_FONT_MISSING", "KOREAN_FONT_MISSING"}


class EastAsiaFontTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)

    def tearDown(self):
        self.temporary_directory.cleanup()

    def test_a_named_east_asian_font_passes_the_check(self):
        korean_document(self.directory / "named.docx", "바탕")
        self.assertEqual(font_codes(run_office(["check", "named.docx"], self.directory)), set())

    def test_a_missing_east_asian_font_is_reported_once(self):
        korean_document(self.directory / "unnamed.docx", None)
        envelope = run_office(["check", "unnamed.docx"], self.directory)
        self.assertEqual([issue["code"] for issue in envelope["issues"] if issue["code"] == "EAST_ASIA_FONT_MISSING"], ["EAST_ASIA_FONT_MISSING"])

if __name__ == "__main__":
    unittest.main()
