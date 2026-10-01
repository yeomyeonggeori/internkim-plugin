import datetime
import unittest

from openpyxl import load_workbook

from sheet_fixture import WorkbookFixture, run_office, run_office_python, stored_cells, write_json

FIXTURE_WORKBOOK = """
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

workbook = Workbook()
sales = workbook.active
sales.title = "Sales"
sales["A1"] = 10
sales["A2"] = 0
sales["B1"] = "=A1/A2"
sales["B2"] = "=Gone!A1"
sales["C1"] = 123456789
sales["C1"].number_format = "#,##0"
sales["D1"] = "{{customer_name}} 님"
sales["E1"] = "=SUM(A1:A2)"
workbook.defined_names["Orphan"] = DefinedName("Orphan", attr_text="Gone!$A$1")
workbook.defined_names["Fine"] = DefinedName("Fine", attr_text="Sales!$A$1")
workbook.save("fixture.xlsx")
"""


class SheetCheckTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        run_office_python(FIXTURE_WORKBOOK, self.directory)

    def check(self):
        return run_office(["sheet", "check", "fixture.xlsx"], self.directory)

    def findings(self):
        return {(issue["code"], issue["location"]) for issue in self.check()["issues"]}

    def test_each_problem_is_reported_once_where_it_is(self):
        self.assertEqual(self.findings(), {
            ("FORMULA_ERROR", "Sales!B1"),
            ("MISSING_SHEET_REFERENCE", "Sales!B2"),
            ("BROKEN_DEFINED_NAME", "Orphan"),
            ("NUMBER_TOO_WIDE", "Sales!C1"),
            ("PLACEHOLDER_LEFT", "Sales!D1"),
        })

    def test_the_computed_error_is_named(self):
        message = next(issue["message"] for issue in self.check()["issues"] if issue["code"] == "FORMULA_ERROR")
        self.assertIn("#DIV/0!", message)

    def test_the_suggested_width_fixes_the_narrow_column(self):
        fix = next(issue["fix"] for issue in self.check()["issues"] if issue["code"] == "NUMBER_TOO_WIDE")
        self.assertEqual([{key: operation[key] for key in ("op", "sheet", "column")} for operation in fix], [{"op": "set_column_width", "sheet": "Sales", "column": "C"}])
        self.assertEqual(self.apply(fix, name="fixture.xlsx")["status"], "ok")
        self.assertNotIn(("NUMBER_TOO_WIDE", "Sales!C1"), self.findings())

    def test_a_number_that_fits_is_left_alone(self):
        write_json(self.directory / "ops.json", [{"op": "set_column_width", "column": "C", "width": 14}])
        run_office(["sheet", "apply", "fixture.xlsx", "ops.json"], self.directory)
        self.assertNotIn(("NUMBER_TOO_WIDE", "Sales!C1"), self.findings())

    def test_a_created_workbook_sizes_its_columns_for_its_numbers(self):
        self.create_workbook([{"title": "표", "rows": [["이름", "금액"], ["박예시", 1500000]]}])
        envelope = run_office(["sheet", "check", "book.xlsx"], self.directory)
        self.assertEqual(envelope["issues"], [])

    def test_validate_no_longer_reports_formula_errors(self):
        codes = {issue["code"] for issue in run_office(["sheet", "validate", "fixture.xlsx"], self.directory)["issues"]}
        self.assertNotIn("FORMULA_ERROR_MARKER", codes)


if __name__ == "__main__":
    unittest.main()


TAMPER_CACHED_VALUES = r"""
import re
import zipfile

with zipfile.ZipFile("book.xlsx") as source:
    entries = {info.filename: source.read(info.filename) for info in source.infolist()}
sheet = entries["xl/worksheets/sheet1.xml"].decode()
sheet = re.sub(r'(<c r="C2"[^>]*><f>[^<]*</f><v>)[^<]*', r'\g<1>25', sheet)
sheet = re.sub(r'(<c r="D2"[^>]*><f>[^<]*</f><v>)[^<]*', r'\g<1>stale', sheet)
entries["xl/worksheets/sheet1.xml"] = sheet.encode()
with zipfile.ZipFile("book.xlsx", "w") as target:
    for name, content in entries.items():
        target.writestr(name, content)
"""


class StaleCachedValueTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        self.create_workbook([{"title": "Sales", "rows": [["a", "b", "double", "label"], [10, 5, "=A2*2", '="n"&B2'], [1, 2, "=A3*2", '="n"&B3']]}])

    def findings(self):
        envelope = run_office(["sheet", "check", "book.xlsx"], self.directory)
        return [(issue["code"], issue["location"]) for issue in envelope["issues"]]

    def test_a_workbook_written_here_has_no_stale_value(self):
        self.assertEqual(self.findings(), [])

    def test_a_stored_value_that_differs_from_the_computed_one_is_reported_and_recalculated(self):
        run_office_python(TAMPER_CACHED_VALUES, self.directory)
        envelope = run_office(["sheet", "check", "book.xlsx"], self.directory)
        issue = next(issue for issue in envelope["issues"] if issue["code"] == "STALE_CACHED_VALUE")
        self.assertEqual(issue["location"], "Sales!C2")
        self.assertIn("Sales!C2 stores 25 but computes 20", issue["message"])
        self.assertIn("Sales!D2", issue["message"])
        self.assertEqual(self.apply(issue["fix"])["status"], "ok")
        self.assertEqual(self.findings(), [])
        self.assertEqual(stored_cells(self.directory / "book.xlsx")["C2"]["value"], "20")


CHART_FIXTURE = """
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

workbook = Workbook()
sales = workbook.active
sales.title = "Sales"
for row in [["month", "amount"], ["1월", 10], ["2월", 20]]:
    sales.append(row)
chart = BarChart()
chart.add_data(Reference(sales, min_col=2, min_row=1, max_row=3), titles_from_data=True)
sales.add_chart(chart, "E2")
empty = BarChart()
empty.add_data(Reference(sales, min_col=8, min_row=1, max_row=3), titles_from_data=True)
sales.add_chart(empty, "E20")
workbook.save("charts.xlsx")
"""


class ChartReferenceTest(WorkbookFixture):
    def test_a_chart_reading_an_empty_range_is_reported_and_one_with_data_is_not(self):
        run_office_python(CHART_FIXTURE, self.directory)
        issues = run_office(["sheet", "check", "charts.xlsx"], self.directory)["issues"]
        self.assertEqual([(issue["code"], issue["location"]) for issue in issues], [("CHART_REFERENCE_BROKEN", "Sales chart 1")])
        self.assertIn("$H$2:$H$3", issues[0]["message"])


class ChartValueTest(WorkbookFixture):
    def test_a_series_that_reads_only_text_is_reported(self):
        run_office_python("""
            from openpyxl import Workbook
            from openpyxl.chart import BarChart, Reference
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "실적"
            for row in [["담당자", "지역", "1월"], ["이샘플", "수도권", 120], ["박예시", "영남", 90]]:
                sheet.append(row)
            chart = BarChart()
            chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=3), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
            sheet.add_chart(chart, "E2")
            workbook.save("charts.xlsx")
        """, self.directory)
        issues = run_office(["sheet", "check", "charts.xlsx"], self.directory)["issues"]
        self.assertEqual([(issue["code"], issue["location"]) for issue in issues], [("CHART_REFERENCE_BROKEN", "실적 chart 0")])
        self.assertIn("$B$2:$B$3, which holds text and no number", issues[0]["message"])


class FormulaVisibilityTest(WorkbookFixture):
    def create(self):
        self.create_workbook([
            {"title": "Data", "rows": [["key", "amount"], ["a", 1], ["b", 2], ["a", 3]]},
            {"title": "T", "rows": [
                ["what", "formula"],
                ["indirect", '=INDIRECT("Data!B2")'],
                ["unique", "=ROWS(UNIQUE(Data!A2:A4))"],
                ["offset", "=SUM(OFFSET(Data!B2,0,0,3,1))"],
                ["unsupported", '=WEBSERVICE("https://example.com")'],
            ]},
        ])

    def issues(self):
        return {issue["code"]: issue for issue in run_office(["sheet", "check", "book.xlsx"], self.directory)["issues"]}

    def test_indirect_and_rows_of_an_array_compute_beside_a_formula_that_cannot(self):
        self.create()
        issues = self.issues()
        self.assertNotIn("FORMULA_ERROR", issues)
        self.assertEqual(issues["FORMULA_NOT_EVALUATED"]["message"], "1 formula cells have no computed value: T!B5")
        self.assertEqual(self.apply([{"op": "recalculate"}])["status"], "warning")
        self.assertEqual(issues["FORMULA_NOT_EVALUATED"]["message"], self.issues()["FORMULA_NOT_EVALUATED"]["message"])

    def test_read_shows_a_dynamic_array_formula_as_written(self):
        self.create()
        cells = run_office(["sheet", "read", "book.xlsx", "--sheet", "T", "--where", "formula"], self.directory)["details"]["range"]["cells"]
        self.assertEqual([(cell["formula"], cell["value"]) for cell in cells[:3]], [('=INDIRECT("Data!B2")', 1), ("=ROWS(UNIQUE(Data!A2:A4))", 2), ("=SUM(OFFSET(Data!B2,0,0,3,1))", 6)])

    def test_a_spec_whose_formula_reads_its_own_cell_creates_nothing(self):
        write_json(self.directory / "spec.json", {"sheets": [{"title": "S", "rows": [["total", "=SUM(B1:B2)"], ["x", 5]], "autoFilter": False}]})
        envelope = run_office(["sheet", "create", "book.xlsx", "--spec", "spec.json"], self.directory)
        self.assertEqual((envelope["status"], envelope["outputPath"], [issue["location"] for issue in envelope["issues"]]), ("error", None, ["S!B1"]))
        self.assertFalse((self.directory / "book.xlsx").exists())

    def test_a_spec_formula_that_does_not_parse_creates_nothing(self):
        write_json(self.directory / "spec.json", {"sheets": [{"title": "S", "rows": [["a", "b"], [1, 2], ["total", "=SUM(B2:B2))"]]}]})
        envelope = run_office(["sheet", "create", "book.xlsx", "--spec", "spec.json"], self.directory)
        self.assertEqual([(issue["code"], issue["location"]) for issue in envelope["issues"]], [("FORMULA_SYNTAX", "spec.sheets[0].rows[2][1]")])
        self.assertFalse((self.directory / "book.xlsx").exists())

    def test_a_formula_reading_its_own_cell_is_reported_as_circular(self):
        run_office_python("""
            from openpyxl import Workbook
            workbook = Workbook()
            workbook.active.title = "S"
            workbook.active.append(["total", "=SUM(B1:B2)"])
            workbook.active.append(["x", 5])
            workbook.save("book.xlsx")
        """, self.directory)
        self.assertEqual([(code, issue["location"]) for code, issue in self.issues().items()], [("CIRCULAR_REFERENCE", "S!B1")])


class FormulaNameTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        self.create_workbook([{"title": "실적", "rows": [["담당", "매출"], ["이샘플", 1200], ["박예시", 980]]}])

    def test_an_unknown_function_and_a_missing_sheet_are_refused_where_they_are_written(self):
        envelope = self.apply([
            {"op": "set_cell", "sheet": "실적", "cell": "C2", "value": "=SUMM(B2:B3)"},
            {"op": "set_cell", "sheet": "실적", "cell": "C3", "value": "=Missing!A1"},
        ])
        self.assertEqual(envelope["status"], "error")
        self.assertEqual([(issue["code"], issue["location"]) for issue in envelope["issues"]], [("MISSING_SHEET_REFERENCE", "실적!C3"), ("UNKNOWN_FUNCTION", "실적!C2")])
        self.assertIn("did you mean 'SUM'?", envelope["issues"][1]["message"])
        self.assertEqual(envelope["issues"][1]["suggestion"], "write SUM in place of SUMM")
        self.assertNotIn("C2", self.cells())

    def test_a_spec_formula_reading_a_later_sheet_and_a_let_function_are_written(self):
        write_json(self.directory / "spec.json", {"sheets": [
            {"title": "요약", "rows": [["합계", "=SUM(실적!B2:B3)"], ["두배", "=LET(double,LAMBDA(x,x*2),double(B1))"]]},
            {"title": "실적", "rows": [["담당", "매출"], ["이샘플", 1200], ["박예시", 980]]},
        ]})
        envelope = run_office(["sheet", "create", "later.xlsx", "--spec", "spec.json"], self.directory)
        self.assertNotIn("UNKNOWN_FUNCTION", [issue["code"] for issue in envelope["issues"]])
        self.assertNotEqual(envelope["status"], "error", envelope["issues"])

    def test_a_name_already_broken_in_the_workbook_does_not_stop_another_edit(self):
        run_office_python(FIXTURE_WORKBOOK, self.directory)
        envelope = self.apply([{"op": "set_cell", "sheet": "Sales", "cell": "F1", "value": "=Gone!B1+A1"}], name="fixture.xlsx")
        self.assertNotEqual(envelope["status"], "error", envelope["issues"])
        codes = {issue["code"]: issue for issue in run_office(["sheet", "check", "fixture.xlsx"], self.directory)["issues"]}
        self.assertEqual(codes["MISSING_SHEET_REFERENCE"]["location"], "Sales!B2")


class TextValueTest(WorkbookFixture):
    def check(self):
        return run_office(["sheet", "check", "book.xlsx"], self.directory)

    def text_issues(self):
        return [issue for issue in self.check()["issues"] if issue["code"] == "VALUE_STORED_AS_TEXT"]

    def test_numbers_and_a_formula_written_as_text_are_reported_with_operations_that_fix_them(self):
        self.create_workbook([{"title": "매출", "rows": [["월", "매출", "비율"], ["1월", "1,200", "12.5%"], ["2월", "₩1,350", 0.2], ["합계", "SUM(B2:B3)", None]]}])
        issues = self.text_issues()
        self.assertEqual([issue["location"] for issue in issues], ["매출!B2", "매출!B3", "매출!B4", "매출!C2"])
        self.assertEqual(issues[0]["fix"], [{"op": "set_cell", "sheet": "매출", "cell": "B2", "value": 1200}, {"op": "format_range", "sheet": "매출", "range": "B2", "numberFormat": "#,##0"}])
        self.assertEqual(issues[2]["fix"], [{"op": "set_cell", "sheet": "매출", "cell": "B4", "value": "=SUM(B2:B3)"}])
        self.assertIn("missing its =", issues[2]["message"])
        self.assertEqual(issues[3]["fix"][0]["value"], 0.125)
        envelope = self.apply([operation for issue in issues for operation in issue["fix"]])
        self.assertEqual(envelope["status"], "ok", envelope)
        self.assertEqual(self.text_issues(), [])
        values = run_office(["sheet", "read", "book.xlsx", "--range", "B2:C4"], self.directory)["details"]["range"]["values"]
        self.assertEqual(values, [[1200, 0.125], [1350, 0.2], [2550, None]])

    def test_a_column_whose_every_cell_reads_as_a_number_is_reported_without_a_number_beside_it(self):
        self.create_workbook([{"title": "실적", "rows": [["담당", "건수", "메모"], ["이샘플", "12", "가"], ["박예시", "9", "나"], ["최견본", "7", "다"]]}])
        issues = self.text_issues()
        self.assertEqual([issue["location"] for issue in issues], ["실적!B2:B4"])
        self.assertEqual(issues[0]["fix"], [{"op": "set_range", "sheet": "실적", "cell": "B2", "values": [[12], [9], [7]]}])
        self.assertEqual(self.apply(issues[0]["fix"])["status"], "ok")
        self.assertEqual(self.text_issues(), [])

    def test_codes_with_leading_zeros_and_labels_in_a_text_column_are_left_alone(self):
        self.create_workbook([{"title": "S", "rows": [["코드", "이름", "수량"], ["007", "이샘플", 3], ["2026", "박예시", "4"]]}])
        self.assertEqual([issue["location"] for issue in self.text_issues()], ["S!C3"])

    def test_a_spec_date_is_a_date_like_a_csv_date_unless_the_type_says_text(self):
        self.create_workbook([{"title": "S", "rows": [["일자", "금액"], ["2026-09-01", 5]]}])
        cell = load_workbook(self.directory / "book.xlsx")["S"]["A2"]
        self.assertEqual((cell.value, cell.number_format), (datetime.datetime(2026, 9, 1), "yyyy-mm-dd"))
        self.apply([{"op": "set_cell", "sheet": "S", "cell": "A3", "value": "2026-09-02", "type": "text"}, {"op": "set_cell", "sheet": "S", "cell": "A4", "value": "2026.09.03"}])
        self.assertEqual(load_workbook(self.directory / "book.xlsx")["S"]["A3"].value, "2026-09-02")
        issues = self.text_issues()
        self.assertEqual([(issue["location"], issue["fix"][0]["value"]) for issue in issues], [("S!A4", "2026-09-03")])

    def test_a_date_written_by_row_csv_or_append_is_the_same_date(self):
        (self.directory / "data.csv").write_text("일자,금액\n2026-01-06,200\n", encoding="utf-8")
        created = run_office(["sheet", "create", "book.xlsx", "--title", "S", "--row", "일자,금액", "--row", "2026-01-05,100"], self.directory)
        self.assertEqual(created["status"], "ok", created)
        self.assertEqual(run_office(["sheet", "edit", "book.xlsx", "--row", "2026-01-07,300"], self.directory)["status"], "ok")
        self.assertEqual(run_office(["convert", "data.csv", "csv.xlsx"], self.directory)["status"], "ok")
        appended = load_workbook(self.directory / "book.xlsx")["S"]
        converted = load_workbook(self.directory / "csv.xlsx").active
        cells = [appended["A2"], appended["A3"], converted["A2"]]
        self.assertEqual([(cell.value, cell.number_format) for cell in cells], [
            (datetime.datetime(2026, 1, 5), "yyyy-mm-dd"),
            (datetime.datetime(2026, 1, 7), "yyyy-mm-dd"),
            (datetime.datetime(2026, 1, 6), "yyyy-mm-dd"),
        ])
        self.assertEqual(appended["B2"].value, 100)

