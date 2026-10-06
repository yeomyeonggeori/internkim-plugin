import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from doc_fixture import OFFICE_ENTRY, run_office, run_office_python, write_json
from task_context_fixture import environment_with_context, write_context_at


FORM_DOCUMENT = """
from docx import Document
document = Document()
approval = document.add_table(rows=2, cols=3)
for cell, text in zip(approval.rows[0].cells, ["담당", "팀장", "대표"]):
    cell.text = text
header = document.add_table(rows=2, cols=4)
for row, texts in zip(header.rows, [["문서번호", "", "작성일", ""], ["거래처", "", "작성자", ""]]):
    for cell, text in zip(row.cells, texts):
        cell.text = text
document.add_paragraph("품목")
items = document.add_table(rows=5, cols=5)
for cell, text in zip(items.rows[0].cells, ["No", "품명", "수량", "단가", "금액"]):
    cell.text = text
for number, row in enumerate(items.rows[1:4], start=1):
    row.cells[0].text = str(number)
total = items.rows[4].cells
total[0].merge(total[3]).text = "합계"
document.save("form.docx")
"""

FORM_WORKBOOK = """
from openpyxl import Workbook
workbook = Workbook()
sheet = workbook.active
sheet.title = "Plan"
sheet["A1"], sheet["C1"] = "Project", "Owner"
sheet.append([])
sheet.append(["No", "Name", "Months", "Rate", "Cost"])
for number in range(1, 4):
    sheet.append([number, None, None, None, f"=C{number + 3}*D{number + 3}"])
sheet.append(["Total", None, None, None, "=SUM(E4:E6)"])
workbook.save("form.xlsx")
"""


def runtime_context(directory):
    facts = {"requester": {"name": "이샘플"}, "today": "2026-10-04", "company": {"ko": {"name": "샘플테크"}}, "registeredDocuments": [{"documentNumber": "SAMPLE-7"}]}
    return write_context_at(Path(directory, "context.json"), facts)


def merge(directory, schema, values, output):
    write_json(Path(directory, "values.json"), values)
    environment = environment_with_context(runtime_context(directory))
    completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "merge", schema, "values.json", output], capture_output=True, text=True, cwd=directory, env=environment)
    return json.loads(completed.stdout)


def typed(schema_path, changes):
    schema = json.loads(Path(schema_path).read_text())
    for field in schema["fields"]:
        field.update(changes.get(field["name"], {}))
        for child in field.get("fields") or []:
            child.update(changes.get(f"{field['name']}.{child['name']}", {}))
    Path(schema_path).write_text(json.dumps(schema, ensure_ascii=False))
    return schema


class DocxFormTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        run_office_python(FORM_DOCUMENT, self.directory.name)

    def derive(self):
        result = run_office(["convert", "form.docx", "form.schema.json"], self.directory.name)
        self.assertEqual(result["status"], "ok", result)
        return json.loads(Path(self.directory.name, "form.schema.json").read_text())

    def test_every_empty_cell_beside_or_under_a_label_becomes_a_field(self):
        names = [field["name"] for field in self.derive()["fields"]]
        self.assertEqual(names, ["담당", "팀장", "대표", "문서번호", "작성일", "거래처", "작성자", "품목", "합계"])

    def test_a_table_of_empty_numbered_rows_becomes_a_list_named_by_the_text_above_it(self):
        items = next(field for field in self.derive()["fields"] if field["type"] == "list")
        self.assertEqual(items["label"], "품목")
        self.assertEqual([child["name"] for child in items["fields"]], ["품명", "수량", "단가", "금액"])
        self.assertEqual(items["at"]["rows"], 3)

    def test_the_filled_form_keeps_its_layout_and_writes_only_its_cells(self):
        self.derive()
        typed(Path(self.directory.name, "form.schema.json"), {
            "담당": {"type": "handwritten"}, "팀장": {"type": "handwritten"}, "대표": {"type": "handwritten"},
            "문서번호": {"known": "document.number"}, "작성일": {"known": "today"}, "작성자": {"known": "requester"},
            "품목.수량": {"type": "quantity"}, "품목.단가": {"type": "amount", "unit": "KRW"},
            "품목.금액": {"type": "amount", "unit": "KRW", "expression": "품목.수량 * 품목.단가"},
            "합계": {"type": "amount", "unit": "KRW", "expression": "sum(품목.금액)"},
        })
        values = {"거래처": None, "품목": [{"품명": "의자", "수량": 2, "단가": 50000}, {"품명": "책상", "수량": 1, "단가": 120000}]}
        result = merge(self.directory.name, "form.schema.json", values, "filled.docx")
        self.assertEqual(result["status"], "ok", result)
        self.assertEqual([blank["label"] for blank in result["details"]["blanks"]], ["거래처"])
        tables = run_office(["read", "filled.docx"], self.directory.name)["details"]["blocks"]
        cells = [block["cells"] for block in tables if block["kind"] == "table"]
        self.assertEqual(cells[0][1], ["", "", ""])
        self.assertEqual(cells[1], [["문서번호", "SAMPLE-7", "작성일", "2026년 10월 4일"], ["거래처", "", "작성자", "이샘플"]])
        self.assertEqual(cells[2][1], ["1", "의자", "2", "50,000원", "100,000원"])
        self.assertEqual(cells[2][4][-1], "220,000원")



SETTLEMENT_FORM = """
from docx import Document
document = Document()
header = document.add_table(rows=1, cols=4)
for cell, text in zip(header.rows[0].cells, ["소    속", "", "Unit price (USD)", ""]):
    pass
header.add_row()
for cell, text in zip(header.rows[1].cells, ["합    계", "", "성명          (인)", ""]):
    cell.text = text
for cell, text in zip(header.rows[0].cells, ["소    속", "", "Unit price (USD)", ""]):
    cell.text = text
document.add_paragraph("1. 경비 내역 (영수증 첨부)")
items = document.add_table(rows=4, cols=3)
for cell, text in zip(items.rows[0].cells, ["No", "사용내역", "금액(원)"]):
    cell.text = text
for number, row in enumerate(items.rows[1:], start=1):
    row.cells[0].text = str(number)
signed = document.add_table(rows=1, cols=1)
signed.rows[0].cells[0].text = "작성일: 20   년   월   일"
document.add_paragraph("위와 같이 출장비를 정산하오니 확인 바랍니다.\\n\\n20    년     월     일\\n")
document.add_paragraph("신청인:                         (인)")
document.add_paragraph("1.  두 칸 띄운 제목")
document.save("settlement.docx")
"""


class FieldIdentityAndParagraphBlankTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        run_office_python(SETTLEMENT_FORM, self.directory.name)
        result = run_office(["convert", "settlement.docx", "settlement.schema.json"], self.directory.name)
        self.assertEqual(result["status"], "ok", result)
        self.schema = json.loads(Path(self.directory.name, "settlement.schema.json").read_text())

    def every_field(self):
        for field in self.schema["fields"]:
            yield field
            yield from field.get("fields") or []

    def test_every_field_has_an_id_an_expression_can_name_and_keeps_its_label(self):
        identifier = __import__("re").compile(r"[^\W\d]\w*")
        for field in self.every_field():
            self.assertRegex(field["name"], identifier)
            self.assertIsNotNone(identifier.fullmatch(field["name"]), field)
        labels = {field["label"] for field in self.every_field()}
        self.assertTrue({"금액(원)", "Unit price (USD)", "1. 경비 내역 (영수증 첨부)"} <= labels, labels)

    def test_blanks_inside_paragraph_text_become_fields(self):
        by_label = {field["label"]: field for field in self.every_field()}
        self.assertEqual(by_label["신청인"]["type"], "text")
        dates = [field for field in self.every_field() if field["type"] == "date"]
        self.assertEqual(len(dates), 2, [field["label"] for field in dates])
        self.assertIn("작성일", {field["label"] for field in dates})
        self.assertNotIn("1.", {field["label"] for field in self.every_field()})
        self.assertEqual(by_label["소 속"]["at"], {"table": 0, "row": 0, "column": 1})
        self.assertEqual(by_label["합 계"]["at"], {"table": 0, "row": 1, "column": 1})
        self.assertEqual(by_label["성명"]["at"]["blank"], 0)

    def test_a_derived_id_works_in_an_expression_and_the_paragraph_blanks_are_filled_in_place(self):
        items = next(field for field in self.schema["fields"] if field["type"] == "list")
        amount = next(child for child in items["fields"] if child["label"] == "금액(원)")
        dates = [field["name"] for field in self.every_field() if field["type"] == "date"]
        applicant = next(field["name"] for field in self.every_field() if field["label"] == "신청인")
        changes = {f"{items['name']}.{amount['name']}": {"type": "amount", "unit": "KRW"}, applicant: {"known": "requester"}}
        changes.update({name: {"known": "today"} for name in dates})
        self.schema = typed(Path(self.directory.name, "settlement.schema.json"), changes)
        self.schema["derived"] = [{"name": "총액", "type": "amount", "unit": "KRW", "expression": f"sum({items['name']}.{amount['name']})"}]
        Path(self.directory.name, "settlement.schema.json").write_text(json.dumps(self.schema, ensure_ascii=False))
        values = {items["name"]: [{amount["name"]: 12000}, {amount["name"]: 8000}]}
        result = merge(self.directory.name, "settlement.schema.json", values, "filled.docx")
        self.assertEqual(result["status"], "ok", result)
        blocks = run_office(["read", "filled.docx"], self.directory.name)["details"]["blocks"]
        texts = [block.get("text", "") for block in blocks if block["kind"] == "paragraph"]
        tables = [block["cells"] for block in blocks if block["kind"] == "table"]
        self.assertIn("2026년 10월 4일", "".join(texts))
        self.assertIn("이샘플", next(text for text in texts if text.startswith("신청인:")))
        self.assertTrue(next(text for text in texts if text.startswith("신청인:")).rstrip().endswith("(인)"))
        self.assertEqual(tables[2][0][0], "작성일: 2026년 10월 4일")
        self.assertEqual(tables[1][1][2], "12,000원")

    def test_a_schema_field_whose_name_an_expression_cannot_read_is_refused_naming_its_label(self):
        self.schema["fields"][0]["name"] = "소속(부서)"
        Path(self.directory.name, "settlement.schema.json").write_text(json.dumps(self.schema, ensure_ascii=False))
        result = merge(self.directory.name, "settlement.schema.json", {}, "filled.docx")
        self.assertEqual(result["status"], "error")
        self.assertIn("소속(부서)", json.dumps(result["issues"], ensure_ascii=False))


class XlsxFormTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        run_office_python(FORM_WORKBOOK, self.directory.name)

    def test_a_column_the_form_computes_stays_its_own_formula(self):
        run_office(["convert", "form.xlsx", "form.schema.json"], self.directory.name)
        schema = json.loads(Path(self.directory.name, "form.schema.json").read_text())
        plan = next(field for field in schema["fields"] if field["type"] == "list")
        self.assertEqual({child["name"]: child["type"] for child in plan["fields"]}, {"Name": "text", "Months": "text", "Rate": "text", "Cost": "ignore"})
        typed(Path(self.directory.name, "form.schema.json"), {"Name.Months": {"type": "quantity"}, f"{plan['name']}.Months": {"type": "quantity"}, f"{plan['name']}.Rate": {"type": "amount"}})
        values = {"Project": "Pilot", "Owner": None, plan["name"]: [{"Name": "박예시", "Months": 2, "Rate": 100}]}
        result = merge(self.directory.name, "form.schema.json", values, "filled.xlsx")
        self.assertEqual(result["status"], "ok", result)
        cells = run_office(["read", "filled.xlsx", "--range", "A1:E7", "--where", "formula"], self.directory.name)["details"]["range"]["cells"]
        self.assertEqual({cell["cell"]: cell["value"] for cell in cells}["E4"], 200)
        rows = run_office(["read", "filled.xlsx", "--range", "A1:E4"], self.directory.name)["details"]["range"]["values"]
        self.assertEqual(rows[0][1], "Pilot")
        self.assertEqual(rows[3][:4], [1, "박예시", 2, 100])


    def test_a_blank_inside_a_cell_s_text_is_its_own_field_and_is_filled_in_place(self):
        code = "from openpyxl import Workbook\nworkbook = Workbook()\nsheet = workbook.active\nsheet.title = 'Plan'\nsheet['A1'] = '작성일: 20   년   월   일'\nworkbook.save('dated.xlsx')\n"
        run_office_python(code, self.directory.name)
        run_office(["convert", "dated.xlsx", "dated.schema.json"], self.directory.name)
        schema = json.loads(Path(self.directory.name, "dated.schema.json").read_text())
        self.assertEqual([(field["label"], field["type"]) for field in schema["fields"]], [("작성일", "date")])
        typed(Path(self.directory.name, "dated.schema.json"), {schema["fields"][0]["name"]: {"known": "today"}})
        result = merge(self.directory.name, "dated.schema.json", {}, "filled.xlsx")
        self.assertEqual(result["status"], "ok", result)
        rows = run_office(["read", "filled.xlsx", "--range", "A1:A1"], self.directory.name)["details"]["range"]["values"]
        self.assertEqual(rows[0][0], "작성일: 2026년 10월 4일")


if __name__ == "__main__":
    unittest.main()
