from pathlib import Path

from deck_fixture import DeckFixture, SCRIPTS_PATH, run_office_python_json


NOTES_DECK = """
<section data-slide-role="summary"><h2>매출이 늘었습니다</h2><p>본문</p><aside class="notes">첫 슬라이드 노트
둘째 줄 &amp; 끝</aside></section>
<section data-slide-role="risk"><h2>노트가 없는 슬라이드입니다</h2><p>본문</p></section>
<section data-slide-role="approval"><h2>마지막 슬라이드입니다</h2><p>본문</p><aside class="notes">마지막 노트</aside></section>
"""
EXPECTED_NOTES = ["첫 슬라이드 노트\n둘째 줄 & 끝", None, "마지막 노트"]

WRITE_AND_READ_BACK = """
import json, sys
from pathlib import Path
sys.path[:0] = [{scripts_path!r}, {tests_path!r}]
from deck.editable_pptx import read_text_layers, write_editable_pptx
from png_fixture import write_png
from pptx import Presentation
from deck.slide_source import split_slide_sources
from deck.slide_structure import extract_notes

notes = [extract_notes(slide_source) for slide_source in split_slide_sources({deck!r})]
layers_path = Path("review/pptx-layers")
layers_path.mkdir(parents=True)
run = {{"text": "본문", "fontFamily": "PaperlogyLocal", "fontWeight": 400, "italic": False, "sizePx": 32, "color": "rgb(0, 0, 0)", "opacity": 1, "letterSpacingPx": 0, "underline": False, "strike": False, "baseline": "", "href": None}}
paragraph = {{"alignment": "l", "lineHeightPx": 40, "spaceBeforePx": 0, "bullet": None, "runs": [run], "lines": [{{"widthPx": 64, "sizePx": 32, "text": "본문"}}]}}
block = {{"box": {{"left": 80, "top": 80, "right": 800, "bottom": 120}}, "insets": {{"left": 0, "top": 0, "right": 0, "bottom": 0}}, "anchor": "t", "noWrap": False, "keepWords": True, "firstLineHalfLeading": 4, "paragraphs": [paragraph]}}
slides = [{{"width": 1600, "height": 900, "visibleText": "본문", "pictureTexts": [], "blocks": [block], "shapes": [], "boxesKeptAsPicture": 0}} for _ in notes]
(layers_path / "layout.json").write_text(json.dumps({{"language": "ko", "slides": slides}}), encoding="utf-8")
for number in range(1, len(notes) + 1):
    write_png(layers_path / f"background.{{number:03}}.png", 32, 18, [[(255, 255, 255, 255)] * 32 for _ in range(18)])
write_editable_pptx(read_text_layers(Path("review"), len(notes)), notes, Path("editable.pptx"))
print(json.dumps([slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else None for slide in Presentation("editable.pptx").slides], ensure_ascii=False))
"""


class NotesRoundTripTest(DeckFixture):
    def read_back(self, deck: str):
        code = WRITE_AND_READ_BACK.format(scripts_path=str(SCRIPTS_PATH), tests_path=str(Path(__file__).resolve().parent), deck=deck)
        return run_office_python_json(code, self.directory)

    def test_the_pptx_carries_the_slide_notes_python_pptx_reads_back(self):
        self.assertEqual(self.read_back(NOTES_DECK), EXPECTED_NOTES)

    def test_a_deck_without_notes_gets_no_notes_parts(self):
        self.assertEqual(self.read_back('<section><h2>노트 없는 덱입니다</h2></section>'), [None])
