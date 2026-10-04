import base64
import io
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest

from PIL import Image

from doc_fixture import OFFICE_ENTRY, SCRIPTS_PATH, run_office, run_office_python
from render_fixture import assert_pages_drawn, can_render, pdf_page_count


REPORT = """
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

document = Document()
footer = document.sections[0].footer.paragraphs[0]
run = footer.add_run()
for kind, text in (("begin", None), ("instr", " PAGE "), ("separate", None), ("text", "1"), ("end", None)):
    node = OxmlElement("w:instrText" if kind == "instr" else "w:t" if kind == "text" else "w:fldChar")
    if text is None:
        node.set(qn("w:fldCharType"), kind)
    else:
        node.text = text
    run._r.append(node)
footer.add_run(" 쪽")
document.sections[0].header.paragraphs[0].text = "예시상사 내부 보고"
document.add_heading("분기 보고서", 0)
for item in ("매출 요약", "지역별 실적", "다음 분기 계획"):
    document.add_paragraph(item, style="List Number")
table = document.add_table(rows=3, cols=3)
table.style = "Table Grid"
table.cell(0, 0).merge(table.cell(0, 2)).text = "병합 머리글"
table.cell(1, 0).merge(table.cell(2, 0)).text = "세로"
paragraph = document.add_paragraph("수정 전 문장")
paragraph._p.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="1" w:author="이샘플" w:date="2026-09-01T09:00:00Z"><w:r><w:t xml:space="preserve"> 추가된 문장</w:t></w:r></w:ins>'))
for number in range(40):
    document.add_paragraph(f"{number + 1}번째 문단입니다. 본 문단은 페이지 나눔을 확인하기 위해 충분히 길게 작성한 한국어 문장으로, 줄바꿈이 여러 번 일어나도록 이어집니다.")
breaker = document.add_paragraph("나눔 앞 문장")
breaker.runs[0].add_break(WD_BREAK.PAGE)
breaker.add_run("나눔 뒤 문장")
document.save("보고서.docx")
"""


def page_sections(html):
    return re.findall(r'<section data-page="(\d+)" style="([^"]*)">(.*?)</section>', html, re.S)


class DocumentPreviewTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary_directory = tempfile.TemporaryDirectory()
        cls.directory = Path(cls.temporary_directory.name)
        run_office_python(REPORT, cls.directory)
        cls.envelope = run_office(["render", "보고서.docx"], cls.directory)
        cls.html = (cls.directory / "보고서-preview" / "preview.html").read_text(encoding="utf-8")
        cls.pages = page_sections(cls.html)

    @classmethod
    def tearDownClass(cls):
        cls.temporary_directory.cleanup()

    def test_each_page_is_one_letter_section_with_inline_styles_only(self):
        self.assertGreaterEqual(len(self.pages), 3)
        self.assertEqual([int(number) for number, _, _ in self.pages], list(range(1, len(self.pages) + 1)))
        self.assertTrue(all("width:816px;height:1056px" in style for _, style, _ in self.pages))
        for forbidden in ("<style", "<script", "::before", "::after", "counter(", ":has("):
            self.assertNotIn(forbidden, self.html)
        self.assertEqual(self.envelope["details"]["pageCount"], len(self.pages))

    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_every_page_is_drawn_as_an_image_and_a_pdf_page(self):
        details = self.envelope["details"]
        self.assertEqual(self.envelope["issues"], [])
        assert_pages_drawn(self, details, self.directory, len(self.pages), (816, 1056))
        self.assertEqual(pdf_page_count(self.directory / details["pdf"]), len(self.pages))

    def test_fonts_name_the_files_the_layout_measured_with(self):
        fonts = self.envelope["details"]["previewFonts"]
        self.assertTrue(fonts)
        self.assertTrue(all(Path(font["path"]).is_file() and font["weight"] in (400, 700) for font in fonts))
        families = {font["family"] for font in fonts}
        for declaration in set(re.findall(r"font-family:(.*?);", self.html)):
            self.assertLessEqual(set(re.findall(r"&quot;([^&]*)&quot;", declaration)), families, declaration)

    def test_headers_footers_numbering_merges_and_tracked_insertions_show(self):
        first, last = self.pages[0][2], self.pages[-1][2]
        self.assertIn("예시상사 내부 보고", first)
        self.assertRegex(first, r">1</span><span[^>]*> 쪽<")
        self.assertRegex(last, rf">{len(self.pages)}</span><span[^>]*> 쪽<")
        self.assertIn(">1.</span>", first)
        self.assertIn("grid-column:1 / span 3", first)
        self.assertIn("grid-row:2 / span 2", first)
        self.assertRegex(first, r"text-decoration:underline;color:#1d4ed8\">[^<]*추가된 문장")

    def test_a_page_break_starts_the_text_after_it_on_a_new_page(self):
        before = next(index for index, (_, _, body) in enumerate(self.pages) if "나눔 앞 문장" in body)
        after = next(index for index, (_, _, body) in enumerate(self.pages) if "나눔 뒤 문장" in body)
        self.assertEqual(after, before + 1)

    def test_every_page_holds_no_more_than_its_body_area(self):
        code = (
            "import json, sys\n"
            "from pathlib import Path\n"
            "from doc.preview.layout import Layout\n"
            "from doc.preview.pagination import Paginator\n"
            "from doc.preview.document import DocxModelBuilder\n"
            "from fonts.preview import FontRegistry\n"
            "pages = Paginator(Layout(FontRegistry())).paginate(DocxModelBuilder(Path('보고서.docx')).sections())\n"
            "print(json.dumps([(sum(item.height for item in page.placed), page.body_bottom - page.body_top) for page in pages]))\n"
        )
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "-c", code], capture_output=True, text=True, cwd=self.directory, check=True)
        for used, available in json.loads(completed.stdout):
            self.assertLessEqual(used, available + 0.5)


GRIDLESS_TABLE_AND_BLANK_PAGE = """
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn

document = Document()
document.add_paragraph("표 위 문장")
table = document.add_table(rows=3, cols=4)
table.style = "Table Grid"
for row, values in zip(table.rows, [["구분", "상반기", "", "하반기"], ["매출", "100", "200", "300"], ["비용", "50", "60", "70"]]):
    for cell, value in zip(row.cells, values):
        cell.text = value
table.cell(0, 1).merge(table.cell(0, 2))
grid = table._tbl.find(qn("w:tblGrid"))
grid.getparent().remove(grid)
breaks = document.add_paragraph()
breaks.add_run().add_break(WD_BREAK.PAGE)
breaks.add_run().add_break(WD_BREAK.PAGE)
document.add_paragraph("마지막 쪽 문장")
document.save("빈쪽.docx")
"""


class PreviewDefectTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary_directory = tempfile.TemporaryDirectory()
        cls.directory = Path(cls.temporary_directory.name)
        run_office_python(GRIDLESS_TABLE_AND_BLANK_PAGE, cls.directory)
        cls.envelope = run_office(["render", "빈쪽.docx"], cls.directory)
        cls.html = (cls.directory / "빈쪽-preview" / "preview.html").read_text(encoding="utf-8")

    @classmethod
    def tearDownClass(cls):
        cls.temporary_directory.cleanup()

    def test_a_table_without_a_grid_takes_its_columns_from_every_row(self):
        columns = re.search(r"grid-template-columns:([^;]*);", self.html).group(1).split()
        self.assertEqual(len(columns), 4)
        self.assertEqual(len(set(columns)), 1, columns)

    def test_a_page_with_nothing_in_its_body_is_reported_by_number(self):
        blank = [issue for issue in self.envelope["issues"] if issue["code"] == "BLANK_PAGE"]
        self.assertEqual([issue["location"] for issue in blank], ["page 2"])



OVER_WIDE_TABLE = """
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Twips

document = Document()
section = document.sections[0]
table = document.add_table(rows=2, cols=6)
table.style = "Table Grid"
widths = [4535, 1644, 1644, 1644, 1644, 1644]
for column, width in zip(table._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths):
    column.set(qn("w:w"), str(width))
for row in table.rows:
    for cell, width in zip(row.cells, widths):
        cell.width = Twips(width)
for cell, text in zip(table.rows[0].cells, ["출장비 정산서", "결재", "담당", "팀장", "본부장", "대표이사"]):
    cell.text = text
document.save("넓은표.docx")
print((section.page_width - section.left_margin - section.right_margin) / 635)
"""


class OverWideTableTest(unittest.TestCase):
    def test_a_table_wider_than_the_text_area_is_scaled_to_it_keeping_its_proportions(self):
        with tempfile.TemporaryDirectory() as name:
            directory = Path(name)
            completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "-c", OVER_WIDE_TABLE], capture_output=True, text=True, cwd=directory, check=True)
            text_width_twips = float(completed.stdout.strip())
            run_office(["render", "넓은표.docx"], directory)
            html = (directory / "넓은표-preview" / "preview.html").read_text(encoding="utf-8")
        columns = [float(width.removesuffix("px")) for width in re.search(r"grid-template-columns:([^;]*);", html).group(1).split()]
        text_width = text_width_twips / 15
        self.assertLessEqual(sum(columns), text_width + 0.5)
        self.assertGreater(sum(columns), text_width - 0.5)
        self.assertAlmostEqual(columns[0] / columns[1], 4535 / 1644, places=2)


TWO_SECTIONS_AND_A_LOGO = """
from docx import Document
from PIL import Image
Image.new("RGB", (400, 200), (29, 78, 216)).save("logo.png")
document = Document()
for number in range(3):
    document.add_paragraph(f"첫 구역 문단 {number + 1}")
document.add_paragraph("둘째 구역 문단")
document.save("구역.docx")
"""


class WatermarkAndPageNumberingTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary_directory = tempfile.TemporaryDirectory()
        cls.directory = Path(cls.temporary_directory.name)
        run_office_python(TWO_SECTIONS_AND_A_LOGO, cls.directory)
        cls.apply([{"op": "set_footer", "text": "- {PAGE} -", "align": "center"}, {"op": "insert_section_break", "after": 2}])
        cls.apply([{"op": "set_page_setup", "section": 1, "pageNumberStart": 7}, {"op": "set_watermark", "image": "logo.png"}])
        cls.envelope = run_office(["render", "구역.docx"], cls.directory)
        cls.pages = page_sections((cls.directory / "구역-preview" / "preview.html").read_text(encoding="utf-8"))

    @classmethod
    def apply(cls, operations):
        (cls.directory / "ops.json").write_text(json.dumps(operations, ensure_ascii=False), encoding="utf-8")
        envelope = run_office(["apply", "구역.docx", "ops.json"], cls.directory)
        assert envelope["status"] == "ok", envelope

    @classmethod
    def tearDownClass(cls):
        cls.temporary_directory.cleanup()

    def test_a_section_numbered_from_a_start_value_shows_it_on_its_first_page(self):
        self.assertEqual(self.envelope["status"], "ok", self.envelope["issues"])
        self.assertEqual([number for number, _, _ in self.pages], ["1", "2"])
        self.assertEqual([re.findall(r">(\d+)<", body)[-1] for _, _, body in self.pages], ["1", "7"])

    def test_a_picture_watermark_is_drawn_washed_out_behind_every_page_and_not_in_the_header(self):
        for _, _, body in self.pages:
            sources = re.findall(r'<img src="data:image/png;base64,([^"]+)"', body)
            self.assertEqual(len(sources), 1)
            self.assertTrue(body.index("<img") < body.index("문단"), "the watermark is drawn before the text, so the text covers it")
            self.assertGreater(min(washed_pixel(sources[0])), 160)

    def test_a_picture_watermark_keeps_its_colors_without_washout_and_takes_a_scale(self):
        (self.directory / "plain.json").write_text(json.dumps([{"op": "set_watermark", "image": "logo.png", "washout": False, "scale": 50}]), encoding="utf-8")
        envelope = run_office(["apply", "구역.docx", "plain.json", "--output", "원색.docx"], self.directory)
        self.assertEqual(envelope["status"], "ok", envelope["issues"])
        run_office(["render", "원색.docx"], self.directory)
        _, _, body = page_sections((self.directory / "원색-preview" / "preview.html").read_text(encoding="utf-8"))[0]
        source = re.search(r'<img src="data:image/png;base64,([^"]+)" style="([^"]*)"', body)
        self.assertEqual(washed_pixel(source.group(1)), (29, 78, 216))
        self.assertEqual(source.group(2).replace(" ", ""), "width:266.67px;height:133.33px")

    def test_a_watermark_takes_text_or_an_image_but_not_both(self):
        (self.directory / "both.json").write_text(json.dumps([{"op": "set_watermark", "text": "대외비", "image": "logo.png"}]), encoding="utf-8")
        envelope = run_office(["apply", "구역.docx", "both.json", "--dry-run"], self.directory)
        self.assertEqual([(issue["code"], issue["location"]) for issue in envelope["issues"]], [("INVALID_VALUE", "ops[0]")])


STRANDED_HEADING = """
from docx import Document
from docx.shared import Pt

document = Document()
document.add_paragraph("머리 문단").paragraph_format.space_after = Pt(20)
for index in range(23):
    document.add_paragraph(f"본문 문단 {index}")
heading = document.add_heading("다음 단계", 2)
heading.paragraph_format.keep_with_next = False
document.add_paragraph("세부 계획은 착수 전에 확정합니다.")
document.save("떨어진제목.docx")
"""

KEPT_HEADING = """
from docx import Document
from docx.shared import Pt

document = Document()
for index in range(24):
    document.add_paragraph(f"본문 문단 {index}")
document.add_heading("다음 단계", 2)
document.add_paragraph("세부 계획은 착수 전에 확정합니다.").paragraph_format.space_after = Pt(24)
for item in ("첫째", "둘째"):
    document.add_paragraph(item, style="List Bullet")
document.save("빠듯한.docx")
"""


class PaginationTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.directory = Path(self.temporary_directory.name)

    def page_texts(self, name):
        run_office(["render", name], self.directory)
        html = (self.directory / f"{Path(name).stem}-preview" / "preview.html").read_text(encoding="utf-8")
        return [re.sub(r"<[^>]+>", "", body) for _, _, body in page_sections(html)]

    def test_a_heading_left_at_a_page_bottom_is_reported_with_keep_with_next_as_its_fix(self):
        run_office_python(STRANDED_HEADING, self.directory)
        issues = [issue for issue in run_office(["check", "떨어진제목.docx"], self.directory)["issues"] if issue["code"] == "HEADING_STRANDED"]
        self.assertEqual([(issue["location"], issue["fix"]) for issue in issues], [("block 24", [{"op": "set_paragraph_format", "block": 24, "keepWithNext": True}])])
        (self.directory / "fix.json").write_text(json.dumps(issues[0]["fix"]), encoding="utf-8")
        self.assertEqual(run_office(["apply", "떨어진제목.docx", "fix.json"], self.directory)["status"], "ok")
        self.assertNotIn("HEADING_STRANDED", [issue["code"] for issue in run_office(["check", "떨어진제목.docx"], self.directory)["issues"]])
        pages = self.page_texts("떨어진제목.docx")
        self.assertIn("다음 단계", pages[1])
        self.assertNotIn("다음 단계", pages[0])

    def test_a_kept_heading_stays_with_a_last_line_whose_space_after_runs_past_the_page(self):
        run_office_python(KEPT_HEADING, self.directory)
        pages = self.page_texts("빠듯한.docx")
        heading_page = next(index for index, page in enumerate(pages) if "다음 단계" in page)
        self.assertIn("세부 계획은 착수 전에 확정합니다.", pages[heading_page])

    def test_list_items_of_one_style_drop_the_space_between_them(self):
        run_office_python(KEPT_HEADING, self.directory)
        reader = "import json; from pathlib import Path; from doc.preview.document import DocxModelBuilder; blocks = DocxModelBuilder(Path('빠듯한.docx')).sections()[0].blocks; print(json.dumps([[block.space_before, block.space_after] for block in blocks[-2:]]))"
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "-c", reader], capture_output=True, text=True, check=True, cwd=self.directory, env={**os.environ, "PYTHONPATH": str(SCRIPTS_PATH)})
        first_item, second_item = json.loads(completed.stdout)
        self.assertEqual((first_item[1], second_item[0]), (0, 0))


def washed_pixel(base64_png):
    with Image.open(io.BytesIO(base64.b64decode(base64_png))) as picture:
        return picture.convert("RGB").getpixel((0, 0))


if __name__ == "__main__":
    unittest.main()
