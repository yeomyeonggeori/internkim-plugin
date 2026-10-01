import unittest

from pptx.util import Pt

from test_deck_pptx_editing import KoreanDeckFixture, codes


class RunStyleTest(KoreanDeckFixture):
    def runs(self, shape, paragraph=0):
        frame = self.presentation().slides[1].shapes[shape].text_frame
        return frame.paragraphs[paragraph].runs

    def test_find_styles_only_the_matching_text_and_keeps_the_rest(self):
        envelope = self.apply([{"op": "set_text_style", "slide": 2, "shape": 2, "find": "128억", "color": "FFD966", "size": 32, "underline": True}])
        self.assertEqual(envelope["status"], "ok", envelope)
        runs = self.runs(2)
        self.assertEqual([run.text for run in runs], ["3분기 매출 ", "128억", " 원"])
        middle = runs[1].font
        self.assertEqual((str(middle.color.rgb), middle.size, middle.underline, middle.bold), ("FFD966", Pt(32), True, True))
        for outer in (runs[0].font, runs[2].font):
            self.assertEqual((str(outer.color.rgb), outer.size, outer.underline, outer.bold), ("FFFFFF", Pt(24), None, True))

    def test_find_styles_every_occurrence_across_paragraphs(self):
        self.apply([{"op": "set_text", "slide": 2, "shape": 3, "text": "고객 42곳, 고객 만족\n재구매 고객 68%"}])
        self.apply([{"op": "set_text_style", "slide": 2, "shape": 3, "find": "고객", "bold": True}])
        bold_texts = [run.text for paragraph in (0, 1) for run in self.runs(3, paragraph) if run.font.bold]
        self.assertEqual(bold_texts, ["고객", "고객", "고객"])

    def test_find_within_one_paragraph_leaves_the_others(self):
        self.apply([{"op": "set_text_style", "slide": 2, "shape": 3, "paragraph": 1, "find": "68%", "italic": True}])
        self.assertEqual([run.font.italic for run in self.runs(3, 0)], [None])
        self.assertEqual([run.text for run in self.runs(3, 1) if run.font.italic], ["68%"])

    def test_text_that_is_not_there_is_refused(self):
        envelope = self.apply([{"op": "set_text_style", "slide": 2, "shape": 2, "find": "129억", "bold": False}])
        self.assertEqual(codes(envelope), ["TARGET_NOT_FOUND"])
        self.assertEqual(envelope["issues"][0]["location"], "ops[0].find")

    def test_a_link_on_part_of_the_text_still_splits_its_run(self):
        envelope = self.apply([{"op": "set_link", "slide": 2, "shape": 2, "text": "128억", "url": "https://example.com"}])
        self.assertEqual(envelope["status"], "ok", envelope)
        self.assertEqual([run.text for run in self.runs(2)], ["3분기 매출 ", "128억", " 원"])


if __name__ == "__main__":
    unittest.main()
