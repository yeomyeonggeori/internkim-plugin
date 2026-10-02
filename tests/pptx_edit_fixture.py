import io
from pathlib import Path
import re
import zipfile

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt


SLIDE_WIDTH = 12192000
SLIDE_HEIGHT = 6858000
NAMESPACES = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
UNKNOWN_EXTENSION_URI = "{8C1F3C2B-7A55-4E1A-9F3E-5B2D7A0C1E99}"
CUSTOM_PART_NAME = "/customXml/item1.xml"
CUSTOM_RELATIONSHIP_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
DRAWING_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/main"
POWERPOINT_DECLARATION = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'


def build_korean_deck(path):
    presentation = Presentation()
    widen_template(presentation, SLIDE_WIDTH / presentation.slide_width)
    presentation.slide_width = Emu(SLIDE_WIDTH)
    presentation.slide_height = Emu(SLIDE_HEIGHT)
    add_cover(presentation)
    add_metrics(presentation)
    add_chart_slide(presentation)
    add_table_slide(presentation)
    add_closing(presentation)
    add_sections(presentation)
    add_custom_part(presentation)
    buffer = io.BytesIO()
    presentation.save(buffer)
    Path(path).write_bytes(as_powerpoint_wrote_it(buffer.getvalue()))


def widen_template(presentation, ratio):
    for master in presentation.slide_masters:
        for owner in (master, *master.slide_layouts):
            for shape in owner.shapes:
                transform = shape._element.spPr.find(f"{{{DRAWING_NAMESPACE}}}xfrm")
                if transform is not None:
                    offset, extent = transform.find(f"{{{DRAWING_NAMESPACE}}}off"), transform.find(f"{{{DRAWING_NAMESPACE}}}ext")
                    offset.set("x", str(round(int(offset.get("x")) * ratio)))
                    extent.set("cx", str(round(int(extent.get("cx")) * ratio)))


def add_cover(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "2026년 3분기 경영 실적 보고"
    slide.placeholders[1].text = "주식회사 예시 · 경영지원팀 이샘플"
    slide.notes_slide.notes_text_frame.text = "인사 후 3분기 요약으로 시작합니다."


def add_metrics(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "매출은 목표를 넘었고 비용은 줄었습니다"
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(609600), Emu(1828800), Emu(3352800), Emu(1371600))
    card.fill.solid()
    card.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    card.line.fill.background()
    label = slide.shapes.add_textbox(Emu(762000), Emu(1981200), Emu(3048000), Emu(457200))
    label.text_frame.text = "3분기 매출 128억 원"
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    label.text_frame.paragraphs[0].runs[0].font.bold = True
    label.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    body = slide.shapes.add_textbox(Emu(4572000), Emu(1828800), Emu(6858000), Emu(1371600))
    body.text_frame.word_wrap = True
    body.text_frame.text = "신규 고객 42곳 확보"
    body.text_frame.add_paragraph().text = "재구매율 68%로 상승"
    for paragraph in body.text_frame.paragraphs:
        paragraph.runs[0].font.size = Pt(20)
    slide.shapes.add_picture(io.BytesIO(sample_photo(834, 400)), Emu(609600), Emu(3505200), Emu(4572000), Emu(2743200))
    picture = slide.shapes[-1]
    picture.crop_left = 0.1
    picture.crop_right = 0.1
    slide.notes_slide.notes_text_frame.text = "매출과 비용을 함께 설명합니다."


def add_chart_slide(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "분기별 매출 추이"
    chart_data = CategoryChartData()
    chart_data.categories = ["1분기", "2분기", "3분기"]
    chart_data.add_series("매출", (96, 110, 128))
    chart_data.add_series("비용", (80, 84, 79))
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(609600), Emu(1524000), Emu(10972800), Emu(4572000), chart_data)
    frame.chart.has_legend = True


def add_table_slide(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "사업부별 목표 대비 실적"
    rows = [["사업부", "목표", "실적"], ["플랫폼", "50억", "58억"], ["컨설팅", "40억", "37억"]]
    table = slide.shapes.add_table(len(rows), 3, Emu(609600), Emu(1524000), Emu(6096000), Emu(1371600)).table
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    first = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(7315200), Emu(1524000), Emu(914400), Emu(914400))
    second = slide.shapes.add_textbox(Emu(8382000), Emu(1676400), Emu(3048000), Emu(609600))
    second.text_frame.text = "담당 박예시"
    group = slide.shapes.add_group_shape([first, second])
    group.name = "담당자 묶음"
    add_fade_animation(slide, group.shape_id)


def add_closing(presentation):
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "4분기에는 컨설팅 회복에 집중합니다"
    slide.placeholders[1].text = "문의 최견본 · team@example.com"
    extension_list = parse_xml(f'<p:extLst {NAMESPACES}><p:ext uri="{UNKNOWN_EXTENSION_URI}"><a:unknownThing value="keep me"/></p:ext></p:extLst>')
    slide._element.append(extension_list)


def add_fade_animation(slide, shape_id):
    timing = parse_xml(f"""<p:timing {NAMESPACES}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="5" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set><p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="7" dur="500"/><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst><p:bldP spid="{shape_id}" grpId="0"/></p:bldLst></p:timing>""")
    slide._element.append(timing)


def add_sections(presentation):
    slide_ids = [slide.slide_id for slide in presentation.slides]
    section_ids = [(slide_ids[:2], "도입"), (slide_ids[2:], "본론")]
    sections = "".join(
        f'<p14:section name="{name}" id="{{{index:08d}-0000-4000-8000-000000000000}}"><p14:sldIdLst>'
        + "".join(f'<p14:sldId id="{slide_id}"/>' for slide_id in identifiers)
        + "</p14:sldIdLst></p14:section>"
        for index, (identifiers, name) in enumerate(section_ids, start=1)
    )
    extension_list = parse_xml(
        f'<p:extLst {NAMESPACES}><p:ext uri="{{521415D9-36F7-43E2-AB2F-B90AF26B5E84}}">'
        f'<p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">{sections}</p14:sectionLst>'
        "</p:ext></p:extLst>"
    )
    presentation.part._element.append(extension_list)


def add_custom_part(presentation):
    blob = '<?xml version="1.0" encoding="UTF-8"?><sample xmlns="urn:example:sample"><owner>이샘플</owner></sample>'.encode("utf-8")
    part = Part(PackURI(CUSTOM_PART_NAME), "application/xml", presentation.part.package, blob)
    presentation.part.relate_to(part, CUSTOM_RELATIONSHIP_TYPE)


def sample_photo(width=640, height=400):
    image = Image.new("RGB", (width, height), (32, 96, 160))
    drawing = ImageDraw.Draw(image)
    drawing.rectangle((width // 4, height // 4, width * 3 // 4, height * 3 // 4), fill=(240, 200, 80))
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def as_powerpoint_wrote_it(package):
    source = zipfile.ZipFile(io.BytesIO(package))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.endswith((".xml", ".rels")):
                data = POWERPOINT_DECLARATION + re.sub(rb"^<\?xml[^>]*\?>\s*", b"", data)
            target.writestr(info, data)
    return output.getvalue()


def part_contents(path):
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


def canonical_xml(data):
    return etree.tostring(etree.fromstring(data), method="c14n")
