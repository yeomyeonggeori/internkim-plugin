from pathlib import Path
import unittest

from openpyxl import load_workbook

from doc_fixture import run_office
from render_fixture import can_render
from test_workbook_declaration import DeclaredWorkbookFixture


def table(rows, dimensions, measure="Units"):
    return {"name": "Data", "columns": [{"name": name} for name in dimensions] + [{"name": measure, "type": "quantity"}], "rows": rows}


def grid(first, second, third=None):
    rows = []
    for one_index, one in enumerate(first):
        for two_index, two in enumerate(second):
            for three_index, three in enumerate(third or [None]):
                rows.append([one, two] + ([three] if third else []) + [10 + one_index * 7 + two_index * 3 + three_index])
    return rows


def declaration(data, views, charts=()):
    return {"kind": "workbook", "language": "en", "tables": [data], "views": views, "charts": list(charts)}


def view(title, rows, columns, add=(), measure="Units"):
    return {"sheet": "Summary", "title": title, "rows": rows, "columns": columns, "measure": measure, "totals": True, "add": list(add)}


STORES = ["North store", "South store", "Harbor store"]
MONTHS = ["Jan", "Feb", "Mar", "Apr"]
YEARS = [2022, 2023, 2024]
TEAMS = ["Alpha", "Beta", "Gamma", "Delta"]


class LayoutFixture(DeclaredWorkbookFixture):
    def build(self, content):
        envelope = self.create(content)
        self.assertEqual(envelope["status"], "ok", envelope)
        return load_workbook(Path(self.directory.name) / "book.xlsx")["Summary"]

    def render_pages(self):
        envelope = run_office(["render", "book.xlsx"], self.directory.name)
        self.assertNotEqual(envelope["status"], "error", envelope)
        return envelope["details"]["pageContents"]

    def header_rows(self, sheet, title):
        title_row = next(row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == title)
        has_group_row = any(merged.min_row == title_row + 1 for merged in sheet.merged_cells.ranges)
        return title_row, title_row + 1 if has_group_row else None, title_row + 1 + has_group_row


@unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
class ChartPageTest(LayoutFixture):
    def assert_charts_whole(self, content, chart_count):
        self.build(content)
        drawings = [drawing for page in self.render_pages() for drawing in page.get("drawings", [])]
        self.assertEqual(len(drawings), chart_count)
        self.assertEqual([drawing["shown"] for drawing in drawings], ["whole"] * chart_count, drawings)

    def weekly(self, weeks, charts):
        data = table(grid([f"Week {number}" for number in range(1, weeks + 1)], YEARS), ["Week", "Year"])
        return declaration(data, [view("By week", ["Week"], "Year")], charts)

    def test_a_chart_that_would_start_near_the_foot_of_a_page_prints_whole(self):
        self.assert_charts_whole(self.weekly(33, [{"view": "By week", "type": "line"}]), 1)

    def test_a_chart_under_a_longer_view_prints_whole(self):
        self.assert_charts_whole(self.weekly(36, [{"view": "By week", "type": "column"}]), 1)

    def test_every_chart_of_several_prints_whole(self):
        self.assert_charts_whole(self.weekly(10, [{"view": "By week", "type": kind} for kind in ("line", "column", "bar")]), 3)

    def test_a_chart_after_a_short_view_prints_whole(self):
        data = table(grid(STORES, MONTHS), ["Store", "Month"])
        self.assert_charts_whole(declaration(data, [view("By store", ["Store"], "Month")], [{"view": "By store", "type": "column"}]), 1)


class GroupedHeaderTest(LayoutFixture):
    def merged_headers(self, sheet):
        return {str(merged): sheet.cell(merged.min_row, merged.min_col).value for merged in sheet.merged_cells.ranges}

    def assert_header_names_the_member_once(self, sheet, group_row, header_row, group, members):
        spans = [merged for merged in sheet.merged_cells.ranges if merged.min_row == group_row and sheet.cell(group_row, merged.min_col).value == group]
        self.assertEqual(len(spans), 1)
        columns = range(spans[0].min_col, spans[0].max_col + 1)
        self.assertEqual([sheet.cell(header_row, column).value for column in columns], members)

    def test_a_change_added_down_the_rows_of_a_view_with_columns_gets_one_group_over_its_columns(self):
        data = table(grid(MONTHS, STORES), ["Month", "Store"])
        sheet = self.build(declaration(data, [view("Units", ["Month"], "Store", [{"name": "Month on month", "expression": "change(Units, Month)"}])]))
        title_row, group_row, header_row = self.header_rows(sheet, "Units")
        self.assertEqual(sheet.cell(group_row, 5).value, "Month on month")
        self.assert_header_names_the_member_once(sheet, group_row, header_row, "Month on month", STORES)

    def test_a_share_added_to_a_view_with_year_columns_gets_one_group(self):
        data = table(grid(TEAMS, YEARS), ["Team", "Year"])
        sheet = self.build(declaration(data, [view("Units", ["Team"], "Year", [{"name": "Share of team", "expression": "share(Units, Team)"}])]))
        _, group_row, header_row = self.header_rows(sheet, "Units")
        self.assert_header_names_the_member_once(sheet, group_row, header_row, "Share of team", YEARS)

    def test_a_change_across_several_years_gets_one_group_of_the_later_years(self):
        data = table(grid(STORES, YEARS), ["Store", "Year"])
        sheet = self.build(declaration(data, [view("Units", ["Store"], "Year", [{"name": "Growth", "expression": "percentChange(Units, Year)"}])]))
        _, group_row, header_row = self.header_rows(sheet, "Units")
        self.assert_header_names_the_member_once(sheet, group_row, header_row, "Growth", YEARS[1:])

    def test_the_added_name_is_written_once_however_many_columns_it_adds(self):
        data = table(grid(MONTHS, STORES, YEARS), ["Month", "Store", "Year"])
        sheet = self.build(declaration(data, [view("Units", ["Month"], "Store", [{"name": "Month on month", "expression": "change(Units, Month)"}])]))
        title_row, _, header_row = self.header_rows(sheet, "Units")
        names = [sheet.cell(row, column).value for row in range(title_row + 1, header_row + 1) for column in range(1, sheet.max_column + 1)]
        self.assertEqual(names.count("Month on month"), 1)

    def test_a_single_added_column_keeps_its_plain_header_and_no_group_row(self):
        data = table(grid(MONTHS, STORES[:1]), ["Month", "Store"])
        sheet = self.build(declaration(data, [view("Units", ["Month"], "Store", [{"name": "Month on month", "expression": "change(Units, Month)"}])]))
        title_row, group_row, _ = self.header_rows(sheet, "Units")
        self.assertIsNone(group_row)
        self.assertEqual([sheet.cell(title_row + 1, column).value for column in range(1, 4)], ["Month", "North store", "Month on month"])
        self.assertEqual(sheet.merged_cells.ranges, set())

    def test_the_values_under_a_group_are_the_same_as_without_one(self):
        data = table(grid(MONTHS, STORES), ["Month", "Store"])
        sheet = self.build(declaration(data, [view("Units", ["Month"], "Store", [{"name": "Month on month", "expression": "change(Units, Month)"}])]))
        shown = self.create(declaration(data, [view("Units", ["Month"], "Store", [{"name": "Month on month", "expression": "change(Units, Month)"}])]))["details"]["views"][0]["rows"]
        self.assertEqual(shown[0][4:], ["Month on month", "", ""])
        self.assertEqual(shown[1][4:], STORES)
        self.assertEqual(shown[3][4:], ["7", "7", "7"])


class HeaderWidthTest(LayoutFixture):
    LONG = "Change against the same period of the previous year"

    def stacked(self, header):
        data = {"name": "Data", "columns": [{"name": "Month"}, {"name": "Store"}, {"name": "Units", "type": "quantity"}, {"name": "Plan", "type": "quantity"}], "rows": [[month, store, 10 + index, 12] for index, month in enumerate(MONTHS) for store in STORES]}
        views = [view("Plain", ["Month"], "Store"), {"sheet": "Summary", "title": "Wide header", "rows": ["Month"], "measures": ["Units", "Plan"], "add": [{"name": header, "expression": "Units - Plan"}]}]
        return self.build(declaration(data, views))

    def header_height(self, sheet):
        title_row = next(row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "Wide header")
        return sheet.row_dimensions[title_row + 1].height

    def test_a_long_header_wraps_in_its_own_view_without_widening_the_view_above(self):
        sheet = self.stacked(self.LONG)
        widest = max(sheet.column_dimensions[column].width for column in "ABCD")
        self.assertLessEqual(widest, len(self.LONG) // 2 + 3)
        self.assertGreater(self.header_height(sheet), 15)

    def test_a_header_that_fits_keeps_one_line_and_no_taller_row(self):
        self.assertIsNone(self.header_height(self.stacked("Gap")))

    def test_a_single_long_word_is_never_cut(self):
        word = "Extraordinarily"
        self.assertGreaterEqual(self.stacked(word).column_dimensions["D"].width, len(word) + 2)

    def test_body_text_still_sets_the_width_of_its_column(self):
        data = {"name": "Data", "columns": [{"name": "Department"}, {"name": "Units", "type": "quantity"}], "rows": [["Customer relationship management", 3], ["Legal", 4]]}
        sheet = self.build(declaration(data, [{"sheet": "Summary", "title": "By department", "rows": ["Department"], "measures": ["Units"]}]))
        self.assertGreaterEqual(sheet.column_dimensions["A"].width, len("Customer relationship management") + 2)


if __name__ == "__main__":
    unittest.main()
