from pathlib import Path
import re
import sys
import tempfile
import unittest
import zipfile

from doc_fixture import SCRIPTS_PATH, run_office, run_office_python, write_json

sys.path.insert(0, str(SCRIPTS_PATH))

from fonts.registry import SANS_BODY, default_family  # noqa: E402

DOCUMENT_FONT = default_family(SANS_BODY).name


def package_part(document_path, part_name):
    with zipfile.ZipFile(document_path) as archive:
        return archive.read(part_name).decode("utf-8")


def style_xml(styles, style_id):
    return re.search(rf'<w:style [^>]*w:styleId="{style_id}".*?</w:style>', styles, re.S).group(0)


class GeneratedDocumentTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.directory = Path(self.temporary_directory.name)

    def created(self, blocks=None, **specification):
        write_json(self.directory / "spec.json", {"title": "제목", "blocks": blocks or [{"type": "paragraph", "text": "본문"}], **specification})
        envelope = run_office(["create", "created.docx", "spec.json"], self.directory)
        self.assertEqual(envelope["status"], "ok", envelope)
        return self.directory / "created.docx"

    def exported(self, markdown):
        (self.directory / "source.md").write_text(markdown, encoding="utf-8")
        envelope = run_office(["create", "exported.docx", "source.md"], self.directory)
        self.assertNotEqual(envelope["status"], "error", envelope)
        return self.directory / "exported.docx"


class KoreanLanguageTest(GeneratedDocumentTest):
    def assert_korean_defaults(self, document_path):
        styles = package_part(document_path, "word/styles.xml")
        settings = package_part(document_path, "word/settings.xml")
        self.assertRegex(styles, r'<w:docDefaults>.*<w:lang [^>]*w:eastAsia="ko-KR"')
        self.assertRegex(settings, r'<w:themeFontLang [^>]*w:eastAsia="ko-KR"')
        for style_id in ("Title", "Heading1", "Heading2", "Heading3"):
            fonts = re.search(r"<w:rFonts [^>]*/>", style_xml(styles, style_id)).group(0)
            self.assertIn(f'w:eastAsia="{DOCUMENT_FONT}"', fonts, style_id)
            self.assertNotIn("Theme=", fonts, style_id)

    def test_created_document_is_korean_by_default(self):
        self.assert_korean_defaults(self.created([{"type": "heading", "text": "개요", "level": 1}]))

    def test_exported_document_is_korean_by_default(self):
        self.assert_korean_defaults(self.exported("# 개요\n\n본문\n"))


def list_paragraphs(document_path):
    document = package_part(document_path, "word/document.xml")
    return [
        (int(re.search(r'<w:ilvl w:val="(\d+)"', properties).group(1)), int(re.search(r'<w:numId w:val="(\d+)"', properties).group(1)))
        for properties in re.findall(r"<w:numPr>.*?</w:numPr>", document, re.S)
    ]


def number_format(document_path, number_id, level):
    numbering = package_part(document_path, "word/numbering.xml")
    abstract_id = re.search(rf'<w:num w:numId="{number_id}"[^>]*><w:abstractNumId w:val="(\d+)"', numbering).group(1)
    definition = re.search(rf'<w:abstractNum [^>]*w:abstractNumId="{abstract_id}".*?</w:abstractNum>', numbering, re.S).group(0)
    return re.search(rf'<w:lvl w:ilvl="{level}".*?<w:numFmt w:val="(\w+)"', definition, re.S).group(1)


class ListLevelTest(GeneratedDocumentTest):
    def test_exported_nested_lists_use_list_levels_and_each_numbered_list_restarts(self):
        path = self.exported("- 하나\n  - 둘\n    - 셋\n- 넷\n\n문단\n\n1. 가\n   1. 나\n2. 다\n\n문단\n\n1. 라\n")
        paragraphs = list_paragraphs(path)
        self.assertEqual([level for level, _ in paragraphs], [0, 1, 2, 0, 0, 1, 0, 0])
        bullet_list, first_numbered, second_numbered = paragraphs[0][1], paragraphs[4][1], paragraphs[7][1]
        self.assertEqual({number for _, number in paragraphs[:4]}, {bullet_list})
        self.assertEqual({number for _, number in paragraphs[4:7]}, {first_numbered})
        self.assertNotEqual(first_numbered, second_numbered)
        self.assertEqual([number_format(path, bullet_list, level) for level in range(3)], ["bullet"] * 3)
        self.assertEqual([number_format(path, first_numbered, level) for level in range(3)], ["decimal", "lowerLetter", "lowerRoman"])
        self.assertEqual(len(re.findall(r"<w:startOverride", package_part(path, "word/numbering.xml"))), 6)
        self.assertNotRegex(package_part(path, "word/document.xml"), r'w:pStyle w:val="List(Bullet|Number)')

    def test_created_lists_are_real_lists(self):
        path = self.created([
            {"type": "bullets", "items": ["하나", "둘"]},
            {"type": "numbered", "items": ["가", "나"]},
            {"type": "numbered", "items": ["다"]},
        ])
        paragraphs = list_paragraphs(path)
        self.assertEqual([level for level, _ in paragraphs], [0] * 5)
        self.assertEqual(len({number for _, number in paragraphs}), 3)
        self.assertEqual(number_format(path, paragraphs[0][1], 0), "bullet")
        self.assertEqual(number_format(path, paragraphs[2][1], 0), "decimal")


def page_setup(document_path):
    document = package_part(document_path, "word/document.xml")
    size = re.search(r"<w:pgSz ([^>]*)/>", document).group(1)
    margins = re.search(r"<w:pgMar ([^>]*)/>", document).group(1)
    return (
        {name: int(value) for name, value in re.findall(r'w:(w|h)="(\d+)"', size)},
        {name: int(value) for name, value in re.findall(r'w:(top|right|bottom|left)="(\d+)"', margins)},
        "landscape" in size,
    )


class PageTest(GeneratedDocumentTest):
    A4 = {"w": 11906, "h": 16838}
    ONE_INCH = {"top": 1440, "right": 1440, "bottom": 1440, "left": 1440}

    def test_created_page_is_a4_with_one_inch_margins(self):
        self.assertEqual(page_setup(self.created()), (self.A4, self.ONE_INCH, False))

    def test_exported_page_is_a4_with_one_inch_margins(self):
        self.assertEqual(page_setup(self.exported("본문\n")), (self.A4, self.ONE_INCH, False))

    def test_explicit_page_values_are_kept(self):
        size, margins, is_landscape = page_setup(self.created(page={"orientation": "landscape", "marginInches": 0.5}))
        self.assertEqual((size, set(margins.values()), is_landscape), ({"w": 16838, "h": 11906}, {720}, True))


def first_table(document_path):
    return re.search(r"<w:tbl>.*?</w:tbl>(<w:p>.*?</w:p>)?", package_part(document_path, "word/document.xml"), re.S)


class TableTest(GeneratedDocumentTest):
    ROWS = [["항목", "값"], ["매출", "100"]]

    def assert_default_table(self, document_path):
        table = first_table(document_path).group(0)
        first_row, second_row = re.findall(r"<w:tr[ >].*?</w:tr>", table, re.S)
        margins = re.search(r"<w:tblCellMar>(.*?)</w:tblCellMar>", table, re.S).group(1)
        self.assertRegex(margins, r'<w:top w:w="[1-9]\d*"')
        self.assertRegex(margins, r'<w:bottom w:w="[1-9]\d*"')
        self.assertIn("<w:tblHeader/>", first_row)
        self.assertEqual(len(re.findall(r'<w:shd [^>]*w:fill="EAF1F8"', first_row)), 2)
        self.assertEqual(len(re.findall(r"<w:b/>", first_row)), 2)
        self.assertNotIn("<w:shd", second_row)
        self.assertNotIn("<w:b/>", second_row)
        self.assertLess(table.index("<w:tblBorders>"), table.index("<w:tblCellMar>"))
        self.assertLess(table.index("<w:tblCellMar>"), table.index("<w:tblLook"))
        self.assertEqual(table.count("<w:tcW"), 4)
        spacer = first_table(document_path).group(1)
        self.assertRegex(spacer or "", r'<w:spacing [^>]*w:after="1[0-9]{2}"')

    def test_created_table_has_margins_a_shaded_bold_header_and_space_after(self):
        self.assert_default_table(self.created([{"type": "table", "rows": self.ROWS}]))

    def test_exported_table_has_margins_a_shaded_bold_header_and_space_after(self):
        self.assert_default_table(self.exported("| 항목 | 값 |\n| --- | --- |\n| 매출 | 100 |\n\n뒤\n"))


class ImageTest(GeneratedDocumentTest):
    def test_markdown_alt_text_becomes_the_picture_description(self):
        from PIL import Image
        Image.new("RGB", (40, 20)).save(self.directory / "chart.png")
        path = self.exported("![세그먼트별 이탈률 막대 차트](chart.png)\n")
        self.assertRegex(package_part(path, "word/document.xml"), r'<wp:docPr [^>]*descr="세그먼트별 이탈률 막대 차트"')


class KoreanLanguageCheckTest(GeneratedDocumentTest):
    def codes(self, path):
        return [issue["code"] for issue in run_office(["check", path.name], self.directory)["issues"]]

    def test_a_generated_document_passes(self):
        self.assertNotIn("EAST_ASIA_LANGUAGE_NOT_KOREAN", self.codes(self.created()))
        self.assertNotIn("EAST_ASIA_LANGUAGE_NOT_KOREAN", self.codes(self.exported("본문\n")))

    def test_a_run_or_style_tagged_for_another_language_is_flagged_and_fixed(self):
        path = self.created([{"type": "paragraph", "text": "본문"}, {"type": "heading", "text": "제목", "level": 1}])
        run_office_python("""
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            document = Document("created.docx")
            run = document.paragraphs[-2].runs[0]._r
            language = OxmlElement("w:lang")
            language.set(qn("w:eastAsia"), "ja-JP")
            run.get_or_add_rPr().append(language)
            document.save("created.docx")
        """, self.directory)
        envelope = run_office(["check", path.name], self.directory)
        issue = next(issue for issue in envelope["issues"] if issue["code"] == "EAST_ASIA_LANGUAGE_NOT_KOREAN")
        write_json(self.directory / "fix.json", issue["fix"])
        self.assertEqual(run_office(["apply", path.name, "fix.json"], self.directory)["status"], "ok")
        self.assertNotIn("EAST_ASIA_LANGUAGE_NOT_KOREAN", self.codes(path))
        self.assertNotIn("ja-JP", package_part(path, "word/document.xml"))

    def test_a_style_tag_that_overrides_the_default_is_flagged(self):
        path = self.created()
        run_office_python("""
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            document = Document("created.docx")
            language = OxmlElement("w:lang")
            language.set(qn("w:eastAsia"), "zh-CN")
            document.styles["Normal"].element.get_or_add_rPr().append(language)
            document.save("created.docx")
        """, self.directory)
        self.assertIn("EAST_ASIA_LANGUAGE_NOT_KOREAN", self.codes(path))


if __name__ == "__main__":
    unittest.main()
