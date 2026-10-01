import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest


OFFICE_ENTRY = Path(__file__).resolve().parents[1] / "skills" / "office" / "scripts" / "office"
READERS = {
    "docx": ("doc read", "doc check", "doc validate", "doc render"),
    "xlsx": ("sheet read", "sheet check", "sheet validate", "sheet render"),
    "pptx": ("deck read", "deck validate", "deck check"),
    "pdf": ("pdf read", "pdf render", "pdf validate"),
}


def run_office(arguments, working_directory):
    completed = subprocess.run([str(OFFICE_ENTRY), *arguments], cwd=working_directory, capture_output=True, text=True, timeout=300)
    return json.loads(completed.stdout), completed.stderr


def write_inputs(directory):
    script = """
import sys
from pathlib import Path
import docx, openpyxl, pptx
from pypdf import PdfWriter
directory = Path(sys.argv[1])
docx.Document().save(directory / "plain.docx")
openpyxl.Workbook().save(directory / "plain.xlsx")
pptx.Presentation().save(directory / "plain.pptx")
writer = PdfWriter(); writer.add_blank_page(200, 200); writer.write(directory / "plain.pdf")
writer = PdfWriter(); writer.add_blank_page(200, 200); writer.encrypt("sample-password"); writer.write(directory / "locked.pdf")
(directory / "notes.txt").write_text("이샘플의 메모", encoding="utf-8")
(directory / "legacy.xls").write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + bytes(504))
"""
    subprocess.run([str(OFFICE_ENTRY), "python", "-c", script, str(directory)], check=True, capture_output=True, timeout=300)


class InputBoundaryTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary_directory = tempfile.TemporaryDirectory()
        cls.directory = Path(cls.temporary_directory.name)
        write_inputs(cls.directory)

    @classmethod
    def tearDownClass(cls):
        cls.temporary_directory.cleanup()

    def assert_refused(self, command, input_name, code):
        envelope, stderr = run_office([*command.split(), input_name], self.directory)
        self.assertNotIn("Traceback", stderr)
        self.assertEqual(envelope["status"], "error", f"{command} {input_name}")
        self.assertEqual(envelope["issues"][0]["code"], code, f"{command} {input_name}")
        return envelope["issues"][0]

    def test_every_reader_names_the_command_for_a_file_of_another_kind(self):
        for kind, commands in READERS.items():
            other_kind = "xlsx" if kind != "xlsx" else "docx"
            for command in commands:
                with self.subTest(command=command):
                    issue = self.assert_refused(command, f"plain.{other_kind}", "WRONG_INPUT_FORMAT")
                    self.assertIn("office doc read" if other_kind == "docx" else "office sheet read", issue["suggestion"])

    def test_every_reader_answers_a_missing_file(self):
        for commands in READERS.values():
            for command in commands:
                with self.subTest(command=command):
                    self.assert_refused(command, "missing.docx", "INPUT_NOT_FOUND")

    def test_text_and_legacy_files_are_refused_with_what_they_are(self):
        self.assertIn("neither an Office file nor a PDF", self.assert_refused("doc read", "notes.txt", "WRONG_INPUT_FORMAT")["message"])
        self.assertIn("office convert", self.assert_refused("sheet read", "legacy.xls", "WRONG_INPUT_FORMAT")["suggestion"])

    def test_a_password_protected_pdf_is_named_by_every_pdf_reader_and_convert(self):
        for command in READERS["pdf"]:
            with self.subTest(command=command):
                self.assert_refused(command, "locked.pdf", "PDF_PASSWORD_REQUIRED")
        envelope, stderr = run_office(["convert", "locked.pdf", "locked.docx"], self.directory)
        self.assertNotIn("Traceback", stderr)
        self.assertEqual(envelope["issues"][0]["code"], "PDF_PASSWORD_REQUIRED")

    def test_apply_checks_its_input_before_reading_the_operations(self):
        (self.directory / "ops.json").write_text("[]", encoding="utf-8")
        for command, other in (("doc apply", "plain.xlsx"), ("sheet apply", "plain.pptx"), ("deck apply", "plain.docx")):
            with self.subTest(command=command):
                envelope, stderr = run_office([*command.split(), other, "ops.json"], self.directory)
                self.assertNotIn("Traceback", stderr)
                self.assertEqual(envelope["issues"][0]["code"], "WRONG_INPUT_FORMAT")


if __name__ == "__main__":
    unittest.main()
