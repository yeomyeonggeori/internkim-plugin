import datetime
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from pdf_fixture import copy_pdf_fixture
from sheet_fixture import SCRIPTS_PATH, run_office, run_office_python
from xlsb_fixture import write_xlsb



class ConversionFixture(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.directory = Path(self.temporary_directory.name)

    def convert(self, source, target):
        return run_office(["convert", source, target], self.directory)


class DelimitedRouteTest(ConversionFixture):
    def test_the_sheet_reference_sends_a_csv_through_create_before_any_guide(self):
        reference = (SCRIPTS_PATH.parent / "references" / "sheet.md").read_text(encoding="utf-8")
        self.assertLess(reference.index("office create ~/documents/<title>.xlsx <data.csv>"), reference.index("office guide"))
        self.assertNotIn("Python's `csv`", reference)
        self.assertNotIn("newest", reference)

    def test_excel_unicode_text_in_utf16_converts_whatever_its_extension(self):
        exported = "\ufeff지역\t실적\n서울\t6200\n".encode("utf-16-le")
        for name in ("유니코드.tsv", "유니코드.csv"):
            with self.subTest(name=name):
                (self.directory / name).write_bytes(exported)
                self.assertEqual(run_office(["create", "유니코드.xlsx", name], self.directory)["status"], "ok")
                values = run_office(["read", "유니코드.xlsx"], self.directory)["details"]["range"]["values"]
                self.assertEqual(values, [["지역", "실적"], ["서울", 6200]])

    def test_a_converted_csv_takes_a_summary_sheet_from_one_apply(self):
        (self.directory / "판매.csv").write_text("월,지역,실적\n2026-04,서울,6200\n2026-04,경기,4140\n2026-05,서울,6280\n", encoding="utf-8")
        self.assertEqual(run_office(["create", "판매.xlsx", "판매.csv"], self.directory)["status"], "ok")
        (self.directory / "ops.json").write_text("""[
            {"op": "add_sheet", "name": "요약"},
            {"op": "set_range", "sheet": "요약", "cell": "A1", "values": [["지역", "실적"], ["서울", "=SUMIFS(판매!C:C,판매!B:B,A2)"]]}
        ]""", encoding="utf-8")
        self.assertEqual(run_office(["apply", "판매.xlsx", "ops.json"], self.directory)["status"], "ok")
        values = run_office(["read", "판매.xlsx", "--sheet", "요약"], self.directory)["details"]["range"]["values"]
        self.assertEqual(values, [["지역", "실적"], ["서울", 12480]])


class BinaryWorkbookTest(ConversionFixture):
    def test_an_xlsb_becomes_a_workbook_with_types_merges_and_sheets(self):
        write_xlsb(self.directory / "실적.xlsb", [
            ("매출", [["지역", "날짜", "금액", "확정"], ["서울", ("date", 45306), 1200.5, True], ["부산", None, 3000, False], ["합계", None, None, None]], [(3, 3, 0, 1)]),
            ("메모", [["내용"], ["첫 줄"]], []),
        ])
        envelope = self.convert("실적.xlsb", "실적.xlsx")
        self.assertEqual(envelope["status"], "warning", envelope)
        self.assertEqual(envelope["details"]["route"], "xlsb -> xlsx")
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["CONVERSION_APPROXIMATED"])
        workbook = load_workbook(self.directory / "실적.xlsx")
        self.assertEqual(workbook.sheetnames, ["매출", "메모"])
        sales = workbook["매출"]
        self.assertEqual([cell.value for cell in sales[2]], ["서울", datetime.datetime(2024, 1, 15), 1200.5, True])
        self.assertEqual([cell.value for cell in sales[3]], ["부산", None, 3000, False])
        self.assertEqual(sales["B2"].number_format, "yyyy-mm-dd")
        self.assertEqual([str(merged) for merged in sales.merged_cells.ranges], ["A4:B4"])


class PdfTablesTest(ConversionFixture):
    def test_pdf_tables_become_typed_sheets_and_a_repeated_header_continues_one(self):
        copy_pdf_fixture("tables.pdf", self.directory)
        envelope = self.convert("tables.pdf", "tables.xlsx")
        self.assertEqual(envelope["details"]["route"], "pdf -> xlsx")
        self.assertEqual(envelope["details"]["tables"], [
            {"sheet": "Page 1", "pages": [1, 2], "rows": 5, "columns": 4, "header": True},
            {"sheet": "Page 3", "pages": [3], "rows": 2, "columns": 4, "header": False},
        ])
        self.assertIn("page 4 has no table", envelope["issues"][0]["message"])
        workbook = load_workbook(self.directory / "tables.xlsx")
        first = workbook["Page 1"]
        self.assertEqual([cell.value for cell in first[1]], ["지점", "매출", "증감률", "기준일"])
        self.assertEqual([cell.value for cell in first[2]], ["서울 본점", 1200000, 0.125, datetime.datetime(2024, 1, 15)])
        self.assertEqual([cell.value for cell in first[3]][1:3], [-3000, -0.04])
        self.assertEqual([cell.value for cell in first[4]][1:], [850500, 0, datetime.datetime(2024, 3, 31)])
        self.assertEqual(first["B4"].number_format, "\"₩\"#,##0")
        self.assertEqual([cell.value for cell in first[5]][1:], ["007", 3.25, "미정"])
        self.assertEqual((first["B2"].number_format, first["C2"].number_format, first["D2"].number_format), ("#,##0", "0.0%", "yyyy-mm-dd"))
        self.assertEqual(first.freeze_panes, "A2")
        third = workbook["Page 3"]
        self.assertEqual([cell.value for cell in third[1]], [1, 2, 3, 4])
        self.assertIsNone(third.freeze_panes)

    def test_a_table_laid_out_without_lines_becomes_a_typed_sheet(self):
        copy_pdf_fixture("statement.pdf", self.directory)
        envelope = self.convert("statement.pdf", "statement.xlsx")
        self.assertEqual(envelope["status"], "ok", envelope["issues"])
        self.assertEqual(envelope["details"]["tables"], [{"sheet": "Page 1", "pages": [1], "rows": 7, "columns": 5, "header": True}])
        sheet = load_workbook(self.directory / "statement.xlsx")["Page 1"]
        self.assertEqual([cell.value for cell in sheet[1]], ["일자", "품목", "수량", "단가", "금액"])
        self.assertEqual([cell.value for cell in sheet[4]], [datetime.datetime(2026, 9, 16), "A4 복사 용지 (박스)", 45, 26500, 1192500])
        self.assertEqual([cell.value for cell in sheet[7]], ["합계", None, None, None, 4024500])
        self.assertEqual(sheet["E2"].number_format, "#,##0")

    def test_a_pdf_without_tables_is_refused_with_what_to_do(self):
        copy_pdf_fixture("text.pdf", self.directory)
        envelope = self.convert("text.pdf", "text.xlsx")
        self.assertEqual(envelope["status"], "error")
        self.assertEqual(envelope["issues"][0]["code"], "TABLE_NOT_FOUND")
        self.assertFalse((self.directory / "text.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
