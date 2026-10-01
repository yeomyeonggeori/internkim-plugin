import sys
import unittest
from pathlib import Path

from pptx.util import Pt

from test_deck_pptx_editing import KoreanDeckFixture, codes


SCRIPTS_PATH = Path(__file__).resolve().parents[1] / "skills" / "office" / "scripts"
sys.path.insert(0, str(SCRIPTS_PATH))

from deck.table_styles import TABLE_STYLES  # noqa: E402


A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
TABLE = {"slide": 4, "shape": 1}


class TableStyleTest(KoreanDeckFixture):
    def table(self):
        return self.presentation().slides[3].shapes[1].table

    def test_style_and_parts_are_written_and_read_back(self):
        envelope = self.apply([{"op": "set_table_style", **TABLE, "style": "Medium Style 2 - Accent 2", "bandRow": False, "firstCol": True}])
        self.assertEqual(envelope["status"], "ok", envelope)
        described = self.slide(4, "--detail")["shapes"][1]
        self.assertEqual(described["style"], "Medium Style 2 - Accent 2")
        self.assertEqual(described["styleParts"], ["firstRow", "firstCol"])
        properties = self.table()._tbl.tblPr
        self.assertEqual(properties.find(f"{A}tableStyleId").text, "{21E4AEA4-8DFA-4A89-87EB-49C32662AFE0}")

    def test_every_style_name_has_its_own_identifier(self):
        self.assertEqual(len({style.identifier for style in TABLE_STYLES}), len(TABLE_STYLES))

    def test_a_style_change_with_no_field_is_refused(self):
        self.assertEqual(codes(self.apply([{"op": "set_table_style", **TABLE}])), ["INVALID_VALUE"])

    def test_a_style_change_on_a_text_box_names_the_table(self):
        envelope = self.apply([{"op": "set_table_style", "slide": 4, "shape": 0, "firstRow": False}])
        self.assertEqual(codes(envelope), ["OPERATION_NOT_APPLICABLE"])
        self.assertIn("shape 1", envelope["issues"][0]["suggestion"])


class TableCellFormatTest(KoreanDeckFixture):
    def table(self):
        return self.presentation().slides[3].shapes[1].table

    def test_a_header_row_takes_fill_text_style_alignment_and_edges(self):
        envelope = self.apply([{"op": "format_table_cells", **TABLE, "row": 0, "fill": "1F4E79", "bold": True, "color": "FFFFFF", "size": 14, "align": "center", "anchor": "middle", "borderColor": "0F2740", "borderWidth": 2}])
        self.assertEqual(envelope["status"], "ok", envelope)
        self.assertIn("3 cells", envelope["details"]["changes"][0]["change"])
        table = self.table()
        for column in range(3):
            cell = table.cell(0, column)
            run = cell.text_frame.paragraphs[0].runs[0]
            self.assertEqual(str(cell.fill.fore_color.rgb), "1F4E79")
            self.assertEqual((run.font.bold, str(run.font.color.rgb), run.font.size), (True, "FFFFFF", Pt(14)))
            self.assertEqual(cell.text_frame.paragraphs[0].alignment, 2)
            self.assertEqual(cell._tc.tcPr.get("anchor"), "ctr")
            edges = [cell._tc.tcPr.find(f"{A}{tag}") for tag in ("lnL", "lnR", "lnT", "lnB")]
            self.assertEqual([edge.get("w") for edge in edges], [str(Pt(2))] * 4)
            self.assertEqual(list(cell._tc.tcPr).index(edges[-1]) < list(cell._tc.tcPr).index(cell._tc.tcPr.find(f"{A}solidFill")), True)
        self.assertIsNone(table.cell(1, 0).fill._xPr.find(f"{A}solidFill"))

    def test_a_block_of_cells_is_formatted_and_nothing_else(self):
        self.apply([{"op": "format_table_cells", **TABLE, "row": 1, "rows": 2, "column": 2, "italic": True}])
        table = self.table()
        self.assertTrue(table.cell(2, 2).text_frame.paragraphs[0].runs[0].font.italic)
        self.assertIsNone(table.cell(2, 1).text_frame.paragraphs[0].runs[0].font.italic)

    def test_a_block_past_the_table_is_refused(self):
        envelope = self.apply([{"op": "format_table_cells", **TABLE, "row": 2, "rows": 3, "bold": True}])
        self.assertEqual(codes(envelope), ["TARGET_NOT_FOUND"])
        self.assertEqual(envelope["issues"][0]["location"], "ops[0].rows")

    def test_a_border_width_without_a_color_is_refused(self):
        self.assertEqual(codes(self.apply([{"op": "format_table_cells", **TABLE, "borderWidth": 2}])), ["INVALID_VALUE"])


class TableSizeTest(KoreanDeckFixture):
    def test_a_column_width_change_resizes_the_frame_by_the_difference(self):
        before = self.slide(4)["shapes"][1]["box"]
        envelope = self.apply([{"op": "set_table_column_width", **TABLE, "column": 0, "width": "3in"}, {"op": "set_table_row_height", **TABLE, "row": 0, "height": "0.75in"}])
        self.assertEqual(envelope["status"], "ok", envelope)
        table = self.presentation().slides[3].shapes[1].table
        self.assertEqual((table.columns[0].width, table.rows[0].height), (3 * 914400, 685800))
        after = self.slide(4)["shapes"][1]["box"]
        self.assertEqual(after["w"] - before["w"], 3 * 914400 - 2032000)
        self.assertEqual(after["h"] - before["h"], 685800 - 457200)


if __name__ == "__main__":
    unittest.main()
