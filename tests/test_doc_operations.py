import json
import re
import subprocess
import sys
import unittest
import zipfile

from doc_fixture import OFFICE_ENTRY, SCRIPTS_PATH, ContractFixture, block_texts, read_details, run_office, run_office_python, write_json


RICH_ADDITIONS = """
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image

Image.new("RGB", (320, 120), (40, 90, 160)).save("logo.png")
document = Document("contract.docx")
section = document.sections[0]
section.header.paragraphs[0].text = "주식회사 예시상사"
section.footer.paragraphs[0].text = "대외비 문서"
document.paragraphs[1]._p.append(parse_xml(
    f'<w:hyperlink {nsdecls("w")} w:anchor="purpose"><w:r><w:t>(목적 참조)</w:t></w:r></w:hyperlink>'
))
document.paragraphs[2]._p.insert(1, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="7" w:name="purpose"/>'))
document.paragraphs[2]._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="7"/>'))
document.add_comment(document.paragraphs[5].runs, text="금액 확인 필요", author="최견본")
document.add_picture("logo.png")
schedule = document.add_table(rows=2, cols=3)
schedule.style = "Table Grid"
schedule.cell(0, 0).merge(schedule.cell(0, 2)).text = "일정"
for cell, value in zip(schedule.rows[1].cells, ["착수", "중간", "완료"]):
    cell.text = value
document.save("contract.docx")
"""

CANONICAL_PARTS = """
import json
import sys
import zipfile
from lxml import etree

def canonical(data):
    return etree.tostring(etree.fromstring(data), method="c14n")

def body_blocks(data):
    body = etree.fromstring(data).find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
    return [etree.tostring(child, method="c14n").decode() for child in body]

before, after = zipfile.ZipFile(sys.argv[1]), zipfile.ZipFile(sys.argv[2])
changed = sorted(
    name for name in before.namelist()
    if name.endswith((".xml", ".rels")) and name != "word/document.xml" and canonical(before.read(name)) != canonical(after.read(name))
)
binary_changed = sorted(name for name in before.namelist() if not name.endswith((".xml", ".rels")) and before.read(name) != after.read(name))
original_blocks = body_blocks(before.read("word/document.xml"))
edited_blocks = set(body_blocks(after.read("word/document.xml")))
kept = [index for index, block in enumerate(original_blocks) if block in edited_blocks]
print(json.dumps({"changedParts": changed, "changedBinaryParts": binary_changed, "keptBlocks": kept, "added": sorted(set(after.namelist()) - set(before.namelist()))}))
"""


class RichDocumentFixture(ContractFixture):
    def setUp(self):
        super().setUp()
        run_office_python(RICH_ADDITIONS, self.directory)

    def apply(self, operations, *flags, output="edited.docx"):
        write_json(self.directory / "ops.json", operations)
        envelope = run_office(["apply", "contract.docx", "ops.json", "--output", output, *flags], self.directory)
        self.assertEqual(envelope["status"], "ok", envelope["issues"])
        return envelope

    def compare(self, edited="edited.docx"):
        (self.directory / "compare.py").write_text(CANONICAL_PARTS, encoding="utf-8")
        completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "compare.py", "contract.docx", edited], capture_output=True, text=True, check=True, cwd=self.directory)
        return json.loads(completed.stdout)


class ContentOperationTest(RichDocumentFixture):
    def test_text_and_paragraph_formatting_touch_only_their_blocks(self):
        self.apply([
            {"op": "format_text", "find": "일천만 원", "bold": True, "color": "#C00000"},
            {"op": "set_paragraph_format", "block": 7, "align": "justify", "spaceAfterPoints": 12},
            {"op": "insert_link", "block": 3, "find": "홈페이지", "url": "https://www.example.com"},
        ])
        comparison = self.compare()
        self.assertEqual(comparison["changedParts"], ["word/_rels/document.xml.rels"])
        self.assertEqual(comparison["changedBinaryParts"], [])
        self.assertEqual(comparison["keptBlocks"], [0, 1, 2, 4, 6, 8, 9, 10, 11, 12])
        self.assertEqual(block_texts(self.directory, "edited.docx"), block_texts(self.directory, "contract.docx"))
        with zipfile.ZipFile(self.directory / "edited.docx") as archive:
            document_xml = archive.read("word/document.xml").decode()
        self.assertIn('<w:color w:val="C00000"/>', document_xml)
        self.assertIn('<w:jc w:val="both"/>', document_xml)

    def test_paragraphs_insert_at_the_start_and_the_end_by_name(self):
        self.apply([
            {"op": "insert_paragraph", "at": "end", "text": "끝 하나"},
            {"op": "insert_paragraph", "at": "start", "text": "처음"},
            {"op": "insert_paragraph", "at": "end", "text": "끝 둘"},
            {"op": "insert_table_column", "block": 8, "at": "start", "cells": ["번호", "1", "2"]},
        ])
        texts = [text for _, text in block_texts(self.directory, "edited.docx")]
        self.assertEqual((texts[0], texts[-2], texts[-1]), ("처음", "끝 하나", "끝 둘"))
        self.assertEqual(read_details(self.directory, "edited.docx")["blocks"][9]["cells"][0][0], "번호")
        write_json(self.directory / "ops.json", [{"op": "insert_paragraph", "after": -1, "text": "x"}])
        refused = run_office(["apply", "contract.docx", "ops.json"], self.directory)
        self.assertEqual(refused["status"], "error")

    def test_notes_bookmarks_and_cross_references_read_back(self):
        self.apply([
            {"op": "insert_footnote", "block": 5, "text": "부가가치세는 세금계산서 발행일에 청구한다.", "afterText": "별도로 한다."},
            {"op": "add_bookmark", "block": 6, "name": "지급조건"},
            {"op": "insert_cross_reference", "block": 9, "bookmark": "지급조건"},
        ])
        details = read_details(self.directory, "edited.docx")
        self.assertEqual(details["notes"], [{"kind": "footnote", "id": 1, "text": "부가가치세는 세금계산서 발행일에 청구한다."}])
        self.assertEqual(details["bookmarks"], ["purpose", "지급조건"])
        self.assertTrue(details["blocks"][9]["text"].endswith("보관한다.제3조 (지급 조건)"))
        comparison = self.compare()
        self.assertIn("word/footnotes.xml", comparison["added"])
        codes = {issue["code"] for issue in run_office(["check", "edited.docx"], self.directory)["issues"]}
        self.assertNotIn("BROKEN_INTERNAL_REFERENCE", codes)

    def test_a_table_of_contents_lists_the_headings_and_asks_word_to_refresh(self):
        self.apply([{"op": "insert_table_of_contents", "after": 1, "title": "목차", "levels": 1}])
        texts = [text for _, text in block_texts(self.directory, "edited.docx")]
        self.assertEqual(texts[2:6], ["목차", "제1조 (목적)", "제2조 (계약 금액)", "제3조 (지급 조건)"])
        codes = {issue["code"] for issue in run_office(["check", "edited.docx"], self.directory)["issues"]}
        self.assertNotIn("STALE_TABLE_OF_CONTENTS", codes)
        self.assertEqual(self.compare()["changedParts"], ["word/settings.xml"])

    def test_tracked_formatting_can_be_rejected(self):
        self.apply([{"op": "format_text", "block": 5, "find": "일천만 원", "bold": True}], "--track", "--author", "박예시")
        revisions = read_details(self.directory, "edited.docx", "--revisions")["revisions"]
        self.assertIn(("formatting", "박예시", ["bold"]), [(revision["type"], revision["author"], revision.get("changed")) for revision in revisions])
        write_json(self.directory / "reject.json", [{"op": "reject_revisions", "author": "박예시"}])
        run_office(["apply", "edited.docx", "reject.json"], self.directory)
        with zipfile.ZipFile(self.directory / "edited.docx") as archive:
            amount_paragraph = re.search(r"<w:p[ >](?:(?!</w:p>).)*일천만 원(?:(?!</w:p>).)*</w:p>", archive.read("word/document.xml").decode()).group(0)
        self.assertNotIn("<w:b/>", amount_paragraph)
        self.assertNotIn("rPrChange", amount_paragraph)


class TableOperationTest(RichDocumentFixture):
    def test_columns_insert_and_delete_and_cells_merge_and_shade(self):
        self.apply([
            {"op": "insert_table_column", "block": 8, "after": 1, "cells": ["지급 시기", "착수 시", "검수 후"]},
            {"op": "format_cells", "block": 8, "row": 0, "column": 0, "toColumn": 1, "fill": "#DCE6F1", "bold": True},
            {"op": "merge_cells", "block": 11, "row": 1, "column": 1, "toColumn": 2},
        ])
        details = read_details(self.directory, "edited.docx")
        self.assertEqual(details["blocks"][8]["cells"], [["구분", "금액", "지급 시기"], ["착수금", "3,000,000", "착수 시"], ["잔금", "7,000,000", "검수 후"]])
        self.assertEqual(details["blocks"][11]["cells"][1], ["착수", "중간\n완료", "중간\n완료"])
        write_json(self.directory / "ops.json", [{"op": "delete_table_column", "block": 11, "column": 0}])
        refused = run_office(["apply", "edited.docx", "ops.json"], self.directory)
        self.assertEqual([issue["code"] for issue in refused["issues"]], ["OPERATION_NOT_APPLICABLE"])
        write_json(self.directory / "ops.json", [{"op": "delete_table_column", "block": 8, "column": 2}])
        self.assertEqual(run_office(["apply", "edited.docx", "ops.json"], self.directory)["status"], "ok")
        self.assertEqual(read_details(self.directory, "edited.docx")["blocks"][8]["columns"], 2)
        with zipfile.ZipFile(self.directory / "edited.docx") as archive:
            self.assertIn('w:fill="DCE6F1"', archive.read("word/document.xml").decode())


class PageOperationTest(RichDocumentFixture):
    def test_page_setup_sections_headers_and_watermark(self):
        self.apply([
            {"op": "set_page_setup", "paper": "A4", "marginInches": 0.8},
            {"op": "insert_section_break", "after": 9, "orientation": "landscape"},
            {"op": "set_footer", "text": "{PAGE} / {NUMPAGES}", "align": "center"},
            {"op": "set_watermark", "text": "대외비"},
        ])
        details = read_details(self.directory, "edited.docx")
        self.assertEqual([(section["orientation"], section["marginsInches"]) for section in details["sections"]], [("portrait", [0.8] * 4), ("landscape", [0.8] * 4)])
        self.assertEqual(details["sections"][0]["header"], "주식회사 예시상사")
        with zipfile.ZipFile(self.directory / "edited.docx") as archive:
            parts = {name: archive.read(name).decode() for name in archive.namelist() if name.startswith("word/header") or name.startswith("word/footer")}
        self.assertTrue(any('string="대외비"' in text for text in parts.values()))
        self.assertTrue(any('w:instr=" NUMPAGES "' in text for text in parts.values()))

    def test_an_image_scales_to_the_given_width(self):
        envelope = self.apply([{"op": "insert_image", "after": 1, "path": "logo.png", "widthInches": 2, "description": "회사 로고"}])
        self.assertIn("2.00 x 0.75 inches", envelope["details"]["changes"][0]["change"])
        self.assertTrue(read_details(self.directory, "edited.docx")["blocks"][2]["picture"])


class SectionPropertyOrderTest(unittest.TestCase):
    def test_page_numbering_is_inserted_before_what_python_docx_orders_after_it(self):
        sys.path.insert(0, str(SCRIPTS_PATH))
        from docx.oxml.section import CT_SectPr
        from doc.operations.pages import SECTION_PROPERTIES_AFTER_PAGE_NUMBERING
        page_margin_successors = CT_SectPr._insert_pgMar.__closure__[0].cell_contents._successors
        after_page_numbering = page_margin_successors[page_margin_successors.index("w:pgNumType") + 1:]
        self.assertEqual(SECTION_PROPERTIES_AFTER_PAGE_NUMBERING, after_page_numbering)


if __name__ == "__main__":
    unittest.main()
