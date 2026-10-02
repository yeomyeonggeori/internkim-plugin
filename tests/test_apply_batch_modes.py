import unittest

from openpyxl import load_workbook

from doc_fixture import DocumentFixture, block_texts, run_office, write_json
from sheet_fixture import WorkbookFixture
from test_deck_pptx_editing import KoreanDeckFixture, codes


class DocumentBatchModeTest(DocumentFixture):
    OPERATIONS = [
        {"op": "set_text", "block": 1, "text": "바뀐 첫 문단"},
        {"op": "set_cell", "block": 3, "row": 9, "column": 0, "text": "없음"},
        {"op": "set_text", "block": 2, "text": "바뀐 둘째 문단"},
    ]

    def apply(self, operations, *options):
        write_json(self.directory / "ops.json", operations)
        return run_office(["apply", "fixture.docx", "ops.json", *options], self.directory)

    def paragraphs(self):
        return [text for _, text in block_texts(self.directory, "fixture.docx")[1:3]]

    def test_the_default_mode_writes_nothing_when_one_operation_fails(self):
        original = (self.directory / "fixture.docx").read_bytes()
        envelope = self.apply(self.OPERATIONS)
        self.assertEqual(codes(envelope), ["TARGET_NOT_FOUND"])
        self.assertEqual((self.directory / "fixture.docx").read_bytes(), original)

    def test_best_effort_writes_every_operation_that_applies_and_reports_the_refused_one(self):
        envelope = self.apply(self.OPERATIONS, "--mode", "best-effort")
        self.assertEqual(envelope["status"], "warning", envelope)
        self.assertEqual(codes(envelope), ["OPERATION_REFUSED"])
        self.assertIn("TARGET_NOT_FOUND", envelope["issues"][0]["message"])
        self.assertEqual(envelope["issues"][0]["location"], "ops[1].row")
        self.assertEqual(envelope["details"]["refused"], [{"index": 1, "op": "set_cell", "codes": ["TARGET_NOT_FOUND"]}])
        self.assertEqual([change["index"] for change in envelope["details"]["changes"]], [0, 2])
        self.assertIn("2 of 3 operations", envelope["summary"])
        self.assertEqual(self.paragraphs(), ["바뀐 첫 문단", "바뀐 둘째 문단"])

    def test_stop_on_error_writes_the_operations_before_the_first_failure(self):
        envelope = self.apply(self.OPERATIONS, "--mode", "stop-on-error")
        self.assertEqual(codes(envelope), ["OPERATION_REFUSED"])
        self.assertEqual(envelope["details"]["skipped"], [2])
        self.assertEqual(self.paragraphs(), ["바뀐 첫 문단", "둘째 문단"])

    def test_an_operation_with_a_wrong_field_is_refused_alone(self):
        envelope = self.apply([{"op": "set_text", "block": 1, "text": "바뀐 첫 문단"}, {"op": "set_text", "block": 2, "txt": "오타"}], "--mode", "best-effort")
        self.assertEqual(envelope["details"]["refused"][0]["codes"], ["UNKNOWN_FIELD"])
        self.assertEqual(self.paragraphs()[0], "바뀐 첫 문단")

    def test_a_batch_where_nothing_applies_is_an_error_and_writes_nothing(self):
        original = (self.directory / "fixture.docx").read_bytes()
        envelope = self.apply([self.OPERATIONS[1]], "--mode", "best-effort")
        self.assertEqual(envelope["status"], "error")
        self.assertEqual(codes(envelope), ["TARGET_NOT_FOUND"])
        self.assertEqual((self.directory / "fixture.docx").read_bytes(), original)


class WorkbookBatchModeTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        self.create_workbook([{"title": "Sales", "rows": [["item", "amount"], ["A", 10]]}])

    def test_best_effort_applies_in_order_around_a_refused_operation(self):
        envelope = self.apply([
            {"op": "set_cell", "sheet": "Sales", "cell": "C1", "value": "note"},
            {"op": "set_cell", "sheet": "Missing", "cell": "A1", "value": 1},
            {"op": "set_cell", "sheet": "Sales", "cell": "C2", "value": "kept"},
        ], "--mode", "best-effort")
        self.assertEqual(codes(envelope), ["OPERATION_REFUSED"], envelope)
        worksheet = load_workbook(self.directory / "book.xlsx")["Sales"]
        self.assertEqual((worksheet["C1"].value, worksheet["C2"].value), ("note", "kept"))

    def test_stop_on_error_leaves_the_rest_untried(self):
        envelope = self.apply([
            {"op": "set_cell", "sheet": "Sales", "cell": "C1", "value": "note"},
            {"op": "set_cell", "sheet": "Missing", "cell": "A1", "value": 1},
            {"op": "set_cell", "sheet": "Sales", "cell": "C2", "value": "skipped"},
        ], "--mode", "stop-on-error")
        self.assertEqual(envelope["details"]["skipped"], [2], envelope)
        worksheet = load_workbook(self.directory / "book.xlsx")["Sales"]
        self.assertEqual((worksheet["C1"].value, worksheet["C2"].value), ("note", None))


class DeckBatchModeTest(KoreanDeckFixture):
    def test_best_effort_keeps_the_layout_review_of_the_slides_it_edited(self):
        envelope = self.apply([
            {"op": "set_text", "slide": 1, "shape": 0, "text": "새 표지 제목"},
            {"op": "set_transform", "slide": 2, "shape": 1, "w": "-1in"},
            {"op": "set_text", "slide": 9, "shape": 0, "text": "없는 슬라이드"},
        ], "--mode", "best-effort")
        self.assertEqual(codes(envelope)[:2], ["OPERATION_REFUSED", "OPERATION_REFUSED"], envelope)
        self.assertEqual([refusal["index"] for refusal in envelope["details"]["refused"]], [1, 2])
        self.assertEqual(envelope["details"]["auditedSlides"], [1])
        self.assertEqual(self.slide(1)["shapes"][0]["text"], "새 표지 제목")


if __name__ == "__main__":
    unittest.main()
