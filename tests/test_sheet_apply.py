import unittest
import zipfile

from openpyxl import load_workbook

from sheet_fixture import WorkbookFixture, run_office, run_office_python, write_json

FIXTURE_WORKBOOK = """
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName

workbook = Workbook()
sales = workbook.active
sales.title = "Sales"
for row in [["item", "amount"], ["A", 10], ["B", 20], ["C", 30]]:
    sales.append(row)
sales["B5"] = "=SUM(B2:B4)"
sales["D1"] = "=$B$5*2+B2"
sales["D2"] = "=AVERAGE(B2:B4)"
sales.auto_filter.ref = "A1:B4"
sales.merge_cells("D3:E3")
sales.column_dimensions["B"].width = 20
chart = BarChart()
chart.add_data(Reference(sales, min_col=2, min_row=1, max_row=4), titles_from_data=True)
chart.set_categories(Reference(sales, min_col=1, min_row=2, max_row=4))
sales.add_chart(chart, "G2")
summary = workbook.create_sheet("Summary")
summary["A1"] = "=Sales!B5"
summary["A2"] = "=SUM(Sales!$B$2:$B$4)"
summary["A3"] = "=Total*2"
summary["A4"] = "=Sales!B3"
summary["A5"] = '="Sales!B5 is "&Sales!B5'
workbook.defined_names["Total"] = DefinedName("Total", attr_text="Sales!$B$5")
workbook.save("fixture.xlsx")
"""


class WorkbookEditTest(WorkbookFixture):
    def setUp(self):
        super().setUp()
        run_office_python(FIXTURE_WORKBOOK, self.directory)

    def edit(self, operations, *extra_arguments):
        envelope = self.apply(operations, *extra_arguments, name="fixture.xlsx")
        self.assertEqual(envelope["status"], "ok", envelope)
        return envelope

    def formulas(self, sheet):
        envelope = run_office(["read", "fixture.xlsx", "--sheet", sheet, "--where", "formula"], self.directory)
        return {cell["cell"]: cell["formula"] for cell in envelope["details"]["range"]["cells"]}

    def values(self, sheet):
        envelope = run_office(["read", "fixture.xlsx", "--sheet", sheet, "--where", "formula"], self.directory)
        return {cell["cell"]: cell["value"] for cell in envelope["details"]["range"]["cells"]}

    def sheet_info(self, sheet="Sales"):
        envelope = run_office(["read", "fixture.xlsx"], self.directory)
        return next(info for info in envelope["details"]["sheets"] if info["name"] == sheet), envelope["details"]["definedNames"]

    def chart_references(self):
        series = load_workbook(self.directory / "fixture.xlsx")["Sales"]._charts[0].series[0]
        return series.tx.strRef.f, series.cat.numRef.f if series.cat.numRef else series.cat.strRef.f, series.val.numRef.f


class MisspelledInputTest(WorkbookEditTest):
    def test_a_misspelled_operation_field_or_sheet_names_the_closest_one(self):
        envelope = self.apply([{"op": "set_cel", "cell": "A1"}, {"op": "format_range", "range": "A1", "fontcolor": "FF0000"}], name="fixture.xlsx")
        self.assertEqual([issue["suggestion"] for issue in envelope["issues"]], ['use "op": "set_cell"', "rename the field to 'fontColor'"])
        envelope = self.apply([{"op": "set_cell", "sheet": "sales", "cell": "A1"}], name="fixture.xlsx")
        self.assertEqual(envelope["issues"][0]["suggestion"], 'use "sheet": "Sales"')
        self.assertIn("it has Sales, Summary", envelope["issues"][0]["message"])

    def test_a_short_korean_sheet_name_with_one_wrong_letter_names_the_sheet(self):
        self.create_workbook([{"title": "실적", "rows": [["a", "b"], [1, 2]]}, {"title": "분석", "rows": [["c"], [3]]}])
        issue = self.apply([{"op": "set_cell", "sheet": "실젹", "cell": "A9", "value": 1}])["issues"][0]
        self.assertEqual((issue["code"], issue["suggestion"]), ("TARGET_NOT_FOUND", 'use "sheet": "실적"'))

    def test_a_sheet_name_close_to_none_lists_the_names(self):
        self.create_workbook([{"title": "실적", "rows": [["a", "b"], [1, 2]]}, {"title": "분석", "rows": [["c"], [3]]}])
        issue = self.apply([{"op": "set_cell", "sheet": "요약표", "cell": "A9", "value": 1}])["issues"][0]
        self.assertEqual(issue["suggestion"], "use one of the sheet names: 실적, 분석")


class SpecificSuggestionTest(WorkbookEditTest):
    def refusal(self, operations):
        envelope = self.apply(operations, name="fixture.xlsx")
        self.assertEqual(envelope["status"], "error", envelope)
        return envelope["issues"][0]

    def test_an_address_names_what_is_wrong_with_it(self):
        issue = self.refusal([{"op": "set_cell", "sheet": "Sales", "cell": "ZZZZ99999999", "value": 1}])
        self.assertIn("column ZZZZ is past XFD and row 99999999 is outside 1 to 1048576", issue["message"])
        issue = self.refusal([{"op": "add_chart", "sheet": "Sales", "type": "bar", "range": "A1:B4", "anchor": "ZZ"}])
        self.assertEqual((issue["location"], issue["suggestion"]), ("ops[0].anchor", 'use a cell such as "ZZ1"'))
        issue = self.refusal([{"op": "add_chart", "type": "line", "range": "Sales!A1:B4"}])
        self.assertEqual(issue["suggestion"], 'put the sheet in its own field and the cells here: "sheet": "Sales", "range": "A1:B4"')
        envelope = run_office(["read", "fixture.xlsx", "--range", "A1:ZZZ"], self.directory)
        self.assertIn("in the corner ZZZ of 'A1:ZZZ', column ZZZ is past XFD and the row number is missing", envelope["summary"])
        envelope = run_office(["read", "fixture.xlsx", "--range", "XFE1"], self.directory)
        self.assertEqual(envelope["status"], "error")

    def test_a_taken_sheet_name_suggests_a_free_one(self):
        issue = self.refusal([{"op": "add_sheet", "name": "Sales"}])
        self.assertEqual(issue["suggestion"], 'pick a name no sheet has, such as "name": "Sales (2)"')
        issue = self.refusal([{"op": "rename_sheet", "sheet": "Summary", "name": "sales"}])
        self.assertEqual(issue["suggestion"], 'pick a name no sheet has, such as "name": "sales (2)"')

    def test_a_field_written_under_another_word_names_the_field_it_means(self):
        issues = self.apply([{"op": "add_chart", "sheet": "Sales", "type": "line", "data": "A1:B4"}], name="fixture.xlsx")["issues"]
        self.assertEqual([(issue["location"], issue["suggestion"]) for issue in issues], [("ops[0].data", "rename the field to 'range'")])
        write_json(self.directory / "spec.json", {"sheets": [{"name": "매출", "data": [["월", "매출"], ["1월", 5]]}]})
        issues = run_office(["create", "book.xlsx", "spec.json"], self.directory)["issues"]
        self.assertEqual([issue["suggestion"] for issue in issues], ["rename the field to 'title'", "rename the field to 'rows'"])

    def test_rows_written_as_objects_are_answered_with_the_lists_they_mean(self):
        write_json(self.directory / "spec.json", {"sheets": [{"title": "매출", "rows": [{"담당자": "이샘플", "매출": 5}, {"매출": 6, "담당자": "박예시"}]}]})
        issue = run_office(["create", "book.xlsx", "spec.json"], self.directory)["issues"][0]
        self.assertEqual((issue["code"], issue["location"]), ("WRONG_TYPE", "spec.sheets[0].rows"))
        self.assertIn("spec.sheets[0].rows[1] lists its keys in another order", issue["message"])
        self.assertIn('[["담당자", "매출"], ["이샘플", 5], ["박예시", 6]]', issue["suggestion"])


class TextLimitTest(WorkbookEditTest):
    def test_text_past_an_excel_limit_is_refused_with_the_limit(self):
        issue = self.apply([{"op": "add_chart", "sheet": "Sales", "type": "bar", "range": "A1:B4", "title": "가" * 3000}], name="fixture.xlsx")["issues"][0]
        self.assertEqual((issue["code"], issue["location"]), ("INVALID_VALUE", "ops[0].title"))
        self.assertIn("more than the 255", issue["message"])
        issue = self.apply([{"op": "set_cell", "sheet": "Sales", "cell": "A9", "value": "x" * 32768}], name="fixture.xlsx")["issues"][0]
        self.assertIn("an Excel cell holds at most 32767", issue["message"])
        write_json(self.directory / "spec.json", {"sheets": [{"title": "2026년 3분기 영업 실적 지역별 담당자별 상세 분석 보고서", "rows": [["a"]]}]})
        issue = run_office(["create", "book.xlsx", "spec.json"], self.directory)["issues"][0]
        self.assertEqual(issue["location"], "spec.sheets[0].title")
        self.assertIn("at most 31 characters", issue["message"])
        self.assertFalse((self.directory / "book.xlsx").exists())


class InsertAndDeleteTest(WorkbookEditTest):
    def test_inserted_rows_shift_every_reference_that_points_past_them(self):
        self.edit([{"op": "insert_rows", "sheet": "Sales", "at": 3, "count": 2}])
        self.assertEqual(self.formulas("Sales"), {"D1": "=$B$7*2+B2", "D2": "=AVERAGE(B2:B6)", "B7": "=SUM(B2:B6)"})
        self.assertEqual(self.formulas("Summary"), {"A1": "=Sales!B7", "A2": "=SUM(Sales!$B$2:$B$6)", "A3": "=Total*2", "A4": "=Sales!B5", "A5": '="Sales!B5 is "&Sales!B7'})

    def test_inserted_rows_move_filters_merged_ranges_names_and_charts(self):
        self.edit([{"op": "insert_rows", "sheet": "Sales", "at": 3, "count": 2}])
        sales, defined_names = self.sheet_info()
        self.assertEqual((sales["autoFilter"], sales["mergedCells"]), ("A1:B6", ["D5:E5"]))
        self.assertEqual(defined_names, [{"name": "Total", "value": "Sales!$B$7", "scope": None}])
        self.assertEqual(self.chart_references(), ("'Sales'!B1", "'Sales'!$A$2:$A$6", "'Sales'!$B$2:$B$6"))

    def test_recomputed_values_follow_the_shifted_formulas(self):
        self.edit([{"op": "insert_rows", "sheet": "Sales", "at": 3}, {"op": "set_cell", "sheet": "Sales", "cell": "B3", "value": 5}])
        self.assertEqual(self.values("Sales")["B6"], 65)
        self.assertEqual(self.values("Summary")["A3"], 130)

    def test_deleted_rows_shrink_ranges_and_break_references_into_them(self):
        self.edit([{"op": "delete_rows", "sheet": "Sales", "at": 3}])
        self.assertEqual(self.formulas("Sales"), {"D1": "=$B$4*2+B2", "D2": "=AVERAGE(B2:B3)", "B4": "=SUM(B2:B3)"})
        self.assertEqual(self.formulas("Summary")["A4"], "=Sales!#REF!")
        self.assertEqual(self.formulas("Summary")["A1"], "=Sales!B4")
        sales, defined_names = self.sheet_info()
        self.assertEqual((sales["autoFilter"], sales["mergedCells"]), ("A1:B3", []))
        self.assertEqual(defined_names[0]["value"], "Sales!$B$4")

    def test_inserted_columns_shift_columns_widths_and_charts(self):
        self.edit([{"op": "insert_columns", "sheet": "Sales", "at": "B"}])
        self.assertEqual(self.formulas("Sales"), {"E1": "=$C$5*2+C2", "E2": "=AVERAGE(C2:C4)", "C5": "=SUM(C2:C4)"})
        self.assertEqual(self.formulas("Summary")["A2"], "=SUM(Sales!$C$2:$C$4)")
        sales, _ = self.sheet_info()
        self.assertEqual((sales["autoFilter"], sales["mergedCells"]), ("A1:C4", ["E3:F3"]))
        self.assertEqual(load_workbook(self.directory / "fixture.xlsx")["Sales"].column_dimensions["C"].width, 20)
        self.assertEqual(self.chart_references()[2], "'Sales'!$C$2:$C$4")

    def test_deleted_columns_break_references_to_them_and_shift_the_rest(self):
        self.edit([{"op": "delete_columns", "sheet": "Sales", "at": "A"}])
        self.assertEqual(self.formulas("Sales"), {"C1": "=$A$5*2+A2", "C2": "=AVERAGE(A2:A4)", "A5": "=SUM(A2:A4)"})
        self.edit([{"op": "delete_columns", "sheet": "Sales", "at": "A"}])
        self.assertEqual(self.formulas("Summary")["A1"], "=Sales!#REF!")
        self.assertEqual(self.formulas("Summary")["A3"], "=Total*2")


class RenameTest(WorkbookEditTest):
    def test_renaming_a_sheet_rewrites_formulas_names_and_charts_but_not_text(self):
        self.edit([{"op": "rename_sheet", "sheet": "Sales", "name": "판매 현황"}])
        summary = self.formulas("Summary")
        self.assertEqual(summary["A1"], "='판매 현황'!B5")
        self.assertEqual(summary["A2"], "=SUM('판매 현황'!$B$2:$B$4)")
        self.assertEqual(summary["A5"], '="Sales!B5 is "&\'판매 현황\'!B5')
        _, defined_names = self.sheet_info("판매 현황")
        self.assertEqual(defined_names[0]["value"], "'판매 현황'!$B$5")
        renamed_chart = load_workbook(self.directory / "fixture.xlsx")["판매 현황"]._charts[0]
        self.assertEqual(renamed_chart.series[0].val.numRef.f, "'판매 현황'!$B$2:$B$4")

    def test_a_renamed_sheet_keeps_its_values(self):
        self.edit([{"op": "rename_sheet", "sheet": "Sales", "name": "Q1"}])
        self.assertEqual(self.values("Summary")["A1"], 60)

    def test_a_name_another_sheet_has_is_refused(self):
        envelope = self.apply([{"op": "rename_sheet", "sheet": "Sales", "name": "summary"}], name="fixture.xlsx")
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["INVALID_VALUE"])


class BatchTest(WorkbookEditTest):
    def test_each_operation_sees_what_the_earlier_ones_did(self):
        self.edit([
            {"op": "add_sheet", "name": "Extra"},
            {"op": "set_cell", "sheet": "Extra", "cell": "A1", "value": "=Sales!B5+1"},
            {"op": "rename_sheet", "sheet": "Extra", "name": "Later"},
            {"op": "set_cell", "sheet": "Later", "cell": "A2", "value": 7},
        ])
        envelope = run_office(["read", "fixture.xlsx", "--sheet", "Later"], self.directory)
        self.assertEqual(envelope["details"]["range"]["values"], [[61], [7]])

    def test_one_bad_operation_leaves_the_file_byte_identical(self):
        original = (self.directory / "fixture.xlsx").read_bytes()
        envelope = self.apply([
            {"op": "insert_rows", "sheet": "Sales", "at": 2},
            {"op": "set_cell", "sheet": "Missing", "cell": "A1", "value": 1},
        ], name="fixture.xlsx")
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["TARGET_NOT_FOUND"])
        self.assertEqual(envelope["issues"][0]["location"], "ops[1].sheet")
        self.assertEqual((self.directory / "fixture.xlsx").read_bytes(), original)
        self.assertEqual(sorted(path.name for path in self.directory.iterdir() if path.name.startswith(".office-")), [])

    def test_a_formula_that_reads_its_own_cell_writes_nothing(self):
        original = (self.directory / "fixture.xlsx").read_bytes()
        envelope = self.apply([
            {"op": "set_cell", "sheet": "Sales", "cell": "A2", "value": "renamed"},
            {"op": "set_cell", "sheet": "Sales", "cell": "B6", "value": "=SUM(B2:B6)"},
        ], name="fixture.xlsx")
        self.assertEqual((envelope["status"], envelope["outputPath"]), ("error", None))
        self.assertEqual([(issue["code"], issue["location"]) for issue in envelope["issues"]], [("CIRCULAR_REFERENCE", "Sales!B6")])
        self.assertIn("nothing was written", envelope["summary"])
        self.assertEqual((self.directory / "fixture.xlsx").read_bytes(), original)

    def test_a_circular_reference_the_workbook_already_had_does_not_block_another_edit(self):
        run_office_python("""
            from openpyxl import load_workbook
            workbook = load_workbook("fixture.xlsx")
            workbook["Summary"]["B1"] = "=B1+1"
            workbook.save("fixture.xlsx")
        """, self.directory)
        envelope = self.apply([{"op": "set_cell", "sheet": "Sales", "cell": "A2", "value": "renamed"}], name="fixture.xlsx")
        self.assertEqual(envelope["status"], "ok", envelope)
        self.assertEqual(load_workbook(self.directory / "fixture.xlsx")["Sales"]["A2"].value, "renamed")

    def test_a_formula_that_does_not_parse_is_refused_where_it_is(self):
        original = (self.directory / "fixture.xlsx").read_bytes()
        cases = (
            ({"op": "set_cell", "sheet": "Sales", "cell": "A3", "value": "=SUM(Z1:Z3"}, "ops[0].value", "leaves SUM( open"),
            ({"op": "set_range", "sheet": "Sales", "cell": "F1", "values": [[1], ["=B2+"]]}, "ops[0].values[1][0]", "ends with the operator +"),
            ({"op": "set_cell", "sheet": "Sales", "cell": "A3", "value": "=SUM(B2:B4))"}, "ops[0].value", "a ) that closes nothing"),
            ({"op": "set_cell", "sheet": "Sales", "cell": "A3", "value": "=B2+*B3"}, "ops[0].value", "nothing before the operator *"),
            ({"op": "add_conditional_format", "sheet": "Sales", "range": "B2:B4", "rule": "formula", "formula": '=$B2>"a'}, "ops[0].formula", "never closes"),
        )
        for operation, location, problem in cases:
            with self.subTest(operation=operation):
                envelope = self.apply([operation], name="fixture.xlsx")
                self.assertEqual([(issue["code"], issue["location"]) for issue in envelope["issues"]], [("FORMULA_SYNTAX", location)])
                self.assertIn(problem, envelope["summary"])
        self.assertEqual((self.directory / "fixture.xlsx").read_bytes(), original)
        self.edit([{"op": "set_cell", "sheet": "Sales", "cell": "A3", "value": "=SUM(Z1:Z3", "type": "text"}])

    def test_a_dry_run_reports_each_change_and_writes_nothing(self):
        original = (self.directory / "fixture.xlsx").read_bytes()
        envelope = self.apply([{"op": "insert_rows", "sheet": "Sales", "at": 2}, {"op": "delete_columns", "sheet": "Sales", "at": "A"}], "--dry-run", name="fixture.xlsx")
        self.assertEqual(envelope["status"], "ok")
        self.assertTrue(envelope["details"]["dryRun"])
        self.assertEqual([change["op"] for change in envelope["details"]["changes"]], ["insert_rows", "delete_columns"])
        self.assertEqual((self.directory / "fixture.xlsx").read_bytes(), original)

    def test_output_leaves_the_source_alone(self):
        original = (self.directory / "fixture.xlsx").read_bytes()
        envelope = self.apply([{"op": "set_cell", "sheet": "Sales", "cell": "A1", "value": "renamed"}], "--output", "copy.xlsx", name="fixture.xlsx")
        self.assertTrue(envelope["outputPath"].endswith("copy.xlsx"))
        self.assertEqual((self.directory / "fixture.xlsx").read_bytes(), original)
        self.assertEqual(load_workbook(self.directory / "copy.xlsx")["Sales"]["A1"].value, "renamed")

    def test_a_column_beyond_the_sheet_is_refused_before_anything_is_written(self):
        envelope = self.apply([{"op": "insert_columns", "sheet": "Sales", "at": "ZZZZ"}], name="fixture.xlsx")
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["INVALID_VALUE"])

    def test_an_unknown_field_is_named(self):
        envelope = self.apply([{"op": "set_cell", "sheet": "Sales", "cell": "A1", "value": 1, "colour": "red"}], name="fixture.xlsx")
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["UNKNOWN_FIELD"])


class CellAndStyleTest(WorkbookEditTest):
    def test_text_starting_with_an_equals_sign_is_a_formula_unless_the_type_says_text(self):
        self.edit([
            {"op": "set_cell", "sheet": "Sales", "cell": "F1", "value": "=B2+1"},
            {"op": "set_cell", "sheet": "Sales", "cell": "F2", "value": "=B2+1", "type": "text"},
            {"op": "set_range", "sheet": "Sales", "cell": "F3", "values": [[1, "=F3*2"], ["=x", None]], "type": "auto"},
        ])
        envelope = run_office(["read", "fixture.xlsx", "--range", "F1:G4"], self.directory)
        selected = envelope["details"]["range"]
        self.assertEqual(selected["formulas"], [["=B2+1", None], [None, None], [None, "=F3*2"], ["=x", None]])
        self.assertEqual(selected["values"][0][0], 11)
        self.assertEqual(selected["values"][1][0], "=B2+1")
        self.assertEqual(selected["values"][2], [1, 2])

    def test_format_range_changes_only_the_properties_it_names(self):
        self.edit([{"op": "format_range", "sheet": "Sales", "range": "A1:B1", "bold": True, "fill": "DCEAF7", "alignment": "center", "numberFormat": "#,##0", "fontColor": "FF0000", "wrapText": True}])
        cell = load_workbook(self.directory / "fixture.xlsx")["Sales"]["B1"]
        self.assertTrue(cell.font.b)
        self.assertEqual(cell.fill.fgColor.rgb, "00DCEAF7")
        self.assertEqual((cell.alignment.horizontal, cell.alignment.wrap_text), ("center", True))
        self.assertEqual((cell.number_format, cell.font.color.rgb), ("#,##0", "00FF0000"))
        self.edit([{"op": "format_range", "sheet": "Sales", "range": "B1", "bold": False}])
        cell = load_workbook(self.directory / "fixture.xlsx")["Sales"]["B1"]
        self.assertFalse(cell.font.b)
        self.assertEqual(cell.fill.fgColor.rgb, "00DCEAF7")

    def test_a_color_that_is_not_six_hex_digits_is_refused(self):
        envelope = self.apply([{"op": "format_range", "sheet": "Sales", "range": "A1", "fill": "red"}], name="fixture.xlsx")
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["INVALID_VALUE"])

    def test_column_width_freeze_and_filter(self):
        self.edit([
            {"op": "set_column_width", "sheet": "Sales", "column": "A", "width": 30},
            {"op": "freeze_panes", "sheet": "Sales", "cell": "B2"},
            {"op": "set_auto_filter", "sheet": "Sales", "range": "A1:B4"},
            {"op": "set_auto_filter", "sheet": "Summary", "range": ""},
        ])
        workbook = load_workbook(self.directory / "fixture.xlsx")
        self.assertEqual(workbook["Sales"].column_dimensions["A"].width, 30)
        self.assertEqual(workbook["Sales"].column_dimensions["B"].width, 20)
        self.assertEqual(workbook["Sales"].freeze_panes, "B2")
        self.assertEqual(workbook["Sales"].auto_filter.ref, "A1:B4")
        self.assertFalse(workbook["Summary"].auto_filter.ref)

    def test_a_chart_is_added_from_a_block_with_its_header_and_categories(self):
        self.edit([{"op": "add_chart", "sheet": "Sales", "type": "line", "range": "A1:B4", "title": "Amounts", "anchor": "G20"}])
        sales, _ = self.sheet_info()
        self.assertEqual([chart["type"] for chart in sales["charts"]], ["bar", "line"])
        chart = load_workbook(self.directory / "fixture.xlsx")["Sales"]._charts[1]
        self.assertEqual(type(chart).__name__, "LineChart")
        self.assertEqual(chart.series[0].val.numRef.f, "'Sales'!$B$2:$B$4")

    def test_a_chart_needs_a_header_row_and_a_series_column(self):
        envelope = self.apply([{"op": "add_chart", "sheet": "Sales", "type": "bar", "range": "A1:A4"}], name="fixture.xlsx")
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["OPERATION_NOT_APPLICABLE"])


class MacroWorkbookTest(WorkbookFixture):
    def build_macro_workbook(self):
        run_office_python("""
        import zipfile
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = "S"
        workbook.active["A1"] = 1
        workbook.active["A2"] = "=A1*2"
        workbook.save("plain.xlsx")
        source = zipfile.ZipFile("plain.xlsx")
        with zipfile.ZipFile("macro.xlsm", "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "[Content_Types].xml":
                    text = data.decode().replace("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml", "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
                    text = text.replace("<Default ", '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/><Default ', 1)
                    data = text.encode()
                if info.filename == "xl/_rels/workbook.xml.rels":
                    data = data.decode().replace("</Relationships>", '<Relationship Id="rId99" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>').encode()
                target.writestr(info, data)
            target.writestr("xl/vbaProject.bin", b"FAKE-VBA-PAYLOAD")
        """, self.directory)

    def test_a_macro_workbook_keeps_its_macros_through_an_edit(self):
        self.build_macro_workbook()
        envelope = self.apply([{"op": "set_cell", "cell": "A1", "value": 5}], name="macro.xlsm")
        self.assertEqual(envelope["status"], "ok", envelope)
        with zipfile.ZipFile(self.directory / "macro.xlsm") as archive:
            self.assertEqual(archive.read("xl/vbaProject.bin"), b"FAKE-VBA-PAYLOAD")
            self.assertIn(b"macroEnabled", archive.read("[Content_Types].xml"))
        self.assertEqual(self.cells("macro.xlsm")["A2"]["value"], "10")

    def test_a_macro_workbook_keeps_its_macros_when_rows_are_appended(self):
        self.build_macro_workbook()
        envelope = self.apply([{"op": "append_rows", "rows": [[3, "=A3*2"]]}], name="macro.xlsm")
        self.assertEqual(envelope["status"], "ok", envelope)
        with zipfile.ZipFile(self.directory / "macro.xlsm") as archive:
            self.assertEqual(archive.read("xl/vbaProject.bin"), b"FAKE-VBA-PAYLOAD")


if __name__ == "__main__":
    unittest.main()
