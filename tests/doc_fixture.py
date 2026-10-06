import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import textwrap
import unittest


SCRIPTS_PATH = Path(__file__).resolve().parents[1] / "skills" / "office" / "scripts"
if str(SCRIPTS_PATH) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_PATH))
OFFICE_ENTRY = SCRIPTS_PATH / "office"

from task_context_fixture import CONTEXT_VARIABLE, write_context_at  # noqa: E402

FIXTURE_DOCUMENT = """
from docx import Document
document = Document()
document.add_heading("개요", 1)
document.add_paragraph("첫 문단 {{ customer_name }} 입니다.")
document.add_paragraph("둘째 문단")
table = document.add_table(rows=2, cols=2)
for row, values in zip(table.rows, [["항목", "값"], ["매출", "100"]]):
    for cell, value in zip(row.cells, values):
        cell.text = value
document.add_heading("결론", 1)
document.save("fixture.docx")
"""


RUNTIME_CONTEXT_FILE = "task-context.json"


def run_office(arguments, working_directory):
    completed = subprocess.run([sys.executable, str(OFFICE_ENTRY), *arguments], capture_output=True, text=True, cwd=working_directory, env=office_environment(working_directory))
    return json.loads(completed.stdout)


def office_environment(working_directory):
    environment = {name: value for name, value in os.environ.items() if name != CONTEXT_VARIABLE}
    context_path = Path(working_directory) / RUNTIME_CONTEXT_FILE
    return environment | ({CONTEXT_VARIABLE: str(context_path)} if context_path.is_file() else {})


def write_runtime_context(directory, profile_path):
    facts = {"requester": {"name": "이샘플", "email": "sample@example.com"}, "today": "2026-10-04",
             "company": {"ko": str(profile_path), "en": str(profile_path)}}
    write_context_at(Path(directory) / RUNTIME_CONTEXT_FILE, facts)


def run_office_python(code, working_directory):
    subprocess.run([sys.executable, str(OFFICE_ENTRY), "python", "-c", textwrap.dedent(code)], check=True, cwd=working_directory)


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def write_form_values(path, values):
    form_values = dict(values)
    profile = form_values.pop("profile", None)
    if profile is not None:
        profile_path = Path(path).resolve().parent / "company-profile.json"
        write_json(profile_path, profile)
        write_runtime_context(profile_path.parent, profile_path)
    write_json(path, form_values)


def block_texts(working_directory, document_name):
    envelope = run_office(["read", document_name], working_directory)
    return [(block["kind"], block.get("text", block.get("cells"))) for block in envelope["details"]["blocks"]]


class DocumentFixture(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)
        run_office_python(FIXTURE_DOCUMENT, self.directory)

    def tearDown(self):
        self.temporary_directory.cleanup()


CONTRACT_DOCUMENT = """
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

document = Document()
document.add_heading("용역 계약서", 0)
document.add_paragraph('주식회사 예시상사(이하 "갑")와 이샘플(이하 "을")은 다음과 같이 용역 계약을 체결한다.')
document.add_heading("제1조 (목적)", 1)
document.add_paragraph("본 계약은 갑이 을에게 홈페이지 개편 용역을 위탁하고 을이 이를 수행하는 데 필요한 사항을 정한다.")
document.add_heading("제2조 (계약 금액)", 1)
document.add_paragraph("계약 금액은 금 일천만 원(₩10,000,000)이며 부가가치세는 별도로 한다.")
document.add_heading("제3조 (지급 조건)", 1)
document.add_paragraph("갑은 검수 완료일로부터 30일 이내에 을에게 계약 금액을 지급한다.")
table = document.add_table(rows=3, cols=2)
table.style = "Table Grid"
for row, values in zip(table.rows, [["구분", "금액"], ["착수금", "3,000,000"], ["잔금", "7,000,000"]]):
    for cell, value in zip(row.cells, values):
        cell.text = value
document.add_paragraph("본 계약의 성립을 증명하기 위하여 계약서 2부를 작성하여 각 1부씩 보관한다.")
purpose = document.paragraphs[3]._p
purpose.append(parse_xml(
    f'<w:ins {nsdecls("w")} w:id="901" w:author="이샘플" w:date="2026-09-01T09:00:00Z">'
    '<w:r><w:t xml:space="preserve"> 세부 범위는 별첨 과업지시서에 따른다.</w:t></w:r></w:ins>'
))
heading_run = document.paragraphs[2].runs[0]._r
heading_run.get_or_add_rPr().append(parse_xml(
    f'<w:b {nsdecls("w")}/>'
))
heading_run.rPr.append(parse_xml(
    f'<w:rPrChange {nsdecls("w")} w:id="902" w:author="이샘플" w:date="2026-09-01T09:05:00Z"><w:rPr/></w:rPrChange>'
))
document.save("contract.docx")
"""


def read_details(working_directory, document_name, *flags):
    return run_office(["read", document_name, *flags], working_directory)["details"]


class ContractFixture(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)
        run_office_python(CONTRACT_DOCUMENT, self.directory)

    def tearDown(self):
        self.temporary_directory.cleanup()
