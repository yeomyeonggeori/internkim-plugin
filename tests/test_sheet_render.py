import datetime
import math
from pathlib import Path
import re
import sys
import tempfile
import unittest

from render_fixture import can_render, pdf_page_count, png_size
from sheet_fixture import SCRIPTS_PATH, run_office, run_office_python

sys.path.insert(0, str(SCRIPTS_PATH))

from core.number_format import displayed  # noqa: E402
from render.office_preview import pixels  # noqa: E402
from core.page_sizes import DEFAULT_PAPER  # noqa: E402
from core.units import inches_to_pixels  # noqa: E402


WORKBOOK = """
import datetime
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

book = Workbook()
sheet = book.active
sheet.title = "실적"
sheet.merge_cells("A1:D1")
sheet["A1"] = "3분기 실적"
sheet.append(["지역", "매출", "달성률", "작성일"])
for region, sales, rate in (("서울", 548300000, 1.054), ("부산", 287450000, 0.927)):
    sheet.append([region, sales, rate, datetime.date(2026, 10, 1)])
for row in (3, 4):
    sheet.cell(row=row, column=2).number_format = "#,##0"
    sheet.cell(row=row, column=3).number_format = "0.0%"
    sheet.cell(row=row, column=4).number_format = 'yyyy"년" m"월" d"일"'
sheet["E3"] = 123456789
sheet["E3"].number_format = "#,##0"
sheet.column_dimensions["B"].width = 14
sheet.column_dimensions["D"].width = 18
sheet.column_dimensions["E"].width = 5
sheet.conditional_formatting.add("C3:C4", CellIsRule(operator="lessThan", formula=["1"], font=Font(color="C00000"), fill=PatternFill("solid", bgColor="FCE4E4")))
chart = BarChart()
chart.add_data(Reference(sheet, min_col=2, min_row=2, max_row=4), titles_from_data=True)
chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=4))
chart.x_axis.title = "지역"
chart.y_axis.title = "매출액"
sheet.add_chart(chart, "A6")
detail = book.create_sheet("내역")
detail.append(["번호", "금액"])
for number in range(1, 151):
    detail.append([number, number * 1000])
detail.print_title_rows = "1:1"
detail.oddFooter.center.text = "&P / &N"
detail.page_setup.orientation = "landscape"
book.save("실적.xlsx")
"""


class SheetPreviewTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary_directory = tempfile.TemporaryDirectory()
        cls.directory = Path(cls.temporary_directory.name)
        run_office_python(WORKBOOK, cls.directory)
        cls.envelope = run_office(["render", "실적.xlsx"], cls.directory)
        cls.html = (cls.directory / "실적-preview" / "preview.html").read_text(encoding="utf-8")
        cls.pages = re.findall(r'<section data-page="(\d+)" style="([^"]*)">(.*?)</section>', cls.html, re.S)

    @classmethod
    def tearDownClass(cls):
        cls.temporary_directory.cleanup()

    @unittest.skipUnless(can_render(), "needs bun, or node 18 or newer")
    def test_every_printed_page_is_drawn_as_an_image_and_a_pdf_page(self):
        details = self.envelope["details"]
        self.assertNotIn("PAGES_NOT_RENDERED", [issue["code"] for issue in self.envelope["issues"]])
        section_sizes = [tuple(math.ceil(float(value)) for value in re.search(r"width:([\d.]+)px;height:([\d.]+)px", style).groups()) for _, style, _ in self.pages]
        self.assertTrue(details["seen"])
        self.assertEqual([png_size(self.directory / page) for page in details["pages"]], section_sizes)
        self.assertEqual(pdf_page_count(self.directory / details["pdf"]), len(self.pages))

    def test_cells_show_their_displayed_values_and_conditional_colors(self):
        first = self.pages[0][2]
        for text in ("548,300,000", "105.4%", "2026년 10월 1일", "3분기 실적"):
            self.assertIn(f">{text}<", first)
        self.assertRegex(first, r"color:#c00000[^>]*><span>92\.7%<")
        self.assertRegex(first, r">#{2,}<")
        self.assertIn("grid-column:1 / span 4", first)
        self.assertIn("<svg", first)

    def test_chart_axis_titles_are_drawn_below_and_beside_the_plot(self):
        first = self.pages[0][2]
        self.assertIn(">지역</text>", first)
        self.assertRegex(first, r'<g transform="rotate\(-90[^"]*"><text[^>]*>매출액</text></g>')

    def test_long_sheets_paginate_with_title_rows_and_page_numbers(self):
        landscape_width = pixels(inches_to_pixels(max(DEFAULT_PAPER.inches)))
        detail_pages = [body for _, style, body in self.pages if f"width:{landscape_width}" in style]
        self.assertGreaterEqual(len(detail_pages), 2)
        self.assertTrue(all(">번호<" in body for body in detail_pages))
        self.assertIn(f">{len(self.pages)} / {len(self.pages)}<", self.pages[-1][2])
        for forbidden in ("<style", "<script", "::before", "counter("):
            self.assertNotIn(forbidden, self.html)
        self.assertTrue(all(Path(font["path"]).is_file() for font in self.envelope["details"]["previewFonts"]))


class NumberFormatTest(unittest.TestCase):
    def test_common_formats_display_as_excel_shows_them(self):
        cases = [
            (1234567.891, "#,##0", "1,234,568"),
            (0.1234, "0.0%", "12.3%"),
            (-1500, "#,##0_);[Red](#,##0)", "(1,500)"),
            (1500, '"₩"#,##0', "₩1,500"),
            (1500, "[$₩-412]#,##0", "₩1,500"),
            (0, '#,##0;-#,##0;"-"', "-"),
            (1234567, '#,##0,"천원"', "1,235천원"),
            (12345678, "0.00E+00", "1.23E+07"),
            (46296, 'yyyy"년" m"월" d"일"', "2026년 10월 1일"),
            (datetime.datetime(2026, 10, 1, 14, 5), "h:mm AM/PM", "2:05 오후"),
            (datetime.date(2026, 10, 1), "mm-dd-yy", "2026-10-01"),
            (3.14159, "General", "3.14159"),
            (14.2, '#,##0.0"%"', "14.2%"),
            (5, "0\\%", "5%"),
        ]
        self.assertEqual([displayed(value, number_format).text for value, number_format, _ in cases], [expected for _, _, expected in cases])
        self.assertEqual(displayed(-3, "#,##0;[Red]-#,##0").color, "#ff0000")


if __name__ == "__main__":
    unittest.main()
