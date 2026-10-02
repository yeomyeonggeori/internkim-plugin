import unittest
import zipfile

from doc_fixture import DocumentFixture, block_texts, read_details, run_office, run_office_python, write_json


class CheckTest(DocumentFixture):
    def test_each_finding_suggests_the_operation_that_fixes_it(self):
        run_office_python("""
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            document = Document("fixture.docx")
            fonts = document.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:rFonts"))
            fonts.attrib.pop(qn("w:eastAsiaTheme"))
            contents = document.paragraphs[0].insert_paragraph_before()
            run = contents.add_run()
            for kind, text in (("begin", None), ("instr", " TOC \\\\o "), ("separate", None), ("text", "개요"), ("end", None)):
                node = OxmlElement("w:instrText" if kind == "instr" else "w:t" if kind == "text" else "w:fldChar")
                if text is None:
                    node.set(qn("w:fldCharType"), kind)
                else:
                    node.text = text
                run._r.append(node)
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), "missing_bookmark")
            document.paragraphs[-1]._p.append(link)
            document.save("fixture.docx")
        """, self.directory)
        envelope = run_office(["check", "fixture.docx"], self.directory)
        findings = {issue["code"]: issue["fix"] for issue in envelope["issues"]}
        self.assertEqual(set(findings), {"PLACEHOLDER_LEFT", "BROKEN_INTERNAL_REFERENCE", "STALE_TABLE_OF_CONTENTS", "EAST_ASIA_FONT_MISSING", "EAST_ASIA_LANGUAGE_NOT_KOREAN"})
        fixes = [findings["PLACEHOLDER_LEFT"][0] | {"replace": "박예시"}, *findings["STALE_TABLE_OF_CONTENTS"], *findings["EAST_ASIA_FONT_MISSING"], *findings["EAST_ASIA_LANGUAGE_NOT_KOREAN"]]
        write_json(self.directory / "fixes.json", fixes)
        self.assertEqual(run_office(["apply", "fixture.docx", "fixes.json"], self.directory)["status"], "ok")
        remaining = {issue["code"] for issue in run_office(["check", "fixture.docx"], self.directory)["issues"]}
        self.assertEqual(remaining, {"BROKEN_INTERNAL_REFERENCE"})
        with zipfile.ZipFile(self.directory / "fixture.docx") as archive:
            self.assertIn(b'w:updateFields w:val="true"', archive.read("word/settings.xml"))

    def test_draft_text_is_a_placeholder_left_like_template_syntax(self):
        run_office_python("""
            from docx import Document
            document = Document("fixture.docx")
            document.add_paragraph("납기: TBD")
            document.add_paragraph("계약일: [insert date]")
            document.add_paragraph("Lorem ipsum dolor sit amet")
            document.save("fixture.docx")
        """, self.directory)
        envelope = run_office(["check", "fixture.docx"], self.directory)
        found = sorted(operation["find"] for issue in envelope["issues"] if issue["code"] == "PLACEHOLDER_LEFT" for operation in issue["fix"])
        self.assertEqual(found, ["Lorem ipsum", "TBD", "[insert date]", "{{ customer_name }}"])

    def test_apply_refuses_a_fix_whose_fill_in_was_not_written(self):
        envelope = run_office(["check", "fixture.docx"], self.directory)
        fix = next(issue["fix"] for issue in envelope["issues"] if issue["code"] == "PLACEHOLDER_LEFT")
        write_json(self.directory / "fixes.json", fix)
        original = (self.directory / "fixture.docx").read_bytes()
        refused = run_office(["apply", "fixture.docx", "fixes.json"], self.directory)
        self.assertEqual([(issue["code"], issue["location"]) for issue in refused["issues"]], [("FILL_IN_LEFT", "ops[0].replace")])
        self.assertEqual((self.directory / "fixture.docx").read_bytes(), original)


class MergeTest(DocumentFixture):
    def test_a_missing_value_writes_nothing(self):
        write_json(self.directory / "values.json", {"customer": "박예시"})
        envelope = run_office(["merge", "fixture.docx", "values.json", "merged.docx"], self.directory)
        self.assertEqual([(issue["code"], issue["location"]) for issue in envelope["issues"]], [("UNRESOLVED_PLACEHOLDER", "body")])
        self.assertFalse((self.directory / "merged.docx").exists())

    def test_values_fill_placeholders_and_unused_names_warn(self):
        write_json(self.directory / "values.json", {"customer_name": "박예시", "custmer": "오타"})
        envelope = run_office(["merge", "fixture.docx", "values.json", "merged.docx"], self.directory)
        self.assertEqual([(issue["code"], issue["location"]) for issue in envelope["issues"]], [("UNUSED_VALUE", "values.custmer")])
        self.assertIn(("paragraph", "첫 문단 박예시 입니다."), block_texts(self.directory, "merged.docx"))

    def test_a_placeholder_split_across_runs_and_an_item_path_fill(self):
        run_office_python("""
            from docx import Document
            document = Document("fixture.docx")
            paragraph = document.add_paragraph()
            for text, bold in (("계약자: {{ cus", False), ("tomer.na", True), ("me }}, 첫 품목: {{ items.0.name }}", False)):
                paragraph.add_run(text).bold = bold
            document.save("fixture.docx")
        """, self.directory)
        write_json(self.directory / "values.json", {"customer_name": "박예시", "customer": {"name": "이샘플"}, "items": [{"name": "연간 유지보수"}]})
        envelope = run_office(["merge", "fixture.docx", "values.json", "merged.docx"], self.directory)
        self.assertEqual(envelope["status"], "ok")
        self.assertIn(("paragraph", "계약자: 이샘플, 첫 품목: 연간 유지보수"), block_texts(self.directory, "merged.docx"))

    def test_a_table_row_naming_a_list_repeats_once_per_item_as_in_a_deck(self):
        run_office_python("""
            from docx import Document
            document = Document("fixture.docx")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = "품목", "금액"
            run = table.cell(1, 0).paragraphs[0]
            run.add_run("{{ items.na"); run.add_run("me }}")
            table.cell(1, 1).text = "{{ items.amount }}"
            document.save("fixture.docx")
        """, self.directory)
        write_json(self.directory / "values.json", {"customer_name": "박예시", "items": [{"name": "노트북", "amount": 1200000}, {"name": "모니터", "amount": 300000}]})
        envelope = run_office(["merge", "fixture.docx", "values.json", "merged.docx"], self.directory)
        self.assertEqual(envelope["status"], "ok", envelope["issues"])
        tables = [block for block in read_details(self.directory, "merged.docx")["blocks"] if block.get("cells")]
        self.assertEqual(tables[-1]["cells"], [["품목", "금액"], ["노트북", "1200000"], ["모니터", "300000"]])

    def test_a_paragraph_naming_a_list_repeats_and_an_empty_list_leaves_it_out(self):
        run_office_python("""
            from docx import Document
            document = Document("fixture.docx")
            document.add_paragraph("품목: {{ items.name }}")
            document.add_paragraph("비고: {{ notes }}")
            document.add_paragraph("첨부: {{ attachments }}")
            document.save("fixture.docx")
        """, self.directory)
        write_json(self.directory / "values.json", {"customer_name": "박예시", "items": [{"name": "노트북"}, {"name": "모니터"}], "notes": ["납기 엄수"], "attachments": []})
        envelope = run_office(["merge", "fixture.docx", "values.json", "merged.docx"], self.directory)
        self.assertEqual(envelope["status"], "ok", envelope["issues"])
        texts = [text for kind, text in block_texts(self.directory, "merged.docx") if kind == "paragraph"]
        self.assertEqual(texts[-3:], ["품목: 노트북", "품목: 모니터", "비고: 납기 엄수"])
        self.assertFalse(any(text.startswith("첨부") for text in texts))

    def test_an_object_where_one_value_is_written_names_the_placeholder(self):
        write_json(self.directory / "values.json", {"customer_name": {"first": "예시"}})
        envelope = run_office(["merge", "fixture.docx", "values.json", "merged.docx"], self.directory)
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["LIST_NEEDS_A_ROW"])
        self.assertFalse((self.directory / "merged.docx").exists())

    def test_a_statement_tag_is_refused_by_name(self):
        run_office_python("""
            from docx import Document
            document = Document("fixture.docx")
            document.add_paragraph("{% if customer_name %}귀하{% endif %}")
            document.save("fixture.docx")
        """, self.directory)
        write_json(self.directory / "values.json", {"customer_name": "박예시"})
        envelope = run_office(["merge", "fixture.docx", "values.json", "merged.docx"], self.directory)
        self.assertEqual([issue["code"] for issue in envelope["issues"]], ["TEMPLATE_SYNTAX_ERROR", "TEMPLATE_SYNTAX_ERROR"])
        self.assertIn("{% if customer_name %}", envelope["issues"][0]["message"])
        self.assertFalse((self.directory / "merged.docx").exists())


if __name__ == "__main__":
    unittest.main()
